"""Score a generated IIG sheet row against a real one (stdlib + openpyxl).

    python scripts/iig_scorecard.py --generated out/<feed>/framework/<feed>_IIG.xlsx \
        --real "<real IIG rows>.xlsx" --sheet ADLS_DELTA_INGESTION_DETAILS \
        --real-table <TGT_TABLE_NAME> [--generated-table <TGT_TABLE_NAME>] \
        [--review <feed>_IIG_REVIEW.xlsx] [--alias generated=real ...]

Rows are matched by TGT_TABLE_NAME. One line per column of the REAL sheet:

* MATCH — same value (a blank generated cell matches a real NULL / blank);
* ALIAS — equal only once the fixture's anonymisation is undone (--alias
  pairs, applied to the generated value, case-insensitive); a case-only
  difference is a DIFF marked "(case only)"; a fixture's masked sequence id
  (``SYN-OBJ-003``) against the generated ``3`` is ALIAS too;
* SRC_ / TGT_DATA_TYPE compare case-insensitively (MATCH with a note);
  LOB is ALIAS when the generated row is the anonymised fixture (its table
  name differs from --real-table);
* OPEN  — the generated cell is blank; owner from the review copy's
  REVIEW_SUMMARY (sheet + column);
* DIFF  — real vs generated, values cut to 60 characters; comma-list
  columns are compared by length + first 3 items.

CREATED_BY / UPDATED_BY values are never printed. The last line:
"filled N/<cols>, matched M, alias A, open K (BSA x, Engineer y,
Engineer-confirm z, CI/CD w), diff D".

ALL-SHEETS mode (multi-table step 7):

    python scripts/iig_scorecard.py --all-sheets \
        --generated out/<feed>/framework/<feed>_IIG.xlsx --real <real IIG>.xlsx \
        [--review <feed>_IIG_REVIEW.xlsx] [--alias generated=real ...] \
        [--summary-only | --matches]

scores every sheet of the REAL workbook (iig_v2: all eight; sheets named
``_…`` and ``REVIEW_SUMMARY`` are skipped) against the generated sheet of
the same name, row pair by row pair, with the same per-column rules
(``score``). ``--real-table`` / ``--sheet`` are not used. Rows pair by KEY:

* DATA_FACTORY_PIPELINE_SCHEDULE, DATABRICKS_NOTEBOOK_DETAILS — PIPELINE_NAME;
* ADLS_DELTA_INGESTION_DETAILS — SRC_FILE_NAME (one row per file);
* STGDELTA_STDDELTA_INGESTION_DET — TGT_TABLE_NAME (one row per table);
* DATA_QUALITY_RULES — (file, RULE_CLASS), the file being the SRC_FILE_NAME
  of the ADLS row of the SAME workbook whose OBJECT_ID is the DQ row's
  (OBJECT_IDs are never compared across workbooks: a golden masks them
  ``SYN-OBJ-00n``, a run numbers them 1..n); equal keys pair in SEQUENCE_NO
  order;
* ADLS_FIXED_WIDTH_HANDLER — SEGMENT; EMAIL_TEMPLATE_CONFIG — STATUS;
* FILE_ADLS_INGESTION_DETAILS (and any sheet not listed) — position.

Keys compare case-insensitively, first as written, then with the generated
key de-aliased (``--alias``, as ``_find``). A BLANK key cell (an open cell)
never decides a pair: such rows pair in sheet order with the rows left over
whose non-blank key cells agree. A row left without a partner is reported
as ``real-only`` / ``generated-only`` (its key, never its values) and each
of its cells — one per column of the real sheet — counts as UNMATCHED.

Output per sheet: a ``== <SHEET>`` line; per row pair a ``-- real r<n> /
generated r<m>: <key>`` line (worksheet rows), its non-MATCH column lines
(``--matches`` adds the MATCH lines) and the per-row summary in the
single-sheet format; a ``-- real r<n>: real-only`` / ``generated-only`` line
per unmatched row; then the sheet's score line. ``--summary-only`` prints
the score lines only (and the sheets not scored). The score line (per sheet,
and a final TOTAL line):
``<SHEET>: score S; paired P, real-only R, generated-only G; cells C,
filled F, matched M, alias A, open K (BSA x, …), diff D, unmatched U`` with
score = 100 × (M + A) / C, one decimal (C = the cells of the paired rows +
the unmatched cells, one per real column; OPEN counts against the score). A
sheet with no row on either side scores 100.0. Generated sheets the real
workbook lacks are listed, not scored. CREATED_BY / UPDATED_BY values are
withheld here too.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

WITHHELD = {"CREATED_BY", "UPDATED_BY", "CRETAED_BY"}
LIST_COLUMNS = {"SRC_COLUMNS", "SRC_DATA_TYPE", "TGT_COLUMN_NAMES", "TGT_DATA_TYPE",
                "MANDATORY_FIELD_LIST", "TGT_PRIMARY_KEY", "SRC_COL_LNGTH",
                "SRC_COL_STRT_END_INDX"}
NULLS = {"", "null", "none"}
# The CV golden fixture's anonymisation of the Socially Determined feeds.
DEFAULT_ALIASES = ["cv_=sd_", "sdh=sdoh", "civic_vantage=socially_determined",
                   "civic vantage=socially determined"]
# Type lists: the real rows mix 'String' / 'string' — compared case-insensitively.
CASELESS_COLUMNS = {"SRC_DATA_TYPE", "TGT_DATA_TYPE"}
# Columns the CV fixture anonymised outright (a value, not a name fragment):
# a difference there is ALIAS when the generated row IS the fixture.
FIXTURE_ANONYMISED = {"LOB"}
# A fixture's masked sequence id (SYN-OBJ-001) = a run's position number (1).
MASKED_SEQUENCE = re.compile(r"SYN-[A-Z]+-(\d+)")
OWNER_BUCKETS = {"BSA": "BSA", "Engineer": "Engineer", "Engineer (confirm)": "Engineer-confirm",
                 "Set at load (CI/CD)": "CI/CD"}
WIDTH = 60


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _is_null(value) -> bool:
    return _text(value).lower() in NULLS


def _cut(value: str) -> str:
    return value if len(value) <= WIDTH else value[:WIDTH - 3] + "..."


def _rows(path: Path, sheet: str | None) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    ws = workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    headers = [_text(h) for h in rows[0]]
    return [{h: v for h, v in zip(headers, r, strict=False) if h} for r in rows[1:]]


def _owners(review: Path | None, sheet: str) -> dict[str, str]:
    if review is None or not review.is_file():
        return {}
    owners: dict[str, str] = {}
    for row in _rows(review, "REVIEW_SUMMARY"):
        if _text(row.get("Sheet")) == sheet:
            owners.setdefault(_text(row.get("Column")), _text(row.get("Owner")))
    return owners


def _dealias(value: str, aliases: list[tuple[str, str]]) -> str:
    out = value
    for generated, real in aliases:
        out = re.sub(re.escape(generated), real, out, flags=re.IGNORECASE)
    return out.lower()


def _masked_sequence(value: str) -> int | None:
    """``n`` of a fixture-masked sequence id ``SYN-<WORD>-<n>`` (the pair-1
    golden's OBJECT_IDs, ``SYN-OBJ-001``), else None."""
    match = MASKED_SEQUENCE.fullmatch(value)
    return int(match.group(1)) if match else None


def _items(value: str) -> list[str]:
    return [i.strip() for i in value.split(",") if i.strip()]


def _diff_text(column: str, real: str, generated: str) -> str:
    if column in LIST_COLUMNS:
        r, g = _items(real), _items(generated)
        text = (f"real {len(r)} items {[_cut(i) for i in r[:3]]} | "
                f"generated {len(g)} items {[_cut(i) for i in g[:3]]}")
        first = next((i for i, (a, b) in enumerate(zip(r, g, strict=False)) if a != b), None)
        if first is not None and first >= 3:
            text += (f"; first difference at item {first + 1}: {_cut(r[first])!r} vs "
                     f"{_cut(g[first])!r}")
        return text
    return f"real {_cut(real)!r} | generated {_cut(generated)!r}"


def _find(rows: list[dict], table: str, aliases) -> dict | None:
    for row in rows:
        if _text(row.get("TGT_TABLE_NAME")).lower() == table.lower():
            return row
    for row in rows:
        if _dealias(_text(row.get("TGT_TABLE_NAME")), aliases) == table.lower():
            return row
    return None


def score(generated: dict, real: dict, owners: dict[str, str],
          aliases: list[tuple[str, str]], fixture: bool = False) -> tuple[list[str], dict]:
    """``fixture`` = the generated row is the anonymised fixture (its table
    name differs from the real one)."""
    lines, counts, _open_by = _score(generated, real, owners, aliases, fixture)
    return lines, counts


def _score(generated: dict, real: dict, owners: dict[str, str],
           aliases: list[tuple[str, str]], fixture: bool) -> tuple[list[str], dict, dict]:
    """``score`` plus the OPEN cells per owner bucket (the all-sheets totals)."""
    lines: list[str] = []
    counts = {"filled": 0, "match": 0, "alias": 0, "open": 0, "diff": 0}
    open_by: dict[str, int] = {}
    columns = list(real)
    for column in columns:
        real_value, gen_value = _text(real.get(column)), _text(generated.get(column))
        withheld = column in WITHHELD
        if gen_value:
            counts["filled"] += 1
        if not gen_value:
            if _is_null(real_value):
                status, detail = "MATCH", "(both blank / NULL)"
                counts["match"] += 1
            else:
                owner = owners.get(column) or ("not in the template" if column not in generated
                                               else "no review entry")
                bucket = OWNER_BUCKETS.get(owner, owner)
                open_by[bucket] = open_by.get(bucket, 0) + 1
                status, detail = "OPEN", f"owner: {owner}"
                counts["open"] += 1
        elif gen_value == real_value or (_is_null(gen_value) and _is_null(real_value)):
            status, detail = "MATCH", ""
            counts["match"] += 1
        elif column in CASELESS_COLUMNS and gen_value.lower() == real_value.lower():
            status = "MATCH"
            detail = "(case-insensitive; " + _diff_text(column, real_value, gen_value) + ")"
            counts["match"] += 1
        elif _masked_sequence(real_value) is not None and gen_value.isdecimal() and (
                int(gen_value) == _masked_sequence(real_value)):
            status = "ALIAS"
            detail = "" if withheld else f"(masked sequence) generated {_cut(gen_value)!r}"
            counts["alias"] += 1
        elif fixture and column in FIXTURE_ANONYMISED:
            status = "ALIAS"
            detail = f"(anonymised in the fixture) generated {_cut(gen_value)!r}"
            counts["alias"] += 1
        elif (_dealias(gen_value, aliases) == real_value.lower()
              and _dealias(gen_value, aliases) != gen_value.lower()):
            # equal only once the fixture's anonymisation is undone
            status, detail = "ALIAS", ("" if withheld else f"generated {_cut(gen_value)!r}")
            counts["alias"] += 1
        else:
            status = "DIFF"
            detail = "(values withheld)" if withheld else _diff_text(column, real_value,
                                                                     gen_value)
            if not withheld and _dealias(gen_value, aliases) == real_value.lower():
                detail = "(case only) " + detail
            counts["diff"] += 1
        lines.append(f"{status:<5} {column:<26} {detail}".rstrip())
    summary = (f"filled {counts['filled']}/{len(columns)}, matched {counts['match']}, "
               f"alias {counts['alias']}, open {counts['open']} ({_open_parts(open_by)}), "
               f"diff {counts['diff']}")
    return lines + [summary], counts, open_by


KNOWN_OWNERS = ["BSA", "Engineer", "Engineer-confirm", "CI/CD"]


def _open_parts(open_by: dict[str, int]) -> str:
    parts = [f"{name} {open_by.get(name, 0)}" for name in KNOWN_OWNERS]
    parts += [f"{name} {n}" for name, n in open_by.items() if name not in KNOWN_OWNERS]
    return ", ".join(parts)


# -- all-sheets mode (multi-table step 7) ------------------------------------------ #

ADLS_SHEET = "ADLS_DELTA_INGESTION_DETAILS"
DQ_SHEET = "DATA_QUALITY_RULES"
DQ_FILE = "file"           # the DQ key's pseudo column: the ADLS SRC_FILE_NAME of OBJECT_ID
# How rows pair, per sheet (the module docstring says why); () = by position.
SHEET_KEYS: dict[str, tuple[str, ...]] = {
    "DATA_FACTORY_PIPELINE_SCHEDULE": ("PIPELINE_NAME",),
    "FILE_ADLS_INGESTION_DETAILS": (),
    ADLS_SHEET: ("SRC_FILE_NAME",),
    "STGDELTA_STDDELTA_INGESTION_DET": ("TGT_TABLE_NAME",),
    "ADLS_FIXED_WIDTH_HANDLER": ("SEGMENT",),
    "DATABRICKS_NOTEBOOK_DETAILS": ("PIPELINE_NAME",),
    DQ_SHEET: (DQ_FILE, "RULE_CLASS"),
    "EMAIL_TEMPLATE_CONFIG": ("STATUS",),
}
SKIPPED_SHEETS = {"REVIEW_SUMMARY"}
COUNT_KEYS = ("filled", "match", "alias", "open", "diff", "unmatched", "cells")


@dataclass
class Sheet:
    headers: list[str]
    rows: list[tuple[int, dict]]        # (worksheet row number, {header: value})


@dataclass
class SheetScore:
    name: str
    keys: tuple[str, ...]
    pairs: list[tuple[int, int]] = field(default_factory=list)   # (real row, generated row)
    real_only: list[int] = field(default_factory=list)
    generated_only: list[int] = field(default_factory=list)
    counts: dict[str, int] = field(default_factory=lambda: dict.fromkeys(COUNT_KEYS, 0))
    open_by: dict[str, int] = field(default_factory=dict)
    lines: list[str] = field(default_factory=list)

    @property
    def score(self) -> float:
        return _percent(self.counts)


def _percent(counts: dict[str, int]) -> float:
    if not counts["cells"]:
        return 100.0                     # no row on either side: the sheets agree
    return 100.0 * (counts["match"] + counts["alias"]) / counts["cells"]


def _workbook(path: Path) -> dict[str, Sheet]:
    """Every sheet: headers (blank ones dropped) and its non-empty rows, each
    padded to the header width."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: dict[str, Sheet] = {}
        for ws in workbook.worksheets:
            values = list(ws.iter_rows(values_only=True))
            headers = [_text(h) for h in values[0]] if values else []
            rows = []
            for number, row in enumerate(values[1:], start=2):
                if all(_text(v) == "" for v in row):
                    continue
                rows.append((number, {h: (row[i] if i < len(row) else None)
                                      for i, h in enumerate(headers) if h}))
            sheets[ws.title] = Sheet([h for h in headers if h], rows)
        return sheets
    finally:
        workbook.close()


def _owners_by_sheet(review: Path | None) -> dict[str, dict[str, str]]:
    if review is None or not review.is_file():
        return {}
    owners: dict[str, dict[str, str]] = {}
    for row in _rows(review, "REVIEW_SUMMARY"):
        sheet = owners.setdefault(_text(row.get("Sheet")), {})
        sheet.setdefault(_text(row.get("Column")), _text(row.get("Owner")))
    return owners


def _row_keys(name: str, sheet: Sheet, book: dict[str, Sheet]) -> list[tuple[str, ...]]:
    keys = SHEET_KEYS.get(name, ())
    files: dict[str, str] = {}
    if DQ_FILE in keys and ADLS_SHEET in book:
        for _number, row in book[ADLS_SHEET].rows:      # OBJECT_ID -> file, same workbook
            files.setdefault(_text(row.get("OBJECT_ID")), _text(row.get("SRC_FILE_NAME")))
    out = []
    for _number, row in sheet.rows:
        object_id = _text(row.get("OBJECT_ID"))
        out.append(tuple((files.get(object_id, "") if object_id else "") if k == DQ_FILE
                         else _text(row.get(k)) for k in keys))
    return out


def _order(name: str, sheet: Sheet) -> list[tuple]:
    """Pairing order: sheet order; DQ rows by SEQUENCE_NO first (equal keys)."""
    out = []
    for index, (_number, row) in enumerate(sheet.rows):
        if name == DQ_SHEET:
            seq = _text(row.get("SEQUENCE_NO"))
            try:
                out.append((0, float(seq), "", index))       # '2' and a numeric cell 2.0
            except ValueError:
                out.append((1, 0.0, seq, index))
        else:
            out.append((index,))
    return out


def _same(real: str, generated: str, aliases) -> bool:
    return real.lower() == generated.lower() or _dealias(generated, aliases) == real.lower()


def pair_rows(real_keys: list[tuple[str, ...]], gen_keys: list[tuple[str, ...]],
              aliases, real_order: list | None = None, gen_order: list | None = None,
              ) -> tuple[list[tuple[int, int]], list[int], list[int]]:
    """Index pairs (real, generated) + the leftovers of each side.

    1. keys equal as written (case-insensitive); 2. equal once the generated
    key is de-aliased; 3. rows with a blank key cell pair, in order, with the
    rows left over whose non-blank key cells agree. Phases 1–2 only pair rows
    whose every key cell is filled; a keyless sheet pairs in phase 3 only."""
    real_order = real_order or list(range(len(real_keys)))
    gen_order = gen_order or list(range(len(gen_keys)))
    real_seq = sorted(range(len(real_keys)), key=lambda i: real_order[i])
    gen_seq = sorted(range(len(gen_keys)), key=lambda i: gen_order[i])
    used_real: set[int] = set()
    used_gen: set[int] = set()
    pairs: list[tuple[int, int]] = []

    def full(key) -> bool:
        return bool(key) and all(key)

    def agree(real, generated, same) -> bool:
        return all(same(a, b) for a, b in zip(real, generated, strict=True))

    def exact(a, b) -> bool:
        return a.lower() == b.lower()

    def aliased(a, b) -> bool:
        return _same(a, b, aliases)

    def blank_or_aliased(a, b) -> bool:
        return not a or not b or _same(a, b, aliases)

    phases = [
        lambda r, g: full(r) and full(g) and agree(r, g, exact),
        lambda r, g: full(r) and full(g) and agree(r, g, aliased),
        lambda r, g: agree(r, g, blank_or_aliased),
    ]
    for accepts in phases:
        for i in real_seq:
            if i in used_real:
                continue
            for j in gen_seq:
                if j not in used_gen and accepts(real_keys[i], gen_keys[j]):
                    pairs.append((i, j))
                    used_real.add(i)
                    used_gen.add(j)
                    break
    pairs.sort()
    return (pairs, [i for i in range(len(real_keys)) if i not in used_real],
            [j for j in range(len(gen_keys)) if j not in used_gen])


def _key_text(keys: tuple[str, ...], values: tuple[str, ...]) -> str:
    if not keys:
        return "by position"
    return ", ".join(f"{k} " + ("(withheld)" if k in WITHHELD else repr(_cut(v)) if v
                                else "(blank)") for k, v in zip(keys, values, strict=True))


def score_sheet(name: str, real: Sheet, generated: Sheet | None, real_book: dict,
                gen_book: dict, owners: dict[str, str], aliases,
                show_matches: bool = False, summary_only: bool = False) -> SheetScore:
    generated = generated or Sheet([], [])
    result = SheetScore(name, SHEET_KEYS.get(name, ()))
    real_keys = _row_keys(name, real, real_book)
    gen_keys = _row_keys(name, generated, gen_book)
    pairs, real_only, gen_only = pair_rows(real_keys, gen_keys, aliases,
                                           _order(name, real), _order(name, generated))
    width = len(real.headers)
    key_label = ", ".join(result.keys) or "position"
    if DQ_FILE in result.keys:
        key_label = key_label.replace(DQ_FILE, "file (ADLS SRC_FILE_NAME of OBJECT_ID)", 1)
    if not summary_only:
        result.lines.append(f"== {name}: rows by {key_label}; real {len(real.rows)}, "
                            f"generated {len(generated.rows)}"
                            + ("" if name in gen_book else " (sheet missing from generated)"))
    for i, j in pairs:
        (real_no, real_row), (gen_no, gen_row) = real.rows[i], generated.rows[j]
        result.pairs.append((real_no, gen_no))
        real_table, gen_table = _text(real_row.get("TGT_TABLE_NAME")), _text(
            gen_row.get("TGT_TABLE_NAME"))
        fixture = bool(real_table and gen_table and real_table.lower() != gen_table.lower())
        lines, counts, open_by = _score(gen_row, real_row, owners, aliases, fixture)
        for key in ("filled", "match", "alias", "open", "diff"):
            result.counts[key] += counts[key]
        result.counts["cells"] += width
        for bucket, n in open_by.items():
            result.open_by[bucket] = result.open_by.get(bucket, 0) + n
        if summary_only:
            continue
        result.lines.append(f"-- real r{real_no} / generated r{gen_no}: "
                            + _key_text(result.keys, real_keys[i]))
        result.lines += [line for line in lines[:-1]
                         if show_matches or not line.startswith("MATCH")]
        result.lines.append(lines[-1])
    for side, leftovers, sheet, keys in (("real", real_only, real, real_keys),
                                         ("generated", gen_only, generated, gen_keys)):
        for index in leftovers:
            number = sheet.rows[index][0]
            getattr(result, f"{side}_only").append(number)
            result.counts["unmatched"] += width
            result.counts["cells"] += width
            if not summary_only:
                result.lines.append(f"-- {side} r{number}: {side}-only ("
                                    f"{_key_text(result.keys, keys[index])}); {width} cells "
                                    "unmatched")
    result.lines.append(f"{name}: " + _summary([result]))
    return result


def _summary(results: list[SheetScore]) -> str:
    """'score S; paired P, real-only R, generated-only G; cells C, filled F,
    matched M, alias A, open K (owners), diff D, unmatched U' over ``results``."""
    counts = dict.fromkeys(COUNT_KEYS, 0)
    open_by: dict[str, int] = {}
    for result in results:
        for key in COUNT_KEYS:
            counts[key] += result.counts[key]
        for bucket, n in result.open_by.items():
            open_by[bucket] = open_by.get(bucket, 0) + n
    paired = sum(len(r.pairs) for r in results)
    real_only = sum(len(r.real_only) for r in results)
    generated_only = sum(len(r.generated_only) for r in results)
    return (f"score {_percent(counts):.1f}; paired {paired}, real-only {real_only}, "
            f"generated-only {generated_only}; cells {counts['cells']}, "
            f"filled {counts['filled']}, matched {counts['match']}, alias {counts['alias']}, "
            f"open {counts['open']} ({_open_parts(open_by)}), diff {counts['diff']}, "
            f"unmatched {counts['unmatched']}")


def score_all(generated: Path, real: Path, review: Path | None, aliases,
              show_matches: bool = False, summary_only: bool = False,
              ) -> tuple[list[SheetScore], list[str]]:
    """Every sheet of the REAL workbook scored; the output lines last."""
    real_book, gen_book = _workbook(real), _workbook(generated)
    owners = _owners_by_sheet(review)
    results = [score_sheet(name, sheet, gen_book.get(name), real_book, gen_book,
                           owners.get(name, {}), aliases, show_matches, summary_only)
               for name, sheet in real_book.items()
               if not name.startswith("_") and name not in SKIPPED_SHEETS]
    lines = [line for result in results for line in result.lines]
    extra = [name for name in gen_book
             if name not in real_book and not name.startswith("_")
             and name not in SKIPPED_SHEETS]
    if extra:
        lines.append("not in the real workbook (not scored): " + ", ".join(extra))
    lines.append(f"TOTAL ({len(results)} sheets): " + _summary(results))
    return results, lines


def total_score(results: list[SheetScore]) -> float:
    counts = dict.fromkeys(COUNT_KEYS, 0)
    for result in results:
        for key in COUNT_KEYS:
            counts[key] += result.counts[key]
    return _percent(counts)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--generated", required=True, type=Path)
    parser.add_argument("--real", required=True, type=Path)
    parser.add_argument("--sheet", default="ADLS_DELTA_INGESTION_DETAILS")
    parser.add_argument("--real-sheet", default=None,
                        help="sheet of the real workbook (default: --sheet, else the first)")
    parser.add_argument("--real-table", default=None,
                        help="TGT_TABLE_NAME of the real row (single-sheet mode: required)")
    parser.add_argument("--generated-table", default=None)
    parser.add_argument("--review", type=Path, default=None,
                        help="review copy (default: <generated>_REVIEW.xlsx next to it)")
    parser.add_argument("--alias", action="append", default=None,
                        help="generated=real anonymisation pair (repeatable)")
    parser.add_argument("--all-sheets", action="store_true",
                        help="score every sheet of the real workbook, rows paired by key")
    parser.add_argument("--summary-only", action="store_true",
                        help="all-sheets: only the per-sheet and TOTAL lines")
    parser.add_argument("--matches", action="store_true",
                        help="all-sheets: print the MATCH column lines too")
    args = parser.parse_args(argv)

    aliases = [tuple(a.split("=", 1)) for a in (args.alias or DEFAULT_ALIASES)]
    review = args.review or args.generated.with_name(
        args.generated.name.replace("_IIG.xlsx", "_IIG_REVIEW.xlsx"))
    if args.all_sheets:
        results, lines = score_all(args.generated, args.real, review, aliases,
                                   show_matches=args.matches, summary_only=args.summary_only)
        if not results:
            print(f"no IIG sheet in {args.real}", file=sys.stderr)
            return 2
        print("\n".join(lines))
        return 0
    if not args.real_table:
        parser.error("--real-table is required (or pass --all-sheets)")
    real = _find(_rows(args.real, args.real_sheet or args.sheet), args.real_table, [])
    if real is None:
        print(f"real row {args.real_table!r} not found", file=sys.stderr)
        return 2
    generated = _find(_rows(args.generated, args.sheet),
                      args.generated_table or args.real_table, aliases)
    if generated is None:
        print(f"generated row {args.generated_table or args.real_table!r} not found",
              file=sys.stderr)
        return 2
    fixture = _text(generated.get("TGT_TABLE_NAME")).lower() != args.real_table.lower()
    lines, _counts = score(generated, real, _owners(review, args.sheet), aliases, fixture)
    print(f"{args.sheet}: real {args.real_table} vs generated "
          f"{_text(generated.get('TGT_TABLE_NAME'))}")
    print("\n".join(lines))
    return 0


if __name__ == "__main__":
    sys.exit(main())
