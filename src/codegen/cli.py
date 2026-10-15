"""codegen CLI — generate Databricks ingestion pipelines from contracts.

Commands:
  generate      one FRD+STTM pair (optionally one feed of it)
  generate-all  every pair listed in config.contracts.pairs
  extract-sttm  extract an STTM mapping contract from a client workbook
                paired with its FRD feed contract (deterministic, no LLM)
  extract-frd   extract an FRD feed contract from a client FRD .docx
                (families F1/F2, deterministic, stdlib docx reading)
  sharepoint-fetch    library -> local input dir (workbooks + contracts)
  sharepoint-publish  one feed's generated artifacts -> library output folder
  databricks-fetch    UC volumes -> local input dir (FRDs + STTM workbooks)
  databricks-publish  one feed's generated artifacts -> UC output volume
                      (human-gated; WRITABLE_PREFIX enforced in code)
  demo-source-files   the demo UI's source-files display JSON (pure read)

The two sharepoint-* commands are the transport seam at the edges; the
generation path between them never opens a socket (codegen/sharepoint.py).

Per feed: resolve -> compile rules -> Layer 2 over unmapped rules -> render
templates -> write candidates artifact -> gate -> report. A feed that fails
resolution or rendering is reported FAIL and does not stop other feeds.
Exit code is 0 only when no feed FAILed (see EXIT_CODES).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from codegen.config import Config, load_config, load_dotenv
from codegen.contracts.resolved import ResolvedFeedSpec
from codegen.emit.context import TemplateGapError, build_context
from codegen.emit.emitter import emit_feed
from codegen.faq import faq_for_spec, with_faq_answers
from codegen.gate import compute_verdict, run_generated_tests, run_preflight
from codegen.gate.verdict import GateResult
from codegen.reasoning import build_provider, run_reasoning
from codegen.reasoning.engine import RuleCandidate, segmented_review_items
from codegen.report import console_summary, write_generation_report
from codegen.report.generation_report import all_failed_line
from codegen.resolve.resolver import ContractMismatchError, normalize_feed_name, resolve_pair
from codegen.rules.compiler import compile_rules
from codegen.run_errors import Coord, report_error


def _write_candidates_artifact(candidates: list[RuleCandidate], feed_dir: Path) -> Path | None:
    """Layer-2 output goes here — a review artifact, never a generated module."""
    if not candidates:
        return None
    artifact_dir = feed_dir / "candidates"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    artifact_path = artifact_dir / "candidates.json"
    payload = []
    for candidate in candidates:
        entry = candidate.model_dump()
        if entry.get("kind") == "layer2":
            # Default-kind entries serialize exactly as before the segmented
            # dialect landed — flat candidates.json stays byte-identical.
            entry.pop("kind", None)
            entry.pop("detail", None)
            entry.pop("citation", None)
        payload.append(entry)
    artifact_path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8", newline="\n")
    return artifact_path


def _generate_feed(
    spec: ResolvedFeedSpec,
    config: Config,
    *,
    dry_run: bool,
    skip_tests: bool,
    output_mode: str | list[str] | None = None,
    extra_flags: list[str] | None = None,
    conventions_profile: str | None = None,
    iig_template: str | None = None,
    playbook_template: str | None = None,
    layout_usage=None,
    faq_answers: dict[str, str] | None = None,
) -> GateResult:
    out_root = Path(config.output.dir)
    reports_dir = Path(config.output.reports_dir)
    feed_dir = out_root / spec.feed_slug

    outcomes = compile_rules(spec)
    provider = build_provider(config, dry_run)
    candidates = run_reasoning(spec, outcomes, provider)
    from codegen.reasoning.raw_log import log_layer2_failures

    # 2026-10-14: every failed Layer-2 attempt on the log (reason + RAW, D1-only)
    log_layer2_failures(provider, spec.feed_slug, config, Path("."))
    from codegen.reasoning.mode import DRY_RUN_OVER_LIVE, provider_mode

    if dry_run and provider_mode(config).live:
        # The report says it too: never a dry run that reads like a live one.
        extra_flags = [f"provider_overridden: {DRY_RUN_OVER_LIVE}", *(extra_flags or [])]
    from codegen.reasoning.usage import layer2_usage

    model_usage = [*([layout_usage] if layout_usage is not None else []),
                   layer2_usage(provider, candidates)]
    # M3: STTM-vs-VDD cross-check — one flag per mismatch citing both cells;
    # a fixed-width FRD with no VDD positions is a failed gate check.
    # M4: resolver provenance flags (facts taken from the STTM because the
    # FRD named none) and the drag-fill detector ride the same list.
    from codegen.gate.derivations import sibling_type_flags
    from codegen.gate.drag_fill import drag_fill_flags
    from codegen.gate.vdd_check import vdd_cross_check

    vdd_flags, vdd_check = vdd_cross_check(spec, config)
    # M7.1: correctness gates are GLOBAL — sibling-type consistency is a flag
    # (never FAIL); the cell / literal / cap checks join the gate after the
    # framework emit for every profile.
    # (de-duplicated: the UI passes the layout stage's flags, and the FRD
    # contract the layout stage wrote carries the same ones — M9.3)
    extra_flags = [*dict.fromkeys([*(extra_flags or []), *spec.provenance_flags]),
                   *drag_fill_flags(spec),
                   *sibling_type_flags(spec, config), *vdd_flags]
    # Segmented-extraction review items (assumption/conflict cards) ride the
    # same review artifact and decision flow as Layer-2 candidates.
    candidates = [*segmented_review_items(spec), *candidates]
    # Three-input model: file answers (fixtures/faq/<slug>.faq.yaml) plus
    # contract prefills; missing file => all defaults, flagged by the gate.
    faq = with_faq_answers(faq_for_spec(spec, config), faq_answers)

    duplicate_outcome = next(
        (o for o in outcomes if o.feature == "allow_duplicate_file_name"), None
    )
    context = build_context(
        spec,
        config,
        # True both ways today: re-ingest is idempotent (MERGE), and the only
        # known duplicate-file rule turns the check off. The flag exists so a
        # future "check on" rule can flip it.
        allow_duplicate_file_name=True,
        duplicate_rule_text=(
            duplicate_outcome.rule_text if duplicate_outcome is not None else None
        ),
        notification_rule_texts=[
            o.rule_text for o in outcomes if o.classification == "notification"
        ],
        faq=faq,
    )
    # Option A ("notebook") is today's path, byte for byte. Option B
    # ("framework") renders the SAME pipeline into a scratch tree so the
    # gate checks stay identical, but persists only ddl/ + framework/;
    # Notebook + Framework persists everything. MIRRORED in
    # service._generate_feed. (codegen.output_modes: the two outputs.)
    from codegen.output_modes import output_parts

    parts = output_parts(output_mode or config.output.mode, source="generate")
    framework_artefacts = None
    if parts == ["framework"]:
        written, checks, tests_skipped, ddl_sources = _emit_framework_only(
            context, spec, config, feed_dir, skip_tests
        )
        framework_artefacts = _run_emit_framework(
            spec, faq, ddl_sources, config, out_root, outcomes, base_dir=None,
            conventions_profile=conventions_profile, iig_template=iig_template,
        )
        written = [*written, *framework_artefacts.files]
    else:
        written = emit_feed(context, out_root)
        checks = None  # computed below, exactly as before
        if "framework" in parts:
            ddl_sources = _read_ddl_sources(feed_dir)
            framework_artefacts = _run_emit_framework(
                spec, faq, ddl_sources, config, out_root, outcomes, base_dir=None,
                conventions_profile=conventions_profile, iig_template=iig_template,
            )
            written = [*written, *framework_artefacts.files]
    if framework_artefacts is not None:
        extra_flags = [*extra_flags, *framework_artefacts.flags]
    _write_candidates_artifact(candidates, feed_dir)

    if checks is None:
        checks = run_preflight(feed_dir, config)
        tests_skipped = skip_tests or not config.gate.run_generated_tests
        if not tests_skipped:
            checks = [*checks, run_generated_tests(feed_dir, config.gate.pytest_tail_lines)]
    if vdd_check is not None:
        checks = [*checks, vdd_check]
    if framework_artefacts is not None:
        checks = [*checks, *framework_artefacts.checks]   # M7.1 derivation gate (global)

    gate = compute_verdict(
        spec.feed_id,
        outcomes,
        candidates,
        checks,
        tests_skipped,
        faq=faq,
        standards=config.engineering_standards,
        extra_flags=extra_flags,
    )
    write_generation_report(
        spec,
        written,
        outcomes,
        candidates,
        gate,
        reports_dir,
        out_root,
        inputs_summary=context["provenance"]["inputs"],
        model_usage=model_usage,
    )
    if dry_run and provider_mode(config).live:
        # The report's FIRST line too (2026-10-13): a D1 report must never look
        # live when --dry-run answered from the mock.
        report = reports_dir / f"{spec.feed_slug}.md"
        report.write_text(f"> **WARNING** {DRY_RUN_OVER_LIVE}\n\n"
                          + report.read_text(encoding="utf-8"), encoding="utf-8", newline="\n")
    if framework_artefacts is not None:
        # Appended AFTER the standard report so report/ stays untouched and
        # notebook-mode reports remain byte-identical.
        from codegen.emit.framework import report_section

        with open(reports_dir / f"{spec.feed_slug}.md", "a",
                  encoding="utf-8", newline="\n") as handle:
            handle.write(report_section(framework_artefacts))
    print(console_summary(spec, gate))
    for usage in model_usage:
        print(f"{'MODEL USAGE':<15} {usage.stage}: {usage.label}")
        if usage.all_failed:
            print(f"{'WARNING':<15} {all_failed_line(usage)}")
    return gate


def _read_ddl_sources(feed_dir: Path) -> list[tuple[str, str]]:
    return sorted(
        (p.name, p.read_text(encoding="utf-8"))
        for p in (feed_dir / "ddl").glob("*.sql")
    )


def _emit_framework_only(context, spec, config, feed_dir, skip_tests):
    """Framework mode: full render + gate checks in a scratch tree (so the
    verdict is IDENTICAL to notebook mode), persisting only ddl/."""
    import shutil
    import tempfile

    with tempfile.TemporaryDirectory() as tmp:
        tmp_root = Path(tmp)
        emit_feed(context, tmp_root, framework_only=True)
        tmp_feed_dir = tmp_root / spec.feed_slug
        checks = run_preflight(tmp_feed_dir, config)
        tests_skipped = skip_tests or not config.gate.run_generated_tests
        if not tests_skipped:
            checks = [*checks,
                      run_generated_tests(tmp_feed_dir, config.gate.pytest_tail_lines)]
        ddl_sources = _read_ddl_sources(tmp_feed_dir)
        feed_dir.mkdir(parents=True, exist_ok=True)
        shutil.copytree(tmp_feed_dir / "ddl", feed_dir / "ddl", dirs_exist_ok=True)
    written = [feed_dir / "ddl" / name for name, _ in ddl_sources]
    return written, checks, tests_skipped, ddl_sources


def _run_emit_framework(spec, faq, ddl_sources, config, out_root, outcomes,
                        base_dir, conventions_profile=None, iig_template=None):
    from codegen.emit.framework import emit_framework

    contracts_dir = Path(config.contracts.dir)
    if base_dir is not None:
        contracts_dir = base_dir / contracts_dir
    frd_path = contracts_dir / spec.frd_contract_name
    artefacts = emit_framework(
        spec, faq, ddl_sources, config, out_root,
        base_dir=base_dir,
        frd_path=frd_path if frd_path.is_file() else None,
        unmapped_rule_texts={
            o.rule_text for o in outcomes if o.classification == "unmapped"
        },
        conventions_profile=conventions_profile,
        iig_template=iig_template,
    )
    from codegen.emit.emitter import unknown_audit_columns

    unknown = unknown_audit_columns(spec.audit_columns)
    if unknown:
        # Chunk B (cold drill): carried as the STTM names them; the framework,
        # not the generated notebook pipeline, populates audit columns.
        artefacts.flags.append(
            f"audit_column_framework_populated:{','.join(unknown)} \u2014 the STTM's audit "
            "rows name them; the DDL / IIG carry them as named, no notebook-pipeline "
            "template populates them (the framework does; notebook mode stops on them)")
    return artefacts


# A feed held back because no document states one of its targets (the
# extractor's NEEDS_ANSWERS): not a FAIL — the run stops short of it and names
# the answers-file key. Exit code when nothing FAILED but a feed needs answers.
EXIT_OK = 0
EXIT_FAILED = 1
EXIT_NEEDS_ANSWERS = 3
EXIT_CODES = f"""exit codes:
  {EXIT_OK}  ok — every feed it produced or was asked for is done (PASS / PASS_WITH_FLAGS);
     a sheet held back for an answer still prints its QUESTION line
  {EXIT_FAILED}  failed — a feed FAILed its gate, or a command could not complete
  {EXIT_NEEDS_ANSWERS}  NEEDS_ANSWERS — layout --require-complete: columns left unplaced / questions
     open; extract-sttm: NO feed produced a usable contract (none written);
     generate: a feed it was asked for (--feed) has no contract
     pair: the FRD is not decided (a QUESTION pair.frd line)
  Every answer OWED (the run cannot complete a feed without it) is one line,
  label padded to 15 columns, then the answers.yaml key, an em-dash, the reason
  (informational items — optional columns unplaced, FRD fields not read, the
  VDD pairing — use the same form with the label NOTE; the harness ignores it):
     QUESTION       feeds[2].stage_target.tables — <reason>
     UNRESOLVED     <sheet>/<layer>/<role> — <reason>
  Answer them (`gaps:` / `answers:` in answers.yaml) and re-run
  (`codegen layout --answers answers.yaml`).
"""

def _fail(stage: str, text: str, reason: str, coord: Coord | None = None, *,
          exc: BaseException | None = None, checks=(), path=None) -> None:
    """One FAIL line and its error marker (D1 post51, 2026-10-15). The line
    (stdout) keeps the quoted text — :func:`codegen.run_errors.first_error_line_full`
    reads it, inside ACFC only; the ``CODEGEN-ERROR`` marker (stderr) carries the
    stage, a reason code and coordinates only — what
    :func:`codegen.run_errors.first_error_line` (the harness, a pushed report)
    reads. An exception's own coordinates (``run_errors.located``) win; ``path``
    becomes ``<workdir>/pair_NN/…`` there (``run_errors.path_hint``)."""
    if isinstance(exc, FileNotFoundError) and reason == "unclassified":
        reason = "input.not_found"
    print(f"{'FAIL':<15} {text}")
    report_error(stage, reason, coord, exc=exc, checks=checks, path=path)


# The ACFC harness (codegen-watch, docs/acfc/HARNESS_EXIT_CODES.md on
# origin/acfc/harness-exit3) reads the needed answer keys from these lines:
# the key is the text between the padded label and the em-dash separator.
ANSWER_LABEL_WIDTH = 15


def answer_line(label: str, key: str, reason: str) -> str:
    """One missing answer in the form the harness parses (its _NEEDED_KEY_RE
    captures everything between the padded label and ' \u2014 '): ``label``
    padded to 15 columns, the answers.yaml key VERBATIM (a sheet name's spaces
    kept; only a line break becomes a space), ' \u2014 ', the reason on one line."""
    def one(text) -> str:
        return " ".join(str(text).split())

    verbatim = re.sub(r"[\r\n]+", " ", str(key)).strip()
    return f"{label:<{ANSWER_LABEL_WIDTH}}{verbatim} \u2014 {one(reason)}"


# FRD fields the run cannot proceed without (the resolver's hard stops): an
# unread one is an answer owed only while the pair-resolved FRD contract still
# lacks it. Every other unread FRD field is informational.
_FRD_REQUIRED_FIELDS = ("file_format", "stage_target.load_strategy", "lobs",
                        "file_name_patterns", "file_patterns")


def _owed_key(document: str, key: str, kind: str = "role", frd_contract=None) -> bool:
    """True when the run cannot complete a feed without an answers.yaml entry
    for ``key`` — a QUESTION / UNRESOLVED line; False = informational, a NOTE
    line (the harness's regex does not match it). STTM: a REQUIRED role. VDD:
    never (its gaps are gate flags). FRD: a field of ``_FRD_REQUIRED_FIELDS``
    the pair-resolved contract still lacks. A gap question (choice / layer /
    text): yes, except the VDD pairing (a run proceeds without a VDD)."""
    from codegen.layout.profile import is_required_role

    if kind != "role":
        return key != "pair.vdd"
    if document == "sttm":
        sheet_layer, _sep, role = key.rpartition("/")
        return is_required_role(sheet_layer.rpartition("/")[2], role)
    if document == "frd" and frd_contract is not None:
        match = re.match(r"^feeds\[(\d+)\]\.(.+)$", key)
        if match is None or match.group(2) not in _FRD_REQUIRED_FIELDS:
            return False
        index = int(match.group(1))
        if index >= len(frd_contract.feeds):
            return False
        from codegen.layout.resolve import _feed_get

        return _feed_get(frd_contract.feeds[index], match.group(2)) in (None, "", [])
    return False


def _frd_answer_hint(key: str, frd_contract=None, gap_fills: list[dict] | None = None) -> str:
    """The remedy an FRD field's answer line names (Chunk B review B2). A
    field the pair-resolved FRD contract already carries (an answer applied,
    another document filled it, the STTM band named it) needs no answer: the
    line says so and where the value came from. A field no answer sets (the
    feed name) names no gaps: key. Else: answer the value under gaps:."""
    from codegen.layout.resolve import _feed_get, answerable_frd_field, unanswerable_frd_reason

    match = re.match(r"^feeds\[(\d+)\]\.(.+)$", key)
    if match is not None and frd_contract is not None \
            and int(match.group(1)) < len(frd_contract.feeds) \
            and _feed_get(frd_contract.feeds[int(match.group(1))],
                          match.group(2)) not in (None, "", []):
        fill = next((f for f in gap_fills or [] if f.get("field") == key), None)
        where = ("" if fill is None else ", from the answers file" if fill["source"] == "user"
                 else f", from the {fill['source']} ({fill['cell']})")
        return f" (frd; the pair-resolved FRD contract carries it{where})"
    if match is not None and not answerable_frd_field(key):
        return f" (frd; {unanswerable_frd_reason(key)})"
    return " (frd; answer the value under gaps: in answers.yaml)"


def _print_elimination_notes(profile) -> None:
    """One NOTE line per band whose layer is source by elimination (the one
    band left of both evidenced target bands) — informational, never owed."""
    from codegen.layout.discover import ELIMINATION_NOTE

    for sheet in profile.mapping_sheets:
        for note in sheet.notes:
            if ELIMINATION_NOTE not in note:
                continue
            ref, _sep, rest = note.partition(" ")
            print(answer_line("NOTE", f"{sheet.name}/{ref}/layer",
                              f"{ref} {rest} — informational, not required to run"))


def _print_file_elimination_notes(flags: list[str]) -> None:
    """One NOTE line per feed whose file was taken by strict elimination
    (layout.resolve.split_frd_feeds_by_sttm), and per blank STTM Schema column
    the FRD's one schema filled — informational, never owed."""
    from codegen.layout.resolve import TARGET_FROM_FRD

    for flag in flags:
        match = (re.match(r"^frd_unstated:(feeds\[\d+\]\.file_name_patterns) (.*\(by "
                          r"elimination: .*)$", flag)
                 # 2026-10-13: a file taken by the sheet's best name match
                 or re.match(r"^frd_unstated:(feeds\[\d+\]\.file_name_patterns) (.*\(best "
                             r"name match: .*)$", flag)
                 # round 4: a blank STTM Schema column filled from the FRD's one schema
                 or re.match(rf"^{TARGET_FROM_FRD}:(feeds\[\d+\]\.\w+_target\.schema) (.*)$",
                             flag)
                 # round 7: file format from the file pattern's extension (the
                 # FRD's format cell names an ingestion type, not a format)
                 or re.match(r"^format_from_extension:(feeds\[\d+\]\.file_format) (.*)$", flag))
        if match:
            print(answer_line("NOTE", match.group(1),
                              f"{match.group(2)} — informational, not required to run"))


def _template_path(args: argparse.Namespace, beside: str | None) -> Path | None:
    """Where the answers template goes: ``--answers-template``, else beside
    the stage's own output file (the run's folder)."""
    from codegen.layout.answers_template import TEMPLATE_NAME

    chosen = getattr(args, "answers_template", None)
    if chosen:
        return Path(chosen)
    return Path(beside).parent / TEMPLATE_NAME if beside else None


def _write_template(target: Path | None, items: list, stage: str, *, merge: bool) -> None:
    """2026-10-13 (D1 pair_11): every open question in one file — key, why,
    candidates, an empty value — filled in the workspace and passed back with
    --answers (codegen.layout.answers_template)."""
    if target is None or not items:
        return
    from codegen.layout.answers_template import write_answers_template

    written = write_answers_template(target, items, stage=stage, merge=merge)
    inputs = sum(1 for i in items if i.section == "inputs")
    print(f"{'ANSWERS FILE':<15} {written} — {len(items) - inputs} question(s), {inputs} "
          "input(s) needed, with their "
          "candidates and an empty value; fill it and rerun with --answers "
          f"{written}")


def _pending_items(pending) -> list:
    from codegen.extract.generic import answers_section
    from codegen.layout.answers_template import TemplateItem

    return [TemplateItem(p.key, f"{p.reason} (feed {p.feed_name}, sheet {p.sheet})",
                         list(p.candidates), answers_section(p.key)) for p in pending]


def _print_needs_answers(pending) -> None:
    """One line per answer a held-back feed needs: QUESTION for a `gaps:` key,
    UNRESOLVED for a band whose layer is open (Chunk A; under `answers:`)."""
    from codegen.extract.generic import answers_section

    if pending:
        feeds = list(dict.fromkeys(p.feed_name for p in pending))
        print(f"{'NEEDS_ANSWERS':<15} {len(pending)} answer(s) for {len(feeds)} held-back "
              f"feed(s): {', '.join(feeds)} — the QUESTION / UNRESOLVED lines below")
    for p in pending:
        section = answers_section(p.key)
        label = "UNRESOLVED" if section == "answers" else "QUESTION"
        print(answer_line(label, p.key, f"{p.reason} (feed {p.feed_name}, sheet {p.sheet}; "
                                        f"answer under {section}: in answers.yaml)"))


def _run_pairs(
    pairs: list[tuple[Path, Path]],
    config: Config,
    *,
    only_feed: str | None,
    dry_run: bool,
    skip_tests: bool,
    output_mode: str | list[str] | None = None,
    vdd_path: Path | None = None,
    conventions_profile: str | None = None,
    iig_template: str | None = None,
    playbook_template: str | None = None,
    answers_path: Path | None = None,
    template_path: Path | None = None,
    command: str = "generate",
) -> int:
    from codegen import owed_inputs as owed_mod
    from codegen import run_stages
    from codegen.resolve.resolver import pending_answers

    inputs = _answers_inputs(answers_path)
    gaps = _answers_gaps(answers_path)
    owed_items: list = []
    from codegen.environments import EnvironmentSelectionError

    try:
        config, run_items, env_questions = _run_environment(config, inputs, output_mode,
                                                            iig_template, dry_run=dry_run)
    except EnvironmentSelectionError as exc:
        _fail(command, f"environment — {exc}", "environment", Coord("environment"), exc=exc)
        return 1
    owed_items += run_items
    failed = False
    matched_feed = False
    # 2026-10-15: a required environment value unset is a QUESTION (exit 3)
    needs_answers = env_questions
    # D1 post51 (2026-10-15): CODEGEN_STAGES.json records the generate outcome.
    # CRASHED (an exception mid-write): the feed's partial emit files are deleted,
    # the emit stages SKIPPED. GATE_FAIL (complete, the gate said no): the files
    # are kept under <feed>/quarantine/ + GATE_FAIL.md, the emit stages PASS.
    # FAIL (nothing written: a mismatch, no --feed match): emit stages SKIPPED.
    out_root = Path(config.output.dir)
    kept: dict[str, list[Path]] = {stage: [] for stage in run_stages.EMIT_STAGES}
    feed_entries: list[dict] = []
    crashed = gate_failed = hard_failed = False
    quarantined: list[str] = []
    for pair_no, (frd_path, sttm_path) in enumerate(pairs, 1):
        try:
            specs = resolve_pair(frd_path, sttm_path, config, vdd_path=vdd_path, gaps=gaps)
            pending = pending_answers(sttm_path)
        except (ContractMismatchError, ValueError) as exc:
            _fail(command, f"{frd_path.name} + {sttm_path.name} — {exc}",
                  "contract.mismatch" if isinstance(exc, ContractMismatchError)
                  else "contract.invalid", Coord("pair", n=pair_no, of=len(pairs)), exc=exc)
            failed = hard_failed = True
            continue
        # Every feed WITH a contract is generated; a held-back feed's answers are
        # printed. Exit 3 only when a feed asked for (--feed) is one of them.
        asked = [p for p in pending if only_feed is not None and (
            config.feed_aliases.get(p.feed_name) or normalize_feed_name(p.feed_name)) == only_feed]
        _print_needs_answers(asked if only_feed is not None else pending)
        if asked:
            needs_answers = matched_feed = True
        for feed_no, spec in enumerate(specs, 1):
            if only_feed is not None and spec.feed_id != only_feed:
                continue
            matched_feed = True
            index = owed_mod.frd_feed_index(frd_path, spec.feed_id, config)
            faq_answers = owed_mod.answers_for(inputs, index, spec.feed_id)
            feed_dir = out_root / spec.feed_slug
            run_stages.clear_quarantine(feed_dir)
            before = run_stages.snapshot(feed_dir)
            where = Coord("feed", n=feed_no, of=len(specs))
            try:
                gate = _generate_feed(spec, config, dry_run=dry_run,
                                      skip_tests=skip_tests, output_mode=output_mode,
                                      conventions_profile=conventions_profile,
                                      iig_template=iig_template,
                                      playbook_template=playbook_template,
                                      faq_answers=faq_answers)
                owed = owed_mod.owed_inputs(
                    spec, config, index=index, output_mode=output_mode,
                    iig_template=iig_template,
                    faq=with_faq_answers(faq_for_spec(spec, config), faq_answers))
                for item in owed:
                    print(owed_mod.input_line(item))
                owed_items += owed_mod.template_items(owed)
                # 2026-10-14 (D1 pair_12): files the LOB lookup keyed to no single
                # LOB — ONE blocking question for the feed (its LOB cells are open)
                lob_item = _lob_question(spec)
                if lob_item is not None:
                    owed_items.append(lob_item)
                    needs_answers = True
            except TemplateGapError as exc:
                _fail(command, f"{spec.feed_id} — template gap: {exc}", "template.gap", where,
                      exc=exc)
                failed = crashed = True
                removed = run_stages.remove_partial(feed_dir, before)
                feed_entries.append(run_stages.feed_entry(
                    out_root, feed_no, len(specs), "CRASHED", {}, removed))
                continue
            except Exception:
                # CRASHED: an uncaught error mid-feed — no partial workbook survives
                # it, the manifest says so, the error propagates (exit 1)
                removed = run_stages.remove_partial(feed_dir, before)
                feed_entries.append(run_stages.feed_entry(
                    out_root, feed_no, len(specs), "CRASHED", {}, removed))
                run_stages.write_stage_manifest(
                    out_root, command=command, exit_code=1,
                    statuses=run_stages.stage_statuses("CRASHED", kept), feeds=feed_entries)
                raise
            if gate.verdict == "FAIL":
                # GATE_FAIL: complete artefacts, the gate said no — kept for review
                failed = gate_failed = True
                report_error(command, "gate.verdict_fail", where, checks=[
                    c.name for c in gate.checks if not c.passed and not c.not_run])
                moved = run_stages.quarantine(feed_dir, before)
                q_dir = feed_dir / run_stages.QUARANTINE_DIR
                q_dir.mkdir(parents=True, exist_ok=True)
                (q_dir / run_stages.GATE_FAIL_FILE).write_text(
                    run_stages.gate_fail_note(gate, feed_no, len(specs), len(moved)),
                    encoding="utf-8", newline="\n")
                written = run_stages.by_stage(moved)
                for stage, paths in written.items():
                    kept[stage] += paths
                quarantined.append(q_dir.relative_to(out_root).as_posix())
                feed_entries.append(run_stages.feed_entry(
                    out_root, feed_no, len(specs), "GATE_FAIL", written,
                    quarantine_dir=q_dir))
                print(f"{'GATE_FAIL':<15} {spec.feed_id} \u2014 {run_stages.REVIEW_ONLY}: "
                      f"{len(moved)} file(s) moved to {q_dir} ({run_stages.GATE_FAIL_FILE})")
                continue
            written = run_stages.written_since(feed_dir, before)
            for stage, paths in written.items():
                kept[stage] += paths
            feed_entries.append(run_stages.feed_entry(
                out_root, feed_no, len(specs), gate.verdict, written))
    if only_feed is not None and not matched_feed:
        _fail(command, f"no resolved feed matches --feed {only_feed!r}", "feed.no_match")
        failed = hard_failed = True
    _write_template(template_path or Path(config.output.dir) / "answers_template.yaml",
                    owed_items, "generate", merge=True)
    # 2026-10-14 (D1 post44): INPUT_NEEDED never blocks — only a held-back feed
    # asked for (a QUESTION / UNRESOLVED line) or a required environment value
    # (2026-10-15) exits 3
    code = 1 if failed else EXIT_NEEDS_ANSWERS if needs_answers else 0
    outcome =("CRASHED" if crashed else "FAIL" if hard_failed else
               "GATE_FAIL" if gate_failed else
               "NEEDS_ANSWERS" if code == EXIT_NEEDS_ANSWERS else "PASS")
    statuses = run_stages.stage_statuses(outcome, kept, quarantined)
    for stage in run_stages.EMIT_STAGES:
        if outcome in ("CRASHED", "FAIL"):
            print(f"{'SKIPPED':<15} {stage} — {run_stages.SKIPPED_UPSTREAM} (exit code "
                  "unchanged; a crashed feed's partial emit artefacts are removed)")
    run_stages.write_stage_manifest(out_root, command=command, exit_code=code,
                                    statuses=statuses, feeds=feed_entries)
    return code


def _run_environment(config: Config, inputs: dict[str, str], output_mode, iig_template,
                     *, dry_run: bool = False):
    """2026-10-14: the run's environment. An answers-file ``inputs:
    run.environment`` selects one when no flag did; a run on the default
    prints ONE INPUT_NEEDED line naming the environment it assumed (and returns
    its answers-template item). Only framework runs whose IIG template has
    environment cells use one.

    2026-10-14 (post47): a custom environment's unset keys may be given as
    ``inputs.run: environment.<key>``; every run-level input (the environment,
    a custom environment's unset keys, the run-scoped owed cells) is printed
    ONCE, in one block, before any feed — and returned as answers-template
    items (``inputs.run``). A named environment left incomplete raises
    EnvironmentSelectionError (the caller prints FAIL).

    2026-10-15 (D1 post50 item 2): ``inputs.run: environment.<key>`` fills any
    environment's key (a named profile's too); the run banner line ``NOTE
    run.environment — environment: <name> (complete, from <origin>; <source>)``
    is printed first; a REQUIRED key still unset is a QUESTION line (the third
    return value True = exit 3)."""
    from codegen import environments as envs
    from codegen import owed_inputs as owed_mod
    from codegen.layout.answers_template import TemplateItem
    from codegen.output_modes import output_parts

    tpl = config.metadata.templates.get(iig_template or config.metadata.template)
    if "framework" in output_parts(output_mode if output_mode else config.output.mode) \
            and tpl is not None and tpl.environment_cells:
        chosen = config.run_environment
        if inputs.get(envs.INPUT_KEY) and (chosen is None or chosen.source == "default"):
            config = envs.with_environment(config, envs.select_environment(
                config, name=inputs[envs.INPUT_KEY], live=_live_run(config, dry_run)))
        env = envs.run_environment(config)
        given = {key[len(envs.INPUT_KEY) + 1:]: value for key, value in inputs.items()
                 if key.startswith(f"{envs.INPUT_KEY}.") and str(value).strip()}
        if env is not None and given:
            config = envs.with_environment(config, env.model_copy(update={
                "entry": envs.overlay_values(env.entry, given),
                "origin": f"{env.origin} + the answers file"}))
    env = envs.run_environment(config) if tpl is not None and tpl.environment_cells else None
    questions = envs.required_items(config) if env is not None else []
    if env is not None:
        print(answer_line("NOTE", envs.INPUT_KEY, f"environment: {envs.banner(config, env)}"))
    for key, reason in questions:
        print(answer_line("QUESTION", key, reason))
    items = [*owed_mod.environment_inputs(config, output_mode=output_mode,
                                          iig_template=iig_template),
             *owed_mod.run_owed_inputs(config, inputs, output_mode=output_mode,
                                       iig_template=iig_template)]
    for line in owed_mod.run_block_lines(items):
        print(line)
    return config, [*(TemplateItem(key, reason, [], owed_mod.INPUTS_SECTION)
                      for key, reason in questions),
                    *(TemplateItem(i.key, i.question,
                                   list(config.environments) if i.key == envs.INPUT_KEY else [],
                                   owed_mod.INPUTS_SECTION) for i in items)], bool(questions)


def _live_run(config: Config, dry_run: bool) -> bool:
    """post47b: the run is live — provider=live and no --dry-run (which asks
    for the mock): a named environment with a key unset is then fatal."""
    from codegen.reasoning.mode import provider_mode

    return provider_mode(config).live and not dry_run


def _lob_question(spec):
    """The feed's ONE LOB question (resolver ``lob_unkeyed``): printed as a
    QUESTION, returned as an answers-template item whose ``files:`` block
    lists every unkeyed file (candidates = the FRD's LOB codes)."""
    if not spec.lob_unkeyed or spec.lob_question_key is None:
        return None
    from codegen.layout.answers_template import TemplateItem
    from codegen.resolve.resolver import lob_question

    reason = lob_question(len(spec.lob_unkeyed), list(spec.lobs))
    print(answer_line("QUESTION", spec.lob_question_key,
                      f"{reason} (feed {spec.feed_id}; their LOB cells are written OPEN — "
                      "answer one LOB per file under gaps: files: in answers.yaml)"))
    return TemplateItem(spec.lob_question_key, f"{reason} (feed {spec.feed_id})",
                        list(spec.lobs), "gaps", files=list(spec.lob_unkeyed))


def _answers_gaps(path: Path | None) -> dict[str, str]:
    """``generate --answers``: the answers file's ``gaps:`` values as text
    (the LOB lookup's ``feeds[<i>].lob_by_file``)."""
    if path is None:
        return {}
    from codegen.layout.answers import load_answers

    return {key: str(value.get("value") or "")
            for key, value in load_answers(path).gaps.items() if isinstance(value, dict)}


def _answers_inputs(path: Path | None) -> dict[str, str]:
    """``generate --answers``: the answers file's ``inputs:`` values (the
    non-blocking owed cells, codegen.owed_inputs)."""
    if path is None:
        return {}
    from codegen.layout.answers import load_answers

    return load_answers(path).inputs


def _extract_sttm(args: argparse.Namespace, config: Config) -> int:
    from codegen.extract import (
        ExtractionError,
        NeedsAnswersError,
        UnresolvedLayoutError,
        WorkbookParseError,
        extract_to_file,
    )

    try:
        workbook_path = Path(args.workbook)
        if not workbook_path.is_file():
            raise FileNotFoundError(f"workbook not found: {workbook_path}")
        layout = None
        width_answers = None
        audit_answers = None
        gaps = None
        if getattr(args, "answers", None):
            # M9.2: feeds[i].fields[<name>].width answers (gaps) ride on the fields.
            from codegen.extract.generic import parse_audit_answers
            from codegen.layout.answers import load_answers
            from codegen.resolve.widths import parse_width_answers

            gaps = load_answers(Path(args.answers)).gaps
            width_answers = parse_width_answers(gaps)
            # Chunk B: feeds[i].audit_columns for a sheet the STTM gives no audit row.
            audit_answers = parse_audit_answers(gaps)
        if getattr(args, "layout", None):
            from codegen.layout.profile import LayoutProfile

            layout = LayoutProfile.model_validate_json(
                Path(args.layout).read_text(encoding="utf-8"))
        elif getattr(args, "answers", None):
            layout = _layout_from_answers(workbook_path, Path(args.answers), config)
        contract = extract_to_file(
            workbook_path,
            _contract_path(args.frd_contract, config),
            Path(args.out),
            config,
            contract_name=args.contract_name,
            generated_date=args.generated_date,
            layout=layout,
            require_complete=bool(getattr(args, "require_complete", False)),
            width_answers=width_answers,
            audit_answers=audit_answers,
            gaps=gaps,
        )
    except NeedsAnswersError as exc:
        # NO feed produced a usable contract: nothing written, exit 3.
        _print_needs_answers(exc.pending)
        print(f"{'NOT EXTRACTED':<15} {args.out} — every feed needs an answer (no contract "
              "written)")
        _write_template(_template_path(args, args.out), _pending_items(exc.pending),
                        "extract-sttm", merge=True)
        return EXIT_NEEDS_ANSWERS
    except UnresolvedLayoutError as exc:
        # --require-complete and the layout leaves roles open: the answers file
        # places them; nothing extracted, exit 3.
        from openpyxl.utils import get_column_letter

        from codegen.layout.answers_template import TemplateItem

        for role in exc.unresolved:
            print(answer_line("UNRESOLVED", f"{role.sheet}/{role.layer}/{role.role}",
                              f"{role.reason} (place it under answers: in answers.yaml)"))
        print(f"{'NOT EXTRACTED':<15} {args.out} — {len(exc.unresolved)} layout role(s) open "
              "(--require-complete)")
        _write_template(_template_path(args, args.out), [
            TemplateItem(f"{r.sheet}/{r.layer}/{r.role}", r.reason,
                         [get_column_letter(c) for c in r.candidates], "answers")
            for r in exc.unresolved], "extract-sttm", merge=True)
        return EXIT_NEEDS_ANSWERS
    except (WorkbookParseError, ExtractionError, FileNotFoundError, ValueError) as exc:
        _fail("extract-sttm", f"extract-sttm — {exc}", "sttm.unreadable", Coord("STTM"),
              exc=exc)
        return 1
    feeds = ", ".join(f"{f.feed_id} ({f.field_count} fields)" for f in contract.feeds)
    if contract.frd_split is not None:
        # 2026-10-13: a one-feed FRD contract split per mapping sheet HERE.
        print(answer_line("NOTE", "frd.split", contract.frd_split.flags[0]))
    print(f"{'EXTRACTED':<15} {args.out} — {len(contract.feeds)} feed(s): {feeds}")
    # Some feeds produced a contract: it is written and the exit is 0; the
    # held-back sheets still print their QUESTION lines (and the template).
    _print_needs_answers(contract.needs_answers)
    _write_template(_template_path(args, args.out), _pending_items(contract.needs_answers),
                    "extract-sttm", merge=True)
    return EXIT_OK


BUILD_STAMP_FILE = "CODEGEN_BUILD.json"


def _write_build_stamp(out_root: Path) -> None:
    """``<output root>/CODEGEN_BUILD.json`` (2026-10-12): which build produced
    the run — version + commit / branch / commit time (or the src tree id) —
    at the run level, outside every feed folder (those stay byte-stable).
    tools/iig_diff.py refuses to score an agent workbook without one."""
    from codegen.build_info import build_identity, build_text, codegen_version

    out_root.mkdir(parents=True, exist_ok=True)
    stamp = {"version": codegen_version(), **build_identity()}
    stamp["line"] = f"codegen {stamp['version']} · {build_text(build_identity())}"
    # Written silently: the console's line order (source line, then the
    # verdict headline) is a contract tests and the harness read.
    (out_root / BUILD_STAMP_FILE).write_text(json.dumps(stamp, indent=2) + "\n",
                                             encoding="utf-8", newline="\n")


def _outputs_through_storage(config: Config):
    """(config, push). Under the default outputs role (``output.dir`` in the
    checkout) nothing changes. With ``storage.outputs`` /
    ``CODEGEN_STORAGE_OUTPUTS`` set, generation writes into the role's local
    working directory and ``push()`` sends the tree up afterwards (Workspace
    API / Files API — never a /Volumes or /Workspace path)."""
    from codegen.storage import StorageError, open_storage

    outputs = open_storage(config, Path(".")).outputs
    if outputs.is_local and outputs.workdir == Path(".") / config.output.dir:
        return config, lambda: None
    scoped = config.model_copy(update={"output": config.output.model_copy(update={
        "dir": str(outputs.workdir), "reports_dir": str(outputs.workdir / "reports")})})

    def push() -> None:
        try:
            sent = outputs.push_tree("")
        except StorageError as exc:
            print(f"{'WARN':<15} outputs not stored: {exc} — they remain in {outputs.workdir}")
            return
        if sent:
            print(f"{'STORED':<15} {len(sent)} file(s) -> {outputs.uri()}")

    return scoped, push


def _answer_notes(notes: list[str], seen: set[str], gaps: bool = True,
                  profile=None) -> list[str]:
    """The answers-file notes worth printing after the last pass: a "matches
    no open question" note is dropped for a key that WAS open in an earlier
    pass (the answer applied there), and for a band layer the final
    ``profile`` shows applied (a band of that sheet carries evidence
    "answer" - e.g. read back from the runtime cache the layout stage
    wrote); ``gaps=False`` drops every gaps: note (another stage consumes
    them)."""
    from codegen.layout.profile import BAND_LAYER_KEY_RE

    answered_sheets = {s.name for s in (profile.sheets if profile is not None else [])
                       if any(b.layer_evidence == "answer" for b in s.bands)}
    kept = []
    for note in notes:
        if not gaps and note.startswith("gaps["):
            continue
        if "matches no open question" in note and any(
                f"['{k}']" in note or f"({k})" in note for k in seen):
            continue
        key = note[note.find("(") + 1:note.find(")")] if "(" in note else ""
        band = BAND_LAYER_KEY_RE.match(key)
        if band is not None and band.group("sheet") in answered_sheets:
            continue
        kept.append(note)
    return kept


# Passes an answers file is applied over (Chunk B): a band layer, then its
# columns, then the questions those open (the file for a table) - each pass
# can open the next question only once.
ANSWER_ROUNDS = 4


def _layout_from_answers(workbook: Path, answers_path: Path, config: Config):
    """The workbook's layout profile with the answers file applied — cache and
    synonyms first, then the file's placements (source=user); never a model."""
    from codegen.layout.answers import apply_answers, load_answers
    from codegen.layout.resolve import resolve_workbook
    from codegen.storage import runtime_layout_cache

    runtime_cache, push_cache = runtime_layout_cache(config, Path("."))
    doc, _wb = resolve_workbook(workbook, config, provider=None,
                                runtime_cache_dir=runtime_cache)
    # M9.1: an answer may set any role, open or not — the file is applied even
    # when nothing is open (it then overrides a synonym / cached placement).
    # Chunk B: re-applied while an answer opens the next question (a band's
    # layer, then its columns).
    answers_file = load_answers(answers_path)
    placed: dict = {}
    notes: list[str] = []
    seen: set[str] = set()
    for _round in range(ANSWER_ROUNDS):
        seen |= {q.key for q in doc.questions}
        found, notes = apply_answers(answers_file, doc.questions, {"sttm": workbook.name},
                                     documents={"sttm": (doc.profile, _wb)})
        merged = {**placed, **found["sttm"]}
        if merged == placed:
            break
        placed = merged
        doc, _wb = resolve_workbook(workbook, config, provider=None,
                                    runtime_cache_dir=runtime_cache, answers=placed, prior=doc)
    # gaps: answers are the layout stage's (applied to the FRD contract it wrote).
    for note in _answer_notes(notes, seen, gaps=False, profile=doc.profile):
        print(f"{'NOTE':<15} {note}")
    if placed:
        print(f"{'ANSWERS':<15} {len(placed)} answer(s) applied from "
              f"{answers_path} (source=user)")
        push_cache()
    for question in doc.questions:
        owed = _owed_key("sttm", question.key)
        print(answer_line("UNRESOLVED" if owed else "NOTE", question.key, question.reason
                          + ("" if owed else " — optional role, not required to run")))
    _print_elimination_notes(doc.profile)
    return doc.profile


def _pair(args: argparse.Namespace, config: Config) -> int:
    """Pair an STTM with its FRD / VDD among a folder's documents BY CONTENT
    (codegen.pairing). An undecided pairing prints its candidates; settle it
    with `pairing:` in the answers file."""
    from codegen.layout.answers import AnswersFileError, load_answers
    from codegen.pairing import (
        OVERRIDE_FILE_NAMES,
        PairOverrides,
        PairOverridesError,
        classify_folder,
        configured_overrides_path,
        load_pair_overrides,
        pair_by_content,
        vdd_by_content_note,
    )

    sttm = Path(args.sttm)
    if not sttm.is_file():
        _fail("pair", f"pair — workbook not found: {sttm}", "workbook.not_found",
              Coord("STTM"))
        return 1
    folders = [Path(d) for d in (args.candidates or [str(sttm.parent)])]
    everything = [p for d in folders if d.is_dir() for p in sorted(d.iterdir()) if p.is_file()]
    files = [p for p in everything if not p.name.startswith("~$") and p != sttm]
    # Every file with its role and the rule that assigned it (2026-10-13): a
    # role token ANYWHERE in the name ("<title>_FRD.docx" is an FRD); an IIG is
    # never an input; a workbook without a role token is a dictionary only on
    # positive content evidence. Nothing is skipped silently.
    roles = classify_folder(files, config)
    # EVERY file of the folder (2026-10-13): the STTM being paired and the
    # Office lock files too — a D1 report's role table had no STTM row.
    print(f"{'ROLE':<15} {sttm.name} — sttm: the STTM being paired (--sttm)")
    for role in roles:
        print(f"{'ROLE':<15} {role.name} — {role.role}: {role.rule}")
    for lock in (p for p in everything if p.name.startswith("~$")):
        print(f"{'ROLE':<15} {lock.name} — ignored: an Office lock file (~$), not a document")
    frds = {r.name: r.path for r in roles if r.role == "frd"}
    vdds = {r.name: r.path for r in roles if r.role == "vdd"}
    try:
        overrides = PairOverrides()
        configured = configured_overrides_path(config, Path("."))
        found = [Path(args.pairs_file)] if args.pairs_file else (
            [configured] if configured is not None else
            [p for p in files if p.name.lower() in OVERRIDE_FILE_NAMES][:1])
        for path in found:
            overrides = overrides.merged(load_pair_overrides(path))
            print(f"{'OVERRIDES':<15} {path.name} — {len(overrides.groups)} pair(s) listed")
    except PairOverridesError as exc:
        _fail("pair", f"pair — pairs override file: {exc}", "pairs.file", Coord("pairs"),
              exc=exc)
        return 1
    other_sttms = [r.name for r in roles if r.role == "sttm"]
    if other_sttms and not (overrides.lookup(sttm.name, "frd")
                            or overrides.lookup(sttm.name, "vdd")):
        # More than one STTM and no override pins this one: the choice is said.
        print(answer_line("NOTE", "pair.sttm",
                          f"the folder holds {len(other_sttms) + 1} STTM workbooks "
                          f"({', '.join([sttm.name, *other_sttms])}); pairing for "
                          f"{sttm.name} (--sttm) — pin each STTM's companions with a pairs "
                          f"override file ({OVERRIDE_FILE_NAMES[0]})"))
    try:
        chosen = (load_answers(Path(args.answers)).pairing.get(sttm.name, {})
                  if args.answers else {})
    except (AnswersFileError, OSError) as exc:
        _fail("pair", f"pair — answers file: {exc}", "answers.file", Coord("answers"),
              exc=exc)
        return 1
    failed = owed = False
    by_name = {r.name: r for r in roles}
    for kind, candidates, explicit in (("frd", frds, config.demo.pairing_map),
                                       ("vdd", vdds, config.demo.vdd_pairing_map)):
        if chosen.get(kind):
            known = chosen[kind] in candidates
            print(f"{'PAIRED' if known else 'FAIL':<15} {kind} {chosen[kind]} — answers file"
                  + ("" if known else " names a document that is not among the candidates"))
            if not known:
                report_error("pair", "answers.file", Coord("answers"))
            failed = failed or not known
            continue
        # An override may name a file the role gate left out — the person's call.
        wanted = overrides.lookup(sttm.name, kind)
        if wanted in by_name and wanted not in candidates and by_name[wanted].path:
            candidates = {**candidates, wanted: by_name[wanted].path}
        decision = pair_by_content(kind, sttm, candidates, config, Path("."),
                                   explicit_map=explicit, overrides=overrides)
        if decision.chosen:
            print(f"{'PAIRED':<15} {kind} {decision.chosen} — rule {decision.rule}: "
                  f"{decision.reason}")
            if decision.rule == "name_fuzzy":
                print(answer_line("NOTE", f"pair.{kind}", decision.reason))
            note = vdd_by_content_note(decision)
            if note:
                print(answer_line("NOTE", f"pair.{kind}", note))
        elif decision.ambiguous:
            # An undecided FRD is an answer owed (exit 3, NEEDS_ANSWERS); an
            # undecided VDD is not (the run proceeds without one) — a NOTE the
            # harness does not record.
            is_owed = _owed_key("pair", f"pair.{kind}", "choice")
            owed = owed or is_owed
            print(answer_line("QUESTION" if is_owed else "NOTE", f"pair.{kind}", decision.reason))
            for candidate in decision.candidates:
                print(f"{'CANDIDATE':<15} {candidate.name} score {candidate.score:g} — "
                      f"{candidate.summary()}")
            print(f"{'REMEDY':<15} answers file: pairing: {{\"{sttm.name}\": "
                  f"{{{kind}: <candidate>}}}} — or a pairs override file "
                  f"({OVERRIDE_FILE_NAMES[0]} next to the STTM, or --pairs-file)")
        elif kind == "vdd" and not candidates:
            left_out = [f"{r.name} ({r.role})" for r in roles
                        if r.role in ("iig", "unclassified") and r.name.lower().endswith(".xlsx")]
            print(f"{'NONE':<15} vdd — no VDD: no workbook carries a 'vdd' token or reads as a "
                  "dictionary by content" + (f"; not candidates: {', '.join(left_out)}"
                                             if left_out else ""))
        else:
            print(f"{'NONE':<15} {kind} — {decision.reason}")
    return EXIT_FAILED if failed else EXIT_NEEDS_ANSWERS if owed else EXIT_OK


def load_workbook_for_answers(path: Path, stop_after: int):
    from codegen.layout.extent import load_document

    return load_document(path, stop_after)


def _unreadable_reason(exc: Exception, workbook: Path, frd: Path | None) -> str:
    """One line for a document no reader can open: which document, why, what
    to do (the structural message only - sheet names and header texts)."""
    from codegen.extract.frd_docx import FrdDocxError

    message = " ".join(str(exc).split())
    if isinstance(exc, FrdDocxError):
        name = frd.name if frd is not None else "the FRD"
        return (f"FRD {name} cannot be read: {message} - supply the FRD in a documented "
                "family (F1 metadata section tables / F2 Solution Requirement tables) or "
                "its contract JSON")
    return (f"STTM {workbook.name} cannot be read: {message} - no sheet carries a mapping "
            "header the synonyms place; add the client's header texts to "
            "extractor.discovery in a config overlay, or run the layout model "
            "(layout.provider: live)")


def _unreadable_where(exc: Exception) -> tuple[str, Coord]:
    """(reason, coordinates) of a document no reader can open."""
    from codegen.extract.frd_docx import FrdDocxError

    if isinstance(exc, FrdDocxError):
        return "frd.unreadable", Coord("FRD")
    return "sttm.unreadable", Coord("STTM")


def _layout(args: argparse.Namespace, config: Config) -> int:
    """Resolve and print the layout of a workbook (and optionally its FRD
    document): source per role, confidences, unresolved roles, the
    provider call count and — for a pair — the cross-document checks."""
    from codegen.layout.model import build_layout_provider
    from codegen.layout.resolve import resolve_pair

    workbook = Path(args.workbook)
    if not workbook.is_file():
        _fail("layout", f"layout — workbook not found: {workbook}", "workbook.not_found",
              Coord("STTM"))
        return 1
    from codegen.layout.answers import (
        AnswersFileError,
        apply_answers,
        load_answers,
        unresolved_report,
    )
    from codegen.storage import StorageError, runtime_layout_cache

    frd = Path(args.frd) if args.frd else None
    vdd = Path(args.vdd) if args.vdd else None
    provider = build_layout_provider(config, dry_run=args.dry_run)
    try:
        # M8.3: the runtime profile cache lives in the state role
        # (storage.state) — the checkout by default, a workspace folder / a
        # volume when configured; never under fixtures/.
        runtime_cache, push_cache = runtime_layout_cache(config, Path("."))
    except StorageError as exc:
        _fail("layout", f"layout — state storage: {exc}", "storage", Coord("storage"),
              exc=exc)
        return 1

    def resolve(answers: dict | None, refresh: bool = False, prior=None):
        return resolve_pair(workbook, frd, config, vdd_path=vdd, provider=provider,
                            runtime_cache_dir=runtime_cache, use_cache=not args.no_cache,
                            answers=answers, refresh=refresh, prior=prior)

    from codegen.extract.frd_docx import FrdDocxError
    from codegen.extract.workbook import WorkbookParseError
    from codegen.layout.discover import NoLayoutError

    unreadable = (FrdDocxError, NoLayoutError, WorkbookParseError)
    try:
        result = resolve(None, refresh=args.refresh)
    except unreadable as exc:
        # Chunk B (cold drill): a document no reader can open is a FAIL that
        # names the document and the remedy, never a traceback.
        _fail("layout", f"layout — {_unreadable_reason(exc, workbook, frd)}",
              *_unreadable_where(exc), exc=exc)
        return 1
    names = {"sttm": workbook.name, "frd": frd.name if frd else "", "vdd": vdd.name if vdd else ""}
    if args.refresh:
        print(f"{'REFRESH':<15} caches bypassed; runtime entries overwritten "
              f"({result.provider_calls} provider call(s))")
    if args.answers:
        # M9.1: applied even when nothing is open — an answer may override a
        # synonym / model / cached placement.
        try:
            stop_after = config.extractor.used_range_empty_rows
            answers_file = load_answers(Path(args.answers))
            sttm_book = load_workbook_for_answers(workbook, stop_after)
            vdd_book = (load_workbook_for_answers(vdd, stop_after)
                        if vdd is not None and result.vdd is not None else None)
        except (AnswersFileError, OSError) as exc:
            _fail("layout", f"layout — answers file: {exc}", "answers.file", Coord("answers"),
                  exc=exc)
            return 1
        answers: dict = {"sttm": {}, "frd": {}, "vdd": {}, "gaps": {}}
        notes: list[str] = []
        seen: set[str] = set()
        # Chunk B (cold drill): one answer can open the next question - a
        # band's layer, then the columns it has no synonym for; a table's
        # columns, then which file feeds it. The file is applied again to
        # every pass's questions until nothing new applies (earlier answers
        # kept), so ONE answers file settles the whole chain. Each pass
        # CONTINUES from the previous one (M9.1b): no cache read, no second
        # model call.
        for _round in range(ANSWER_ROUNDS):
            documents = {"sttm": (result.sttm.profile, sttm_book)}
            if vdd_book is not None and result.vdd is not None:
                documents["vdd"] = (result.vdd.profile, vdd_book)
            seen |= {q.key for q in result.questions}
            try:
                found, notes = apply_answers(answers_file, result.questions, names,
                                             documents=documents)
            except AnswersFileError as exc:
                _fail("layout", f"layout — answers file: {exc}", "answers.file",
                      Coord("answers"), exc=exc)
                return 1
            merged = {k: {**answers.get(k, {}), **found.get(k, {})} for k in found}
            if merged == answers:
                break
            answers = merged
            try:
                result = resolve(answers, refresh=args.refresh, prior=result)
            except unreadable as exc:
                _fail("layout", f"layout — {_unreadable_reason(exc, workbook, frd)}",
                      *_unreadable_where(exc), exc=exc)
                return 1
        for note in _answer_notes(notes, seen, profile=result.sttm.profile):
            print(f"{'NOTE':<15} {note}")
        placed = sum(len(answers[k]) for k in ("sttm", "vdd", "gaps"))
        if placed:
            print(f"{'ANSWERS':<15} {placed} answer(s) applied from {args.answers} (source=user)")
    try:
        push_cache()
    except StorageError as exc:
        print(f"{'WARN':<15} layout — profile cache not stored: {exc}")
    if args.report_unresolved:
        target = Path(args.report_unresolved)
        layout_report = result.report()
        rejected = [line for document in ("sttm", "frd", "vdd")
                    for line in ((layout_report.get(document) or {}).get("rejections") or [])]
        target.write_text(unresolved_report(result.questions, names, rejected),
                          encoding="utf-8", newline="\n")
        print(f"{'REPORT':<15} {target} — {len(result.questions)} unresolved item(s), "
              "structural labels only")
    report = result.report()
    printed: set[str] = set()                # UNRESOLVED keys already on stdout
    if args.json:
        print(json.dumps(report, indent=2, ensure_ascii=False))
    else:
        for document in ("sttm", "frd"):
            doc = report[document]
            if doc is None:
                continue
            print(f"{document.upper():<15} source={doc['source']} cache_hit={doc['cache_hit']} "
                  f"provider_calls={doc['provider_calls']} roles_by_source="
                  f"{doc['roles_by_source']}")
            for item in doc["rejections"]:
                print(f"{'REJECTED':<15} {item}")
            for item in doc["unresolved"]:
                key, _sep, reason = item.partition(": ")
                owed = _owed_key(document, key, "role", result.frd_contract)
                # Chunk B: an FRD field is answered with its VALUE under gaps:
                # (a QUESTION); only an STTM column placement is UNRESOLVED.
                label = "NOTE" if not owed else "QUESTION" if document == "frd" else "UNRESOLVED"
                where = (_frd_answer_hint(key, result.frd_contract, result.gap_fills)
                         if document == "frd" else "")
                print(answer_line(label, key, (reason or item) + where + (
                    "" if owed else " — informational, not required to run")))
                printed.add(key)
            if document == "sttm":
                _print_elimination_notes(result.sttm.profile)
            else:
                _print_file_elimination_notes(result.flags)
        profile = result.sttm.profile
        for key in sorted(profile.confidence):
            print(f"{'ROLE':<15} {key} conf={profile.confidence[key]:.2f} "
                  f"source={profile.role_sources.get(key, profile.source)}")
        for check in report["cross_checks"]:
            print(f"{'CROSSCHECK':<15} {check}")
        from codegen.reasoning.usage import layout_usage

        print(f"{'MODEL USAGE':<15} layout: {layout_usage(provider).label}")
    if args.profile_out:
        Path(args.profile_out).write_text(
            result.sttm.profile.model_dump_json(indent=2) + "\n", encoding="utf-8", newline="\n")
    if args.frd_contract_out:
        if result.frd_contract is None:
            _fail("layout", "layout — --frd-contract-out needs --frd", "usage")
            return 1
        from codegen.extract.frd_docx import contract_to_json as frd_contract_to_json

        target = Path(args.frd_contract_out)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(frd_contract_to_json(result.frd_contract), encoding="utf-8",
                          newline="\n")
        print(f"{'FRD CONTRACT':<15} {target} — feeds "
              f"{[f.feed_name for f in result.frd_contract.feeds]} "
              f"({len(result.gap_fills)} value(s) taken from the other documents)")
    if args.require_complete and result.questions:
        # Unplaced columns / open questions are answers OWED, not a failure:
        # exit 3 — layout is the first stage, so this is the only way a pair
        # with unplaced columns reads NEEDS_ANSWERS instead of FAILED. One
        # UNRESOLVED line per unplaced column (a role question, placed under
        # `answers:`), one QUESTION line per other question (`gaps:`).
        owed = [q for q in result.questions
                if _owed_key(q.document, q.key, q.kind, result.frd_contract)]
        for question in result.questions:
            if question.key in printed:
                continue
            printed.add(question.key)
            if question not in owed:
                print(answer_line("NOTE", question.key,
                                  f"{question.reason} ({question.document}) — informational, "
                                  "not required to run"))
            elif question.kind == "role" and question.document == "frd":
                print(answer_line("QUESTION", question.key,
                                  f"{question.reason} (frd; answer the value under gaps: in "
                                  "answers.yaml)"))
            elif question.kind == "role":
                print(answer_line("UNRESOLVED", question.key,
                                  f"{question.reason} ({question.document}; place it "
                                  "under answers: in answers.yaml)"))
            else:
                print(answer_line("QUESTION", question.key,
                                  f"{question.reason} ({question.document}); candidates "
                                  f"{question.candidates}"))
        if owed:
            print(f"{'NEEDS_ANSWERS':<15} layout — {len(owed)} answer(s) owed "
                  "(--require-complete); nothing is wrong with the documents")
            _write_template(_template_path(args, args.frd_contract_out or args.profile_out),
                            _question_items(owed), "layout", merge=False)
            return EXIT_NEEDS_ANSWERS
    return 0


def _question_items(questions) -> list:
    """Layout questions as answers-template items: a column placement under
    answers: (its candidate headers), every other question under gaps: (its
    candidate values)."""
    from codegen.layout.answers_template import TemplateItem

    items = []
    for q in questions:
        if q.kind == "role" and q.document != "frd":
            candidates = [str(c.get("header") or c.get("col")) for c in q.candidates]
            items.append(TemplateItem(q.key, q.reason, candidates, "answers"))
        else:
            candidates = [str(c.get("value") if "value" in c else c.get("layer"))
                          for c in q.candidates]
            items.append(TemplateItem(q.key, q.reason, candidates, "gaps"))
    return items


def _extract_vdd(args: argparse.Namespace, config: Config) -> int:
    from codegen.extract.vdd import VddExtractionError, contract_to_json, extract_vdd_contract
    from codegen.layout.discover import NoLayoutError

    try:
        vdd_path = Path(args.vdd)
        if not vdd_path.is_file():
            raise FileNotFoundError(f"VDD workbook not found: {vdd_path}")
        tables: list[str] | None = None
        if args.sttm_contract:
            from codegen.contracts.sttm import SttmContract

            sttm = SttmContract.model_validate_json(
                Path(args.sttm_contract).read_text(encoding="utf-8"))
            tables = sorted({t for f in sttm.feeds
                             for t in (f.stage.table, f.source_table, f.feed_id) if t})
        layout = None
        if args.layout:
            from codegen.layout.profile import LayoutProfile

            layout = LayoutProfile.model_validate_json(
                Path(args.layout).read_text(encoding="utf-8"))
        contract, profile = extract_vdd_contract(
            vdd_path, config, sttm_tables=tables, layout=layout,
            generated_date=args.generated_date, contract_name=args.contract_name)
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(contract_to_json(contract), encoding="utf-8", newline="\n")
        if args.profile:
            Path(args.profile).write_text(profile.model_dump_json(indent=2) + "\n",
                                          encoding="utf-8", newline="\n")
    except (VddExtractionError, NoLayoutError, FileNotFoundError, ValueError) as exc:
        _fail("extract-vdd", f"extract-vdd — {exc}", "vdd.unreadable", Coord("VDD"), exc=exc)
        return 1
    print(f"{'EXTRACTED':<15} {args.out} — {len(contract.files)} file row(s), "
          f"{len(contract.fields)} field(s) on {contract.field_sheets}, "
          f"{len(contract.position_rows)} position row(s); layout source {profile.source}, "
          f"{len(profile.unresolved)} unresolved role(s)")
    return 0


def _extract_frd(args: argparse.Namespace, config: Config) -> int:
    from codegen.extract.frd_docx import FrdDocxError, contract_to_json, extract_frd_contract

    try:
        docx_path = Path(args.docx)
        if not docx_path.is_file():
            raise FileNotFoundError(f"FRD document not found: {docx_path}")
        contract, profile = extract_frd_contract(
            docx_path, config, contract_name=args.contract_name,
            generated_date=args.generated_date,
        )
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(contract_to_json(contract), encoding="utf-8", newline="\n")
        if args.profile:
            profile_path = Path(args.profile)
            profile_path.parent.mkdir(parents=True, exist_ok=True)
            profile_path.write_text(profile.model_dump_json(indent=2) + "\n",
                                    encoding="utf-8", newline="\n")
    except (FrdDocxError, FileNotFoundError, ValueError) as exc:
        _fail("extract-frd", f"extract-frd — {exc}", "frd.docx", Coord("FRD"), exc=exc)
        return 1
    feeds = ", ".join(f.feed_name for f in contract.feeds)
    print(f"{'EXTRACTED':<15} {args.out} — family {profile.family}, layout source "
          f"{profile.source}, {len(contract.feeds)} feed(s): {feeds}; "
          f"{len(profile.unresolved)} unresolved field(s); status {contract.status}")
    # On its own, extract-frd cannot tell an answer owed from a field the STTM
    # or the VDD supplies (the layout stage's gap chain decides): every unread
    # FRD field is informational here — NOTE, never UNRESOLVED.
    for item in profile.unresolved:
        print(answer_line("NOTE", item.field, f"{item.reason} — FRD field not read; the "
                                              "layout stage fills it or asks"))
    return 0


def _sharepoint_fetch(args: argparse.Namespace, config: Config) -> int:
    """Pull the library's workbooks/contracts into a local input directory.

    The read-side edge of the transport seam (codegen.sharepoint's module
    docstring has the full rationale). Deliberately a separate command rather
    than a flag on `extract-sttm`: keeping the network at the edge is what
    lets the generator stay offline and credential-free, and it means a
    re-run after a parse fix does not re-download anything.

    Fetching nothing is an error, not a no-op — an empty library folder and a
    successful fetch must not look the same to whatever runs next.
    """
    from codegen.sharepoint import (
        SUPPORTED_SUFFIXES,
        GraphError,
        SharePointConfigError,
        build_client,
        config_for,
    )

    try:
        cfg = config_for(config.sharepoint)
        client = build_client(cfg)
        print(f"site:    {cfg.host}{cfg.site_path}")
        print(f"library: {cfg.library}  folder: {cfg.input_folder or '<root>'}")
        print(f"dest:    {args.dest}")
        fetched = client.fetch_to_dir(args.dest, suffixes=SUPPORTED_SUFFIXES)
    except (SharePointConfigError, GraphError) as exc:
        _fail("sharepoint-fetch", f"sharepoint-fetch — {exc}", "storage", exc=exc)
        return 1

    for item in fetched:
        print(f"{'FETCHED':<15} {item.name} — {item.size:,} bytes, modified {item.modified}")
    if not fetched:
        print(
            f"{'FAIL':<15} sharepoint-fetch — no supported documents in "
            f"{cfg.library}/{cfg.input_folder or '<root>'} "
            f"(supported: {sorted(SUPPORTED_SUFFIXES)}). Upload an STTM workbook or an "
            f"FRD contract, or check sharepoint.input_folder in config/config.yaml."
        )
        report_error("sharepoint-fetch", "input.not_found", Coord("storage"))
        return 1
    print(f"\n{len(fetched)} file(s) -> {args.dest}")
    return 0


def _publishable(args: argparse.Namespace, config: Config) -> list[Path]:
    """The artifact list for one publish call, with containment enforced.

    `--path` addresses one file relative to the feed's output directory and
    is resolved-and-checked the same way the UI's file reader does: a library
    write must never be steerable outside `out/<feed_slug>/` by a `..`.
    With no `--path`, the deliverables are the feed's generation report and
    its assembled notebook — the two artifacts a reviewer hands over.
    """
    feed_dir = (Path(config.output.dir) / args.feed).resolve()
    if args.path:
        target = (feed_dir / args.path).resolve()
        if not target.is_relative_to(feed_dir):
            raise ValueError(f"path escapes the feed directory: {args.path}")
        return [target]
    return [
        Path(config.output.reports_dir).resolve() / f"{args.feed}.md",
        feed_dir / f"{args.feed}.ipynb",
    ]


def _sharepoint_publish(args: argparse.Namespace, config: Config) -> int:
    """Publish one feed's generated artifacts to the library output folder.

    The write-side edge. Publishing is deliberately NOT folded into
    `generate`: generation re-runs every time a rule or a contract changes,
    and a re-generate is not automatically a re-publish — the human gate sits
    between them. Running this command IS that gate, which is why it takes no
    `--confirm` flag; the UI endpoint, which a stray POST could reach, does
    require one.

    Publishing nothing is an error for the same reason a silent no-op is
    wrong upstream: downstream it is indistinguishable from success.
    """
    from codegen.sharepoint import (
        GraphError,
        SharePointConfigError,
        build_client,
        config_for,
        published_name,
    )

    try:
        artifacts = _publishable(args, config)
    except ValueError as exc:
        _fail("sharepoint-publish", f"sharepoint-publish — {exc}", "publish", exc=exc)
        return 1

    missing = [p for p in artifacts if not p.is_file()]
    if missing:
        print(
            f"{'FAIL':<15} sharepoint-publish — nothing to publish for {args.feed!r}: "
            + ", ".join(str(p) for p in missing)
            + ". Run `codegen generate` (or generate-all) first."
        )
        report_error("sharepoint-publish", "publish", Coord("output"))
        return 1

    try:
        cfg = config_for(config.sharepoint)
        client = build_client(cfg)
        print(f"site:      {cfg.host}{cfg.site_path}")
        print(f"library:   {cfg.library}  folder: {cfg.output_folder or '<root>'}")
        for path in artifacts:
            name = published_name(args.feed, path.name)
            result = client.upload_file(path, name=name)
            print(f"{'PUBLISHED':<15} {name} — {path.stat().st_size:,} bytes "
                  f"-> {result.get('webUrl', '<no url>')}")
    except (SharePointConfigError, GraphError) as exc:
        _fail("sharepoint-publish", f"sharepoint-publish — {exc}", "publish", exc=exc)
        return 1

    print(f"\n{len(artifacts)} artifact(s) published to "
          f"{cfg.library}/{cfg.output_folder or '<root>'}.")
    return 0


def _databricks_publish(args: argparse.Namespace, config: Config) -> int:
    """Publish one feed's generated artifacts to a Unity Catalog volume.

    The UC twin of `sharepoint-publish` — same doctrine: running this command
    IS the human gate (no `--confirm`; the UI endpoint requires one), and
    publishing nothing is an error, never a silent no-op. The target may be
    overridden per call, but `codegen.databricks` refuses anything outside
    WRITABLE_PREFIX in code.
    """
    from codegen.databricks import (
        DatabricksConfigError,
        DatabricksTransportError,
        config_for,
        ensure_volume,
        publish_artifacts,
    )

    try:
        artifacts = _publishable(args, config)
    except ValueError as exc:
        _fail("databricks-publish", f"databricks-publish — {exc}", "publish", exc=exc)
        return 1
    missing = [p for p in artifacts if not p.is_file()]
    if missing:
        print(
            f"{'FAIL':<15} databricks-publish — nothing to publish for {args.feed!r}: "
            + ", ".join(str(p) for p in missing)
            + ". Run `codegen generate` (or generate-all) first."
        )
        report_error("databricks-publish", "publish", Coord("output"))
        return 1

    try:
        cfg = config_for(config.databricks)
        target = ensure_volume(cfg, volume=args.volume or cfg.output_volume,
                               knob="output_volume", catalog=args.catalog,
                               schema=args.schema)
        print(f"volume:    {target['full_name']}"
              + ("  (created)" if target["created"] else ""))
        published = publish_artifacts(
            cfg, args.feed, artifacts, catalog=args.catalog,
            schema=args.schema, volume=args.volume, force=args.force,
        )
        for item in published:
            print(f"{'PUBLISHED':<15} {item['name']} — "
                  f"{item['size_bytes']:,} bytes -> {item['path']}")
    except (DatabricksConfigError, DatabricksTransportError) as exc:
        _fail("databricks-publish", f"databricks-publish — {exc}", "publish", exc=exc)
        return 1

    print(f"\n{len(artifacts)} artifact(s) published to /Volumes/"
          f"{target['full_name'].replace('.', '/')}/{args.feed}/.")
    return 0


# Filename date-token table for the landing seeder. Tokens resolve only at
# word boundaries (so OH / MIDS stay literal); HHMM/HH use a fixed synthetic
# delivery time of 06:00. Anything date-ish left after resolution is flagged
# and kept literal — the seeder never guesses.
_SEED_TOKEN_TABLE: tuple[tuple[str, str], ...] = (
    ("CCYYMMDD", "%Y%m%d"),
    ("YYYYMMDD", "%Y%m%d"),
    ("CCYY", "%Y"),
    ("YYYY", "%Y"),
    ("MM", "%m"),
    ("DD", "%d"),
    ("HHMM", "0600"),
    ("HH", "06"),
)
_SEED_LEFTOVER_RE = None  # compiled lazily in _resolve_pattern_tokens


def _resolve_pattern_tokens(pattern: str, date) -> tuple[str, list[str]]:
    """(resolved filename, flagged leftover tokens)."""
    import re as re_module

    global _SEED_LEFTOVER_RE  # noqa: PLW0603 — lazy compile, module cache
    if _SEED_LEFTOVER_RE is None:
        _SEED_LEFTOVER_RE = re_module.compile(
            r"(?<![A-Za-z])(?:CC|YY|MM|DD|HH|SS)[A-Z]*(?![a-z])"
        )
    resolved = pattern
    for token, replacement in _SEED_TOKEN_TABLE:
        value = date.strftime(replacement) if "%" in replacement else replacement
        resolved = re_module.sub(
            rf"(?<![A-Za-z]){token}(?![A-Za-z])", value, resolved
        )
    flagged = _SEED_LEFTOVER_RE.findall(resolved)
    return resolved, flagged


def _synthetic_file_bytes(spec, filename: str) -> bytes:
    """Deterministic synthetic content: header row from the spec's source
    columns + 3 obviously-fake rows (seeded RNG keyed by filename; values
    like MBR000001 — never real-looking PII, never client data)."""
    import random

    columns = [f.source_column for seg in spec.segments for f in seg.fields]
    delimiter = spec.delimiter or ","
    rng = random.Random(f"{spec.feed_slug}/{filename}")
    rows = [delimiter.join(columns)]
    for _ in range(3):
        rows.append(delimiter.join(
            f"{column[:3].upper()}{rng.randint(0, 999999):06d}"
            for column in columns
        ))
    return ("\n".join(rows) + "\n").encode("utf-8")


def _databricks_seed_landing(args: argparse.Namespace, config: Config) -> int:
    """Seed the landing volume with synthetic files (the seam's first write).

    Dry-run by default; --apply creates the volume, uploads, lists back and
    verifies. The target is read from config ONLY and guarded by
    codegen.databricks.WRITABLE_PREFIX — there is no target argument.
    """
    import hashlib
    from datetime import datetime

    from codegen.databricks import (
        DatabricksConfigError,
        DatabricksTransportError,
        config_for,
        ensure_volume,
        list_landing,
        upload_file,
    )
    from codegen.resolve.resolver import resolve_pair

    date = datetime.strptime(args.date, "%Y%m%d") if args.date else datetime.now()

    contracts_dir = Path(config.contracts.dir)
    frd_path = contracts_dir / config.demo.frd
    sttm_path = contracts_dir / config.demo.sttm
    if not (frd_path.is_file() and sttm_path.is_file()):
        print(f"{'FAIL':<15} databricks-seed-landing — demo contract pair not found "
              f"({frd_path.name}, {sttm_path.name}); synthetic content comes from "
              "the resolved demo spec only.")
        report_error("databricks-seed-landing", "contract.not_found", Coord("contract"))
        return 1
    specs = resolve_pair(frd_path, sttm_path, config)

    files: list[tuple[str, bytes]] = []
    flagged_tokens: list[str] = []
    seen: set[str] = set()
    for spec in specs:
        root = (spec.landing_location or "").replace("\\", "/").strip("/")
        if not root:
            print(f"{'SKIP':<15} {spec.feed_slug} — no landing_location in the FRD")
            continue
        for pattern in spec.file_name_patterns:
            name, flagged = _resolve_pattern_tokens(pattern, date)
            flagged_tokens += [f"{spec.feed_slug}/{pattern}: {t}" for t in flagged]
            relative = f"{root}/{name}"
            if relative in seen:  # CCYY and YYYY variants resolve identically
                continue
            seen.add(relative)
            files.append((relative, _synthetic_file_bytes(spec, name)))

    try:
        cfg = config_for(config.databricks)
        target = f"{cfg.catalog}.{cfg.schema}.{cfg.landing_volume}"
    except DatabricksConfigError as exc:
        _fail("databricks-seed-landing", f"databricks-seed-landing — {exc}", "storage",
              exc=exc)
        return 1

    print(f"target volume: {target}  (the ONLY writable location)")
    for relative, data in files:
        print(f"{'WOULD UPLOAD' if not args.apply else 'QUEUED':<15} "
              f"{relative} — {len(data):,} bytes")
    for note in flagged_tokens:
        print(f"{'FLAGGED':<15} unresolved token left literal: {note}")
    if not args.apply:
        print(f"\nDRY RUN — {len(files)} file(s), nothing written. "
              "Re-run with --apply to create the volume and upload.")
        return 0

    answer = input(
        f"\nType 'yes' to create/seed {target} with {len(files)} synthetic "
        "file(s): "
    )
    if answer.strip() != "yes":
        print(f"{'ABORTED':<15} confirmation not given — nothing written.")
        return 1

    try:
        state = ensure_volume(cfg)
        print(f"{'CREATED' if state['created'] else 'EXISTED':<15} {state['full_name']}")
        for relative, data in files:
            path = upload_file(cfg, relative, data, force=args.force)
            print(f"{'UPLOADED':<15} {path} — {len(data):,} bytes")
        listed = list_landing(cfg)
    except (DatabricksConfigError, DatabricksTransportError) as exc:
        _fail("databricks-seed-landing", f"databricks-seed-landing — {exc}", "storage",
              exc=exc)
        return 1

    # Verify: everything sent must list back with the size we sent.
    listed_sizes = {entry["path"]: entry["size"] for entry in listed}
    mismatches = [
        f"{relative}: sent {len(data):,} B, listed "
        f"{listed_sizes.get(relative, 'MISSING')}"
        for relative, data in files
        if listed_sizes.get(relative) != len(data)
    ]
    if mismatches:
        print(f"{'FAIL':<15} verification mismatch:\n  - " + "\n  - ".join(mismatches))
        report_error("databricks-seed-landing", "verify", Coord("storage"))
        return 1

    manifest = {
        "volume": target,
        "seeded_at": datetime.now().isoformat(timespec="seconds"),
        "date_token": date.strftime("%Y%m%d"),
        "files": [
            {"path": relative, "bytes": len(data),
             "sha256": hashlib.sha256(data).hexdigest()}
            for relative, data in files
        ],
    }
    manifest_path = Path(config.output.dir) / "_seed_manifest.json"
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n",
                             encoding="utf-8", newline="\n")
    print(f"\n{'VERIFIED':<15} {len(files)} file(s) listed back with matching sizes")
    for entry in listed:
        print(f"  {entry['path']}  {entry['size']:,} B  {entry['modified']}")
    print(f"manifest -> {manifest_path}")
    return 0


def _contract_path(value: str, config: Config) -> Path:
    path = Path(value)
    if path.is_file():
        return path
    fallback = Path(config.contracts.dir) / value
    if fallback.is_file():
        return fallback
    raise FileNotFoundError(f"contract not found: {value} (also tried {fallback})")


def _take_provider_flag(argv: list[str]) -> tuple[list[str], str | None]:
    """``--provider mock|live`` anywhere on the command line (2026-10-13): the
    one provider switch, global to every command (codegen.reasoning.mode)."""
    rest: list[str] = []
    value: str | None = None
    items = iter(argv)
    for item in items:
        if item == "--provider":
            value = next(items, None)
        elif item.startswith("--provider="):
            value = item.split("=", 1)[1]
        else:
            rest.append(item)
    return rest, value


def main(argv: list[str] | None = None) -> int:
    """The CLI. A ``--provider`` flag holds for THIS command only (reset on return)."""
    from codegen.reasoning.mode import set_flag
    from codegen.run_errors import silence_pydantic_deprecations

    # D1 post50: a Pydantic deprecation on stderr was read as the run's error
    silence_pydantic_deprecations()
    try:
        return _main(argv)
    except Exception as exc:
        # D1 post51: an uncaught error still gets its coordinates-only marker;
        # the traceback (the quoted tier) follows on stderr, exit code unchanged
        report_error(_stage_of(sys.argv[1:] if argv is None else argv), "uncaught", exc=exc)
        raise
    finally:
        set_flag(None)


def _stage_of(argv: list[str]) -> str:
    """The CLI command named on ``argv`` (the marker's stage)."""
    from codegen.run_errors import STAGES

    return next((a for a in argv if a in STAGES), "codegen")


def _parser() -> argparse.ArgumentParser:
    """The CLI's argument parser (every subcommand)."""
    parser = argparse.ArgumentParser(prog="codegen", description=__doc__, epilog=EXIT_CODES,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    subparsers = parser.add_subparsers(dest="command", required=True)

    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--config", default="config/config.yaml")
    common.add_argument(
        "--dry-run",
        action="store_true",
        help="force the mock Layer-2 provider (no network)",
    )
    common.add_argument(
        "--skip-tests",
        action="store_true",
        help="skip running the generated pytest suites",
    )
    common.add_argument(
        "--output-mode",
        choices=["notebook", "framework"],
        action="append",
        default=None,
        help="override output.mode; repeat for both: notebook (Option A, a "
        "standalone PySpark pipeline) and/or framework (Option B: DDL scripts + "
        "config rows + inserts for the existing ingestion framework)",
    )

    generate = subparsers.add_parser(
        "generate", parents=[common], help="generate one contract pair"
    )
    generate.add_argument("--feed", help="generate only this feed_id from the pair")
    generate.add_argument("--frd-contract", "--frd", dest="frd_contract", required=True,
                          help="FRD contract JSON (the pair's FRD side)")
    generate.add_argument("--sttm-contract", "--sttm", dest="sttm_contract", required=True,
                          help="STTM mapping contract JSON (the pair's STTM side)")
    generate.add_argument("--vdd-contract", "--vdd", dest="vdd_contract",
                          help="VDD contract JSON (codegen extract-vdd) — the pair's third input")
    generate.add_argument("--profile", dest="conventions_profile", default=None,
                          help="conventions profile (config conventions.profiles; default "
                               "conventions.profile — edo_sfmc = today's output)")
    generate.add_argument("--iig-template", dest="iig_template", default=None,
                          help="IIG workbook template version (iig_v1 = demo.metadata_sheet, "
                               "or a key of config metadata.templates)")
    generate.add_argument("--playbook-template", dest="playbook_template", default=None,
                          help="retired with the RFC package output: accepted, has no "
                               "effect (outputs are notebook and / or framework)")
    generate.add_argument("--answers", default=None,
                          help="answers file: its inputs: values fill the non-blocking owed "
                               "IIG cells (INPUT_NEEDED lines; ACTIVE_START_DATE = "
                               "feeds[<i>].active_start_date)")
    generate.add_argument("--answers-template", dest="answers_template", default=None,
                          help="where the INPUT_NEEDED items' answers template goes (default: "
                               "answers_template.yaml in the output root)")
    # 2026-10-14: the run's environment (containers, storage account ...) is an input
    environment = generate.add_mutually_exclusive_group()
    environment.add_argument("--environment", default=None,
                             help="the deployment environment profile: D1, A2, PR ... — "
                                  "codegen-state/environments/<name>.yaml over the overlay's "
                                  "environments.<NAME> block; default: the one marked default "
                                  "(an INPUT_NEEDED line records it). A required value left "
                                  "unset is a QUESTION (exit 3)")
    environment.add_argument("--environment-file", dest="environment_file", default=None,
                             help="a YAML block of this run's environment values, taken "
                                  "alone (containers.<layer>, storage_account, catalog_map.*, "
                                  "connection_ids.*, values.<name>); wins over --environment")

    subparsers.add_parser("generate-all", parents=[common], help="generate every configured pair")

    extract = subparsers.add_parser(
        "extract-sttm",
        help="extract an STTM mapping contract from a workbook + FRD contract pair",
    )
    extract.add_argument("--config", default="config/config.yaml")
    extract.add_argument("--workbook", required=True, help="client STTM workbook (.xlsx)")
    extract.add_argument("--frd-contract", required=True)
    extract.add_argument("--out", required=True, help="path for the emitted contract JSON")
    extract.add_argument("--contract-name", help="override the derived contract_name")
    extract.add_argument(
        "--generated-date",
        help="YYYY-MM-DD stamped as generated_date; defaults to today "
        "(inject for byte-reproducible output)",
    )
    extract.add_argument("--layout", help="read through this resolved LayoutProfile JSON "
                                          "instead of discovering the layout")
    extract.add_argument("--answers", help="answers.yaml placing unresolved roles by hand "
                                           "(see `codegen layout --answers`); no model is "
                                           "called")
    extract.add_argument("--require-complete", action="store_true",
                         help="refuse when the layout has unresolved roles")
    extract.add_argument("--answers-template", dest="answers_template", default=None,
                         help="where to write the answers template when answers are owed "
                              "(default: answers_template.yaml beside --out; questions a "
                              "layout run wrote there are kept)")

    layout = subparsers.add_parser(
        "layout",
        help="resolve a workbook's (and optionally its FRD's) layout: cache -> synonyms "
             "-> model -> validate -> user",
    )
    layout.add_argument("--config", default="config/config.yaml")
    layout.add_argument("--workbook", required=True, help="STTM workbook (.xlsx)")
    layout.add_argument("--frd", help="FRD document (.docx) or contract JSON of the pair")
    layout.add_argument("--vdd", help="vendor data dictionary (.xlsx) for the cross-checks")
    layout.add_argument("--dry-run", action="store_true",
                        help="mock provider (answers from fixtures/layout_profiles)")
    layout.add_argument("--no-cache", action="store_true", help="ignore cached profiles")
    layout.add_argument("--frd-contract-out",
                        help="write the FRD contract AS THE PAIR RESOLVED IT (a feed the FRD "
                             "leaves unnamed is named after the STTM stage band, gaps filled "
                             "from the STTM / VDD, each flagged) — use it for extract-sttm / "
                             "generate instead of a separate extract-frd")
    layout.add_argument("--refresh", action="store_true",
                        help="re-resolve past every cache AND overwrite the runtime cache "
                             "entries (a result with open questions tombstones them); the "
                             "reasons a model answer was rejected are written to "
                             "<runtime cache>/rejections/<fingerprint>.json")
    layout.add_argument("--json", action="store_true", help="print the report as JSON")
    layout.add_argument("--profile-out", help="write the resolved STTM profile JSON here")
    layout.add_argument("--answers", help="answers.yaml placing unresolved roles by hand: "
                                          "(document, sheet, role) -> column header text or "
                                          "index; applied with source=user")
    layout.add_argument("--report-unresolved", nargs="?", const="unresolved_headers.md",
                        help="write unresolved_headers.md (or the given path): per unresolved "
                             "role the sheet name, the header row texts and the candidate "
                             "columns — structural labels only, no data rows")
    layout.add_argument("--require-complete", action="store_true",
                        help="print the open questions and exit non-zero when roles remain")
    layout.add_argument("--answers-template", dest="answers_template", default=None,
                        help="with --require-complete: where to write the open questions as an "
                             "answers template (default: answers_template.yaml beside "
                             "--frd-contract-out / --profile-out)")

    pair = subparsers.add_parser(
        "pair",
        help="pair an STTM with its FRD / VDD among a folder's documents by CONTENT (feed "
             "name, target tables, file patterns; the ticket number is one weak signal)",
    )
    pair.add_argument("--config", default="config/config.yaml")
    pair.add_argument("--sttm", required=True, help="STTM workbook (.xlsx)")
    pair.add_argument("--candidates", action="append",
                      help="folder of candidate documents (repeatable; default: the STTM's)")
    pair.add_argument("--answers", help="answers.yaml whose `pairing:` settles an undecided pair")
    pair.add_argument("--pairs-file",
                      help="pairs override file (sttm / frd / vdd groups) that wins over "
                           "every automatic rule; default: pairing_overrides.yaml next to "
                           "the STTM, or inputs.pairing.overrides_file")

    extract_vdd = subparsers.add_parser(
        "extract-vdd",
        help="extract a Vendor Data Dictionary contract from a VDD workbook (patterns V1/V2/V3)",
    )
    extract_vdd.add_argument("--config", default="config/config.yaml")
    extract_vdd.add_argument("--vdd", required=True, help="vendor data dictionary (.xlsx)")
    extract_vdd.add_argument("--out", required=True, help="path for the emitted contract JSON")
    extract_vdd.add_argument("--sttm-contract", help="STTM contract JSON whose table names "
                                                     "select the field sheets of a "
                                                     "one-sheet-per-table dictionary")
    extract_vdd.add_argument("--layout", help="read through this resolved LayoutProfile JSON")
    extract_vdd.add_argument("--profile", help="also write the resolved layout profile JSON here")
    extract_vdd.add_argument("--contract-name", help="override the derived contract_name")
    extract_vdd.add_argument("--generated-date", help="YYYY-MM-DD stamped as generated_date")

    extract_frd = subparsers.add_parser(
        "extract-frd",
        help="extract an FRD feed contract from a client FRD .docx (families F1/F2)",
    )
    extract_frd.add_argument("--config", default="config/config.yaml")
    extract_frd.add_argument("--docx", required=True, help="client FRD document (.docx)")
    extract_frd.add_argument("--out", required=True, help="path for the emitted contract JSON")
    extract_frd.add_argument("--profile", help="also write the resolved FRD layout profile "
                                               "JSON to this path")
    extract_frd.add_argument("--contract-name", help="override the derived contract_name")
    extract_frd.add_argument(
        "--generated-date",
        help="YYYY-MM-DD stamped as generated_date; defaults to today "
        "(inject for byte-reproducible output)",
    )

    fetch = subparsers.add_parser(
        "sharepoint-fetch",
        help="download STTM workbooks / FRD contracts from the SharePoint library",
    )
    fetch.add_argument("--config", default="config/config.yaml")
    fetch.add_argument(
        "--dest", required=True, help="local directory the documents land in"
    )

    demo_sources = subparsers.add_parser(
        "demo-source-files",
        help="print the demo UI's source-files JSON (pure read, no network)",
    )
    demo_sources.add_argument("--config", default="config/config.yaml")

    seed = subparsers.add_parser(
        "databricks-seed-landing",
        help="seed the landing volume with SYNTHETIC files from the demo spec "
        "(dry-run by default; the ONLY writable location is guarded in code)",
    )
    seed.add_argument("--config", default="config/config.yaml")
    seed.add_argument("--apply", action="store_true",
                      help="actually create the volume and upload (typed 'yes' "
                      "confirmation required); default is a dry-run print")
    seed.add_argument("--force", action="store_true",
                      help="allow overwriting files that already exist in the volume")
    seed.add_argument("--date", help="YYYYMMDD resolved into filename date tokens "
                      "(CCYY/YYYY/MM/DD; HHMM fixed at 0600); default today")

    demo_metadata = subparsers.add_parser(
        "demo-metadata-sheet",
        help="print the metadata-sheet preview JSON, or write it as .xlsx "
        "(display only; pure read, no network)",
    )
    demo_metadata.add_argument("--config", default="config/config.yaml")
    demo_metadata.add_argument("--xlsx", help="write the workbook here instead of printing JSON")

    db_fetch = subparsers.add_parser(
        "databricks-fetch",
        help="download raw FRD/STTM documents from the Unity Catalog volumes "
        "(read-only; requires the [databricks] extra)",
    )
    db_fetch.add_argument("--config", default="config/config.yaml")
    db_fetch.add_argument(
        "--dest", default="inputs/databricks",
        help="local directory the documents land in (default: inputs/databricks)",
    )

    publish = subparsers.add_parser(
        "sharepoint-publish",
        help="publish one feed's generated artifacts to the SharePoint library",
    )
    publish.add_argument("--config", default="config/config.yaml")
    publish.add_argument("--feed", required=True, help="feed_slug to publish")
    publish.add_argument(
        "--path",
        help="one file relative to out/<feed_slug>/ (default: the feed's "
        "report + assembled notebook)",
    )

    db_publish = subparsers.add_parser(
        "databricks-publish",
        help="publish one feed's generated artifacts to a Unity Catalog "
        "volume (human-gated write; targets outside WRITABLE_PREFIX refused)",
    )
    db_publish.add_argument("--config", default="config/config.yaml")
    db_publish.add_argument("--feed", required=True, help="feed_slug to publish")
    db_publish.add_argument(
        "--path",
        help="one file relative to out/<feed_slug>/ (default: the feed's "
        "report + assembled notebook)",
    )
    db_publish.add_argument("--catalog", default="",
                            help="override databricks.catalog for the target")
    db_publish.add_argument("--schema", default="",
                            help="override databricks.schema for the target")
    db_publish.add_argument("--volume", default="",
                            help="override databricks.output_volume")
    db_publish.add_argument("--force", action="store_true",
                            help="overwrite files already in the volume")
    doctor = subparsers.add_parser(
        "doctor", help="print which codegen runs (version, location, overlays, python, "
                       "sys.path) and whether it knows the config's sections")
    doctor.add_argument("--config", default="config/config.yaml")
    return parser


def _main(argv: list[str] | None = None) -> int:
    parser = _parser()
    from codegen.reasoning.mode import (
        ProviderConfigError,
        check_live_config,
        set_flag,
    )

    argv, provider_flag = _take_provider_flag(list(sys.argv[1:] if argv is None else argv))
    if provider_flag not in (None, "mock", "live"):
        parser.error(f"--provider {provider_flag!r}: expected mock | live")
    set_flag(provider_flag)
    args = parser.parse_args(argv)
    load_dotenv()
    # First line of EVERY command: which code runs, with which overlays —
    # printed before load_config, which is what fails on a stale install.
    # Machine-readable output (--json) keeps stdout clean: the line goes to stderr.
    # Chunk D (2026-10-09): it also names the conventions profile and IIG
    # template in effect (the committed overlay makes them acfc_prx / iig_v2),
    # read from the raw config + overlays — still before load_config.
    from codegen.build_info import doctor_lines, source_line
    from codegen.config import active_defaults, overlay_paths

    overlays = [str(p) for p in overlay_paths(args.config)]
    defaults = active_defaults(args.config)
    if args.command == "doctor":
        print("\n".join(doctor_lines(overlays, defaults)))
        return 0
    settings = dict(defaults)
    if getattr(args, "dry_run", False) and settings.get("provider", "").startswith(
            "provider=live"):
        # FIRST line of the run (2026-10-13): live asked for, --dry-run wins.
        from codegen.reasoning.mode import DRY_RUN_OVER_LIVE

        print(f"{'WARNING':<15} {DRY_RUN_OVER_LIVE} ({settings['provider']} was asked for; "
              "this command answers from the mock)",
              file=sys.stderr if getattr(args, "json", False) else sys.stdout, flush=True)
        settings["provider"] = DRY_RUN_OVER_LIVE
    for key, flag_value in (("profile", getattr(args, "conventions_profile", None)),
                            ("iig_template", getattr(args, "iig_template", None))):
        if flag_value:
            settings[key], settings[f"{key}_from_flag"] = flag_value, "1"
    print(source_line(overlays, settings),
          file=sys.stderr if getattr(args, "json", False) else sys.stdout, flush=True)
    config = load_config(args.config)
    # Live asked for and not runnable: ONE line naming the unset setting, at
    # startup — never a run that quietly answers from the mock.
    try:
        check_live_config(config)
    except ProviderConfigError as exc:
        _fail("codegen", f"provider — {exc}", "provider", exc=exc)
        return EXIT_FAILED

    if args.command == "extract-sttm":
        return _extract_sttm(args, config)
    if args.command == "extract-frd":
        return _extract_frd(args, config)
    if args.command == "extract-vdd":
        return _extract_vdd(args, config)
    if args.command == "layout":
        return _layout(args, config)
    if args.command == "pair":
        return _pair(args, config)

    if args.command == "demo-source-files":
        # Same JSON as GET /api/demo/source-files — display data only.
        from codegen.demo_sources import source_files_payload

        try:
            # shell_mode pinned synthetic: this command stays offline.
            payload = source_files_payload(config, Path.cwd(), shell_mode="synthetic")
        except FileNotFoundError as exc:
            _fail("demo-source-files", f"demo-source-files — {exc}", "input.not_found",
                  exc=exc)
            return 1
        print(json.dumps(payload, indent=2))
        return 0

    if args.command == "demo-metadata-sheet":
        # Same payload as GET /api/demo/metadata-sheet — display data only.
        # The STTM-derived cells resolve from the demo contract pair when its
        # fixture files exist; otherwise the columns tab is empty, honestly.
        from codegen.metadata_sheet import build_workbook, metadata_sheet_payload

        contracts_dir = Path(config.contracts.dir)
        frd_path = contracts_dir / config.demo.frd
        sttm_path = contracts_dir / config.demo.sttm
        specs = unmapped = None
        if frd_path.is_file() and sttm_path.is_file():
            specs = resolve_pair(frd_path, sttm_path, config)
            unmapped = {
                spec.feed_slug: {
                    o.rule_text
                    for o in compile_rules(spec)
                    if o.classification == "unmapped"
                }
                for spec in specs
            }
        try:
            payload = metadata_sheet_payload(
                config, Path.cwd(), specs=specs, unmapped_by_slug=unmapped
            )
        except FileNotFoundError as exc:
            _fail("demo-metadata-sheet", f"demo-metadata-sheet — {exc}", "input.not_found",
                  exc=exc)
            return 1
        if args.xlsx:
            build_workbook(payload).save(args.xlsx)
            coverage = payload["coverage"]
            print(f"{'WRITTEN':<15} {args.xlsx} — {coverage['derived']}/{coverage['total']} "
                  "cells derived from documents")
        else:
            print(json.dumps(payload, indent=2))
        return 0

    if args.command == "sharepoint-fetch":
        return _sharepoint_fetch(args, config)

    if args.command == "databricks-seed-landing":
        return _databricks_seed_landing(args, config)

    if args.command == "databricks-fetch":
        # Read-side edge of the Databricks volumes seam — same posture as
        # sharepoint-fetch: fetch to local disk, then generate from disk.
        from codegen.databricks import (
            DatabricksConfigError,
            DatabricksTransportError,
            config_for,
            fetch_document,
            list_documents,
        )

        try:
            cfg = config_for(config.databricks)
            listing = list_documents(cfg)
            print(f"volumes: {cfg.catalog}.{cfg.schema}.{cfg.frd_volume} + "
                  f".{cfg.sttm_volume}  (profile {cfg.profile})")
            fetched = 0
            for kind, files in listing.items():
                volume = cfg.frd_volume if kind == "frd" else cfg.sttm_volume
                for item in files:
                    local = fetch_document(cfg, volume, item["name"], args.dest)
                    print(f"{'FETCHED':<15} {local.name} — {item['size']:,} bytes")
                    fetched += 1
        except (DatabricksConfigError, DatabricksTransportError) as exc:
            _fail("databricks-fetch", f"databricks-fetch — {exc}", "storage", exc=exc)
            return 1
        if not fetched:
            print(f"{'FAIL':<15} databricks-fetch — no supported documents in "
                  "either volume. Upload an STTM workbook or an FRD document, or "
                  "check the `databricks:` section of config/config.yaml.")
            report_error("databricks-fetch", "input.not_found", Coord("storage"))
            return 1
        print(f"\n{fetched} file(s) -> {args.dest}")
        return 0

    if args.command == "sharepoint-publish":
        return _sharepoint_publish(args, config)

    if args.command == "databricks-publish":
        return _databricks_publish(args, config)

    if args.command == "generate":
        paths: list[Path] = []
        for role, value in (("FRD", args.frd_contract), ("STTM", args.sttm_contract)):
            try:
                paths.append(_contract_path(value, config))
            except FileNotFoundError as exc:
                _fail("generate", str(exc), "contract.not_found", Coord(role), exc=exc,
                      path=value)
                return 1
        pairs = [(paths[0], paths[1])]
        only_feed = args.feed
    else:
        if not config.contracts.pairs:
            # Empty since 2026-08-22 (contract fixtures removed from the repo).
            # Refuse loudly rather than report a silent zero-feed success.
            print(
                f"{'FAIL':<15} config.contracts.pairs is empty -- nothing to generate. "
                "Restore anonymized contract pairs in config/config.yaml or run "
                "`codegen generate --frd-contract X --sttm-contract Y`."
            )
            report_error(args.command, "config", Coord("config"))
            return 1
        contracts_dir = Path(config.contracts.dir)
        pairs = [(contracts_dir / p.frd, contracts_dir / p.sttm) for p in config.contracts.pairs]
        only_feed = None

    try:
        config, push_outputs = _outputs_through_storage(config)
    except Exception as exc:  # noqa: BLE001 — a storage URI / auth problem, named
        _fail(args.command, f"outputs storage — {exc}", "storage", Coord("output"), exc=exc)
        return 1
    _write_build_stamp(Path(config.output.dir))
    # 2026-10-14: the run's environment (--environment / --environment-file; else
    # the default, recorded by an INPUT_NEEDED line in _run_pairs)
    from codegen.environments import (
        EnvironmentSelectionError,
        select_environment,
        with_environment,
    )

    try:
        env_file = getattr(args, "environment_file", None)
        config = with_environment(config, select_environment(
            config, name=getattr(args, "environment", None),
            file=Path(env_file) if env_file else None,
            live=_live_run(config, getattr(args, "dry_run", False))))
    except EnvironmentSelectionError as exc:
        _fail(args.command, f"environment — {exc}", "environment", Coord("environment"),
              exc=exc)
        return 1
    # M9: `generate` parsed --vdd / --profile / --iig-template /
    # --playbook-template and then dropped them (the notebook run inside ACFC
    # had to edit config.yaml to select acfc_prx, and its VDD never reached
    # the gate). generate-all has none of these flags: getattr -> None.
    vdd_contract = getattr(args, "vdd_contract", None)
    try:
        vdd_path = _contract_path(vdd_contract, config) if vdd_contract else None
    except FileNotFoundError as exc:
        _fail(args.command, str(exc), "contract.not_found", Coord("VDD"), exc=exc,
              path=vdd_contract)
        return 1
    code = _run_pairs(
        pairs, config, only_feed=only_feed, dry_run=args.dry_run,
        skip_tests=args.skip_tests, output_mode=args.output_mode, vdd_path=vdd_path,
        conventions_profile=getattr(args, "conventions_profile", None),
        iig_template=getattr(args, "iig_template", None),
        playbook_template=getattr(args, "playbook_template", None),
        answers_path=Path(args.answers) if getattr(args, "answers", None) else None,
        template_path=(Path(args.answers_template)
                       if getattr(args, "answers_template", None) else None),
        command=args.command,
    )
    push_outputs()
    return code


if __name__ == "__main__":
    sys.exit(main())
