"""post128: the DDL test cases (Nikshit's review, 2026-10-23; Venu, 2026-10-27).

A SEPARATE framework artefact next to the combined DDL, never inside it, named
like the DDL with the profile's ``ddl_test_cases_suffix`` before the extension
(``ACCUM_DDL.txt`` -> ``ACCUM_DDL_TestCases.txt``). It is read-only SQL and
``--`` comments the ENGINEER runs after the DDL in a Databricks SQL editor; the
agent writes it and never runs it. Every statement is a SELECT or a SHOW (the
gate check ``ddl_tests_read_only`` proves it), so the file may be pasted whole.

Every value comes from the resolved feed (the FRD / STTM / VDD it was built
from) or from the DDL the same run wrote; a value no document states is OPEN in
the text and raises one ``ddl_tests_open:<case>`` flag - never a guess. The
cases and their order are the profile's ``ddl_test_cases`` list.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from codegen.contracts.resolved import ResolvedFeedSpec

OPEN = "OPEN (no document states it)"
READ_ONLY_CHECK = "ddl_tests_read_only"
# the Spark CSV reader's default, said once in the mapping_mode case
_SPARK_POSITIONAL = ("Spark applies a given schema to a delimited file BY POSITION and ignores "
                     "the header names unless the header is validated (enforceSchema=false); "
                     "a fixed-width file is positional by definition.")


@dataclass(frozen=True)
class DdlTestCases:
    text: str
    flags: list[str] = field(default_factory=list)
    # post154: the numbered cases (id, title, lines, sql) the workbook is built from
    cases: list[dict] = field(default_factory=list)


# post154: what each FRD load strategy (contracts.frd.LoadStrategy) does when a
# SECOND file is loaded, the row count the engineer enters to check it, and the
# failure
_SECOND_LOAD = {
    "Truncate and Load": ("the table holds only the second file's rows (the first file's "
                          "rows are gone)",
                          "the second file's data record count minus the rows the load "
                          "rejected",
                          "rows of the first file still in the table -> FAIL"),
    "Append": ("the table keeps every row of the first file and adds the second file's",
               "both files' data record counts added up, minus the rows the loads rejected",
               "rows of the first file missing, or a file loaded twice -> FAIL"),
    "Upsert": ("a key present in both files is updated, never duplicated; a new key is added",
               "the number of distinct keys across both files",
               "a key in the table twice -> FAIL (the key case checks it too)"),
}


def ddl_test_cases_name(ddl_name: str, suffix: str) -> str:
    """``ACCUM_DDL.txt`` + ``_TestCases`` -> ``ACCUM_DDL_TestCases.txt``."""
    stem, dot, extension = ddl_name.rpartition(".")
    return f"{stem}{suffix}.{extension}" if dot else f"{ddl_name}{suffix}"


def _prose(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def _line(value: str) -> str:
    """One comment line: inner whitespace collapsed, a leading indent kept."""
    indent = value[:len(value) - len(value.lstrip(" "))]
    return indent + _prose(value)


def _literal(value: str) -> str:
    from codegen.emit.framework import _sql_comment

    return "'" + _sql_comment(value) + "'"


def _parts(qualified: str) -> tuple[str, str, str] | None:
    parts = qualified.split(".")
    return (parts[0], parts[1], parts[2]) if len(parts) == 3 else None


def _fields_of(spec: ResolvedFeedSpec, segments: tuple[str, ...]):
    return [f for s in spec.segments if s.segment in segments for f in s.fields]


def _tables(spec: ResolvedFeedSpec, definitions, profile) -> list[dict]:
    """Every table block the DDL wrote, with what its cases need."""
    tables = []
    required_sources = set(spec.not_null_columns)
    for definition in definitions:
        fields = _fields_of(spec, definition.segments)
        for block in (definition.stage, definition.standard):
            if block is None or not block.qualified:
                continue
            names = {name for name, _type in block.columns}
            if block.layer == "stage":
                column = [f.stage_column for f in fields]
                keyed = [f.stage_column for f in fields if f.primary_key]
                strategy = spec.stage_load_strategy
                properties = profile.stage_tblproperties or {}
            else:
                column = [f.standard_column for f in fields]
                keyed = [f.standard_column for f in fields
                         if (f.standard_primary_key if f.standard_primary_key is not None
                             else f.primary_key)]
                strategy = spec.standard_load_strategy
                properties = profile.standard_tblproperties or {}
            required = [c for f, c in zip(fields, column, strict=True)
                        if c and c in names and (f.source_column in required_sources
                                                 or c in keyed)]
            tables.append({
                "layer": block.layer, "qualified": block.qualified,
                "parts": _parts(block.qualified), "columns": list(block.columns),
                "comment_count": len(block.comments), "table_comment": block.table_comment,
                "key": [c for c in keyed if c in names],
                "required": list(dict.fromkeys(required)),
                "strategy": strategy, "properties": list(properties.items()),
                "mapping": [(f.source_column, f.stage_column, f.standard_column)
                            for f in fields] if block.layer == "stage" else [],
            })
    return tables


def _case(title: str, lines: list[str], sql: list[tuple[str, list[str]]] | None = None) -> dict:
    return {"title": title, "lines": lines,
            "sql": [{"statement": s, "expected": e} for s, e in (sql or [])]}


# the case ID inside a check's SQL, and a parameter the engineer enters (one per
# case: expected_rows_tc15), both set when the cases are numbered
_ID = "'__TEST_ID__'"
_ROWS = "__EXPECTED_ROWS__"


def _check(source: str, passed: str, found: str, expected: str,
           empty: str | None = None) -> tuple[str, list[str]]:
    """post134: a self-checking SELECT - one row: test_id, result (PASS / FAIL, or
    NO DATA when ``empty`` holds: the table has no rows to test), found, expected."""
    no_data = f"WHEN {empty} THEN 'NO DATA' " if empty else ""
    return (f"SELECT {_ID} AS test_id, CASE {no_data}WHEN {passed} THEN 'PASS' ELSE 'FAIL' "
            f"END AS result, CAST({found} AS STRING) AS found, {expected} AS expected "
            f"FROM {source}", [])


def catalog_type(ddl_type: str, profile) -> str:
    """The catalog's spelling of a DDL type (``information_schema.columns.
    full_data_type``): lower case, no spaces, aliases renamed, bare defaults filled."""
    spelled = re.sub(r"\s+", "", ddl_type).lower()
    if spelled in (profile.ddl_test_type_defaults or {}):
        return profile.ddl_test_type_defaults[spelled]
    base, paren, rest = spelled.partition("(")
    base = (profile.ddl_test_type_names or {}).get(base, base)
    if not paren and base in (profile.ddl_test_type_defaults or {}):
        return profile.ddl_test_type_defaults[base]
    return base + paren + rest


def build(spec: ResolvedFeedSpec, definitions, profile, ddl_name: str,
          banner: list[tuple[str, str]], tags: tuple[str, str] | None,
          email_statuses: list[str] | None = None,
          table_source: bool = False) -> DdlTestCases:
    """The test-cases text for one feed's combined DDL (``ddl_name``).
    ``email_statuses`` = the EMAIL_TEMPLATE_CONFIG statuses the same run's IIG
    writes (None = the run wrote no such sheet). ``table_source`` (post156): the
    source is a database table - the cases speak of loads from the source table,
    and the file-only cases (location, delimiter / position, several files) are
    not written."""
    file_only = ("file_location", "mapping_mode", "column_shift", "multiple_files")
    from codegen.emit.emitter import _environment

    flags: list[str] = []
    tables = _tables(spec, definitions, profile)
    stage_tables = [t for t in tables if t["layer"] == "stage"]
    patterns = ", ".join(repr(p) for p in spec.file_name_patterns)
    cases: list[dict] = []

    def open_value(case: str, value, reason: str) -> str:
        if value:
            return _prose(value)
        flags.append(f"ddl_tests_open:{case} — {reason}; the case says OPEN")
        return OPEN

    for name in profile.ddl_test_cases:
        if table_source and name in file_only:
            continue                                  # no file: nothing to check
        if table_source and name == "file_available":
            frequency = open_value(name, spec.frequency, "no document states the frequency")
            cases.append(_case("Source table available", [
                f"Check: the source table {patterns} has rows for the run (frequency: "
                f"{frequency}).",
                "Negative: the source table is empty or missing -> the load stops and reports "
                "it; nothing is loaded.",
                "From: the STTM source table; the FRD frequency."]))
        elif name == "file_available":
            frequency = open_value(name, spec.frequency, "no document states the frequency")
            cases.append(_case("File available", [
                f"Check: a file matching {patterns} arrives for the run (frequency: {frequency}).",
                "Negative: no file for the run -> the load stops and reports the missing file; "
                "nothing is loaded.",
                "From: the FRD / STTM file name pattern; the FRD frequency."]))
        elif name == "file_location":
            landing = open_value(name, spec.landing_location,
                                 "no document states the landing location")
            cases.append(_case("File in the right location", [
                f"Check: the file is in the landing location: {landing}.",
                "Negative: the file sits in another folder -> it is reported missing, never "
                "picked up from there.",
                "From: the FRD landing location."]))
        elif name == "column_mapping":
            test_input = "run a test load" if table_source else "load a test file"
            missing_input = ("the source table without one column" if table_source
                             else "a test file missing one source column")
            for table in stage_tables:
                cases.append(_case(f"Column mapping — {table['qualified']}", [
                    f"Check: {test_input}; each source column lands in the column the STTM "
                    "maps it to (source -> stage -> standard):",
                    *(f"  {s} -> {st} -> {sd or '(stage only)'}" for s, st, sd in
                      table["mapping"]),
                    f"Negative: {missing_input} -> the load fails or "
                    "rejects the rows; no column is left silently empty.",
                    "From: the STTM mapping rows."]))
        elif name == "columns_populated":
            for table in tables:
                if not table["required"]:
                    cases.append(_case(f"Required columns populated — {table['qualified']}", [
                        "No document marks a column of this table required or as a key: no "
                        "check."]))
                    continue
                missing = " OR ".join(f"{c} IS NULL" for c in table["required"])
                cases.append(_case(f"Required columns populated — {table['qualified']}", [
                    f"Check: every row fills {', '.join(table['required'])}.",
                    f"Negative: a test file with {table['required'][0]} blank -> the row is "
                    "rejected or flagged, never loaded with a NULL.",
                    "From: the STTM / FRD required (not-null) and primary-key columns."],
                    [_check(table["qualified"],
                            f"SUM(CASE WHEN {missing} THEN 1 ELSE 0 END) = 0",
                            f"SUM(CASE WHEN {missing} THEN 1 ELSE 0 END)",
                            "'0 rows with a required column empty'",
                            empty="COUNT(*) = 0")]))
        elif name == "column_shift":
            for table in stage_tables:
                columns = [c for c, _t in table["columns"]]
                cases.append(_case(f"Column shift — {table['qualified']}", [
                    f"Check: the file's {len(table['mapping'])} source columns arrive in the "
                    f"STTM order, first {table['mapping'][0][0]!r}, last "
                    f"{table['mapping'][-1][0]!r}; the table has {len(columns)} columns "
                    "(audit columns included).",
                    "Negative: a test file with two columns swapped, or one inserted -> the load "
                    "fails or rejects; no value ever lands one column over.",
                    "From: the STTM column order (and the VDD positions for a fixed-width "
                    "file). How this behaves depends on the next case."]))
        elif name == "mapping_mode":
            delimiter = (repr(spec.delimiter) if spec.delimiter
                         else "none (fixed width)")
            cases.append(_case("Columns matched by name or by position?", [
                f"File: format {_prose(spec.file_format)}, delimiter {delimiter}.",
                f"Note: {_SPARK_POSITIONAL}",
                "Question for the framework owners: does the framework validate the header row? "
                "If not, the column shift case loads shifted values silently - its negative "
                "test is the one that matters most."]))
        elif name == "record_count":
            for table in stage_tables:
                source_rows = ("the source table's row count" if table_source
                               else "the file's data record count")
                cases.append(_case(f"Record count — {table['qualified']}", [
                    "Check: every " + ("row of the source table" if table_source
                                       else "data record of the file")
                    + " is loaded (or rejected and counted).",
                    f"Enter {_ROWS} when asked: {source_rows} minus the "
                    "rows the load rejected, after a first load into the empty table.",
                    "From: " + ("the source table at test time." if table_source else
                                "the file itself at test time (its trailer count where it "
                                "has one).")],
                    [_check(table["qualified"],
                            f"COUNT(*) = CAST(:{_ROWS} AS BIGINT)", "COUNT(*)",
                            f"CAST(:{_ROWS} AS STRING)")]))
        elif name == "key_unique":
            for table in tables:
                if not table["key"]:
                    cases.append(_case(f"Primary key unique — {table['qualified']}", [
                        "No document states a primary key for this table: no check."]))
                    continue
                key = ", ".join(table["key"])
                strategy = _prose(table["strategy"]) if table["strategy"] else OPEN
                cases.append(_case(f"Primary key unique — {table['qualified']}", [
                    f"Check: the primary key ({key}) is unique (load strategy: {strategy}).",
                    "Negative: load the same file twice -> the key stays unique (an Upsert "
                    "updates the row; an Append is stopped by the processed-files check).",
                    "From: the STTM Primary Key cells; the FRD load strategy."],
                    [_check(f"(SELECT {key} FROM {table['qualified']} GROUP BY {key} "
                            "HAVING COUNT(*) > 1) AS duplicates", "COUNT(*) = 0", "COUNT(*)",
                            "'0 duplicated keys'",
                            empty=f"(SELECT COUNT(*) FROM {table['qualified']}) = 0")]))
        elif name in ("table_exists", "columns_match", "comments_present", "table_tags"):
            for table in tables:
                case = _schema_case(name, table, tags, profile)
                if case is not None:
                    cases.append(case)
                elif table["parts"] is None:
                    flags.append(f"ddl_tests_unqualified:{table['qualified']} — not a "
                                 f"three-part name; no {name} case")
        elif name == "table_properties":
            for table in tables:
                if not table["properties"]:
                    continue
                cases.append(_case(f"Table properties — {table['qualified']}", [
                    "Check: the TBLPROPERTIES the DDL sets are on the deployed table.",
                    "Checked by eye: SHOW returns no PASS / FAIL row (a query cannot read "
                    "table properties).",
                    "From: the DDL (the profile's table properties)."],
                    [(f"SHOW TBLPROPERTIES {table['qualified']}",
                      [f"includes {k} = {v}" for k, v in table["properties"]])]))
        elif name == "load_strategy":
            for table in tables:
                label = f"{table['qualified']} ({table['layer']})"
                strategy = table["strategy"]
                if not strategy:
                    open_value(name, None, f"no load strategy stated for {table['qualified']}")
                    cases.append(_case(f"Load strategy on a second file — {label}", [
                        f"Load strategy: {OPEN}: no check until the FRD states it."]))
                    continue
                if strategy not in _SECOND_LOAD:
                    flags.append(f"ddl_tests_load_strategy_unread:{table['qualified']} — "
                                 f"{strategy!r} is not a strategy this case knows; no check")
                    cases.append(_case(f"Load strategy on a second file — {label}", [
                        f"Load strategy: {_prose(strategy)}: no check written."]))
                    continue
                behaviour, enter, negative = _SECOND_LOAD[strategy]
                if table_source:
                    behaviour, enter, negative = (
                        behaviour.replace("second file's", "second load's")
                        .replace("first file's", "first load's"),
                        enter.replace("the second file's data record count",
                                      "the source table's row count at the second load")
                        .replace("both files' data record counts added up",
                                 "both loads' row counts added up")
                        .replace("both files", "both loads"),
                        negative.replace("first file", "first load")
                        .replace("a file loaded twice", "a load counted twice"))
                first_then = ("load once, then again after the source table changes"
                              if table_source else "load a first test file, then a second")
                cases.append(_case(f"Load strategy on a second {'load' if table_source else 'file'}"
                                   f" — {label}", [
                    f"Check: {first_then}; with the load strategy "
                    f"{strategy}, {behaviour}.",
                    f"Enter {_ROWS} when asked: {enter}.",
                    f"Negative: {negative}.",
                    f"From: the FRD load strategy ({table['layer']})."],
                    [_check(table["qualified"], f"COUNT(*) = CAST(:{_ROWS} AS BIGINT)",
                            "COUNT(*)", f"CAST(:{_ROWS} AS STRING)")]))
        elif name == "second_file" and table_source:
            cases.append(_case("A second load", [
                "Check: after the source table changes, load again and run every case "
                "above; each still passes.",
                "Negative: a case that passed on the first load fails on the second -> FAIL.",
                "From: the STTM source table."]))
        elif name == "second_file":
            cases.append(_case("A second sample file", [
                f"Check: load another sample file matching {patterns} (a later business "
                "date) and run every case above again; each still passes.",
                "Negative: a case that passed on the first file fails on the second -> FAIL.",
                "From: the FRD / STTM file name pattern."]))
        elif name == "multiple_files":
            landing = open_value(name, spec.landing_location,
                                 "no document states the landing location")
            processed = getattr(spec.processed_files_table, "qualified", None)
            cases.append(_case("Several files at once", [
                f"Check: place two files matching {patterns} in {landing} at the same time; "
                "both are loaded, each once"
                + (f", and each appears once in {processed}." if processed else "."),
                "Negative: a file skipped, or loaded twice -> FAIL.",
                "From: the FRD file name pattern and landing location"
                + ("; the processed-files table." if processed else ".")]))
        elif name == "notification_emails":
            if email_statuses is None:
                continue                                  # no email sheet in this IIG
            if not email_statuses:
                flags.append("ddl_tests_open:notification_emails — the IIG configures no "
                             "email status; the case says OPEN")
                cases.append(_case("Email notifications", [f"Statuses: {OPEN}."]))
                continue
            triggers = profile.ddl_test_email_triggers or {}
            for status in email_statuses:
                when = triggers.get(status)
                cases.append(_case(f"Email notification — {status}", [
                    f"Check: the framework sends the {status!r} email"
                    + (f" when {when}" if when else "") + ", to the recipients the IIG row "
                    "states.",
                    f"Negative: {when or 'the event happens'} and no {status!r} email "
                    "arrives -> FAIL.",
                    "From: the IIG's EMAIL_TEMPLATE_CONFIG row for this status."]))
        elif name == "ddl_rerun":
            cases.append(_case("The DDL is safe to run again", [
                f"Check: run {ddl_name} a second time against the deployed tables -> no error, "
                f"and every table keeps its rows (each statement is "
                f"{profile.create_statement}).",
                "Negative: none - a re-run never drops or empties a table.",
                "From: the DDL."]))
    for number, case in enumerate(cases, start=1):
        case["id"] = f"TC{number:02d}"
        # every comment line ONE physical line: a document value spanning lines (an
        # STTM cell "Name\n(Type)") must never leave its '--' (the 12-pair check)
        case["title"] = _prose(case["title"])
        rows = f"expected_rows_{case['id'].lower()}"
        case["lines"] = [_line(line).replace(_ROWS, rows) for line in case["lines"]]
        for q in case["sql"]:
            q["statement"] = (_prose(q["statement"]).replace(_ID, f"'{case['id']}'")
                              .replace(_ROWS, rows))
            q["expected"] = [_line(line) for line in q["expected"]]
    text = _environment().get_template("framework/ddl_test_cases.txt.j2").render(
        ddl_name=ddl_name, feed=spec.feed_slug, banner=banner, cases=cases,
        table_source=table_source)
    return DdlTestCases(text=text, flags=flags, cases=cases)


def workbook_rows(cases: list[dict], profile, table_source: bool = False) -> list[list[str]]:
    """post154: the cases as rows in the client's test-record layout
    (``ddl_test_workbook_headers``; each column filled by its
    ``ddl_test_workbook_columns`` kind). The tester fills the blank columns."""
    rows = []
    for case in cases:
        lines = case["lines"]
        source = [ln for ln in lines if ln.startswith("From:")]
        negative = [ln for ln in lines if ln.startswith("Negative:")]
        body = [ln for ln in lines if ln not in source and ln not in negative]
        expected = []
        if any(q["statement"].startswith("SELECT") for q in case["sql"]):
            expected.append("Each query returns result PASS (NO DATA = "
                            + ("run a test load" if table_source else "load a test file")
                            + " first).")
        expected += [ln for q in case["sql"] for ln in q["expected"]] + negative
        fill = {
            "id": case["id"],
            "case": "\n".join([case["title"], *body]),
            "expected": "\n".join(expected),
            "sql": "".join(q["statement"] + ";\n" for q in case["sql"]).rstrip("\n"),
            "from": "\n".join(ln.removeprefix("From:").strip() for ln in source),
            "blank": "",
        }
        rows.append([fill[kind] for kind in profile.ddl_test_workbook_columns])
    return rows


def workbook(sheets: list[tuple[str, list[list[str]]]], profile, banner):
    """post154: the test-record workbook - one sheet per feed (title, rows), the
    headers bold and frozen, then a _provenance sheet (the same banner as the
    text file). Byte-stable through metadata_sheet.stable_workbook_bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    from codegen.metadata_sheet import _pin_workbook_properties

    book = Workbook()
    _pin_workbook_properties(book)
    book.remove(book.active)
    headers = profile.ddl_test_workbook_headers
    widths = {"id": 12, "case": 60, "expected": 60, "sql": 80, "from": 40, "blank": 18}
    for title, rows in sheets:
        sheet = book.create_sheet(title=title[:31])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for column, kind in enumerate(profile.ddl_test_workbook_columns, start=1):
            letter = sheet.cell(row=1, column=column).column_letter
            sheet.column_dimensions[letter].width = widths[kind]
            for cell in sheet[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
    provenance = book.create_sheet(title="_provenance")
    provenance.append(["input", "value"])
    for key, value in banner:
        provenance.append([key, value])
    return book


def _schema_case(name: str, table: dict, tags: tuple[str, str] | None, profile) -> dict | None:
    """A structure case read from the catalog's information_schema; Unity Catalog
    stores schema and table names in lower case, so they are compared that way."""
    if table["parts"] is None:
        return None
    catalog, schema, table_name = table["parts"]
    where = (f"WHERE table_schema = {_literal(schema.lower())} AND table_name = "
             f"{_literal(table_name.lower())}")
    label = f"{table['qualified']} ({table['layer']})"
    if name == "table_exists":
        return _case(f"Table exists — {label}", [
            "Check: the DDL created the table in the stated catalog and schema.",
            "From: the DDL."],
            [_check(f"{catalog}.information_schema.tables {where}", "COUNT(*) = 1",
                    "COUNT(*)", "'1'")])
    if name == "columns_match":
        columns = table["columns"]
        expected = ", ".join(f"({i}, {_literal(c.lower())}, {_literal(catalog_type(t, profile))})"
                             for i, (c, t) in enumerate(columns, start=1))
        # post153: the catalog counts from ddl_test_ordinal_base (Databricks: 0); the
        # DDL's columns are numbered from 1 above, so a FAIL reads the same way
        base = profile.ddl_test_ordinal_base
        position = ("ordinal_position" if base is None or base == 1
                    else f"ordinal_position + {1 - base}")
        deployed = (f"(SELECT {position} AS col_position, lower(column_name) AS col_name, "
                    f"replace(lower(full_data_type), ' ', '') AS col_type FROM "
                    f"{catalog}.information_schema.columns {where}) AS deployed")
        source = (f"{deployed} FULL OUTER JOIN VALUES {expected} AS ddl(col_position, col_name, "
                  "col_type) ON deployed.col_position = ddl.col_position AND deployed.col_name "
                  "= ddl.col_name AND deployed.col_type = ddl.col_type WHERE deployed.col_name "
                  "IS NULL OR ddl.col_name IS NULL")
        found = ("CASE WHEN COUNT(*) = 0 THEN '0 differences' ELSE concat_ws(', ', "
                 "array_sort(collect_list(concat(CASE WHEN ddl.col_name IS NULL THEN 'deployed "
                 "' ELSE 'DDL ' END, CAST(coalesce(deployed.col_position, ddl.col_position) AS "
                 "STRING), ' ', coalesce(deployed.col_name, ddl.col_name), ' ', "
                 "coalesce(deployed.col_type, ddl.col_type))))) END")
        return _case(f"Columns and types match the DDL — {label}", [
            f"Check: the deployed table has the DDL's {len(columns)} columns, in its order, "
            "with its types (names and types compared in lower case; a type in the catalog's "
            "spelling, e.g. INTEGER = int).",
            "A FAIL lists each difference: 'DDL n name type' is missing or different on the "
            "table, 'deployed n name type' is what the table has instead.",
            "From: the DDL (the STTM columns and types)."],
            [_check(source, "COUNT(*) = 0", found, "'0 differences'")])
    if name == "comments_present":
        sql = [_check(f"{catalog}.information_schema.columns {where} AND comment IS NOT NULL",
                      f"COUNT(*) = {table['comment_count']}", "COUNT(*)",
                      f"'{table['comment_count']}'")]
        if table["table_comment"]:
            comment = _literal(table["table_comment"])
            sql.append(_check(f"{catalog}.information_schema.tables {where}",
                              f"max(comment) = {comment}",
                              "coalesce(max(comment), '(no table comment)')", comment))
        return _case(f"Comments present — {label}", [
            "Check: the column and table COMMENTs the DDL writes are on the deployed table.",
            "From: the DDL (the STTM descriptions, primary-key notes, the FRD load "
            "strategy)."], sql)
    if name == "table_tags":
        if tags is None:
            return None
        domain, subdomain = _literal(tags[0]), _literal(tags[1])
        return _case(f"Tags — {label}", [
            "Check: the DOMAIN / SUBDOMAIN tags the DDL sets are on the deployed table.",
            "From: the FRD Domain / Sub-Domain."],
            [_check(f"{catalog}.information_schema.table_tags WHERE schema_name = "
                    f"{_literal(schema.lower())} AND table_name = "
                    f"{_literal(table_name.lower())}",
                    f"SUM(CASE WHEN (tag_name = 'DOMAIN' AND tag_value = {domain}) OR "
                    f"(tag_name = 'SUBDOMAIN' AND tag_value = {subdomain}) THEN 1 ELSE 0 END) "
                    "= 2",
                    "coalesce(concat_ws(', ', array_sort(collect_list(concat(tag_name, ' = ', "
                    "tag_value)))), '')",
                    _literal(f"DOMAIN = {tags[0]}, SUBDOMAIN = {tags[1]}"))])
    return None
