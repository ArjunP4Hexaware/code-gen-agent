"""Option B: DDL scripts + config rows + SQL inserts for the ACFC framework.

The deck's "CodeGen Agent Architecture" slide: the approved STTM becomes
deployable artefacts — Option A a PySpark notebook (fresh pipeline, ~10% of
runs), or Option B an *addition* to the existing ingestion framework (~90%):
DDL scripts as a reviewable workbook, config rows as the approval artefact,
and SQL insert statements loaded into the framework database only after a
person approves. This module renders Option B into
``out/<feed_slug>/framework/``.

Row building is NOT duplicated here: ``codegen.metadata_sheet`` (the layout
named by ``framework.metadata_layout``) is the single source of truth — the
same tabs, headers, provenance badges and coverage the preview panel shows
are what gets approved and inserted. Framework-assigned ID columns
(``always_blank``) render as ``framework.id_placeholder`` and are flagged,
never invented. DDL stays engineer-run: nothing here creates tables, and
the inserts are plain INSERTs — idempotency belongs to the framework's own
load path.

Multi-table step 6: the combined CREATE text is one block per table and
layer (``table_definitions``), and ``metadata_inserts.sql``
(``codegen.emit.metadata_inserts``) carries the metadata rows that define those
tables — the client's "DDL" — from the same payload as the IIG, on EVERY
framework run (2026-10-08). The profile's ``emit_dml`` + ``dml.enabled``
switches now gate only the runner notebooks that execute it; the
per-environment ``config_inserts_<env>.sql`` it replaced is retired.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

from codegen.config import Config
from codegen.contracts.resolved import ResolvedFeedSpec, ResolvedTable
from codegen.emit.metadata_inserts import FILE_NAME as METADATA_INSERTS_FILE
from codegen.emit.metadata_inserts import emit_metadata_inserts
from codegen.faq import LoadPatternFaq, summarize
from codegen.gate.derivations import (
    check_iig_derivations,
    check_iig_workbooks_single_line,
    check_iig_workbooks_single_value,
    check_sql_literals,
    join_path,
    normalise_iig_line_breaks,
    normalise_path,
    path_normalised_flag,
    path_whitespace_flags,
)
from codegen.gate.preflight import GateCheck
from codegen.metadata_sheet import (
    BADGE_LABELS,
    DERIVED_BADGES,
    build_workbook,
    metadata_sheet_payload,
)

_ADDITION_TEMPLATE = """\
# Framework addition — {slug}

{banner}

An **addition to the existing ingestion framework**, not a standalone
pipeline. The agent writes the files below; it never edits a notebook,
never creates a table and never applies anything itself.

| File | What it is |
| --- | --- |
{ddl_row}{iig_rows}
| `ADDITION.md` | This manifest. |

## Review flow

{review_flow}

## Framework-assigned IDs left blank

The framework assigns these; the agent never invents them. Each is left as
`{id_placeholder}`:

{blank_columns}
{segmented_block}
Layout note: the tab names and column headers are the client IIG template
(anonymized reference — `fixtures/reference/SFMC_IIG.xlsx`).
"""


# M7 §4: artefacts labelled by the system they run against.
TARGET_DDL = "DDL — Databricks (Unity Catalog)"
TARGET_DML = "DML — SQL Server metadata DB (run from notebook)"
TARGET_REVIEW = "Review sheets"
TARGET_NOTES = "Notes"
DDL_TARGET_HEADER = "-- TARGET SYSTEM = Databricks (Unity Catalog) — run in SQL editor/notebook"


def artefact_group(name: str) -> str:
    if name.endswith("_table_creation.txt") or name.endswith("_DDL.txt"):
        return TARGET_DDL
    if name == METADATA_INSERTS_FILE:
        # Multi-table step 6: the client's "DDL" (rule 8) — the metadata rows
        # that define the Unity Catalog tables, beside their CREATE reference.
        return TARGET_DDL
    if name.startswith("Insert_scripts_config_table_") and name.endswith(".py"):
        return TARGET_DML
    if name.endswith(".xlsx"):
        return TARGET_REVIEW
    return TARGET_NOTES


def artefact_groups(files: list[Path]) -> dict[str, list[str]]:
    """target system -> file names, groups in a fixed order."""
    out: dict[str, list[str]] = {}
    for group in (TARGET_DDL, TARGET_DML, TARGET_REVIEW, TARGET_NOTES):
        names = [p.name for p in files if artefact_group(p.name) == group]
        if names:
            out[group] = names
    return out


@dataclass(frozen=True)
class FrameworkArtefacts:
    """What one feed's Option B emit produced (for reports and the UI)."""

    files: list[Path]
    row_counts: dict[str, int]
    coverage: dict[str, int]
    flagged_blank_columns: list[str] = field(default_factory=list)
    # Segmented-extraction surfacing (empty/None on flat feeds): the ASSUMED
    # discriminator line and the held-back workbook Standard layer.
    assumed_notes: list[str] = field(default_factory=list)
    held_back: list[str] = field(default_factory=list)
    # M4: gate flags this emit raised (blank-and-flag IIG columns, a DDL
    # file name derived from the slug) — joined to the verdict by the caller.
    flags: list[str] = field(default_factory=list)
    # M7: gate CHECKS this emit ran (derived names / paths, SQL literals,
    # three-part qualification) — a failed one FAILs the feed.
    checks: list[GateCheck] = field(default_factory=list)
    # M7: the framework payload (IIG rows) the DML emitter renders from.
    payload: dict | None = None
    # M7 §4: file names grouped by target system (artefact_groups).
    groups: dict[str, list[str]] = field(default_factory=dict)
    # IIG-first: why no DML was written (None = it was). Shown in the report.
    dml_disabled_reason: str | None = None
    # post66: IIG sheets left empty BY RULE -> the reason (MFT / delimited),
    # in sheet order — the App's review summary lists them.
    empty_by_rule: dict[str, str] = field(default_factory=dict)


def _sanitize_abbrev(slug: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", slug).strip("_").upper()


def _config_default_catalog(config: Config | None, profile, layer: str) -> str | None:
    """The catalog chain's last link: the profile's default_catalog, else the
    global conventions.default_catalog (M7.1). 2026-10-14 (post47): the run's
    environment's default_catalog first (codegen.environments)."""
    from codegen.environments import default_catalog

    env = default_catalog(config, layer)
    if env:
        return env
    own = profile.default_catalog.get(layer) if profile is not None else None
    if own or config is None:
        return own
    return config.conventions.default_catalog.get(layer)


def _custom_env_without_catalog(config: Config | None, layer: str) -> bool:
    """2026-10-14 (post47): the run's environment owes the catalog for
    ``layer`` — a custom one states none, or (post47b, provider=mock) a named
    one's is a placeholder — an INPUT_NEEDED (run.environment.<key>), never a
    FAIL: the DDL waits for the value, the run proceeds."""
    from codegen.environments import default_catalog_unset

    return default_catalog_unset(config, layer)


def _env_catalog_unset_flag(config: Config, layer: str, name: str) -> str:
    from codegen.environments import run_environment

    env = run_environment(config)
    return (f"catalog_env_unset:{layer} — {name}: environment {env.name} ({env.source}) leaves "
            f"the {layer} catalog unset (catalog_map / default_catalog.{layer}; never taken "
            f"from another environment); the {layer} DDL is not written until it is given "
            "(INPUT_NEEDED run.environment.*)")


def _qualify(table, layer: str, profile, spec: ResolvedFeedSpec,
             flags: list[str], config: Config) -> str | None:
    """catalog.schema.table for a CREATE, or None when the profile requires
    three parts and no source states the catalog / schema. Chain: the
    resolved table's catalog (FRD label, else the STTM band — the resolver
    carries whichever stated one) -> profile / conventions default_catalog
    (provenance config_default); each fallback is a provenance flag;
    nothing is invented."""
    catalog, schema, name = table.catalog, table.schema_name, table.table
    default = _config_default_catalog(config, profile, layer)
    if not catalog and default:
        catalog = default
        flags.append(f"catalog_from_config:{layer} — {schema}.{name}: no FRD / STTM catalog; "
                     f"provenance config_default (conventions default_catalog[{layer}] = "
                     f"{catalog!r})")
    if profile.require_qualified_names and not (catalog and schema):
        missing = "catalog" if not catalog else "schema"
        if missing == "catalog" and _custom_env_without_catalog(config, layer):
            flags.append(_env_catalog_unset_flag(config, layer, f"{schema or '?'}.{name}"))
            return None
        flags.append(f"{missing}_unstated:{layer} — {schema or '?'}.{name}: sources checked: "
                     "FRD Target Catalog and Schema label, STTM target band catalog column, "
                     f"conventions profile default_catalog[{layer}] — none states one; the "
                     f"{layer} DDL is not written (a two-part name is never written)")
        return None
    return ".".join(p for p in (catalog, schema, name) if p)


def _sttm_row(f) -> str:
    """Where a field's mapping row sits in the STTM, for a flag's citation."""
    if f.provenance is None:
        return "STTM (row not recorded)"
    return f"STTM {f.provenance.sheet} row {f.provenance.row}"


def _standard_from_stage_type(f, stage_type: str, flags: list[str]) -> str:
    """The standard-table type of one column under ``standard_from_stage``:
    the STTM standard band's type, never the stage band's.

    * Same type as the stage column once both are normalised through
      ``sql_type`` -> the stage column's text, unchanged (a feed whose bands
      agree, pair 1, keeps its golden byte for byte).
    * A different type -> ``sql_type`` of the standard band's type
      (``Date`` -> ``DATE``, ``Decimal(18,4)`` -> ``DECIMAL(18,4)``).
    * Blank -> the stage type, flagged ``standard_type_blank``.
    * No SQL mapping -> written as the STTM states it (never guessed),
      flagged ``standard_type_unmapped``."""
    from codegen.emit.context import TemplateGapError, sql_type

    stated = (f.standard_datatype or "").strip()
    if not stated:
        flags.append(f"standard_type_blank:{f.stage_column} — {_sttm_row(f)} states no "
                     f"standard-band data type; the standard table uses the stage type "
                     f"{stage_type!r}")
        return stage_type
    try:
        standard = sql_type(stated)
    except TemplateGapError:
        flags.append(f"standard_type_unmapped:{f.stage_column} — {_sttm_row(f)} standard-band "
                     f"data type {stated!r} maps to no SQL type; written as stated, not "
                     "guessed")
        return stated
    try:
        stage = sql_type(stage_type)
    except TemplateGapError:
        stage = None
    return stage_type if stage == standard else standard


@dataclass(frozen=True)
class TableBlock:
    """One table definition of one layer (multi-table step 6): the resolved
    table, its MAPPED three-part name (``_qualify``; None = the profile
    requires three parts and no source states the catalog / schema — no
    CREATE is written for it), its business columns and its audit columns."""

    layer: str                               # "stage" | "standard"
    table: ResolvedTable
    qualified: str | None
    business: tuple[tuple[str, str], ...]    # (column, type), STTM order
    audit: tuple[tuple[str, str], ...]
    # post104: column -> its DDL COMMENT text (unescaped; a column absent has
    # none) and the table COMMENT, under the profile's comment knobs.
    comments: tuple[tuple[str, str], ...] = ()
    table_comment: str | None = None

    @property
    def label(self) -> str:
        """The name a reader sees: the mapped three-part name, else the
        resolved table's own (two-part) name."""
        return self.qualified or self.table.qualified_name

    @property
    def columns(self) -> list[tuple[str, str]]:
        return [*self.business, *self.audit]


@dataclass(frozen=True)
class TableDefinition:
    """One table the feed defines (rule 1: one per distinct stage triple):
    its stage definition and, when its rows carry a Standard band, its
    standard definition."""

    segments: tuple[str, ...]
    stage: TableBlock
    standard: TableBlock | None


def _column_comment(f, layer: str, profile) -> str | None:
    """post104: a column's DDL COMMENT text under the profile's knobs - the
    STTM field description (``column_comments``), then ``primary_key_note``
    when the layer's Primary Key cell marks the column; None = nothing."""
    parts = []
    if profile.column_comments and f.ddl_description:
        parts.append(f.ddl_description)
    key = (f.standard_primary_key if layer == "standard" and f.standard_primary_key is not None
           else f.primary_key)
    if key and profile.primary_key_note:
        parts.append(profile.primary_key_note)
    return profile.column_comment_separator.join(parts) or None


def _comment_blank_flag(layer: str, table: str, blank: list[str]) -> str:
    """One flag per table block: the columns (audit ones included) the DDL
    leaves without a COMMENT."""
    return (f"ddl_comment_blank:{layer}:{table} — {len(blank)} column(s) without a COMMENT: "
            "the STTM states no Field Description / Column Description / Description / "
            f"Data Definition for {', '.join(blank)} (a Comments cell is rule text, never a "
            "description)")


def _load_strategy_comment(profile, strategy, layer_label: str) -> str | None:
    if not profile.load_strategy_comment or not strategy:
        return None
    return profile.load_strategy_comment.format(strategy=strategy, layer=layer_label)


def table_definitions(spec: ResolvedFeedSpec, profile, flags: list[str] | None = None,
                      config: Config | None = None) -> list[TableDefinition]:
    """Every table the feed defines, one entry per distinct stage (catalog,
    schema, table) — ``metadata_template._table_groups`` — in first-seen
    order. The single derivation behind the combined CREATE reference text
    AND the metadata inserts' table index (so the two never disagree).

    Per table:

    * stage columns = the table's own rows (a column shared by its segments
      once, STTM order), typed as the STTM stage band states them under
      ``typed_stage`` (else ``STRING``);
    * standard columns = the rows of the same table that carry a Standard
      band (a row whose Standard band is blank is not carried), named as
      that band names them; the type is the Standard band's — under
      ``standard_from_stage`` through ``_standard_from_stage_type`` (the
      stage text when both normalise to the same SQL type), else as stated;
      a table NONE of whose rows carries a Standard band keeps the earlier
      ``standard_from_stage`` reading (the stage list, typed per column);
    * audit columns = the table's segments' own audit rows when every
      segment states them, else the feed's, in the profile's casing;
    * names = ``_qualify`` (the resolved, catalog-MAPPED table; the
      ``default_catalog`` fallback and the three-part rule raise their
      flags here, once per table).

    A single-table feed whose Standard band names every stage column as the
    stage band does (pair 1) is therefore exactly the pre-step-6 output."""
    from codegen.metadata_template import _table_groups

    flags = flags if flags is not None else []
    out: list[TableDefinition] = []
    for group in _table_groups(spec):
        stage_columns: list[tuple[str, str]] = []
        standard_columns: list[tuple[str, str]] = []
        stage_comments: dict[str, str] = {}
        standard_comments: dict[str, str] = {}
        stage_blank: list[str] = []
        standard_blank: list[str] = []
        # A table NONE of whose rows carries a Standard band (the standard
        # table named by the FRD only) keeps the standard_from_stage
        # convention as it was: the stage list, typed per column.
        no_band = not any(f.standard_column is not None for f in group.fields)
        for f in group.fields:
            stage_type = f.stage_datatype if profile.typed_stage else "STRING"
            stage_columns.append((f.stage_column, stage_type))
            _note_comment(f, "stage", f.stage_column, profile, stage_comments, stage_blank)
            if group.standard is None:
                continue    # Chunk B: a stage-only table (no standard flags for it)
            if no_band and profile.standard_from_stage:
                standard_columns.append(
                    (f.stage_column, _standard_from_stage_type(f, stage_type, flags)))
                _note_comment(f, "standard", f.stage_column, profile, standard_comments,
                              standard_blank)
                continue
            if f.standard_column is None:
                continue
            _note_comment(f, "standard", f.standard_column, profile, standard_comments,
                          standard_blank)
            if profile.standard_from_stage:
                standard_columns.append(
                    (f.standard_column, _standard_from_stage_type(f, stage_type, flags)))
            else:
                standard_columns.append((f.standard_column, f.standard_datatype or ""))
        if all(s.audit_columns for s in group.segments):
            audit_rows = [a for s in group.segments for a in s.audit_columns]
        else:
            audit_rows = list(spec.audit_columns)
        seen: dict[str, str] = {}
        audit_descriptions: dict[str, str | None] = {}
        for a in audit_rows:
            seen.setdefault(a.column, profile.audit_type_casing.get(a.datatype, a.datatype))
            if not audit_descriptions.get(a.column):
                audit_descriptions[a.column] = a.description
        audit = tuple(seen.items())
        # post104: audit columns are commented from their own STTM rows
        audit_comments = {c: d for c, d in audit_descriptions.items()
                          if d and profile.column_comments}
        audit_blank = ([c for c, d in audit_descriptions.items() if not d]
                       if profile.column_comments else [])
        stage = TableBlock(
            layer="stage", table=group.stage,
            qualified=_qualify(group.stage, "stage", profile, spec, flags, config),
            business=tuple(stage_columns), audit=audit,
            comments=(*stage_comments.items(), *audit_comments.items()),
            table_comment=_load_strategy_comment(profile, spec.stage_load_strategy, "STG"))
        if stage_blank or audit_blank:
            flags.append(_comment_blank_flag("stage", stage.label, stage_blank + audit_blank))
        standard = None
        if group.standard is not None:
            standard = TableBlock(
                layer="standard", table=group.standard,
                qualified=_qualify(group.standard, "standard", profile, spec, flags, config),
                business=tuple(standard_columns), audit=audit,
                comments=(*standard_comments.items(), *audit_comments.items()),
                table_comment=_load_strategy_comment(profile, spec.standard_load_strategy,
                                                     "STD"))
            if standard_blank or audit_blank:
                flags.append(_comment_blank_flag("standard", standard.label,
                                                 standard_blank + audit_blank))
        out.append(TableDefinition(segments=tuple(s.segment for s in group.segments),
                                   stage=stage, standard=standard))
    return out


def _note_comment(f, layer: str, column: str, profile, comments: dict[str, str],
                  blank: list[str]) -> None:
    """Record one column's COMMENT (first row of a column wins, as the column
    list does) and, under ``column_comments``, a column without a description."""
    if column in comments or column in blank:
        return
    text = _column_comment(f, layer, profile)
    if text:
        comments[column] = text
    if profile.column_comments and not f.ddl_description:
        blank.append(column)


def _combined_ddl_text(spec: ResolvedFeedSpec, profile, flags: list[str] | None = None,
                       config: Config | None = None,
                       definitions: list[TableDefinition] | None = None) -> str | None:
    """The combined-layout deployment DDL (M4, ``acfc_prx``) — the CREATE
    TABLE REFERENCE text (the client's "DDL" is the metadata INSERT rows,
    ``metadata_inserts.sql``; MULTI_TABLE_DESIGN rule 8).

    Multi-table step 6: ONE CREATE block per table and layer, each named by
    its mapped three-part name and carrying its own columns from its own
    band (``table_definitions``): every table's stage block, then every
    table's standard block, in table order. Whitespace comes from the
    profile; a single-table feed is byte-identical to the pre-step-6 file
    (pair 1's golden). M7: every CREATE is three-part or omitted
    (``_qualify``); None when no block qualifies. ``definitions`` lets the
    caller pass the list it already computed (flags raised once)."""
    from codegen.emit.emitter import _environment

    if definitions is None:
        definitions = table_definitions(spec, profile, flags, config)
    blocks = []
    for layer, banner in (("stage", profile.stage_banner),
                          ("standard", profile.standard_banner)):
        for definition in definitions:
            block = definition.stage if layer == "stage" else definition.standard
            if block is None or not block.qualified:
                continue
            comments = dict(block.comments)
            blocks.append({
                "banner": banner, "qualified": block.qualified,
                "columns": [(name, dtype, _sql_comment(comments[name]) if name in comments
                             else None) for name, dtype in block.columns],
                "using_prefix": profile.block_using_prefix.get(layer, ""),
                "table_comment": (_sql_comment(block.table_comment) if block.table_comment
                                  else None),
                "tblproperties": (list((profile.standard_tblproperties or {}).items())
                                  if layer == "standard" else [])})
    if not blocks:
        return None
    template = _environment().get_template("framework/combined_ddl.txt.j2")
    return template.render(ddl=profile, blocks=blocks)


def _provenance_banner_rows(spec: ResolvedFeedSpec, faq: LoadPatternFaq,
                            config: Config, faq_sha: str) -> list[tuple[str, str]]:
    """The provenance banner as key/value rows — same inputs the generated
    files' banners carry, plus the FAQ and standards (inputs of this mode)."""
    faq_summary = summarize(faq)
    return [
        ("FRD contract", f"{spec.frd_contract_name} sha256 {spec.frd_contract_sha256}"),
        ("STTM contract", f"{spec.sttm_contract_name} sha256 {spec.sttm_contract_sha256}"),
        ("Load-pattern FAQ", f"sha256 {faq_sha} — {faq_summary['answered']} answered, "
                             f"{faq_summary['unknown']} unknown"),
        ("Engineering standards", config.engineering_standards.status),
        ("Layout", "client IIG template (anonymized reference — "
                   "fixtures/reference/SFMC_IIG.xlsx); values carry over"),
    ]


def _faq_sha(spec: ResolvedFeedSpec, config: Config, base_dir: Path | None) -> str:
    faq_dir = Path(config.load_pattern_faq.per_feed_dir)
    if base_dir is not None and not faq_dir.is_absolute():
        faq_dir = base_dir / faq_dir
    path = faq_dir / f"{spec.feed_slug}.faq.yaml"
    if not path.is_file():
        return "defaults (no FAQ file)"
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def _filter_payload_for_feed(payload: dict, slug: str) -> dict:
    """This feed's rows only, coverage recomputed for the filtered set."""
    tabs = {}
    derived = synthetic = needs_template = total = 0
    for name, tab in payload["tabs"].items():
        rows = [r for r in tab["rows"] if r.get("feed_slug") == slug]
        tabs[name] = {**tab, "rows": rows}
        for row in rows:
            for entry in row["badges"].values():
                total += 1
                badge = entry["badge"]
                if badge in DERIVED_BADGES:
                    derived += 1
                elif badge == "synthetic":
                    synthetic += 1
                else:
                    needs_template += 1
    return {
        **payload,
        "tabs": tabs,
        "coverage": {"derived": derived, "synthetic": synthetic,
                     "needs_template": needs_template, "total": total},
    }


# -- <slug>_stage/standard_table_creation.txt ---------------------------------- #
# Deployment-team DDL as two plain .txt files (run in the data lake by the
# deployment team), conformant to the client's reference goldens
# (fixtures/reference/SFMC_*_table_creation.txt). The .sql sources under
# out/<slug>/ddl/ stay untouched — these files are re-styled from them.


def _classify_ddl(file_name: str) -> tuple[str, str, str, str]:
    """(layer, schema, table, purpose) from the emitted DDL file name."""
    stem = file_name.removesuffix(".sql")
    if stem.endswith(".standard"):
        schema, _, table = stem.removesuffix(".standard").partition(".")
        return "standard", schema, table, "standard target table"
    schema, _, table = stem.partition(".")
    if table.lower().endswith("_errors"):
        return "stage", schema, table, "errors side-table"
    if table.lower().endswith("_processed_files"):
        return "stage", schema, table, "processed-files ledger"
    if table.lower().endswith("_recycle"):
        return "stage", schema, table, "recycle store"
    return "stage", schema, table, "stage target table"


# Section order inside each .txt: main target table first, side-tables after.
_DDL_PURPOSE_ORDER = {
    "stage target table": 0,
    "errors side-table": 1,
    "processed-files ledger": 2,
    "recycle store": 3,
    "standard target table": 0,
}

_CREATE_RE = re.compile(r"CREATE TABLE IF NOT EXISTS\s+(\S+)\s*\(")
_COLUMN_RE = re.compile(r"^\s*`(?P<name>[^`]+)`\s+(?P<dtype>[^,]+(?:\([^)]*\))?),?\s*$")


def _parse_ddl_statement(statement: str) -> tuple[str, list[tuple[str, str]]]:
    """(qualified table name, [(column, dtype)]) from an emitted .sql source."""
    match = _CREATE_RE.search(statement)
    if match is None:
        raise ValueError("DDL source has no CREATE TABLE statement")
    columns: list[tuple[str, str]] = []
    for line in statement.splitlines():
        col = _COLUMN_RE.match(line)
        if col:
            columns.append((col.group("name"), col.group("dtype").rstrip().rstrip(",")))
    return match.group(1), columns


def _sql_comment(text: str) -> str:
    # Single-line comment literals (the reference goldens' style): collapse
    # any embedded newlines/whitespace runs, escape quotes. post104: the
    # Databricks SQL escape is a backslash (O\'Connell), as ACFC's prod DDLs
    # write it; a doubled quote is no escape there (DBR 18 reads 'a''b' as two
    # chained literals, "ab"), so backslashes are escaped first, then quotes.
    return re.sub(r"\s+", " ", text).strip().replace("\\", "\\\\").replace("'", "\\'")


def _description_map(spec: ResolvedFeedSpec) -> dict[str, str]:
    """column name -> STTM description, for the per-column COMMENT clauses."""
    descriptions: dict[str, str] = {}
    for segment in spec.segments:
        for f in segment.fields:
            if f.description:
                descriptions.setdefault(f.stage_column, f.description)
                if f.standard_column:
                    descriptions.setdefault(f.standard_column, f.description)
    return descriptions


def _synthetic_location(spec: ResolvedFeedSpec, config: Config, table: str,
                        flags: list[str] | None = None) -> str:
    # M7: assembled from normalized segments (gate.derivations.join_path) —
    # never a raw string concatenation. First ACFC run: the FRD landing's own
    # punctuation (a trailing '.') is normalised and flagged path_normalised,
    # so the literal gate checks what the path IS, not the input's typo.
    raw = join_path(config.framework.synthetic_location_prefix,
                    spec.landing_location or spec.feed_slug, table)
    value, changes = normalise_path(raw)
    if changes and flags is not None:
        flags.append(path_normalised_flag(f"DDL LOCATION ({table})", raw, value, changes))
    return value


def _source_fragment(spec: ResolvedFeedSpec) -> str:
    """" (source: <vendor>)" for the table COMMENT — the vendor / source-
    system value only (Data Source / Vendor Metadata labels); omitted when
    unstated, multi-line or a label-prefixed description (M7 §2 h)."""
    source = (spec.source_system or "").strip()
    if not source or source.lower() == "unstated" or "\n" in source or "=" in source:
        return ""
    return f" (source: {source})"


def _table_creation_text(layer: str, entries: list[tuple[str, str, str]],
                         spec: ResolvedFeedSpec, config: Config,
                         banner: list[tuple[str, str]],
                         extra_banner_lines: list[str] | None = None,
                         profile=None, flags: list[str] | None = None) -> str:
    """Render one deployment-team .txt file for a layer's DDL entries.
    M7: under a profile requiring qualified names, an entry whose CREATE is
    not three-part is left out (flagged catalog_unstated:<layer>)."""
    from codegen.emit.emitter import _environment

    flags = flags if flags is not None else []
    descriptions = _description_map(spec)
    tables = []
    for _file_name, statement, purpose in sorted(
            entries, key=lambda e: (_DDL_PURPOSE_ORDER.get(e[2], 9), e[0])):
        qualified, columns = _parse_ddl_statement(statement)
        if profile is not None and len(qualified.split(".")) < 3:
            default = _config_default_catalog(config, profile, layer)
            if default:
                flags.append(f"catalog_from_config:{layer} — {qualified}: no FRD / STTM catalog; "
                             f"provenance config_default (conventions default_catalog[{layer}] "
                             f"= {default!r})")
                qualified = f"{default}.{qualified}"
            elif profile.require_qualified_names and _custom_env_without_catalog(config, layer):
                flags.append(_env_catalog_unset_flag(config, layer, qualified))
                continue
            elif profile.require_qualified_names:
                flags.append(f"catalog_unstated:{layer} — {qualified}: sources checked: FRD "
                             "Target Catalog and Schema label, STTM target band catalog "
                             f"column, conventions profile default_catalog[{layer}] — none "
                             f"states one; the {layer} DDL is not written (a two-part name is "
                             f"never written); answer `feeds[<n>].{layer}_target.catalog` "
                             "under gaps: (n = the feed's FRD index) and re-run layout, or set "
                             "the profile's default_catalog in a config overlay")
                continue
        rendered_columns = [
            {
                "name": name,
                "dtype": dtype,
                "comment": _sql_comment(
                    descriptions.get(name, name.replace("_", " ").title())),
            }
            for name, dtype in columns
        ]
        tables.append({
            "qualified": qualified,
            "columns": rendered_columns,
            "table_comment": _sql_comment(
                f"{purpose.capitalize()} for feed {spec.feed_name}"
                f"{_source_fragment(spec)}"),
            "location": (_synthetic_location(spec, config, qualified.rsplit(".", 1)[-1],
                                             flags)
                         if layer == "stage" else None),
            "tags": ((spec.domain, spec.sub_domain)
                     if spec.domain and spec.sub_domain else None),
        })
    banner_lines = [f"{key}: {value}" for key, value in banner
                    if key != "Layout"]
    banner_lines.append("DDL is engineer-run; the agent never creates target "
                        "tables (create_tables: false).")
    banner_lines.extend(extra_banner_lines or [])
    template = _environment().get_template("framework/table_creation.txt.j2")
    return template.render(layer=layer, tables=tables, banner_lines=banner_lines)


# -- config_inserts.sql -------------------------------------------------------- #


def _sql_value(value, dialect: str) -> str:
    if isinstance(value, bool):
        return ("1" if value else "0") if dialect == "sqlserver" else str(value).upper()
    if isinstance(value, (int, float)):
        return str(value)
    text = str(value)
    escaped = text.replace("'", "''")
    return f"N'{escaped}'" if dialect == "sqlserver" else f"'{escaped}'"


def _sql_identifier(name: str, dialect: str) -> str:
    if dialect == "sqlserver":
        return f"[{name}]"
    return '"' + name.replace('"', '""') + '"'


def _sql_table(name: str, dialect: str) -> str:
    return ".".join(_sql_identifier(part, dialect) for part in name.split("."))


def _insert_statement(headers: list[str], row: dict, table: str, dialect: str,
                      placeholder: str, always_blank: set[str]) -> str:
    values = []
    for header in headers:
        if header in always_blank:
            values.append(placeholder)
            continue
        value = row["values"][header]
        if value == "" or value is None:
            values.append("NULL")
        else:
            values.append(_sql_value(value, dialect))
    columns = ", ".join(_sql_identifier(h, dialect) for h in headers)
    return (f"INSERT INTO {_sql_table(table, dialect)} ({columns}) "
            f"VALUES ({', '.join(values)});")


def _config_inserts_workbook(payload: dict, spec: ResolvedFeedSpec,
                             config: Config, banner: list[tuple[str, str]]):
    """config_inserts.xlsx: one sheet per populated tab, value rows on the
    real IIG layout, the generated INSERT statement as the final column."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from codegen.metadata_sheet import _pin_workbook_properties

    dialect = config.framework.sql_dialect
    placeholder = config.framework.id_placeholder
    always_blank = set(config.demo.metadata_sheet.always_blank)

    workbook = Workbook()
    _pin_workbook_properties(workbook)
    workbook.remove(workbook.active)
    # Excel's hard per-cell limit is 32,767 characters; a wide feed's INSERT
    # (SRC_COLUMNS/TGT_COLUMN_NAMES lists) can exceed it, so oversize
    # statements continue in INSERT_STATEMENT_PART<n> columns — lossless,
    # never silently truncated.
    cell_limit = 32000
    for tab_name, tab in payload["tabs"].items():
        if not tab["rows"]:
            continue
        table = config.framework.tables.get(tab_name, tab_name)
        sheet = workbook.create_sheet(title=tab_name[:31])
        rendered = []
        max_parts = 1
        for row in tab["rows"]:
            display = [placeholder if h in always_blank else row["values"][h]
                       for h in tab["headers"]]
            statement = _insert_statement(tab["headers"], row, table, dialect,
                                          placeholder, always_blank)
            parts = [statement[i:i + cell_limit]
                     for i in range(0, len(statement), cell_limit)]
            max_parts = max(max_parts, len(parts))
            rendered.append((display, parts))
        headers = [*tab["headers"], "INSERT_STATEMENT",
                   *(f"INSERT_STATEMENT_PART{n}" for n in range(2, max_parts + 1))]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for display, parts in rendered:
            padded = parts + [""] * (max_parts - len(parts))
            sheet.append([*display, *padded])
        for index in range(1, len(tab["headers"]) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 24
        for index in range(len(tab["headers"]) + 1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 100

    provenance = workbook.create_sheet(title="_provenance")
    provenance.append(["input", "value"])
    for cell in provenance[1]:
        cell.font = Font(bold=True)
    for key, value in banner:
        provenance.append([key, value])
    provenance.append(["dialect", dialect])
    provenance.append(["note", "Run only after the config rows workbook is "
                               "approved; the client's ADF metadata path "
                               "inserts these rows. Plain INSERTs — "
                               "idempotency belongs to the framework's load "
                               "path."])
    return workbook


# -- the emit ------------------------------------------------------------------ #


def emit_framework(
    spec: ResolvedFeedSpec,
    faq: LoadPatternFaq,
    ddl_sources: list[tuple[str, str]],
    config: Config,
    out_root: Path,
    base_dir: Path | None = None,
    frd_path: Path | None = None,
    unmapped_rule_texts: set[str] | None = None,
    conventions_profile: str | None = None,
    iig_template: str | None = None,
) -> FrameworkArtefacts:
    """Render one feed's Option B artefacts into ``out/<slug>/framework/``.

    ``conventions_profile`` / ``iig_template`` (M4) select the DDL layout
    and the IIG workbook version; the defaults reproduce today's output."""
    framework_dir = out_root / spec.feed_slug / "framework"
    framework_dir.mkdir(parents=True, exist_ok=True)
    root = base_dir if base_dir is not None else Path(".")
    profile = config.conventions.get(conventions_profile)
    template_name, template_cfg = config.metadata.resolve(iig_template)
    flags: list[str] = []
    if template_cfg is not None and template_cfg.environment_cells:
        # 2026-10-14: which environment the run used (the report header names it)
        from codegen.environments import environment_flags

        flags += environment_flags(config)

    payload = metadata_sheet_payload(
        config, root,
        specs=[spec],
        unmapped_by_slug={spec.feed_slug: unmapped_rule_texts or set()},
        run_label=spec.feed_slug,
        faq_by_slug={spec.feed_slug: faq},
        frd_path=frd_path,
        template=template_name,
        conventions_profile=conventions_profile,
    )
    payload = _filter_payload_for_feed(payload, spec.feed_slug)
    # 2026-10-13 (pair_12): every IIG cell is ONE line before anything is written
    # from the payload (workbooks, metadata_inserts.sql); each change is a flag.
    flags.extend(normalise_iig_line_breaks(payload))
    # post57 (D1 post55 pair_01): whitespace in a stated path is kept byte for
    # byte; under path_whitespace: keep each path cell gets ONE flag (coordinates
    # + kind), under fail the derivations check fails it
    path_whitespace = getattr(profile, "path_whitespace", "keep")
    if path_whitespace == "keep":
        flags.extend(path_whitespace_flags(payload, config.gate.derivations))
    from codegen.metadata_template import blank_flags, shape_flags

    if template_cfg is not None:
        flags.extend(blank_flags(payload))
    # Path notes (iig_path_shape_on_uri / path_normalised) for every template.
    flags.extend(shape_flags(payload))
    # 2026-10-14 (framework owner's rule 3): an MFT-landed feed's FILE_ADLS sheet
    # is written with its headers only — said once
    from codegen.metadata_template import fixed_width_handler_empty_reason

    reasons: dict[str, str] = {}
    file_adls = payload["tabs"].get("FILE_ADLS_INGESTION_DETAILS")
    if template_cfg is not None and file_adls is not None and not file_adls["rows"]:
        from codegen.environments import apply_environment
        from codegen.metadata_sheet import _feed_from_spec
        from codegen.metadata_template import file_adls_mft_note

        note = file_adls_mft_note(_feed_from_spec(spec), config, spec,
                                  apply_environment(template_cfg, config))
        if note is not None:
            flags.append(note)
            reasons["FILE_ADLS_INGESTION_DETAILS"] = "MFT"
    # post66: a delimited file's fixed-width handler sheet is empty by rule too
    handler = payload["tabs"].get("ADLS_FIXED_WIDTH_HANDLER")
    if handler is not None and not handler["rows"]:
        reason = fixed_width_handler_empty_reason(spec, config)
        if reason is not None:
            reasons["ADLS_FIXED_WIDTH_HANDLER"] = reason
    empty_by_rule = {name: reasons[name] for name in payload["tabs"] if name in reasons}
    banner = _provenance_banner_rows(spec, faq, config,
                                     _faq_sha(spec, config, base_dir))

    files: list[Path] = []

    from codegen.metadata_sheet import stable_workbook_bytes

    # Deployment-team DDL: two plain .txt files (stage / standard), run in the
    # data lake by the deployment team; conformant to the reference goldens.
    stage_entries: list[tuple[str, str, str]] = []
    standard_entries: list[tuple[str, str, str]] = []
    omitted_side_tables: list[str] = []
    side_table_purposes = ("errors side-table", "processed-files ledger")
    for file_name, statement in ddl_sources:
        layer, _schema, table, purpose = _classify_ddl(file_name)
        if (not config.framework.deployment_ddl_include_side_tables
                and purpose in side_table_purposes):
            # The deployment .txt carries exactly the FRD's Target Table
            # Name list; the side-tables are CodeGen conventions and stay in
            # out/<slug>/ddl/ (and in notebook mode) untouched.
            omitted_side_tables.append(table)
            continue
        target = stage_entries if layer == "stage" else standard_entries
        target.append((file_name, statement, purpose))
    frd_tables = [s.stage_table.table for s in spec.segments]
    if spec.recycle is not None:
        frd_tables.append(spec.recycle.recycle_table.table)
    extra_banner_lines: list[str] = []
    if omitted_side_tables:
        extra_banner_lines.append(
            f"side-tables omitted from deployment DDL — not in FRD Target "
            f"Table Name (FRD lists: {', '.join(frd_tables)}); omitted: "
            f"{', '.join(sorted(omitted_side_tables))} (CodeGen conventions, "
            "kept in ddl/)")
    if profile.ddl_layout == "combined":
        abbrev_answer = getattr(faq, "feed_abbreviation", None)
        if abbrev_answer is not None and abbrev_answer.source != "unknown":
            abbrev = str(abbrev_answer.value)
        else:
            abbrev = _sanitize_abbrev(spec.feed_slug)
            flags.append(f"ddl_file_name_from_slug: no feed_abbreviation in the load-pattern "
                         f"FAQ; the combined DDL is named {abbrev!r} from the feed slug")
        ddl_names = [profile.ddl_file_name_pattern.format(feed_abbrev=abbrev)]
        definitions = table_definitions(spec, profile, flags, config)
        combined = _combined_ddl_text(spec, profile, flags, config, definitions=definitions)
        if combined is not None:
            if profile.target_system_header:
                combined = DDL_TARGET_HEADER + "\n" + combined
            txt_path = framework_dir / ddl_names[0]
            txt_path.write_text(combined, encoding="utf-8", newline="")
            files.append(txt_path)
        else:
            ddl_names = []
    else:
        ddl_names = []
        for layer, entries in (("stage", stage_entries), ("standard", standard_entries)):
            txt_path = framework_dir / f"{spec.feed_slug}_{layer}_table_creation.txt"
            text = _table_creation_text(layer, entries, spec, config, banner,
                                        extra_banner_lines=extra_banner_lines,
                                        profile=profile, flags=flags)
            if profile.target_system_header:
                text = DDL_TARGET_HEADER + "\n" + text
            txt_path.write_text(text, encoding="utf-8", newline="\n")
            files.append(txt_path)
            ddl_names.append(txt_path.name)
        definitions = None
    dml_disabled_reason = _dml_disabled_reason(config, profile, conventions_profile)
    # Multi-table step 6 (2026-10-08: on EVERY framework run): the metadata
    # INSERT rows that define the tables — the client's "DDL", rule 8 — from
    # the same payload as the IIG. The table index reuses the CREATE text's
    # definitions (flags raised once); another layout computes them for the
    # index only, its flags discarded (its own DDL path raised them).
    if definitions is None:
        definitions = table_definitions(spec, profile, [], config)
    metadata_inserts = emit_metadata_inserts(
        spec, payload, definitions, config, framework_dir, banner,
        ddl_name=" / ".join(ddl_names) or None, faq=faq)
    files.append(metadata_inserts.path)
    flags.extend(metadata_inserts.flags)
    # M7 §3, retired SQL 2026-10-08: behind the DML switches, only the runner
    # notebooks that execute metadata_inserts.sql over JDBC.
    dml_artefacts = None
    if dml_disabled_reason is None:
        from codegen.emit.dml import emit_dml

        dml_artefacts = emit_dml(config, framework_dir, metadata_inserts.path.name)
        files.extend(dml_artefacts.files)
        flags.extend(dml_artefacts.flags)
    # M7 gate checks: unstated catalog / schema is a FAIL, not a flag.
    checks: list[GateCheck] = []
    if dml_artefacts is not None:
        checks.extend(dml_artefacts.checks)
    checks.append(metadata_inserts.check)
    unqualified = [f for f in flags if f.startswith(("catalog_unstated:", "schema_unstated:"))]
    if unqualified:
        checks.append(GateCheck(name="qualified_names", passed=False,
                                details="; ".join(unqualified)))
    ddl_texts = [(p.name, p.read_text(encoding="utf-8")) for p in files if p.suffix == ".txt"]
    checks.append(check_sql_literals(ddl_texts, config, path_whitespace))
    checks.append(check_iig_derivations(payload, config, path_whitespace))

    rows_workbook = build_workbook(payload)
    inputs_sheet = rows_workbook.create_sheet(title="_inputs")
    inputs_sheet.append(["input", "value"])
    for key, value in banner:
        inputs_sheet.append([key, value])
    rows_path = framework_dir / "config_rows.xlsx"
    rows_path.write_bytes(stable_workbook_bytes(rows_workbook))
    files.append(rows_path)

    inserts_path = framework_dir / "config_inserts.xlsx"
    inserts_workbook = _config_inserts_workbook(payload, spec, config, banner)
    if dml_disabled_reason is None:
        # M7 §4: sheet 1 says this workbook is the REVIEW copy; the executable
        # script is metadata_inserts.sql next to it (config_inserts_<env>.sql
        # retired 2026-10-08).
        readme = inserts_workbook.create_sheet(title="README", index=0)
        readme.append([f"review copy — executable script is {metadata_inserts.path.name} "
                       "(fill every <<COLUMN#n>> placeholder first); run it from "
                       f"{config.dml.notebook_file_name_pattern.format(env='<env>')} "
                       f"(environments: {', '.join(config.dml.environments)}) against the SQL "
                       "Server metadata DB after the config rows are approved"])
        readme.column_dimensions["A"].width = 140
    inserts_path.write_bytes(stable_workbook_bytes(inserts_workbook))
    files.append(inserts_path)

    blank_list = (template_cfg.always_blank if template_cfg is not None
                  else config.demo.metadata_sheet.always_blank)
    always_blank = [
        header
        for tab in payload["tabs"].values()
        for header in tab["headers"]
        if header in set(blank_list)
    ]
    flagged = sorted(set(always_blank))
    # IIG-first M3: the BSA's review copy and the clean copy CI/CD loads —
    # behind the profile's emit_iig_review switch (off: output as before).
    iig_rows = ""
    review_name = clean_name = ""
    if profile.emit_iig_review:
        from codegen.iig_review import write_iig_workbooks

        review_path, clean_path, _open = write_iig_workbooks(
            payload, blank_list, config, framework_dir, spec.feed_slug, template_name,
            empty_by_rule)
        files.extend([review_path, clean_path])
        review_name, clean_name = review_path.name, clean_path.name
        iig_rows = _iig_addition_rows(review_name, clean_name, config)
    # A line break that slipped past the boundary FAILS with its cell address.
    checks.append(check_iig_workbooks_single_line(
        [rows_path] + ([clean_path] if profile.emit_iig_review else [])))
    # 2026-10-14 (D1 pair_12): a single-code column (LOB ...) never holds a list.
    checks.append(check_iig_workbooks_single_value(
        [rows_path] + ([clean_path] if profile.emit_iig_review else []),
        config.gate.derivations))
    # 2026-10-16: ADDITION.md describes what is presented — the DDL, the IIG
    # and its review copy — and the review flow (IIG -> BSA review -> CI/CD
    # applies the DDL); the other framework files are written, not described.
    ddl_list = " / ".join(f"`{name}`" for name in ddl_names) or "the DDL"
    ddl_row = (
        f"| {ddl_list} | The Databricks (Unity Catalog) table DDL, in the client's "
        "reference format. **Applied by the CI/CD pipeline after the IIG is approved** — "
        "the agent never creates target tables. |"
    )
    review_flow = _review_flow(ddl_list, review_name, clean_name, profile.emit_iig_review,
                               conventions_profile or config.conventions.profile)
    assumed_notes: list[str] = []
    held_back: list[str] = []
    seg = spec.segmented_extraction
    if seg is not None:
        ident = seg.identification
        assumed_notes.append(
            f"record identification ({ident.method}): trailer marker "
            f"{ident.trailer_marker!r}, header = {ident.header_rule} — "
            f"evidence: {ident.citation}")
        assumed_notes += [
            f"{entry.note} — evidence: {entry.citation}"
            for entry in seg.provenance_notes
        ]

    segmented_lines = [f"- **PROVENANCE**: {note}" for note in assumed_notes]
    segmented_lines += [f"- **HELD BACK**: {held}" for held in held_back]
    segmented_block = ("\n" + "\n".join(segmented_lines) + "\n"
                       if segmented_lines else "")
    addition_path = framework_dir / "ADDITION.md"
    addition_path.write_text(
        _ADDITION_TEMPLATE.format(
            slug=spec.feed_slug,
            banner="\n".join(f"- {k}: {v}" for k, v in banner),
            ddl_row=ddl_row,
            iig_rows=iig_rows,
            review_flow=review_flow,
            id_placeholder=config.framework.id_placeholder,
            blank_columns="\n".join(f"- `{c}`" for c in flagged) or "- none",
            segmented_block=segmented_block,
        ),
        encoding="utf-8", newline="\n",
    )
    files.append(addition_path)

    return FrameworkArtefacts(
        files=files,
        row_counts={name: len(tab["rows"]) for name, tab in payload["tabs"].items()},
        coverage=payload["coverage"],
        flagged_blank_columns=flagged,
        assumed_notes=assumed_notes,
        held_back=held_back,
        flags=flags,
        checks=checks,
        payload=payload,
        groups=artefact_groups(files),
        dml_disabled_reason=dml_disabled_reason,
        empty_by_rule=empty_by_rule,
    )


def _iig_addition_rows(review_name: str, clean_name: str, config: Config) -> str:
    """IIG-first: the ADDITION.md rows for the two IIG workbooks (only when
    the profile emits them)."""
    return (
        f"\n| `{clean_name}` | The IIG: the template's sheets and columns, values "
        "only — the file the BSA certifies and the CI/CD pipeline loads. |"
        f"\n| `{review_name}` | The BSA's review copy of the IIG: sheet "
        f"`{config.iig_review.summary_sheet}` first (one row per sheet / column / "
        "reason / owner group, and one row per sheet left empty by rule saying why), "
        "then the template sheets with every open cell "
        "highlighted by owner and commented with its reason, citation and owner. |"
    )


def _review_flow(ddl_list: str, review_name: str, clean_name: str, emits_iig: bool,
                 profile_name: str) -> str:
    """ADDITION.md's review flow: IIG -> BSA review -> CI/CD applies the DDL."""
    if emits_iig:
        first = (f"1. **IIG** — the agent writes `{clean_name}` and the BSA's review copy "
                 f"`{review_name}`: every cell it could not fill is open, highlighted and "
                 "commented with its reason and owner.")
        second = (f"2. **BSA review** — the BSA fills or corrects the open cells in "
                  f"`{review_name}` and certifies `{clean_name}`. The BSA-annotated IIG is "
                  "the approval artefact.")
    else:
        first = ("1. **IIG** — this profile writes no IIG workbook "
                 f"(`conventions.profiles.{profile_name}.emit_iig_review` is off).")
        second = ("2. **BSA review** — the BSA reviews and certifies the IIG; the "
                  "BSA-annotated IIG is the approval artefact.")
    third = (f"3. **CI/CD applies the DDL** — after approval, the CI/CD pipeline applies "
             f"{ddl_list} and loads the approved IIG. Nothing is applied before approval.")
    return "\n".join([first, second, third])


def _dml_disabled_reason(config: Config, profile, profile_name: str | None) -> str | None:
    """None when the runner notebooks are written; else which switch is off.
    ``metadata_inserts.sql`` is written regardless (2026-10-08)."""
    if not config.dml.enabled:
        return "dml.enabled is false"
    if not profile.emit_dml:
        name = profile_name or config.conventions.profile
        return f"conventions profile {name!r} has emit_dml: false"
    return None


def presented_files(files: list[Path]) -> list[Path]:
    """The framework files the App and the report present, in that order: the
    DDL, the IIG, its review copy and ADDITION.md (2026-10-16). The other
    framework files are still written; they are not listed."""
    tests = (
        lambda name: name.endswith(("_DDL.txt", "_table_creation.txt")),
        lambda name: name.endswith("_IIG.xlsx"),
        lambda name: name.endswith("_IIG_REVIEW.xlsx"),
        lambda name: name == "ADDITION.md",
    )
    return [path for test in tests for path in files if test(path.name)]


def report_section(artefacts: FrameworkArtefacts) -> str:
    """The "Framework artefacts" section appended to the generation report;
    2026-10-16 (post63): the same files and approval sentence as the App's
    artefacts panel."""
    coverage = artefacts.coverage
    lines = [
        "",
        "## Framework artefacts",
        "",
        "Under `framework/` — an addition to the existing ingestion framework.",
        "The BSA-annotated IIG is the approval artefact; the CI/CD pipeline",
        "applies the DDL after approval. The agent never applies anything itself.",
        "",
        "| Artefact | Detail |",
        "| --- | --- |",
    ]
    for path in presented_files(artefacts.files):
        lines.append(f"| `{path.name}` | framework/{path.name} |")
    for tab, count in artefacts.row_counts.items():
        lines.append(f"| IIG rows: {tab} | {count} |")
    lines.append(
        "| framework-assigned IDs | "
        + (", ".join(f"`{c}`" for c in artefacts.flagged_blank_columns) or "none")
        + " — left blank, never invented |"
    )
    for note in artefacts.assumed_notes:
        lines.append(f"| **PROVENANCE** | {note} |")
    for held in artefacts.held_back:
        lines.append(f"| **HELD BACK** | {held} |")
    # post66: the derived ratio lives here (the App's panel dropped it), labelled
    lines += ["", derived_line(coverage)]
    return "\n".join(lines) + "\n"


def derived_line(coverage: dict[str, int]) -> str:
    """The labelled derived ratio of the report's framework section."""
    return (f"cells derived from documents: {coverage['derived']:,} of "
            f"{coverage['total']:,}; the rest are profile defaults, framework-assigned "
            "blanks, or open inputs")


# Badge labels re-exported for the report/UI so the vocabulary stays single-
# sourced from metadata_sheet.
__all__ = ["FrameworkArtefacts", "emit_framework", "presented_files", "report_section",
           "BADGE_LABELS"]
