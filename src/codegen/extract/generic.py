"""Content-driven STTM extraction: read a workbook THROUGH a layout profile.

The profile (``codegen.layout``) says where things are; this module copies
what the cells say — verbatim — into the same STTM mapping contract the
resolver already consumes. Nothing here guesses: a role the profile could
not place reads as empty (and the profile's ``unresolved`` list says why),
a row that carries no field name is skipped and logged, a segment spelling
the vocabulary does not know is carried raw and refused loudly at contract
time, and every emitted field records the cell it came from.

Two stages:

- :func:`read_workbook` — profile + workbook → :class:`GenericIR` (meta
  facts, per-sheet field rows with provenance, audit rows, auxiliary
  sheets, skip log). No FRD needed; this is what the layout tests pin.
- :func:`build_generic_contract` — GenericIR + FRD contract → SttmContract
  (pairing by stage-table membership, then by file pattern, then a single
  feed to a single sheet — each recorded as a note).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field, replace
from pathlib import Path

from openpyxl.utils import get_column_letter

from codegen.config import Config, DiscoveryConfig
from codegen.contracts.frd import FrdContract, FrdFeed
from codegen.contracts.sttm import (
    AuditColumn,
    AuxiliarySheet,
    FieldProvenance,
    LayoutSummary,
    LoadRules,
    RecycleSpec,
    SourceFile,
    SttmContract,
    SttmFeed,
    SttmField,
    TableRef,
)
from codegen.formats import split_type_format
from codegen.layout.discover import (
    NONE_SEGMENT,
    Discovery,
    none_segment_spellings,
    normalize,
    text,
)
from codegen.layout.profile import LayoutProfile, Role, SheetProfile

_AUDIT_DATATYPES = {"string": "String", "timestamp": "Timestamp"}
_CANONICAL_SEGMENTS = ("Header", "Detail", "Trailer")
_WINDOW_DAYS_RE = re.compile(r"(\d+)\s*days?", re.IGNORECASE)
_RECYCLE_FLAG_RE = re.compile(r"\bY\s*\(")


@dataclass(frozen=True)
class FieldRow:
    sheet: str
    row: int
    col: int                          # the field-name cell's column (provenance)
    segment_raw: str | None
    segment: str | None               # canonical Header/Detail/Trailer, else None
    values: dict[str, str | None]     # "<layer>.<role>" -> cell text


@dataclass(frozen=True)
class AuditRow:
    sheet: str
    row: int
    segment: str | None
    column: str
    datatype_raw: str
    table: str | None
    # M11 item 13: the ONE layer the row names its column in, when only one.
    layer: str | None = None


@dataclass
class SheetData:
    profile: SheetProfile
    meta: dict[str, str]                       # resolved meta key -> value
    meta_raw: list[tuple[str, str | None]]     # every label:value row as seen
    fields: list[FieldRow] = field(default_factory=list)
    audit: list[AuditRow] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    # M11 item 13: rows the scope column marks out of scope, and that column.
    out_of_scope: list[int] = field(default_factory=list)
    scope_column: int | None = None

    def rows_per_segment(self) -> dict[str, int]:
        counts: dict[str, int] = {}
        for row in self.fields:
            key = row.segment or row.segment_raw or "-"
            counts[key] = counts.get(key, 0) + 1
        return counts


@dataclass
class GenericIR:
    workbook_name: str
    profile: LayoutProfile
    sheets: list[SheetData]
    auxiliary: list[AuxiliarySheet]
    version: str | None
    diagnostics: list[str]


def layout_summary(profile: LayoutProfile) -> LayoutSummary:
    return LayoutSummary(
        strategy=profile.strategy,
        source=profile.source,
        fingerprint=profile.fingerprint,
        unresolved=[f"{u.sheet}/{u.layer}/{u.role}: {u.reason}" for u in profile.unresolved],
    )


# ----------------------------------------------------------------- reading


def read_workbook(found: Discovery, name: str, config: Config) -> GenericIR:
    disc = config.extractor.discovery
    workbook = found.workbook
    profile = found.profile
    sheets: list[SheetData] = []
    auxiliary: list[AuxiliarySheet] = []
    version: str | None = None
    diagnostics = list(found.diagnostics)
    for sp in profile.sheets:
        ws = workbook[sp.name]
        if sp.kind == "mapping":
            sheets.append(_read_mapping_sheet(ws, sp, disc, config, diagnostics))
        elif sp.kind == "version":
            version = _last_version(ws, sp) or version
        elif sp.kind != "ignore":
            auxiliary.append(_read_auxiliary(ws, sp))
    return GenericIR(workbook_name=name, profile=profile, sheets=sheets, auxiliary=auxiliary,
                     version=version, diagnostics=diagnostics)


def _read_auxiliary(ws, sp: SheetProfile) -> AuxiliarySheet:
    header_row = sp.header_row or 1
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        cells = [text(c) for c in row[: len(sp.headers)]]
        if any(c is not None for c in cells):
            rows.append(cells + [None] * (len(sp.headers) - len(cells)))
    return AuxiliarySheet(kind=sp.kind, sheet=sp.name, headers=list(sp.headers), rows=rows)


def _last_version(ws, sp: SheetProfile) -> str | None:
    header_row = sp.header_row
    if header_row is None:
        return None
    headers = [normalize(c) for c in next(ws.iter_rows(min_row=header_row, max_row=header_row,
                                                        values_only=True))]
    if "version" not in headers:
        return None
    col = headers.index("version")
    versions = [text(r[col]) for r in ws.iter_rows(min_row=header_row + 1, values_only=True)
                if col < len(r) and text(r[col]) is not None]
    return versions[-1] if versions else None


def _typed(raw: str, column: str, sheet: str, row: int, flags: list[str]) -> str:
    """M11 item 13: a field type cell carrying a format -> its type, flagged."""
    base, fmt = split_type_format(raw)
    if fmt is not None:
        flags.append(f"type_format:{column} — the STTM declares {raw!r} (sheet {sheet!r} row "
                     f"{row}): type {base!r}, format {fmt!r}")
        return base
    return raw


def _row_segment(cells: list, sp: SheetProfile, current: str | None) -> str | None:
    if sp.segment_strategy == "column" and sp.segment_column is not None:
        return _cell(cells, sp.segment_column) or current
    return current


def _scope_column(ws, header_row: int, config: Config) -> int | None:
    """M11 item 13: the column whose data values are ALL In Scope / Out of
    scope markers (config ``extractor.scope_in_values`` / ``scope_out_values``)
    — found by value, since the real header embeds a vendor name. At least
    two marked rows, or it is not a scope column."""
    markers = {normalize(v) for v in (*config.extractor.scope_in_values,
                                      *config.extractor.scope_out_values)}
    if not markers:
        return None
    seen: dict[int, list[str]] = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        for col, value in enumerate(row, start=1):
            if value is not None and str(value).strip():
                seen.setdefault(col, []).append(normalize(value))
    candidates = [col for col, values in seen.items()
                  if len(values) >= 2 and all(v in markers for v in values)]
    return candidates[0] if len(candidates) == 1 else None


def _segment_lookup(disc: DiscoveryConfig) -> dict[str, str]:
    """Spelling -> canonical segment. The ``none`` class is left OUT: its
    spellings resolve to no segment (see ``_none_segments``)."""
    return {normalize(s): canonical for canonical, spellings in disc.segment_synonyms.items()
            if canonical != NONE_SEGMENT for s in spellings}


def _none_segments(sheet: SheetData, disc: DiscoveryConfig) -> list[str]:
    """M11: the spellings this sheet uses that mean "single record type"
    (config ``extractor.discovery.segment_synonyms.none``: NA, N/A, -). A
    sheet whose Segment column holds ONLY these is extracted unsegmented —
    the v0.6.2 ACFC run died on `segment spelling(s) ['NA']`."""
    none_values = none_segment_spellings(disc)
    used = {r.segment_raw for r in sheet.fields if r.segment_raw is not None}
    if not used or not all(normalize(v) in none_values for v in used):
        return []
    return sorted(used)


def _read_mapping_sheet(ws, sp: SheetProfile, disc: DiscoveryConfig, config: Config,
                        diagnostics: list[str]) -> SheetData:
    header_row = sp.header_row or 1
    source = sp.band("source")
    stage = sp.band("stage")
    field_col = source.column(Role.FIELD_NAME) if source is not None else None
    stage_column_col = stage.column(Role.COLUMN) if stage is not None else None
    stage_table_col = stage.column(Role.TABLE) if stage is not None else None
    stage_type_col = stage.column(Role.TARGET_TYPE) if stage is not None else None
    segments = _segment_lookup(disc)
    audit_markers = {normalize(m) for m in config.extractor.audit_source_markers}

    meta: dict[str, str] = {}
    meta_raw: list[tuple[str, str | None]] = []
    blank = {normalize(b) for b in disc.meta_blank_values}
    for entry in sp.meta_rows:
        value = None
        if entry.value_col is not None:
            value = text(ws.cell(row=entry.row, column=entry.value_col).value)
        meta_raw.append((entry.label, value))
        if value is not None and normalize(value) in blank:
            # "TBD": a stated blank, not a value — the chain looks further.
            diagnostics.append(f"sheet {sp.name!r} row {entry.row}: meta {entry.label!r} reads "
                               f"{value!r} (a placeholder); read as blank")
            continue
        if entry.key is not None and value is not None and entry.key not in meta:
            meta[entry.key] = value

    data = SheetData(profile=sp, meta=meta, meta_raw=meta_raw)
    data.scope_column = _scope_column(ws, header_row, config)
    scope_out = {normalize(v) for v in config.extractor.scope_out_values}
    audit_rule_markers = [normalize(m) for m in config.extractor.audit_load_rule_markers]
    rules_band = next((b for b in sp.bands if b.layer == "rules"
                       and b.column(Role.LOAD_RULE) is not None), None)
    load_rule_col = (rules_band.column(Role.LOAD_RULE) if rules_band is not None
                     else source.column(Role.LOAD_RULE) if source is not None else None)
    standard = sp.band("standard")
    standard_column_col = standard.column(Role.COLUMN) if standard is not None else None
    standard_type_col = standard.column(Role.TARGET_TYPE) if standard is not None else None
    current_segment_raw: str | None = None
    if sp.segment_strategy == "banner" and sp.band_row is not None and source is not None:
        current_segment_raw = text(ws.cell(row=sp.band_row, column=source.col_start).value)
    elif sp.segment_strategy == "sheet":
        current_segment_raw = sp.name

    for row_number, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                     start=header_row + 1):
        cells = [text(c) for c in row]
        if all(c is None for c in cells):
            continue
        filled = [(i + 1, c) for i, c in enumerate(cells) if c is not None]
        # In-sheet segment banner ("Detail Record"): one text cell inside the
        # source band, nothing on the target side.
        if (sp.segment_strategy == "banner" and len(filled) == 1 and source is not None
                and source.col_start <= filled[0][0] <= source.col_end):
            current_segment_raw = filled[0][1]
            data.skipped.append(f"row {row_number}: segment banner {filled[0][1]!r}")
            continue
        # M9.0: a sheet with a Segment COLUMN may still introduce each record
        # block with a banner row ("Details"). The column states the segment;
        # the banner is structure, logged as such (not as a nameless field).
        if (sp.segment_strategy == "column" and len(filled) == 1 and source is not None
                and source.col_start <= filled[0][0] <= source.col_end
                and normalize(filled[0][1]) in segments):
            data.skipped.append(f"row {row_number}: segment banner {filled[0][1]!r} "
                                "(the Segment column states the segment)")
            continue
        if data.scope_column is not None and normalize(
                _cell(cells, data.scope_column)) in scope_out:
            data.out_of_scope.append(row_number)
            continue
        field_name = _cell(cells, field_col)
        stage_column = _cell(cells, stage_column_col)
        source_values = [c for i, c in filled if source is not None
                         and source.col_start <= i <= source.col_end]
        load_rule = normalize(_cell(cells, load_rule_col))
        if audit_rule_markers and any(load_rule.startswith(m) for m in audit_rule_markers):
            # M11 item 13: the Load Rules cell says AUDIT — whatever the source
            # side reads, and even when only ONE layer names the column.
            std_column = _cell(cells, standard_column_col)
            column = stage_column or std_column
            if column is None:
                data.skipped.append(f"row {row_number}: audit row names no column; skipped")
                continue
            layer = None
            if stage_column is None:
                layer = "standard"
            elif standard is not None and std_column is None:
                layer = "stage"
            data.audit.append(AuditRow(
                sheet=sp.name, row=row_number,
                segment=segments.get(normalize(segment_raw)) if (
                    segment_raw := _row_segment(cells, sp, current_segment_raw)) else None,
                column=column,
                datatype_raw=(_cell(cells, stage_type_col) if stage_column is not None
                              else _cell(cells, standard_type_col)) or "",
                table=_cell(cells, stage_table_col), layer=layer))
            continue
        if sp.segment_strategy == "column" and sp.segment_column is not None:
            # An audit row leaves its Segment cell empty: it belongs to the
            # segment block it sits in (the last stated segment).
            stated = _cell(cells, sp.segment_column)
            if stated is not None:
                current_segment_raw = stated
            segment_raw = stated if stated is not None else current_segment_raw
        else:
            segment_raw = current_segment_raw
        is_audit = (field_name is not None and normalize(field_name) in audit_markers) or (
            field_name is None and stage_column is not None
            and all(normalize(v) in audit_markers for v in source_values))
        if is_audit:
            if stage_column is None:
                data.skipped.append(f"row {row_number}: audit marker without a stage column; "
                                    "skipped")
                continue
            data.audit.append(AuditRow(
                sheet=sp.name, row=row_number,
                segment=segments.get(normalize(segment_raw)) if segment_raw else None,
                column=stage_column, datatype_raw=_cell(cells, stage_type_col) or "",
                table=_cell(cells, stage_table_col)))
            continue
        if field_name is None:
            reason = ("field_name role unresolved" if field_col is None
                      else "no field name in the source band")
            data.skipped.append(f"row {row_number}: {reason}; skipped")
            continue
        values: dict[str, str | None] = {}
        for band in sp.bands:
            for role, col in band.roles.items():
                values[f"{band.layer}.{role}"] = _cell(cells, col)
        data.fields.append(FieldRow(
            sheet=sp.name, row=row_number, col=field_col or 0,
            segment_raw=segment_raw,
            segment=segments.get(normalize(segment_raw)) if segment_raw else None,
            values=values))
    for entry in data.skipped:
        diagnostics.append(f"sheet {sp.name!r}: {entry}")
    return data


def _cell(cells: list[str | None], col: int | None) -> str | None:
    if col is None or col < 1 or col > len(cells):
        return None
    return cells[col - 1]


# --------------------------------------------------------------- contract


class GenericExtractionError(ValueError):
    """The generic IR cannot be paired/emitted as a contract; names the spot."""


def _yes(value: str | None, disc: DiscoveryConfig) -> bool:
    return normalize(value) in {normalize(v) for v in disc.yes_values}


def _first(row: FieldRow, *keys: str) -> str | None:
    for key in keys:
        value = row.values.get(key)
        if value is not None:
            return value
    return None


def _dominant(values: list[str | None]) -> str | None:
    counts: dict[str, int] = {}
    for value in values:
        if value is not None:
            counts[value] = counts.get(value, 0) + 1
    return max(counts, key=counts.get) if counts else None


def _band_constant(sheet: SheetData, key: str) -> tuple[str | None, str]:
    """(value, provenance) of a band-level constant — schema, table, catalog.

    A band constant is stated once per sheet: repeated on every data row, or
    once in a merged cell whose anchor is the only non-empty cell (openpyxl
    reads the rest of a merged range as None). Either way it is the FIRST
    non-empty cell of the column, and the provenance names that cell. A
    column that holds SEVERAL distinct values is not a constant (one table
    per segment, pair 8): the dominant value stands for the sheet, as before,
    and the provenance says so."""
    layer, role = key.split(".")
    band = sheet.profile.band(layer)  # type: ignore[arg-type]
    col = band.column(role) if band is not None else None
    stated = [(r.row, r.values.get(key)) for r in sheet.fields if r.values.get(key) is not None]
    if col is None:
        return None, f"{layer} {role} role is not placed on sheet {sheet.profile.name!r}"
    if not stated:
        return None, (f"{sheet.profile.name}!{get_column_letter(col)}: every data row of the "
                      f"{layer} {role} column is empty")
    distinct = list(dict.fromkeys(v for _r, v in stated))
    if len(distinct) == 1:
        row, value = stated[0]
        return value, (f"{sheet.profile.name}!{get_column_letter(col)}{row} (first non-empty "
                       f"cell of the {layer} {role} column; {len(stated)} row(s) state it)")
    value = _dominant([v for _r, v in stated])
    row = next(r for r, v in stated if v == value)
    return value, (f"{sheet.profile.name}!{get_column_letter(col)}{row} (dominant of "
                   f"{len(distinct)} values in the {layer} {role} column)")


def _schema_chain(layer: str, band_value: str | None, band_provenance: str, frd_value: str | None,
                  config: Config, sheet: str, notes: list[str], flags: list[str]) -> str | None:
    """STTM target band -> FRD -> conventions.default_schema[layer]; every
    link after the first is a provenance note AND a flag. None = no source."""
    if band_value is not None:
        notes.append(f"sheet {sheet!r}: {layer} schema {band_value!r} — {band_provenance}")
        return band_value
    default = config.conventions.default_schema.get(layer)
    for value, source in ((frd_value, "FRD 'Target Catalog and Schema'"),
                          (default, f"config_default (conventions.default_schema[{layer}])")):
        if value:
            flags.append(f"sttm_unstated:{layer}.schema source_used:{source}: {value!r} — "
                         f"{band_provenance}")
            notes.append(f"sheet {sheet!r}: {layer} schema not stated in the STTM "
                         f"({band_provenance}); taken from {source}: {value!r}")
            return value
    return None


def _unmapped_markers(config: Config) -> set[str]:
    return {normalize(m) for m in config.extractor.unmapped_markers}


def _canonical_file_name(name: str) -> str:
    return name.strip().lower().replace("ccyy", "yyyy")


def _match_feed(sheet: SheetData, stage_table: str | None, frd: FrdContract,
                ir: GenericIR, notes: list[str]) -> FrdFeed:
    if stage_table is not None:
        matches = [f for f in frd.feeds if stage_table in f.stage_target.tables]
        if len(matches) == 1:
            return matches[0]
    patterns = _sheet_file_patterns(sheet, ir)
    canonical = {_canonical_file_name(p) for p in patterns}
    matches = [f for f in frd.feeds
               if canonical & {_canonical_file_name(p) for p in f.file_name_patterns}]
    if len(matches) == 1:
        notes.append(f"sheet {sheet.profile.name!r} paired to FRD feed "
                     f"{matches[0].feed_name!r} by file pattern")
        return matches[0]
    if len(frd.feeds) == 1 and len(ir.sheets) == 1:
        notes.append(f"sheet {sheet.profile.name!r} paired to the FRD's only feed "
                     f"{frd.feeds[0].feed_name!r} (single sheet, single feed)")
        return frd.feeds[0]
    raise GenericExtractionError(
        f"sheet {sheet.profile.name!r} (stage table {stage_table!r}, file patterns "
        f"{sorted(patterns)}) matches {len(matches)} FRD feeds "
        f"{[f.feed_name for f in frd.feeds]}; expected exactly one by stage table or "
        "file pattern"
    )


def _sheet_file_patterns(sheet: SheetData, ir: GenericIR) -> list[str]:
    patterns: list[str] = []
    for key in ("file_names", "file_name_example"):
        value = sheet.meta.get(key)
        if value:
            patterns += [p.strip() for p in re.split(r"[\n;,]+", value) if p.strip()]
    for aux in ir.auxiliary:
        if aux.kind != "file_details":
            continue
        headers = [normalize(h) for h in aux.headers]
        col = next((i for i, h in enumerate(headers) if h in ("filename", "file name",
                                                              "inbound file name")), None)
        if col is not None:
            patterns += [r[col] for r in aux.rows if col < len(r) and r[col]]
    return patterns


def _rule_columns(fields, segmented: bool, predicate) -> list[str]:
    """Source columns satisfying ``predicate``. Flat sheets: every field (as
    before). Segmented sheets (M4): the DETAIL segment's fields only, deduped
    — header/trailer fields are not row keys of the detail table (the
    segmented extractor does the same)."""
    if not segmented:
        return [f.source_column for f in fields if predicate(f)]
    detail = [f for f in fields if f.record_segment == "Detail"] or fields
    return list(dict.fromkeys(f.source_column for f in detail if predicate(f)))


def _file_details_row(feed: FrdFeed, ir: GenericIR) -> dict[str, str] | None:
    canonical = {_canonical_file_name(p) for p in feed.file_name_patterns}
    for aux in ir.auxiliary:
        if aux.kind != "file_details":
            continue
        headers = [normalize(h) for h in aux.headers]
        for row in aux.rows:
            entry = {h: v for h, v in zip(headers, row, strict=False) if v}
            name = entry.get("filename") or entry.get("file name") or entry.get("inbound file name")
            if name and _canonical_file_name(name) in canonical:
                return entry
    return None


def _recycle(sheet: SheetData, config: Config, disc: DiscoveryConfig) -> RecycleSpec | None:
    flagged = [(row, row.values.get("rules.recycle_flag")) for row in sheet.fields
               if row.values.get("rules.recycle_flag")]
    if not flagged:
        return None
    row, raw = flagged[0]
    match = _RECYCLE_FLAG_RE.search(raw or "")
    if match is None:
        return None
    validation = raw[match.end():].strip().rstrip(")").rstrip()
    if not validation:
        return None
    window = _WINDOW_DAYS_RE.search(raw)
    days = int(window.group(1)) if window else config.defaults.recycle_window_days
    return RecycleSpec(
        applies_to=_first(row, "source.field_name") or "",
        enabled=True,
        validation=validation,
        on_match=config.extractor.recycle_on_match.format(window_days=days),
        on_no_match=config.extractor.recycle_on_no_match.format(window_days=days),
        recycle_window_days=days,
    )


def build_generic_contract(ir: GenericIR, frd: FrdContract, config: Config, *,
                           contract_name: str | None = None,
                           generated_date: str,
                           width_answers: dict[int, dict[str, int]] | None = None
                           ) -> SttmContract:
    disc = config.extractor.discovery
    notes: list[str] = []
    feeds: list[SttmFeed] = []
    for sheet in ir.sheets:
        feeds.append(_build_feed(sheet, ir, frd, config, disc, notes, width_answers or {}))
    return SttmContract(
        contract_name=contract_name or f"STTM mapping contract extracted from {ir.workbook_name}",
        generated_from_workbook=ir.workbook_name,
        sttm_version=ir.version or "unversioned (no version sheet recognised)",
        generated_date=generated_date,
        notes=[
            "Extracted deterministically by `codegen extract-sttm` (content-driven layout, "
            f"strategy {ir.profile.strategy!r}, source {ir.profile.source!r}) from the "
            f"workbook named above, paired with FRD contract '{frd.contract_name}'.",
            "Column facts are verbatim workbook values read through the layout profile; "
            "every field records the cell it came from.",
            *notes,
        ],
        feeds=feeds,
        auxiliary_sheets=ir.auxiliary,
        layout=layout_summary(ir.profile),
    )


def _width_facts(row: FieldRow, field_name: str, sp: SheetProfile, flags: list[str]
                 ) -> tuple[int | None, str | None]:
    """(source_width, source_precision) for one field row — the STTM links of
    the width chain (codegen.resolve.widths), flagged with their cells. The
    width is returned only when it is NOT the plain integer length."""
    from codegen.resolve.widths import as_integer, as_precision, sttm_width

    length = _first(row, "source.length", "source.field_length")
    start, end = _first(row, "source.start"), _first(row, "source.end")
    if as_integer(start) is None or as_integer(length) is not None:
        return None, None            # not positional, or the length IS the width
    source = sp.band("source")
    label = " ".join(field_name.split())

    def cell(*roles: str) -> str:
        col = next((source.column(r) for r in roles if source and source.column(r)), None)
        return f"{sp.name}!{get_column_letter(col)}{row.row}" if col else f"{sp.name} row {row.row}"

    precision = as_precision(length)
    if precision is not None:
        flags.append(f"length_is_precision:{label} — STTM {cell('length', 'field_length')} reads "
                     f"{length!r}: a precision, kept as the source type precision "
                     f"{precision}; it is not a byte width and none is derived from it")
    elif length is not None:
        flags.append(f"length_not_a_width:{label} — STTM {cell('length', 'field_length')} reads "
                     f"{length!r}, which is not an integer byte width")
    width, link = sttm_width(length, start, end)
    if link == "sttm_span":
        flags.append(f"width_from_sttm_span:{label} — STTM {cell('start')} / {cell('end')}: "
                     f"end {end} - start {start} + 1 = {width}")
        return width, precision
    return None, precision


def _build_feed(sheet: SheetData, ir: GenericIR, frd: FrdContract, config: Config,
                disc: DiscoveryConfig, notes: list[str],
                width_answers: dict[int, dict[str, int]] | None = None) -> SttmFeed:
    from codegen.resolve.resolver import _FORMAT_DELIMITERS, normalize_feed_name
    from codegen.resolve.widths import normalize_field_name

    sp = sheet.profile
    name = sp.name
    if not sheet.fields:
        raise GenericExtractionError(
            f"sheet {name!r}: no field rows could be read "
            f"(skipped: {sheet.skipped[:3]}{'…' if len(sheet.skipped) > 3 else ''}; "
            f"unresolved: {[u.role for u in ir.profile.unresolved_for(name)]})")
    missing = [role for role in ("table", "column", "target_type")
               if all(r.values.get(f"stage.{role}") is None for r in sheet.fields)]
    if missing:
        raise GenericExtractionError(
            f"sheet {name!r}: the stage band has no values for {missing} (roles unresolved or "
            "empty) — a contract needs stage table, column and data type per field")
    flags: list[str] = []
    stage_table, table_provenance = _band_constant(sheet, "stage.table")
    stage_catalog, _ = _band_constant(sheet, "stage.catalog")
    if stage_table is None:
        raise GenericExtractionError(f"sheet {name!r}: stage table cells are empty — "
                                     f"{table_provenance}")
    notes.append(f"sheet {name!r}: stage table {stage_table!r} — {table_provenance}")
    feed = _match_feed(sheet, stage_table, frd, ir, notes)
    feed_id = config.feed_aliases.get(feed.feed_name) or normalize_feed_name(feed.feed_name)
    # The schema chain (M9.2): STTM band -> FRD -> config default, then the
    # hard stop — an empty schema never reaches a contract.
    band_schema, schema_provenance = _band_constant(sheet, "stage.schema")
    stage_schema = _schema_chain("stage", band_schema, schema_provenance,
                                 feed.stage_target.schema_name, config, name, notes, flags)
    if stage_schema is None:
        raise GenericExtractionError(
            f"sheet {name!r}: no source states the stage schema — {schema_provenance}; the FRD "
            "'Target Catalog and Schema' is blank and conventions.default_schema.stage is unset. "
            "Place the schema role (`codegen layout --answers`), or state it in the FRD / config")

    has_standard = any(
        r.values.get("standard.column") for r in sheet.fields
        if normalize(r.values.get("standard.column")) not in _unmapped_markers(config))
    standard_table, _ = _band_constant(sheet, "standard.table")
    standard_catalog, _ = _band_constant(sheet, "standard.catalog")
    standard_schema = None
    if has_standard:
        band_schema, schema_provenance = _band_constant(sheet, "standard.schema")
        standard_schema = _schema_chain("standard", band_schema, schema_provenance,
                                        feed.standard_target.schema_name, config, name, notes,
                                        flags)

    segmented = any(r.segment_raw is not None for r in sheet.fields)
    none_spellings = _none_segments(sheet, config.extractor.discovery)
    if none_spellings:
        # Every Segment cell says "not segmented": one record type, extracted
        # flat, and the sheet's own spelling is on the record.
        segmented = False
        flags.append(f"segments_none:{name}: the Segment column reads "
                     f"{', '.join(repr(v) for v in none_spellings)} — a single record type "
                     "(extractor.discovery.segment_synonyms.none); extracted unsegmented")
        notes.append(f"sheet {name!r}: segment column reads "
                     f"{', '.join(repr(v) for v in none_spellings)} — single record type")
    elif segmented:
        unknown = sorted({r.segment_raw for r in sheet.fields if r.segment is None} - {None})
        if unknown:
            raise GenericExtractionError(
                f"sheet {name!r}: segment spelling(s) {unknown} are outside the "
                f"{'/'.join(_CANONICAL_SEGMENTS)} vocabulary (extractor.discovery."
                "segment_synonyms) and none of them means 'single record type' "
                "(the `none` class: "
                f"{sorted(none_segment_spellings(config.extractor.discovery)) or 'unset'}); "
                "the contract dialect cannot carry them")

    fields: list[SttmField] = []
    markers = _unmapped_markers(config)
    stage_band = sp.band("stage")
    stage_column_col = stage_band.column("column") if stage_band is not None else None
    for row in sheet.fields:
        field_name = _first(row, "source.field_name")
        assert field_name is not None
        stage_column = _first(row, "stage.column")
        stage_type = _first(row, "stage.target_type")
        if stage_column is not None and normalize(stage_column) in markers:
            # The STTM says this source field is not mapped: left out of both
            # layers, flagged with the cell that says so (never a note only).
            cell = f"{name}!{get_column_letter(stage_column_col or 1)}{row.row}"
            label = " ".join(field_name.split())
            flags.append(f"field_unmapped:{label} — STTM {cell} reads {stage_column!r}; the "
                         "field is in neither the stage nor the standard table")
            notes.append(f"sheet {name!r} row {row.row}: field {label!r} skipped — stage column "
                         f"cell {cell} reads {stage_column!r}")
            continue
        if stage_column is None or stage_type is None:
            raise GenericExtractionError(
                f"sheet {name!r} row {row.row}: field {field_name!r} has no stage "
                f"column/data type (column={stage_column!r}, type={stage_type!r})")
        stage_type = _typed(stage_type, stage_column, name, row.row, flags)
        # A NULLABLE column says "NULL"/"nullable"/"yes" for nullable and
        # "NOT NULL"/"no" for required — the null-ish spellings of no_values
        # (meant for yes/no columns) do not apply to it.
        nullable_text = normalize(_first(row, "source.nullable"))
        no_words = {normalize(v) for v in disc.no_values if "null" not in normalize(v)}
        not_null = (_yes(_first(row, "rules.not_null", "source.not_null"), disc)
                    or normalize(_first(row, "source.null_check")).startswith("not null")
                    or nullable_text.startswith("not null")
                    or (bool(nullable_text) and nullable_text in no_words))
        mandatory = _yes(_first(row, "source.required", "source.mandatory",
                                "rules.required", "stage.mandatory_column"), disc)
        pii = _yes(_first(row, "source.pii", "rules.pii"), disc)
        rule_parts = [v for v in (_first(row, "rules.load_rule", "source.load_rule"),
                                  _first(row, "stage.transformation"),
                                  _first(row, "source.business_rule")) if v]
        value_spec = "; ".join(rule_parts) if rule_parts else _first(row, "source.comments")
        source_width, source_precision = _width_facts(row, field_name, sp, flags)
        answered = (width_answers or {}).get(frd.feeds.index(feed), {}).get(
            normalize_field_name(field_name))
        fields.append(SttmField(
            source_column=field_name,
            description=_first(row, "source.description", "rules.data_definition",
                               "source.comments", "stage.field_description"),
            sample_value=_first(row, "source.sample_value", "source.example_value"),
            source_datatype=_first(row, "source.source_type") or "unstated",
            nullable=not (not_null or mandatory),
            phi=pii,
            mandatory=mandatory,
            stage_column=stage_column,
            stage_datatype=stage_type,
            standard_column=_first(row, "standard.column") if has_standard else None,
            standard_datatype=(_typed(_first(row, "standard.target_type"),
                                      _first(row, "standard.column") or stage_column, name,
                                      row.row, flags)
                               if has_standard and _first(row, "standard.target_type") else None),
            value_spec=value_spec,
            source_length=_first(row, "source.length", "source.field_length"),
            source_start=_first(row, "source.start"),
            source_end=_first(row, "source.end"),
            source_width=source_width,
            source_precision=source_precision,
            width_answer=answered,
            record_segment=row.segment if segmented else None,  # type: ignore[arg-type]
            record_segment_label=row.segment_raw if segmented else None,
            stage_table=(_first(row, "stage.table") or stage_table) if segmented else None,
            standard_table=(_first(row, "standard.table") if segmented and has_standard
                            else None),
            provenance=FieldProvenance(
                sheet=row.sheet, row=row.row, col=row.col,
                source=ir.profile.role_source(row.sheet, "source", "field_name")),
        ))

    audit: list[AuditColumn] = []
    seen: set[str] = set()
    for entry in sheet.audit:
        base, fmt = split_type_format(entry.datatype_raw)
        if fmt is not None:
            flags.append(f"type_format:{entry.column} — the STTM declares "
                         f"{entry.datatype_raw!r} (sheet {name!r} row {entry.row}): type "
                         f"{base!r}, format {fmt!r}")
            entry = replace(entry, datatype_raw=base)
        if entry.layer is not None:
            flags.append(f"audit_column_one_layer:{entry.column} — the STTM names it in the "
                         f"{entry.layer} layer only (sheet {name!r} row {entry.row}, Load Rules "
                         "marks it an audit column); carried as a feed audit column")
        datatype = _AUDIT_DATATYPES.get(normalize(entry.datatype_raw))
        if datatype is None:
            # M11: a type outside String / Timestamp is the client's word,
            # carried verbatim and flagged — not a reason to refuse the STTM.
            datatype = (entry.datatype_raw or "").strip()
            if not datatype:
                raise GenericExtractionError(
                    f"sheet {name!r} row {entry.row}: audit column {entry.column!r} states no "
                    "datatype (the cell is empty)")
            flags.append(f"audit_type_nonstandard:{entry.column} — the STTM declares "
                         f"{datatype!r} (sheet {name!r} row {entry.row}); "
                         f"{sorted(_AUDIT_DATATYPES.values())} are the standard audit types")
        if entry.column not in seen:
            seen.add(entry.column)
            audit.append(AuditColumn(column=entry.column, datatype=datatype))  # type: ignore[arg-type]
    if not audit:
        raise GenericExtractionError(
            f"sheet {name!r}: no audit rows (rows whose source-side field name is an audit "
            f"marker {config.extractor.audit_source_markers} with a stage column) — the "
            "contract requires at least one audit column; add them to the STTM")

    if sheet.out_of_scope:
        rows = sheet.out_of_scope
        span = (f"rows {rows[0]}-{rows[-1]}" if len(rows) > 3 else
                "rows " + ", ".join(str(r) for r in rows))
        flags.append(f"rows_out_of_scope:{name} — {len(rows)} row(s) marked out of scope in "
                     f"column {get_column_letter(sheet.scope_column or 1)} ({span}); skipped")
    source_kind = None
    source_band = sp.band("source")
    if source_band is not None:
        rdbms_roles = [r for r in ("server", "source_database", "source_schema", "source_table")
                       if source_band.column(r) is not None]
        if len(rdbms_roles) >= 2:
            source_kind = "rdbms"
            flags.append(f"source_kind_rdbms:{name} — the STTM's source band describes a "
                         f"database table ({', '.join(rdbms_roles)} columns), not a file: the "
                         "DDL is generated; the DML writes a REVIEW block for the file-source "
                         "tables; the generated notebook pipeline reads FILES and does not "
                         "apply")
    details = _file_details_row(feed, ir)
    patterns = _sheet_file_patterns(sheet, ir)
    frd_canonical = {_canonical_file_name(p): p for p in feed.file_name_patterns}
    name_pattern = next((p for p in patterns if _canonical_file_name(p) in frd_canonical), None)
    if name_pattern is None and not feed.file_name_patterns and patterns:
        # M4: a docx-extracted FRD may name no file pattern; the workbook's
        # meta rows / FILE_DETAILS then supply it (the resolver flags the
        # provenance: file_pattern_from_sttm).
        name_pattern = patterns[0]
        notes.append(f"feed {feed.feed_name!r}: the FRD names no file pattern; taken from the "
                     f"workbook ({name_pattern!r})")
    if name_pattern is None and feed.file_name_patterns:
        name_pattern = feed.file_name_patterns[0]
        notes.append(f"feed {feed.feed_name!r}: file pattern taken from the FRD "
                     f"({name_pattern!r}); the workbook names none that matches")
    elif name_pattern is None:
        # M9.1b: not a hard stop any more. The chain FRD -> STTM meta rows /
        # File Details -> VDD FILES sheet -> the `file_patterns` gap answer
        # continues at resolve time; only `generate` stops when it ends empty.
        notes.append(f"feed {feed.feed_name!r}: neither the FRD nor the workbook (meta rows "
                     "'File Names' / 'File Name Example', File Details sheet) names a file "
                     "pattern; left open for the VDD FILES sheet / the file_patterns answer")
    from codegen.formats import is_spreadsheet

    spreadsheet = is_spreadsheet(feed.file_format, config,
                                 [*feed.file_name_patterns, *patterns])
    delimiter = "" if spreadsheet else (
        feed.delimiter or sheet.meta.get("delimiter")
        or _FORMAT_DELIMITERS.get((feed.file_format or "").lower()))
    # M11: which worksheet an inbound spreadsheet's data sits on (STTM
    # meta row); unstated, the layout stage asks and the gate flags it.
    sheet_name = sheet.meta.get("sheet_name") if spreadsheet else None
    frequency = (details or {}).get("frequency") or sheet.meta.get("frequency") or feed.frequency
    source_system = (details or {}).get("vendor") or sheet.meta.get("file_generator") \
        or feed.source_system

    try:
        return SttmFeed(
            feed_id=feed_id,
            source_system=source_system,
            mapping_sheet=name,
            source_file=SourceFile(name_pattern=name_pattern, format=feed.file_format,
                                   delimiter=delimiter, frequency=frequency,
                                   sheet_name=sheet_name),
            stage=TableRef(schema=stage_schema, table=stage_table, catalog=stage_catalog),
            standard=(TableRef(schema=standard_schema or "", table=standard_table or "",
                               catalog=standard_catalog) if has_standard else None),
            load_rules=LoadRules(
                not_null_columns=_rule_columns(fields, segmented, lambda f: not f.nullable),
                mandatory_columns=_rule_columns(fields, segmented, lambda f: f.mandatory),
                phi_columns=_rule_columns(fields, segmented, lambda f: f.phi),
                recycle=_recycle(sheet, config, disc),
            ),
            audit_columns=audit,
            field_count=len(fields),
            fields=fields,
            meta_rows=dict(sheet.meta),
            source_table=_dominant([r.values.get("source.source_table") for r in sheet.fields]),
            extraction_flags=flags,
            source_kind=source_kind,
        )
    except ValueError as exc:
        raise GenericExtractionError(
            f"feed {feed.feed_name!r}: extracted feed fails contract validation:\n{exc}"
        ) from exc


def extract_generic_contract(found: Discovery, workbook_path: Path, frd: FrdContract,
                             config: Config, *, contract_name: str | None,
                             generated_date: str,
                             width_answers: dict[int, dict[str, int]] | None = None
                             ) -> SttmContract:
    ir = read_workbook(found, workbook_path.name, config)
    return build_generic_contract(ir, frd, config, contract_name=contract_name,
                                  generated_date=generated_date, width_answers=width_answers)


__all__ = [
    "AuditRow",
    "FieldRow",
    "GenericExtractionError",
    "GenericIR",
    "SheetData",
    "build_generic_contract",
    "extract_generic_contract",
    "layout_summary",
    "read_workbook",
]
