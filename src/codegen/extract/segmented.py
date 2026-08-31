"""Segmented (Header/Detail/Trailer) STTM workbook extraction — v2 dialect.

The layout family the flat extractor refuses (one wide mapping sheet:
key:value metadata block, band-label row discovered by scan, header row
beneath it, per-row Segment column, per-segment audit rows with empty source
cells — docs/SEGMENTED_MODE_DESIGN.md) parses here into the SAME flat feed
shape the rest of the pipeline consumes: Detail rows are the payload and
become the column mappings; Header/Trailer rows are file envelope and become
``EnvelopeEntry`` records (report-only, never table DDL); everything else the
workbook declared rides in the feed's ``segmented`` block.

Governance boundaries — DECLARED, never inferred:

1. The H/D/T discriminator values are stated nowhere in the workbook. The
   run proceeds only when the feed's FAQ declares
   ``record_type_discriminators``; absent, the v1 refusal fires unchanged
   plus a remedy line. ``assumed_*`` status gates a review item.
2. The FRD contract governs target layers. A workbook Standard layer the FRD
   does not scope is parsed and HELD (``held_standard`` — table + column
   count), never emitted and never deleted.
3. (Discovered on the real workbook:) the Mandatory/Primary Key columns may
   carry no signal at all. The natural key then comes only from the FAQ's
   ``natural_key_columns`` declaration; absent, loud refusal.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from codegen.config import Config, SegmentedExtractorConfig
from codegen.contracts.frd import FrdContract, FrdFeed
from codegen.contracts.sttm import (
    AuditColumn,
    EnvelopeEntry,
    HeldStandardTable,
    LoadRules,
    SegmentedExtraction,
    SourceFile,
    SttmContract,
    SttmFeed,
    SttmField,
    TableRef,
)
from codegen.extract.workbook import SegmentedWorkbookError, WorkbookParseError, _norm, _text
from codegen.faq import load_faq
from codegen.resolve.resolver import _FORMAT_DELIMITERS, normalize_feed_name

_AUDIT_DATATYPES = {"string": "String", "timestamp": "Timestamp"}
_YES_VALUES = {"yes", "y", "true"}
_DELIMITER_RE = re.compile(r"delimited\s*(\S)")


@dataclass(frozen=True)
class _SegRow:
    row_number: int
    segment: str | None  # canonical Header/Detail/Trailer; None = audit row
    values: tuple


@dataclass(frozen=True)
class _Blocks:
    """Resolved column indexes for the three bands."""

    source: dict[str, int]
    stage: dict[str, int]
    standard: dict[str, int] | None


def _canonical_file_name(name: str) -> str:
    return name.strip().lower().replace("ccyy", "yyyy")


def _find_mapping_sheet(workbook, seg_config: SegmentedExtractorConfig,
                        band_vocabulary: set[str], name: str):
    """(worksheet, band_row_index, rows) — the one sheet matching the family
    signature (band row + metadata block above and/or Segment header below)."""
    candidates = []
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        for index, row in enumerate(rows):
            found = {v for v in (_norm(c) for c in row) if v in band_vocabulary}
            if len(found) < 2:
                continue
            metadata_above = any(
                r[0] is not None and len(r) > 1 and r[1] is not None
                for r in rows[:index]
            )
            header_row = rows[index + 1] if index + 1 < len(rows) else ()
            segment_header = any(_norm(c).startswith("segment") for c in header_row)
            if metadata_above or segment_header:
                candidates.append((ws, index, rows))
            break
    if not candidates:
        raise WorkbookParseError(
            f"{name}: no sheet matches the segmented layout family "
            "(band row + metadata block / Segment header)"
        )
    if len(candidates) > 1:
        raise WorkbookParseError(
            f"{name}: {len(candidates)} sheets match the segmented layout family "
            f"({[c[0].title for c in candidates]}); expected exactly one mapping sheet"
        )
    return candidates[0]


def _parse_metadata_block(rows, band_row: int, seg_config: SegmentedExtractorConfig,
                          sheet: str) -> dict[str, str]:
    wanted = {_norm(label): logical for logical, label in seg_config.metadata_keys.items()}
    facts: dict[str, str] = {}
    for row in rows[:band_row]:
        if not row or row[0] is None:
            continue
        key = _norm(row[0])
        value = _text(row[1]) if len(row) > 1 else None
        if key in wanted and value is not None:
            logical = wanted[key]
            if logical in facts:
                raise WorkbookParseError(
                    f"{sheet}: metadata key {row[0]!r} appears more than once")
            facts[logical] = value
    missing = {"files", "generator", "file_type"} - set(facts)
    if missing:
        raise WorkbookParseError(
            f"{sheet}: metadata block is missing key(s) "
            f"{sorted(seg_config.metadata_keys[m] for m in missing)}"
        )
    return facts


def _resolve_blocks(header: list, band_starts: dict[str, int], max_col: int,
                    seg_config: SegmentedExtractorConfig, sheet: str) -> _Blocks:
    source_start = band_starts["source"]
    stage_start = band_starts["stage"]
    standard_start = band_starts.get("standard")

    source_synonyms = {
        logical: {_norm(s) for s in names}
        for logical, names in seg_config.source_headers.items()
    }
    source: dict[str, int] = {}
    for index in range(source_start, stage_start):
        text = _norm(header[index]) if index < len(header) else ""
        if not text:
            continue
        matches = [
            logical for logical, names in source_synonyms.items()
            if text in names or any(text.startswith(n.split(" (")[0]) for n in names)
        ]
        if not matches:
            raise WorkbookParseError(
                f"{sheet}: unrecognized source-block header {header[index]!r} "
                f"(column {index + 1}); known: {seg_config.source_headers}"
            )
        source.setdefault(matches[0], index)
    required = {"field_name", "datatype", "segment"} - set(source)
    if required:
        raise WorkbookParseError(
            f"{sheet}: source block is missing required column(s) {sorted(required)}")

    def _table_block(start: int, end: int, band: str) -> dict[str, int]:
        known = {_norm(h): h for h in seg_config.table_headers}
        resolved: dict[str, int] = {}
        for index in range(start, end):
            text = _norm(header[index]) if index < len(header) else ""
            if not text:
                continue
            if text not in known:
                raise WorkbookParseError(
                    f"{sheet}: unrecognized {band}-block header {header[index]!r} "
                    f"(column {index + 1}); expected one of {seg_config.table_headers}"
                )
            resolved.setdefault(text, index)
        missing = {"schema", "tablename", "columnname", "datatype"} - set(resolved)
        if missing:
            raise WorkbookParseError(
                f"{sheet}: {band} block is missing column(s) {sorted(missing)}")
        return resolved

    stage = _table_block(stage_start, standard_start or (max_col + 1), "stage")
    standard = (
        _table_block(standard_start, max_col + 1, "standard")
        if standard_start is not None
        else None
    )
    return _Blocks(source=source, stage=stage, standard=standard)


def _segment_rows(rows, header_row: int, blocks: _Blocks,
                  seg_config: SegmentedExtractorConfig, sheet: str) -> list[_SegRow]:
    canonical = {k: v for k, v in seg_config.segment_names.items()}
    out: list[_SegRow] = []
    for row_number, row in enumerate(rows[header_row + 1:], start=header_row + 2):
        if all(_text(cell) is None for cell in row):
            continue
        segment_raw = _cell(row, blocks.source["segment"])
        field_name = _cell(row, blocks.source["field_name"])
        segment = canonical.get(_norm(segment_raw)) if segment_raw is not None else None
        if segment_raw is not None and segment is None:
            raise WorkbookParseError(
                f"{sheet} row {row_number}: unrecognized Segment value {segment_raw!r} "
                f"(known: {sorted(seg_config.segment_names.values())})"
            )
        if segment is None and field_name is not None:
            raise WorkbookParseError(
                f"{sheet} row {row_number}: field {field_name!r} carries no Segment value"
            )
        out.append(_SegRow(row_number=row_number, segment=segment, values=row))
    if not out:
        raise WorkbookParseError(f"{sheet}: no data rows below the header row")
    return out


def _cell(row: tuple, index: int | None) -> str | None:
    if index is None or index >= len(row):
        return None
    return _text(row[index])


def _is_yes(value: str | None) -> bool:
    return _norm(value) in _YES_VALUES


def extract_segmented_contract(
    workbook_path: Path,
    frd: FrdContract,
    config: Config,
    *,
    refusal_evidence: str,
    contract_name: str | None = None,
    generated_date: str | None = None,
) -> SttmContract:
    """Parse a segmented workbook + its FRD feed into the STTM contract.

    ``refusal_evidence`` is the family-detection message the flat path
    raised — re-used verbatim when the declaration gate refuses, so the
    evidence list stays identical to v1's refusal plus the remedy line.
    """
    from codegen.extract.extractor import ExtractionError  # local: avoid cycle

    seg_config = config.extractor.segmented
    labels = config.extractor.band_labels
    band_vocabulary = (
        {_norm(labels.source), _norm(labels.stage), _norm(labels.standard)}
        | {_norm(v) for v in seg_config.source_band_variants}
    )

    workbook = load_workbook(workbook_path, data_only=True)
    ws, band_row, rows = _find_mapping_sheet(
        workbook, seg_config, band_vocabulary, workbook_path.name)

    band_starts: dict[str, int] = {}
    for index, cell in enumerate(rows[band_row]):
        text = _norm(cell)
        if not text:
            continue
        if text in {_norm(labels.source)} | {_norm(v) for v in seg_config.source_band_variants}:
            band_starts["source"] = index
        elif text == _norm(labels.stage):
            band_starts["stage"] = index
        elif text == _norm(labels.standard):
            band_starts["standard"] = index
    if "source" not in band_starts or "stage" not in band_starts:
        raise WorkbookParseError(
            f"{ws.title}: band row {band_row + 1} lacks source/stage band labels")

    facts = _parse_metadata_block(rows, band_row, seg_config, ws.title)
    max_col = max(len(r) for r in rows) - 1
    header = list(rows[band_row + 1])
    blocks = _resolve_blocks(header, band_starts, max_col, seg_config, ws.title)
    seg_rows = _segment_rows(rows, band_row + 1, blocks, seg_config, ws.title)

    frd_feed = _match_frd_feed(frd, facts, workbook_path.name)
    feed_id = config.feed_aliases.get(frd_feed.feed_name) or normalize_feed_name(
        frd_feed.feed_name)

    # ---- governance boundary 1: discriminators must be DECLARED ----------
    faq = load_faq(normalize_feed_name(feed_id), config)
    discriminators = faq.record_type_discriminators
    if discriminators is None:
        raise SegmentedWorkbookError(
            refusal_evidence
            + f"\nREMEDY: declare record_type_discriminators (with status "
            f"assumed_pending_source_team or confirmed) in the feed's FAQ file "
            f"(fixtures/faq/{normalize_feed_name(feed_id)}.faq.yaml) to proceed "
            "under an explicitly surfaced assumption — the workbook states the "
            "H/D/T values nowhere, so the agent will not infer them."
        )

    # ---- group rows: Detail payload, H/T envelope, per-segment audit ------
    detail_rows: list[_SegRow] = []
    envelope: list[EnvelopeEntry] = []
    audit_by_segment: dict[str, list[tuple[str, str]]] = {}
    row_counts: dict[str, int] = {}
    segments_found: list[str] = []
    current_segment: str | None = None
    for seg_row in seg_rows:
        row = seg_row.values
        if seg_row.segment is not None:
            current_segment = seg_row.segment
            if seg_row.segment not in segments_found:
                segments_found.append(seg_row.segment)
            row_counts[seg_row.segment] = row_counts.get(seg_row.segment, 0) + 1
            if seg_row.segment == "Detail":
                detail_rows.append(seg_row)
            else:
                envelope.append(EnvelopeEntry(
                    segment=seg_row.segment,  # type: ignore[arg-type]
                    field_name=_cell(row, blocks.source["field_name"]) or "",
                    datatype=_cell(row, blocks.source.get("datatype")),
                    stage_table=_cell(row, blocks.stage.get("tablename")),
                    description=_cell(row, blocks.source.get("comments")),
                    rule=_cell(row, blocks.source.get("business_rule")),
                ))
            continue
        # audit row: empty source cell + empty Segment, inside a segment block
        column = _cell(row, blocks.stage.get("columnname"))
        datatype_raw = _cell(row, blocks.stage.get("datatype")) or ""
        if column is None or current_segment is None:
            raise WorkbookParseError(
                f"{ws.title} row {seg_row.row_number}: row carries neither a Segment "
                "value nor an audit-column identity"
            )
        datatype = _AUDIT_DATATYPES.get(_norm(datatype_raw))
        if datatype is None:
            raise WorkbookParseError(
                f"{ws.title} row {seg_row.row_number}: audit column {column!r} has "
                f"datatype {datatype_raw!r}; expected one of {sorted(_AUDIT_DATATYPES)}"
            )
        audit_by_segment.setdefault(current_segment, []).append((column, datatype))

    if not detail_rows:
        raise WorkbookParseError(f"{ws.title}: no Detail-segment rows found")

    # ---- Detail payload -> flat-shaped fields ----------------------------
    stage_tables = {
        _cell(r.values, blocks.stage["tablename"]) for r in detail_rows
    } - {None}
    if len(stage_tables) != 1:
        raise WorkbookParseError(
            f"{ws.title}: Detail rows name {len(stage_tables)} stage tables "
            f"{sorted(t for t in stage_tables if t)}; expected exactly one"
        )
    stage_table = next(iter(stage_tables))
    stage_schema = _cell(detail_rows[0].values, blocks.stage["schema"])
    if stage_schema is None:
        raise WorkbookParseError(f"{ws.title}: Detail rows carry no stage Schema")

    # ---- governance boundary 2: the FRD governs target layers -------------
    frd_has_standard = bool(frd_feed.standard_target.tables)
    held_standard: list[HeldStandardTable] = []
    if blocks.standard is not None and not frd_has_standard:
        counts: dict[tuple[str | None, str | None, str], int] = {}
        for seg_row in seg_rows:
            if seg_row.segment is None:
                continue
            table = _cell(seg_row.values, blocks.standard["tablename"])
            if table is None:
                continue
            key = (
                _cell(seg_row.values, blocks.standard.get("catalog")),
                _cell(seg_row.values, blocks.standard["schema"]),
                table,
            )
            counts[key] = counts.get(key, 0) + 1
        held_standard = [
            HeldStandardTable(
                catalog=catalog, schema=schema, table=table, column_count=count,
                reason=(
                    f"workbook defines a Standard layer; FRD contract "
                    f"'{frd.contract_name}' scopes this feed stage-only — held "
                    "for source-team ruling, not emitted"
                ),
            )
            for (catalog, schema, table), count in sorted(
                counts.items(), key=lambda kv: kv[0][2])
        ]

    emit_standard = blocks.standard is not None and frd_has_standard

    fields: list[SttmField] = []
    mandatory_sources: list[str] = []
    for seg_row in detail_rows:
        row = seg_row.values
        field_name = _cell(row, blocks.source["field_name"])
        assert field_name is not None  # _segment_rows guarantees it
        mandatory = _is_yes(_cell(row, blocks.stage.get("mandatory column")))
        primary_key = _is_yes(_cell(row, blocks.stage.get("primary key")))
        if mandatory or primary_key:
            mandatory_sources.append(field_name)
        value_parts = [
            part for part in (
                _cell(row, blocks.stage.get("transformations/data quality")),
                _cell(row, blocks.source.get("business_rule")),
            ) if part
        ]
        fields.append(SttmField(
            source_column=field_name,
            description=_cell(row, blocks.source.get("comments")),
            sample_value=None,
            source_datatype=_cell(row, blocks.source.get("datatype")) or "unstated",
            nullable=not (mandatory or primary_key),
            phi=_is_yes(_cell(row, blocks.source.get("pii"))),
            mandatory=mandatory,
            stage_column=_require(ws.title, seg_row.row_number,
                                  _cell(row, blocks.stage["columnname"]),
                                  "stage ColumnName"),
            stage_datatype=_require(ws.title, seg_row.row_number,
                                    _cell(row, blocks.stage["datatype"]),
                                    "stage DataType"),
            standard_column=(_cell(row, blocks.standard["columnname"])
                             if emit_standard else None),
            standard_datatype=(_cell(row, blocks.standard["datatype"])
                               if emit_standard else None),
            value_spec="; ".join(value_parts) if value_parts else None,
        ))

    # ---- natural key: workbook signal, else FAQ declaration ---------------
    notes: list[str] = []
    if mandatory_sources:
        not_null = mandatory_sources
        natural_key_declared: list[str] = []
    elif faq.natural_key_columns:
        source_names = {f.source_column for f in fields}
        unknown = [c for c in faq.natural_key_columns if c not in source_names]
        if unknown:
            raise ExtractionError(
                f"feed {frd_feed.feed_name!r}: FAQ natural_key_columns "
                f"{unknown} are not Detail-segment source fields"
            )
        not_null = list(faq.natural_key_columns)
        natural_key_declared = list(faq.natural_key_columns)
        notes.append(
            "natural key: the workbook's Mandatory/Primary Key columns carry no "
            "signal; using the FAQ-declared natural_key_columns "
            f"{natural_key_declared} (engineer declaration, not workbook evidence)"
        )
    else:
        raise ExtractionError(
            f"feed {frd_feed.feed_name!r}: the workbook's Mandatory/Primary Key "
            "columns carry no signal and the FAQ declares no natural_key_columns "
            "— declare the natural-key source columns in "
            f"fixtures/faq/{normalize_feed_name(feed_id)}.faq.yaml; the agent "
            "will not invent a merge key"
        )

    if frd_feed.recycle_rule:
        # The workbook layout has no Recycle Flag column; the FRD's recycle
        # rule text exists but nothing structured backs it. Surface, don't
        # fabricate a RecycleSpec.
        notes.append(
            f"recycle: FRD states {frd_feed.recycle_rule!r} but this workbook "
            "layout carries no recycle validation column — no recycle module is "
            "generated; the FRD-declared recycle stage table is held with the "
            "envelope (source-team input needed)"
        )

    detail_audit = audit_by_segment.get("Detail", [])
    if not detail_audit:
        raise ExtractionError(
            f"feed {frd_feed.feed_name!r}: no Detail-segment audit rows found "
            "(rows with empty source cells inside the Detail block)"
        )

    segmented = SegmentedExtraction(
        segments_found=segments_found,
        row_counts=row_counts,
        envelope=envelope,
        discriminators=discriminators,
        natural_key_declared=natural_key_declared,
        held_standard=held_standard,
        notes=notes,
    )

    file_pattern = _match_file_pattern(facts["files"], frd_feed, workbook_path.name)
    delimiter = frd_feed.delimiter or _metadata_delimiter(facts) or \
        _FORMAT_DELIMITERS.get(frd_feed.file_format.lower())
    if delimiter is None:
        raise ExtractionError(
            f"feed {frd_feed.feed_name!r}: no delimiter in the FRD, the metadata "
            f"block ({facts['file_type']!r}), or implied by format "
            f"{frd_feed.file_format!r}"
        )

    feed = SttmFeed(
        feed_id=feed_id,
        source_system=facts["generator"],
        mapping_sheet=ws.title,
        source_file=SourceFile(
            name_pattern=file_pattern,
            format=frd_feed.file_format,
            delimiter=delimiter,
            frequency=facts.get("frequency"),
        ),
        stage=TableRef(schema=stage_schema, table=stage_table),
        standard=(
            TableRef(
                schema=_cell(detail_rows[0].values, blocks.standard["schema"]) or "",
                table=_cell(detail_rows[0].values, blocks.standard["tablename"]) or "",
            )
            if emit_standard else None
        ),
        load_rules=LoadRules(
            not_null_columns=not_null,
            mandatory_columns=mandatory_sources or list(natural_key_declared),
            phi_columns=[f.source_column for f in fields if f.phi],
            recycle=None,
        ),
        audit_columns=[AuditColumn(column=c, datatype=d) for c, d in detail_audit],
        field_count=len(fields),
        fields=fields,
        segmented=segmented,
    )

    return SttmContract(
        contract_name=contract_name
        or f"STTM mapping contract extracted from {workbook_path.name}",
        generated_from_workbook=workbook_path.name,
        sttm_version="unversioned (segmented workbook carries no VERSION_HISTORY sheet)",
        generated_date=generated_date or _today(),
        notes=[
            "Extracted deterministically by `codegen extract-sttm` (segmented "
            f"dialect) from the workbook named above, paired with FRD contract "
            f"'{frd.contract_name}'.",
            "Detail-segment rows are the payload; Header/Trailer rows are file "
            "envelope (see the feed's segmented.envelope — report-only, never DDL).",
            "Record-type discriminators are DECLARED in the feed's FAQ "
            f"(status: {discriminators.status}) — the workbook states them nowhere.",
        ],
        feeds=[feed],
    )


def _require(sheet: str, row_number: int, value: str | None, what: str) -> str:
    if value is None:
        raise WorkbookParseError(f"{sheet} row {row_number}: {what} is empty")
    return value


def _today() -> str:
    import datetime

    return datetime.date.today().isoformat()


def _match_frd_feed(frd: FrdContract, facts: dict[str, str], name: str) -> FrdFeed:
    """Pair by file patterns (canonicalized), same rule as the flat path."""
    workbook_patterns = {
        _canonical_file_name(line)
        for line in re.split(r"[\n,;]+", facts["files"])
        if line.strip()
    }
    matches = [
        feed for feed in frd.feeds
        if workbook_patterns & {_canonical_file_name(p) for p in feed.file_name_patterns}
    ]
    if len(matches) != 1:
        raise WorkbookParseError(
            f"{name}: {len(matches)} FRD feeds match the metadata File(s) patterns "
            f"{sorted(workbook_patterns)} (FRD feeds: {[f.feed_name for f in frd.feeds]})"
        )
    return matches[0]


def _match_file_pattern(files_value: str, frd_feed: FrdFeed, name: str) -> str:
    frd_canonical = {_canonical_file_name(p): p for p in frd_feed.file_name_patterns}
    for line in re.split(r"[\n,;]+", files_value):
        canonical = _canonical_file_name(line)
        if canonical in frd_canonical:
            return line.strip()
    raise WorkbookParseError(
        f"{name}: none of the metadata File(s) entries match the FRD's "
        f"file_name_patterns {frd_feed.file_name_patterns}"
    )


def _metadata_delimiter(facts: dict[str, str]) -> str | None:
    match = _DELIMITER_RE.search(facts.get("file_type", ""))
    if match and match.group(1) not in (")",):
        return match.group(1)
    return None
