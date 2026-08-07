"""Parse a client-authored STTM workbook (.xlsx) into an intermediate form.

Layout knowledge lives in ``config.extractor`` (sheet names, band labels,
header synonyms, audit markers) — nothing here is positional. Every mapping
sheet is a merged three-band table (Source File Layout | Stage Layer |
Standard Layer): row 1 carries the band labels, row 2 the real headers, data
starts at row 3. Header resolution is fuzzy (lowercased, whitespace
collapsed) because one golden workbook already ships three source-block
header dialects — see docs/EXTRACTOR_RECON.md §3a.

Trailing rows whose source-side cell is an audit marker ("NA") are peeled
into audit columns, never returned as mapping rows. Flat dialect only: a
sheet mapping to more than one stage table (or naming a segment column) is
the segmented Header/Detail/Trailer dialect, which is out of scope — it
raises :class:`SegmentedWorkbookError` (decision D4).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from codegen.config import ExtractorConfig

# Audit rows carry lowercase workbook datatypes; the contract's AuditColumn
# is a Literal["String", "Timestamp"], so normalize casing here.
_AUDIT_DATATYPES = {"string": "String", "timestamp": "Timestamp"}

# Row-2 headers whose normalized text names a segment column mark the
# segmented dialect even before stage tables are compared.
_SEGMENT_HEADERS = {"segment", "record segment"}

_WINDOW_DAYS_RE = re.compile(r"(\d+)\s*days?", re.IGNORECASE)
_RECYCLE_FLAG_RE = re.compile(r"\bY\s*\(")


class WorkbookParseError(ValueError):
    """The workbook layout/content cannot be parsed; message names the spot."""


class SegmentedWorkbookError(WorkbookParseError):
    """Segmented (Header/Detail/Trailer) dialect detected — v1 is flat-only."""


def _norm(value: object) -> str:
    """Lowercase, collapse whitespace runs, strip — the comparison form for
    every header/label/marker in this module."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def _text(value: object) -> str | None:
    """Cell value as stripped text; None when empty."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


@dataclass(frozen=True)
class MappingRow:
    row_number: int
    source_column: str
    description: str | None
    sample_value: str | None
    source_datatype: str
    nullable: bool
    phi: bool
    mandatory: bool
    stage_column: str
    stage_datatype: str
    standard_column: str | None
    standard_datatype: str | None
    value_spec: str | None
    recycle_text: str | None


@dataclass(frozen=True)
class AuditRow:
    column: str
    datatype: str  # normalized: "String" | "Timestamp"


@dataclass(frozen=True)
class SheetIR:
    sheet_name: str
    stage_schema: str
    stage_table: str
    standard_schema: str | None
    standard_table: str | None
    rows: tuple[MappingRow, ...]
    audit: tuple[AuditRow, ...]


@dataclass(frozen=True)
class FileDetailsRow:
    vendor: str
    file_name: str
    frequency: str | None


@dataclass(frozen=True)
class RecycleText:
    """One flagged row's verbatim recycle cell, pre-digested."""

    applies_to: str
    validation: str  # verbatim inner text after the Y(...) flag
    window_days: int | None


@dataclass(frozen=True)
class WorkbookIR:
    workbook_name: str
    version: str
    file_details: tuple[FileDetailsRow, ...]
    sheets: tuple[SheetIR, ...]


def parse_workbook(path: Path, config: ExtractorConfig) -> WorkbookIR:
    workbook = load_workbook(path, data_only=True)
    sheet_names = workbook.sheetnames

    mapping_sheets = [n for n in sheet_names if n.startswith(config.mapping_sheet_prefix)]
    if not mapping_sheets:
        _reject_segmented_family(workbook, config, path.name)
        raise WorkbookParseError(
            f"{path.name}: no mapping sheets found (prefix {config.mapping_sheet_prefix!r}); "
            f"sheets present: {sheet_names}"
        )

    return WorkbookIR(
        workbook_name=path.name,
        version=_parse_version_history(workbook, config, path.name),
        file_details=_parse_file_details(workbook, config, path.name),
        sheets=tuple(_parse_mapping_sheet(workbook[n], config) for n in mapping_sheets),
    )


# How many leading rows to scan for a band-label row when checking whether a
# prefix-less workbook belongs to the segmented layout family (its bands sit
# below a key:value metadata block, row 9 in the observed real workbook).
_FAMILY_SCAN_ROWS = 30

# The segmented family's band vocabulary includes a shorter source label.
_SOURCE_BAND_VARIANTS = {"source layout"}


def _reject_segmented_family(workbook, config: ExtractorConfig, name: str) -> None:
    """Content-based recognition of the segmented (CAQH-style) layout
    family, so it gets an accurate diagnostic instead of the misleading
    'no mapping sheets found'. Signature: some row's cells resolve to >= 2
    band labels (including the 'Source Layout' variant), with a key:value
    metadata block above it and/or a Segment column in the header row
    beneath it."""
    labels = config.band_labels
    band_vocabulary = {
        _norm(labels.source),
        _norm(labels.stage),
        _norm(labels.standard),
    } | _SOURCE_BAND_VARIANTS

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        rows = [
            row
            for _, row in zip(
                range(_FAMILY_SCAN_ROWS), ws.iter_rows(values_only=True), strict=False
            )
        ]
        for index, row in enumerate(rows):
            found = {v for v in (_norm(c) for c in row) if v in band_vocabulary}
            if len(found) < 2:
                continue
            metadata_above = any(
                r[0] is not None and len(r) > 1 and r[1] is not None for r in rows[:index]
            )
            header_row = rows[index + 1] if index + 1 < len(rows) else ()
            segment_header = any(_norm(c).startswith("segment") for c in header_row)
            if metadata_above or segment_header:
                raise SegmentedWorkbookError(
                    f"{name}: sheet {sheet_name!r} matches the segmented (metadata-block"
                    f" + band-row) layout family — band labels {sorted(found)} on row "
                    f"{index + 1}"
                    + (", key:value metadata block above" if metadata_above else "")
                    + (", per-row Segment column in the header" if segment_header else "")
                    + ". This layout family is recognized but unsupported in v1 "
                    "(flat dialect only) — see docs/SEGMENTED_MODE_DESIGN.md"
                )


# ------------------------------------------------------------- metadata sheets


def _parse_file_details(workbook, config: ExtractorConfig, name: str) -> tuple[FileDetailsRow, ...]:
    sheet_name = config.file_details_sheet
    if sheet_name not in workbook.sheetnames:
        raise WorkbookParseError(f"{name}: required sheet {sheet_name!r} is missing")
    ws = workbook[sheet_name]

    header = [_norm(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    def _find(logical: str, synonyms: list[str]) -> int:
        normalized = [_norm(s) for s in synonyms]
        hits = [i for i, h in enumerate(header) if h in normalized]
        if len(hits) != 1:
            raise WorkbookParseError(
                f"{sheet_name}: expected exactly one {logical!r} header among {synonyms} "
                f"in row 1, found {len(hits)} (headers: {header})"
            )
        return hits[0]

    fd_headers = config.file_details_headers
    vendor_col = _find("vendor", fd_headers.vendor)
    file_col = _find("file_name", fd_headers.file_name)
    freq_col = _find("frequency", fd_headers.frequency)

    rows: list[FileDetailsRow] = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vendor, file_name = _text(row[vendor_col]), _text(row[file_col])
        if vendor is None and file_name is None:
            continue
        if vendor is None or file_name is None:
            raise WorkbookParseError(
                f"{sheet_name} row {row_number}: Vendor and FileName must both be present "
                f"(got vendor={vendor!r}, file_name={file_name!r})"
            )
        rows.append(
            FileDetailsRow(vendor=vendor, file_name=file_name, frequency=_text(row[freq_col]))
        )
    if not rows:
        raise WorkbookParseError(f"{sheet_name}: no data rows")
    return tuple(rows)


def _parse_version_history(workbook, config: ExtractorConfig, name: str) -> str:
    """The sheet's content block floats (the golden workbook anchors it at
    E6), so locate the header row by content, then take the LAST stated
    version."""
    sheet_name = config.version_history_sheet
    if sheet_name not in workbook.sheetnames:
        raise WorkbookParseError(f"{name}: required sheet {sheet_name!r} is missing")
    ws = workbook[sheet_name]

    version_col: int | None = None
    header_row: int | None = None
    for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
        normalized = [_norm(c) for c in row]
        if "version" in normalized and "date" in normalized:
            version_col = normalized.index("version")
            header_row = row_number
            break
    if version_col is None or header_row is None:
        raise WorkbookParseError(
            f"{sheet_name}: no header row containing both 'Version' and 'Date' found"
        )

    versions = [
        _text(row[version_col])
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True)
        if _text(row[version_col]) is not None
    ]
    if not versions:
        raise WorkbookParseError(f"{sheet_name}: header found but no version rows below it")
    return versions[-1]  # type: ignore[return-value]  # None filtered above


# ------------------------------------------------------------- mapping sheets


@dataclass(frozen=True)
class _Bands:
    """Column index ranges (0-based, inclusive start / exclusive end)."""

    source: tuple[int, int]
    stage: tuple[int, int]
    standard: tuple[int, int] | None


def _find_bands(ws: Worksheet, config: ExtractorConfig) -> _Bands:
    labels = config.band_labels
    row1 = [_norm(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    def _start(label: str) -> int | None:
        hits = [i for i, v in enumerate(row1) if v == _norm(label)]
        if len(hits) > 1:
            raise WorkbookParseError(
                f"{ws.title}: band label {label!r} appears {len(hits)} times in row 1"
            )
        return hits[0] if hits else None

    source = _start(labels.source)
    stage = _start(labels.stage)
    standard = _start(labels.standard)
    if source is None or stage is None:
        raise WorkbookParseError(
            f"{ws.title}: row 1 must carry the {labels.source!r} and {labels.stage!r} "
            f"band labels; found {[v for v in row1 if v]}"
        )
    if not source < stage or (standard is not None and not stage < standard):
        raise WorkbookParseError(
            f"{ws.title}: band labels out of order (source at {source}, stage at {stage}, "
            f"standard at {standard})"
        )
    stage_end = standard if standard is not None else ws.max_column
    return _Bands(
        source=(source, stage),
        stage=(stage, stage_end),
        standard=(standard, ws.max_column) if standard is not None else None,
    )


def _find_recycle_column(
    ws: Worksheet, headers: list[str | None], config: ExtractorConfig
) -> int | None:
    """The per-row Recycle Flag column can sit anywhere — the golden
    workbook parks it beyond the standard band (col R) — so scan the whole
    header row."""
    prefix = _norm(config.recycle_header_prefix)
    hits = [i for i, h in enumerate(headers) if _norm(h).startswith(prefix)]
    if len(hits) > 1:
        raise WorkbookParseError(
            f"{ws.title}: {len(hits)} row-2 headers start with "
            f"{config.recycle_header_prefix!r} (columns {[i + 1 for i in hits]}); "
            "expected at most one"
        )
    return hits[0] if hits else None


def _resolve_source_headers(
    ws: Worksheet,
    headers: list[str | None],
    band: tuple[int, int],
    config: ExtractorConfig,
    recycle_col: int | None,
) -> dict[str, int]:
    """Map logical column names -> column index within the source band."""
    synonyms = {
        logical: [_norm(s) for s in names] for logical, names in config.header_synonyms.items()
    }
    resolved: dict[str, int] = {}
    for index in range(*band):
        if index == recycle_col:
            continue
        header = _norm(headers[index])
        if not header:
            continue
        if header in _SEGMENT_HEADERS:
            raise SegmentedWorkbookError(
                f"{ws.title}: header {headers[index]!r} names a record segment column — "
                "segmented (Header/Detail/Trailer) STTM workbooks are not supported by "
                "this extractor (flat dialect only; see CLAUDE.md)"
            )
        matches = [logical for logical, names in synonyms.items() if header in names]
        if not matches:
            raise WorkbookParseError(
                f"{ws.title}: unrecognized source-block header {headers[index]!r} "
                f"(column {index + 1}, row 2); known synonyms: {config.header_synonyms}"
            )
        logical = matches[0]
        if logical in resolved:
            raise WorkbookParseError(
                f"{ws.title}: source-block header for {logical!r} appears twice "
                f"(columns {resolved[logical] + 1} and {index + 1})"
            )
        resolved[logical] = index

    missing = {
        "source_column",
        "description",
        "sample_value",
        "source_datatype",
        "null_check",
        "phi",
        "mandatory",
    } - set(resolved)
    if missing:
        raise WorkbookParseError(
            f"{ws.title}: source block is missing required column(s) {sorted(missing)}; "
            f"row-2 headers seen: {[headers[i] for i in range(*band)]}"
        )
    return resolved


def _resolve_table_headers(
    ws: Worksheet,
    headers: list[str | None],
    band: tuple[int, int],
    band_name: str,
    config: ExtractorConfig,
    recycle_col: int | None,
) -> dict[str, int]:
    """Map the Schema/TableName/ColumnName/DataType headers of a stage or
    standard band -> column index. Blank spacer columns are skipped; the
    per-row recycle column may sit inside the band's range."""
    known = [_norm(h) for h in config.table_block_headers]
    resolved: dict[str, int] = {}
    for index in range(*band):
        if index == recycle_col:
            continue
        header = _norm(headers[index])
        if not header:
            continue
        if header.startswith(_norm(config.recycle_header_prefix)):
            continue
        if header not in known:
            raise WorkbookParseError(
                f"{ws.title}: unrecognized {band_name}-block header {headers[index]!r} "
                f"(column {index + 1}, row 2); expected one of {config.table_block_headers}"
            )
        if header in resolved:
            raise WorkbookParseError(
                f"{ws.title}: {band_name}-block header {headers[index]!r} appears twice"
            )
        resolved[header] = index
    missing = set(known) - set(resolved)
    if missing:
        raise WorkbookParseError(
            f"{ws.title}: {band_name} block is missing column(s) {sorted(missing)}"
        )
    return resolved


def _parse_bool(
    ws: Worksheet, row_number: int, value: object, kind: str, true_values: tuple[str, ...],
    false_values: tuple[str, ...],
) -> bool:
    normalized = _norm(value)
    if normalized in true_values:
        return True
    if normalized in false_values:
        return False
    raise WorkbookParseError(
        f"{ws.title} row {row_number}: unparseable {kind} value {value!r} "
        f"(expected one of {true_values + false_values})"
    )


def _required_text(ws: Worksheet, row_number: int, value: object, column_name: str) -> str:
    text = _text(value)
    if text is None:
        raise WorkbookParseError(f"{ws.title} row {row_number}: {column_name} is empty")
    return text


def _parse_mapping_sheet(ws: Worksheet, config: ExtractorConfig) -> SheetIR:
    bands = _find_bands(ws, config)
    headers = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    recycle_col = _find_recycle_column(ws, headers, config)
    source_cols = _resolve_source_headers(ws, headers, bands.source, config, recycle_col)
    stage_cols = _resolve_table_headers(ws, headers, bands.stage, "stage", config, recycle_col)
    standard_cols = (
        _resolve_table_headers(ws, headers, bands.standard, "standard", config, recycle_col)
        if bands.standard is not None
        else None
    )

    audit_markers = tuple(_norm(m) for m in config.audit_source_markers)
    rows: list[MappingRow] = []
    audit: list[AuditRow] = []
    stage_refs: set[tuple[str, str]] = set()
    standard_refs: set[tuple[str, str]] = set()

    for row_number, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if all(_text(cell) is None for cell in row):
            continue
        source_column = row[source_cols["source_column"]]

        if _norm(source_column) in audit_markers:
            audit.append(_parse_audit_row(ws, row_number, row, stage_cols))
            continue

        source_text = _required_text(ws, row_number, source_column, "source column")
        stage_ref = (
            _required_text(ws, row_number, row[stage_cols["schema"]], "stage Schema"),
            _required_text(ws, row_number, row[stage_cols["tablename"]], "stage TableName"),
        )
        stage_refs.add(stage_ref)
        standard_column = standard_datatype = None
        if standard_cols is not None:
            standard_refs.add(
                (
                    _required_text(ws, row_number, row[standard_cols["schema"]], "standard Schema"),
                    _required_text(
                        ws, row_number, row[standard_cols["tablename"]], "standard TableName"
                    ),
                )
            )
            standard_column = _required_text(
                ws, row_number, row[standard_cols["columnname"]], "standard ColumnName"
            )
            standard_datatype = _required_text(
                ws, row_number, row[standard_cols["datatype"]], "standard DataType"
            )

        null_check = _norm(row[source_cols["null_check"]])
        if not null_check:
            raise WorkbookParseError(f"{ws.title} row {row_number}: NULL-check cell is empty")
        nullable = not null_check.startswith("not null")

        rows.append(
            MappingRow(
                row_number=row_number,
                source_column=source_text,
                description=_text(row[source_cols["description"]]),
                sample_value=_text(row[source_cols["sample_value"]]),
                source_datatype=_required_text(
                    ws, row_number, row[source_cols["source_datatype"]], "source DataType"
                ),
                nullable=nullable,
                phi=_parse_bool(ws, row_number, row[source_cols["phi"]], "PHI", ("yes",), ("no",)),
                mandatory=_parse_bool(
                    ws, row_number, row[source_cols["mandatory"]], "Mandatory", ("yes",), ("no",)
                ),
                stage_column=_required_text(
                    ws, row_number, row[stage_cols["columnname"]], "stage ColumnName"
                ),
                stage_datatype=_required_text(
                    ws, row_number, row[stage_cols["datatype"]], "stage DataType"
                ),
                standard_column=standard_column,
                standard_datatype=standard_datatype,
                value_spec=(
                    _text(row[source_cols["value_spec"]]) if "value_spec" in source_cols else None
                ),
                recycle_text=_text(row[recycle_col]) if recycle_col is not None else None,
            )
        )

    if not rows:
        raise WorkbookParseError(f"{ws.title}: no mapping rows found")
    if len(stage_refs) > 1:
        raise SegmentedWorkbookError(
            f"{ws.title}: mapping rows land in {len(stage_refs)} stage tables "
            f"{sorted(t for _, t in stage_refs)} — this is the segmented "
            "(Header/Detail/Trailer) STTM dialect, which is not supported by this "
            "extractor (flat dialect only; see CLAUDE.md)"
        )
    if len(standard_refs) > 1:
        raise WorkbookParseError(
            f"{ws.title}: mapping rows name {len(standard_refs)} standard tables "
            f"{sorted(t for _, t in standard_refs)}; expected exactly one"
        )

    stage_schema, stage_table = next(iter(stage_refs))
    standard_schema, standard_table = (
        next(iter(standard_refs)) if standard_refs else (None, None)
    )
    return SheetIR(
        sheet_name=ws.title,
        stage_schema=stage_schema,
        stage_table=stage_table,
        standard_schema=standard_schema,
        standard_table=standard_table,
        rows=tuple(rows),
        audit=tuple(audit),
    )


def _parse_audit_row(
    ws: Worksheet, row_number: int, row: tuple, stage_cols: dict[str, int]
) -> AuditRow:
    column = _text(row[stage_cols["columnname"]])
    if column is None or _norm(column) == "na":
        raise WorkbookParseError(
            f"{ws.title} row {row_number}: audit row carries no stage ColumnName "
            f"(got {column!r}) — audit column identities are unrecoverable"
        )
    raw_datatype = _text(row[stage_cols["datatype"]]) or ""
    datatype = _AUDIT_DATATYPES.get(_norm(raw_datatype))
    if datatype is None:
        raise WorkbookParseError(
            f"{ws.title} row {row_number}: audit column {column!r} has datatype "
            f"{raw_datatype!r}; expected one of {sorted(_AUDIT_DATATYPES)}"
        )
    return AuditRow(column=column, datatype=datatype)


# ------------------------------------------------------------------- recycle


def parse_recycle_text(sheet: SheetIR) -> RecycleText | None:
    """Digest the per-row Recycle Flag cells of one sheet. The flagged row's
    cell looks like:

        Recycle Flag ( Enabled for 7 Days)
        Y ( Check with SUBS_ID from ... FOR GRP_CK = 47, if available ... )

    The text inside the ``Y ( ... )`` wrapper is kept VERBATIM as the
    validation text (decision D3 — the resolver accepts the client phrasing;
    no canonical rewriting here)."""
    flagged = [row for row in sheet.rows if row.recycle_text is not None]
    if not flagged:
        return None
    if len(flagged) > 1:
        raise WorkbookParseError(
            f"{sheet.sheet_name}: {len(flagged)} rows carry recycle text "
            f"({[r.source_column for r in flagged]}); expected at most one"
        )
    row = flagged[0]
    text = row.recycle_text
    assert text is not None
    match = _RECYCLE_FLAG_RE.search(text)
    if match is None:
        raise WorkbookParseError(
            f"{sheet.sheet_name} row {row.row_number}: recycle cell has no 'Y (...)' "
            f"flag marker: {text!r}"
        )
    validation = text[match.end() :].strip()
    if validation.endswith(")"):
        validation = validation[:-1].rstrip()
    if not validation:
        raise WorkbookParseError(
            f"{sheet.sheet_name} row {row.row_number}: recycle flag carries no validation text"
        )
    window_match = _WINDOW_DAYS_RE.search(text)
    return RecycleText(
        applies_to=row.source_column,
        validation=validation,
        window_days=int(window_match.group(1)) if window_match else None,
    )
