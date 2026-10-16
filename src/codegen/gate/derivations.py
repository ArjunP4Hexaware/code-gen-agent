"""Derivation gate (M7 §2 f–i): every derived NAME, PATH and SQL LITERAL is
checked before it ships, and the sibling-type consistency of the standard
band is flagged.

* IIG cells: no cell may contain a line break; identifier columns
  (``derivations.name_columns``) must match the name charset and the length
  cap; path columns (``derivations.path_column_suffixes``) may hold no
  whitespace, no empty segment, no segment ending in punctuation, exactly
  one separator between segments, and stay under the path cap. A
  violation is a FAILED gate check naming the field, the value (truncated)
  and the provenance the cell was built from. post57: whitespace is the
  conventions profile's ``path_whitespace`` — ``keep`` (default): the path is
  written as the document states it, byte for byte, and each path cell gets
  ONE ``iig_path_whitespace`` flag naming its coordinates and the kind of
  whitespace (double space / leading / trailing / tab); ``fail``: whitespace
  fails the check as before.
* SQL literals in the deployment DDL (``COMMENT '…'``, ``LOCATION '…'``,
  TBLPROPERTIES values): no newline / carriage return, balanced quotes
  after escaping, no leading / trailing whitespace; ``LOCATION`` literals
  are paths and follow the path rules too.
* Path assembly (:func:`join_path`) joins NORMALIZED segments; callers never
  string-concatenate a raw value.
* Name tokens (:func:`short_token`): a component of a WF_/NB_ name comes
  from a short token or stays blank — never a slug of prose.
* Sibling types: standard-band columns grouped by trailing suffix
  (``_pct``, ``_amt`` … from config, plus any suffix ≥ ``sibling_min_group``
  columns share); a group mixing base types is flagged
  ``sibling_type_mismatch:<suffix>`` citing the minority rows' STTM cells.
  The type is never altered.

Caps are config (``gate.derivations``): the framework's own limits are not
documented, so the shipped values are conservative defaults, stated as such.
"""

from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from codegen.config import Config, DerivationsConfig
from codegen.contracts.resolved import ResolvedFeedSpec
from codegen.gate.preflight import GateCheck

_PATH_SEPARATORS = re.compile(r"[\\/]+")
_TRAILING_PUNCT = ".,;:"
# M10.2: a value with one of these schemes is a LOCATION URI, not a folder
# path — validated as a URI (scheme, authority, no whitespace, its path
# segments by the folder rules) and carried through unchanged.
LOCATION_SCHEMES = ("abfss", "abfs", "wasbs", "wasb", "dbfs", "s3", "s3a", "adl", "gs")
_SCHEME_RE = re.compile(r"^\s*([A-Za-z][A-Za-z0-9+.-]*):/")


def location_scheme(value: str | None) -> str | None:
    """``abfss`` for ``abfss://c@a.example.invalid/x``, ``dbfs`` for
    ``dbfs:/mnt/x``; None for a folder path (or an unknown scheme, which the
    folder rules then reject — ``https://`` is not a storage location)."""
    match = _SCHEME_RE.match(value or "")
    scheme = match.group(1).lower() if match else None
    return scheme if scheme in LOCATION_SCHEMES else None


def split_location(value: str) -> tuple[str, str, str]:
    """(scheme, authority, path) of a location URI. ``dbfs:/mnt/x`` has no
    authority: ('dbfs', '', '/mnt/x')."""
    text = value.strip()
    scheme, _, rest = text.partition(":")
    if rest.startswith("//"):
        authority, _, tail = rest[2:].partition("/")
        return scheme.lower(), authority, "/" + tail if tail or rest[2:].endswith("/") else ""
    return scheme.lower(), "", rest


def uri_violations(value: str, cap: int) -> list[str]:
    """The location-URI rule: scheme + authority (except dbfs:/), no line
    break / whitespace, length cap, and every path segment by the folder
    rules (no empty segment, none ending in punctuation)."""
    problems: list[str] = []
    if "\n" in value or "\r" in value:
        problems.append("line break")
    if re.search(r"\s", value):
        problems.append("whitespace")
    if len(value) > cap:
        problems.append(f"length {len(value)} > {cap}")
    scheme, authority, path = split_location(value)
    if scheme != "dbfs" and not authority:
        problems.append(f"no authority after {scheme}://")
    if authority and authority[-1] in _TRAILING_PUNCT:
        problems.append(f"authority {authority!r} ends in punctuation")
    inner = path.strip("/")
    for segment in inner.split("/") if inner else []:
        if segment == "":
            problems.append("empty segment")
        elif segment[-1] in _TRAILING_PUNCT:
            problems.append(f"segment {segment!r} ends in punctuation")
    return problems


def join_location(uri: str, *parts: str | None) -> str:
    """A location URI + relative segments: the URI's own text is kept, the
    parts are normalized like ``join_path`` and appended with one separator
    each. The scheme / authority are never re-spelt."""
    base = uri.strip().rstrip("/")
    tail = join_path(*parts)
    return f"{base}/{tail}" if tail else base + "/"
_LITERAL_RE = re.compile(r"(COMMENT|LOCATION)\s+'((?:[^']|'')*)'", re.IGNORECASE)
_TBLPROP_RE = re.compile(r"'((?:[^']|'')*)'\s*=\s*'((?:[^']|'')*)'")


def _truncate(value: str, limit: int = 80) -> str:
    text = value.replace("\n", "\\n").replace("\r", "\\r")
    return text if len(text) <= limit else text[:limit] + "…"


# ----------------------------------------------------------------- paths ---


def path_segments(value: str) -> list[str]:
    """Segments of a path in either separator style, whitespace-stripped;
    the empty segments a doubled separator leaves are KEPT so the check can
    name them (assembly drops them, checking reports them)."""
    text = value.strip()
    if text.startswith(("abfss://", "https://", "http://", "wasbs://", "dbfs:/")):
        scheme, _, rest = text.partition("://")
        return [f"{scheme}://" + rest.split("/", 1)[0]] + (
            rest.split("/", 1)[1].split("/") if "/" in rest else [])
    return re.split(r"[\\/]", text.strip("\\/")) if text.strip("\\/") else []


PATH_WHITESPACE_MODES = ("keep", "fail")


def whitespace_kinds(value: str) -> list[str]:
    """The kinds of whitespace a path carries that a folder name's single
    inner space is not: ``double space`` (a run of spaces), ``leading`` /
    ``trailing`` (at the value's ends or beside a separator), ``tab``, ``other``
    (any other non-line-break whitespace). [] for a clean path."""
    kinds: list[str] = []
    if re.search(r" {2,}", value):
        kinds.append("double space")
    if value[:1].isspace() or re.search(r"[\\/][ \t]+\S", value):
        kinds.append("leading")
    if value[-1:].isspace() or re.search(r"\S[ \t]+[\\/]", value):
        kinds.append("trailing")
    if "\t" in value:
        kinds.append("tab")
    if re.search(r"[^\S \t\r\n]", value):
        kinds.append("other")
    return kinds


def path_violations(value: str, cap: int, whitespace: str = "fail") -> list[str]:
    """The folder-path rule. ``whitespace`` = the profile's ``path_whitespace``:
    ``fail`` = any whitespace is a violation; ``keep`` = none is (the cell's
    ``iig_path_whitespace`` flag names it instead)."""
    problems: list[str] = []
    if "\n" in value or "\r" in value:
        problems.append("line break")
    if whitespace != "keep" and re.search(r"\s", value):
        problems.append("whitespace")
    if len(value) > cap:
        problems.append(f"length {len(value)} > {cap}")
    if re.search(r"(?<!:)//|\\\\", value.replace("://", "", 1)):
        problems.append("doubled separator")
    for segment in path_segments(value):
        if segment == "":
            problems.append("empty segment")
        elif segment[-1] in _TRAILING_PUNCT:
            problems.append(f"segment {segment!r} ends in punctuation")
    return problems


def join_path(*parts: str | None, separator: str = "/") -> str:
    """Assemble a path from NORMALIZED segments: every part is split on both
    separator styles, blanks are dropped, each segment is stripped, and the
    result carries exactly one separator between segments (a scheme prefix
    such as ``abfss://host`` is kept as the first segment). Segments are not
    otherwise altered — a segment ending in punctuation is the gate's job."""
    segments: list[str] = []
    for part in parts:
        if not part:
            continue
        text = str(part).strip()
        if "://" in text:
            scheme, _, rest = text.partition("://")
            host, _, tail = rest.partition("/")
            segments.append(f"{scheme}://{host}")
            text = tail
        segments.extend(s.strip() for s in _PATH_SEPARATORS.split(text) if s.strip())
    return separator.join(segments)


def normalise_path(value: str, *, lowercase: bool = False) -> tuple[str, list[str]]:
    """(normalised value, what changed) for a DERIVED landing / target path.

    Input documents carry punctuation a path cannot (first ACFC run: an FRD
    ADLS Location ending ``…/accum.``, ``inbound//vendor``): every segment is
    stripped of trailing ``.,;:``, a doubled or backslash separator becomes one
    ``/``, a segment that was only punctuation (or blank) is dropped, and —
    when the convention says so — the segments are lowercased. Whitespace is
    never touched (post57): the path keeps it byte for byte and the profile's
    ``path_whitespace`` decides (flag or fail). A location URI keeps its scheme
    and authority verbatim; only its path is normalised. The leading / trailing separator is kept.
    Nothing changed → the value is returned exactly as given, with ``[]``,
    so a clean path is byte-identical. The caller flags ``path_normalised``
    with the before / after; the gate then checks the normalised value."""
    text = value
    head, path = "", text
    if location_scheme(text):
        _scheme, _authority, path = split_location(text)
        head = text[: len(text) - len(path)] if path and text.endswith(path) else text
        if not path:
            return value, []
    changes: list[str] = []
    if "\\" in path:
        changes.append("backslash separator(s) made '/'")
    if re.search(r"[\\/]{2,}", path):
        changes.append("doubled separator collapsed")
    inner = path.strip("\\/")
    lead, trail = path[:1] in ("/", "\\"), path[-1:] in ("/", "\\") and bool(inner)
    segments: list[str] = []
    for raw in re.split(r"[\\/]", inner) if inner else []:
        segment = raw                  # post57: whitespace kept byte for byte
        if segment.strip() == "":
            if segment:
                changes.append(f"blank segment {raw!r} dropped")
            continue                   # a doubled separator / blank segment (noted above)
        stripped = segment.rstrip(_TRAILING_PUNCT)
        if stripped.strip() == "":
            changes.append(f"punctuation-only segment {segment!r} dropped")
            continue
        if stripped != segment:
            changes.append(f"trailing punctuation stripped from {segment!r}")
        if lowercase and stripped != stripped.lower():
            changes.append(f"segment {stripped!r} lowercased (convention)")
            stripped = stripped.lower()
        segments.append(stripped)
    if not changes:
        return value, []
    return head + ("/" if lead else "") + "/".join(segments) + ("/" if trail else ""), \
        list(dict.fromkeys(changes))


def path_normalised_flag(where: str, before: str, after: str, changes: list[str]) -> str:
    """The ``path_normalised`` flag: the cell / literal, before and after."""
    return (f"path_normalised:{where} — before {before!r} after {after!r} "
            f"({'; '.join(changes)})")


# ----------------------------------------------------------------- names ---


def short_token(value: str | None, max_length: int) -> str | None:
    """The value as a name token when it IS one (single line, no longer than
    the cap after sanitizing); None for prose, which never becomes a name."""
    if not value or "\n" in value or "\r" in value:
        return None
    token = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    if not token or len(token) > max_length:
        return None
    return token


# ------------------------------------------------------------- IIG cells ---


def _is_path_column(header: str, cfg: DerivationsConfig) -> bool:
    upper = header.upper()
    return any(upper.endswith(s.upper()) for s in cfg.path_column_suffixes)


def violation_kind(problem: str) -> str:
    """A path / URI / name problem phrase as its rule KIND (D1 post51 pair_01:
    counted per kind for the coordinates tier) — never the value it quotes."""
    if problem.startswith("length "):
        return "length"
    if problem.startswith("segment ") and problem.endswith("ends in punctuation"):
        return "segment_punctuation"
    if problem.startswith("authority ") and problem.endswith("ends in punctuation"):
        return "authority_punctuation"
    if problem.startswith("no authority"):
        return "no_authority"
    kind = re.sub(r"[^a-z0-9]+", "_", problem.lower()).strip("_")
    return kind if re.fullmatch(r"[a-z][a-z0-9_]{0,40}", kind or "") else "other"


def _check_cell(header: str, value: str, where: str, cfg, pattern, problems: list[str],
                kinds: Counter | None = None, whitespace: str = "fail") -> int:
    """One value (or one segment of a list column) against its column's rule;
    returns 1 when it was a location URI (counted in the details)."""
    kinds = kinds if kinds is not None else Counter()
    if header in cfg.name_columns:
        cap = cfg.name_max_length.get(header, cfg.name_max_length.get("default", 128))
        charset = cfg.name_patterns.get(header, cfg.name_pattern)
        if not (pattern if charset == cfg.name_pattern else re.compile(charset)).match(value):
            problems.append(f"{where}: {_truncate(value)!r} is not an identifier ({charset})")
            kinds["not_identifier"] += 1
        if len(value) > cap:
            problems.append(f"{where}: length {len(value)} > {cap}")
            kinds["name_length"] += 1
    elif _is_path_column(header, cfg):
        # M10.2: which rule applied is part of the finding.
        if location_scheme(value):
            for problem in uri_violations(value, cfg.path_max_length):
                problems.append(f"{where}: {problem} in {_truncate(value)!r} "
                                "(location URI rule)")
                kinds["uri_" + violation_kind(problem)] += 1
            return 1
        for problem in path_violations(value, cfg.path_max_length, whitespace):
            problems.append(f"{where}: {problem} in {_truncate(value)!r}")
            kinds["path_" + violation_kind(problem)] += 1
    return 0


def check_iig_derivations(payload: dict, config: Config,
                          path_whitespace: str = "keep") -> GateCheck:
    """Every cell of the framework config rows: single line; identifier
    columns match the charset + cap; path columns follow the path rules
    (whitespace by ``path_whitespace``: keep = never a violation)."""
    cfg = config.gate.derivations
    pattern = re.compile(cfg.name_pattern)
    problems: list[str] = []
    kinds: Counter = Counter()
    uris = 0
    for tab_name, tab in payload.get("tabs", {}).items():
        for row_index, row in enumerate(tab.get("rows", []), start=2):
            for header, value in row.get("values", {}).items():
                if value in ("", None) or not isinstance(value, str):
                    continue
                badge = row.get("badges", {}).get(header, {})
                source = badge.get("badge", "?")
                if badge.get("tooltip"):
                    source += f" — {badge['tooltip']}"
                where = f"{tab_name}!{header} row {row_index} ({source})"
                if "\n" in value or "\r" in value:
                    address = _cell_address(tab, header, row_index)
                    problems.append(f"{where}: line break"
                                    + (f" at {tab_name}!{address}" if address else "")
                                    + f" in {_truncate(value)!r}")
                    kinds["line_break"] += 1
                    continue
                # 2026-10-14: a list column's comma-separated value (the corpus's
                # two-notebook form) is checked segment by segment
                segments = ([s.strip() for s in value.split(",")]
                            if header in cfg.list_columns and "," in value else [value])
                if len(segments) > 1 and not all(segments):
                    problems.append(f"{where}: an empty segment in the list "
                                    f"{_truncate(value)!r}")
                    kinds["list_empty_segment"] += 1
                    continue
                for segment in segments:
                    uris += _check_cell(header, segment, where, cfg, pattern, problems, kinds,
                                        path_whitespace)
    if problems:
        return GateCheck(name="derivations", passed=False, details="; ".join(problems),
                         violations=dict(sorted(kinds.items())))
    details = "every derived name and path is single-line and within its cap"
    if uris:
        details += f"; {uris} location URI cell(s) checked by the location URI rule"
    return GateCheck(name="derivations", passed=True, details=details)


# ------------------------------------------------ path whitespace (post57) ---


def path_whitespace_flag(where: str, kinds: list[str]) -> str:
    return (f"iig_path_whitespace:{where} — {', '.join(kinds)}; kept as the source "
            "document states it (BSA to confirm)")


def path_whitespace_flags(payload: dict, cfg: DerivationsConfig) -> list[str]:
    """D1 post55 pair_01: ONE flag per path cell (``<SHEET>!<A1> (<COLUMN>)``)
    whose value carries whitespace a folder name's single inner space is not —
    the cell itself is never changed (``path_whitespace: keep``). A location
    URI keeps its own rule; a list column is read segment by segment."""
    flags: list[str] = []
    for tab_name, tab in payload.get("tabs", {}).items():
        for row_index, row in enumerate(tab.get("rows", []), start=2):
            for header, value in row.get("values", {}).items():
                if not isinstance(value, str) or not value or not _is_path_column(header, cfg):
                    continue
                parts = value.split(",") if header in cfg.list_columns else [value]
                if len(parts) > 1:
                    parts = [p[1:] if p.startswith(" ") and i else p
                             for i, p in enumerate(parts)]   # the list's own ", "
                kinds = list(dict.fromkeys(k for p in parts if not location_scheme(p)
                                           for k in whitespace_kinds(p)))
                if kinds:
                    address = _cell_address(tab, header, row_index) or f"row {row_index}"
                    flags.append(path_whitespace_flag(f"{tab_name}!{address} ({header})",
                                                      kinds))
    return flags


# ------------------------------------------------- single-line IIG cells ---
# 2026-10-13 (pair_12, CAQH on D1): a two-line FRD cell reached 22 SUBDOMAIN
# cells. Every IIG cell is ONE line: the payload is normalised at the emitter
# boundary (before any workbook or metadata_inserts.sql is written), each
# change a flag naming the cell; the written workbooks are then checked, and a
# line break that slips through FAILS with its address.

_LINE_BREAK_RE = re.compile(r"[ \t]*(?:\r\n|\r|\n)+[ \t]*")


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, rest = divmod(index - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


def _cell_address(tab: dict, header: str, row_index: int) -> str | None:
    headers = tab.get("headers") or []
    if header not in headers:
        return None
    return f"{_column_letter(headers.index(header) + 1)}{row_index}"


def normalise_iig_line_breaks(payload: dict) -> list[str]:
    """Rewrite every payload cell holding a line break as one line (the break
    and the blanks around it become one space), in place; one flag per cell:
    ``iig_line_break_normalised:<SHEET>!<A1> (<COLUMN>) — <value> written as
    one line``. Row 1 of every sheet is the header row."""
    flags: list[str] = []
    for tab_name, tab in payload.get("tabs", {}).items():
        for row_index, row in enumerate(tab.get("rows", []), start=2):
            values = row.get("values", {})
            for header, value in list(values.items()):
                if not isinstance(value, str) or not ("\n" in value or "\r" in value):
                    continue
                values[header] = _LINE_BREAK_RE.sub(" ", value).strip()
                address = _cell_address(tab, header, row_index) or f"row {row_index}"
                flags.append(f"iig_line_break_normalised:{tab_name}!{address} ({header}) — "
                             f"'{_truncate(value)}' written as one line")
    return flags


def check_iig_workbooks_single_line(paths: list[Path]) -> GateCheck:
    """Every string cell of the written IIG workbooks is one line; a line
    break FAILS with ``<file> <SHEET>!<A1> (<header>)``. Sheets whose name
    starts with ``_`` (provenance / inputs) are not IIG sheets."""
    from openpyxl import load_workbook

    problems: list[str] = []
    cells = 0
    for path in paths:
        workbook = load_workbook(path, read_only=True)
        try:
            for sheet in workbook.worksheets:
                if sheet.title.startswith("_"):
                    continue
                headers: list = []
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if row_index == 1:
                        headers = list(row)
                    for col_index, value in enumerate(row, start=1):
                        if not isinstance(value, str):
                            continue
                        cells += 1
                        if "\n" in value or "\r" in value:
                            header = (headers[col_index - 1] if col_index <= len(headers)
                                      else None)
                            problems.append(
                                f"{path.name} {sheet.title}!{_column_letter(col_index)}"
                                f"{row_index} ({header}): line break in '{_truncate(value)}'")
        finally:
            workbook.close()
    if problems:
        return GateCheck(name="iig_single_line", passed=False, details="; ".join(problems))
    names = ", ".join(p.name for p in paths)
    return GateCheck(name="iig_single_line", passed=True,
                     details=f"{cells} text cell(s) in {names}: every one single-line")


# ------------------------------------------------ single-value IIG cells ---
# 2026-10-14 (D1 pair_12): the FRD's whole Region / LOB table reached every
# ADLS row's LOB cell, comma-joined. A column holding ONE code
# (gate.derivations.single_value_columns) never holds a list or a table's
# header words: the written workbooks are checked, a violation FAILS the run
# with its address.

_LIST_SEPARATORS = re.compile(r"[,;\r\n]")


def _header_word(value: str, words: list[str]) -> str | None:
    for word in words:
        # a word on its own ('LOB_A' is a code, not the header 'LOB'); a word
        # ending in a mark ('REG#') may run into its number ('REG#1')
        after = r"(?![A-Za-z0-9_])" if word[-1:].isalnum() else r"(?![A-Za-z_])"
        if re.search(rf"(?<![A-Za-z0-9_]){re.escape(word)}{after}", value, re.IGNORECASE):
            return word
    return None


def single_value_problem(header: str, value: str, cfg: DerivationsConfig) -> str | None:
    """Why ``value`` cannot sit in the single-code column ``header``, else None."""
    import fnmatch

    if not isinstance(value, str) or not value.strip() or not any(
            fnmatch.fnmatchcase(header, p) for p in cfg.single_value_columns):
        return None
    if _LIST_SEPARATORS.search(value):
        return "a list (comma / semicolon / line break)"
    word = _header_word(value, [*cfg.single_value_header_words, header.replace("_", " ")])
    if word is not None:
        return f"a header word ({word!r})"
    return None


def check_iig_workbooks_single_value(paths: list[Path], cfg: DerivationsConfig) -> GateCheck:
    """Every cell of a single-code column (LOB, CLAIM_TYPE*, the flags ...) of
    the written IIG workbooks holds one value; a list or a header word FAILS
    with ``<file> <SHEET>!<A1> (<header>)``. ``_``-sheets are not IIG sheets."""
    from openpyxl import load_workbook

    problems: list[str] = []
    cells = 0
    for path in paths:
        workbook = load_workbook(path, read_only=True)
        try:
            for sheet in workbook.worksheets:
                if sheet.title.startswith("_"):
                    continue
                headers: list = []
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if row_index == 1:
                        headers = [str(h) if h is not None else "" for h in row]
                        continue
                    for col_index, value in enumerate(row, start=1):
                        header = headers[col_index - 1] if col_index <= len(headers) else ""
                        why = single_value_problem(header, value, cfg)
                        if header and isinstance(value, str) and value.strip():
                            cells += 1 if why is None else 0
                        if why is not None:
                            problems.append(
                                f"{path.name} {sheet.title}!{_column_letter(col_index)}"
                                f"{row_index} ({header}): {why} in '{_truncate(value)}'")
        finally:
            workbook.close()
    if problems:
        return GateCheck(name="iig_single_value", passed=False,
                         details="a single-code column holds more than one value: "
                                 + "; ".join(problems))
    return GateCheck(name="iig_single_value", passed=True,
                     details="every single-code cell (" + ", ".join(cfg.single_value_columns)
                             + ") holds one value")


# ---------------------------------------------------------- SQL literals ---


def literal_violations(literal: str, *, is_path: bool, cfg: DerivationsConfig,
                       path_whitespace: str = "fail") -> list[str]:
    problems: list[str] = []
    if "\n" in literal or "\r" in literal:
        problems.append("line break")
    if literal != literal.strip():
        problems.append("leading/trailing whitespace")
    if literal.replace("''", "").count("'"):
        problems.append("unbalanced quote")
    if is_path:
        problems.extend(uri_violations(literal, cfg.path_max_length)
                        if location_scheme(literal)
                        else path_violations(literal, cfg.path_max_length, path_whitespace))
    return problems


def check_sql_literals(ddl_texts: list[tuple[str, str]], config: Config,
                       path_whitespace: str = "keep") -> GateCheck:
    """COMMENT / LOCATION / TBLPROPERTIES literals of every emitted DDL text
    (a LOCATION's inner whitespace by ``path_whitespace``)."""
    cfg = config.gate.derivations
    problems: list[str] = []
    for name, text in ddl_texts:
        for match in _LITERAL_RE.finditer(text):
            kind, literal = match.group(1).upper(), match.group(2)
            for problem in literal_violations(literal, is_path=kind == "LOCATION", cfg=cfg,
                                              path_whitespace=path_whitespace):
                problems.append(f"{name}: {kind} literal {_truncate(literal)!r}: {problem}")
        for match in _TBLPROP_RE.finditer(text):
            for literal in match.groups():
                for problem in literal_violations(literal, is_path=False, cfg=cfg):
                    problems.append(f"{name}: TBLPROPERTIES literal {_truncate(literal)!r}: "
                                    f"{problem}")
    return GateCheck(name="sql_literals", passed=not problems,
                     details="; ".join(problems) if problems
                     else "every COMMENT / LOCATION / TBLPROPERTIES literal is single-line")


# --------------------------------------------------------- sibling types ---


def _base_type(dtype: str | None) -> str:
    return re.sub(r"\(.*\)", "", dtype or "").strip().lower()


def sibling_type_flags(spec: ResolvedFeedSpec, config: Config) -> list[str]:
    """One ``sibling_type_mismatch:<suffix>`` flag per suffix group whose
    standard-band base types differ, citing the minority rows' STTM cells.
    Groups: suffixes from config, plus any suffix ``sibling_min_group`` or
    more columns share. The types are transcribed as stated, never altered."""
    cfg = config.gate.derivations
    configured = {s.lower() for s in cfg.sibling_suffixes}
    flags: list[str] = []
    for segment in spec.segments:
        groups: dict[str, list] = {}
        for f in segment.fields:
            name = f.standard_column or f.stage_column
            if not f.standard_datatype or "_" not in name:
                continue
            suffix = "_" + name.rsplit("_", 1)[1].lower()
            groups.setdefault(suffix, []).append(f)
        for suffix, fields in sorted(groups.items()):
            if suffix not in configured and len(fields) < cfg.sibling_min_group:
                continue
            counts = Counter(_base_type(f.standard_datatype) for f in fields)
            if len(counts) < 2:
                continue
            majority = counts.most_common(1)[0][0]
            minority = [f for f in fields if _base_type(f.standard_datatype) != majority]
            cited = ", ".join(
                f"{f.standard_column or f.stage_column} {f.standard_datatype}"
                + (f" ({f.provenance.sheet} row {f.provenance.row})" if f.provenance else "")
                for f in minority)
            flags.append(
                f"sibling_type_mismatch:{suffix} — {segment.stage_table.table}: "
                f"{len(fields)} columns end in {suffix}, {counts[majority]} typed "
                f"{majority}, minority: {cited}; transcribed as stated, confirm with the STTM "
                "author")
    return flags


__all__ = [
    "check_iig_derivations",
    "check_iig_workbooks_single_line",
    "check_iig_workbooks_single_value",
    "check_sql_literals",
    "join_path",
    "literal_violations",
    "normalise_iig_line_breaks",
    "path_whitespace_flag",
    "path_whitespace_flags",
    "whitespace_kinds",
    "normalise_path",
    "path_normalised_flag",
    "path_segments",
    "path_violations",
    "short_token",
    "sibling_type_flags",
]
