"""Gate check ``xlsx_no_formulas`` (P0-5, 2026-10-30): no emitted workbook holds
a formula cell.

Every workbook the agent writes is a transcription of document text; a formula
cell (openpyxl ``data_type == 'f'``) in one is document text Excel would run
(``codegen.xlsx_safe``). The check reopens each written ``.xlsx`` and FAILS
naming ``<file> <SHEET>!<A1> (<header>)`` — coordinates and the header row's
label only, never the cell's text. Always on: a security invariant has no knob.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from codegen.gate.preflight import GateCheck

NAME = "xlsx_no_formulas"


def check_xlsx_no_formulas(paths: list[Path]) -> GateCheck:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    problems: list[str] = []
    cells = 0
    for path in paths:
        # openpyxl 3.1.2 keeps a load_workbook(path) open - read from bytes
        workbook = load_workbook(BytesIO(Path(path).read_bytes()))
        for sheet in workbook.worksheets:
            headers = {cell.column: cell.value for cell in sheet[1]} if sheet.max_row else {}
            for cell in sheet._cells.values():
                cells += 1
                if cell.data_type == "f":
                    header = headers.get(cell.column) if cell.row > 1 else None
                    problems.append(f"{Path(path).name} {sheet.title}!"
                                    f"{get_column_letter(cell.column)}{cell.row} ({header})")
    if problems:
        return GateCheck(name=NAME, passed=False,
                         details="formula cell(s) in an emitted workbook: " + "; ".join(problems))
    names = ", ".join(Path(p).name for p in paths) or "no workbook"
    return GateCheck(name=NAME, passed=True,
                     details=f"{cells} cell(s) in {names}: no formula")
