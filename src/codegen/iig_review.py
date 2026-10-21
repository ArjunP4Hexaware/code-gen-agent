"""IIG-first M3: the two IIG workbooks a framework run hands to the BSA.

* ``<feed>_IIG_REVIEW.xlsx`` — a summary sheet first (one row per sheet /
  column / reason / owner group: cell count, example cell, example
  citation), then the template's sheets with every OPEN cell (blank, or
  filled with a synthetic stand-in, or carrying a derivation note)
  highlighted in its owner's colour and commented with the reason, the
  citation and the owner.
* ``<feed>_IIG.xlsx`` — the clean copy: the template's sheets and columns in
  order, values only (no fills, comments or summary). Its data cells are the
  same values ``config_rows.xlsx`` carries.

Both are built from the payload the framework emitter already produced
(codegen.metadata_sheet / codegen.metadata_template) — this module never
derives a value. The reason of an open cell comes from the cell itself; its
owner from the ``iig_review`` config table (sheet + column, else the
reason's default), so the client reassigns columns without a code change.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from codegen.config import Config, IigReviewConfig

# Both template versions word their synthetic path tooltips this way
# (metadata_template._PATH_TOOLTIP, metadata_sheet's reference-workbook path).
PATH_TOOLTIP_PREFIX = "synthetic path shape"
SYNTHETIC_BADGES = ("synthetic", "needs_template")
SUMMARY_HEADERS = ["Sheet", "Column", "Reason", "Owner", "Cells", "Example cell",
                   "Example citation"]
COMMENT_AUTHOR = "CodeGen"


@dataclass(frozen=True)
class OpenCell:
    sheet: str
    row: int            # 1-based worksheet row (header = row 1)
    col: int            # 1-based worksheet column
    column: str         # header
    reason: str
    owner: str
    citation: str

    @property
    def ref(self) -> str:
        return f"{get_column_letter(self.col)}{self.row}"


def _blank(value) -> bool:
    return value is None or value == ""


def _reason(column: str, value, entry: dict, always_blank: set[str],
            review: IigReviewConfig) -> str | None:
    """Why a cell needs a person, or None when an input states it."""
    badge = entry.get("badge")
    tooltip = entry.get("tooltip") or ""
    if _blank(value):
        if entry.get("deliberate_blank"):
            return None                  # a family convention leaves it blank (e.g. LOB)
        if column in review.audit_date_columns:
            return "audit_date"
        if column in review.audit_by_columns:
            return "audit_by"
        if column in always_blank:
            return "framework_assigned"
        return "unstated"
    if entry.get("note") or entry.get("path_note"):
        return "flagged_note"
    if badge in SYNTHETIC_BADGES:
        return "synthetic_path" if tooltip.startswith(PATH_TOOLTIP_PREFIX) else \
            "template_constant"
    return None


def _citation(column: str, reason: str, entry: dict, template: str) -> str:
    tooltip = (entry.get("tooltip") or "").strip()
    if reason == "flagged_note":
        return "; ".join(n for n in (entry.get("note"), entry.get("path_note")) if n)
    if reason in ("synthetic_path", "template_constant"):
        return tooltip or f"template {template} constant"
    if reason == "framework_assigned":
        return (tooltip or "assigned by the ACFC framework / manual by client instruction") \
            + f" (template {template}: always_blank)"
    if reason == "audit_date":
        return f"set when the config tables are loaded (template {template}: always_blank)"
    if reason == "audit_by":
        return ("RFC<rfc_number> once the load-pattern FAQ answers rfc_number; no FRD field "
                "or STTM cell states it")
    return f"no FRD field, STTM cell or FAQ answer states {column} (template {template})"


def open_cells(payload: dict, always_blank: list[str], config: Config,
               template: str) -> list[OpenCell]:
    """Every open cell of the payload, in sheet / row / column order."""
    review = config.iig_review
    blank_set = set(always_blank)
    cells: list[OpenCell] = []
    for sheet, tab in payload["tabs"].items():
        headers = tab["headers"]
        for row_index, row in enumerate(tab["rows"], start=2):
            for col_index, column in enumerate(headers, start=1):
                entry = row["badges"].get(column) or {}
                reason = _reason(column, row["values"].get(column), entry, blank_set, review)
                if reason is None:
                    continue
                cells.append(OpenCell(
                    sheet=sheet, row=row_index, col=col_index, column=column, reason=reason,
                    owner=review.owner_for(sheet, column, reason),
                    citation=_citation(column, reason, entry, template)))
    return cells


def dq_review_entries(payload: dict, config: Config) -> list[dict]:
    """One 'additional DQ rules — Engineer' summary entry per FILE (decision
    2026-10-07): only the header/trailer split and the STTM-derived rules are
    generated, so the DQ sheet is visibly — never silently — incomplete. The
    files are the ADLS_DELTA_INGESTION_DETAILS rows (OBJECT_ID, SRC_FILE_NAME);
    no cell to fill, so Cells = 0."""
    tabs = payload.get("tabs", {})
    if "DATA_QUALITY_RULES" not in tabs or "ADLS_DELTA_INGESTION_DETAILS" not in tabs:
        return []
    review = config.iig_review
    entries = []
    for row in tabs["ADLS_DELTA_INGESTION_DETAILS"]["rows"]:
        values = row["values"]
        name = values.get("SRC_FILE_NAME") or values.get("OBJECT_NAME") or "?"
        entries.append({
            "sheet": "DATA_QUALITY_RULES", "column": "(rule rows)",
            "reason": "additional_dq_rules",
            "owner": review.owner_for("DATA_QUALITY_RULES", "(rule rows)",
                                      "additional_dq_rules"),
            "cells": 0, "example_cell": "",
            "example_citation": (f"file {name} (OBJECT_ID {values.get('OBJECT_ID') or '?'}): "
                                 "only the header/trailer split and the rules the STTM states "
                                 "are generated — add the feed's other DQ rule classes")})
    return entries


def review_groups(cells: list[OpenCell]) -> list[dict]:
    """One entry per (sheet, column, reason, owner), first-seen order."""
    groups: dict[tuple[str, str, str, str], dict] = {}
    for cell in cells:
        key = (cell.sheet, cell.column, cell.reason, cell.owner)
        group = groups.get(key)
        if group is None:
            groups[key] = {"sheet": cell.sheet, "column": cell.column, "reason": cell.reason,
                           "owner": cell.owner, "cells": 1, "example_cell": cell.ref,
                           "example_citation": cell.citation}
        else:
            group["cells"] += 1
    return list(groups.values())


def _template_sheets(workbook: Workbook, payload: dict) -> None:
    for sheet_name, tab in payload["tabs"].items():
        sheet = workbook.create_sheet(title=sheet_name)
        headers = tab["headers"]
        sheet.append(headers)
        for row in tab["rows"]:
            sheet.append([row["values"][h] for h in headers])


def _new_workbook() -> Workbook:
    from codegen.metadata_sheet import _pin_workbook_properties

    workbook = Workbook()
    _pin_workbook_properties(workbook)
    workbook.remove(workbook.active)
    return workbook


def clean_workbook(payload: dict) -> Workbook:
    """The template's sheets and columns, values only."""
    workbook = _new_workbook()
    _template_sheets(workbook, payload)
    return workbook


def review_workbook(payload: dict, cells: list[OpenCell], config: Config,
                    empty_by_rule: dict[str, str] | None = None) -> Workbook:
    """Summary sheet first, then the template sheets with open cells
    highlighted by owner and commented (reason, citation, owner). post89: a
    sheet in ``empty_by_rule`` ({sheet: reason}) gets one summary row saying why
    it has no rows."""
    review = config.iig_review
    workbook = _new_workbook()
    summary = workbook.create_sheet(title=review.summary_sheet)
    summary.append(SUMMARY_HEADERS)
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for group in [*review_groups(cells), *dq_review_entries(payload, config)]:
        summary.append([group["sheet"], group["column"], review.reason_labels[group["reason"]],
                        review.owner_labels[group["owner"]], group["cells"],
                        group["example_cell"], group["example_citation"]])
    for sheet, reason in (empty_by_rule or {}).items():
        summary.append([sheet, "", review.empty_by_rule_label, review.empty_by_rule_owner, 0,
                        "", review.empty_by_rule_reasons.get(reason, reason)])
    _template_sheets(workbook, payload)
    for cell in cells:
        target = workbook[cell.sheet].cell(row=cell.row, column=cell.col)
        fill = review.owner_fills[cell.owner]
        target.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
        target.comment = Comment(
            f"Reason: {review.reason_labels[cell.reason]}\n"
            f"Citation: {cell.citation}\n"
            f"Owner: {review.owner_labels[cell.owner]}", COMMENT_AUTHOR)
    return workbook


def write_iig_workbooks(payload: dict, always_blank: list[str], config: Config,
                        framework_dir: Path, feed: str, template: str,
                        empty_by_rule: dict[str, str] | None = None,
                        ) -> tuple[Path, Path, list[OpenCell]]:
    """Write the review copy and the clean copy; return their paths + the
    open cells (for the report / tests)."""
    from codegen.metadata_sheet import stable_workbook_bytes

    review = config.iig_review
    cells = open_cells(payload, always_blank, config, template)
    review_path = framework_dir / review.review_file_pattern.format(feed=feed)
    clean_path = framework_dir / review.clean_file_pattern.format(feed=feed)
    review_path.write_bytes(stable_workbook_bytes(review_workbook(payload, cells, config,
                                                                  empty_by_rule)))
    clean_path.write_bytes(stable_workbook_bytes(clean_workbook(payload)))
    return review_path, clean_path, cells


# --------------------------------------------------------- cell counts ---
# 2026-10-15 (D1 post55): the harness's "unpopulated IIG cells" read one number
# per sheet from a workbook it globbed out of a shared folder — two pairs (4
# feeds / 11 columns and 1 feed / 115 columns) printed the same nine numbers.
# Each feed now records its own counts in CODEGEN_STAGES.json, from the payload
# its clean workbook was written from.

# An open cell the inputs panel asks about: its NOTE is an owed-cell /
# environment NOTE (each names an INPUT_NEEDED key; codegen.owed_inputs).
ASKED_NOTE_KINDS = ("iig_owed", "iig_owed_run", "iig_env", "iig_env_placeholder")
_INPUT_RE = re.compile(r"INPUT_NEEDED ((?:run|feeds)\.[\w.<>]+?)(?=[:\s]|$)")
_ENV_REF_RE = re.compile(r"states no ([\w.]+)")


def asked_input(note: str | None, feed_id: str) -> str | None:
    """The inputs-panel key an open cell's NOTE asks for, or None when the
    NOTE is no owed / environment NOTE (``run.email_to``,
    ``feeds.<feed>.active_start_date``, ``run.environment.storage_account``)."""
    text = note or ""
    if text.split(":", 1)[0] not in ASKED_NOTE_KINDS:
        return None
    match = _INPUT_RE.search(text)
    if match:
        return match.group(1).replace("<feed>", feed_id)
    ref = _ENV_REF_RE.search(text)
    return f"run.environment.{ref.group(1)}" if ref else "run.environment"


def used_range_blanks(path: Path) -> dict[str, int]:
    """Per sheet: the empty cells of the sheet's used range (rows 1..max_row x
    columns 1..max_column) — the harness's pre-post56 "unpopulated" number,
    kept as ``template_blanks`` for one release."""
    from io import BytesIO

    from openpyxl import load_workbook

    # read from memory: openpyxl 3.1.2 (the D1 pin) keeps the file open
    # otherwise, and the GATE_FAIL quarantine move then fails on Windows
    workbook = load_workbook(BytesIO(Path(path).read_bytes()))
    try:
        return {ws.title: sum(1 for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                                                           max_col=ws.max_column)
                              for cell in row if _blank(cell.value))
                for ws in workbook.worksheets}
    finally:
        workbook.close()


def iig_cell_counts(payload: dict, feed_id: str, clean_path: Path | None = None) -> dict:
    """One feed's IIG cell counts, per sheet, over the data rows the generator
    WROTE (the payload rows, up to the last row with a filled cell): an empty
    cell is ``open_for_bsa`` when its NOTE asks the inputs panel for it
    (:func:`asked_input`), else ``empty_other``; ``template_blanks`` is the
    old used-range count of the written clean workbook. ``open_inputs`` lists
    the inputs behind the open_for_bsa cells (one input may open several: a
    status row each, a flow's schedule + notebook rows)."""
    old = used_range_blanks(clean_path) if clean_path is not None and clean_path.is_file() \
        else {}
    sheets: dict[str, dict[str, int]] = {}
    inputs: set[str] = set()
    for sheet, tab in payload.get("tabs", {}).items():
        headers = tab["headers"]
        rows = list(tab["rows"])
        while rows and all(_blank(rows[-1]["values"].get(h)) for h in headers):
            rows.pop()
        asked = other = 0
        for row in rows:
            for header in headers:
                if not _blank(row["values"].get(header)):
                    continue
                key = asked_input((row["badges"].get(header) or {}).get("note"), feed_id)
                if key is None:
                    other += 1
                else:
                    asked += 1
                    inputs.add(key)
        sheets[sheet] = {"rows": len(rows), "open_for_bsa": asked, "empty_other": other,
                         "template_blanks": old.get(sheet, 0)}
    return {"sheets": sheets, "open_inputs": sorted(inputs),
            "totals": {k: sum(s[k] for s in sheets.values())
                       for k in ("open_for_bsa", "empty_other", "template_blanks")}}


__all__ = ["ASKED_NOTE_KINDS", "OpenCell", "asked_input", "clean_workbook",
           "dq_review_entries", "iig_cell_counts", "open_cells", "review_groups",
           "review_workbook", "used_range_blanks", "write_iig_workbooks"]
