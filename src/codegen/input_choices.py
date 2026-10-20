"""The BSA's choice for an owed input that has a default (post59).

PIPELINE_NAME (one per flow row: grand_master / master / file_to_stg /
stg_to_std) and the defaulted run inputs (APPLICATION_NAME, EMAIL_TO and the
email status set, when a default exists) are a three-way choice:

- ``use_default`` — the agent's proposal: for PIPELINE_NAME the conventions
  profile's per-flow template (``pipeline_name_templates``) rendered with the
  feed's {domain} {subdomain} {source} {frequency} {layer} {load_mode}; for the
  others the owed rule's ``default``;
- ``naming_standard`` — the BSA's own template in the same placeholder syntax
  (for PIPELINE_NAME one template for the feed renders all four names);
- ``leave_blank`` — the cell stays empty.

Proposed is not applied (Aruna's rule holds for every unattended run): an
input whose choice is not stated explicitly — in the answers file
(``choice:`` / ``standard:``) or by the BSA's Confirm in the App — is written
BLANK, and its INPUT_NEEDED line shows the proposal and its template. A
default that collides with a PIPELINE_NAME of the loaded goldens / corpus
(``pipeline_name_corpus``) is never proposed: the line asks instead.
use_default and naming_standard values carry the "taken from other documents"
gate flag (``frd_unstated:<TAB>.<HEADER> source_used:...``). CODEGEN_STAGES.json
records each input's ``source``: default | standard | blank | answer (a value
typed into the answers file).

post61 — two kinds of default (``Defaulted.kind``):

- ``profile_default``: the value comes from the conventions profile or the
  ACFC overlay — the email status set (``email_status_default`` / the rule's
  ``default``), EMAIL_TO, the notebook NAME per leg (``corpus_cells``
  DATABRICKS_NOTEBOOK_NAME by notebook kind) and the ADLS SRC_FILE_NAME style
  (``feed_overrides.<feed>.src_file_name`` / ``family_conventions``). Applied
  when unanswered — ``source`` profile, no gate flag, no open line — and still
  shown (CLI NOTE line, answers template, the App's panel) as changeable.
- ``agent_proposal``: the value is rendered or derived by the agent —
  PIPELINE_NAME from its template, APPLICATION_NAME from {source} (and
  ACTIVE_START_DATE, which has none): shown, never applied without the BSA's
  confirmation (Aruna's rule above).

An owed rule's kind is its ``default_kind``; unset, a default with a {token}
is an agent_proposal, a literal one a profile_default.

Used by the CLI (``codegen generate``) and the App."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

CHOICES = ("use_default", "naming_standard", "leave_blank")
SOURCE_OF = {"use_default": "default", "naming_standard": "standard", "leave_blank": "blank"}
PLACEHOLDERS = ("domain", "subdomain", "source", "frequency", "layer", "load_mode",
                # post79: the IIG's own cell values, so a template can follow the prod
                # convention (docs/acfc/OVERLAY_FIELDS.md, pipeline_name_templates)
                "lob", "object", "object_code", "src_format", "source_code",
                "schema_domain", "frequency_code")
FLOWS = ("grand_master", "master", "file_to_stg", "stg_to_std")
PIPELINE = "pipeline_name"
PROFILE_DEFAULT = "profile_default"
AGENT_PROPOSAL = "agent_proposal"
# post61: the profile defaults that are not owed cells
NOTEBOOK = "notebook_name"
NOTEBOOK_TAB = "DATABRICKS_NOTEBOOK_DETAILS"
NOTEBOOK_HEADER = "DATABRICKS_NOTEBOOK_NAME"
SRC_FILE_NAME_STYLE = "src_file_name_style"
STYLES = ("pattern", "prefix_star", "prefix_middle")
# the evidence prefix a chosen FAQ answer carries (metadata_template._owed reads it)
EVIDENCE = "choice:"
_PLACEHOLDER_RE = re.compile(r"\{([^{}]*)\}")
_NOT_NAME_RE = re.compile(r"[^A-Za-z0-9_]+")


@dataclass(frozen=True)
class Choice:
    """An explicit choice for one input (answers file / App Confirm)."""

    choice: str
    standard: str = ""


@dataclass
class Proposal:
    """What use_default would write: the rendered value and its template;
    ``refused`` = why there is none (no template, a collision)."""

    value: str = ""
    template: str = ""
    refused: str = ""

    @property
    def placeholder(self) -> bool:
        """Rendered from a SYNTHETIC template (``SYN_``…): no naming standard
        is loaded — the shipped placeholder, never the client's."""
        return is_placeholder_template(self.template)


def default_kind(rule) -> str:
    """An owed rule's kind of default (post61): its ``default_kind``, else a
    default with a {token} = agent_proposal, a literal one = profile_default."""
    kind = getattr(rule, "default_kind", None)
    if kind:
        return kind
    default = str(getattr(rule, "default", "") or "")
    return AGENT_PROPOSAL if _PLACEHOLDER_RE.search(default) else PROFILE_DEFAULT


SYNTHETIC_PREFIX = "SYN_"
PLACEHOLDER_LABEL = "no naming standard loaded — placeholder"


def is_placeholder_template(template: str) -> bool:
    """A template whose name starts with ``SYN_`` (config/config.yaml ships
    only these; the real ones come from codegen-state/overlays/acfc_prx.yaml)."""
    return str(template or "").strip().upper().startswith(SYNTHETIC_PREFIX)


@dataclass
class Resolved:
    """One input's outcome for a feed: the FAQ answer to apply (None = not
    applied: the cell stays blank and the input stays open) and the manifest
    record."""

    name: str                     # the FAQ companion (pipeline_name_file_to_stg)
    key: str                      # the input key (pipeline_name.file_to_stg / application_name)
    source: str                   # default | standard | blank | answer
    chosen: bool
    value: str = ""
    template: str = ""
    proposal: Proposal = field(default_factory=Proposal)
    note: str = ""
    kind: str = AGENT_PROPOSAL

    @property
    def applied(self) -> bool:
        """A profile default written unanswered (post61): not an open input."""
        return self.source == "profile" and bool(self.value)


def parse_choice(key: str, entry) -> Choice | None:
    """An answers-file ``inputs`` entry's choice (None: no choice stated).
    ValueError on an unknown choice or a naming_standard without a template is
    raised by the caller's validation (:func:`validate_choice`)."""
    if not isinstance(entry, dict):
        return None
    choice = str(entry.get("choice") or "").strip()
    standard = str(entry.get("standard") or "").strip()
    if not choice and not standard:
        return None
    if not choice and standard:
        choice = "naming_standard"
    return Choice(choice, standard)


def validate_choice(key: str, choice: Choice) -> str | None:
    """The problem with ``choice`` (None: valid)."""
    if choice.choice not in CHOICES:
        return f"inputs[{key!r}].choice must be one of {', '.join(CHOICES)}"
    if key.endswith(SRC_FILE_NAME_STYLE):
        # post61: a style, never blank
        if choice.choice == "leave_blank":
            return f"inputs[{key!r}]: the SRC_FILE_NAME style cannot be left blank"
        if choice.standard and choice.standard not in STYLES:
            return f"inputs[{key!r}].standard must be one of {', '.join(STYLES)}"
        return None
    if choice.standard:
        unknown = [p for p in _PLACEHOLDER_RE.findall(choice.standard) if p not in PLACEHOLDERS]
        if unknown:
            return (f"inputs[{key!r}].standard names unknown placeholder(s) {unknown}; "
                    f"use {', '.join('{' + p + '}' for p in PLACEHOLDERS)}")
    return None


def render(template: str, tokens: dict[str, str], *, name: bool = True,
           case: str = "upper") -> str:
    """``template`` with its {placeholders} filled from ``tokens``. ``name``:
    a name — every run of characters outside [A-Za-z0-9_] becomes one ``_``
    (an empty token leaves no doubled ``_``), upper-cased when ``case`` is
    ``upper``; else the text as filled (a list, an address)."""
    filled = _PLACEHOLDER_RE.sub(lambda m: str(tokens.get(m.group(1)) or "").strip(), template)
    if not name:
        return filled.strip()
    text = _NOT_NAME_RE.sub("_", filled)
    text = re.sub(r"_+", "_", text).strip("_")
    return text.upper() if case == "upper" else text


def evidence(source: str, template: str = "") -> str:
    """The FAQ answer's evidence for a chosen value (read back by
    :func:`chosen_of`)."""
    return f"{EVIDENCE}{source}|{template}"


def chosen_of(answer) -> tuple[str, str] | None:
    """(source, template) of a chosen FAQ answer; None for any other answer."""
    text = str(getattr(answer, "evidence", "") or "")
    if not text.startswith(EVIDENCE):
        return None
    source, _bar, template = text[len(EVIDENCE):].partition("|")
    return source, template


def flows(tpl) -> list[str]:
    """The template's flows (its pipeline roles, as the corpus names them)."""
    from codegen.metadata_template import _ROLE_FLOWS

    return list(dict.fromkeys(_ROLE_FLOWS.get(r.role, r.role) for r in tpl.pipeline_roles))


def _strategy(value) -> str:
    return str(getattr(value, "value", value) or "")


def feed_tokens(spec, tpl, faq=None, config=None) -> dict[str, str]:
    """{domain} {subdomain} {source} {frequency} of ``spec`` (no layer), and
    (post79) the IIG's own cells a pipeline name may carry: {lob} (the first
    file's ADLS LOB cell), {object} / {object_code} (its OBJECT_NAME / that
    name's first word), {src_format} (its file extension), {source_code} (the
    SOURCE cell's first word), {schema_domain} (the first word the stage and
    standard schema names share, else the standard schema's first word) and
    {frequency_code} (the frequency's ``engineering_standards.
    frequency_abbreviations`` code). A token the run cannot compute is ''."""
    from codegen.metadata_sheet import _feed_from_spec
    from codegen.metadata_template import _frequency, _source_value

    feed = _feed_from_spec(spec)
    frequency = str(_frequency(feed, faq, tpl)["value"] or "").strip()
    tokens = {"domain": spec.domain or "", "subdomain": spec.sub_domain or "",
              "source": _source_value(feed), "frequency": frequency}
    tokens.update(_cell_tokens(spec, tpl, config, frequency))
    tokens["source_code"] = _words(tokens["source"])[0] if _words(tokens["source"]) else ""
    return tokens


def _words(text) -> list[str]:
    return [w for w in re.split(r"[^A-Za-z0-9]+", str(text or "")) if w]


def _schema_domain(stage, standard) -> str:
    """The word the stage and standard schema names share (the domain code of a
    ``<x>_<DOM>`` / ``<DOM>_<y>`` pair), else the standard schema's first word."""
    first = _words(getattr(stage, "schema_name", ""))
    second = _words(getattr(standard, "schema_name", ""))
    shared = [w for w in first if w.casefold() in {x.casefold() for x in second}]
    return shared[0] if shared else (second[0] if second else (first[0] if first else ""))


def _cell_tokens(spec, tpl, config, frequency: str) -> dict[str, str]:
    from pathlib import PurePosixPath

    segments = list(getattr(spec, "segments", None) or [])
    detail = next((g for g in segments if g.segment == "Detail"), segments[0] if segments else None)
    out = {"schema_domain": _schema_domain(detail.stage_table if detail else None,
                                           spec.standard_table),
           "frequency_code": "", "lob": "", "object": "", "object_code": "", "src_format": ""}
    if config is not None:
        codes = {k.casefold(): v for k, v in
                 config.engineering_standards.frequency_abbreviations.items()}
        out["frequency_code"] = str(codes.get(frequency.casefold(), "") or "")
        from codegen.metadata_template import (
            _feed_files,
            _file_lob_cell,
            _object_name,
            _wildcarded,
        )

        files = _feed_files(spec, config)
        if files:
            pattern = _wildcarded(files[0].pattern, tpl)
            out["object"] = _object_name(pattern, tpl.object_name_strip_tokens,
                                         tpl.object_name_bound_tokens)
            out["object_code"] = (_words(out["object"]) or [""])[0]
            out["src_format"] = PurePosixPath(pattern).suffix.lstrip(".")
            lob = _file_lob_cell(tpl, "ADLS_DELTA_INGESTION_DETAILS", files[0], spec)
            out["lob"] = str(lob.get("value") or "")
    return out


def flow_tokens(base: dict[str, str], spec, profile, flow: str) -> dict[str, str]:
    """``base`` + the flow's {layer} (``pipeline_flow_layers``) and {load_mode}
    (the stage load strategy for file_to_stg, the standard one for stg_to_std;
    none for the master rows)."""
    load = (spec.stage_load_strategy if flow == "file_to_stg" else
            spec.standard_load_strategy if flow == "stg_to_std" else "")
    return {**base, "layer": profile.pipeline_flow_layers.get(flow, ""),
            "load_mode": _strategy(load)}


_CORPUS_CACHE: dict[tuple, frozenset[str]] = {}


def _corpus_files(paths: list[str], base: Path | None = None) -> list[Path]:
    """The corpus files ``paths`` name: a folder's ``*.xlsx`` / ``*.txt``
    (recursively), a file itself; missing paths contribute nothing."""
    files: list[Path] = []
    for raw in paths:
        path = Path(raw)
        if not path.is_absolute() and base is not None:
            path = base / path
        if path.is_dir():
            files += sorted(p for p in path.rglob("*") if p.suffix.lower() in (".xlsx", ".txt"))
        elif path.is_file():
            files.append(path)
    return files


def corpus_pipeline_names(paths: list[str], base: Path | None = None) -> frozenset[str]:
    """Every PIPELINE_NAME (case-folded) in the loaded goldens / corpus: each
    path a folder (its ``*.xlsx``, read recursively), an ``.xlsx`` (every sheet
    with a PIPELINE_NAME header) or a ``.txt`` (one name per line). Missing
    paths contribute nothing. Cached per file state."""
    files = _corpus_files(paths, base)
    stamp = tuple((str(p), p.stat().st_mtime_ns, p.stat().st_size) for p in files)
    if stamp in _CORPUS_CACHE:
        return _CORPUS_CACHE[stamp]
    names: set[str] = set()
    for path in files:
        if path.suffix.lower() == ".txt":
            names |= {ln.strip().casefold() for ln in
                      path.read_text(encoding="utf-8", errors="replace").splitlines()
                      if ln.strip()}
            continue
        names |= _workbook_pipeline_names(path)
    _CORPUS_CACHE[stamp] = frozenset(names)
    return _CORPUS_CACHE[stamp]


def _workbook_pipeline_names(path: Path) -> set[str]:
    from io import BytesIO

    from openpyxl import load_workbook

    try:
        book = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — an unreadable corpus file adds nothing
        return set()
    names: set[str] = set()
    try:
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None) or ()
            columns = [i for i, h in enumerate(header)
                       if str(h or "").strip().upper() == "PIPELINE_NAME"]
            if not columns:
                continue
            for row in rows:
                for i in columns:
                    if i < len(row) and str(row[i] or "").strip():
                        names.add(str(row[i]).strip().casefold())
    finally:
        book.close()
    return names


@dataclass(frozen=True)
class Defaulted:
    """One input that can carry a proposal."""

    name: str          # FAQ companion
    key: str           # input key under inputs.run / inputs.feeds.<feed>
    scope: str         # run | feed
    rule: object       # the owed rule (OwedCellConfig) / the email status rule
    flow: str = ""     # a PIPELINE_NAME flow
    statuses: bool = False
    kind: str = AGENT_PROPOSAL     # post61: profile_default | agent_proposal
    leg: str = ""      # a notebook NAME's flow (file_to_stg / stg_to_std)
    style: bool = False            # the ADLS SRC_FILE_NAME style
    pk: bool = False               # post70: TGT_PRIMARY_KEY (codegen.primary_key)


def _notebook_legs(tpl) -> list[str]:
    """The flows of the template's notebook rows (post61)."""
    from codegen.metadata_template import _ROLE_FLOWS

    rows = tpl.template_rows.get(NOTEBOOK_TAB) or []
    return list(dict.fromkeys(_ROLE_FLOWS.get(str(r.get("_pipeline_role") or ""), "")
                              for r in rows if r.get("_pipeline_role")))


def _has_header(tpl, header: str) -> bool:
    return any(header in tab.headers for tab in tpl.tabs.values())


def defaulted_inputs(tpl) -> list[Defaulted]:
    """The template's inputs that offer the three-way choice: every per-flow
    PIPELINE_NAME, every owed cell with a ``default`` and the email status set;
    post61: the profile defaults — the notebook NAME per leg (when
    ``corpus_cells`` names one) and the SRC_FILE_NAME style."""
    out: dict[str, Defaulted] = {}
    for rules in tpl.owed_cells.values():
        for rule in rules.values():
            if rule.per_flow:
                for flow in flows(tpl):
                    name = f"{rule.faq}_{flow}"
                    out.setdefault(name, Defaulted(name, f"{rule.faq}.{flow}", rule.scope,
                                                   rule, flow))
            elif rule.default:
                out.setdefault(rule.faq, Defaulted(rule.faq, rule.faq, rule.scope, rule,
                                                   kind=default_kind(rule)))
    if tpl.email_status_input is not None:
        rule = tpl.email_status_input
        kind = rule.default_kind or PROFILE_DEFAULT
        out.setdefault(rule.faq, Defaulted(rule.faq, rule.faq, "run", rule, statuses=True,
                                           kind=kind))
    notebook = (tpl.corpus_cells.get(NOTEBOOK_TAB) or {}).get(NOTEBOOK_HEADER)
    if notebook is not None:
        for leg in _notebook_legs(tpl):
            name = f"{NOTEBOOK}_{leg}"
            out.setdefault(name, Defaulted(name, f"{NOTEBOOK}.{leg}", "feed", notebook,
                                           kind=PROFILE_DEFAULT, leg=leg))
    if _has_header(tpl, "SRC_FILE_NAME"):
        out.setdefault(SRC_FILE_NAME_STYLE, Defaulted(
            SRC_FILE_NAME_STYLE, SRC_FILE_NAME_STYLE, "feed", tpl.family_conventions,
            kind=PROFILE_DEFAULT, style=True))
    from codegen import primary_key

    if primary_key.enabled(tpl):
        out.setdefault(primary_key.KEY, Defaulted(primary_key.KEY, primary_key.KEY, "feed",
                                                  None, kind=AGENT_PROPOSAL, pk=True))
    return list(out.values())


def _key_proposal(spec, faq, profile, base: Path | None) -> Proposal:
    """post70: the feed's TGT_PRIMARY_KEY proposal — an explicit key signal
    only (codegen.primary_key); none = refused, never the not-null list."""
    from codegen import primary_key

    corpus = primary_key.corpus_keys(list(getattr(profile, "pipeline_name_corpus", []) or []),
                                     base)
    offer = primary_key.feed_proposal(spec, faq, corpus)
    if not offer:
        return Proposal(refused="no explicit key signal (no Primary Key cell, FRD key "
                                "statement, FAQ natural_key_columns or corpus row)")
    return Proposal(offer.value, offer.source)


def _notebook_proposal(item: Defaulted, spec, tpl, config) -> Proposal:
    """The notebook NAME the profile writes for ``item.leg`` of ``spec``: the
    ``corpus_cells`` value for the feed's notebook kind (post61)."""
    from codegen.config import is_placeholder
    from codegen.metadata_sheet import _feed_from_spec
    from codegen.metadata_template import _corpus_key, _notebook_kind

    rule = item.rule
    if config is None:
        return Proposal(refused="the notebook kind needs the run's config")
    kind = _notebook_kind(item.leg, _feed_from_spec(spec), spec, tpl, config)
    shape = rule.value if rule.by is None else _corpus_key(rule.values, kind)
    template = f"corpus_cells {NOTEBOOK_HEADER} ({rule.by or 'value'} {kind})"
    if not shape:
        return Proposal(template=template, refused=f"the profile names no notebook for {kind}")
    unset = [t for t in _PLACEHOLDER_RE.findall(shape)
             if not str(tpl.corpus_tokens.get(t) or "").strip()]
    value = _PLACEHOLDER_RE.sub(lambda m: str(tpl.corpus_tokens.get(m.group(1)) or "")
                                .strip(), shape)
    if unset or is_placeholder(shape) or is_placeholder(value):
        return Proposal(template=template, refused="a value the D1 overlay sets")
    return Proposal(value, template)


def _style_proposal(spec, tpl) -> Proposal:
    """The SRC_FILE_NAME style the profile applies to ``spec`` (post61)."""
    from codegen.metadata_template import _feed_override

    override = _feed_override(tpl, spec.feed_id)
    if override is not None and override.src_file_name:
        return Proposal(override.src_file_name,
                        f"feed_overrides.{spec.feed_id}.src_file_name")
    return Proposal(tpl.family_conventions.src_file_name, "family_conventions.src_file_name")


def proposal(item: Defaulted, spec, profile, tpl, tokens: dict[str, str],
             taken: frozenset[str], config=None, faq=None, base: Path | None = None) -> Proposal:
    """What use_default would write for ``item`` of ``spec``."""
    from codegen.config import is_placeholder
    from codegen.metadata_template import (
        _owed_default,
        _pick_statuses,
        email_status_default_label,
    )

    if item.leg:
        return _notebook_proposal(item, spec, tpl, config)
    if item.pk:
        return _key_proposal(spec, faq, profile, base)
    if item.style:
        return _style_proposal(spec, tpl)
    if item.statuses:
        statuses = [str(r.get("STATUS") or "") for r in
                    tpl.template_rows.get("EMAIL_TEMPLATE_CONFIG") or []]
        chosen = _pick_statuses(statuses, tpl.email_status_default) or statuses
        return Proposal(", ".join(chosen), email_status_default_label(tpl))
    if item.flow:
        template = profile.pipeline_name_templates.get(item.flow, "")
        if not template:
            return Proposal(refused="the conventions profile has no pipeline_name_templates "
                                    f"entry for {item.flow}")
        if is_placeholder_template(template):
            # post79: a SYN_ placeholder is never a name to write — not even on
            # use_default (the Pyx runs wrote SYN_PL_... names that way)
            return Proposal(template=template, refused=f"{PLACEHOLDER_LABEL}: set "
                            "conventions.profiles.<profile>.pipeline_name_templates in the "
                            "ACFC overlay")
        value = render(template, flow_tokens(tokens, spec, profile, item.flow),
                       case=profile.pipeline_name_case)
        if not value:
            return Proposal(template=template, refused="the template renders empty")
        if value.casefold() in taken:
            return Proposal(template=template,
                            refused="the proposal collides with a PIPELINE_NAME of the loaded "
                                    "goldens / corpus (pipeline_name_corpus)")
        return Proposal(value, template)
    template = str(item.rule.default or "")
    value = _owed_default(item.rule, {**tpl.corpus_tokens, **tokens})
    if not value or is_placeholder(value):
        return Proposal(template=template, refused="no default for this run (a value the "
                                                   "environment's overlay sets)")
    return Proposal(value, template)


def _first(mapping: dict, keys: list[str]):
    return next((mapping[k] for k in keys if k in mapping), None)


def resolve_feed(spec, *, profile, tpl, inputs: dict[str, str],
                 choices: dict[str, Choice], index: int | str = "", faq=None,
                 base: Path | None = None, config=None) -> list[Resolved]:
    """Every defaulted input of ``spec``: the value it gets (an agent proposal
    only on an explicit choice or a typed value; a profile default also when
    unanswered — post61) and its proposal."""
    items = defaulted_inputs(tpl)
    if not items:
        return []
    taken = (corpus_pipeline_names(profile.pipeline_name_corpus, base)
             if any(i.flow for i in items) else frozenset())
    tokens = feed_tokens(spec, tpl, faq, config)
    out: list[Resolved] = []
    for item in items:
        prefixes = (["run."] if item.scope == "run"
                    else [f"feeds.{spec.feed_id}.", f"feeds[{index}]."])
        offer = proposal(item, spec, profile, tpl, tokens, taken, config, faq, base)
        typed = _first(inputs, [p + item.key for p in prefixes])
        if typed is not None and str(typed).strip():
            out.append(Resolved(item.name, item.key, "answer", True, str(typed).strip(),
                                proposal=offer, kind=item.kind))
            continue
        own = _first(choices, [p + item.key for p in prefixes])
        shared = (_first(choices, [p + item.rule.faq for p in prefixes]) if item.flow
                  else None)
        if item.pk and offer.refused and own is None and shared is None:
            continue           # post70: no signal — the blocking QUESTION asks instead
        choice = own or shared
        kind = item.kind
        if item.flow and profile.pipeline_name_autoset and offer.value:
            kind = PROFILE_DEFAULT      # post79: the convention's name, written unasked
        if kind == PROFILE_DEFAULT and (choice is None or choice.choice == "use_default"):
            # post61: a profile default — written unanswered, no gate flag (the
            # writer applies the profile's own value), still changeable
            if offer.value:
                out.append(Resolved(item.name, item.key, "profile", choice is not None,
                                    offer.value, offer.template, offer, kind=kind))
            else:
                out.append(Resolved(item.name, item.key, "blank", False, proposal=offer,
                                    kind=kind))
        elif choice is None:
            out.append(Resolved(item.name, item.key, "blank", False, proposal=offer))
        elif choice.choice == "leave_blank":
            out.append(Resolved(item.name, item.key, "blank", True, proposal=offer,
                                note="left blank by choice (leave_blank)", kind=kind))
        elif choice.choice == "use_default":
            if offer.value:
                out.append(Resolved(item.name, item.key, "default", True, offer.value,
                                    offer.template, offer))
            else:
                out.append(Resolved(item.name, item.key, "blank", False, proposal=offer,
                                    note=f"use_default chosen, nothing to apply: "
                                         f"{offer.refused}"))
        else:
            template = choice.standard or (shared.standard if shared is not None else "")
            values = flow_tokens(tokens, spec, profile, item.flow) if item.flow else tokens
            value = (render(template, {**tpl.corpus_tokens, **values},
                            name=bool(item.flow) and not item.pk,
                            case=profile.pipeline_name_case) if template else "")
            if value:
                clash = bool(item.flow) and value.casefold() in taken
                out.append(Resolved(item.name, item.key, "standard", True, value, template,
                                    offer, "the naming standard's name collides with a "
                                           "PIPELINE_NAME of the loaded goldens / corpus"
                                    if clash else "", kind=kind))
            else:
                out.append(Resolved(item.name, item.key, "blank", False, proposal=offer,
                                    note="naming_standard chosen without a template that "
                                         "renders a value", kind=kind))
    return out


def faq_answers(resolved: list[Resolved]) -> dict:
    """The FAQ answers a run applies: every explicit choice (a typed value
    rides the answers file's own path)."""
    from codegen.faq import FaqAnswer

    out = {r.name: FaqAnswer(value=r.value, source="engineer",
                             evidence=evidence(r.source, r.template))
           for r in resolved if r.chosen and r.source not in ("answer", "profile")}
    # post79: an autoset PIPELINE_NAME is written from the FAQ like a chosen one
    # (the writer fills that column from answers only); evidence "profile"
    out.update({r.name: FaqAnswer(value=r.value, source="engineer",
                                  evidence=evidence("profile", r.template))
                for r in resolved if r.source == "profile" and r.value
                and r.key.startswith(PIPELINE + ".")})
    return out


def records(resolved: list[Resolved]) -> dict[str, dict]:
    """CODEGEN_STAGES.json: per input key, ``{kind, source, chosen, proposed,
    template}`` (+ ``refused`` / ``note`` when one applies). post61: source
    ``profile`` = a profile default written (unanswered or use_default)."""
    out: dict[str, dict] = {}
    for r in resolved:
        entry = {"kind": r.kind, "source": r.source, "chosen": r.chosen,
                 "proposed": r.proposal.value, "template": r.template or r.proposal.template}
        if is_placeholder_template(entry["template"]):
            entry["placeholder"] = True
        if r.proposal.refused:
            entry["refused"] = r.proposal.refused
        if r.note:
            entry["note"] = r.note
        out[r.key] = entry
    return out


HOW = ("choose use_default / naming_standard / leave_blank (answers file `choice:`, or "
       "Confirm in the App)")


def shared_template_items(feed_id: str, owed: list) -> list:
    """The answers template's feed-level ``pipeline_name`` entry: one choice
    (or ONE naming-standard template) for every PIPELINE_NAME row of the feed —
    a flow's own entry wins."""
    from codegen.layout.answers_template import TemplateItem

    prefix = f"feeds.{feed_id}.{PIPELINE}."
    flow_items = [o for o in owed if o.options and o.key.startswith(prefix)]
    if not flow_items:
        return []
    return [TemplateItem(
        f"feeds.{feed_id}.{PIPELINE}",
        "one choice for every PIPELINE_NAME row of this feed (a flow's own entry wins); "
        "naming_standard: ONE template renders all four names", [], "inputs",
        options=list(CHOICES), shared=True,
        proposed="; ".join(f"{o.key[len(prefix):]}={o.proposed}" for o in flow_items
                           if o.proposed),
        proposed_template=" | ".join(dict.fromkeys(o.template for o in flow_items
                                                   if o.template)))]


HOW_PROFILE = ("change it: answers file `choice: naming_standard` + `standard:` (your "
               "value), or in the App")


def proposal_text(r: Resolved) -> str:
    """The INPUT_NEEDED line's addition for an open defaulted input (post61:
    the NOTE line's for an applied profile default)."""
    if r.applied:
        return (f"profile default applied: {r.value!r} ({r.template or r.proposal.template}) "
                f"— {HOW_PROFILE}")
    if r.proposal.value:
        label = f"; {PLACEHOLDER_LABEL}" if r.proposal.placeholder else ""
        return (f"proposed {r.proposal.value!r} (template {r.proposal.template!r}{label}) — "
                f"not written until chosen: {HOW}")
    if r.proposal.refused:
        return f"no proposal: {r.proposal.refused} — {HOW}"
    return HOW
