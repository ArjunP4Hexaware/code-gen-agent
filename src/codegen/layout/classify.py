"""What KIND of workbook is this — by what it says, never by its name (M9.3).

The document chooser used to list every ``.xlsx`` as an STTM and again as a
possible Vendor Data Dictionary; inside a workspace folder of pairs
(``…/pair_1/{FRD.docx, STTM.xlsx, VDD.xlsx}``) that doubles every list and
lets a dictionary be picked as the mapping. The discovery signatures already
tell the two apart:

* **sttm** — a mapping sheet is discoverable: a band row carrying a stage AND
  a standard token (content-driven), ``MAPPING-`` sheets, or the segmented
  family (``codegen.layout.discover.discover``) — and (Chunk A, 2026-10-09)
  one of its bands carries EVIDENCE of a target layer; target-shaped columns
  nothing names are ``unclassified`` with a reason starting ``confirm:``;
* **vdd** — a FILES sheet or a field sheet by header signature
  (``discover_vdd``; checked second — an STTM's "Layout" sheet may carry
  field-sheet-like headers);
* **iig** / **iig_review** (2026-10-16) — checked FIRST, when the caller passes
  the IIG template's sheet names (``iig_sheets``): at least two of the
  workbook's sheets are IIG template tabs (DATA_FACTORY_PIPELINE_SCHEDULE …);
  with the review copy's summary sheet (``review_sheet``) as well it is the
  BSA's review copy. Never an input — the App's chooser leaves both out of
  its lists. Callers that pass no sheet names (the CLI's pair folders) get
  the three kinds below, as before;
* **unclassified** — neither. Never hidden: the chooser lists it with a badge,
  because a workbook the signatures do not know is exactly the one a person
  must be able to pick.

Pure and offline: a path in, a verdict and its reason (structural labels
only — sheet names and the strategy) out.
"""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from codegen.config import ExtractorConfig
from codegen.layout.discover import (
    CANDIDATE_SHEET_NOTE,
    NoLayoutError,
    discover,
    discover_vdd,
)

WorkbookKind = Literal["sttm", "vdd", "iig", "iig_review", "unclassified"]
# the kinds that are never an input (an IIG is the agent's OUTPUT / a golden)
IIG_KINDS = ("iig", "iig_review")
# an IIG has at least this many template tabs among its sheets
IIG_MIN_TABS = 2


@dataclass(frozen=True)
class WorkbookClass:
    kind: WorkbookKind
    reason: str

    def as_dict(self) -> dict[str, str]:
        return {"kind": self.kind, "reason": self.reason}


def _sheet_names(path: Path) -> list[str]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def iig_kind(sheets: list[str], iig_sheets: Iterable[str],
             review_sheet: str = "") -> WorkbookClass | None:
    """The workbook is an IIG / an IIG review copy by its SHEET NAMES (None:
    it is not): at least ``IIG_MIN_TABS`` of them are IIG template tabs;
    with ``review_sheet`` too, the review copy."""
    tabs = {str(s).strip().upper() for s in iig_sheets if str(s).strip()}
    matched = [s for s in sheets if s.strip().upper() in tabs]
    if not tabs or len(matched) < IIG_MIN_TABS:
        return None
    shown = ", ".join(matched[:3]) + (f" … +{len(matched) - 3}" if len(matched) > 3 else "")
    review = review_sheet.strip().upper()
    if review and any(s.strip().upper() == review for s in sheets):
        return WorkbookClass("iig_review", f"an IIG review copy: sheet {review_sheet} and "
                                           f"{len(matched)} IIG template sheet(s) ({shown})")
    return WorkbookClass("iig", f"an IIG: {len(matched)} IIG template sheet(s) ({shown})")


def classify_workbook(path: Path, extractor: ExtractorConfig, *,
                      iig_sheets: Iterable[str] = (), review_sheet: str = "") -> WorkbookClass:
    unconfirmed: list[str] = []
    candidates: list[str] = []
    if iig_sheets:
        try:
            iig = iig_kind(_sheet_names(path), iig_sheets, review_sheet)
        except Exception as exc:  # noqa: BLE001 — a workbook that cannot be opened is a fact to show
            return WorkbookClass("unclassified", f"could not be read: {type(exc).__name__}: {exc}")
        if iig is not None:
            return iig
    try:
        found = discover(path, extractor)
        # Chunk A: an STTM only when a band carries EVIDENCE of a target layer
        # (a stage / standard band). Target-shaped columns nothing names — a
        # dictionary or IIG sheet looks the same — are "unclassified —
        # confirm", never silently an STTM.
        # Chunk B: a CANDIDATE mapping sheet (pass 3 - read by its band title
        # spans or shared header qualifiers, no band row, no label group) is
        # never an STTM by itself: an IIG / dictionary sheet ("SRC_… / TGT_…")
        # has the same shape. "Confirm" - a person picks it as the STTM.
        candidates = [s.name for s in found.profile.mapping_sheets
                      if any(n.startswith(CANDIDATE_SHEET_NOTE) for n in s.notes)]
        mapping = [s.name for s in found.profile.mapping_sheets
                   if any(b.layer in ("stage", "standard") for b in s.bands)
                   and s.name not in candidates]
        if mapping:
            return WorkbookClass("sttm", f"mapping sheet(s) {mapping} "
                                         f"({found.profile.strategy} discovery)")
        unconfirmed = [s.name for s in found.profile.mapping_sheets if s.name not in candidates]
    except NoLayoutError:
        pass
    except Exception as exc:  # noqa: BLE001 — a workbook that cannot be opened is a fact to show
        return WorkbookClass("unclassified", f"could not be read: {type(exc).__name__}: {exc}")
    try:
        found = discover_vdd(path, extractor)
        sheets = [f"{s.name} ({s.kind})" for s in found.profile.sheets
                  if s.kind in ("vdd_files", "vdd_fields")]
        return WorkbookClass("vdd", f"dictionary sheet(s) {sheets}")
    except NoLayoutError:
        pass
    except Exception as exc:  # noqa: BLE001
        return WorkbookClass("unclassified", f"could not be read: {type(exc).__name__}: {exc}")
    if candidates and not unconfirmed:
        return WorkbookClass(
            "unclassified",
            "confirm: no band row with stage + standard labels (STTM) and no FILES / "
            f"field-sheet header (Vendor Data Dictionary); sheet(s) {candidates} read as "
            "CANDIDATE mapping sheets only (band title spans / shared header qualifiers) - an "
            "IIG or dictionary sheet has the same shape; pick it as the STTM to read it as one")
    if unconfirmed:
        return WorkbookClass(
            "unclassified",
            f"confirm: sheet(s) {unconfirmed} carry target-shaped columns (schema / table / "
            "column / data type) but no band title, header word, catalog or schema value names "
            "a stage or standard layer — pick it as the STTM to answer the band questions")
    return WorkbookClass(
        "unclassified",
        "no band row with stage + standard labels (STTM) and no FILES / field-sheet header "
        "(Vendor Data Dictionary)")


__all__ = ["IIG_KINDS", "WorkbookClass", "WorkbookKind", "classify_workbook", "iig_kind"]
