"""Layout discovery — every strategy ends in a :class:`LayoutProfile`.

Three deterministic strategies, tried in order; the first that applies
wins, so the two legacy paths keep today's behaviour byte for byte:

1. ``mapping_prefix`` — the ``MAPPING-`` sheet-name path (FILE_DETAILS +
   VERSION_HISTORY + per-sheet band row 1 / header row 2), resolved with
   the legacy functions in :mod:`codegen.extract.workbook`.
2. ``segmented_family`` — the CAQH-shaped family (key:value metadata block,
   band row with the ``Source Layout`` variant, per-row Segment column),
   resolved with the legacy block resolution in
   :mod:`codegen.extract.segmented`. Applies only when that resolution
   succeeds; otherwise the sheet falls through to strategy 3 and the
   profile notes say so.
3. ``content`` — content-driven discovery (M1): a mapping sheet is any
   sheet with a row carrying both a stage token and a standard token
   (``extractor.discovery.band_tokens``) and a header row beneath it (or
   the same row, when the layer prefixes live in the header texts);
   band spans come from merged ranges, else from band-text positions;
   header→role, meta-row and segment vocabularies come from
   ``extractor.discovery`` in config; auxiliary sheets are recognised by
   header signature; everything unresolved is listed, never guessed.

Every profile carries ``source="synonyms"`` here — a model or a person
(M2.5) can only ever refine what these tables leave unresolved.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from codegen.config import DiscoveryConfig, ExtractorConfig
from codegen.layout.fingerprint import fingerprint
from codegen.layout.profile import (
    REQUIRED_ROLES,
    BandProfile,
    LayoutProfile,
    MetaRow,
    Role,
    SheetProfile,
    UnresolvedRole,
    confidence_key,
)

# Synonym-table hits are exact normalized matches: full confidence.
_SYNONYM_CONFIDENCE = 1.0
# Rows scanned for an auxiliary sheet's header signature.
_AUX_SCAN_ROWS = 12
# A header row must carry at least this many text cells.
_MIN_HEADER_CELLS = 3
# A band row carries a few group labels, never a full header's worth.
_MAX_BAND_LABELS = 5


def normalize(value: object) -> str:
    """Comparison form for every header, label, token and marker: lowercase,
    every run of non-alphanumerics (space, _ / - ( ) ? : # . …) → one
    space, trimmed. "Target_Column_Name_in_DL" → "target column name in dl"."""
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def text(value: object) -> str | None:
    if value is None:
        return None
    stripped = str(value).strip()
    return stripped or None


@dataclass
class Discovery:
    """A discovered layout plus the loaded workbook (loaded once, reused by
    the reader) and the diagnostics discovery logged on the way."""

    profile: LayoutProfile
    workbook: object
    diagnostics: list[str] = field(default_factory=list)


class NoLayoutError(ValueError):
    """No discovery strategy applies to this workbook."""


def discover(path: Path, config: ExtractorConfig) -> Discovery:
    workbook = load_workbook(path, data_only=True)
    digest = fingerprint(workbook)
    diagnostics: list[str] = []
    for strategy in (_mapping_prefix, _segmented_family, _content):
        profile = strategy(workbook, config, path.name, digest, diagnostics)
        if profile is not None:
            return Discovery(profile=profile, workbook=workbook, diagnostics=diagnostics)
    raise NoLayoutError(
        f"{path.name}: no mapping sheets found — no discovery strategy applies "
        f"(prefix {config.mapping_sheet_prefix!r}, segmented family, content-driven); "
        f"sheets present: {workbook.sheetnames}; diagnostics: {diagnostics}"
    )


# ------------------------------------------------ strategy 1: MAPPING- prefix

def _mapping_prefix(workbook, config: ExtractorConfig, name: str, digest: str,
                    diagnostics: list[str]) -> LayoutProfile | None:
    from codegen.extract import workbook as legacy

    mapping = [n for n in workbook.sheetnames if n.startswith(config.mapping_sheet_prefix)]
    if not mapping:
        return None
    sheets: list[SheetProfile] = []
    confidence: dict[str, float] = {}
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        if sheet_name in mapping:
            bands = legacy._find_bands(ws, config)
            headers = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
            recycle_col = legacy._find_recycle_column(ws, headers, config)
            source = legacy._resolve_source_headers(ws, headers, bands.source, config, recycle_col)
            stage = legacy._resolve_table_headers(
                ws, headers, bands.stage, "stage", config, recycle_col)
            standard = (
                legacy._resolve_table_headers(
                    ws, headers, bands.standard, "standard", config, recycle_col)
                if bands.standard is not None else None
            )
            source_roles = legacy.LEGACY_SOURCE_ROLES
            table_roles = legacy.LEGACY_TABLE_ROLES
            profiles = [
                BandProfile(layer="source", col_start=bands.source[0] + 1,
                            col_end=bands.source[1],
                            roles={source_roles[k].value: v + 1 for k, v in source.items()},
                            label=config.band_labels.source),
                BandProfile(layer="stage", col_start=bands.stage[0] + 1, col_end=bands.stage[1],
                            roles={table_roles[k].value: v + 1 for k, v in stage.items()},
                            label=config.band_labels.stage),
            ]
            if standard is not None and bands.standard is not None:
                profiles.append(BandProfile(
                    layer="standard", col_start=bands.standard[0] + 1,
                    col_end=bands.standard[1],
                    roles={table_roles[k].value: v + 1 for k, v in standard.items()},
                    label=config.band_labels.standard))
            if recycle_col is not None:
                # The per-row recycle column may sit anywhere (the golden
                # workbook parks it beyond the standard band); attach it to
                # the band whose span holds it, else as a trailing band.
                col = recycle_col + 1
                holder = next((b for b in profiles if b.col_start <= col <= b.col_end), None)
                if holder is not None:
                    profiles[profiles.index(holder)] = holder.model_copy(
                        update={"roles": {**holder.roles, Role.RECYCLE_FLAG.value: col}})
                else:
                    profiles.append(BandProfile(layer="rules", col_start=col, col_end=col,
                                                roles={Role.RECYCLE_FLAG.value: col}))
            for band in profiles:
                for role in band.roles:
                    confidence[confidence_key(sheet_name, band.layer, role)] = _SYNONYM_CONFIDENCE
            sheets.append(SheetProfile(name=sheet_name, kind="mapping", header_row=2,
                                       band_row=1, bands=profiles))
        elif sheet_name == config.file_details_sheet:
            headers = [str(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
                       if c is not None]
            sheets.append(SheetProfile(name=sheet_name, kind="file_details", header_row=1,
                                       headers=headers))
        elif sheet_name == config.version_history_sheet:
            sheets.append(SheetProfile(name=sheet_name, kind="version"))
        else:
            diagnostics.append(f"sheet {sheet_name!r}: not a MAPPING- sheet; ignored")
            sheets.append(SheetProfile(name=sheet_name, kind="ignore"))
    return LayoutProfile(fingerprint=digest, sheets=sheets, confidence=confidence,
                         source="synonyms", strategy="mapping_prefix")


# ------------------------------------------ strategy 2: segmented (CAQH) family

# Segmented logical names (codegen.extract.segmented) -> roles.
SEGMENTED_SOURCE_ROLES = {
    "ordinal": Role.ORDINAL,
    "field_name": Role.FIELD_NAME,
    "datatype": Role.SOURCE_TYPE,
    "length": Role.LENGTH,
    "fixed_width_length": Role.FIELD_LENGTH,
    "fixed_width_start": Role.START,
    "fixed_width_end": Role.END,
    "segment": Role.SEGMENT,
    "pii": Role.PII,
    "comments": Role.COMMENTS,
    "business_rule": Role.BUSINESS_RULE,
}
SEGMENTED_TABLE_ROLES = {
    "catalog": Role.CATALOG,
    "schema": Role.SCHEMA,
    "tablename": Role.TABLE,
    "columnname": Role.COLUMN,
    "datatype": Role.TARGET_TYPE,
    "mandatory column": Role.MANDATORY_COLUMN,
    "primary key": Role.PRIMARY_KEY,
    "field description": Role.FIELD_DESCRIPTION,
    "table description": Role.TABLE_DESCRIPTION,
    "transformations/data quality": Role.TRANSFORMATION,
}
SEGMENTED_FAMILY_NOTE = (
    "matches the segmented (metadata-block + band-row) layout family — see "
    "docs/SEGMENTED_MODE_DESIGN.md"
)


def _segmented_family(workbook, config: ExtractorConfig, name: str, digest: str,
                      diagnostics: list[str]) -> LayoutProfile | None:
    from codegen.extract import segmented as legacy
    from codegen.extract.workbook import WorkbookParseError, _norm

    seg = config.segmented
    labels = config.band_labels
    band_vocabulary = (
        {_norm(labels.source), _norm(labels.stage), _norm(labels.standard)}
        | {_norm(v) for v in seg.source_band_variants}
    )
    try:
        ws, band_index, rows = legacy._find_mapping_sheet(workbook, seg, band_vocabulary, name)
    except WorkbookParseError as exc:
        if "no sheet matches" in str(exc):
            return None
        diagnostics.append(f"segmented family signature: {exc}")
        return None
    diagnostics.append(f"sheet {ws.title!r}: {SEGMENTED_FAMILY_NOTE}")
    band_starts: dict[str, int] = {}
    source_variants = {_norm(labels.source)} | {_norm(v) for v in seg.source_band_variants}
    for index, cell in enumerate(rows[band_index]):
        value = _norm(cell)
        if not value:
            continue
        if value in source_variants:
            band_starts["source"] = index
        elif value == _norm(labels.stage):
            band_starts["stage"] = index
        elif value == _norm(labels.standard):
            band_starts["standard"] = index
    try:
        facts = legacy._parse_metadata_block(rows, band_index, seg, ws.title)
        max_col = max(len(r) for r in rows) - 1
        header = list(rows[band_index + 1])
        blocks = legacy._resolve_blocks(header, band_starts, max_col, seg, ws.title)
    except WorkbookParseError as exc:
        diagnostics.append(
            f"sheet {ws.title!r}: segmented family signature matched but the segmented "
            f"extractor cannot resolve it ({exc}); routed to content-driven discovery")
        return None
    del facts
    stage_end = band_starts.get("standard", max_col + 1)
    bands = [
        BandProfile(layer="source", col_start=band_starts["source"] + 1,
                    col_end=band_starts["stage"],
                    roles={SEGMENTED_SOURCE_ROLES[k].value: v + 1
                           for k, v in blocks.source.items()},
                    label=text(rows[band_index][band_starts["source"]])),
        BandProfile(layer="stage", col_start=band_starts["stage"] + 1, col_end=stage_end,
                    roles={SEGMENTED_TABLE_ROLES[k].value: v + 1 for k, v in blocks.stage.items()},
                    label=text(rows[band_index][band_starts["stage"]])),
    ]
    if blocks.standard is not None and "standard" in band_starts:
        bands.append(BandProfile(
            layer="standard", col_start=band_starts["standard"] + 1, col_end=max_col + 1,
            roles={SEGMENTED_TABLE_ROLES[k].value: v + 1 for k, v in blocks.standard.items()},
            label=text(rows[band_index][band_starts["standard"]])))
    wanted = {_norm(label): logical for logical, label in seg.metadata_keys.items()}
    meta_rows = []
    for row_index, row in enumerate(rows[:band_index], start=1):
        label = text(row[0]) if row else None
        if label is None:
            continue
        meta_rows.append(MetaRow(row=row_index, col=1, label=label, key=wanted.get(_norm(label)),
                                 value_col=2))
    confidence = {confidence_key(ws.title, b.layer, r): _SYNONYM_CONFIDENCE
                  for b in bands for r in b.roles}
    sheets = []
    for sheet_name in workbook.sheetnames:
        if sheet_name == ws.title:
            sheets.append(SheetProfile(
                name=sheet_name, kind="mapping", header_row=band_index + 2,
                band_row=band_index + 1, bands=bands, meta_rows=meta_rows,
                segment_strategy="column", segment_column=blocks.source["segment"] + 1,
                notes=[SEGMENTED_FAMILY_NOTE]))
        else:
            diagnostics.append(f"sheet {sheet_name!r}: not the segmented mapping sheet; ignored")
            sheets.append(SheetProfile(name=sheet_name, kind="ignore"))
    return LayoutProfile(fingerprint=digest, sheets=sheets, confidence=confidence,
                         source="synonyms", strategy="segmented_family",
                         notes=[f"{ws.title}: {SEGMENTED_FAMILY_NOTE}"])


# -------------------------------------------- strategy 3: content-driven

_LAYER_ORDER = ("stage", "standard", "rules", "source")


def _band_layer(label: str, tokens: dict[str, list[str]]) -> str | None:
    normalized = normalize(label)
    for layer in _LAYER_ORDER:
        for token in tokens.get(layer, []):
            if normalize(token) in normalized:
                return layer
    return None


def _has_token(cells: list[str], tokens: list[str]) -> bool:
    return any(normalize(t) in c for c in cells for t in tokens if c)


def _content(workbook, config: ExtractorConfig, name: str, digest: str,
             diagnostics: list[str]) -> LayoutProfile | None:
    disc = config.discovery
    if not disc.roles or not disc.band_tokens:
        diagnostics.append("content-driven discovery: extractor.discovery vocabulary is empty")
        return None
    sheets: list[SheetProfile] = []
    confidence: dict[str, float] = {}
    unresolved: list[UnresolvedRole] = []
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        mapping = _discover_mapping_sheet(ws, disc, config, confidence, unresolved, diagnostics)
        if mapping is not None:
            sheets.append(mapping)
            continue
        sheets.append(_classify_auxiliary(ws, disc, diagnostics))
    mapping_sheets = [s for s in sheets if s.kind == "mapping"]
    if not mapping_sheets:
        return None
    sheets = _apply_sheet_segments(sheets, disc)
    return LayoutProfile(fingerprint=digest, sheets=sheets, confidence=confidence,
                         source="synonyms", strategy="content", unresolved=unresolved,
                         notes=list(diagnostics))


def _discover_mapping_sheet(ws, disc: DiscoveryConfig, config: ExtractorConfig,
                            confidence: dict[str, float], unresolved: list[UnresolvedRole],
                            diagnostics: list[str]) -> SheetProfile | None:
    rows = [list(r) for _, r in zip(range(disc.scan_rows),
                                    ws.iter_rows(values_only=True), strict=False)]
    band_index: int | None = None
    header_index: int | None = None
    field_tokens = disc.roles.get("source", {}).get(Role.FIELD_NAME.value, [])
    for index, row in enumerate(rows):
        cells = [normalize(c) for c in row]
        if not (_has_token(cells, disc.band_tokens.get("stage", []))
                and _has_token(cells, disc.band_tokens.get("standard", []))):
            continue
        below = rows[index + 1] if index + 1 < len(rows) else []
        below_cells = [normalize(c) for c in below]
        filled = sum(1 for c in cells if c)
        # (a) band row + header row beneath carrying a field-name header;
        # (b) the token row IS the header row (layer prefixes in the texts);
        # (c) a sparse band row (a few labels) over a wide header row whose
        #     field-name header is not in the vocabulary (pinned unresolved).
        if _has_token(below_cells, field_tokens):
            band_index, header_index = index, index + 1
            break
        if _has_token(cells, field_tokens):
            header_index = index
            break
        if (filled <= _MAX_BAND_LABELS
                and sum(1 for c in below_cells if c) >= _MIN_HEADER_CELLS):
            band_index, header_index = index, index + 1
            break
    if header_index is None:
        return None
    header = rows[header_index]
    last_col = max((i + 1 for i, c in enumerate(header) if text(c) is not None), default=0)
    if band_index is not None:
        bands = _bands_from_band_row(ws, rows[band_index], band_index + 1, last_col, disc)
    else:
        bands = _bands_from_header_prefixes(header, last_col, disc)
    notes: list[str] = []
    if band_index is not None:
        notes.append(f"band row {band_index + 1}, header row {header_index + 1}")
    else:
        notes.append(f"no band row; layer prefixes in header row {header_index + 1}")
    if not _has_token([normalize(c) for c in header],
                      disc.roles.get("source", {}).get(Role.FIELD_NAME.value, [])):
        notes.append("header row carries no field-name synonym; field_name unresolved")
    resolved_bands: list[BandProfile] = []
    for band in bands:
        resolved_bands.append(_resolve_band_roles(ws.title, band, header, disc, confidence,
                                                  unresolved, diagnostics,
                                                  bands_before=list(resolved_bands)))
    meta_rows = _meta_rows(rows[: (band_index if band_index is not None else header_index)], disc,
                           ws.title, diagnostics)
    source = next((b for b in resolved_bands if b.layer == "source"), None)
    segment_column = source.column(Role.SEGMENT) if source is not None else None
    if segment_column is not None:
        strategy = "column"
    elif _has_banner_rows(ws, header_index + 1, resolved_bands, disc):
        strategy = "banner"
        notes.append("segments introduced by in-sheet label rows spanning the source group")
    else:
        strategy = "none"
    for layer, required in REQUIRED_ROLES.items():
        band = next((b for b in resolved_bands if b.layer == layer), None)
        if band is None:
            continue
        for role in required:
            if band.column(role) is None and not any(
                    u.sheet == ws.title and u.layer == layer and u.role == role.value
                    for u in unresolved):
                unresolved.append(UnresolvedRole(
                    sheet=ws.title, layer=layer, role=role.value,
                    reason="required role has no header matching its synonyms",
                    candidates=[c for c in range(band.col_start, band.col_end + 1)
                                if c not in band.roles.values()]))
    for entry in notes:
        diagnostics.append(f"sheet {ws.title!r}: {entry}")
    return SheetProfile(name=ws.title, kind="mapping", header_row=header_index + 1,
                        band_row=None if band_index is None else band_index + 1,
                        bands=resolved_bands, meta_rows=meta_rows,
                        segment_strategy=strategy, segment_column=segment_column, notes=notes)


def _bands_from_band_row(ws, band_row: list, band_row_number: int, last_col: int,
                         disc: DiscoveryConfig) -> list[BandProfile]:
    labels = [(i + 1, text(c)) for i, c in enumerate(band_row) if text(c) is not None]
    merged = {r.min_col: r.max_col for r in ws.merged_cells.ranges
              if r.min_row == band_row_number and r.max_row == band_row_number}
    bands: list[BandProfile] = []
    for position, (col, label) in enumerate(labels):
        if col in merged:
            end = merged[col]
        elif position + 1 < len(labels):
            end = labels[position + 1][0] - 1
        else:
            end = max(last_col, col)
        layer = _band_layer(label, disc.band_tokens)
        if layer is None:
            layer = "source" if position == 0 else "rules"
        bands.append(BandProfile(layer=layer, col_start=col, col_end=end, label=label))
    covered_to = max((b.col_end for b in bands), default=0)
    if last_col > covered_to:
        bands.append(BandProfile(layer="rules", col_start=covered_to + 1, col_end=last_col,
                                 label=None))
    return bands


def _bands_from_header_prefixes(header: list, last_col: int,
                                disc: DiscoveryConfig) -> list[BandProfile]:
    layers: list[str | None] = []
    for col in range(1, last_col + 1):
        value = normalize(header[col - 1]) if col - 1 < len(header) else ""
        layer = None
        for candidate in ("stage", "standard"):
            if any(value.startswith(normalize(t)) for t in disc.band_tokens.get(candidate, [])):
                layer = candidate
                break
        layers.append(layer)
    first_stage = next((i for i, layer in enumerate(layers) if layer == "stage"), None)
    last_target = max((i for i, layer in enumerate(layers) if layer in ("stage", "standard")),
                      default=-1)
    bands: list[BandProfile] = []
    if first_stage is None:
        return bands
    bands.append(BandProfile(layer="source", col_start=1, col_end=first_stage, label=None))
    for layer in ("stage", "standard"):
        cols = [i + 1 for i, item in enumerate(layers) if item == layer]
        if cols:
            bands.append(BandProfile(layer=layer, col_start=min(cols), col_end=max(cols),
                                     label=None))
    if last_target + 1 < last_col:
        bands.append(BandProfile(layer="rules", col_start=last_target + 2, col_end=last_col,
                                 label=None))
    return bands


def _roles_group(band: BandProfile, bands_before: list[BandProfile]) -> str:
    if band.layer in ("stage", "standard"):
        return "target"
    if band.layer == "source":
        return "source"
    # A rules band AFTER a target band is the trailing group (recycle flag,
    # DQ mandatory list, comments); before it, the data-rules band.
    return "trailing" if any(b.layer in ("stage", "standard") for b in bands_before) else "rules"


def _resolve_band_roles(sheet: str, band: BandProfile, header: list, disc: DiscoveryConfig,
                        confidence: dict[str, float], unresolved: list[UnresolvedRole],
                        diagnostics: list[str], bands_before: list[BandProfile] | None = None
                        ) -> BandProfile:
    group = _roles_group(band, bands_before or [])
    table = {role: {normalize(s) for s in spellings}
             for role, spellings in disc.roles.get(group, {}).items()}
    roles: dict[str, int] = {}
    for col in range(band.col_start, band.col_end + 1):
        raw = header[col - 1] if col - 1 < len(header) else None
        value = normalize(raw)
        if not value:
            continue
        matches = [role for role, spellings in table.items() if value in spellings]
        if len(matches) == 1:
            role = matches[0]
            if role in roles:
                diagnostics.append(
                    f"sheet {sheet!r}: header {raw!r} (column {col}) also matches "
                    f"{role!r}, already taken by column {roles[role]}; kept the first")
                continue
            roles[role] = col
            confidence[confidence_key(sheet, band.layer, role)] = _SYNONYM_CONFIDENCE
        elif len(matches) > 1:
            unresolved.append(UnresolvedRole(
                sheet=sheet, layer=band.layer, role="|".join(sorted(matches)),
                reason=f"header {raw!r} (column {col}) is ambiguous between {sorted(matches)}",
                candidates=[col]))
        else:
            diagnostics.append(
                f"sheet {sheet!r}: {band.layer} header {raw!r} (column {col}) matches no "
                f"{group} synonym; left unresolved")
    return band.model_copy(update={"roles": roles})


def _meta_rows(rows_above: list[list], disc: DiscoveryConfig, sheet: str,
               diagnostics: list[str]) -> list[MetaRow]:
    synonyms = {normalize(s): key for key, spellings in disc.meta_synonyms.items()
                for s in spellings}
    out: list[MetaRow] = []
    for index, row in enumerate(rows_above, start=1):
        cells = [(i + 1, text(c)) for i, c in enumerate(row) if text(c) is not None]
        if not cells:
            continue
        label_col, label = cells[0]
        value_col = cells[1][0] if len(cells) > 1 else None
        key = synonyms.get(normalize(label))
        if key is None:
            diagnostics.append(f"sheet {sheet!r} row {index}: meta label {label!r} unrecognised")
        out.append(MetaRow(row=index, col=label_col, label=label, key=key, value_col=value_col))
    return out


def _has_banner_rows(ws, header_row: int, bands: list[BandProfile],
                     disc: DiscoveryConfig) -> bool:
    source = next((b for b in bands if b.layer == "source"), None)
    if source is None:
        return False
    spellings = {normalize(s) for values in disc.segment_synonyms.values() for s in values}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        filled = [(i + 1, text(c)) for i, c in enumerate(row) if text(c) is not None]
        if len(filled) != 1:
            continue
        col, value = filled[0]
        if source.col_start <= col <= source.col_end and normalize(value) in spellings:
            return True
    return False


def _apply_sheet_segments(sheets: list[SheetProfile], disc: DiscoveryConfig) -> list[SheetProfile]:
    """One-sheet-per-segment: every mapping sheet's name spells a segment."""
    mapping = [s for s in sheets if s.kind == "mapping"]
    if len(mapping) < 2:
        return sheets
    spellings = {normalize(s) for values in disc.segment_synonyms.values() for s in values}
    if all(normalize(s.name) in spellings or normalize(s.name).split(" ")[-1] in spellings
           for s in mapping):
        return [s.model_copy(update={"segment_strategy": "sheet"}) if s.kind == "mapping"
                and s.segment_strategy == "none" else s for s in sheets]
    return sheets


def _classify_auxiliary(ws, disc: DiscoveryConfig, diagnostics: list[str]) -> SheetProfile:
    rows = [list(r) for _, r in zip(range(_AUX_SCAN_ROWS), ws.iter_rows(values_only=True),
                                    strict=False)]
    for kind, alternatives in disc.auxiliary_sheets.items():
        for index, row in enumerate(rows, start=1):
            cells = [normalize(c) for c in row if text(c) is not None]
            if len(cells) < 1:
                continue
            for tokens in alternatives:
                if all(any(normalize(t) in c for c in cells) for t in tokens):
                    headers = [str(text(c)) for c in row if text(c) is not None]
                    diagnostics.append(f"sheet {ws.title!r}: auxiliary {kind} (header row {index})")
                    return SheetProfile(name=ws.title, kind=kind, header_row=index,  # type: ignore[arg-type]
                                        headers=headers)
    diagnostics.append(f"sheet {ws.title!r}: no mapping structure or auxiliary signature; ignored")
    return SheetProfile(name=ws.title, kind="ignore")


__all__ = [
    "Discovery",
    "NoLayoutError",
    "SEGMENTED_FAMILY_NOTE",
    "SEGMENTED_SOURCE_ROLES",
    "SEGMENTED_TABLE_ROLES",
    "discover",
    "normalize",
    "text",
]
