"""A worksheet's REAL used range (M11 item 11).

A sheet may carry a styled but EMPTY cell a million rows down (Excel keeps
formatting a user dragged to the bottom). openpyxl then reports
``max_row = 1,048,538`` — and every ``iter_rows`` over the sheet, in normal
mode, CREATES a cell object for every row up to it: 34 columns x a million
rows. That, not the data, is what ran the v0.6.2 parser past its budget on a
sheet holding 336 real rows (SHAPES_ROUND2 §3).

So documents are loaded through :func:`load_document`, which trims each
worksheet to its populated extent right after the load: the last row with a
value before the first run of ``stop_after`` consecutive empty rows (config
``extractor.used_range_empty_rows``). openpyxl derives ``max_row`` from the
cells it holds, so every scan downstream is bounded without a change at the
call site. Nothing with a value is dropped silently: rows with values beyond
such a gap are counted and reported (``TrimReport``).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class TrimReport:
    """Per sheet: (declared max_row, used max_row, valued rows beyond the gap)."""

    sheets: dict[str, tuple[int, int, int]] = field(default_factory=dict)

    @property
    def trimmed(self) -> dict[str, tuple[int, int, int]]:
        return {name: v for name, v in self.sheets.items() if v[0] != v[1]}

    def notes(self) -> list[str]:
        out = []
        for name, (declared, used, beyond) in self.trimmed.items():
            note = (f"sheet {name!r}: used range {used} row(s) (the sheet reached row "
                    f"{declared:,}, formatting only)")
            if beyond:
                note += (f"; {beyond} row(s) WITH values lie beyond a run of empty rows and "
                         "were not read — raise extractor.used_range_empty_rows to read them")
            out.append(note)
        return out


def used_max_row(valued_rows: list[int], stop_after: int) -> tuple[int, int]:
    """(last row before the first gap of >= ``stop_after`` empty rows, the
    number of valued rows beyond it). ``valued_rows`` sorted ascending."""
    if not valued_rows:
        return 0, 0
    last = valued_rows[0]
    for index, row in enumerate(valued_rows[1:], start=1):
        if stop_after > 0 and row - last - 1 >= stop_after:
            return last, len(valued_rows) - index
        last = row
    return last, 0


def trim_worksheet(ws, stop_after: int) -> tuple[int, int, int]:
    """Drop the cells beyond the used range. (declared, used, beyond)."""
    cells = getattr(ws, "_cells", None)
    if cells is None:                                    # a read-only sheet: nothing held
        return ws.max_row or 0, ws.max_row or 0, 0
    declared = ws.max_row or 0
    valued = sorted({row for (row, _col), cell in cells.items()
                     if cell.value is not None and cell.value != ""})
    used, beyond = used_max_row(valued, stop_after)
    # Only a real dead tail (>= stop_after rows past the data) is trimmed: a
    # sheet with a few formatted empty rows keeps its exact max_row, which is
    # fingerprint input — so no cached layout profile is invalidated.
    if stop_after > 0 and used and declared - used >= stop_after:
        for key in [k for k in cells if k[0] > used]:
            del cells[key]
        dims = getattr(ws, "row_dimensions", None)
        if dims is not None:
            for row in [r for r in list(dims) if isinstance(r, int) and r > used]:
                del dims[row]
    return declared, (ws.max_row or 0), beyond


def trim_workbook(workbook, stop_after: int) -> TrimReport:
    report = TrimReport()
    for ws in workbook.worksheets:
        report.sheets[ws.title] = trim_worksheet(ws, stop_after)
    return report


def load_document(path: Path, stop_after: int):
    """``load_workbook(path, data_only=True)`` trimmed to the used range —
    the ONE way a document workbook is opened for discovery and extraction."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    workbook.trim_report = trim_workbook(workbook, stop_after)   # type: ignore[attr-defined]
    return workbook
