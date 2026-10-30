"""Workbook fingerprint — the header REGION only, hashed.

The fingerprint identifies a layout, not a document: sheet names, merged
ranges, row/column counts and the first ``header_rows`` rows of every sheet
as text with cell coordinates. Data rows below the header region never
enter it, so two workbooks with the same layout and different content share
a fingerprint (the layout-profile cache key), and a header change breaks
it. In M2.5 this rendering is also the workbook content a model sees — with
one exception (post111): a CAPPED region (no header row found within
``header_rows``) may be data rows, so the model gets that sheet's summary line
only (``model_region``); the fingerprint still hashes the full capped region.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, replace

from openpyxl.utils import get_column_letter

HEADER_ROWS = 25
# A header row carries at least this many cells; a band row at most this many.
_MIN_HEADER_CELLS = 3
_MAX_BAND_LABELS = 5


@dataclass(frozen=True)
class SheetRegion:
    name: str
    rows: int
    cols: int
    merged: tuple[str, ...]
    cells: tuple[tuple[str, str], ...]   # (coordinate, text) for non-empty header-region cells
    last_row: int                        # the region's last row (the header row when bounded)
    capped: bool                         # header region not bounded within HEADER_ROWS
    withheld: bool = False               # model_region: a capped region's cells not sent


def header_bound(ws, header_rows: int = HEADER_ROWS) -> tuple[int, bool]:
    """(last row of the header region, capped?) by a cheap structural scan:
    the first row with >= 3 non-empty cells whose next row is populated too
    is the header row — unless it is a sparse band row (<= 5 labels) over a
    wider header row, in which case the header row is the one beneath. Meta
    rows (label:value) and blank rows above are part of the region; data
    rows below it are not. No such row within ``header_rows`` → capped."""
    counts = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, header_rows + 1),
                            values_only=True):
        counts.append(sum(1 for c in row if c is not None and str(c).strip()))
    for index, count in enumerate(counts[:header_rows]):
        if count < _MIN_HEADER_CELLS:
            continue
        below = counts[index + 1] if index + 1 < len(counts) else 0
        if below < _MIN_HEADER_CELLS:
            continue
        if count <= _MAX_BAND_LABELS and below > count:
            return index + 2, False      # band row over a wider header row
        return index + 1, False
    return min(ws.max_row, header_rows), True


def sheet_region(ws, header_rows: int = HEADER_ROWS) -> SheetRegion:
    last_row, capped = header_bound(ws, header_rows)
    cells: list[tuple[str, str]] = []
    for row in ws.iter_rows(min_row=1, max_row=last_row):
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                cells.append((f"{get_column_letter(cell.column)}{cell.row}", str(cell.value)))
    return SheetRegion(
        name=ws.title,
        rows=ws.max_row,
        cols=ws.max_column,
        merged=tuple(sorted(str(r) for r in ws.merged_cells.ranges)),
        cells=tuple(cells),
        last_row=last_row,
        capped=capped,
    )


def clip_text(value: str, max_chars: int | None) -> str:
    """``value`` cut to ``max_chars`` characters plus a marker naming the full
    length (P1-7: the model's copy only - never a hash input, never an artefact)."""
    if max_chars is None or len(value) <= max_chars:
        return value
    return f"{value[:max_chars]}... [clipped, {len(value)} chars]"


def model_region(ws, header_rows: int = HEADER_ROWS, *, max_cell_chars: int | None = None,
                 withhold: bool = False) -> SheetRegion:
    """What the layout MODEL sees of one sheet: its header region, except a
    capped one. A sheet with no header row within ``header_rows`` (a two-column
    lookup, valid-values or member list) has no header region, and its first rows
    may be data rows, so none of its cells are sent - only the summary line
    (name, size, merged ranges). post111 (codeflow audit #9): up to 25 rows of such
    a sheet used to reach the model.

    P1-7 (red team LM-5): ``withhold`` sends the summary line only for a bounded
    sheet too (a sheet with no open role), and ``max_cell_chars`` clips each sent
    cell's text (:func:`clip_text`). Neither touches :func:`sheet_region`, which
    the fingerprint - the cache key - hashes."""
    region = sheet_region(ws, header_rows)
    if region.capped or withhold:
        return replace(region, cells=(), withheld=True)
    if max_cell_chars is not None:
        region = replace(region, cells=tuple((coordinate, clip_text(value, max_cell_chars))
                                             for coordinate, value in region.cells))
    return region


def render_region(region: SheetRegion) -> str:
    """The canonical text of one sheet's header region (hash input and, in
    M2.5, prompt material)."""
    lines = [
        f"sheet {region.name!r} rows={region.rows} cols={region.cols} "
        f"header_region=1..{region.last_row}"
        + ((" (capped, cells withheld)" if region.withheld else " (capped)")
           if region.capped else (" (cells withheld)" if region.withheld else "")),
        "merged: " + (", ".join(region.merged) or "-"),
    ]
    lines += [f"{coordinate}: {text}" for coordinate, text in region.cells]
    return "\n".join(lines)


def fingerprint(workbook, header_rows: int = HEADER_ROWS) -> str:
    """sha256 over every sheet's rendered header region, in sheet order."""
    digest = hashlib.sha256()
    for ws in workbook.worksheets:
        digest.update(render_region(sheet_region(ws, header_rows)).encode("utf-8"))
        digest.update(b"\n---\n")
    return digest.hexdigest()


__all__ = ["HEADER_ROWS", "SheetRegion", "clip_text", "fingerprint", "header_bound",
           "model_region", "render_region", "sheet_region"]
