"""Layout resolution — cache → synonyms → model → validate → user, per
document and for the (STTM, FRD[, VDD]) PAIR (M2.5 §3, §7, §9).

Order, enforced in code:

a. **Cache**: ``layout.cache_dirs`` (repo) then the runtime cache, keyed by
   fingerprint (pair fingerprint first for a pair). Hit → source ``cache``,
   zero model calls. M9.1: a RUNTIME entry is also keyed by the vocabulary
   hash (synonym tables + role vocabulary + required roles) — an entry made
   under other tables is stale and ignored; a cached profile that lacks a
   required role is never trusted (re-resolved) and never written; ``refresh``
   bypasses the cache and overwrites the entry.
b. **Synonyms**: the deterministic discovery of M1 / M2. Every required role
   resolved → done, source ``synonyms``.
c. **Model**: ONE call per new document, carrying the fingerprint material
   only (header regions / table labels), the role vocabulary, the synonym
   hints, the partial profile and the unresolved remainder. Only the
   remainder is merged out of the answer: a role the synonyms placed, or a
   column another role already claims, is never overwritten.
d. **Validate**: every merged claim is checked against the document
   (:mod:`codegen.layout.validate`); a failed claim drops to unresolved
   with its reason.
e. **User**: what is still unresolved comes back as questions (role,
   sheet, the header row rendered, candidate columns) — a missing REQUIRED
   role is always among them. Answers merge with source ``user`` and
   confidence 1.0 and may place ANY role, open or not: an answer that
   conflicts with a synonym / model / cache placement wins. A complete
   profile is saved to the runtime cache.
f. A profile with unresolved roles still extracts — those values read
   empty and the gate flags them.

Pair level: unresolved roles of both documents form ONE question list;
cross-document checks (feed name, target catalog/schema, format /
delimiter / frequency vs the STTM meta rows and the VDD FILES sheet) only
adjust confidence and raise flags — they never change a value.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath

from openpyxl.utils import get_column_letter
from pydantic import ValidationError

from codegen.config import Config
from codegen.contracts.frd import FrdContract
from codegen.layout.discover import (
    Discovery,
    discover,
    discover_vdd,
    is_blank_placeholder,
    normalize,
    stale_value_evidence,
    text,
)
from codegen.layout.extent import load_document
from codegen.layout.fingerprint import fingerprint, model_region, render_region, sheet_region
from codegen.layout.frd_profile import (
    FrdFieldSource,
    FrdLayoutProfile,
    FrdUnresolved,
    frd_fingerprint,
)
from codegen.layout.hints import frd_field_help, role_help
from codegen.layout.model import (
    RAW_LOG_CHARS,
    LayoutModelProvider,
    LayoutProviderError,
    build_frd_request,
    build_sttm_request,
)
from codegen.layout.profile import (
    BAND_LAYER_CHOICES,
    BAND_LAYER_ROLE,
    BAND_REF_RE,
    REQUIRED_ROLES,
    BandProfile,
    LayoutProfile,
    MetaRow,
    Role,
    SheetProfile,
    UnresolvedRole,
    confidence_key,
    is_band_layer_key,
    missing_required_roles,
)
from codegen.layout.validate import Rejection, validate_profile
from codegen.run_errors import Coord, located

_STEM_MIN = 4


# --------------------------------------------------------------- results


@dataclass(frozen=True)
class LayoutQuestion:
    document: str                  # "sttm" | "frd"
    sheet: str | None
    layer: str | None
    role: str
    reason: str
    header: list[str] = field(default_factory=list)          # "C: Field Name" …
    candidates: list[dict] = field(default_factory=list)     # {col, header} / {table,row,col,label}
    # Plain-language help (layout/hints.py): what is asked, where it usually
    # sits in the document, and the candidate the dialog pre-selects.
    title: str = ""
    hint: str = ""
    suggested: int | None = None
    # WHY that candidate is pre-selected (a deterministic rule, named) — shown
    # next to the pre-selection and given to the model so its advice can
    # confirm or reject it on the same basis.
    suggested_reason: str = ""
    # role: place a column / an FRD table cell (candidates {col,header} or
    # {table,row,col,label}); choice: pick between documents' values
    # (candidates {value, source, cell}); layer: apply a single stated load
    # strategy to stage / standard / both (candidates {layer, value}).
    kind: str = "role"

    @property
    def key(self) -> str:
        # A role question of a workbook is keyed by where the role sits; every
        # other kind (choice / layer / text) carries its own key in ``role``.
        if self.document in ("sttm", "vdd") and self.kind == "role":
            return f"{self.sheet}/{self.layer}/{self.role}"
        return self.role

    def as_dict(self) -> dict:
        return {"document": self.document, "key": self.key, "sheet": self.sheet,
                "layer": self.layer, "role": self.role, "reason": self.reason,
                "header": self.header, "candidates": self.candidates,
                "title": self.title, "hint": self.hint, "suggested": self.suggested,
                "suggested_reason": self.suggested_reason, "kind": self.kind}


@dataclass(frozen=True)
class CrossCheck:
    name: str
    status: str                     # agree | disagree | unchecked
    frd_citation: str
    sttm_citation: str
    detail: str

    def render(self) -> str:
        return f"{self.name}: {self.status} — FRD {self.frd_citation}; STTM {self.sttm_citation}" \
               f"{' — ' + self.detail if self.detail else ''}"


@dataclass
class DocumentResolution:
    document: str
    profile: LayoutProfile | FrdLayoutProfile
    questions: list[LayoutQuestion] = field(default_factory=list)
    rejections: list[Rejection] = field(default_factory=list)
    provider_calls: int = 0
    cache_hit: bool = False
    fingerprint: str = ""

    @property
    def sources(self) -> dict[str, int]:
        counts = {"synonyms": 0, "model": 0, "user": 0, "cache": 0}
        profile = self.profile
        if isinstance(profile, LayoutProfile):
            keys = [confidence_key(s.name, b.layer, r) for s in profile.sheets
                    for b in s.bands for r in b.roles]
            for key in keys:
                counts[profile.role_sources.get(key, profile.source)] += 1
        else:
            for key in profile.fields:
                counts[profile.field_sources.get(key, profile.source)] += 1
        return counts

    @property
    def complete(self) -> bool:
        return not self.profile.unresolved


@dataclass
class PairResolution:
    sttm: DocumentResolution
    frd: DocumentResolution | None
    frd_contract: FrdContract | None
    cross_checks: list[CrossCheck] = field(default_factory=list)
    flags: list[str] = field(default_factory=list)
    pair_fingerprint: str | None = None
    pair_cache_hit: bool = False
    # M3: the Vendor Data Dictionary, the pair's third document.
    vdd: DocumentResolution | None = None
    # FRD fields taken from another document (or the person's choice):
    # [{field, title, value, source, cell}] — shown as "Taken from other
    # documents"; each one also a frd_unstated flag.
    gap_fills: list[dict] = field(default_factory=list)
    # M9.2: the person's feeds[i].fields[<name>].width answers, parsed —
    # {feed index: {normalized field name: width}}; handed to extract-sttm.
    width_answers: dict = field(default_factory=dict)

    @property
    def documents(self) -> list[DocumentResolution]:
        return [d for d in (self.sttm, self.frd, self.vdd) if d is not None]

    @property
    def questions(self) -> list[LayoutQuestion]:
        return [q for d in self.documents for q in d.questions]

    @property
    def provider_calls(self) -> int:
        return sum(d.provider_calls for d in self.documents)

    @property
    def rejections(self) -> list[Rejection]:
        return [r for d in self.documents for r in d.rejections]

    def report(self) -> dict:
        return {
            "pair_fingerprint": self.pair_fingerprint,
            "pair_cache_hit": self.pair_cache_hit,
            "provider_calls": self.provider_calls,
            "sttm": _document_report(self.sttm),
            "frd": _frd_report(self.frd) if self.frd else None,
            "vdd": _document_report(self.vdd) if self.vdd else None,
            "gap_fills": list(self.gap_fills),
            "cross_checks": [c.render() for c in self.cross_checks],
            "flags": list(self.flags),
            "questions": [q.as_dict() for q in self.questions],
        }


def _frd_report(doc: DocumentResolution) -> dict:
    """The FRD's report adds its family and, for F2, which signal found the
    Solution Requirement tables (D1 post51: caption or header-row fingerprint)."""
    from codegen.extract.frd_docx import F2_SIGNAL_NOTE

    return {**_document_report(doc), "family": doc.profile.family,
            "family_signals": [n for n in doc.profile.notes if n.startswith(F2_SIGNAL_NOTE)]}


def _document_report(doc: DocumentResolution) -> dict:
    return {
        "fingerprint": doc.fingerprint,
        "source": doc.profile.source,
        "cache_hit": doc.cache_hit,
        "provider_calls": doc.provider_calls,
        "roles_by_source": doc.sources,
        "rejections": [r.render() for r in doc.rejections],
        "unresolved": [f"{u.sheet}/{u.layer}/{u.role}: {u.reason}"
                       if isinstance(u, UnresolvedRole) else f"{u.field}: {u.reason}"
                       for u in doc.profile.unresolved],
    }


# ------------------------------------------------------------------ cache


def vocabulary_hash(config: Config) -> str:
    """The second half of a runtime cache key (M9.1): every table discovery
    reads (``extractor:`` — header / label / segment synonyms, band tokens,
    thresholds), the role vocabulary and the required roles. A profile cached
    under other tables answers a different question and is ignored."""
    material = {
        "extractor": config.extractor.model_dump(mode="json"),
        "roles": [r.value for r in Role],
        "required": {layer: [r.value for r in roles] for layer, roles in REQUIRED_ROLES.items()},
    }
    return hashlib.sha256(json.dumps(material, sort_keys=True, ensure_ascii=False)
                          .encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class _Caches:
    """The curated repo cache (``layout.cache_dirs`` — tracked, test-pinned to
    a fresh build, also the mock provider's answers, so plain profiles) and
    the RUNTIME cache, whose entries carry the vocabulary hash."""

    repo: list[Path]
    runtime: Path | None
    vocabulary: str


def _cache_dirs(config: Config, base_dir: Path, runtime_cache_dir: Path | None,
                cache_dirs: list[Path] | None) -> _Caches:
    repo = ([Path(d) for d in cache_dirs] if cache_dirs is not None
            else [base_dir / d for d in config.layout.cache_dirs])
    runtime = Path(runtime_cache_dir) if runtime_cache_dir is not None else None
    return _Caches(repo=repo, runtime=runtime, vocabulary=vocabulary_hash(config))


def _load_cached(fingerprint_value: str, caches: _Caches, prefix: str = "") -> dict | None:
    """The cached payload for a fingerprint, the ``vocabulary`` key removed. A
    runtime entry made under another vocabulary, or tombstoned by a refresh
    that could not complete, is skipped."""
    for directory in [*caches.repo, *([caches.runtime] if caches.runtime else [])]:
        if not directory.is_dir():
            continue
        runtime = caches.runtime is not None and directory == caches.runtime
        for path in sorted(directory.glob(f"{prefix}*.json")):
            try:
                payload = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                continue
            if not isinstance(payload, dict):
                continue
            if payload.get("fingerprint") == fingerprint_value or (
                    prefix and payload.get("pair_fingerprint") == fingerprint_value):
                if runtime and (payload.get("invalidated")
                                or payload.get("vocabulary") != caches.vocabulary):
                    continue
                return {k: v for k, v in payload.items() if k != "vocabulary"}
    return None


def _stale_reason(profile: LayoutProfile) -> str | None:
    missing = missing_required_roles(profile)
    if not missing:
        return None
    return f"cached profile lacks required role(s) {missing}; ignored and re-resolved"


def _invalidate_runtime(caches: _Caches, name: str, key: str, value: str) -> None:
    """A refresh that could not complete tombstones the runtime entry (an
    overwrite — the storage roles have no delete), so the old profile is not
    served again."""
    if caches.runtime is None or not (caches.runtime / name).is_file():
        return
    _save_runtime({key: value, "invalidated": True, "vocabulary": caches.vocabulary},
                  caches.runtime, name)


# Keys a schema-error location may name: the profile schemas' own field names
# and the role vocabulary. Anything else is a string the model invented.
_SCHEMA_KEYS = (
    {name for model in (LayoutProfile, UnresolvedRole, FrdLayoutProfile, FrdFieldSource,
                        FrdUnresolved) for name in model.model_fields}
    | {"bands", "meta_rows", "roles", "layer", "col_start", "col_end", "label", "row", "col",
       "key", "value_col", "name", "kind", "header_row", "band_row", "segment_strategy",
       "segment_column", "headers"}
    | {r.value for r in Role})


def _schema_errors(exc: ValidationError) -> list[dict]:
    return [{"type": e["type"], "msg": e["msg"],
             "loc": [p if isinstance(p, int) or p in _SCHEMA_KEYS else "<key>"
                     for p in e["loc"]]}
            for e in exc.errors(include_input=False, include_url=False)]


def _canonical(key: str) -> str:
    from codegen.answer_keys import canonical_key

    return canonical_key(key)


def _cached_schema_reason(exc: ValidationError) -> str:
    """post133 (audit #9): pydantic's text quotes the input values (labels a
    document holds) - only the masked error locations and messages are kept."""
    errors = _schema_errors(exc)
    first = f"; first: {_error_line(errors[0])}" if errors else ""
    return f"cached profile failed schema validation: {len(errors)} error(s){first}"


def _error_line(error: dict) -> str:
    where = ".".join(str(p) for p in error["loc"]) or "<answer>"
    return f"model answer schema error at {where}: {error['msg']}"


def _salvage_sheet(raw) -> tuple[SheetProfile | None, int]:
    """(the sheet's valid part, bands dropped). A band that fails the schema
    is dropped, the rest of the sheet kept; a sheet whose own fields fail is
    dropped whole."""
    if not isinstance(raw, dict):
        return None, 0
    bands, dropped = [], 0
    for band in raw.get("bands") or []:
        try:
            bands.append(BandProfile.model_validate(band))
        except ValidationError:
            dropped += 1
    meta_rows = []
    for row in raw.get("meta_rows") or []:
        try:
            meta_rows.append(MetaRow.model_validate(row))
        except ValidationError:
            continue
    whole = {**raw, "bands": bands, "meta_rows": meta_rows}
    slim = {k: raw[k] for k in ("name", "kind", "header_row", "band_row") if k in raw}
    for candidate in (whole, {**slim, "bands": bands, "meta_rows": meta_rows}):
        try:
            return SheetProfile.model_validate(candidate), dropped
        except ValidationError:
            continue
    return None, dropped + len(bands)


def _drop_null_roles(answer):
    """A role the model maps to ``null`` is a role it did not find — the
    prompt's "leave a role out", said another way; it is no schema error."""
    if not isinstance(answer, dict) or not isinstance(answer.get("sheets"), list):
        return answer
    sheets = []
    for sheet in answer["sheets"]:
        if isinstance(sheet, dict) and isinstance(sheet.get("bands"), list):
            sheet = {**sheet, "bands": [
                {**b, "roles": {k: v for k, v in b["roles"].items() if v is not None}}
                if isinstance(b, dict) and isinstance(b.get("roles"), dict) else b
                for b in sheet["bands"]]}
        sheets.append(sheet)
    return {**answer, "sheets": sheets}


def salvage_layout_answer(answer, partial: LayoutProfile
                          ) -> tuple[LayoutProfile | None, list[str], list[dict]]:
    """(profile, report lines, schema errors) for a model's layout answer.

    A schema-valid answer is used as before (no line). Otherwise EVERY
    validation error is a report line (location keys outside the profile
    schema masked ``<key>``, never an input value) and the answer is
    salvaged part by part: each sheet whose own fields validate is kept with
    its valid bands; the top-level fields come from the partial profile. The
    kept part then meets the same workbook validation as a whole answer —
    nothing the model says reaches a profile unchecked. None when no sheet
    survives."""
    answer = _drop_null_roles(answer)
    try:
        return LayoutProfile.model_validate(answer), [], []
    except ValidationError as exc:
        errors = _schema_errors(exc)
    raw_sheets = answer.get("sheets") if isinstance(answer, dict) else None
    kept: list[SheetProfile] = []
    dropped_bands = 0
    for raw in raw_sheets if isinstance(raw_sheets, list) else []:
        sheet, dropped = _salvage_sheet(raw)
        dropped_bands += dropped
        if sheet is not None:
            kept.append(sheet)
    total = len(raw_sheets) if isinstance(raw_sheets, list) else 0
    outcome = (f"kept {len(kept)} of {total} sheet(s), {dropped_bands} band(s) dropped — "
               "the kept part is validated against the workbook like a whole answer"
               if kept else "nothing usable kept")
    lines = [f"model response failed schema validation: {len(errors)} error(s); {outcome}",
             *(_error_line(e) for e in errors)]
    if not kept:
        return None, lines, errors
    profile = LayoutProfile(fingerprint=partial.fingerprint, sheets=kept, source="model",
                            strategy=partial.strategy)
    return profile, lines, errors


def _log_rejections(caches: _Caches, digest: str, document: str, rejections: list[Rejection],
                    schema_errors: list[dict]) -> Path | None:
    """``<runtime cache>/rejections/<fingerprint>.json`` — why a model answer
    was rejected, in full (the report line keeps only the first line of a
    schema failure). Reasons and error locations only: a location key outside
    the profile schema / role vocabulary is written as ``<key>``; no workbook
    cell and no model string is stored."""
    if caches.runtime is None or not (rejections or schema_errors):
        return None
    payload = {
        "fingerprint": digest,
        "document": document,
        "vocabulary": caches.vocabulary,
        "rejections": [r.render() for r in rejections],
        "schema_errors": schema_errors,
    }
    return _save_runtime(payload, caches.runtime / "rejections", f"{digest}.json")


def _log_raw_responses(caches: _Caches, digest: str, document: str,
                       raw: list[str]) -> Path | None:
    """D1-only diagnostics (2026-10-13, pair_11 "Expecting value: line 1
    column 1"): the first ``RAW_LOG_CHARS`` characters of every model response
    that was rejected, on stderr and in ``<runtime cache>/rejections/
    <fingerprint>.raw.json`` (the state role — never a tracked file, a
    contract or a report). A response can echo the header labels it was
    shown, so this stays on the workspace that ran it."""
    if not raw:
        return None
    for number, response in enumerate(raw, 1):
        print(f"RAW (D1 only)  {document} model response {number}/{len(raw)}, first "
              f"{RAW_LOG_CHARS} chars: {response!r}", file=sys.stderr)
    if caches.runtime is None:
        return None
    payload = {"fingerprint": digest, "document": document, "chars": RAW_LOG_CHARS,
               "responses": raw}
    return _save_runtime(payload, caches.runtime / "rejections", f"{digest}.raw.json")


def _save_runtime(payload: dict, runtime_cache_dir: Path | None, name: str) -> Path | None:
    if runtime_cache_dir is not None and "fixtures" in Path(runtime_cache_dir).parts:
        # A runtime profile carries the document's REAL sheet names; fixtures/
        # is tracked. Fixture profiles are built by scripts/build_layout_profiles.py.
        raise ValueError(f"runtime layout profiles are never written under fixtures/ "
                         f"(got {runtime_cache_dir})")
    if runtime_cache_dir is None:
        return None
    runtime_cache_dir.mkdir(parents=True, exist_ok=True)
    path = runtime_cache_dir / name
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
                    encoding="utf-8", newline="\n")
    return path


def _as_cache(profile: LayoutProfile) -> LayoutProfile:
    keys = [confidence_key(s.name, b.layer, r) for s in profile.sheets for b in s.bands
            for r in b.roles]
    return profile.model_copy(update={"source": "cache",
                                      "role_sources": {k: "cache" for k in keys}})


class _StaleCache(Exception):
    """A cached profile that must not be trusted (a required role missing)."""


# ------------------------------------------------------------- STTM side


def _header_strip(ws, header_row: int) -> list[str]:
    cells = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    return [f"{get_column_letter(i)}: {text(c)}" for i, c in enumerate(cells, start=1)
            if text(c) is not None]


def band_layer_question(document: str, sheet: str, layer: str, reason: str,
                        header: list[str] | None = None) -> LayoutQuestion:
    """Chunk A: the question for a band nothing names — answered with a LAYER
    (source | stage | standard), never a column. Also what the App lists for
    a feed held back on ``<sheet>/band[<n>]/layer`` (ui/backend/demo.py)."""
    return LayoutQuestion(
        document=document, sheet=sheet, layer=layer, role=BAND_LAYER_ROLE,
        reason=reason, header=list(header or []),
        candidates=[{"value": choice, "label": choice} for choice in BAND_LAYER_CHOICES],
        title=f"Which layer is {layer} of sheet {sheet!r}?",
        hint=("These columns are shaped like a target band (schema / table / column / "
              "data type) but nothing in the document names their layer: no band "
              "title over them, no layer word in their headers, no catalog or schema "
              "value that says stage or standard. Column order is not evidence. "
              "Answer source, stage or standard; the sheet's feed waits for it."))


def _questions_for(profile: LayoutProfile, workbook,
                   document: str = "sttm") -> list[LayoutQuestion]:
    questions: list[LayoutQuestion] = []
    for item in profile.unresolved:
        sheet = profile.sheet(item.sheet)
        if sheet is None or sheet.header_row is None or item.sheet not in workbook.sheetnames:
            continue
        ws = workbook[item.sheet]
        header = next(ws.iter_rows(min_row=sheet.header_row, max_row=sheet.header_row,
                                   values_only=True))
        if BAND_REF_RE.match(str(item.layer)) and item.role == BAND_LAYER_ROLE:
            # Chunk A: a band nothing names - answered with a LAYER, not a column.
            questions.append(band_layer_question(document, item.sheet, str(item.layer),
                                                 item.reason,
                                                 _header_strip(ws, sheet.header_row)))
            continue
        band = sheet.band(item.layer)  # type: ignore[arg-type]
        claimed = set(band.roles.values()) if band else set()
        span = range(band.col_start, band.col_end + 1) if band else range(1, len(header) + 1)
        candidates = [{"col": c, "header": text(header[c - 1])} for c in span
                      if c not in claimed and c - 1 < len(header) and text(header[c - 1])]
        title, hint, suggested = role_help(item.role, item.layer, candidates)
        questions.append(LayoutQuestion(
            document=document, sheet=item.sheet, layer=item.layer, role=item.role,
            reason=item.reason, header=_header_strip(ws, sheet.header_row),
            candidates=candidates, title=title, hint=hint, suggested=suggested))
    return questions


def _merge_columns(profile: LayoutProfile, claims: dict[str, int], source: str,
                   confidence: float, *, override: bool = False) -> LayoutProfile:
    """Merge ``{"<sheet>/<layer>/<role>": col}`` claims. A model's claims fill
    ONLY roles the profile has not placed, onto columns no other role of that
    band claims. A person's (``override=True``, M9.1) may set ANY role: the
    answer replaces a synonym / model / cache placement of that role, and a
    role that held the answered column gives it up (a required one is then
    listed as unresolved again); each displacement is a profile note."""
    sheets = []
    conf = dict(profile.confidence)
    role_sources = dict(profile.role_sources)
    notes = list(profile.notes)
    vocabulary = {r.value for r in Role}
    for sp in profile.sheets:
        bands = []
        for band in sp.bands:
            roles = dict(band.roles)
            for key, col in claims.items():
                parts = key.split("/")
                if len(parts) != 3 or parts[0] != sp.name or parts[1] != band.layer:
                    continue
                role = parts[2]
                if not isinstance(col, int) or role not in vocabulary:
                    continue
                if roles.get(role) == col:
                    if override:
                        conf[key] = confidence
                        role_sources[key] = source
                    continue
                if not override and (role in roles or col in roles.values()):
                    continue
                if role in roles:
                    was = role_sources.get(key, profile.source)
                    notes.append(f"{key}: user answer moved the role from column {roles[role]} "
                                 f"({was}) to column {col}")
                for other, other_col in list(roles.items()):
                    if other_col == col and other != role:
                        other_key = confidence_key(sp.name, band.layer, other)
                        notes.append(f"{other_key}: column {col} given to {role!r} by a user "
                                     "answer; the role is open again")
                        del roles[other]
                        conf.pop(other_key, None)
                        role_sources.pop(other_key, None)
                roles[role] = col
                conf[key] = confidence
                role_sources[key] = source
            bands.append(band.model_copy(update={"roles": roles}))
        source_band = next((b for b in bands if b.layer == "source"), None)
        segment_column = sp.segment_column
        strategy = sp.segment_strategy
        if (source_band is not None and source_band.column(Role.SEGMENT) is not None
                and segment_column is None):
            segment_column = source_band.column(Role.SEGMENT)
            strategy = "column"
        sheets.append(sp.model_copy(update={"bands": bands, "segment_column": segment_column,
                                            "segment_strategy": strategy}))
    resolved_keys = {confidence_key(s.name, b.layer, r) for s in sheets for b in s.bands
                     for r in b.roles}
    unresolved = [u for u in profile.unresolved
                  if confidence_key(u.sheet, u.layer, u.role) not in resolved_keys]
    merged = profile.model_copy(update={"sheets": sheets, "confidence": conf,
                                        "role_sources": role_sources, "unresolved": unresolved,
                                        "notes": notes})
    return _list_missing_required(merged) if override else merged


def _list_missing_required(profile: LayoutProfile) -> LayoutProfile:
    """Every REQUIRED role a band does not place is in ``unresolved`` — and so
    in the question list — whoever produced the profile (M9.1: a model-made
    profile once omitted the schema AND the note that it was missing)."""
    listed = {confidence_key(u.sheet, u.layer, u.role) for u in profile.unresolved}
    added = []
    for key in missing_required_roles(profile):
        if key in listed:
            continue
        sheet_name, layer, role = key.rsplit("/", 2)
        band = profile.sheet(sheet_name).band(layer)  # type: ignore[union-attr,arg-type]
        added.append(UnresolvedRole(
            sheet=sheet_name, layer=layer, role=role,  # type: ignore[arg-type]
            reason="required role is not placed",
            candidates=[c for c in range(band.col_start, band.col_end + 1)
                        if c not in band.roles.values()]))
    if not added:
        return profile
    return profile.model_copy(update={"unresolved": [*profile.unresolved, *added]})


def _strip_unknown_roles(model_profile: LayoutProfile) -> tuple[LayoutProfile, int]:
    """Drop role names outside the vocabulary (and every note) from a model
    answer BEFORE anything reads it — nothing textual the model invented is
    ever kept, quoted or reported."""
    vocabulary = {r.value for r in Role}
    dropped = 0
    sheets = []
    for sp in model_profile.sheets:
        bands = []
        for band in sp.bands:
            roles = {r: c for r, c in band.roles.items() if r in vocabulary}
            dropped += len(band.roles) - len(roles)
            bands.append(band.model_copy(update={"roles": roles, "label": None,
                                                 "layer_evidence": None}))
        sheets.append(sp.model_copy(update={"bands": bands, "notes": []}))
    return model_profile.model_copy(update={"sheets": sheets, "notes": [], "unresolved": []}), \
        dropped


def _claims_from_model(model_profile: LayoutProfile) -> dict[str, int]:
    return {confidence_key(s.name, b.layer, r): c
            for s in model_profile.sheets for b in s.bands for r, c in b.roles.items()}


def resolve_workbook(path: Path, config: Config, *, provider: LayoutModelProvider | None = None,
                     cache_dirs: list[Path] | None = None, runtime_cache_dir: Path | None = None,
                     answers: dict[str, int] | None = None, base_dir: Path | None = None,
                     use_cache: bool = True, document: str = "sttm",
                     sttm_tables: list[str] | None = None,
                     refresh: bool = False,
                     prior: DocumentResolution | None = None
                     ) -> tuple[DocumentResolution, object]:
    """Resolve a workbook's layout — an STTM (``document="sttm"``) or a Vendor
    Data Dictionary (``document="vdd"``, discovered by ``discover_vdd`` and
    narrowed to ``sttm_tables``); the rest of the path is identical.
    ``refresh`` (M9.1) bypasses every cache and OVERWRITES the runtime entry:
    a complete result replaces it, an incomplete one tombstones it.
    ``prior`` (M9.1b) is this document's resolution from an earlier pass of
    the SAME invocation (the CLI resolves once to learn the questions, then
    again with the answers file): the second pass continues from it — no
    cache read, no second model call, every role keeps the source that really
    placed it (a refresh used to come back ``source=cache``: it re-read the
    entry its own first pass had just written)."""
    base = base_dir if base_dir is not None else Path(".")
    caches = _cache_dirs(config, base, runtime_cache_dir, cache_dirs)
    workbook = load_document(path, config.extractor.used_range_empty_rows)
    digest = fingerprint(workbook)
    rejections: list[Rejection] = []
    schema_errors: list[dict] = []
    calls = 0
    # Chunk A: an answer to <sheet>/band[<n>]/layer is a LAYER, applied by
    # discovering again with it (the band then exists and its roles resolve);
    # every other answer places a column.
    band_answers = {k: v for k, v in (answers or {}).items() if is_band_layer_key(k)}
    answers = {k: v for k, v in (answers or {}).items() if not is_band_layer_key(k)} or None
    band_layers = {k.rsplit("/", 1)[0]: str(v) for k, v in band_answers.items()}

    if (prior is not None and isinstance(prior.profile, LayoutProfile)
            and prior.fingerprint == digest and not band_answers):
        profile = prior.profile
        rejections = list(prior.rejections)
        if answers:
            profile = _merge_columns(profile, answers, "user", 1.0, override=True)
            profile = profile.model_copy(update={"source": "user"})
            profile, user_rejections = validate_profile(
                profile, workbook, config.extractor, check_sources={"user"}, document=document)
            rejections += user_rejections
        profile = _list_missing_required(profile)
        doc = DocumentResolution(document, profile, rejections=rejections,
                                 provider_calls=prior.provider_calls,
                                 cache_hit=prior.cache_hit, fingerprint=digest)
        doc.questions = _questions_for(profile, workbook, document)
        if not profile.unresolved:
            if caches.runtime is not None and (refresh or prior.provider_calls or answers):
                _save_runtime({**profile.model_dump(mode="json"),
                               "vocabulary": caches.vocabulary}, caches.runtime,
                              f"{digest}.json")
        elif refresh:
            _invalidate_runtime(caches, f"{digest}.json", "fingerprint", digest)
        return doc, workbook

    cached = (_load_cached(digest, caches)
              if use_cache and not refresh and not band_answers else None)
    if cached is not None:
        try:
            profile = _as_cache(LayoutProfile.model_validate(cached))
            # A layer read from catalog / schema VALUES is never trusted from
            # a cache (the fingerprint hashes headers, not values): re-derived.
            stale = _stale_reason(profile) or (
                stale_value_evidence(profile, workbook, config.extractor)
                if document == "sttm" else None)
            if stale is not None:
                raise _StaleCache(stale)
            doc = DocumentResolution(document, profile, cache_hit=True, fingerprint=digest)
            if answers:
                profile = _merge_columns(profile, answers, "user", 1.0, override=True)
                profile, user_rejections = validate_profile(
                    profile, workbook, config.extractor, check_sources={"user"},
                    document=document)
                profile = _list_missing_required(profile)
                doc.rejections += user_rejections
                doc.profile = profile
                if not profile.unresolved:
                    _save_runtime({**profile.model_dump(mode="json"),
                                   "vocabulary": caches.vocabulary},
                                  caches.runtime, f"{digest}.json")
            doc.questions = _questions_for(doc.profile, workbook, document)
            return doc, workbook
        except ValidationError as exc:
            rejections.append(Rejection(document, None, None, None,
                                        _cached_schema_reason(exc)))
        except _StaleCache as exc:
            rejections.append(Rejection(document, None, None, None, str(exc)))

    found: Discovery = (discover_vdd(path, config.extractor, sttm_tables=sttm_tables)
                        if document == "vdd"
                        else discover(path, config.extractor, band_layers=band_layers or None))
    profile = found.profile
    keys = [confidence_key(s.name, b.layer, r) for s in profile.sheets for b in s.bands
            for r in b.roles]
    profile = profile.model_copy(update={"role_sources": {k: "synonyms" for k in keys}})
    model_allowed = True
    if (band_answers and prior is not None and isinstance(prior.profile, LayoutProfile)
            and prior.fingerprint == digest):
        # The second pass of one invocation, now with band layers answered: no
        # second model call - the first pass's model placements carry over.
        carried = {confidence_key(s.name, b.layer, r): c
                   for s in prior.profile.sheets for b in s.bands for r, c in b.roles.items()
                   if prior.profile.role_sources.get(confidence_key(s.name, b.layer, r))
                   == "model"}
        if carried:
            profile = _merge_columns(profile, carried, "model", config.layout.placed_confidence)
        rejections = list(prior.rejections)
        calls = prior.provider_calls
        model_allowed = False

    model_from = len(rejections)
    # A band's LAYER is never the model's to place (evidence only): the model
    # is asked about column roles only.
    role_gaps = [u for u in profile.unresolved
                 if not (BAND_REF_RE.match(str(u.layer)) and u.role == BAND_LAYER_ROLE)]
    if role_gaps and provider is not None and model_allowed:
        regions = [render_region(model_region(workbook[name])) for name in workbook.sheetnames]
        request = build_sttm_request(digest, regions, profile, role_gaps, config)
        if document == "vdd":
            request["kind"] = "vdd_layout"
            request["synonym_hints"] = {"files": config.extractor.vdd.files_roles,
                                        "fields": config.extractor.vdd.field_roles}
        calls += 1
        model_profile = None
        try:
            answer = provider.complete_layout(request)
        except LayoutProviderError as exc:
            rejections.append(Rejection(document, None, None, None, f"model: {exc}"))
            _log_raw_responses(caches, digest, document, exc.raw)
        else:
            # 2026-10-13 (D1 pair_11): one schema error no longer discards the
            # whole answer — every error is reported, the valid sheets / bands
            # go on to the same workbook validation as before.
            model_profile, error_lines, schema_errors = salvage_layout_answer(answer, profile)
            rejections += [Rejection(document, None, None, None, line) for line in error_lines]
            if error_lines:
                _log_raw_responses(caches, digest, document,
                                   [json.dumps(answer, ensure_ascii=False)[:RAW_LOG_CHARS]])
        if model_profile is not None:
            # Every claim in the answer — sheets, header/band rows, spans,
            # roles — is validated against the workbook FIRST; only the
            # surviving column claims are merged, onto roles the synonyms
            # left open, and the merged profile is validated once more
            # against the synonyms' own band spans.
            model_profile, dropped = _strip_unknown_roles(model_profile)
            if dropped:
                # Never echo a model-invented name into a report.
                rejections.append(Rejection(document, None, None, None,
                                            f"{dropped} role name(s) outside the vocabulary "
                                            "dropped from the model answer"))
            known = {confidence_key(s.name, b.layer, r): c for s in profile.sheets
                     for b in s.bands for r, c in b.roles.items()}
            restructured = {
                s.name for s in model_profile.sheets
                if profile.sheet(s.name) is not None and (
                    (s.header_row, s.band_row)
                    != (profile.sheet(s.name).header_row, profile.sheet(s.name).band_row))}
            answer_keys = _claims_from_model(model_profile)
            answer_profile = model_profile.model_copy(update={
                "source": "model",
                # The delta is the model's claim: a role the synonyms did not
                # place, a different column for one they did, or any role on
                # a sheet whose header/band rows the answer moved.
                "role_sources": {
                    k: ("synonyms" if known.get(k) == c and k.split("/")[0] not in restructured
                        else "model")
                    for k, c in answer_keys.items()}})
            checked, answer_rejections = validate_profile(
                answer_profile, workbook, config.extractor, check_sources={"model"},
                document=document)
            rejections += answer_rejections
            profile = _merge_columns(profile, _claims_from_model(checked), "model",
                                     config.layout.placed_confidence)
            profile = profile.model_copy(update={"source": "model"})
            profile, model_rejections = validate_profile(
                profile, workbook, config.extractor, check_sources={"model"}, document=document)
            rejections += [r for r in model_rejections if r not in rejections]

    if calls:
        # The full reasons a model answer was rejected (the report line keeps
        # one) — before the person's answers add rejections of their own.
        _log_rejections(caches, digest, document, rejections[model_from:], schema_errors)

    if answers:
        profile = _merge_columns(profile, answers, "user", 1.0, override=True)
        profile = profile.model_copy(update={"source": "user"})
        profile, user_rejections = validate_profile(profile, workbook, config.extractor,
                                                    check_sources={"user"}, document=document)
        rejections += user_rejections

    profile = _list_missing_required(profile)
    doc = DocumentResolution(document, profile, rejections=rejections, provider_calls=calls,
                             fingerprint=digest)
    doc.questions = _questions_for(profile, workbook, document)
    if not profile.unresolved:
        # Never reached by a profile with a required role missing (it is in
        # ``unresolved``). A refresh overwrites even a synonyms-only result.
        if caches.runtime is not None and (refresh or calls + len(answers or {})
                                           + len(band_answers)):
            _save_runtime({**profile.model_dump(mode="json"), "vocabulary": caches.vocabulary},
                          caches.runtime, f"{digest}.json")
    elif refresh:
        _invalidate_runtime(caches, f"{digest}.json", "fingerprint", digest)
    return doc, workbook


# -------------------------------------------------------------- FRD side


def _frd_required(profile: FrdLayoutProfile) -> list[FrdUnresolved]:
    return list(profile.unresolved)


def _frd_labels(content) -> list[str]:
    """The fingerprint material of a document: per table, the title and
    the label texts (never a value cell)."""
    lines = []
    for index, rows in enumerate(content.tables):
        title = " | ".join(c.strip() for c in rows[0]) if rows else ""
        lines.append(f"table {index}: {title}")
        first_title = rows[0][0].strip() if rows and rows[0] else ""
        for row_index, row in enumerate(rows[1:], start=1):
            if not row:
                continue
            label = row[0].strip()
            col = 0
            if label == first_title and len(row) > 1:
                label, col = row[1].strip(), 1
            lines.append(f"  row {row_index} col {col}: {label}")
    return lines


def _frd_questions(profile: FrdLayoutProfile, content, config: Config) -> list[LayoutQuestion]:
    from codegen.extract.frd_docx import is_prefix_cell, normalize_label, section_title

    sections = {normalize_label(s): key for key, spellings in
                config.extractor.frd.section_titles.items() for s in spellings}
    used = {(f.table, f.row) for f in profile.fields.values()}
    candidates = []
    for ref in profile.sections:
        rows = content.tables[ref.table]
        titled = section_title(rows, sections)
        token = titled[2] if titled is not None else None
        for row_index, row in enumerate(rows[1:], start=1):
            if (ref.table, row_index) in used or not row:
                continue
            col = 1 if len(row) > 1 and (row[0].strip() == rows[0][0].strip() or (
                token is not None and is_prefix_cell(row[0], ref.title, token))) else 0
            if row[col].strip():
                candidates.append({"table": ref.table, "row": row_index, "col": col,
                                   "label": row[col].strip(), "section": ref.section})
    questions = []
    for item in profile.unresolved:
        title, hint, suggested = frd_field_help(item.field, candidates, config)
        questions.append(LayoutQuestion(document="frd", sheet=None, layer=None, role=item.field,
                                        reason=item.reason, header=[], candidates=candidates,
                                        title=title, hint=hint, suggested=suggested))
    return questions


# post129 (audit #6): the paths whose DQ-section rows read "<label>: <value>" -
# the synonym path's own rule (extract.frd_docx, the positional DQ rule rows)
_LABELLED_RULE_PATH = re.compile(r"\.(validation_rules\[\d+\]|recycle_rule)$")


def _validate_frd(profile: FrdLayoutProfile, content, config: Config,
                  check_sources: set[str]) -> tuple[FrdLayoutProfile, list[Rejection]]:
    from codegen.extract.frd_docx import _inline_pairs, normalize_label

    sections = {normalize_label(s): key for key, spellings in
                config.extractor.frd.section_titles.items() for s in spellings}
    fields = dict(profile.fields)
    confidence = dict(profile.confidence)
    field_sources = dict(profile.field_sources)
    rejections: list[Rejection] = []
    unresolved = list(profile.unresolved)
    for path, source in list(fields.items()):
        if field_sources.get(path, profile.source) not in check_sources:
            continue
        reason = None
        if source.table >= len(content.tables):
            reason = f"table {source.table} does not exist"
        elif source.row >= len(content.tables[source.table]):
            reason = f"table {source.table} has no row {source.row}"
        else:
            row = content.tables[source.table][source.row]
            if source.col >= len(row):
                reason = f"row {source.row} has no column {source.col}"
            elif normalize_label(row[source.col]) != normalize_label(source.label):
                # post133 (audit #9): coordinates only - never the cell, never the claim
                reason = (f"the label cell (table {source.table} row {source.row} col "
                          f"{source.col}) does not read as the claimed label")
            elif source.value_col is not None and source.value_col <= source.col:
                reason = "value column is not to the right of the label"
            elif source.section is not None and source.section not in set(sections.values()):
                reason = "the claimed section is not a known metadata section"
            else:
                # post129 (audit #6): the model said WHERE; what the contract
                # quotes is the document's - its label cell as written, the
                # inline pair as the cell spells it, labelled_rule as the
                # synonym path derives it (never the model's flag).
                update = {"label": row[source.col],
                          "labelled_rule": bool(source.section == "data_quality"
                                                and _LABELLED_RULE_PATH.search(path))}
                if source.inline_label is not None:
                    text = (row[source.value_col] if source.value_col is not None
                            and source.value_col < len(row) else "")
                    stated = next((label for label, _value in _inline_pairs(
                        text, config.extractor.frd)
                        if normalize_label(label) == normalize_label(source.inline_label)),
                        None)
                    if stated is None:
                        reason = (f"the claimed inline label names no 'Label: value' pair "
                                  f"in the value cell (table {source.table} row {source.row} "
                                  f"col {source.value_col})")
                    else:
                        update["inline_label"] = stated
                if reason is None:
                    fields[path] = source.model_copy(update=update)
        if reason:
            rejections.append(Rejection("frd", None, None, path, reason))
            fields.pop(path)
            confidence.pop(path, None)
            field_sources.pop(path, None)
            unresolved.append(FrdUnresolved(field=path, feed_index=source.feed_index,
                                            reason=f"rejected: {reason}"))
    return profile.model_copy(update={"fields": fields, "confidence": confidence,
                                      "field_sources": field_sources,
                                      "unresolved": unresolved}), rejections


def _merge_frd(profile: FrdLayoutProfile, claims: dict[str, dict], source: str,
               confidence: float, *, allowed: set[str] | None = None) -> FrdLayoutProfile:
    """``allowed`` (post133, audit #9): the paths that may be claimed - for the
    model, the fields it was asked about; a claim under any other path (a key
    the model invented) is dropped, never merged nor quoted."""
    from codegen.answer_keys import canonical_key

    fields = dict(profile.fields)
    conf = dict(profile.confidence)
    field_sources = dict(profile.field_sources)
    for claimed, claim in claims.items():
        # post74 (D1 post73 pair_04): a claim under a header label
        # ("feeds[0].Delimiter") is the role's path ("feeds[0].delimiter")
        path = canonical_key(claimed)
        if path in fields or not isinstance(claim, dict):
            continue
        if allowed is not None and path not in allowed:
            continue
        try:
            fields[path] = FrdFieldSource.model_validate(claim)
        except ValidationError:
            continue
        conf[path] = confidence
        field_sources[path] = source
    unresolved = [u for u in profile.unresolved if u.field not in fields]
    return profile.model_copy(update={"fields": fields, "confidence": conf,
                                      "field_sources": field_sources, "unresolved": unresolved})


def resolve_frd(path: Path, config: Config, *, provider: LayoutModelProvider | None = None,
                cache_dirs: list[Path] | None = None, runtime_cache_dir: Path | None = None,
                answers: dict[str, dict] | None = None, base_dir: Path | None = None,
                use_cache: bool = True,
                refresh: bool = False,
                prior: DocumentResolution | None = None
                ) -> tuple[DocumentResolution, object]:
    from codegen.extract.frd_docx import discover_frd, read_docx

    base = base_dir if base_dir is not None else Path(".")
    caches = _cache_dirs(config, base, runtime_cache_dir, cache_dirs)
    content = read_docx(path, prose_ok=True)
    digest = frd_fingerprint(content.tables)
    rejections: list[Rejection] = []
    schema_errors: list[dict] = []
    calls = 0

    if (prior is not None and isinstance(prior.profile, FrdLayoutProfile)
            and prior.fingerprint == digest):
        # Same invocation, second pass (see resolve_workbook): continue from it.
        profile = prior.profile
        rejections = list(prior.rejections)
        if answers:
            profile = _merge_frd(profile, answers, "user", 1.0)
            profile = profile.model_copy(update={"source": "user"})
            profile, user_rejections = _validate_frd(profile, content, config, {"user"})
            rejections += user_rejections
        doc = DocumentResolution("frd", profile, rejections=rejections,
                                 provider_calls=prior.provider_calls,
                                 cache_hit=prior.cache_hit, fingerprint=digest)
        doc.questions = _frd_questions(profile, content, config)
        if not profile.unresolved:
            if caches.runtime is not None and (refresh or prior.provider_calls or answers):
                _save_runtime({**profile.model_dump(mode="json"),
                               "vocabulary": caches.vocabulary}, caches.runtime,
                              f"{digest}.json")
        elif refresh:
            _invalidate_runtime(caches, f"{digest}.json", "fingerprint", digest)
        return doc, content

    cached = _load_cached(digest, caches) if use_cache and not refresh else None
    if cached is not None:
        try:
            profile = FrdLayoutProfile.model_validate(cached)
            profile = profile.model_copy(update={
                "source": "cache", "field_sources": {k: "cache" for k in profile.fields}})
            doc = DocumentResolution("frd", profile, cache_hit=True, fingerprint=digest)
            if answers:
                profile = _merge_frd(profile, answers, "user", 1.0)
                profile, user_rejections = _validate_frd(profile, content, config, {"user"})
                doc.rejections += user_rejections
                doc.profile = profile
            doc.questions = _frd_questions(doc.profile, content, config)
            return doc, content
        except ValidationError as exc:
            rejections.append(Rejection("frd", None, None, None, _cached_schema_reason(exc)))

    profile = discover_frd(content, config.extractor.frd)
    profile = profile.model_copy(update={"field_sources": {k: "synonyms" for k in profile.fields}})

    model_from = len(rejections)
    if profile.unresolved and provider is not None:
        request = build_frd_request(digest, _frd_labels(content), profile.model_dump(mode="json"),
                                    [u.model_dump(mode="json") for u in profile.unresolved],
                                    config)
        calls += 1
        try:
            answer = provider.complete_layout(request)
            model_profile = FrdLayoutProfile.model_validate(answer)
        except LayoutProviderError as exc:
            rejections.append(Rejection("frd", None, None, None, f"model: {exc}"))
            _log_raw_responses(caches, digest, "frd", exc.raw)
        except ValidationError as exc:
            rejections.append(Rejection("frd", None, None, None,
                                        "model response failed schema validation: "
                                        f"{str(exc).splitlines()[0]}"))
            schema_errors = _schema_errors(exc)
            _log_raw_responses(caches, digest, "frd",
                               [json.dumps(answer, ensure_ascii=False)[:RAW_LOG_CHARS]])
        else:
            claims = {k: v.model_dump(mode="json") for k, v in model_profile.fields.items()}
            asked = {u.field for u in profile.unresolved}
            dropped = sum(1 for k in claims
                          if _canonical(k) not in asked and _canonical(k) not in profile.fields)
            if dropped:
                rejections.append(Rejection("frd", None, None, None,
                                            f"{dropped} claim(s) for fields not asked dropped "
                                            "from the model answer (not repeated)"))
            profile = _merge_frd(profile, claims, "model", config.layout.placed_confidence,
                                 allowed=asked)
            profile = profile.model_copy(update={"source": "model"})
            profile, model_rejections = _validate_frd(profile, content, config, {"model"})
            rejections += model_rejections
    if calls:
        _log_rejections(caches, digest, "frd", rejections[model_from:], schema_errors)

    if answers:
        profile = _merge_frd(profile, answers, "user", 1.0)
        profile = profile.model_copy(update={"source": "user"})
        profile, user_rejections = _validate_frd(profile, content, config, {"user"})
        rejections += user_rejections

    doc = DocumentResolution("frd", profile, rejections=rejections, provider_calls=calls,
                             fingerprint=digest)
    doc.questions = _frd_questions(profile, content, config)
    if not profile.unresolved:
        if caches.runtime is not None and (refresh or calls + len(answers or {})):
            _save_runtime({**profile.model_dump(mode="json"), "vocabulary": caches.vocabulary},
                          caches.runtime, f"{digest}.json")
    elif refresh:
        _invalidate_runtime(caches, f"{digest}.json", "fingerprint", digest)
    return doc, content


# ------------------------------------------------------------ cross-checks


def _stem_match(name: str, candidate: str) -> bool:
    a, b = normalize(name), normalize(candidate)
    if not a or not b:
        return False
    if a in b:
        return True
    tokens = [t for t in a.split() if len(t) >= _STEM_MIN]
    words = b.split()
    return bool(tokens) and all(any(w.startswith(t) or t.startswith(w) for w in words
                                    if len(w) >= _STEM_MIN) for t in tokens)


def _dominant(values: list[str | None]) -> str | None:
    counts: dict[str, int] = {}
    for value in values:
        if value:
            counts[value] = counts.get(value, 0) + 1
    return max(counts, key=counts.get) if counts else None


def _own_sheet_files(sp, ws, config: Config) -> list[tuple[str, str]]:
    """The file(s) a mapping sheet's own ``File Name`` meta row states."""
    blanks = config.extractor.discovery.meta_blank_values
    found: list[tuple[str, str]] = []
    for entry in sp.meta_rows:
        if entry.key not in ("file_names", "file_name_example") or entry.value_col is None:
            continue
        value = text(ws.cell(row=entry.row, column=entry.value_col).value)
        if not value or is_blank_placeholder(value, blanks):
            continue
        cell = f"{sp.name}!{get_column_letter(entry.value_col)}{entry.row} ({entry.label!r})"
        found += [(name.strip(), cell) for name in re.split(r"[;\n,]+", value) if name.strip()]
    return found


_FILE_DATE_TOKEN = re.compile(r"^(?:[yc]{2,4}|mm|dd|[yc]{4}mm(?:dd)?|mmdd[yc]{4}|mmdd[yc]{2}|"
                         r"\d{4,8}|\*+)$", re.IGNORECASE)


def sheet_feed_label(entry: dict) -> tuple[str, str]:
    """A per-SHEET feed name for a workbook whose mapping sheets would
    otherwise share one (2026-10-13, the D1 pair_11 'duplicate feed_id'):
    the sheet's own file-name cell (extension and trailing date / wildcard
    tokens dropped) when it states exactly one file, else the sheet name.
    Returns (label, where it came from)."""
    from codegen.resolve.gapfill import blank_sheet_feed_name

    own = entry.get("own_files") or []
    if len(own) == 1:
        name, cell = own[0]
        stem = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", name.strip())
        tokens = [t for t in re.split(r"[_\-.\s]+", stem) if t]
        while len(tokens) > 1 and _FILE_DATE_TOKEN.match(tokens[-1]):
            tokens.pop()
        if tokens:
            return "_".join(tokens), f"the sheet's own file cell {cell}: {name!r}"
    return blank_sheet_feed_name(entry["sheet"]), f"the sheet name {entry['sheet']!r}"


def _sttm_facts(profile: LayoutProfile, workbook, config: Config) -> dict:
    """What the STTM says about the feed, each with a citation."""
    facts: dict = {"texts": [], "schemas": {}, "meta": {}, "files": [], "sheet_tables": [],
                   "blank_table_sheets": [],
                   "file_rows": []}
    for ws in workbook.worksheets:
        region = sheet_region(ws)
        facts["texts"] += [(f"{ws.title}!{coord}", value) for coord, value in region.cells]
    open_band_sheets = {u.sheet for u in profile.unresolved
                        if BAND_REF_RE.match(str(u.layer)) and u.role == BAND_LAYER_ROLE}
    for sp in profile.mapping_sheets:
        # Per mapping sheet: the dominant stage / standard table and stage schema.
        # A sheet with a band whose layer is open says nothing yet (Chunk A):
        # its band question comes first, never a table question built on it.
        if sp.name in open_band_sheets:
            continue
        ws = workbook[sp.name]
        rows = list(ws.iter_rows(min_row=(sp.header_row or 1) + 1, values_only=True))
        standard_band = sp.band("standard")  # type: ignore[arg-type]
        entry = {"sheet": sp.name, "stage_table": None, "standard_table": None,
                 "stage_schema": None, "standard_schema": None,
                 "stage_tables": [], "standard_tables": [],
                 "standard_band": standard_band is not None
                 and standard_band.column(Role.TABLE) is not None,
                 # the TableName column is PLACED (blank cells), not unresolved
                 "stage_table_placed": sp.band("stage") is not None  # type: ignore[arg-type]
                 and sp.band("stage").column(Role.TABLE) is not None,  # type: ignore[arg-type]
                 # 2026-10-13: the file THIS sheet maps, from its own meta row
                 # (family E: one mapping sheet per source file) — (name, cell).
                 "own_files": _own_sheet_files(sp, ws, config),
                 # post156: row by row, which stage table goes with which standard /
                 # source table (the per-table split pairs by these, never by names)
                 "table_rows": [],
                 "segmented": sp.segment_strategy != "none"}
        for layer, role, key in (("stage", Role.TABLE, "stage_table"),
                                 ("standard", Role.TABLE, "standard_table"),
                                 ("stage", Role.SCHEMA, "stage_schema"),
                                 ("standard", Role.SCHEMA, "standard_schema")):
            band = sp.band(layer)  # type: ignore[arg-type]
            col = band.column(role) if band else None
            if col is not None:
                values = [text(r[col - 1]) if col - 1 < len(r) else None for r in rows]
                entry[key] = _dominant(values)
                if role == Role.TABLE:
                    # Chunk B: EVERY table the band names (a segmented sheet
                    # maps a table per record type), in sheet order.
                    entry[f"{layer}_tables"] = list(dict.fromkeys(v for v in values if v))
        def _col(layer: str, role: Role, sp=sp) -> int | None:
            band = sp.band(layer)  # type: ignore[arg-type]
            return band.column(role) if band is not None else None

        stage_col, standard_col = _col("stage", Role.TABLE), _col("standard", Role.TABLE)
        source_col = _col("source", Role.SOURCE_TABLE)
        if stage_col is not None:
            def _cell_at(r, col):
                return text(r[col - 1]) if col is not None and col - 1 < len(r) else None

            entry["table_rows"] = [(_cell_at(r, stage_col), _cell_at(r, standard_col),
                                    _cell_at(r, source_col))
                                   for r in rows if _cell_at(r, stage_col)]
        if entry["stage_table"]:
            facts["sheet_tables"].append(entry)
        else:
            facts["blank_table_sheets"].append(entry)
    for sp in profile.mapping_sheets:
        ws = workbook[sp.name]
        header_row = sp.header_row or 1
        rows = list(ws.iter_rows(min_row=header_row + 1, values_only=True))
        for layer in ("stage", "standard"):
            band = sp.band(layer)  # type: ignore[arg-type]
            if band is None:
                continue
            for role in (Role.SCHEMA, Role.CATALOG):
                col = band.column(role)
                if col is None:
                    continue
                value = _dominant([text(r[col - 1]) if col - 1 < len(r) else None for r in rows])
                if value:
                    facts["schemas"][f"{layer}.{role.value}"] = (
                        value, f"{sp.name}!{get_column_letter(col)} ({role.value} column, "
                               f"dominant value)")
        blanks = config.extractor.discovery.meta_blank_values
        for entry in sp.meta_rows:
            if entry.key in ("file_format", "delimiter", "frequency", "file_names",
                             "target_table_desc", "load_strategy", "lob") \
                    and entry.value_col is not None:
                value = text(ws.cell(row=entry.row, column=entry.value_col).value)
                if value and not is_blank_placeholder(value, blanks):   # "TBD" states nothing
                    facts["meta"][entry.key] = (
                        value, f"{sp.name}!{get_column_letter(entry.value_col)}{entry.row} "
                               f"({entry.label!r})")
    for sp in profile.sheets:
        if sp.kind != "file_details" or sp.header_row is None:
            continue
        from codegen.extract.annotations import annotation_reason

        ws = workbook[sp.name]
        headers = [normalize(c) for c in next(ws.iter_rows(min_row=sp.header_row,
                                                            max_row=sp.header_row,
                                                            values_only=True))]
        freq_col = next((i for i, h in enumerate(headers) if "frequency" in h), None)
        name_col = next((i for i, h in enumerate(headers)
                         if "file" in h and "name" in h and "description" not in h), None)
        for row_index, styled in enumerate(ws.iter_rows(min_row=sp.header_row + 1),
                                           start=sp.header_row + 1):
            row = [c.value for c in styled]
            # Guidance rows ("DataType = …", notes, italic examples) are not
            # files (codegen.extract.annotations; the extractor reports them).
            if annotation_reason([text(v) for v in row], name_col, styled) is not None:
                continue
            if freq_col is not None and freq_col < len(row) and text(row[freq_col]):
                facts["meta"].setdefault("frequency", (
                    text(row[freq_col]),
                    f"{sp.name}!{get_column_letter(freq_col + 1)}{row_index}"))
            if name_col is not None and name_col < len(row) and text(row[name_col]):
                facts["files"].append((text(row[name_col]),
                                       f"{sp.name}!{get_column_letter(name_col + 1)}{row_index}"))
                # M7: the row's own frequency (the FRD's Frequency cell may
                # point at this tab) — resolved PER derived feed.
                frequency = (text(row[freq_col]) if freq_col is not None
                             and freq_col < len(row) else "")
                facts["file_rows"].append({
                    "name": text(row[name_col]),
                    "cell": f"{sp.name}!{get_column_letter(name_col + 1)}{row_index}",
                    "frequency": frequency or None,
                    "frequency_cell": (f"{sp.name}!{get_column_letter(freq_col + 1)}{row_index}"
                                       if freq_col is not None else None),
                })
    return facts


def _vdd_facts(vdd_profile: LayoutProfile | None, workbook) -> dict[str, list[tuple[str, str]]]:
    """The FILES sheet's format / delimiter / cadence values, read through
    the VDD profile's roles, each with its cell."""
    if vdd_profile is None or workbook is None:
        return {}
    out: dict[str, list[tuple[str, str]]] = {}
    for sp in vdd_profile.sheets:
        if sp.kind != "vdd_files" or sp.header_row is None:
            continue
        band = sp.band("files")  # type: ignore[arg-type]
        if band is None:
            continue
        ws = workbook[sp.name]
        wanted = {Role.FORMAT: "file_format", Role.DELIMITER: "delimiter",
                  Role.CADENCE: "frequency", Role.FILE_PATTERN: "file_patterns"}
        for row_index, row in enumerate(ws.iter_rows(min_row=sp.header_row + 1, values_only=True),
                                        start=sp.header_row + 1):
            for role, key in wanted.items():
                col = band.column(role)
                if col is not None and col - 1 < len(row) and text(row[col - 1]):
                    out.setdefault(key, []).append(
                        (text(row[col - 1]), f"{sp.name}!{get_column_letter(col)}{row_index}"))
    return out


def _sttm_tables(profile: LayoutProfile, workbook) -> list[str]:
    """Table names the STTM speaks of (stage tables and database source
    tables), for narrowing a one-sheet-per-table dictionary."""
    tables: list[str] = []
    for sp in profile.mapping_sheets:
        ws = workbook[sp.name]
        rows = list(ws.iter_rows(min_row=(sp.header_row or 1) + 1, values_only=True))
        for layer, role in (("stage", Role.TABLE), ("source", Role.SOURCE_TABLE)):
            band = sp.band(layer)  # type: ignore[arg-type]
            col = band.column(role) if band else None
            if col is None:
                continue
            value = _dominant([text(r[col - 1]) if col - 1 < len(r) else None for r in rows])
            if value:
                tables.append(value)
    return tables


def _frd_citation(contract: FrdContract, index: int, field_name: str) -> str:
    item = contract.field_provenance.get(f"feeds[{index}].{field_name}")
    if item is None:
        return f"contract feeds[{index}].{field_name}"
    return f"table {item.table} row {item.row} ({item.label!r})"


def cross_check(contract: FrdContract, sttm_profile: LayoutProfile, workbook, config: Config,
                vdd_profile: LayoutProfile | None = None, vdd_workbook=None) -> list[CrossCheck]:
    from codegen.resolve.gapfill import same_value

    facts = _sttm_facts(sttm_profile, workbook, config)
    vdd = _vdd_facts(vdd_profile, vdd_workbook)
    checks: list[CrossCheck] = []
    for index, feed in enumerate(contract.feeds):
        # 1. feed / object name against the STTM meta rows and file details.
        hits = [(cit, value) for cit, value in facts["texts"] if _stem_match(feed.feed_name, value)]
        meta_cits = ", ".join(cit for cit, _ in facts["meta"].values()) or "header region"
        named_after = next((e for e in facts["sheet_tables"]
                            if e["stage_table"] == feed.feed_name), None)
        if named_after is not None and not hits:
            # The FRD named no feed; the layout stage named it after this
            # sheet's stage band — nothing independent is left to compare.
            checks.append(CrossCheck(
                "feed_name", "unchecked", _frd_citation(contract, index, "feed_name"),
                f"{named_after['sheet']} stage band",
                f"{feed.feed_name!r} is the STTM's own stage table (the FRD names no feed)"))
        else:
            checks.append(CrossCheck(
                "feed_name", "agree" if hits else "disagree",
                _frd_citation(contract, index, "feed_name"),
                hits[0][0] if hits else meta_cits,
                f"{feed.feed_name!r} ↔ {hits[0][1]!r}" if hits else
                f"{feed.feed_name!r} appears in no STTM meta row / file-details cell"))
        # 2. target catalog / schema against the dominant target-band values.
        for layer, target in (("stage", feed.stage_target), ("standard", feed.standard_target)):
            for role, value in (("schema", target.schema_name), ("catalog", target.catalog)):
                sttm_value = facts["schemas"].get(f"{layer}.{role}")
                citation = _frd_citation(contract, index, f"{layer}_target.{role}")
                if value is None or sttm_value is None:
                    checks.append(CrossCheck(f"{layer}_{role}", "unchecked", citation,
                                             sttm_value[1] if sttm_value else "no value",
                                             "one side states none"))
                    continue
                agree = normalize(value) == normalize(sttm_value[0])
                checks.append(CrossCheck(
                    f"{layer}_{role}", "agree" if agree else "disagree", citation,
                    sttm_value[1], f"{value!r} ↔ {sttm_value[0]!r}"))
        # 3. format / delimiter / frequency against the meta rows and the VDD FILES sheet.
        for key, value in (("file_format", feed.file_format), ("delimiter", feed.delimiter),
                           ("frequency", feed.frequency)):
            candidates = []
            if key in facts["meta"]:
                candidates.append(facts["meta"][key])
            candidates += vdd.get(key, [])
            if value is None or not candidates:
                checks.append(CrossCheck(key, "unchecked", _frd_citation(contract, index, key),
                                         ", ".join(c for _, c in candidates) or "no value",
                                         "one side states none"))
                continue
            # post113 (codeflow audit #8): same_value, so two punctuation
            # delimiters agree only when they are the same character
            matches = [(v, cit) for v, cit in candidates
                       if same_value(v, value, config.extractor.delimiter_words)]
            checks.append(CrossCheck(
                key, "agree" if matches else "disagree", _frd_citation(contract, index, key),
                matches[0][1] if matches else ", ".join(c for _, c in candidates),
                f"{value!r} ↔ {matches[0][0]!r}" if matches else
                f"{value!r} vs {[v for v, _ in candidates]!r}"))
    return checks


def _apply_cross_checks(pair: PairResolution, checks: list[CrossCheck], config: Config) -> None:
    bonus = config.layout.crosscheck_bonus
    profile = pair.sttm.profile
    assert isinstance(profile, LayoutProfile)
    confidence = dict(profile.confidence)
    frd_conf = dict(pair.frd.profile.confidence) if pair.frd else {}
    for check in checks:
        if check.status == "agree":
            for sp in profile.mapping_sheets:
                for layer in ("stage", "standard"):
                    if check.name.startswith(layer):
                        role = check.name.split("_", 1)[1]
                        key = confidence_key(sp.name, layer, role)
                        if key in confidence:
                            confidence[key] = min(1.0, round(confidence[key] + bonus, 3))
                if check.name == "feed_name":
                    confidence[f"{sp.name}/meta/feed_name"] = min(
                        1.0, round(confidence.get(f"{sp.name}/meta/feed_name", 0.9) + bonus, 3))
            for key in list(frd_conf):
                if key.endswith(check.name) or (check.name == "feed_name"
                                                and key.endswith(".feed_name")):
                    frd_conf[key] = min(1.0, round(frd_conf[key] + bonus, 3))
        elif check.status == "disagree":
            pair.flags.append(f"layout_crosscheck:{check.name} — FRD {check.frd_citation} vs "
                              f"STTM {check.sttm_citation}: {check.detail}")
    pair.sttm.profile = profile.model_copy(update={"confidence": confidence})
    if pair.frd is not None:
        pair.frd.profile = pair.frd.profile.model_copy(update={"confidence": frd_conf})


# ------------------------------------------------------------------ pair


def resolve_pair(sttm_path: Path, frd_path: Path | None, config: Config, *,
                 vdd_path: Path | None = None, provider: LayoutModelProvider | None = None,
                 answers: dict | None = None, cache_dirs: list[Path] | None = None,
                 runtime_cache_dir: Path | None = None, base_dir: Path | None = None,
                 use_cache: bool = True, generated_date: str | None = None,
                 refresh: bool = False,
                 prior: PairResolution | None = None) -> PairResolution:
    """Resolve the pair; ``answers`` = ``{"sttm": {key: col}, "frd": {field: {…}}}``.
    ``refresh`` re-resolves every document past the caches and overwrites the
    runtime entries (the pair entry included). ``prior`` = this pair's
    resolution from an earlier pass of the same invocation: each document
    continues from it (no cache read, no second model call, true sources)."""
    answers = answers or {}
    base = base_dir if base_dir is not None else Path(".")
    caches = _cache_dirs(config, base, runtime_cache_dir, cache_dirs)
    frd_is_docx = frd_path is not None and Path(frd_path).suffix.lower() == ".docx"
    has_vdd = vdd_path is not None and Path(vdd_path).is_file()

    # a. pair cache first — a repeat pair costs zero model calls.
    pair_fp = None
    pair_hit = False
    sttm_doc = frd_doc = vdd_doc = None
    workbook = content = vdd_workbook = None
    if frd_path is not None:
        sttm_fp = fingerprint(load_document(sttm_path, config.extractor.used_range_empty_rows))
        if frd_is_docx:
            from codegen.extract.frd_docx import read_docx

            frd_fp = frd_fingerprint(read_docx(frd_path, prose_ok=True).tables)
        else:
            frd_fp = hashlib.sha256(Path(frd_path).read_bytes()).hexdigest()
        vdd_fp = (fingerprint(load_document(vdd_path, config.extractor.used_range_empty_rows))
                  if has_vdd else "")
        pair_fp = hashlib.sha256(f"{sttm_fp}:{frd_fp}:{vdd_fp}".encode()).hexdigest()
        cached = (_load_cached(pair_fp, caches, prefix="pair_")
                  if use_cache and not refresh and prior is None else None)
        if cached is not None and not answers:
            try:
                profile = _as_cache(LayoutProfile.model_validate(cached["sttm"]))
                if _stale_reason(profile) is not None:
                    raise KeyError("stale pair entry: a required role is missing")
                workbook = load_document(sttm_path, config.extractor.used_range_empty_rows)
                if stale_value_evidence(profile, workbook, config.extractor) is not None:
                    # Value-derived band layers are re-derived on every hit.
                    raise KeyError("stale pair entry: a band layer read from values")
                sttm_doc = DocumentResolution("sttm", profile, cache_hit=True, fingerprint=sttm_fp)
                if cached.get("frd") is not None:
                    frd_profile = FrdLayoutProfile.model_validate(cached["frd"])
                    frd_profile = frd_profile.model_copy(update={
                        "source": "cache",
                        "field_sources": {k: "cache" for k in frd_profile.fields}})
                    frd_doc = DocumentResolution("frd", frd_profile, cache_hit=True,
                                                 fingerprint=frd_fp)
                if cached.get("vdd") is not None and has_vdd:
                    vdd_workbook = load_document(vdd_path, config.extractor.used_range_empty_rows)
                    vdd_doc = DocumentResolution(
                        "vdd", _as_cache(LayoutProfile.model_validate(cached["vdd"])),
                        cache_hit=True, fingerprint=vdd_fp)
                pair_hit = True
            except (ValidationError, KeyError):
                sttm_doc = frd_doc = vdd_doc = None

    if sttm_doc is None:
        sttm_doc, workbook = resolve_workbook(
            sttm_path, config, provider=provider, cache_dirs=cache_dirs,
            runtime_cache_dir=runtime_cache_dir, answers=answers.get("sttm"), base_dir=base,
            use_cache=use_cache, refresh=refresh, prior=prior.sttm if prior else None)
    if frd_is_docx and frd_doc is None:
        frd_doc, content = resolve_frd(
            frd_path, config, provider=provider, cache_dirs=cache_dirs,
            runtime_cache_dir=runtime_cache_dir, answers=answers.get("frd"), base_dir=base,
            use_cache=use_cache, refresh=refresh, prior=prior.frd if prior else None)
    if has_vdd and vdd_doc is None:
        assert isinstance(sttm_doc.profile, LayoutProfile)
        vdd_doc, vdd_workbook = resolve_workbook(
            Path(vdd_path), config, provider=provider, cache_dirs=cache_dirs,
            runtime_cache_dir=runtime_cache_dir, answers=answers.get("vdd"), base_dir=base,
            use_cache=use_cache, document="vdd",
            sttm_tables=_sttm_tables(sttm_doc.profile, workbook), refresh=refresh,
            prior=prior.vdd if prior else None)

    frd_contract: FrdContract | None = None
    if frd_path is not None:
        if frd_is_docx:
            from codegen.extract.frd_docx import read_docx, read_frd

            if content is None:
                content = read_docx(frd_path, prose_ok=True)
            assert frd_doc is not None and isinstance(frd_doc.profile, FrdLayoutProfile)
            frd_contract = read_frd(content, frd_doc.profile, config,
                                    document_name=Path(frd_path).name,
                                    generated_date=generated_date or "1970-01-01")
        else:
            frd_contract = FrdContract.model_validate(
                json.loads(Path(frd_path).read_text(encoding="utf-8")))

    # FRD gaps: (1) one metadata block describing several files becomes one
    # feed per STTM mapping sheet (the stage band is authoritative for tables;
    # each feed's file is paired by a unique name match or asked); (2) fill
    # from the STTM meta rows / VDD FILES / FAQ, or ask when sources disagree
    # or the only source is ambiguous (resolve/gapfill.py).
    gap = None
    split_flags: list[str] = []
    split_questions: list[LayoutQuestion] = []
    if frd_contract is not None and isinstance(sttm_doc.profile, LayoutProfile):
        facts = _sttm_facts(sttm_doc.profile, workbook, config)
        remembered = dict(sttm_doc.profile.file_choices)
        frd_contract, split_flags, split_questions = split_frd_feeds_by_sttm(
            frd_contract, facts, answers.get("gaps") or {}, config,
            force=sttm_doc.profile.strategy == "content", remember=remembered)
        if remembered != sttm_doc.profile.file_choices:
            sttm_doc.profile = sttm_doc.profile.model_copy(update={"file_choices": remembered})
        # M7: refused cells (nested tables / per-file blocks) resolve per
        # derived feed by file name; the rest stay unstated + flagged.
        frd_contract, structured_flags = resolve_structured_fields(frd_contract, config)
        split_flags = [*split_flags, *structured_flags]
        gap = fill_frd_gaps(
            frd_contract, frd_doc, facts,
            _vdd_facts(vdd_doc.profile if vdd_doc is not None else None,  # type: ignore[arg-type]
                       vdd_workbook),
            answers.get("gaps") or {}, config, base,
            asked={q.key for q in split_questions})
        frd_contract = gap.contract
        if frd_is_docx:
            # M9.3: what the layout stage DID to the contract (named the feed,
            # split it, took a value from another document or the person)
            # travels with it, so a CLI `generate` on the written contract
            # raises the same flags the UI run does. The "authoritative, never
            # asked" notes stay out: the contract resolver reports those with
            # the value and the cell.
            carried = [f for f in [*split_flags, *gap.flags]
                       if "(authoritative; never asked)" not in f]
            frd_contract = frd_contract.model_copy(update={"extraction_flags": list(
                dict.fromkeys([*frd_contract.extraction_flags, *carried]))})
        if frd_doc is not None:
            frd_doc.questions = [q for q in frd_doc.questions
                                 if q.role not in gap.handled] + split_questions + gap.questions
    # M9.2: a positional field whose length is not an integer ("10,2" — a
    # precision), with no STTM end and no VDD span, needs its byte width from
    # the person: one text question per field, answered under `gaps`.
    from codegen.resolve.widths import parse_width_answers

    width_answers = parse_width_answers(answers.get("gaps"))
    if isinstance(sttm_doc.profile, LayoutProfile):
        sttm_doc.questions = [
            *sttm_doc.questions,
            *_width_questions(sttm_doc.profile, workbook, vdd_path if has_vdd else None,
                              vdd_doc, frd_contract, width_answers, config)]
    pair = PairResolution(sttm=sttm_doc, frd=frd_doc, frd_contract=frd_contract,
                          pair_fingerprint=pair_fp, pair_cache_hit=pair_hit, vdd=vdd_doc,
                          gap_fills=list(gap.fills) if gap else [],
                          width_answers=width_answers)
    pair.flags.extend(split_flags)
    if gap is not None:
        pair.flags.extend(gap.flags)
    handled = gap.handled if gap is not None else set()
    for doc in pair.documents:
        for item in doc.profile.unresolved:
            where = (f"{item.sheet}/{item.layer}/{item.role}" if isinstance(item, UnresolvedRole)
                     else item.field)
            if doc.document == "frd" and where in handled:
                continue  # filled from another document / never asked: flagged above
            pair.flags.append(f"layout_unresolved:{doc.document} {where} — read as empty "
                              f"({item.reason})")
        for rejection in doc.rejections:
            pair.flags.append(f"layout_rejected:{rejection.render()}")
    if frd_contract is not None and isinstance(sttm_doc.profile, LayoutProfile):
        pair.cross_checks = cross_check(
            frd_contract, sttm_doc.profile, workbook, config,
            vdd_profile=vdd_doc.profile if vdd_doc is not None else None,  # type: ignore[arg-type]
            vdd_workbook=vdd_workbook)
        _apply_cross_checks(pair, pair.cross_checks, config)
    if pair_fp and caches.runtime is not None and not pair_hit:
        if not pair.questions and (refresh or pair.provider_calls or answers
                                   or (prior is not None and not prior.pair_cache_hit)):
            _save_runtime({
                "pair_fingerprint": pair_fp,
                "vocabulary": caches.vocabulary,
                "sttm": pair.sttm.profile.model_dump(mode="json"),
                "frd": pair.frd.profile.model_dump(mode="json") if pair.frd else None,
                "vdd": pair.vdd.profile.model_dump(mode="json") if pair.vdd else None,
            }, caches.runtime, f"pair_{pair_fp}.json")
        elif refresh and pair.questions:
            _invalidate_runtime(caches, f"pair_{pair_fp}.json", "pair_fingerprint", pair_fp)
    return pair


def _width_questions(profile: LayoutProfile, workbook, vdd_path, vdd_doc, frd_contract,
                     width_answers: dict, config: Config) -> list[LayoutQuestion]:
    """The last link of the width chain (codegen.resolve.widths), asked: every
    positional field (integer start) whose STTM length is not an integer, whose
    STTM end is missing, and for which the VDD states no span either."""
    from codegen.extract.generic import read_workbook
    from codegen.resolve.widths import as_integer, normalize_field_name, sttm_width, width_key

    if profile.strategy != "content" or any(
            u.role in ("field_name", "start") for u in profile.unresolved):
        return []
    try:
        ir = read_workbook(discovery_for(profile, workbook), "", config)
    except Exception:  # noqa: BLE001 — an unreadable sheet is the extractor's error to raise
        return []
    vdd_fields: list = []
    if vdd_path is not None and vdd_doc is not None and not vdd_doc.profile.unresolved:
        try:
            from codegen.extract.vdd import extract_vdd_contract

            vdd_fields = list(extract_vdd_contract(
                Path(vdd_path), config, layout=vdd_doc.profile,  # type: ignore[arg-type]
                generated_date="1970-01-01")[0].fields)
        except Exception:  # noqa: BLE001 — no VDD facts: the question is asked instead
            vdd_fields = []
    from codegen.extract.generic import _skip_marker, _unmapped_markers

    markers = _unmapped_markers(config)
    feeds = len(frd_contract.feeds) if frd_contract is not None else 1
    questions: list[LayoutQuestion] = []
    for sheet_index, sheet in enumerate(ir.sheets):
        feed_index = sheet_index if feeds == len(ir.sheets) else 0
        segmented = any(r.segment for r in sheet.fields) and any(f.segment for f in vdd_fields)
        source = sheet.profile.band("source")
        for row in sheet.fields:
            name = row.values.get("source.field_name") or ""
            length = row.values.get("source.length") or row.values.get("source.field_length")
            start, end = row.values.get("source.start"), row.values.get("source.end")
            if _skip_marker(row, markers) is not None:
                continue                      # "Do Not Map": in no table, needs no width
            if as_integer(start) is None or sttm_width(length, start, end)[0] is not None:
                continue
            key = normalize_field_name(name)
            if any(normalize_field_name(v.name) == key
                   and (not segmented or (v.segment_canonical or v.segment) == row.segment)
                   and ((v.start is not None and v.end is not None) or v.length)
                   for v in vdd_fields):
                continue                      # the VDD states the span: the resolver takes it
            if key in width_answers.get(feed_index, {}):
                continue
            col = source.column("length") or source.column("field_length") if source else None
            cell = (f"{sheet.profile.name}!{get_column_letter(col)}{row.row}" if col
                    else f"{sheet.profile.name} row {row.row}")
            label = " ".join(name.split())
            questions.append(LayoutQuestion(
                document="sttm", sheet=sheet.profile.name, layer="source",
                role=width_key(feed_index, name), kind="text",
                reason=f"{cell} reads {length!r} — not an integer byte width (a precision is "
                       "never a width); the STTM states no end for the field and the VDD no "
                       "start / end",
                title=f"Byte width of {label}",
                hint="How many bytes the field occupies in the fixed-width record (an "
                     "integer), as the vendor's layout states it."))
    return questions


@dataclass
class GapFillResult:
    contract: FrdContract
    fills: list[dict] = field(default_factory=list)
    questions: list[LayoutQuestion] = field(default_factory=list)
    flags: list[str] = field(default_factory=list)
    handled: set[str] = field(default_factory=set)   # FRD question keys not to ask


_GAP_FIELDS = ("file_format", "delimiter", "frequency", "stage_target.load_strategy",
               "standard_target.load_strategy")
_STTM_AUTHORITATIVE = ("stage_target.schema", "stage_target.tables")
# Round 7: the file-name extensions that state a format (_FeedGapFiller._file_format).
_FORMAT_EXTENSIONS = ("csv", "psv", "tsv")
# A generic "text file" format cell: these extensions decide (2026-10-13).
_TEXT_EXTENSIONS = ("csv", "psv", "tsv", "txt")


def _cell_row(cell: str) -> tuple[str, str]:
    """``"FILES!C4"`` -> ``("FILES", "4")``: the sheet row a VDD fact sits on."""
    sheet, _sep, ref = cell.rpartition("!")
    return sheet, ref.lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ")
# Chunk B (cold drill): further FRD fields a person may answer under gaps:
# when the FRD states them nowhere the reader can place (text values; LOBs a
# list, split on , ; and line breaks).
_ANSWERED_TEXT_FIELDS = ("domain", "sub_domain", "landing_location", "source_system")
_ANSWERED_LIST_FIELDS = ("lobs",)
_TARGET_FIELDS = ("stage_target.schema", "stage_target.tables", "standard_target.schema",
                  "standard_target.tables",
                  # Chunk B: a catalog no document states (and no conventions
                  # default supplies) is the person's to state.
                  "stage_target.catalog", "standard_target.catalog")
# Chunk B review (B2): every FRD field a gaps: VALUE sets - the branches of
# _FeedGapFiller.run, the split step's file choice (file_name_patterns) and
# the file-pattern text answer. The feed NAME is deliberately not one: the
# FRD's Name row names a feed, else split_frd_feeds_by_sttm names it after
# the STTM stage band BEFORE any answer is read - the name is the join key
# the extractor pairs each sheet by (stage table, file pattern, the derived
# name of a blank-TableName sheet), so a stated name could not be honoured
# without breaking that pairing once the FRD splits into several feeds.
_ANSWERABLE_FRD_FIELDS = frozenset((*_GAP_FIELDS, *_ANSWERED_TEXT_FIELDS,
                                    *_ANSWERED_LIST_FIELDS, *_TARGET_FIELDS, "load_strategy",
                                    "file_name_patterns", "file_patterns"))
_UNANSWERABLE_FRD_REASONS = {
    "feed_name": "the FRD's Name row names the feed, else the layout stage names it after the "
                 "STTM stage band - the key each STTM sheet pairs by; an answer cannot rename "
                 "a feed",
}


def answerable_frd_field(key: str) -> bool:
    """True when a ``gaps:`` value for ``feeds[i].<field>`` reaches the
    pair-resolved FRD contract (``_ANSWERABLE_FRD_FIELDS``)."""
    match = re.match(r"^feeds\[\d+\]\.(.+)$", key)
    return match is not None and match.group(1) in _ANSWERABLE_FRD_FIELDS


def unanswerable_frd_reason(key: str) -> str:
    """Why a ``gaps:`` value for this FRD field is not applied."""
    field = key.split("].", 1)[1] if "]." in key else key
    return _UNANSWERABLE_FRD_REASONS.get(
        field, "the layout stage takes no answer for this field")


def _feed_get(feed, dotted: str):
    obj = feed
    for part in dotted.split("."):
        obj = getattr(obj, part if part != "schema" else "schema_name", None)
    return obj


def _feed_set(feed, dotted: str, value):
    head, _sep, tail = dotted.partition(".")
    if not tail:
        return feed.model_copy(update={head: value})
    inner = getattr(feed, head)
    return feed.model_copy(update={head: _feed_set(inner, tail, value)})


_FILE_LIKE = re.compile(r"(\.[a-z0-9]{2,4}$|[*?]|yyyy|ccyy|mmdd)", re.IGNORECASE)
_NOISE_TOKENS = {"sd", "data", "package", "file", "files", "report", "the", "and", "of", "stg",
                 "std", "src", "tgt", "raw", "yyyy", "yyyymmdd", "mm", "dd", "hhmm", "csv",
                 "psv", "txt", "dat", "to", "from"}


_DATE_TOKEN = re.compile(r"^[ymdhc]+$")


def _name_tokens(name: str) -> set[str]:
    return {t for t in re.split(r"[^a-z0-9]+", name.lower()) if len(t) >= 3
            and t not in _NOISE_TOKENS and not _DATE_TOKEN.match(t)}


def _tokens_overlap(a: set[str], b: set[str]) -> bool:
    return any(x == y or x.startswith(y) or y.startswith(x) for x in a for y in b)


def _mutual_unique_matches(tables: list[str], files: list[str]) -> dict[str, str]:
    """table -> file when the two name each other uniquely: the file is the
    only one whose tokens overlap the table's AND the table is the only one
    whose tokens overlap the file's. Tokens every table shares (a family
    suffix such as ``risk``) carry no information and are dropped first."""
    table_tokens = {t: _name_tokens(t) for t in tables}
    if len(tables) > 1:
        common = set.intersection(*table_tokens.values()) if table_tokens else set()
        table_tokens = {t: toks - common for t, toks in table_tokens.items()}
    file_tokens = {f: _name_tokens(f) for f in files}
    forward = {t: [f for f in files if _tokens_overlap(toks, file_tokens[f])]
               for t, toks in table_tokens.items()}
    reverse = {f: [t for t in tables if _tokens_overlap(table_tokens[t], toks)]
               for f, toks in file_tokens.items()}
    return {t: fs[0] for t, fs in forward.items()
            if len(fs) == 1 and reverse[fs[0]] == [t]}


def file_identity(name: str) -> str:
    """One file however the documents spell it (2026-10-13, D1 pair_11): the
    FRD's "X_DNC_Report_MMDDYYYY.csv" and File Details' "X_DNC_Report_*.csv"
    are the same file — case, separators, folder, extension and date /
    wildcard tokens dropped."""
    stem = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", PurePosixPath(name.strip()).name)
    tokens = [t for t in re.split(r"[^a-z0-9*]+", stem.lower()) if t]
    kept = [t for t in tokens if not _FILE_DATE_TOKEN.match(t) and not _DATE_TOKEN.match(t)]
    return "_".join(kept) or normalize(name)


def _same_word(a: str, b: str) -> bool:
    """Equal, or one the other's prefix ("repor" = "report": a sheet title
    cut at 31 characters)."""
    return a == b or (min(len(a), len(b)) >= 4 and (a.startswith(b) or b.startswith(a)))


def _match_score(sheet_tokens: set[str], file_tokens: set[str]) -> float:
    """Jaccard overlap of two token sets (``_same_word`` equality)."""
    if not sheet_tokens or not file_tokens:
        return 0.0
    shared = sum(1 for s in sheet_tokens if any(_same_word(s, f) for f in file_tokens))
    return shared / (len(sheet_tokens) + len(file_tokens) - shared)


def best_file_matches(sheets: dict[str, set[str]], files: list[str], *,
                      shared_sheet: set[str] | None = None, shared_file: set[str] | None = None
                      ) -> tuple[dict[str, tuple[str, float, list[str]]],
                                 dict[str, list[tuple[str, float]]]]:
    """2026-10-13 (D1 pair_11, five file questions): each sheet against the
    files by token overlap — its title, its File Name row, its feed label.
    Tokens every file shares (the vendor prefix) and tokens every sheet shares
    carry nothing and are dropped (``shared_sheet`` / ``shared_file``: computed
    over ALL the workbook's sheets / files when the caller passes the open
    ones only — a family word like "risk" is no evidence). A sheet whose best score is unique and
    above zero takes that file, unless another sheet scores the same file as
    high or higher (then it is asked). Returns (sheet -> (file, score, the
    words that matched)) and (sheet -> its tied best (file, score)) for the
    sheets left to ask."""
    file_tokens = {f: _name_tokens(f) for f in files}
    if shared_file is None:
        shared_file = set.intersection(*file_tokens.values()) if len(files) > 1 else set()
    file_tokens = {f: toks - shared_file for f, toks in file_tokens.items()}
    if shared_sheet is None:
        shared_sheet = set.intersection(*sheets.values()) if len(sheets) > 1 else set()
    sheets = {s: toks - shared_sheet for s, toks in sheets.items()}
    picks: dict[str, tuple[str, float]] = {}
    ties: dict[str, list[tuple[str, float]]] = {}
    for sheet, toks in sheets.items():
        ranked = sorted(((f, _match_score(toks, file_tokens[f])) for f in files),
                        key=lambda item: -item[1])
        top = ranked[0][1] if ranked else 0.0
        best = [(f, s) for f, s in ranked if s == top and s > 0]
        if len(best) == 1:
            picks[sheet] = best[0]
        else:
            ties[sheet] = best
    chosen: dict[str, tuple[str, float, list[str]]] = {}
    for sheet, (name, score) in picks.items():
        if any(s >= score for other, (f, s) in picks.items() if other != sheet and f == name):
            ties[sheet] = [(name, score)]
            continue
        words = sorted(t for t in sheets[sheet]
                       if any(_same_word(t, f) for f in file_tokens[name]))
        chosen[sheet] = (name, score, words)
    return chosen, ties


TARGET_FROM_FRD = "sttm_target_from_frd"


def _frd_cell(contract: FrdContract, dotted: str) -> str:
    # The standard schema is read from the same 'Target Schema' cell as the stage one.
    item = (contract.field_provenance.get(f"feeds[0].{dotted}")
            or contract.field_provenance.get(f"feeds[0].{dotted.replace('standard', 'stage')}"))
    return (f"FRD table {item.table} row {item.row} ({item.label!r})" if item
            else "FRD contract")


def _blank_sheet_targets(index: int, entry: dict, feed, files: list[str],
                         contract: FrdContract, gaps: dict
                         ) -> tuple[list[str], list[LayoutQuestion]]:
    """A split feed whose sheet leaves every Schema / TableName cell blank
    (round 4, 2026-10-12). Schema: the FRD's ONE stated schema per layer
    applies to every feed — a NOTE citing the FRD field; anything else (no
    schema, a per-file-group list) is a QUESTION. Table: no conventions rule
    derives a table name from a file name, so it is never invented — one
    QUESTION per layer, the FRD's file name quoted as the hint."""
    flags: list[str] = []
    questions: list[LayoutQuestion] = []
    layers = ("stage", "standard") if entry.get("standard_band", True) else ("stage",)
    for layer in layers:
        target = getattr(feed, f"{layer}_target")
        key = f"feeds[{index}].{layer}_target.schema"
        if entry.get(f"{layer}_schema") is None and key not in gaps:
            if target.schema_name:
                flags.append(f"{TARGET_FROM_FRD}:{key} source_used:"
                             f"{_frd_cell(contract, f'{layer}_target.schema')}: "
                             f"{target.schema_name!r} — sheet {entry['sheet']!r} leaves every "
                             f"{layer} Schema cell blank; the FRD states one {layer} schema "
                             "for every feed")
            else:
                refused = contract.structured.get(f"feeds[0].{layer}_target.schema") or \
                    contract.structured.get("feeds[0].stage_target.schema")
                why = (f"the FRD's {refused.label!r} cell is a {refused.kind} value "
                       f"({len(refused.text.splitlines())} line(s)), not one schema"
                       if refused is not None else f"the FRD states no {layer} schema")
                questions.append(LayoutQuestion(
                    document="frd", sheet=None, layer=None, role=key,
                    reason=f"sheet {entry['sheet']!r} leaves every {layer} Schema cell blank "
                           f"and {why}", header=[], candidates=[],
                    title=f"{layer.capitalize()} schema of {entry['sheet']}",
                    hint=f"The {layer} schema this sheet's table lands in.", kind="text"))
        key = f"feeds[{index}].{layer}_target.tables"
        if key in gaps or entry.get(f"{layer}_table"):
            continue          # answered, or the sheet's own band states the table
        quoted = ", ".join(repr(f) for f in files) or "no file"
        frd_says = (f"the FRD's 'Target Table Name' names files ({quoted}), not tables" if files
                    else "the FRD names no table for this sheet")
        questions.append(LayoutQuestion(
            document="frd", sheet=None, layer=None, role=key,
            reason=f"sheet {entry['sheet']!r} leaves every {layer} TableName cell blank; "
                   f"{frd_says}, and no conventions rule derives a table name from a file name",
            header=[], candidates=[],
            title=f"{layer.capitalize()} table of {entry['sheet']}",
            hint=f"The {layer} table the file {quoted} (FRD 'Target Table Name') loads into.",
            kind="text"))
    return flags, questions


def _file_kind(cell: str) -> str:
    """A file candidate's source, as ``file_tie_source_priority`` names it."""
    if cell == "FRD file pattern":
        return "frd_pattern"
    if cell.startswith("FRD 'Target Table Name'"):
        return "frd_target_table"
    return "sttm_file_details"


_EXTENSION_RE = re.compile(r"\.[A-Za-z0-9]{2,5}$")


def _shaped(name: str, vendor: set[str]) -> bool:
    """A vendor prefix (the name's first word is a word of the feed's source
    system) + an extension."""
    words = [w for w in re.split(r"[^A-Za-z0-9]+", name) if w]
    return bool(words) and words[0].lower() in vendor and bool(
        _EXTENSION_RE.search(name.rstrip("*")))


def break_file_tie(tied: list[str], kinds: dict[str, set[str]], profile,
                   vendor: set[str]) -> tuple[str, str] | None:
    """post59: ``tied`` files (equal name overlap with a table) — the one the
    conventions profile prefers: the best source (``file_tie_source_priority``:
    the FRD's landed pattern over an STTM FILE_DETAILS label …), then the shape
    (``file_tie_shape``: a vendor prefix + an extension). (file, why) or None
    while still tied."""
    order = list(profile.file_tie_source_priority)
    left = list(tied)
    if order:
        rank = {n: min((order.index(k) for k in kinds.get(n, set()) if k in order),
                       default=len(order)) for n in left}
        best = min(rank.values())
        left = [n for n in left if rank[n] == best]
        if len(left) == 1:
            sources = "/".join(sorted(kinds.get(left[0], set()) & set(order)))
            return left[0], f"source priority: {sources} first ({' > '.join(order)})"
    if profile.file_tie_shape:
        shaped = [n for n in left if _shaped(n, vendor)]
        if len(shaped) == 1:
            return shaped[0], "shape: a vendor prefix + an extension"
    return None


def split_frd_feeds_by_sttm(contract: FrdContract, facts: dict, gaps: dict, config: Config,
                            *, force: bool = False, remember: dict | None = None
                            ) -> tuple[FrdContract, list[str], list[LayoutQuestion]]:
    """A single FRD metadata block that describes several files (its
    'Target Table Name' row lists FILE names, none of them an STTM stage
    table) becomes one FRD feed per STTM mapping sheet: feed name and stage
    / standard tables from the sheet's target bands (authoritative), every
    other fact shared. Each feed's file pattern is the FILE_DETAILS / FRD
    file name whose tokens uniquely match the table's; otherwise a 'choice'
    question asks which file feeds the table. Flagged, never silent.

    ``force`` (2026-10-13, D1 pair_11): a content-strategy STTM with SEVERAL
    mapping sheets and ONE FRD feed always splits — one feed per sheet — even
    when the FRD's tables match the sheets in a way the rules above leave
    alone: unsplit, every sheet pairs with the one feed and extract-sttm
    stops. Feeds are named as above (stage table; per sheet when names
    collide or the FRD names no feed — the sheet's File Name row, else the
    sheet name)."""
    from codegen.extract.frd_docx import is_unnamed_feed
    from codegen.layout.hints import frd_field_help

    sheet_tables = facts.get("sheet_tables") or []
    per_table = _split_per_table(contract, sheet_tables)
    if per_table is not None:
        return per_table
    # Round 4 (2026-10-12): a copy of the STTM whose mapping sheets leave EVERY
    # Schema / TableName cell blank still splits — the mapping sheets alone say
    # how many feeds there are (one per sheet, named after it).
    # An UNRESOLVED TableName column (cold pairs) is a layout question, not blank cells.
    blank = facts.get("blank_table_sheets") or []
    blank_only = (not sheet_tables and len(blank) > 1
                  and all(e.get("stage_table_placed") for e in blank))
    # An UNRESOLVED TableName column stays the layout question it is (asked
    # first); only sheets whose table column is placed count for the split.
    forced = force and len(sheet_tables) + sum(
        1 for e in blank if e.get("stage_table_placed")) > 1
    if len(contract.feeds) != 1 or not (sheet_tables or blank_only or forced):
        return contract, [], []
    feed = contract.feeds[0]
    refused = contract.structured.get("feeds[0].feed_name")
    if refused is not None and refused.feed_name:
        refused = None      # M9.1b: the block itself names the feed (a labelled line / row)
    garbled_name = ("\n" in feed.feed_name or len(feed.feed_name) > 80
                    or refused is not None or is_unnamed_feed(feed.feed_name))
    name_reason = (f"the FRD's Object Name cell is a {refused.kind} value, not a name"
                   if refused is not None else
                   "the FRD names no feed" if is_unnamed_feed(feed.feed_name) else
                   "the FRD's feed name is not a name (line breaks / over 80 characters)")
    stated = {normalize(t) for t in feed.stage_target.tables}
    per_sheet = False
    if any(normalize(e["stage_table"]) in stated for e in sheet_tables):
        # The FRD names the STTM's tables: nothing to split. A single sheet
        # still NAMES a feed the FRD left unnamed (M9.3) — flagged.
        if len(sheet_tables) == 1 and garbled_name:
            entry = sheet_tables[0]
            renamed = feed.model_copy(update={"feed_name": entry["stage_table"]})
            return contract.model_copy(update={"feeds": [renamed]}), [
                f"frd_unstated:feeds[0].feed_name source_used:STTM stage band "
                f"{entry['sheet']!r}: {entry['stage_table']!r} ({name_reason})"], []
        # Chunk B (cold drill): ONE FRD feed names several tables and the STTM
        # maps each on its own sheet (family B: a MAPPING- sheet per table).
        # Every sheet is its own feed - one feed carrying two sheets has no
        # single stage table and reached the contract as a duplicate feed id.
        # 2026-10-13 (D1 pair_11): family E maps each SOURCE FILE on its own
        # sheet, all into the table the FRD names — still one feed per sheet
        # (named per sheet below); one feed carrying five sheets reached the
        # contract as a duplicate feed id.
        sheet_keys = [normalize(e["stage_table"]) for e in sheet_tables]
        shared_table = len(set(sheet_keys)) < len(sheet_keys)
        per_sheet = len(sheet_tables) > 1 and (
            set(sheet_keys) <= stated if shared_table
            else len(set(sheet_keys)) == len(sheet_keys) and set(sheet_keys) == stated)
        if not per_sheet and not forced:
            return contract, [], []
        per_sheet = per_sheet or forced
    file_like = [t for t in feed.stage_target.tables if _FILE_LIKE.search(t)]
    pool = [(name, f"FRD 'Target Table Name' ({name!r})") for name in file_like]
    pool += [(name, cell) for name, cell in facts.get("files", [])]
    if "file_names" in facts.get("meta", {}):
        value, cell = facts["meta"]["file_names"]
        pool += [(n.strip(), cell) for n in re.split(r"[;\n,]+", value) if n.strip()]
    pool += [(p, "FRD file pattern") for p in feed.file_name_patterns]
    seen: set[str] = set()
    candidates = []
    # post59: every source a file appears in (the tie-break's source priority)
    kinds_by_id: dict[str, set[str]] = {}
    meta_cell = facts.get("meta", {}).get("file_names", (None, None))[1]
    for name, cell in pool:
        key = file_identity(name) if normalize(name) else ""
        if key:
            kinds_by_id.setdefault(key, set()).add(
                "sttm_meta" if cell == meta_cell else _file_kind(cell))
    for name, cell in pool:
        # One candidate per FILE (2026-10-13): the FRD's spelling (first in the
        # pool) and File Details' wildcard one made every match ambiguous.
        key = file_identity(name) if normalize(name) else ""
        if key and key not in seen:
            seen.add(key)
            candidates.append((name, cell))
    rename = len(sheet_tables) > 1 or garbled_name
    # First ACFC run: in a split, a sheet whose TableName cells are ALL blank
    # still gets its feed (named after the sheet, no table) so the extractor
    # can hold it back as NEEDS_ANSWERS with a precise key — never dropped.
    from codegen.resolve.gapfill import blank_sheet_feed_name

    blank_sheets = facts.get("blank_table_sheets") or []
    if len(sheet_tables) + len(blank_sheets) > 1 and blank_sheets:
        sheet_tables = [*sheet_tables, *blank_sheets]
        rename = True
    if per_sheet:
        flags = [f"frd_feeds_split_from_sttm: FRD feed {feed.feed_name[:60]!r} names "
                 f"{feed.stage_target.tables} as target tables and the STTM maps each on its "
                 f"own sheet; {len(sheet_tables)} feeds derived, one per mapping sheet "
                 f"{[(e['sheet'], e['stage_table']) for e in sheet_tables]} "
                 "(stage band authoritative)"]
    elif len(sheet_tables) > 1:
        flags = [f"frd_feeds_split_from_sttm: FRD feed {feed.feed_name[:60]!r} names "
                 f"{feed.stage_target.tables} as target tables — none is an STTM stage table; "
                 f"{len(sheet_tables)} feeds derived from the STTM mapping sheets "
                 f"{[(e['sheet'], e['stage_table']) for e in sheet_tables]} "
                 "(stage band authoritative)"]
    else:
        named = sheet_tables[0].get("stage_tables") or [sheet_tables[0]["stage_table"]]
        flags = [f"frd_unstated:feeds[0].stage_target.tables source_used:STTM stage band "
                 f"{sheet_tables[0]['sheet']!r}: "
                 f"{named[0] if len(named) == 1 else named!r} (the FRD names "
                 f"{feed.stage_target.tables or 'no table'})"]
        if garbled_name:
            flags.append(f"frd_unstated:feeds[0].feed_name source_used:STTM stage band "
                         f"{sheet_tables[0]['sheet']!r}: {sheet_tables[0]['stage_table']!r} "
                         f"({name_reason})")
    questions: list[LayoutQuestion] = []
    feeds = []
    # A feed is named after its sheet's stage table (else the sheet). Sheets
    # that would SHARE a name (one table, a file per sheet) are named per
    # sheet instead — their own file cell, else the sheet name — so every
    # feed id is unique; a workbook whose names are already unique keeps them.
    labels: dict[str, str] = {}
    own_file_sheets: set[str] = set()
    bases = {e["sheet"]: e["stage_table"] or blank_sheet_feed_name(e["sheet"])
             for e in sheet_tables}
    counts: dict[str, int] = {}
    for base in bases.values():
        counts[normalize(base)] = counts.get(normalize(base), 0) + 1
    renamed: list[str] = []
    for e in sheet_tables:
        base = bases[e["sheet"]]
        if counts[normalize(base)] == 1 or len(sheet_tables) == 1:
            labels[e["sheet"]] = base
            continue
        name, source = sheet_feed_label(e)
        labels[e["sheet"]] = name
        if source.startswith("the sheet's own file"):
            own_file_sheets.add(e["sheet"])
        renamed.append(f"{e['sheet']!r} -> {name!r} ({source})")
    taken: dict[str, str] = {}
    for e in sheet_tables:              # still a clash: the sheet name, then an ordinal
        key = normalize(labels[e["sheet"]])
        if key in taken:
            labels[e["sheet"]] = blank_sheet_feed_name(e["sheet"])
            own_file_sheets.discard(e["sheet"])
            key = normalize(labels[e["sheet"]])
            if key in taken:
                labels[e["sheet"]] = f"{labels[e['sheet']]}_{len(taken) + 1}"
                key = normalize(labels[e["sheet"]])
        taken[key] = e["sheet"]
    if renamed:
        shared_names = sorted({base for base in bases.values() if counts[normalize(base)] > 1})
        flags.append("feed_names_per_sheet: the mapping sheets would share the feed name(s) "
                     f"{shared_names}; each sheet is its own feed: {'; '.join(renamed)}")

    def label(entry: dict) -> str:
        return labels[entry["sheet"]]

    matched = _mutual_unique_matches([label(e) for e in sheet_tables],
                                     [n for n, _c in candidates])
    # 2026-10-13 (D1 pair_11): the tables left unmatched — each sheet's title,
    # File Name row and label against the files nobody matched; a unique best
    # overlap is taken (a NOTE), a tie is asked with the tied files named.
    scored: dict[str, tuple[str, float, list[str]]] = {}
    tied: dict[str, list[tuple[str, float]]] = {}
    open_entries = [e for e in sheet_tables if label(e) not in matched]
    open_files = [n for n, _c in candidates if n not in matched.values()]
    if len(sheet_tables) > 1 and open_entries and open_files:
        tokens = {e["sheet"]: (_name_tokens(e["sheet"]) | _name_tokens(label(e))
                               | {t for n, _c in e.get("own_files") or []
                                  for t in _name_tokens(n)})
                  for e in sheet_tables}
        every_file = [_name_tokens(n) for n, _c in candidates]
        scored, tied = best_file_matches(
            {e["sheet"]: tokens[e["sheet"]] for e in open_entries}, open_files,
            shared_sheet=set.intersection(*tokens.values()),
            shared_file=set.intersection(*every_file) if len(every_file) > 1 else set())
    # 2026-10-12 (owner decision, fewer questions — was: only the dialog's
    # suggestion): when the unique name matches leave exactly ONE table and
    # ONE file over (which also means the documents list exactly as many files
    # as there are tables), that file is the table's value BY ELIMINATION,
    # flagged `(by elimination ...)` and printed as a NOTE — never asked. Two
    # or more left over on either side stays a question (no suggestion).
    taken_files = {*matched.values(), *(f for f, _s, _w in scored.values())}
    leftover_files = [n for n, _c in candidates if n not in taken_files]
    leftover_tables = [label(e) for e in sheet_tables
                       if label(e) not in matched and e["sheet"] not in scored]
    eliminated = (leftover_files[0] if len(leftover_files) == 1 and len(leftover_tables) == 1
                  else None)
    for index, entry in enumerate(sheet_tables):
        table = label(entry)
        key = f"feeds[{index}].file_name_patterns"
        if entry["stage_table"] is None:
            flags.append(f"frd_unstated:feeds[{index}].stage_target.tables source_used:none — "
                         f"sheet {entry['sheet']!r} leaves every TableName cell blank and the "
                         f"FRD names no table for it; feed named {table!r} after the sheet")
        patterns: list[str] = []
        chosen = gaps.get(key)
        cached = (remember or {}).get(entry["sheet"]) if chosen is None else None
        if chosen is not None and remember is not None and len(sheet_tables) > 1:
            # post59: the BSA's answer is cached with the layout profile
            remember[entry["sheet"]] = chosen["value"]
        if cached is not None and any(n == cached for n, _c in candidates):
            patterns = [cached]
            flags.append(f"frd_unstated:{key} source_used:the BSA's earlier answer (cached in "
                         f"the layout profile): {cached!r}")
        elif chosen is None and entry["sheet"] in own_file_sheets:
            # The sheet states its own file: that file is this feed's.
            name, cell = entry["own_files"][0]
            patterns = [name]
            flags.append(f"frd_unstated:{key} source_used:{cell}: {name!r} (the sheet's own "
                         "file row)")
        elif len(sheet_tables) == 1:
            # One sheet = one feed: every file the documents name belongs to it.
            # File-like 'Target Table Name' entries are its patterns when the
            # FRD states none; the resolver's STTM fallback covers the rest.
            if file_like and not feed.file_name_patterns:
                patterns = list(file_like)
                flags.append(f"frd_unstated:{key} source_used:FRD 'Target Table Name' "
                             f"(file-like entries): {file_like}")
        elif chosen is not None:
            patterns = [chosen["value"]]
            flags.append(f"frd_unstated:{key} source_used:user chosen from {chosen['source']}: "
                         f"{chosen['value']!r}")
        elif len(candidates) == 1:
            patterns = [candidates[0][0]]
            flags.append(f"frd_unstated:{key} source_used:{candidates[0][1]}: "
                         f"{candidates[0][0]!r} (the only file named)")
        elif candidates:
            if table in matched:
                name = matched[table]
                cell = next(c for n, c in candidates if n == name)
                patterns = [name]
                flags.append(f"frd_unstated:{key} source_used:{cell}: "
                             f"{name!r} (unique name match with table {table!r})")
            elif entry["sheet"] in scored:
                name, score, words = scored[entry["sheet"]]
                cell = next(c for n, c in candidates if n == name)
                patterns = [name]
                flags.append(f"frd_unstated:{key} source_used:{cell}: {name!r} (best name "
                             f"match: sheet {entry['sheet']!r} / table {table!r} share "
                             f"{words} with the file, overlap {score:.2f}; no other file "
                             "scores as high)")
            elif eliminated is not None and table == leftover_tables[0]:
                cell = next(c for n, c in candidates if n == eliminated)
                patterns = [eliminated]
                flags.append(f"frd_unstated:{key} source_used:{cell}: {eliminated!r} (by "
                             f"elimination: table {table!r} matches no file by name; the "
                             f"documents list {len(candidates)} files for {len(sheet_tables)} "
                             "tables and every other file pairs with another table)")
            elif (len(tied.get(entry["sheet"], [])) > 1 and (broken := break_file_tie(
                    [n for n, _s in tied[entry["sheet"]]],
                    {n: kinds_by_id.get(file_identity(n), set()) for n, _c in candidates},
                    config.conventions.get(None),
                    {w.lower() for w in re.split(r"[^A-Za-z0-9]+", feed.source_system or "")
                     if len(w) > 1})) is not None):
                name, why = broken
                cell = next(c for n, c in candidates if n == name)
                patterns = [name]
                flags.append(f"frd_unstated:{key} source_used:{cell}: {name!r} (tie between "
                             f"{[n for n, _s in tied[entry['sheet']]]} broken by {why})")
            else:
                title, hint, _s = frd_field_help(key, [], config)
                best = [n for n, _s in tied.get(entry["sheet"], [])]
                why = (f"tie between {best} (equal name overlap)" if len(best) > 1 else
                       f"{best[0]!r} matches another sheet as well or better" if best else
                       "no file name shares a word with the sheet / table")
                ordered = ([(n, c) for n, c in candidates if n in best]
                           + [(n, c) for n, c in candidates if n not in best])
                questions.append(LayoutQuestion(
                    document="frd", sheet=None, layer=None, role=key,
                    reason=f"which file feeds stage table {table!r} (sheet {entry['sheet']!r})? "
                           f"{why}; candidates: {[n for n, _c in ordered]}", header=[],
                    candidates=[{"value": n, "source": "STTM FILE_DETAILS / FRD", "cell": c}
                                for n, c in ordered],
                    title=f"File for table {table}", hint=hint, kind="choice"))
        if entry["stage_table"] is None:
            target_flags, target_questions = _blank_sheet_targets(
                index, entry, feed, patterns or file_like, contract, gaps)
            flags += target_flags
            questions += target_questions
        standard_tables = (list(entry.get("standard_tables") or [entry["standard_table"]])
                           if entry["standard_table"] else list(
            feed.standard_target.tables if not file_like else []))
        standard_update: dict = {"tables": standard_tables}
        sheet_std_schema = entry.get("standard_schema")
        if sheet_std_schema and feed.standard_target.schema_name is None and standard_tables:
            standard_update["schema_name"] = sheet_std_schema
            flags.append(f"frd_unstated:feeds[{index}].standard_target.schema source_used:STTM "
                         f"standard band {entry['sheet']!r}: {sheet_std_schema!r}")
        stage_update: dict = {"tables": list(entry.get("stage_tables") or [entry["stage_table"]])
                              if entry["stage_table"] else []}
        sheet_schema = entry.get("stage_schema")
        frd_schema = feed.stage_target.schema_name
        if sheet_schema and (frd_schema or "").lower() != sheet_schema.lower():
            # One FRD block, several sheets: the sheet's stage band names the
            # schema THIS feed lands in (authoritative, never asked).
            stage_update["schema_name"] = sheet_schema   # field name, not the "schema" alias
            flags.append(f"frd_unstated:feeds[{index}].stage_target.schema source_used:STTM "
                         f"stage band {entry['sheet']!r}: {sheet_schema!r} (the FRD block says "
                         f"{frd_schema!r} for every file)")
        # A feed of a SPLIT whose file is still asked owns NO pattern: the FRD
        # block's list belongs to all of the feeds (2026-10-12 — once the FRD
        # reader lifted the block's files into it, the unpaired feed inherited
        # every one, and extract failed on "N FILE_DETAILS rows match" instead
        # of naming the question). One sheet = one feed keeps the whole list.
        shared = list(feed.file_name_patterns) if len(sheet_tables) == 1 else []
        feeds.append(feed.model_copy(update={
            "feed_name": table if rename else feed.feed_name,
            "file_name_patterns": patterns or shared,
            "stage_target": feed.stage_target.model_copy(update=stage_update),
            "standard_target": feed.standard_target.model_copy(update=standard_update),
        }))
    return contract.model_copy(update={"feeds": feeds}), flags, questions


def _split_per_table(contract: FrdContract, sheet_tables: list[dict]
                     ) -> tuple[FrdContract, list[str], list] | None:
    """post156 (pair_09 shape): ONE FRD feed names N stage tables and ONE STTM
    mapping sheet maps all N - its stage TableName column holds exactly the FRD's
    N tables - on a sheet that is NOT segmented (a header / detail / trailer
    sheet maps a table per record type: that stays one feed). Then each table is
    its own feed, every pairing from the documents' own rows or an exact name,
    never a guess:

    - the feed is named after its stage table (a document value);
    - its standard table = the one standard table its rows name, when exactly one
      and the FRD names it (or names none); else none + a flag;
    - its file pattern = the FRD pattern that IS the table (or its source table)
      by file identity; else none + a flag (a test case then says OPEN).

    None = not this shape (the per-sheet split decides as before)."""
    if len(contract.feeds) != 1 or len(sheet_tables) != 1:
        return None
    entry = sheet_tables[0]
    feed = contract.feeds[0]
    tables = list(entry.get("stage_tables") or [])
    if len(tables) < 2 or entry.get("segmented"):
        return None
    stated = {normalize(t) for t in feed.stage_target.tables}
    if {normalize(t) for t in tables} != stated or len(stated) != len(tables):
        return None
    frd_standard = {normalize(t) for t in feed.standard_target.tables}
    patterns = {file_identity(p): p for p in feed.file_name_patterns if normalize(p)}
    flags = [f"frd_feeds_split_per_table: FRD feed {feed.feed_name[:60]!r} names "
             f"{len(tables)} stage tables and STTM sheet {entry['sheet']!r} maps all of them "
             f"(its stage TableName column); {len(tables)} feeds, one per table, paired by "
             "exact table name"]
    feeds = []
    for index, table in enumerate(tables):
        rows = [r for r in entry.get("table_rows") or [] if normalize(r[0]) == normalize(table)]
        standards = list(dict.fromkeys(s for _t, s, _src in rows if s))
        if len(standards) == 1 and (not frd_standard or normalize(standards[0]) in frd_standard):
            standard = [standards[0]]
        else:
            standard = []
            flags.append(f"table_split_standard_unpaired:{table} — its rows name "
                         f"{len(standards)} standard table(s)"
                         + ("" if len(standards) != 1 else ", one the FRD does not name")
                         + "; no standard table for this feed")
        sources = list(dict.fromkeys(src for _t, _s, src in rows if src))
        keys = {file_identity(table), *(file_identity(s) for s in sources[:1])} - {""}
        hits = list(dict.fromkeys(patterns[k] for k in keys if k in patterns))
        if len(hits) == 1:
            own = hits
        else:
            own = []
            flags.append(f"frd_unstated:feeds[{index}].file_name_patterns — no FRD file pattern "
                         f"is table {table!r}{' or its source table' if sources else ''} "
                         "by name; left open")
        feeds.append(feed.model_copy(update={
            "feed_name": table,
            "file_name_patterns": own,
            "stage_target": feed.stage_target.model_copy(update={"tables": [table]}),
            "standard_target": feed.standard_target.model_copy(update={"tables": standard}),
        }))
    return contract.model_copy(update={"feeds": feeds}), flags, []


# ------------------------------------------------ M7: refused cells, per feed


def _feed_tokens(feed) -> set[str]:
    tokens: set[str] = set()
    for pattern in feed.file_name_patterns:
        tokens |= _name_tokens(pattern)
    for table in feed.stage_target.tables:
        tokens |= _name_tokens(table)
    return tokens


def _match_rows_to_feeds(keys: list[str], feeds: list) -> dict[int, int]:
    """feed index -> row index when a row's key names exactly one feed and
    that feed matches exactly one row (exact file-name match first, then
    the mutual-unique token rule of the file pairing)."""
    canonical = {i: {normalize(p).replace("ccyy", "yyyy") for p in f.file_name_patterns}
                 for i, f in enumerate(feeds)}
    out: dict[int, int] = {}
    for r, key in enumerate(keys):
        exact = [i for i, names in canonical.items()
                 if normalize(key).replace("ccyy", "yyyy") in names]
        if len(exact) == 1:
            out[exact[0]] = r
    all_tokens = {i: _feed_tokens(f) for i, f in enumerate(feeds)}
    if len(feeds) > 1:
        common = set.intersection(*all_tokens.values())
        all_tokens = {i: t - common for i, t in all_tokens.items()}
    row_tokens = {r: _name_tokens(k) for r, k in enumerate(keys)}

    def score(i: int, r: int) -> int:
        # overlapping token pairs (prefix-tolerant, as the file pairing counts them)
        return sum(1 for a in all_tokens[i] for b in row_tokens[r]
                   if a == b or a.startswith(b) or b.startswith(a))

    # Mutual BEST match with a strict margin, by elimination: a pair is
    # taken when the row scores higher for this feed than for any other
    # remaining feed AND this feed scores higher for the row than for any
    # other remaining row; matched pairs drop out and the pass repeats
    # ("Community Demographics" / "Community Risk" share a token — the
    # first pairs on its second token, the second pairs once the first is
    # gone). Ties decide nothing.
    progress = True
    while progress:
        progress = False
        remaining_feeds = [i for i in range(len(feeds)) if i not in out]
        remaining_rows = [r for r in range(len(keys)) if r not in out.values()]
        for i in remaining_feeds:
            scores = {r: score(i, r) for r in remaining_rows}
            best = max(scores.values(), default=0)
            if best == 0 or list(scores.values()).count(best) != 1:
                continue
            r = next(r for r, v in scores.items() if v == best)
            if all(best > score(j, r) for j in remaining_feeds if j != i):
                out[i] = r
                progress = True
                break
    return out


def _path_column(headers: list[str], values: dict[str, str], key: str) -> str | None:
    """The nested row's path cell: the header naming a path / location, else
    the first non-key cell that carries a path separator."""
    for header, value in values.items():
        if any(w in normalize(header) for w in ("path", "location")) and value:
            return value
    for value in values.values():
        if value != key and re.search(r"[\\/]", value):
            return value
    return None


def split_single_frd_feed(frd: FrdContract, profile: LayoutProfile, workbook, config: Config,
                          gaps: dict | None = None):
    """``extract-sttm``'s half of the split (2026-10-13, D1 pair_11): an FRD
    contract with ONE feed for a content-strategy STTM with several mapping
    sheets — an ``extract-frd`` contract, or one written before the layout
    stage split it — is split HERE by the layout stage's own rule
    (:func:`split_frd_feeds_by_sttm`, forced, then the per-feed structured
    fields). Returns a :class:`~codegen.contracts.sttm.FrdFeedSplit` (the
    derived feeds + flags, carried in the STTM contract so ``generate``
    joins the same split), or None when nothing splits."""
    from codegen.contracts.sttm import FrdFeedSplit

    if (len(frd.feeds) != 1 or not isinstance(profile, LayoutProfile)
            or profile.strategy != "content" or len(profile.mapping_sheets) < 2):
        return None
    facts = _sttm_facts(profile, workbook, config)
    split, flags, questions = split_frd_feeds_by_sttm(frd, facts, gaps or {}, config,
                                                      force=True,
                                                      remember=dict(profile.file_choices))
    if len(split.feeds) < 2:
        return None
    split, structured_flags = resolve_structured_fields(split, config)
    source = frd.feeds[0].feed_name
    lead = (f"frd_feeds_split_at_extract: the FRD contract {frd.contract_name!r} has ONE feed "
            f"({source[:60]!r}) and the STTM {len(profile.mapping_sheets)} mapping sheets — "
            f"split by extract-sttm into {len(split.feeds)} feeds "
            f"{[f.feed_name for f in split.feeds]} (the layout stage's rule; "
            "`codegen layout --frd-contract-out` writes the same split)")
    open_files = [f"frd_split_file_open:{q.key} — {q.reason}" for q in questions]
    return FrdFeedSplit(source_feed_name=source, feeds=list(split.feeds),
                        flags=[lead, *flags, *structured_flags, *open_files])


def resolve_structured_fields(contract: FrdContract, config: Config
                              ) -> tuple[FrdContract, list[str]]:
    """Give every (derived) feed its own row of a nested table / its own
    per-file block, matched by file name; everything else refused stays
    unstated and is flagged. Provenance on every value; nothing is guessed."""
    if not contract.structured:
        return contract, []
    flags: list[str] = []
    feeds = list(contract.feeds)
    for path, item in contract.structured.items():
        field = path.split(".", 1)[1] if "." in path else path
        where = f"FRD table {item.table} row {item.row} ({item.label!r})"
        if item.kind == "labelled_files":
            # M9.3: the reader already took the listed files as the feed's
            # patterns and flagged the unnamed feed (contract.extraction_flags).
            continue
        if item.kind in ("nested_table", "per_file_blocks") and item.rows:
            matched = _match_rows_to_feeds([r.key for r in item.rows], feeds)
            for index, feed in enumerate(feeds):
                key = f"feeds[{index}].{field}"
                row_index = matched.get(index)
                if row_index is None:
                    flags.append(f"frd_{item.kind}:{key} — no row of the {item.label!r} cell "
                                 f"names this feed's file uniquely ({where}); unstated")
                    continue
                row = item.rows[row_index]
                if field == "landing_location":
                    value = _path_column(item.headers, row.values, row.key)
                    if value is None:
                        flags.append(f"frd_{item.kind}:{key} — row {row.key!r} carries no path "
                                     f"({where}); unstated")
                        continue
                    feeds[index] = _feed_set(feed, "landing_location", value)
                    flags.append(f"frd_nested:{key} source_used:{where} row {row.key!r}: "
                                 f"{value!r}")
                elif field == "domain":
                    applied = []
                    for name in ("domain", "sub_domain"):
                        if row.values.get(name):
                            feeds[index] = _feed_set(feeds[index], name, row.values[name])
                            applied.append(f"{name}={row.values[name]!r}")
                    if applied:
                        flags.append(f"frd_nested:feeds[{index}].domain source_used:{where} block "
                                     f"{row.key!r}: {', '.join(applied)}")
                else:
                    flags.append(f"frd_{item.kind}:{key} — the {item.label!r} cell is a table "
                                 f"({where}); the feed is named after the STTM stage band")
        else:
            detail = {"pointer": f" → {item.target}",
                      "label_prefixed": f" (label {item.prefix_label!r} is not the field; the "
                                        "vendor label is used instead)",
                      "multiline": " (line breaks)",
                      "nested_table": "", "per_file_blocks": ""}[item.kind]
            for index in range(len(feeds)):
                flags.append(f"frd_{item.kind}:feeds[{index}].{field}{detail} — {where}; "
                             "unstated, resolved from the other documents")
    return contract.model_copy(update={"feeds": feeds}), flags


class _FeedGapFiller:
    """Applies the gap chain to ONE FRD feed (resolve/gapfill.py doctrine)."""

    def __init__(self, contract: FrdContract, index: int, facts: dict, vdd: dict,
                 gaps: dict, config: Config, base_dir: Path, result: GapFillResult,
                 asked: set[str] | None = None) -> None:
        from codegen.layout.hints import frd_field_help

        self.asked = asked or set()
        self.contract = contract
        self.feed = contract.feeds[index]
        self.patched = self.feed
        self.prefix = f"feeds[{index}]."
        self.meta = facts.get("meta", {})
        self.file_rows = facts.get("file_rows", [])
        self.band_schemas = facts.get("schemas", {})
        self.band_tables = facts.get("sheet_tables", [])
        self.vdd = vdd
        self.gaps = gaps
        self.config = config
        self.base_dir = base_dir
        self.result = result
        self._title = lambda path: frd_field_help(path, [], config)[0]

    # -- helpers ------------------------------------------------------------------
    def frd_stmt(self, dotted: str, value):
        from codegen.resolve.gapfill import Statement

        item = self.contract.field_provenance.get(self.prefix + dotted)
        cell = (f"table {item.table} row {item.row} ({item.label!r})" if item
                else "FRD contract")
        return Statement(str(value), "FRD", cell)

    def record_fill(self, dotted: str, statement) -> None:
        from codegen.resolve.gapfill import fill_flag

        key = self.prefix + dotted
        self.patched = _feed_set(self.patched, dotted, statement.value)
        self.result.fills.append({"field": key, "title": self._title(key),
                                  "value": statement.value, "source": statement.source,
                                  "cell": statement.cell})
        self.result.flags.append(fill_flag(key, statement))
        self.result.handled.add(key)

    def ask(self, dotted: str, kind: str, candidates: list[dict], reason: str) -> None:
        from codegen.layout.hints import frd_field_help

        key = self.prefix + dotted
        title, hint, _s = frd_field_help(key, [], self.config)
        self.result.questions.append(LayoutQuestion(
            document="frd", sheet=None, layer=None, role=key, reason=reason, header=[],
            candidates=candidates, title=title, hint=hint, kind=kind))
        self.result.handled.add(key)

    def handled(self, dotted: str) -> bool:
        return self.prefix + dotted in self.result.handled

    def _file_patterns(self) -> None:
        key = self.prefix + "file_patterns"
        if self.feed.file_name_patterns:
            return
        if self.prefix + "file_name_patterns" in self.asked:
            return          # the split asks which file feeds this table: one question, not two
        chosen = self.gaps.get(key)
        if chosen is not None:
            patterns = [p.strip() for p in re.split(r"[;\n,]+", chosen["value"]) if p.strip()]
            if patterns:
                self.patched = _feed_set(self.patched, "file_name_patterns", patterns)
                self.result.fills.append({"field": key, "title": "File name patterns",
                                          "value": "; ".join(patterns), "source": "user",
                                          "cell": "layout dialog / answers file"})
                self.result.flags.append(f"file_pattern_from_user:{key} — the FRD, the STTM "
                                         f"and the VDD name no file pattern; the person "
                                         f"stated {patterns}")
                self.result.handled.add(key)
                return
        sttm_states = bool(self.meta.get("file_names") or self.meta.get("file_name_example")
                           or self.file_rows)
        if sttm_states and len(self.contract.feeds) <= 1:
            # Single-feed FRD: the downstream contract resolver's per-sheet
            # STTM meta matches the global meta — defer safely.
            return          # the STTM states them: taken (and flagged) downstream
        # Multi-feed FRDs: self.meta is global (any mapping sheet) but the
        # downstream resolver uses the per-sheet meta of whichever STTM sheet
        # THIS feed is matched to — that sheet may have no file info.  Fall
        # through to VDD / ask so the layout stage resolves it here.
        stated = list(dict.fromkeys(self.vdd.get("file_patterns", [])))
        if stated:
            patterns = list(dict.fromkeys(v for v, _cell in stated))
            cells = ", ".join(cell for _v, cell in stated)
            self.patched = _feed_set(self.patched, "file_name_patterns", patterns)
            self.result.fills.append({"field": key, "title": "File name patterns",
                                      "value": "; ".join(patterns), "source": "VDD",
                                      "cell": cells})
            self.result.flags.append(f"file_pattern_from_vdd:{key} — the FRD and the STTM name "
                                     f"no file pattern; the VDD FILES sheet states {patterns} "
                                     f"({cells})")
            self.result.handled.add(key)
            return
        self.result.questions.append(LayoutQuestion(
            document="frd", sheet=None, layer=None, role=key, kind="text",
            reason="no document names a file pattern: the FRD states none, the STTM meta rows "
                   "'File Names' / 'File Name Example' and its File Details sheet are blank, "
                   "and no VDD FILES sheet states one",
            header=[], candidates=[], title="File name patterns",
            hint="Type the file name pattern(s) this feed arrives as, separated by ';' "
                 "(wildcards and date placeholders as the vendor writes them)."))
        self.result.handled.add(key)

    def _compound_frequency(self) -> None:
        """Chunk B (cold drill): an FRD Frequency cell naming SEVERAL cadences
        for several files ("Claims daily; reversals weekly") is no single
        value: the File Details row of THIS feed's file decides when it states
        exactly one, flagged with both cells; else the FRD text stays."""
        from codegen.extract.frd_docx import _CADENCE_RE

        current = _feed_get(self.feed, "frequency")
        if self.handled("frequency") or not current:
            return
        if len({m.lower() for m in _CADENCE_RE.findall(current)}) < 2:
            return
        wanted = {normalize(p).replace("ccyy", "yyyy") for p in self.feed.file_name_patterns}
        own = list(dict.fromkeys((r["frequency"], r["frequency_cell"]) for r in self.file_rows
                                 if r["frequency"]
                                 and normalize(r["name"]).replace("ccyy", "yyyy") in wanted))
        if len({value for value, _cell in own}) != 1:
            return
        value, cell = own[0]
        key = self.prefix + "frequency"
        self.patched = _feed_set(self.patched, "frequency", value)
        self.result.fills.append({"field": key, "title": self._title(key), "value": value,
                                  "source": "STTM", "cell": cell})
        self.result.flags.append(
            f"frd_frequency_per_file:{key} — the FRD {self._frequency_cell(current)} names "
            f"several cadences ({current!r}); this feed's File Details row {cell} states "
            f"{value!r}: taken")
        self.result.handled.add(key)

    def _frequency_cell(self, current: str) -> str:
        """The FRD cell the compound Frequency text was read from. A feed the
        split step derived from the FRD's one metadata block carries that
        block's text, recorded under feeds[0]: cite that cell, never a bare
        "FRD contract" (Chunk B review)."""
        item = self.contract.field_provenance.get(self.prefix + "frequency")
        if item is None and self.contract.feeds \
                and _feed_get(self.contract.feeds[0], "frequency") == current:
            item = self.contract.field_provenance.get("feeds[0].frequency")
        if item is None:
            return f"contract field {self.prefix}frequency"
        return f"table {item.table} row {item.row} ({item.label!r})"

    def _delimiter_in_format_text(self) -> bool:
        """Chunk B (cold drill): the FRD states no delimiter cell but its
        FORMAT cell names one ("Pipe delimited (|) text file", "CSV, comma
        delimited"). That is the FRD's statement: taken (flagged, citing the
        cell) when the STTM / VDD agree or are silent, a choice question when
        they state another character - never the STTM value silently. True
        when this step decided (or asked)."""
        from codegen.resolve.gapfill import Statement, delimiter_in_format, fill_flag
        from codegen.resolve.resolver import delimiter_char

        feed = self.feed
        if feed.delimiter or not feed.file_format:
            return False
        words = self.config.extractor.delimiter_words
        char = delimiter_in_format(feed.file_format, words)
        if char is None:
            return False
        frd = self.frd_stmt("file_format", feed.file_format)
        stated = Statement(char, "FRD", f"{frd.cell}: {feed.file_format!r}")
        others = []
        if "delimiter" in self.meta:
            value, cell = self.meta["delimiter"]
            others.append(Statement(value, "STTM", cell))
        others += [Statement(v, "VDD", cell) for v, cell in self.vdd.get("delimiter", [])]
        disagreeing = [o for o in others if delimiter_char(o.value, words) not in (None, char)]
        if disagreeing:
            seen: list = []
            for o in [stated, *disagreeing]:
                if all(delimiter_char(o.value, words) != delimiter_char(s.value, words)
                       for s in seen):
                    seen.append(o)
            self.ask("delimiter", "choice",
                     [{"value": s.value, "source": s.source, "cell": s.cell} for s in seen],
                     "the documents disagree - the FRD's format cell names one delimiter, the "
                     "STTM / VDD another; choose the value to use")
            return True
        self.patched = _feed_set(self.patched, "delimiter", char)
        key = self.prefix + "delimiter"
        self.result.fills.append({"field": key, "title": self._title(key), "value": char,
                                  "source": "FRD", "cell": stated.cell})
        self.result.flags.append(fill_flag(key, stated))
        self.result.handled.add(key)
        return True

    def _extension_formats(self, allowed: tuple[str, ...] = _FORMAT_EXTENSIONS) -> list:
        """One statement per distinct delimiter-naming extension (.csv / .psv /
        .tsv) of this feed's file patterns — a .txt / .dat name says nothing
        about the layout."""
        from codegen.resolve.gapfill import Statement

        out: list = []
        for pattern in _feed_get(self.patched, "file_name_patterns") or []:
            ext = PurePosixPath(pattern.strip()).suffix.lstrip(".").lower()
            if ext in allowed and all(s.value != ext for s in out):
                out.append(Statement(ext, "file name",
                                     f"file name pattern {pattern!r} (extension .{ext})"))
        return out

    def _own_vdd_rows(self, key: str) -> list[tuple[str, str]]:
        """The VDD FILES values of ``key`` on the row(s) whose file pattern is
        one of this feed's (one FILES row per file: the .psv file's row, not
        the first row's csv); every row when no row's pattern matches."""
        rows = self.vdd.get(key, [])
        wanted = {normalize(p).replace("ccyy", "yyyy")
                  for p in _feed_get(self.patched, "file_name_patterns") or []}
        own = {_cell_row(cell) for value, cell in self.vdd.get("file_patterns", [])
               if normalize(value).replace("ccyy", "yyyy") in wanted}
        return [(v, c) for v, c in rows if _cell_row(c) in own] or rows

    def _generic_text_format(self, stated: str) -> bool:
        """2026-10-13 (D1 pair_11): the FRD's format cell says only "Text
        files" — every csv / psv / tsv / txt file is a text file, so it never
        disagrees with an extension. The extension the documents state decides:
        the file patterns', else a STTM / VDD format cell that IS an extension
        ("csv", ".psv") — flag ``format_from_extension`` (a NOTE). Two
        different extensions = a QUESTION listing them; none = the chain as
        before (False)."""
        from codegen.resolve.gapfill import Statement

        frd = self.frd_stmt("file_format", stated)
        found = self._extension_formats(_TEXT_EXTENSIONS)
        cells: list = []
        if "file_format" in self.meta:
            value, cell = self.meta["file_format"]
            cells.append(Statement(value, "STTM", cell))
        cells += [Statement(v, "VDD", cell) for v, cell in self._own_vdd_rows("file_format")]
        for cell in cells:
            ext = normalize(cell.value).lstrip(".")
            if ext in _TEXT_EXTENSIONS:
                if all(s.value != ext for s in found):
                    found.append(Statement(ext, cell.source, f"{cell.cell}: {cell.value!r}"))
            elif cell.value and not any(normalize(cell.value) == normalize(g) for g in
                                        self.config.extractor.frd.generic_text_formats):
                return False                     # another document states a real format
        if not found:
            return False
        key = self.prefix + "file_format"
        if len(found) > 1:
            self.ask("file_format", "choice",
                     [{"value": s.value, "source": s.source, "cell": s.cell} for s in found],
                     f"the FRD format cell ({frd.cell}) says {stated!r} and the documents "
                     f"state {len(found)} extensions ({', '.join(s.value for s in found)}); "
                     "choose the format to use")
            return True
        ext = found[0]
        self.patched = _feed_set(self.patched, "file_format", ext.value)
        self.result.fills.append({"field": key, "title": self._title(key), "value": ext.value,
                                  "source": ext.source, "cell": ext.cell})
        self.result.flags.append(f"format_from_extension:{key} {ext.cell} -> {ext.value!r}; "
                                 f"the FRD format cell ({frd.cell}) says {stated!r}: a text file, "
                                 "not its layout (extractor.frd.generic_text_formats)")
        self.result.handled.add(key)
        return True

    def _file_format(self) -> bool:
        """Round 7 (D1 App, VDD attached, 2026-10-13): the FRD's 'Object/data
        Format' cell read "File Data Ingestion" — an ingestion TYPE, not a file
        format — and was asked as a disagreement with the VDD's csv. A cell that
        names none of ``extractor.frd.format_tokens`` is never a candidate: the
        file patterns' extension decides (.csv -> csv, flag
        ``format_from_extension`` = a NOTE citing the pattern), a STTM / VDD
        format that agrees confirms it, one that implies another delimiter is a
        QUESTION. An FRD format that names a delimiter the extension contradicts
        ("pipe" for a .csv file) is a QUESTION too. False = the chain as before
        (an FRD that states nothing, or a format the extension does not
        contradict)."""
        from codegen.resolve.gapfill import (
            Statement,
            distinct,
            format_delimiter,
            format_statements,
            is_format_text,
            same_value,
        )

        words = self.config.extractor.delimiter_words
        stated = _feed_get(self.feed, "file_format")
        if stated is None:
            return False
        if normalize(stated) in {normalize(g) for g in
                                 self.config.extractor.frd.generic_text_formats}:
            return self._generic_text_format(stated)
        exts = self._extension_formats()
        frd = self.frd_stmt("file_format", stated)
        if is_format_text(stated, self.config.extractor.frd.format_tokens):
            want = format_delimiter(stated, words)
            clash = [e for e in exts
                     if want is not None and format_delimiter(e.value, words) != want]
            if not clash:
                return False
            self.ask("file_format", "choice",
                     [{"value": s.value, "source": s.source, "cell": s.cell}
                      for s in [frd, *clash]],
                     "the FRD's format names one delimiter, the file name's extension "
                     "another; choose the format to use")
            return True
        key = self.prefix + "file_format"
        not_a_format = (f"the FRD's {frd.cell} {stated!r} names no file format (an ingestion "
                        "type, none of extractor.frd.format_tokens) — not a candidate")
        others = []
        if "file_format" in self.meta:
            value, cell = self.meta["file_format"]
            others.append(Statement(value, "STTM", cell))
        others += [Statement(v, "VDD", cell) for v, cell in self._own_vdd_rows("file_format")]
        others = distinct(format_statements(others))
        if not exts:
            self.result.flags.append(f"frd_format_not_a_format:{key} {not_a_format}")
            if len(others) == 1:
                self.record_fill("file_format", others[0])
            elif len(others) > 1:
                self.ask("file_format", "choice",
                         [{"value": s.value, "source": s.source, "cell": s.cell}
                          for s in others],
                         "the FRD states no file format and the other documents disagree")
            return True
        fixed = self.config.extractor.vdd.fixed_width_tokens
        clash = [o for o in others if not same_value(o.value, exts[0].value) and (
            format_delimiter(o.value, words) not in (None, format_delimiter(exts[0].value, words))
            or any(t.lower() in o.value.lower() for t in fixed))]
        if len(exts) > 1 or clash:
            self.ask("file_format", "choice",
                     [{"value": s.value, "source": s.source, "cell": s.cell}
                      for s in [*exts, *clash]],
                     "the file name's extension and the other documents name different "
                     "formats (the FRD's cell names none); choose the format to use")
            return True
        ext = exts[0]
        cell = ext.cell + "".join(f"; confirmed by {o.source} {o.cell}" for o in others)
        self.patched = _feed_set(self.patched, "file_format", ext.value)
        self.result.fills.append({"field": key, "title": self._title(key), "value": ext.value,
                                  "source": ext.source, "cell": cell})
        self.result.flags.append(f"format_from_extension:{key} {cell} -> {ext.value!r}; "
                                 f"{not_a_format}")
        self.result.handled.add(key)
        return True

    # -- the chain ----------------------------------------------------------------
    def run(self):
        from codegen.faq import load_faq
        from codegen.resolve.gapfill import (
            Statement,
            distinct,
            fill_flag,
            format_statements,
            parse_load_strategy_text,
            same_value,
            strategy_from_faq,
        )
        from codegen.resolve.resolver import normalize_feed_name

        feed, prefix, meta = self.feed, self.prefix, self.meta
        # a. the person's earlier answers win.
        for key, choice in self.gaps.items():
            if not key.startswith(prefix):
                continue
            dotted = key[len(prefix):]
            statement = Statement(choice["value"], "user", f"chosen from {choice['source']}")
            if dotted == "load_strategy":
                layers = (("stage", "standard") if choice.get("layer") in (None, "both")
                          else (choice["layer"],))
                for layer in layers:
                    self.record_fill(f"{layer}_target.load_strategy", statement)
                self.result.handled.add(key)
            elif dotted in _GAP_FIELDS or dotted in _ANSWERED_TEXT_FIELDS:
                self.record_fill(dotted, statement)
            elif dotted in _ANSWERED_LIST_FIELDS:
                values = [v.strip() for v in re.split(r"[;\n,]+", statement.value) if v.strip()]
                if values:
                    self.patched = _feed_set(self.patched, dotted, values)
                    self.result.fills.append({"field": key, "title": self._title(key),
                                              "value": "; ".join(values), "source": "user",
                                              "cell": statement.cell})
                    self.result.flags.append(fill_flag(key, statement))
                    self.result.handled.add(key)
            elif dotted in _TARGET_FIELDS:
                # A target the STTM leaves blank (the extractor's NEEDS_ANSWERS
                # key): written into the FRD feed, which the extractor reads next.
                layer_target, _dot, part = dotted.partition(".")
                value = ([t.strip() for t in re.split(r"[;\n,]+", statement.value) if t.strip()]
                         if part == "tables" else statement.value.strip())
                self.patched = _feed_set(self.patched, f"{layer_target}."
                                         f"{'schema_name' if part == 'schema' else part}", value)
                self.result.fills.append({"field": key, "title": self._title(key),
                                          "value": statement.value, "source": "user",
                                          "cell": statement.cell})
                self.result.flags.append(fill_flag(key, statement))
                self.result.handled.add(key)
            # feeds[i].file_name_patterns answers are consumed by the split step.

        # a'. file patterns (M9.1b): FRD -> STTM meta rows / File Details (the
        # extractor + contract resolver take those, flagged file_pattern_from_
        # sttm) -> VDD FILES sheet -> ask. Only the last two are decided here.
        self._file_patterns()

        # b. file format / delimiter: STTM meta row, then VDD FILES.
        for dotted in ("file_format", "delimiter"):
            if self.handled(dotted):
                continue
            if dotted == "delimiter" and self._delimiter_in_format_text():
                continue
            if dotted == "file_format" and self._file_format():
                continue
            others = []
            if dotted in meta:
                value, cell = meta[dotted]
                others.append(Statement(value, "STTM", cell))
            others += [Statement(v, "VDD", cell) for v, cell in self.vdd.get(dotted, [])]
            if dotted == "file_format":
                # ".dat" names the file, not its layout: never a disagreement
                # with a document that states a format.
                others = format_statements(
                    others if _feed_get(feed, dotted) is None
                    else [self.frd_stmt(dotted, _feed_get(feed, dotted)), *others])
                others = [o for o in others if o.source != "FRD"]
            words = self.config.extractor.delimiter_words
            others = distinct(others, words)
            current = _feed_get(feed, dotted)
            if current is not None:
                disagreeing = [o for o in others if not same_value(current, o.value, words)]
                if disagreeing:
                    self.ask(dotted, "choice",
                             [{"value": s.value, "source": s.source, "cell": s.cell}
                              for s in distinct([self.frd_stmt(dotted, current), *disagreeing],
                                                words)],
                             "the documents disagree — choose the value to use")
                continue
            if len(others) == 1:
                self.record_fill(dotted, others[0])
            elif len(others) > 1:
                self.ask(dotted, "choice",
                         [{"value": s.value, "source": s.source, "cell": s.cell} for s in others],
                         "the FRD is silent and the other documents disagree")
            # no statement anywhere: the ordinary FRD question stays.

        # b'. frequency: the STTM File Details row for THIS feed's file (the
        # FRD's cell may point there), then the meta row, then the VDD cadence.
        self._compound_frequency()
        if not self.handled("frequency") and _feed_get(feed, "frequency") is None:
            wanted = {normalize(p).replace("ccyy", "yyyy") for p in feed.file_name_patterns}
            own_rows = [r for r in self.file_rows if r["frequency"]
                        and normalize(r["name"]).replace("ccyy", "yyyy") in wanted]
            others: list[Statement] = []
            if own_rows:
                others = distinct([Statement(r["frequency"], "STTM", r["frequency_cell"])
                                   for r in own_rows])
            elif "frequency" in meta:
                value, cell = meta["frequency"]
                others = [Statement(value, "STTM", cell)]
            others += [Statement(v, "VDD", cell) for v, cell in self.vdd.get("frequency", [])]
            others = distinct(others)
            if len(others) == 1:
                self.record_fill("frequency", others[0])
            elif len(others) > 1:
                self.ask("frequency", "choice",
                         [{"value": s.value, "source": s.source, "cell": s.cell} for s in others],
                         "the FRD states no frequency and the other documents disagree")

        # b''. LOBs (Chunk B): the STTM meta row 'LOB' when the FRD lists none
        # (the resolver stops on a feed without LOBs; the multi-table step
        # already reads the same row for the per-LOB files).
        if not self.handled("lobs") and not feed.lobs and "lob" in meta:
            value, cell = meta["lob"]
            lobs = [v.strip() for v in re.split(r"[;\n,]+", value) if v.strip()]
            if lobs:
                self.patched = _feed_set(self.patched, "lobs", lobs)
                key = prefix + "lobs"
                stated = Statement(value, "STTM", cell)
                self.result.fills.append({"field": key, "title": self._title(key),
                                          "value": "; ".join(lobs), "source": "STTM",
                                          "cell": cell})
                self.result.flags.append(fill_flag(key, stated))
                self.result.handled.add(key)

        # c. load strategies: STTM per layer, else FAQ (stage) / blank (standard).
        text_, sttm_cell = meta.get("load_strategy", (None, "STTM meta row 'Load Strategy'"))
        parsed = parse_load_strategy_text(text_, self.config.extractor.frd.target_schema_markers)
        for layer in ("stage", "standard"):
            dotted = f"{layer}_target.load_strategy"
            if self.handled(dotted):
                continue
            current = _feed_get(feed, dotted)
            per_layer = parsed.get(layer)
            if current is not None:
                if per_layer and not same_value(current, per_layer):
                    self.ask(dotted, "choice",
                             [{"value": current, "source": "FRD",
                               "cell": self.frd_stmt(dotted, current).cell},
                              {"value": per_layer, "source": "STTM", "cell": sttm_cell}],
                             "the FRD and the STTM disagree — choose the value to use")
                continue
            if per_layer:
                self.record_fill(dotted, Statement(per_layer, "STTM", sttm_cell))
        stage_missing = _feed_get(self.patched, "stage_target.load_strategy") is None
        standard_missing = _feed_get(self.patched, "standard_target.load_strategy") is None
        if "any" in parsed and not self.handled("load_strategy") and (
                stage_missing or standard_missing):
            value = parsed["any"]
            self.ask("load_strategy", "layer",
                     [{"layer": "stage", "value": value}, {"layer": "standard", "value": value},
                      {"layer": "both", "value": value}],
                     f"the STTM states one load strategy ({value!r}, {sttm_cell}) without "
                     "saying which layer it applies to")
            self.result.handled.update({prefix + "stage_target.load_strategy",
                                        prefix + "standard_target.load_strategy"})
        if stage_missing and not self.handled("stage_target.load_strategy"):
            faq = load_faq(normalize_feed_name(feed.feed_name), self.config,
                           base_dir=self.base_dir)
            statement = strategy_from_faq(faq)
            if statement is not None:
                self.record_fill("stage_target.load_strategy", statement)

        # d. the STTM stage band is authoritative for catalog / schema / tables:
        # never a question; flagged when the FRD stated nothing.
        for dotted in _STTM_AUTHORITATIVE:
            if self.handled(dotted):
                continue                          # the person answered it (step a)
            band_states = (bool(self.band_tables) if dotted.endswith("tables")
                           else "stage.schema" in self.band_schemas)
            if _feed_get(feed, dotted) in (None, [], "") and band_states:
                self.result.flags.append(f"frd_unstated:{prefix}{dotted} source_used:STTM stage "
                                         "band (authoritative; never asked)")
            elif _feed_get(feed, dotted) in (None, [], ""):
                # First ACFC run: the band is blank too — the extractor's chain
                # (File Details / layer convention) applies, else NEEDS_ANSWERS.
                self.result.flags.append(f"frd_unstated:{prefix}{dotted} source_used:none — the "
                                         "STTM stage band leaves it blank too; answer "
                                         f"`{prefix}{dotted}` under gaps: unless File Details "
                                         "or the layer convention states it")
            self.result.handled.add(prefix + dotted)
        return self.patched


def fill_frd_gaps(contract: FrdContract, frd_doc, facts: dict, vdd: dict, gaps: dict,
                  config: Config, base_dir: Path, asked: set[str] | None = None
                  ) -> GapFillResult:
    """Apply the FRD gap chain to every feed (see resolve/gapfill.py). ``asked``:
    the keys the split step already asks (``feeds[i].file_name_patterns``) —
    never asked a second time in another form."""
    result = GapFillResult(contract=contract)
    feeds = [_FeedGapFiller(contract, index, facts, vdd, gaps, config, base_dir, result,
                            asked or set()).run()
             for index in range(len(contract.feeds))]
    result.contract = contract.model_copy(update={"feeds": feeds})
    return result


def discovery_for(profile: LayoutProfile, workbook) -> Discovery:
    """A Discovery the extractor reads through, from a resolved profile."""
    return Discovery(profile=profile, workbook=workbook, diagnostics=list(profile.notes))


_ANSWER_KEY_RE = re.compile(r"^[^/]+/(source|rules|stage|standard|files|fields)/[a-z_]+$")


def parse_answers(raw: dict) -> dict:
    """Validate the shape of a user's answers payload (values only; the
    merge step re-validates every claim against the document). Answer keys
    carry sheet names: every error is ``located`` at the answers file only."""
    out: dict = {"sttm": {}, "frd": {}, "vdd": {}}

    def bad(message: str) -> BaseException:
        return located(ValueError(message), Coord("answers"), "answers.file")

    for document in ("sttm", "vdd"):
        for key, col in (raw.get(document) or {}).items():
            if document == "sttm" and is_band_layer_key(str(key)):
                # Chunk A: <sheet>/band[<n>]/layer is answered with a layer.
                if col not in BAND_LAYER_CHOICES:
                    raise bad(f"bad STTM answer {key!r}: {col!r} - a band layer is one "
                              f"of {' | '.join(BAND_LAYER_CHOICES)}")
                out[document][str(key)] = col
                continue
            if not _ANSWER_KEY_RE.match(str(key)) or not isinstance(col, int) or col < 1:
                raise bad(f"bad {document.upper()} answer {key!r}: {col!r}")
            out[document][key] = col
    for path, claim in (raw.get("frd") or {}).items():
        if not isinstance(claim, dict):
            raise bad(f"bad FRD answer {path!r}")
        out["frd"][path] = claim
    out["gaps"] = {}
    for key, choice in (raw.get("gaps") or {}).items():
        if not isinstance(choice, dict) or not isinstance(choice.get("value"), str):
            raise bad(f"bad gap answer {key!r}: expected {{value, layer?, source?}}")
        if choice.get("layer") not in (None, "stage", "standard", "both"):
            raise bad(f"bad gap answer {key!r}: layer must be stage | standard | both")
        out["gaps"][str(key)] = {"value": choice["value"], "layer": choice.get("layer"),
                                 "source": str(choice.get("source") or "user")}
    from codegen.resolve.widths import parse_width_answers

    parse_width_answers(out["gaps"])      # a width answer must be a positive integer
    return out


__all__ = [
    "CrossCheck",
    "DocumentResolution",
    "LayoutQuestion",
    "PairResolution",
    "band_layer_question",
    "cross_check",
    "discovery_for",
    "parse_answers",
    "resolve_frd",
    "resolve_pair",
    "resolve_workbook",
    "vocabulary_hash",
]
