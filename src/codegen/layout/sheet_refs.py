"""Sheet references in answer keys (post57, D1 post55 leak report).

All ten leak hits of the post55 12-pair run were sheet names inside the
needed-answer keys the harness reads from stdout (``<sheet>/source/field_name``,
``<sheet>/band[1]/layer``). A key printed on a QUESTION / UNRESOLVED / NOTE /
ROLE line now names a sheet by its POSITION: ``sheets[<i>]`` = the workbook's
i-th sheet, 0-based like ``feeds[<i>]`` (the coordinates tier's ``STTM sheet
<i+1>`` is the same sheet, 1-based). The sheet name stays in the ACFC-side
files only: the answers template's ``sheet_name:`` reading aid and
``unresolved_headers.md``.

Internally nothing changes: layout profiles, caches, the resolver's answers
and the App keep ``<sheet name>/<layer>/<role>``. Translation happens at the
edge — :func:`public_key` on the way out, :func:`sheet_name` on the way in (an
answers-file entry may give ``sheet: "sheets[2]"`` or the name as written).
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Mapping
from io import BytesIO
from pathlib import Path

SHEET_REF_RE = re.compile(r"^sheets\[(\d+)\]$")
_KEY_REF_RE = re.compile(r"^sheets\[(\d+)\]/")
# A note quoting a role key in parentheses: "answers[0] (<sheet>/stage/table) ..."
_NOTE_KEY_RE = re.compile(r"\(([^()]*/[^()/]+/[^()/]+)\)")
UNKNOWN_REF = "sheets[?]"


def sheet_ref(index: int) -> str:
    return f"sheets[{index}]"


def _positions(names: Mapping[str, int | None] | Iterable[str] | None) -> dict[str, int]:
    if names is None:
        return {}
    if isinstance(names, Mapping):
        return {str(k): int(v) for k, v in names.items() if v is not None}
    return {str(n): i for i, n in enumerate(getattr(names, "sheetnames", names))}


def workbook_sheet_names(path: Path | str | None) -> list[str]:
    """A workbook's sheet names in order, read from its bytes (openpyxl 3.1.2
    keeps a file handle open otherwise); [] when it cannot be read."""
    if path is None:
        return []
    try:
        from openpyxl import load_workbook

        book = load_workbook(BytesIO(Path(path).read_bytes()), read_only=True)
    except Exception:       # noqa: BLE001 - a reference never fails a run
        return []
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def public_sheet(sheet: str | None, names) -> str:
    """``sheets[<i>]`` for a sheet name (``UNKNOWN_REF`` when it is not one of
    ``names``); a reference passes through."""
    if sheet is not None and SHEET_REF_RE.match(str(sheet)):
        return str(sheet)
    index = _positions(names).get(str(sheet))
    return UNKNOWN_REF if index is None else sheet_ref(index)


def public_key(key: str, names) -> str:
    """``<sheet name>/<rest>`` -> ``sheets[<i>]/<rest>``. ``names``: the
    workbook's sheet names in order (a list / a workbook) or {name: index}.
    A key that names no listed sheet (``feeds[0].file_format``, a key already
    in reference form) is returned as it is."""
    key = str(key)
    if _KEY_REF_RE.match(key):
        return key
    for name, index in sorted(_positions(names).items(), key=lambda kv: -len(kv[0])):
        if key.startswith(name + "/"):
            return sheet_ref(index) + key[len(name):]
    return key


def public_role_key(key: str, names) -> str:
    """:func:`public_key` for a key that IS a role / band key
    (``<sheet>/<layer>/<role>``): a sheet missing from ``names`` reads
    ``sheets[?]`` rather than its name."""
    out = public_key(key, names)
    if out != key or _KEY_REF_RE.match(key):
        return out
    head, sep, rest = key.rpartition("/")
    sheet, sep2, layer = head.rpartition("/")
    return f"{UNKNOWN_REF}/{layer}/{rest}" if sep and sep2 and sheet else key


def public_note(text: str, names) -> str:
    """A note quoting role keys in parentheses, each made public."""
    return _NOTE_KEY_RE.sub(lambda m: f"({public_key(m.group(1), names)})", str(text))


def sheet_name(value: str | None, names) -> str | None:
    """The sheet an answers entry means: ``sheets[<i>]`` -> the i-th name
    (None when out of range or ``names`` is unknown); a name as written
    passes through."""
    if value is None:
        return None
    match = SHEET_REF_RE.match(str(value).strip())
    if match is None:
        return value
    by_index = {i: n for n, i in _positions(names).items()}
    return by_index.get(int(match.group(1)))


def private_key(key: str, names) -> str:
    """``sheets[<i>]/<rest>`` -> ``<sheet name>/<rest>`` (unchanged when the
    reference is out of range)."""
    match = _KEY_REF_RE.match(str(key))
    if match is None:
        return str(key)
    name = sheet_name(f"sheets[{match.group(1)}]", names)
    return str(key) if name is None else name + str(key)[match.end() - 1:]


__all__ = ["SHEET_REF_RE", "UNKNOWN_REF", "private_key", "public_key", "public_note",
           "public_role_key", "public_sheet", "sheet_name", "sheet_ref",
           "workbook_sheet_names"]
