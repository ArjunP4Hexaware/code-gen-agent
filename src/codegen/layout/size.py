"""How big a workbook is, without reading it (M11).

A real 140k-cell STTM made the document parser child run past its budget and
be killed: the App only learnt "the document was not read within 120s". The
discovery and classification scans are what cost the time — several of them
walk every row of every sheet — so the cheap question has to be asked FIRST:

    how many cells does this workbook declare?

``read_only=True`` answers it from each sheet's ``<dimension>`` tag without
materialising a single row (the full load stays where it is: the fingerprint
hashes merged-cell ranges, which read-only mode does not expose). Over the
cap, the document is ``unreadable`` WITH the count and the cap in the reason
— a verdict, not a killed process.

The declared dimension is a hint, not a promise: some writers declare the
whole 1,048,576-row sheet. ``_sheet_cells`` therefore ignores a dimension
whose row count is the sheet maximum and falls back to the column count.
"""

from __future__ import annotations

from pathlib import Path

# openpyxl's own sheet limits — a dimension claiming these states nothing.
_MAX_ROWS = 1_048_576
_MAX_COLUMNS = 16_384


class WorkbookTooLarge(RuntimeError):
    """The workbook declares more cells than the configured cap."""


def _sheet_cells(ws) -> int:
    rows = ws.max_row or 0
    columns = ws.max_column or 0
    if rows >= _MAX_ROWS or columns >= _MAX_COLUMNS:
        # A writer that declares the whole grid tells us nothing about the
        # data; count the columns only, so such a file is never refused on a
        # declaration alone.
        return columns if columns < _MAX_COLUMNS else 0
    return rows * columns


# A cell of a shared-string sheet costs roughly this many bytes of sheet XML
# (`<c r="A1" t="s"><v>12</v></c>` and friends). Only used when a workbook
# declares no usable dimension, and then the message says "about".
_BYTES_PER_CELL = 40


def workbook_cells(path: Path) -> tuple[int, dict[str, int], bool]:
    """(total cells, per-sheet counts, estimated?).

    Read from each sheet's declared ``<dimension>``. A writer that declares
    none (openpyxl's own write-only mode, some exporters) leaves nothing to
    count, so the total is then ESTIMATED from the uncompressed size of the
    sheet XML in the zip — read from the zip directory, nothing decompressed.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        per_sheet = {ws.title: _sheet_cells(ws) for ws in workbook.worksheets}
    finally:
        workbook.close()
    total = sum(per_sheet.values())
    if total:
        return total, per_sheet, False
    estimated = _estimate_from_zip(path)
    return sum(estimated.values()), estimated, True


def _estimate_from_zip(path: Path) -> dict[str, int]:
    """Per-sheet cell ESTIMATE from the zip directory's uncompressed sizes."""
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            return {
                info.filename.rsplit("/", 1)[-1]: info.file_size // _BYTES_PER_CELL
                for info in archive.infolist()
                if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml")
            }
    except Exception:  # noqa: BLE001 — not this function's verdict to give
        return {}


def check_workbook_size(path: Path, max_cells: int) -> None:
    """Raise ``WorkbookTooLarge`` when the workbook is over the cap. A cap of
    0 (or less) is off. A workbook that cannot even be opened read-only is
    left alone — the reader that follows reports that in its own words."""
    if max_cells <= 0:
        return
    try:
        total, per_sheet, estimated = workbook_cells(Path(path))
    except Exception:  # noqa: BLE001 — not this function's verdict to give
        return
    if total > max_cells:
        biggest = sorted(per_sheet.items(), key=lambda kv: -kv[1])[:3]
        detail = ", ".join(f"{name} {cells:,}" for name, cells in biggest if cells)
        raise WorkbookTooLarge(
            f"unreadable: {'about ' if estimated else ''}{total:,} cells, > cap {max_cells:,} "
            f"(inputs.max_workbook_cells; largest sheets: {detail}) — the document was not "
            "read. Raise the cap for this workspace, or split the workbook.")
