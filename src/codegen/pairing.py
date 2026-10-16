"""Pairing by CONTENT (M8.2): which FRD / Vendor Data Dictionary belongs to
the chosen STTM.

Ticket numbers cannot decide it — in the ACFC inventory four pairs share one
ticket (docs/acfc/RETROFIT_LOG.md §7). So every candidate is scored against
what the STTM itself says, with the same deterministic readers the run uses
(synonyms only — never a model):

  FRD   feed / object name found in the STTM's header region; target tables
        and schemas against the STTM's stage / standard bands; file name
        patterns against the STTM's file-details rows (CCYY → YYYY)
  VDD   FILES-sheet file patterns / format / delimiter / cadence against the
        STTM's file rows and meta rows; one-sheet-per-table dictionaries by
        sheet name against the STTM's tables

plus weak NAME signals — a shared ticket number, the same pairing key
(``name_key``), a token overlap at or above ``inputs.pairing.name_fuzzy_threshold``
(``name_fuzzy``) — worth less than any content signal; an overlap below the
threshold (``name_overlap``) weighs nothing and only ranks a question. A
candidate is paired only when it wins by ``inputs.pairing.margin`` and reaches
``min_score``; otherwise the top candidates are returned as a QUESTION (the
layout dialog / the CLI answers file) and nothing is paired. A pairs OVERRIDE
file (``pairing_overrides.yaml``, 2026-10-13) wins over everything, then an
explicit ``demo.pairing_map`` entry. Every signal carries the cells it rests on;
every pair made by override, key or fuzzy match is logged
(``codegen.pairing``: "paired by override / key / fuzzy").
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field, replace
from pathlib import Path

from codegen.config import Config
from codegen.demo_sources import (
    _pair_key,
    auto_pair_document,
    classify_document_name,
    document_stem,
    match_by_name,
    name_overlap,
    pairing_key,
)
from codegen.layout.discover import normalize

_log = logging.getLogger("codegen.pairing")

_NAME_SIGNALS = ("ticket", "name_key", "name_fuzzy", "name_overlap")
# How a log line / a reason names each rule.
_RULE_WORDS = {"override": "override", "name_key": "key", "name_fuzzy": "fuzzy",
               "pairing_map": "the config pairing map", "ticket": "ticket",
               "content": "content", "same_folder": "the only candidate in the folder"}


@dataclass(frozen=True)
class PairSignal:
    name: str
    weight: float
    detail: str


@dataclass(frozen=True)
class PairCandidate:
    name: str
    score: float
    signals: tuple[PairSignal, ...] = ()
    name_similarity: float = 0.0        # token overlap with the STTM's name (ranks a question)

    def summary(self) -> str:
        return "; ".join(f"{s.name} (+{s.weight:g}): {s.detail}" for s in self.signals) \
            or "no content in common"


@dataclass(frozen=True)
class PairDecision:
    kind: str                       # "frd" | "vdd"
    sttm: str
    chosen: str | None
    rule: str | None    # override | pairing_map | ticket | name_key | name_fuzzy | content | None
    candidates: tuple[PairCandidate, ...] = ()
    reason: str = ""

    min_score: float = 0.0
    # Undecided but owed a question although no candidate scores: the names
    # overlap below the threshold, or an override names a missing document.
    ask: bool = False

    @property
    def ambiguous(self) -> bool:
        """Undecided AND worth asking. An FRD is required, so any scored
        candidate is worth showing. A VDD is optional: a candidate that shares
        SOMETHING (a score above 0) is asked about — M15.4: a best score above
        0 but below ``min_score`` used to be a dead zone, neither paired nor
        asked, and it blocked the same-folder rescue — and, kept as found, so
        is a no-content decision (``min_score`` 0.0 on those branches: the
        dictionary is OFFERED rather than silently dropped). A name overlap
        below the fuzzy threshold is never a pair — it is asked (``ask``)."""
        if self.chosen is not None or not self.candidates:
            return False
        if self.ask:
            return True
        best = self.candidates[0].score
        if self.kind == "frd":
            return best > 0
        return best > 0 or best >= self.min_score

    def question(self) -> dict:
        """The decision as a layout-dialog ``choice`` question (same shape
        ``LayoutQuestion.as_dict`` produces; answered through ``gaps``)."""
        label = "FRD" if self.kind == "frd" else "Vendor Data Dictionary"
        return {
            "document": self.kind, "key": f"pair.{self.kind}", "sheet": None, "layer": None,
            "role": f"pair.{self.kind}", "kind": "choice",
            "title": f"Which {label} belongs to {self.sttm}?",
            "reason": self.reason,
            "hint": ("Scored by content: feed / object name, target schema and tables, file "
                     "patterns; ticket number and file name are weak signals only."),
            "header": [],
            "candidates": [{"value": c.name, "source": f"candidate {i} · score {c.score:g}",
                            "cell": c.summary()}
                           for i, c in enumerate(self.candidates, start=1)],
            "suggested": None, "suggested_reason": "",
        }


# -- facts ----------------------------------------------------------------------------

@dataclass
class _Facts:
    texts: list[tuple[str, str]] = field(default_factory=list)       # (citation, value)
    tables: dict[str, str] = field(default_factory=dict)            # normalized -> citation
    schemas: dict[str, str] = field(default_factory=dict)
    files: dict[str, str] = field(default_factory=dict)             # canonical name -> citation
    meta: dict[str, tuple[str, str]] = field(default_factory=dict)  # key -> (value, citation)
    names: list[str] = field(default_factory=list)                  # feed / object names
    raw_tables: list[str] = field(default_factory=list)             # as written


_cache: dict[tuple, _Facts] = {}


def facts_to_dict(facts: _Facts) -> dict:
    """A document's pairing facts as JSON (M9.3: computed once in the UI's
    background task and kept in the state role)."""
    return {"texts": [list(t) for t in facts.texts], "tables": dict(facts.tables),
            "schemas": dict(facts.schemas), "files": dict(facts.files),
            "meta": {k: list(v) for k, v in facts.meta.items()},
            "names": list(facts.names), "raw_tables": list(facts.raw_tables)}


def facts_from_dict(data: dict) -> _Facts:
    return _Facts(texts=[tuple(t) for t in data.get("texts", [])],
                  tables=dict(data.get("tables", {})), schemas=dict(data.get("schemas", {})),
                  files=dict(data.get("files", {})),
                  meta={k: tuple(v) for k, v in data.get("meta", {}).items()},
                  names=list(data.get("names", [])), raw_tables=list(data.get("raw_tables", [])))


def empty_facts() -> _Facts:
    """The facts of a document that was NOT read (unreadable, or not read within
    its time budget): passed as ``known_facts`` it scores on its name alone and
    ``pair_by_content`` never opens it."""
    return _Facts()


def document_facts(kind: str, path: Path, config: Config, base_dir: Path) -> _Facts:
    """The pairing facts of one document: ``kind`` = sttm | frd | vdd."""
    if kind == "sttm":
        return sttm_facts(path, config, base_dir)
    if kind == "frd":
        return frd_facts(path, config, base_dir)
    return vdd_facts(path, config, base_dir, ())


_NAME_STEM_MIN = 4


def _canonical_file(name: str) -> str:
    """The extractor's comparison form (date placeholders CCYY → YYYY), then
    ``normalize`` — an FRD may write a file name with spaces where the STTM
    uses underscores."""
    return normalize(name.strip().lower().replace("ccyy", "yyyy"))


def _name_match(name: str, candidate: str) -> bool:
    """Stricter than the layout cross-check's stem match, because here a
    false positive pairs the WRONG documents: the whole name inside the text,
    or EVERY token of the name matching a word — short tokens exactly
    ("FEED_8" never matches "Feed Type"), longer ones by shared stem
    ("Accumulators" ↔ "ACCUM")."""
    a, b = normalize(name), normalize(candidate)
    if not a or not b:
        return False
    if a in b:
        return True
    words = b.split()

    def hit(token: str) -> bool:
        if len(token) < _NAME_STEM_MIN:
            return token in words
        return any(len(w) >= _NAME_STEM_MIN and (w.startswith(token) or token.startswith(w))
                   for w in words)

    tokens = a.split()
    return any(len(t) >= _NAME_STEM_MIN for t in tokens) and all(hit(t) for t in tokens)


def _key(path: Path, kind: str, extra: tuple = ()) -> tuple:
    stat = path.stat()
    return (kind, str(path), stat.st_mtime_ns, stat.st_size, *extra)


def sttm_facts(path: Path, config: Config, base_dir: Path) -> _Facts:
    key = _key(path, "sttm")
    if key in _cache:
        return _cache[key]
    from codegen.layout.resolve import _sttm_facts, resolve_workbook

    doc, workbook = resolve_workbook(path, config, provider=None, base_dir=base_dir)
    raw = _sttm_facts(doc.profile, workbook, config)
    facts = _Facts(texts=list(raw["texts"]))
    for entry in raw["sheet_tables"]:
        for layer in ("stage", "standard"):
            if entry.get(f"{layer}_table"):
                facts.raw_tables.append(entry[f"{layer}_table"])
                facts.tables[normalize(entry[f"{layer}_table"])] = (
                    f"{entry['sheet']} {layer} band")
            if entry.get(f"{layer}_schema"):
                facts.schemas[normalize(entry[f"{layer}_schema"])] = (
                    f"{entry['sheet']} {layer} band")
    for name, citation in raw["files"]:
        facts.files[_canonical_file(name)] = citation
    for meta_key, (value, citation) in raw["meta"].items():
        facts.meta[meta_key] = (value, citation)
    _cache[key] = facts
    return facts


def frd_facts(path: Path, config: Config, base_dir: Path) -> _Facts:
    key = _key(path, "frd")
    if key in _cache:
        return _cache[key]
    from codegen.contracts.frd import FrdContract

    if path.suffix.lower() == ".docx":
        from codegen.extract.frd_docx import read_frd
        from codegen.layout.resolve import resolve_frd

        doc, content = resolve_frd(path, config, provider=None, base_dir=base_dir)
        contract = read_frd(content, doc.profile, config, document_name=path.name,
                            generated_date="1970-01-01")
    else:
        contract = FrdContract.model_validate_json(path.read_text(encoding="utf-8"))
    facts = _Facts()
    for index, feed in enumerate(contract.feeds):
        where = f"{path.name} feeds[{index}]"
        if feed.feed_name:
            facts.names.append(feed.feed_name)
        for layer, target in (("stage", feed.stage_target), ("standard", feed.standard_target)):
            if target is None:
                continue
            for table in target.tables or []:
                facts.tables[normalize(table)] = f"{where}.{layer}_target.tables"
            if target.schema_name:
                facts.schemas[normalize(target.schema_name)] = f"{where}.{layer}_target.schema"
        for pattern in feed.file_name_patterns or []:
            facts.files[_canonical_file(pattern)] = f"{where}.file_name_patterns"
        for meta_key in ("file_format", "delimiter", "frequency"):
            value = getattr(feed, meta_key, None)
            if value:
                facts.meta.setdefault(meta_key, (value, f"{where}.{meta_key}"))
    _cache[key] = facts
    return facts


def vdd_facts(path: Path, config: Config, base_dir: Path, sttm_tables: tuple[str, ...]) -> _Facts:
    key = _key(path, "vdd", sttm_tables)
    if key in _cache:
        return _cache[key]
    from openpyxl.utils import get_column_letter

    from codegen.layout.profile import Role
    from codegen.layout.resolve import _vdd_facts, resolve_workbook, text

    doc, workbook = resolve_workbook(path, config, provider=None, base_dir=base_dir,
                                     document="vdd", sttm_tables=list(sttm_tables))
    facts = _Facts()
    for meta_key, values in _vdd_facts(doc.profile, workbook).items():
        if values:
            facts.meta[meta_key] = values[0]
    for sp in doc.profile.sheets:
        if sp.kind == "vdd_fields":
            facts.tables[normalize(sp.name)] = f"{path.name} sheet {sp.name!r}"
        band = sp.band("files") if sp.kind == "vdd_files" and sp.header_row else None
        col = band.column(Role.FILE_PATTERN) if band else None
        if col is None:
            continue
        rows = workbook[sp.name].iter_rows(min_row=sp.header_row + 1, values_only=True)
        for row_index, row in enumerate(rows, start=sp.header_row + 1):
            value = text(row[col - 1]) if col - 1 < len(row) else ""
            if value:
                facts.files[_canonical_file(value)] = (
                    f"{sp.name}!{get_column_letter(col)}{row_index}")
    _cache[key] = facts
    return facts


# -- scoring --------------------------------------------------------------------------

def _name_signals(sttm_name: str, candidate: str, weights: dict[str, float],
                  threshold: float) -> list[PairSignal]:
    """The weak NAME signals: a shared ticket; the same pairing key
    (``name_key``) or an overlap at the threshold (``name_fuzzy``) — both
    weighed by the ``name_stem`` knob; an overlap below it (``name_overlap``,
    weight 0: it never pairs, it ranks the question)."""
    signals = []
    shared = _pair_key(sttm_name) & _pair_key(candidate)
    if shared:
        signals.append(PairSignal("ticket", weights["ticket"],
                                  f"both names carry {', '.join(sorted(shared))}"))
    key = pairing_key(sttm_name)
    similarity, stems = name_overlap(sttm_name, candidate)
    stem_text = ("; stem-matched " + ", ".join(f"{a!r} ~ {b!r}" for a, b in stems)
                 if stems else "")
    if key and key == pairing_key(candidate):
        signals.append(PairSignal("name_key", weights["name_stem"],
                                  f"names share the key {key!r}"))
    elif similarity >= threshold:
        signals.append(PairSignal("name_fuzzy", weights["name_stem"],
                                  f"names overlap {similarity:.2f} ≥ {threshold:g}{stem_text}"))
    elif similarity > 0:
        signals.append(PairSignal("name_overlap", 0.0,
                                  f"names overlap {similarity:.2f} < {threshold:g} — below the "
                                  f"threshold, never paired on it{stem_text}"))
    return signals


def _overlap(a: dict[str, str], b: dict[str, str]) -> list[str]:
    return [f"{k!r} ({a[k]} ↔ {b[k]})" for k in sorted(set(a) & set(b)) if k]


def score_frd(sttm: _Facts, frd: _Facts, weights: dict[str, float]) -> list[PairSignal]:
    signals = []
    for name in frd.names:
        hit = next(((cit, value) for cit, value in sttm.texts if _name_match(name, value)), None)
        if hit:
            signals.append(PairSignal("feed_name", weights["feed_name"],
                                      f"{name!r} ↔ {hit[1]!r} at {hit[0]}"))
            break
    # The one-block / many-files FRD shape lists FILE names under "Target Table
    # Name" (docs/acfc/SHAPES_FOR_PORT.md) — so its tables also meet the files.
    frd_files = {**frd.tables, **frd.files}
    for label, a, b in (("tables", frd.tables, sttm.tables), ("schema", frd.schemas, sttm.schemas),
                        ("file_patterns", frd_files, sttm.files)):
        common = _overlap(a, b)
        if common:
            signals.append(PairSignal(label, weights[label], "; ".join(common[:3])))
    return signals


def score_vdd(sttm: _Facts, vdd: _Facts, weights: dict[str, float]) -> list[PairSignal]:
    signals = []
    for label, a, b in (("file_patterns", vdd.files, sttm.files),
                        ("tables", vdd.tables, sttm.tables)):
        common = _overlap(a, b)
        if common:
            signals.append(PairSignal(label, weights[label], "; ".join(common[:3])))
    for meta_key in ("file_format", "delimiter", "frequency"):
        if meta_key in sttm.meta and meta_key in vdd.meta and (
                normalize(sttm.meta[meta_key][0]) == normalize(vdd.meta[meta_key][0])):
            signals.append(PairSignal(
                f"meta_{meta_key}", weights["meta"],
                f"{sttm.meta[meta_key][0]!r} at {sttm.meta[meta_key][1]} ↔ "
                f"{vdd.meta[meta_key][1]}"))
    return signals


def decide(kind: str, sttm_name: str, scored: list[PairCandidate], config: Config) -> PairDecision:
    knobs = config.inputs.pairing
    threshold = knobs.name_fuzzy_threshold
    ranked = sorted(scored, key=lambda c: (-c.score, -c.name_similarity, c.name))
    top = tuple(ranked[: knobs.top_candidates])
    names = [c.name for c in ranked]
    if not any(s.name not in _NAME_SIGNALS for c in ranked for s in c.signals):
        # No document says anything in common: only the names speak — a unique
        # shared ticket, else the pairing key, else an overlap at the threshold.
        legacy = auto_pair_document(sttm_name, names, threshold=threshold)
        if legacy is not None:
            reason = (match_by_name(sttm_name, names, threshold).reason
                      if legacy[1] in ("name_key", "name_fuzzy")
                      else f"no content signal; names alone: {legacy[1]}")
            return PairDecision(kind, sttm_name, legacy[0], legacy[1], top, reason)
        if any(c.signals for c in ranked):
            # The names say SOMETHING (an overlap below the threshold, two
            # candidates at it, a shared ticket): a question, never a pair.
            return PairDecision(kind, sttm_name, None, None, top,
                                "no content signal and the names do not single one out: "
                                + match_by_name(sttm_name, names, threshold).reason, ask=True)
        return PairDecision(kind, sttm_name, None, None, top, "no candidate shares any content")
    best = ranked[0]
    runner_up = ranked[1].score if len(ranked) > 1 else 0.0
    if best.score >= knobs.min_score and best.score - runner_up >= knobs.margin:
        legacy = auto_pair_document(sttm_name, names, threshold=threshold)
        # The rule names what decided: a name rule when it alone would have
        # picked the same document, "content" when the documents had to speak.
        rule = legacy[1] if legacy is not None and legacy[0] == best.name else "content"
        reason = (f"{best.name} scores {best.score:g}, next {runner_up:g} "
                  f"(margin ≥ {knobs.margin:g})")
        if rule in ("name_key", "name_fuzzy"):
            # The NOTE names the files and the stem-matched words (D1 pair_2's
            # NOTE carried the score only).
            reason += "; " + match_by_name(sttm_name, names, threshold).reason
        return PairDecision(kind, sttm_name, best.name, rule, top, reason)
    return PairDecision(
        kind, sttm_name, None, None, top,
        f"no candidate wins by the margin: best {best.score:g}, next {runner_up:g} "
        f"(need ≥ {knobs.min_score:g} and a lead of {knobs.margin:g})",
        min_score=knobs.min_score)


def _logged(decision: PairDecision) -> PairDecision:
    """Every automatic pair is logged with the rule that made it ("paired by
    override / key / fuzzy / …") and the reason, which names the files."""
    if decision.chosen is not None:
        reason = decision.reason
        prefix = f"paired by {_RULE_WORDS.get(decision.rule or '', decision.rule)}"
        _log.info("%s %s: %s ↔ %s — %s", "NOTE" if decision.rule == "name_fuzzy" else "PAIRED",
                  decision.kind, decision.sttm, decision.chosen,
                  reason if reason.startswith(prefix) else f"{prefix}: {reason}")
    return decision


def pair_by_content(kind: str, sttm_path: Path, candidates: dict[str, Path], config: Config,
                    base_dir: Path, explicit_map: dict[str, str] | None = None,
                    known_facts: dict[str, _Facts] | None = None,
                    overrides: PairOverrides | None = None) -> PairDecision:
    """Decide the ``kind`` ("frd" | "vdd") companion of ``sttm_path`` among
    ``candidates`` (name -> local path). A candidate that cannot be read
    scores on its name alone — one unreadable file never blocks pairing.
    ``known_facts`` (name -> facts already read, M9.3) are used instead of
    opening the document; the STTM's own under its file name. ``overrides``
    (a pairs override file) wins over every rule; an override naming a
    document that is not a candidate is a QUESTION, never a fall-through."""
    sttm_name = sttm_path.name
    names = [n for n in candidates if n != sttm_name]
    by_stem = {document_stem(n): n for n in names}
    missing_override = None
    wanted = overrides.lookup(sttm_name, kind) if overrides is not None else None
    if wanted:
        source = overrides.source(sttm_name)
        if document_stem(wanted) in by_stem:
            chosen = by_stem[document_stem(wanted)]
            return _logged(PairDecision(
                kind, sttm_name, chosen, "override", (),
                f"paired by override: {source} lists {sttm_name!r} ↔ {chosen!r}"))
        missing_override = (f"the override file {source} names {wanted!r} as the "
                            f"{kind.upper()} of {sttm_name!r}, but it is not among the "
                            f"candidates ({', '.join(repr(n) for n in names) or 'none'})")
    explicit = {document_stem(k): document_stem(v) for k, v in (explicit_map or {}).items()}
    mapped = explicit.get(document_stem(sttm_name))
    if mapped and mapped in by_stem and missing_override is None:
        return _logged(PairDecision(kind, sttm_name, by_stem[mapped], "pairing_map", (),
                                    "explicit pairing_map entry"))
    weights = config.inputs.pairing.weights
    threshold = config.inputs.pairing.name_fuzzy_threshold
    known_facts = known_facts or {}
    try:
        sttm = known_facts.get(sttm_name) or sttm_facts(sttm_path, config, base_dir)
    except Exception:  # noqa: BLE001 — an unreadable STTM still pairs by name
        sttm = _Facts()
    scored = []
    for name in names:
        signals: list[PairSignal] = []
        try:
            if kind == "frd":
                signals = score_frd(sttm, known_facts.get(name) or frd_facts(
                    candidates[name], config, base_dir), weights)
            else:
                facts = known_facts.get(name) or vdd_facts(
                    candidates[name], config, base_dir, tuple(sttm.raw_tables))
                signals = score_vdd(sttm, facts, weights)
        except Exception:  # noqa: BLE001, S110 — name signals below still apply
            pass
        signals = [*signals, *_name_signals(sttm_name, name, weights, threshold)]
        scored.append(PairCandidate(name, sum(s.weight for s in signals), tuple(signals),
                                    name_overlap(sttm_name, name)[0]))
    decision = decide(kind, sttm_name, scored, config)
    if missing_override is not None:
        return replace(decision, chosen=None, rule=None, reason=missing_override, ask=True)
    return _logged(decision)


# -- the pairs override file (2026-10-13) ------------------------------------------------
# Never blocked by naming: a person lists the documents that belong together and
# that wins over every automatic rule —
#
#   pairs:
#     - sttm: STTM_<feed>.xlsx
#       frd:  <feed>_FRD.docx
#       vdd:  <feed>_VDD.xlsx        # optional
#
# (JSON with the same shape works too.) Names compare as ``document_stem`` does:
# case-insensitive, upload prefix / copy suffix / extension ignored. Found as
# ``--pairs-file``, ``inputs.pairing.overrides_file`` / CODEGEN_PAIRING_OVERRIDES,
# or a file named below next to the documents (the STTM's folder first).

OVERRIDE_FILE_NAMES = ("pairing_overrides.yaml", "pairing_overrides.yml",
                       "pairing_overrides.json")
_OVERRIDE_KEYS = ("sttm", "frd", "vdd")


class PairOverridesError(ValueError):
    """A pairs override file that cannot be used — named, never ignored."""


@dataclass(frozen=True)
class PairOverrides:
    groups: dict[str, dict[str, str]] = field(default_factory=dict)  # sttm stem -> entry
    sources: dict[str, str] = field(default_factory=dict)            # sttm stem -> file

    def lookup(self, sttm_name: str, kind: str) -> str | None:
        return self.groups.get(document_stem(sttm_name), {}).get(kind)

    def source(self, sttm_name: str) -> str | None:
        return self.sources.get(document_stem(sttm_name))

    def merged(self, other: PairOverrides) -> PairOverrides:
        """Both files' entries; the same STTM listed differently is an error."""
        groups, sources = dict(self.groups), dict(self.sources)
        for stem, entry in other.groups.items():
            if stem in groups and groups[stem] != entry:
                raise PairOverridesError(
                    f"{other.sources[stem]} and {sources[stem]} pair the STTM "
                    f"{entry['sttm']!r} differently")
            groups[stem], sources[stem] = entry, other.sources[stem]
        return PairOverrides(groups, sources)

    def __bool__(self) -> bool:
        return bool(self.groups)


def parse_pair_overrides(text: str, source: str) -> PairOverrides:
    import yaml

    try:
        data = yaml.safe_load(text)
    except yaml.YAMLError as exc:
        raise PairOverridesError(f"{source}: not YAML / JSON: {exc}") from exc
    entries = data.get("pairs") if isinstance(data, dict) else data
    if isinstance(data, dict) and set(data) - {"pairs"}:
        raise PairOverridesError(f"{source}: unknown top-level key(s) "
                                 f"{sorted(set(data) - {'pairs'})} (expected 'pairs')")
    if not isinstance(entries, list):
        raise PairOverridesError(f"{source}: expected 'pairs:' with a list of "
                                 "{sttm, frd, vdd} entries")
    groups: dict[str, dict[str, str]] = {}
    for index, entry in enumerate(entries, start=1):
        where = f"{source} pairs[{index}]"
        if not isinstance(entry, dict):
            raise PairOverridesError(f"{where}: expected a mapping of sttm / frd / vdd")
        unknown = sorted(set(entry) - set(_OVERRIDE_KEYS))
        if unknown:
            raise PairOverridesError(f"{where}: unknown key(s) {unknown} "
                                     f"(allowed: {', '.join(_OVERRIDE_KEYS)})")
        values = {k: str(v).strip() for k, v in entry.items() if v is not None and str(v).strip()}
        if "sttm" not in values:
            raise PairOverridesError(f"{where}: no 'sttm' file name")
        if not {"frd", "vdd"} & set(values):
            raise PairOverridesError(f"{where}: pairs the STTM {values['sttm']!r} with nothing "
                                     "(give 'frd' and / or 'vdd')")
        stem = document_stem(values["sttm"])
        if stem in groups:
            raise PairOverridesError(f"{where}: the STTM {values['sttm']!r} is listed twice")
        groups[stem] = values
    return PairOverrides(groups, {stem: source for stem in groups})


def load_pair_overrides(path: Path) -> PairOverrides:
    try:
        text = Path(path).read_text(encoding="utf-8")
    except OSError as exc:
        raise PairOverridesError(f"{path}: cannot be read: {exc}") from exc
    return parse_pair_overrides(text, Path(path).name)


def configured_overrides_path(config: Config, base_dir: Path) -> Path | None:
    """``CODEGEN_PAIRING_OVERRIDES`` > ``inputs.pairing.overrides_file``
    (relative = against ``base_dir``)."""
    import os

    raw = os.environ.get("CODEGEN_PAIRING_OVERRIDES") or config.inputs.pairing.overrides_file
    if not raw:
        return None
    path = Path(raw)
    return path if path.is_absolute() else base_dir / path


# -- roles of the files in a folder (CLI / notebook fallback) -----------------------------

@dataclass(frozen=True)
class FileRole:
    """One file of a pair folder: its role (frd | sttm | vdd | iig |
    unclassified | override) and the rule that assigned it."""

    name: str
    role: str
    rule: str
    path: Path | None = None

    def as_dict(self) -> dict[str, str]:
        return {"name": self.name, "role": self.role, "rule": self.rule}


def file_role(name: str, content: tuple[str, str] | None = None) -> FileRole:
    """The role of one file: its NAME first (a role token anywhere; ``iig``
    always wins — a validation golden is never an input), then — for a
    workbook whose name carries no role — its CONTENT verdict ``(kind,
    reason)`` from ``classify_workbook`` / the document index: a dictionary
    only on positive structural evidence (a FILES or field sheet by its
    header signature), never "an .xlsx that is not the STTM"."""
    if name.lower() in OVERRIDE_FILE_NAMES:
        return FileRole(name, "override", "the pairs override file")
    verdict = classify_document_name(name)
    if verdict.role == "vdd" and content is not None and content[0] == "sttm":
        return FileRole(name, "sttm", f"{verdict.rule}, but it reads as a mapping workbook: "
                                      f"{content[1]}")
    if verdict.role != "unclassified" or not name.lower().endswith(".xlsx"):
        if verdict.role == "unclassified" and name.lower().endswith(".docx"):
            return FileRole(name, "unclassified",
                            "no frd token in the name (an FRD .docx is recognised by its name)")
        return FileRole(name, verdict.role, verdict.rule)
    if content is None:
        return FileRole(name, "unclassified", f"{verdict.rule}; content not read")
    kind, reason = content
    if kind in ("iig", "iig_review"):
        return FileRole(name, "iig", f"IIG by its sheets: {reason}")
    if kind == "vdd":
        return FileRole(name, "vdd", f"dictionary by content: {reason}")
    if kind == "sttm":
        return FileRole(name, "sttm", f"mapping workbook by content: {reason}")
    return FileRole(name, "unclassified", f"{verdict.rule}; content: {reason}")


def vdd_by_content_note(decision: PairDecision) -> str | None:
    """'VDD by content: <file>' when the chosen dictionary carries no ``vdd``
    token (an override is the person's own choice and says so itself)."""
    if (decision.kind == "vdd" and decision.chosen and decision.rule != "override"
            and classify_document_name(decision.chosen).role != "vdd"):
        return f"VDD by content: {decision.chosen}"
    return None


def classify_folder(paths: list[Path], config: Config) -> list[FileRole]:
    """Every file with its role; workbooks without a role token are read
    (``classify_workbook``) for positive content evidence."""
    from codegen.layout.classify import classify_workbook

    roles = []
    for path in paths:
        content = None
        if (path.suffix.lower() == ".xlsx"
                and classify_document_name(path.name).role in ("unclassified", "vdd")):
            verdict = classify_workbook(path, config.extractor)
            content = (verdict.kind, verdict.reason)
        role = file_role(path.name, content)
        roles.append(replace(role, path=path))
    return roles
