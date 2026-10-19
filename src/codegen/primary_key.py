"""post70: TGT_PRIMARY_KEY — proposed only from an explicit key signal.

A key comes from a signal that says "key", never from the not-null list (a
required column is not a key). The signals, first one that names columns of
the table wins:

1. the STTM's Primary Key / Key cells (``load_rules.natural_key``; per table
   the band's own cells — the header families of docs/acfc/SHAPES_FOR_PORT.md);
2. the FRD's Primary Key / Business Key / Natural Key row
   (``FrdFeed.key_statement``) — the detail table only, names matched to its
   columns;
3. the load-pattern FAQ's ``natural_key_columns`` — the detail table only;
4. a corpus row (the loaded goldens, ``pipeline_name_corpus``) whose
   TGT_TABLE_NAME is the same table — its TGT_PRIMARY_KEY as written.

A proposal is an agent proposal (codegen.input_choices, Aruna's rule): shown on
the ``INPUT_NEEDED feeds.<feed>.primary_key`` line and the cell's NOTE, written
only on an explicit choice (``choice: use_default`` / a typed list / the App's
Confirm) — then flagged ``key_from_contract``. No signal at all and a merge load
(TGT_LOAD_OPTION Upsert / Merge) = ONE blocking ``QUESTION
feeds.<feed>.primary_key`` whose candidates are the not-null columns, none
selected. No signal and no merge = the cell stays blank (no key is needed).

Enabled per IIG template (``metadata.templates.<t>.primary_key_input``)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

KEY = "primary_key"                     # the FAQ companion / input name
HEADER = "TGT_PRIMARY_KEY"
FLAG = "key_from_contract"
QUESTION_NOTE = "iig_key_question"
MERGE_OPTIONS = frozenset({"upsert", "merge"})
_NONE_WORDS = frozenset({"", "none", "n/a", "na", "tbd", "null", "-", "unknown", "nil"})
_SPLIT = re.compile(r"[,;\n]+|\s+and\s+", re.IGNORECASE)


@dataclass(frozen=True)
class KeyProposal:
    """What use_default would write for one table: its columns (the band's
    column names) and the signal that states them; empty = no signal."""

    columns: tuple[str, ...] = ()
    source: str = ""
    kind: str = ""           # sttm | frd | faq | corpus

    def __bool__(self) -> bool:
        return bool(self.columns)

    @property
    def value(self) -> str:
        return ",".join(self.columns)


@dataclass
class TableKey:
    """One IIG row's table: its fields' (stage, standard, source) names."""

    table: str
    layer: str                                  # stage | standard
    fields: list = field(default_factory=list)  # SttmField
    detail: bool = True


def enabled(tpl) -> bool:
    return bool(getattr(tpl, "primary_key_input", False)) and any(
        HEADER in tab.headers for tab in tpl.tabs.values())


def merge_load(strategy) -> bool:
    text = str(getattr(strategy, "value", strategy) or "").strip().lower()
    return text in MERGE_OPTIONS


def merge_legs(spec) -> list[str]:
    """The legs whose load is a merge: ``stage`` (the ADLS row) / ``standard``."""
    legs = []
    if merge_load(spec.stage_load_strategy):
        legs.append("stage")
    if spec.standard_table is not None and merge_load(spec.standard_load_strategy):
        legs.append("standard")
    return legs


def _column(f, layer: str) -> str:
    return (f.stage_column if layer == "stage" else f.standard_column) or ""


def _match(names: list[str], table: TableKey) -> tuple[str, ...]:
    """``names`` (source, stage or standard spellings) as the table's band
    columns; () unless EVERY name is one of its fields."""
    out = []
    for name in names:
        wanted = name.strip().lower()
        hit = next((f for f in table.fields if wanted in {
            (f.source_column or "").lower(), (f.stage_column or "").lower(),
            (f.standard_column or "").lower()}), None)
        if hit is None or not _column(hit, table.layer):
            return ()
        out.append(_column(hit, table.layer))
    return tuple(dict.fromkeys(out))


def statement_names(text: str | None) -> list[str]:
    """The column names a key statement lists (``None`` / ``N/A`` = none)."""
    if not text or text.strip().lower() in _NONE_WORDS:
        return []
    names = [part.strip().strip(".") for part in _SPLIT.split(text)]
    return [n for n in names if n and n.lower() not in _NONE_WORDS]


def corpus_keys(paths: list[str], base: Path | None = None) -> dict[str, str]:
    """TGT_TABLE_NAME (case-folded) -> TGT_PRIMARY_KEY of every corpus row that
    states one (the loaded goldens; a table with two different keys is left
    out — no single statement)."""
    from codegen.input_choices import _corpus_files

    files = [p for p in _corpus_files(paths, base) if p.suffix.lower() == ".xlsx"]
    stamp = tuple((str(p), p.stat().st_mtime_ns, p.stat().st_size) for p in files)
    if stamp not in _CORPUS_CACHE:
        seen: dict[str, set[str]] = {}
        for path in files:
            for table, key in _workbook_keys(path):
                seen.setdefault(table, set()).add(key)
        _CORPUS_CACHE[stamp] = {t: next(iter(k)) for t, k in seen.items() if len(k) == 1}
    return dict(_CORPUS_CACHE[stamp])


_CORPUS_CACHE: dict[tuple, dict[str, str]] = {}


def _workbook_keys(path: Path) -> list[tuple[str, str]]:
    from io import BytesIO

    from openpyxl import load_workbook

    try:
        book = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — an unreadable corpus file adds nothing
        return []
    out = []
    try:
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = [str(h or "").strip().upper() for h in (next(rows, None) or ())]
            if "TGT_TABLE_NAME" not in header or HEADER not in header:
                continue
            t, k = header.index("TGT_TABLE_NAME"), header.index(HEADER)
            for row in rows:
                table = str(row[t] or "").strip() if t < len(row) else ""
                key = str(row[k] or "").strip() if k < len(row) else ""
                if table and key.lower() not in _NONE_WORDS:
                    out.append((table.split(".")[-1].casefold(), key))
    finally:
        book.close()
    return out


def proposal(table: TableKey, spec, faq=None,
             corpus: dict[str, str] | None = None) -> KeyProposal:
    """The table's key proposal from the explicit signals, in order."""
    from codegen.metadata_template import _primary_key

    marked = _primary_key(table.fields, table.layer)
    if marked:
        band = "Stage" if table.layer == "stage" else "Standard"
        return KeyProposal(tuple(marked), f"the STTM's Primary Key cells ({band} band)", "sttm")
    if table.detail:
        statement = getattr(spec, "frd_key_statement", None)
        columns = _match(statement_names(statement), table)
        if columns:
            return KeyProposal(columns, f"the FRD's key statement {statement!r}", "frd")
        declared = list(getattr(faq, "natural_key_columns", None) or [])
        columns = _match(declared, table)
        if columns:
            return KeyProposal(columns, "the load-pattern FAQ natural_key_columns", "faq")
    stated = (corpus or {}).get(table.table.split(".")[-1].casefold())
    if stated:
        return KeyProposal(tuple(statement_names(stated)) or (stated,),
                           f"a corpus row for table {table.table!r}", "corpus")
    return KeyProposal()


def candidates(spec, layer: str = "stage") -> list[str]:
    """The not-null columns shown — none selected — when no signal exists."""
    if layer == "stage":
        return list(spec.not_null_columns)
    detail = {f.stage_column: f.standard_column for f in spec.detail_segment.fields}
    return [detail[c] for c in spec.not_null_columns if detail.get(c)]


def detail_table(spec, layer: str = "stage") -> TableKey:
    """The detail table (the feed's natural key lives there)."""
    seg = spec.detail_segment
    name = (seg.stage_table.table if layer == "stage" else
            (seg.standard_table.table if getattr(seg, "standard_table", None) is not None
             else (spec.standard_table.table if spec.standard_table is not None else "")))
    return TableKey(name, layer, list(seg.fields), True)


def feed_proposal(spec, faq=None, corpus=None) -> KeyProposal:
    """The feed's proposal: the detail table's, stage columns (the input's value)."""
    return proposal(detail_table(spec, "stage"), spec, faq, corpus)


def answered(faq) -> tuple[str, str, str] | None:
    """(how, value, template) of the feed's primary_key answer: ``default`` /
    ``standard`` (an explicit choice), ``blank`` (leave_blank), ``answer`` (a
    typed list); None = unanswered."""
    from codegen.input_choices import chosen_of

    answer = getattr(faq, KEY, None) if faq is not None else None
    if answer is None or answer.source == "unknown":
        return None
    chosen = chosen_of(answer)
    if chosen is not None:
        return chosen[0], str(answer.value or "").strip(), chosen[1]
    value = str(answer.value or "").strip()
    return ("answer", value, "") if value else None


def key_cell(tab: str, table: TableKey, spec, faq, offer: KeyProposal, *,
             merge: bool, load_option: str = "") -> dict | None:
    """The row's TGT_PRIMARY_KEY cell (None: nothing to say — the caller's
    own blank / family convention stands)."""
    from codegen.metadata_template import _cell

    where = f"{tab}.{HEADER}"
    how = answered(faq)
    if how is not None:
        kind, value, template = how
        if kind == "blank":
            cell = _cell("", "needs_template", "left blank by choice (leave_blank)")
            cell["badge_entry"]["note"] = (f"iig_owed_blank:{where} — left blank: the BSA chose "
                                           f"leave_blank for {KEY}")
            return cell
        if kind == "default":
            if not offer:
                return None
            cell = _cell(offer.value, {"sttm": "from_sttm", "frd": "from_frd",
                                       "faq": "from_faq"}.get(offer.kind, "synthetic"),
                         f"proposed from {offer.source}; confirmed (use_default)")
            cell["badge_entry"]["note"] = (f"{FLAG}:{where} source_used:{offer.source}: "
                                           f"{offer.value!r} (chosen: use_default)")
            return cell
        if not table.detail:
            return None                 # a typed key is the detail table's
        names = statement_names(value)
        columns = _match(names, table) or tuple(names)
        given = f"naming_standard {template!r}" if template else "answers file"
        return _cell(",".join(columns), "from_faq", f"the BSA's key ({given})")
    if offer:
        shown = (f"proposed {offer.value!r} from {offer.source}: not written until chosen")
        cell = _cell("", "needs_template", f"owed: the table's key — {shown}")
        cell["badge_entry"]["note"] = (
            f"iig_owed:{where} — the table's key (TGT_PRIMARY_KEY) — {shown} (INPUT_NEEDED "
            f"feeds.<feed>.{KEY}: choose use_default / leave_blank or type the columns under "
            f"inputs.feeds.<feed>)")
        return cell
    if merge and table.detail:
        listed = ", ".join(candidates(spec, table.layer)) or "none"
        cell = _cell("", "needs_template",
                     f"open — a {load_option or 'merge'} load needs a key and no input states "
                     f"one; candidates (not selected): {listed}")
        cell["badge_entry"]["note"] = (
            f"{QUESTION_NOTE}:{where} — {load_option or 'merge'} load, no key signal (no "
            f"Primary Key cell, FRD key statement, FAQ natural_key_columns or corpus row); "
            f"the not-null columns are candidates, none selected: {listed} (QUESTION "
            f"feeds.<feed>.{KEY})")
        return cell
    return None


def question(spec, faq, corpus=None) -> tuple[str, list[str]] | None:
    """(reason, candidates) of the feed's blocking key QUESTION — a merge leg,
    no signal for the detail table, unanswered; None otherwise."""
    legs = merge_legs(spec)
    if not legs or answered(faq) is not None:
        return None
    if any(proposal(detail_table(spec, leg), spec, faq, corpus) for leg in legs):
        return None
    listed = candidates(spec, "stage")
    reason = (f"the {' and '.join(legs)} load is a merge and no input states the table's key "
              "(no Primary Key cell, FRD key statement, FAQ natural_key_columns or corpus "
              "row) — type the key columns; the not-null columns are candidates, none "
              f"selected: {', '.join(listed) or 'none'}")
    return reason, listed
