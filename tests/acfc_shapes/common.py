# ruff: noqa: E501  -- fixture data tables: documented header/label texts kept on one line
"""Shared helpers for the ACFC-shape fixture builders.

Everything here is synthetic vocabulary only. Two output formats:

- ``.xlsx`` via openpyxl, serialized through the repo's byte-stable
  ``stable_workbook_bytes`` (pinned timestamps, pinned zip member dates) so
  the tracked fixture and a rebuild compare byte for byte.
- ``.docx`` via a minimal stdlib OOXML writer (zipfile + literal XML): the
  FRD extractor (M2) reads with stdlib zipfile+XML, so the fixtures are
  written the same way — no python-docx dependency on either side.
"""

from __future__ import annotations

import io
import zipfile
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from codegen.metadata_sheet import _pin_workbook_properties, stable_workbook_bytes

# Fixed zip member timestamp for the docx writer (byte-stable output).
_PINNED = (2026, 1, 1, 0, 0, 0)


# ------------------------------------------------------------------ xlsx


def new_workbook() -> Workbook:
    workbook = Workbook()
    _pin_workbook_properties(workbook)
    workbook.remove(workbook.active)
    return workbook


def write_rows(ws, rows: list[list], start_row: int = 1) -> None:
    """Write rows at ``start_row`` (1-based); ``None`` cells stay empty."""
    for offset, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(row=start_row + offset, column=col, value=value)


def merge_span(ws, row: int, start_col: int, end_col: int) -> None:
    """Merge a horizontal band ``start_col..end_col`` (1-based, inclusive)."""
    if end_col > start_col:
        ws.merge_cells(
            start_row=row, start_column=start_col, end_row=row, end_column=end_col
        )


def band_row(groups: list[tuple[str, int]]) -> list:
    """A band-label row from ``[(label, width), ...]``: the label sits in the
    group's first cell, the rest of the group is blank (merged by caller)."""
    row: list = []
    for label, width in groups:
        row.extend([label] + [None] * (width - 1))
    return row


def merge_bands(ws, row: int, groups: list[tuple[str, int]]) -> None:
    col = 1
    for _label, width in groups:
        merge_span(ws, row, col, col + width - 1)
        col += width


def xlsx_bytes(workbook: Workbook) -> bytes:
    return stable_workbook_bytes(workbook)


def col(letter_index: int) -> str:
    return get_column_letter(letter_index)


# ------------------------------------------------------------------ docx

_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
    "</Types>"
)
_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
    "</Relationships>"
)
_DOC_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
    "</Relationships>"
)
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_STYLES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<w:styles xmlns:w="{_W_NS}">'
    '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading4"><w:name w:val="heading 4"/><w:basedOn w:val="Normal"/></w:style>'
    '<w:style w:type="table" w:styleId="TableGrid"><w:name w:val="Table Grid"/></w:style>'
    "</w:styles>"
)


def _runs(text: str) -> str:
    """Text → runs; an embedded newline becomes a ``<w:br/>`` (the way Word
    stores a Shift+Enter inside a table cell, e.g. "Business \\nRequirement").
    A cell that Word split into several runs is modelled by splitting on
    ``|`` — the reader must join runs with no separator."""
    parts = text.split("\n")
    out: list[str] = []
    for index, part in enumerate(parts):
        if index:
            out.append("<w:r><w:br/></w:r>")
        for run in part.split("|"):
            if run:
                out.append(f'<w:r><w:t xml:space="preserve">{escape(run)}</w:t></w:r>')
    return "".join(out)


def paragraph(text: str, style: str | None = None) -> str:
    props = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{props}{_runs(text)}</w:p>"


def table(rows: list[list]) -> str:
    """A Word table; every row must have the same cell count. An empty
    string is an empty (label-present, value-blank) cell. A cell whose value
    is itself a list of rows renders as a NESTED table inside the cell (the
    way Word stores a table pasted into a metadata cell) followed by the
    empty paragraph Word requires after a nested table."""
    width = max(sum(_span(c) for c in r) for r in rows)
    grid = "".join('<w:gridCol w:w="2400"/>' for _ in range(width))
    body = []
    for row in rows:
        cells = "".join(_tc(value) for value in row)
        body.append(f"<w:tr>{cells}</w:tr>")
    return (
        '<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
        f"<w:tblGrid>{grid}</w:tblGrid>{''.join(body)}</w:tbl>"
    )


def merged(text: str = "", span: int = 1, vmerge: str | None = None) -> dict:
    """A MERGED table cell (v0.7.1, SHAPES_ROUND2 §1): ``span`` grid columns
    wide (``<w:gridSpan>``), and/or part of a vertical merge (``<w:vMerge>``:
    "restart" opens it, "continue" is a covered cell). Word stores such a cell
    as ONE ``<w:tc>`` — which is exactly what the reader must cope with."""
    return {"text": text, "span": span, "vmerge": vmerge}


def _span(value) -> int:
    return value.get("span", 1) if isinstance(value, dict) else 1


def _tc(value) -> str:
    if isinstance(value, list):
        return f"<w:tc>{table(value)}<w:p/></w:tc>"
    if isinstance(value, dict):
        props = ""
        if value.get("span", 1) > 1:
            props += f'<w:gridSpan w:val="{value["span"]}"/>'
        if value.get("vmerge") == "restart":
            props += '<w:vMerge w:val="restart"/>'
        elif value.get("vmerge") == "continue":
            props += "<w:vMerge/>"
        return f"<w:tc><w:tcPr>{props}</w:tcPr>{paragraph(value.get('text', ''))}</w:tc>"
    return f"<w:tc>{paragraph(value)}</w:tc>"


def docx_bytes(body_parts: list[str]) -> bytes:
    """Assemble a .docx from body XML fragments (paragraph()/table())."""
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{_W_NS}"><w:body>{"".join(body_parts)}'
        "<w:sectPr/></w:body></w:document>"
    )
    members = [
        ("[Content_Types].xml", _CONTENT_TYPES),
        ("_rels/.rels", _RELS),
        ("word/_rels/document.xml.rels", _DOC_RELS),
        ("word/styles.xml", _STYLES),
        ("word/document.xml", document),
    ]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, text in members:
            info = zipfile.ZipInfo(name, date_time=_PINNED)
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, text.encode("utf-8"))
    return out.getvalue()
