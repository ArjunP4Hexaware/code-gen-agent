"""Pair 1 — the ONE feed whose fixtures describe a real (scrubbed) golden.

The golden's column names, types, segment layout (fixed-width start/len)
and audit columns are STRUCTURE (docs/acfc/rfc_capture/.../goldens/pair_1).
This module is the single source of truth for that structure: the pair-1
STTM, FRD and VDD fixtures are all rendered from ``COLUMNS`` / ``SEGMENTS``
below, so the three documents can never disagree with each other or with
the aliased golden DDL.

Alias map (fixtures/acfc_shapes/ALIASES.md): applied to the fixtures AND to
the tracked copy of the golden, case-preserving (``prx``→``vnd_p``,
``PRX``→``VND_P``; ``acfc``→``client``, ``ACFC``→``CLIENT``), plus the two
source-system names. The one storage-account VALUE the scrub scanner flagged
in the "scrubbed" golden is replaced by COLUMN (``STORAGE_ACCOUNT_COLUMNS``)
so the raw value is never spelled in a tracked file.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

# Ordered, case-sensitive replacements. Longer/more specific first.
ALIASES: list[tuple[str, str]] = [
    ("Abarca", "VENDOR_A"),
    ("Facets", "SRC_SYS_A"),
    ("PRX", "VND_P"),
    ("Prx", "Vnd_P"),
    ("prx", "vnd_p"),
    ("ACFC", "CLIENT"),
    ("Acfc", "Client"),
    ("acfc", "client"),
]


# IIG columns whose every value is a storage-account name: replaced wholesale
# (the golden carries a real one the upstream scrub missed).
STORAGE_ACCOUNT_COLUMNS = {"TGT_STORAGE_ACCOUNT_NAME"}
STORAGE_ACCOUNT_ALIAS = "syn-storage-001"


def alias(text: str) -> str:
    for raw, replacement in ALIASES:
        text = text.replace(raw, replacement)
    return text


@dataclass(frozen=True)
class Column:
    name: str
    dtype: str          # exactly as the golden DDL spells it
    segment: str        # golden segment vocabulary: HDDR / DET / TRLR
    start: int          # fixed-width start (1-based), from the golden handler
    length: int
    ordinal: int        # position within the segment (1-based)
    required: str = "N"
    load_rule: str = "Load as is"
    description: str = ""


# Transcribed from the golden: ACCUM_DDL.txt (names/types, order) and the
# ADLS_FIXED_WIDTH_HANDLER rows (segment membership, LEN, start_ind).
SEGMENTS: dict[str, list[tuple[str, int, int]]] = {
    "HDDR": [
        ("SEGMENT_IDENTIFIER", 1, 2), ("FILE_TYPE", 3, 1), ("DATA_CATEGORY", 4, 1),
        ("PROCESS_DATE", 5, 8), ("PROCESS_TIME", 13, 6), ("GROUP_ID", 19, 10),
        ("FILLER_1", 29, 20), ("VERSION_RELEASE", 49, 2),
    ],
    "DET": [
        ("SEGMENT_IDENTIFIER", 1, 2), ("CARDHOLDER_ID", 3, 15), ("PLAN_ID", 18, 16),
        ("INDIVIDUAL_DEDUCTIBLE", 34, 17), ("FAMILY_DEDUCTIBLE", 51, 18),
        ("COPAY", 69, 19), ("COINSURANCE", 88, 20), ("INDIVIDUAL_LIMIT", 108, 21),
        ("FAMILY_LIMIT", 129, 22), ("PLAN_YEAR_START_DATE", 151, 23),
        ("PLAN_YEAR_END_DATE", 174, 24), ("FILLER_2", 198, 25),
    ],
    "TRLR": [
        ("SEGMENT_IDENTIFIER", 1, 2), ("LINE_COUNT", 3, 10), ("FILLER_3", 13, 20),
    ],
}

# Golden DDL types, in golden column order (21 business columns).
DDL_TYPES: dict[str, str] = {
    "SEGMENT_IDENTIFIER": "String", "FILE_TYPE": "String", "DATA_CATEGORY": "String",
    "PROCESS_DATE": "Date", "PROCESS_TIME": "String", "GROUP_ID": "String",
    "FILLER_1": "String", "VERSION_RELEASE": "String", "CARDHOLDER_ID": "String",
    "PLAN_ID": "String", "INDIVIDUAL_DEDUCTIBLE": "Decimal(17,2)",
    "FAMILY_DEDUCTIBLE": "Decimal(18,2)", "COPAY": "Decimal(19,2)",
    "COINSURANCE": "Decimal(20,2)", "INDIVIDUAL_LIMIT": "Decimal(21,2)",
    "FAMILY_LIMIT": "Decimal(22,2)", "PLAN_YEAR_START_DATE": "Date",
    "PLAN_YEAR_END_DATE": "Date", "FILLER_2": "String", "LINE_COUNT": "String",
    "FILLER_3": "String",
}
DDL_ORDER = list(DDL_TYPES)

# Audit columns exactly as the golden DDL spells them (uppercase types).
AUDIT_COLUMNS: list[tuple[str, str]] = [
    ("SRC_FILE_NAME", "STRING"),
    ("REC_CREATION_TIME", "TIMESTAMP"),
    ("REC_UPDATED_TIME", "TIMESTAMP"),
]

_DESCRIPTIONS = {
    "SEGMENT_IDENTIFIER": "Record segment marker: 00 header, 10 detail, 99 trailer",
    "FILE_TYPE": "File type indicator",
    "DATA_CATEGORY": "Data category indicator",
    "PROCESS_DATE": "File process date (yyyyMMdd)",
    "PROCESS_TIME": "File process time (HHmmss)",
    "GROUP_ID": "Group identifier",
    "FILLER_1": "Filler",
    "VERSION_RELEASE": "Layout version/release",
    "CARDHOLDER_ID": "Cardholder identifier",
    "PLAN_ID": "Plan identifier",
    "INDIVIDUAL_DEDUCTIBLE": "Individual deductible accumulated amount",
    "FAMILY_DEDUCTIBLE": "Family deductible accumulated amount",
    "COPAY": "Copay accumulated amount",
    "COINSURANCE": "Coinsurance accumulated amount",
    "INDIVIDUAL_LIMIT": "Individual out-of-pocket limit",
    "FAMILY_LIMIT": "Family out-of-pocket limit",
    "PLAN_YEAR_START_DATE": "Plan year start date (yyyyMMdd)",
    "PLAN_YEAR_END_DATE": "Plan year end date (yyyyMMdd)",
    "FILLER_2": "Filler",
    "LINE_COUNT": "Number of detail records in the file",
    "FILLER_3": "Filler",
}
_DATE_RULE = "Convert yyyyMMdd to yyyy-MM-dd"
_REQUIRED = {"SEGMENT_IDENTIFIER", "CARDHOLDER_ID", "PLAN_ID", "PROCESS_DATE", "LINE_COUNT"}


def columns() -> list[Column]:
    """Every (segment, field) row of the pair-1 layout, segment order HDDR,
    DET, TRLR — 23 rows over 21 distinct columns (SEGMENT_IDENTIFIER leads
    every segment and lands in the ONE stage table's single column)."""
    out: list[Column] = []
    for segment, fields in SEGMENTS.items():
        for ordinal, (name, start, length) in enumerate(fields, start=1):
            out.append(Column(
                name=name,
                dtype=DDL_TYPES[name],
                segment=segment,
                start=start,
                length=length,
                ordinal=ordinal,
                required="Y" if name in _REQUIRED else "N",
                load_rule=_DATE_RULE if DDL_TYPES[name] == "Date" else "Load as is",
                description=_DESCRIPTIONS[name],
            ))
    return out


# Feed-level facts (aliased golden vocabulary + the FRD fixture's values).
STAGE_CATALOG = alias("pr_dlk_prx")
STAGE_SCHEMA = alias("stg_prx_accum")
STANDARD_CATALOG = alias("pr_std_prx")
STANDARD_SCHEMA = "accum"
TABLE = alias("prx_accum_acfc")
FEED_NAME = "Accumulators"
PROCESS_NAME = alias("PRX_ACCUMULATOR")
VENDOR_NAME = "VENDOR_A"
VENDOR_ABBREVIATION = alias("PRX")
DOMAIN = "Pharmacy"
SUB_DOMAIN = "Accumulators"
FILE_PATTERNS = [alias(p) for p in (
    "I_ACCUM_*_TO_ACFC_*.csv", "I_ACCUM_*_FROM_ACFC_*.csv",
    "F_ACCUM_*_TO_ACFC_*.csv", "F_ACCUM_*_FROM_ACFC_*.csv",
)]
FREQUENCY = "Daily"
LOB = "ALL"
FILE_FORMAT = "Fixed Width Text"
SOURCE_TYPE = "String"   # golden SRC_DATA_TYPE: every source column is String

# Golden segment vocabulary <-> VDD segment vocabulary (SHAPES_FOR_PORT §3:
# pair-1 VDD identifies HDR/DTL/TRL via a Segment column).
VDD_SEGMENT = {"HDDR": "HDR", "DET": "DTL", "TRLR": "TRL"}
SEGMENT_FILTER = {"HDDR": "value like '00%'", "DET": "value like '10%'", "TRLR": "value like '99%'"}


def alias_golden_ddl(text: str) -> str:
    return alias(text)


def alias_golden_iig(source: Path):
    """The golden IIG with the alias map applied to every string cell,
    re-serialized byte-stably. Sheet order, headers, row counts untouched."""
    from openpyxl import load_workbook

    from .common import xlsx_bytes

    workbook = load_workbook(source)
    for ws in workbook.worksheets:
        header = [c.value for c in ws[1]]
        storage_cols = {i + 1 for i, h in enumerate(header) if h in STORAGE_ACCOUNT_COLUMNS}
        for row in ws.iter_rows():
            for cell in row:
                if cell.row > 1 and cell.column in storage_cols and cell.value is not None:
                    cell.value = STORAGE_ACCOUNT_ALIAS
                elif isinstance(cell.value, str):
                    cell.value = alias(cell.value)
    from codegen.metadata_sheet import _pin_workbook_properties

    _pin_workbook_properties(workbook)
    return xlsx_bytes(workbook)


_RAW_TOKENS = re.compile("|".join(re.escape(raw) for raw, _ in ALIASES))


def has_raw_token(text: str) -> bool:
    """True when any pre-alias token survives in ``text``."""
    return _RAW_TOKENS.search(text) is not None
