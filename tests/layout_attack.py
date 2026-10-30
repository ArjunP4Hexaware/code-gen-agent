"""Test support for the layout red-team hardening (P1-1 / P1-2 / P1-3 / P1-7,
2026-10-30): synthetic copies of the pair-4 STTM and the f1 pair-1 FRD with
attacker-chosen header / label text, and a stub layout provider that returns
attacker-chosen answers. Synthetic fixtures only; nothing here calls a model.
"""

from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
PAIR4_STTM = REPO / "fixtures" / "acfc_shapes" / "sttm" / "pair_4_family_d.xlsx"
F1_FRD = REPO / "fixtures" / "acfc_shapes" / "frd" / "f1_pair_1.docx"
PAIR4_SHEET = "STTM"
PAIR4_HEADER_ROW = 5
# pair 4's header row: stage band 14-19, standard band 20-25
STAGE_CATALOG, STAGE_SCHEMA, STAGE_TABLE = 15, 16, 17
STANDARD_CATALOG, STANDARD_SCHEMA = 21, 22
SOURCE_LEN = 6


def pair4_copy(dst: Path, headers: dict[int, str] | None = None,
               cells: dict[tuple[int, int], object] | None = None,
               column_values: dict[int, object] | None = None) -> Path:
    """A copy of the pair-4 STTM: ``headers`` {column: text} on the header row,
    ``cells`` {(row, column): value}, ``column_values`` {column: value} on every
    data row under the header."""
    shutil.copy(PAIR4_STTM, dst)
    wb = load_workbook(dst)
    ws = wb[PAIR4_SHEET]
    for col, value in (headers or {}).items():
        ws.cell(PAIR4_HEADER_ROW, col, value)
    for (row, col), value in (cells or {}).items():
        ws.cell(row, col, value)
    for col, value in (column_values or {}).items():
        for row in range(PAIR4_HEADER_ROW + 1, ws.max_row + 1):
            if ws.cell(row, 2).value is not None:
                ws.cell(row, col, value)
    wb.save(dst)
    return dst


def frd_with_label(dst: Path, old_label: str, new_label: str, src: Path = F1_FRD) -> Path:
    """A copy of an FRD .docx with one label cell's text replaced (the label
    must be a single run, written once)."""
    with zipfile.ZipFile(src) as zin:
        items = {name: zin.read(name) for name in zin.namelist()}
    xml = items["word/document.xml"].decode("utf-8")
    needle = f">{old_label}<"
    assert xml.count(needle) == 1, (old_label, xml.count(needle))
    items["word/document.xml"] = xml.replace(needle, f">{new_label}<").encode("utf-8")
    with zipfile.ZipFile(dst, "w") as zout:
        for name, data in items.items():
            zout.writestr(name, data)
    return dst


class StubProvider:
    """A compromised / injected layout model: ``answer(request) -> dict`` is
    the attacker's answer (default: the partial profile back, unchanged).
    Records every request."""

    name = "stub"

    def __init__(self, answer=None) -> None:
        self.requests: list[dict] = []
        self._answer = answer

    def complete_layout(self, request: dict) -> dict:
        self.requests.append(request)
        partial = json.loads(json.dumps(request["partial_profile"]))
        return self._answer(request, partial) if self._answer else partial

    def advise_layout(self, request: dict) -> dict:
        return {"advice": []}


def place_sttm(sheet: str, roles: dict[str, dict[str, int]]):
    """An STTM answer: ``roles`` {layer: {role: column}} set on ``sheet``'s
    bands of the partial profile."""
    def answer(_request, partial):
        for sp in partial["sheets"]:
            if sp["name"] != sheet:
                continue
            for band in sp["bands"]:
                band["roles"].update(roles.get(band["layer"], {}))
        return partial
    return answer


def place_frd(fields: dict[str, dict]):
    """An FRD answer: ``fields`` {path: field source} merged into the partial."""
    def answer(_request, partial):
        partial.setdefault("fields", {}).update(fields)
        return partial
    return answer


def with_layout(config, **update):
    return config.model_copy(update={"layout": config.layout.model_copy(update=update)})


__all__ = ["F1_FRD", "PAIR4_HEADER_ROW", "PAIR4_SHEET", "PAIR4_STTM", "SOURCE_LEN",
           "STAGE_CATALOG", "STAGE_SCHEMA", "STAGE_TABLE", "STANDARD_CATALOG",
           "STANDARD_SCHEMA", "StubProvider", "frd_with_label", "pair4_copy", "place_frd",
           "place_sttm", "with_layout"]
