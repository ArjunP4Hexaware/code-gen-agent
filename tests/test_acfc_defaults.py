"""Chunk D (2026-10-09): acfc_prx + iig_v2 are the DEFAULTS wherever the
committed ACFC environment overlay applies (config/overlays/acfc_env.yaml —
load_config applies it first for the App and EVERY CLI run).

* config loaded with the overlay -> conventions.profile acfc_prx,
  metadata.template iig_v2; without it (CODEGEN_SKIP_ENV_OVERLAY=1, the suite)
  -> the shipped edo_sfmc / iig_v1;
* a plain ``codegen generate`` (no --profile / --iig-template) under the
  overlay writes the acfc_prx / iig_v2 artefacts (the one combined DDL file,
  the 8-sheet IIG, metadata_inserts.sql);
* every CLI command's first line and ``codegen doctor`` name the active
  profile and template, so a surprised local run is explained in line one;
* the App's selectors offer the overlay's values as their defaults.
"""

from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from codegen.build_info import source_line
from codegen.config import DEFAULT_OVERLAY, active_defaults, load_config

REPO = Path(__file__).resolve().parents[1]
SRC = REPO / "src"
CONFIG = REPO / "config" / "config.yaml"
OVERLAY = REPO / "config" / DEFAULT_OVERLAY
OVERLAY_SHOWN = str(Path("config") / "overlays" / "acfc_env.yaml")


def _overlay_on(monkeypatch) -> None:
    monkeypatch.delenv("CODEGEN_SKIP_ENV_OVERLAY", raising=False)
    monkeypatch.delenv("CODEGEN_CONFIG_OVERLAYS", raising=False)


def test_the_overlay_makes_acfc_prx_and_iig_v2_the_defaults(monkeypatch):
    _overlay_on(monkeypatch)
    config = load_config(CONFIG)
    assert (config.conventions.profile, config.metadata.template) == ("acfc_prx", "iig_v2")
    assert active_defaults(CONFIG) == {"profile": "acfc_prx", "iig_template": "iig_v2"}
    monkeypatch.setenv("CODEGEN_SKIP_ENV_OVERLAY", "1")             # the suite's setting
    shipped = load_config(CONFIG)
    assert (shipped.conventions.profile, shipped.metadata.template) == ("edo_sfmc", "iig_v1")
    assert active_defaults(CONFIG) == {"profile": "edo_sfmc", "iig_template": "iig_v1"}
    explicit = load_config(CONFIG, overlays=[OVERLAY])
    assert (explicit.conventions.profile, explicit.metadata.template) == ("acfc_prx", "iig_v2")


def test_a_flag_is_marked_on_the_first_line():
    line = source_line(["a.yaml"], {"profile": "edo_sfmc", "profile_from_flag": "1",
                                    "iig_template": "iig_v2"})
    assert line.endswith(" · profile: edo_sfmc (--profile) · IIG template: iig_v2 · "
                         "overlays: a.yaml")


def _env(**extra: str) -> dict[str, str]:
    env = {**os.environ, "PYTHONUTF8": "1",
           "PYTHONPATH": os.pathsep.join([str(SRC), os.environ.get("PYTHONPATH", "")]),
           "CODEGEN_NOTIFICATION_EMAILS": "syn.dl.prodsupport@synthetic.example"}
    env.pop("CODEGEN_SKIP_ENV_OVERLAY", None)
    env.pop("CODEGEN_CONFIG_OVERLAYS", None)
    env.update(extra)
    return env


def _cli(*args: str, env: dict[str, str]) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, "-m", "codegen.cli", *args], cwd=REPO, env=env,
                          text=True, capture_output=True, encoding="utf-8", check=False)


def test_doctor_names_the_active_profile_and_template():
    done = _cli("doctor", env=_env())
    assert done.returncode == 0, done.stderr[-2000:]
    first, profile, template = done.stdout.splitlines()[:3]
    assert first.endswith(f" · profile: acfc_prx · IIG template: iig_v2 · overlays: "
                          f"{OVERLAY_SHOWN}")
    assert profile.startswith("conventions profile: acfc_prx (conventions.profile")
    assert template.startswith("IIG template: iig_v2 (metadata.template")
    skipped = _cli("doctor", env=_env(CODEGEN_SKIP_ENV_OVERLAY="1"))
    assert skipped.stdout.splitlines()[0].endswith(
        " · profile: edo_sfmc · IIG template: iig_v1 · overlays: (none)")


def test_a_plain_generate_under_the_overlay_writes_the_acfc_prx_iig_v2_artefacts(tmp_path):
    """Pair 4 (synthetic) through the CLI with the committed overlay and no
    --profile / --iig-template: the first line says why the artefacts are
    acfc_prx / iig_v2 ones."""
    from acfc_shapes import FIXTURE_ROOT, pair4_nb
    from codegen.extract import contract_to_json, extract_contract

    config = load_config(CONFIG, overlays=[OVERLAY])
    frd = FIXTURE_ROOT / pair4_nb.FRD_PATH
    sttm = tmp_path / "sttm.contract.json"
    sttm.write_text(contract_to_json(extract_contract(
        FIXTURE_ROOT / pair4_nb.STTM_PATH, frd, config, generated_date="2026-10-09")),
        encoding="utf-8")
    scope = tmp_path / "scope.yaml"            # outputs into tmp, on top of the default
    scope.write_text(f"output: {{dir: '{(tmp_path / 'out').as_posix()}', reports_dir: "
                     f"'{(tmp_path / 'reports').as_posix()}'}}\n", encoding="utf-8")
    done = _cli("generate", "--frd-contract", str(frd), "--sttm-contract", str(sttm),
                "--output-mode", "framework", "--dry-run", "--skip-tests",
                env=_env(CODEGEN_CONFIG_OVERLAYS=str(scope)))
    assert done.returncode == 0, (done.stdout[-3000:], done.stderr[-2000:])
    first = done.stdout.splitlines()[0]
    assert " · profile: acfc_prx · IIG template: iig_v2 · overlays: " in first
    assert first.endswith(f"overlays: {OVERLAY_SHOWN} -> {scope}")
    assert "(--profile)" not in first and "(--iig-template)" not in first
    framework = tmp_path / "out" / pair4_nb.FEED / "framework"
    names = sorted(p.name for p in framework.iterdir())
    # acfc_prx: ONE combined DDL file (edo_sfmc writes two *_table_creation.txt files)
    assert "NB_COB_REPORT_DDL.txt" in names
    assert not [n for n in names if n.endswith("_table_creation.txt")]
    # iig_v2: the 8-sheet IIG (iig_v1 is the 7-tab SFMC layout)
    clean = load_workbook(framework / f"{pair4_nb.FEED}_IIG.xlsx")
    assert clean.sheetnames == list(config.metadata.templates["iig_v2"].tabs)
    sql = (framework / "metadata_inserts.sql").read_text(encoding="utf-8")
    assert "INSERT INTO [dbo].[stg_delta_stddelta_ingestion_details] (" in sql


def test_the_app_selectors_default_to_the_overlay(monkeypatch):
    pytest.importorskip("fastapi")
    from ui.backend.demo import DemoRunner

    _overlay_on(monkeypatch)
    runner = SimpleNamespace(_store=SimpleNamespace(config=load_config(CONFIG)),
                             conventions_profile=None, iig_template=None,
                             playbook_template=None)
    options = DemoRunner.generation_options(runner)
    assert options["conventions_profile"]["default"] == "acfc_prx"
    assert options["iig_template"]["default"] == "iig_v2"
    assert "acfc_prx" in options["conventions_profile"]["options"]
    assert "iig_v2" in options["iig_template"]["options"]
