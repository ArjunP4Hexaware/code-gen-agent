"""The first ACFC run on real documents (feature/multi-table) — four defects,
each reproduced here on SYNTHETIC documents (no client data).

1. Key lookup by (segment, field name): a fixed-width STTM whose header,
   detail and trailer all carry a field named ``REC_TYPE`` mapped to three
   different target columns made the resolver map the Detail key to the
   TRAILER's column (one combined lookup, last match wins) and
   ``emit/context.py`` raise ``KeyError: 'REC_TYP_TRLR'``.
2. STTM robustness: (a) FILE_DETAILS annotation rows ("DataType = …", note
   rows, italic guidance) are skipped and reported, never read as files;
   (b) blank Schema / TableName cells take the File Details target / the
   FRD / the layer convention (flag ``sttm_target_missing``), else the feed
   is NEEDS_ANSWERS with the exact answers-file key — the other feeds run.
3. Derived landing / target paths are normalised (trailing punctuation,
   doubled separators, the family's case convention) and flagged
   ``path_normalised`` with before / after; the derivation gate does not
   FAIL on the input's punctuation.
4. ``ruff --fix`` (safe fixes only) runs before the ruff gate; what it fixed
   is the ``ruff_fixed`` flag, so cosmetic lint never FAILs a feed.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl.styles import Font

from acfc_shapes import sttm as sttm_fixtures
from acfc_shapes.common import xlsx_bytes
from codegen import cli
from codegen.contracts.resolved import ResolvedTable, SegmentSpec
from codegen.contracts.sttm import SttmField
from codegen.emit.context import build_context
from codegen.extract import NeedsAnswersError, extract_contract
from codegen.extract import contract_to_json as sttm_to_json
from codegen.extract.annotations import annotation_reason
from codegen.extract.frd_docx import contract_to_json as frd_to_json
from codegen.gate.derivations import check_iig_derivations, normalise_path
from codegen.gate.preflight import GateCheck, _ruff_check
from codegen.gate.verdict import compute_verdict
from codegen.layout.answers import AnswersFile, apply_answers
from codegen.layout.model import MockLayoutProvider
from codegen.layout.resolve import resolve_pair
from codegen.resolve import resolver
from codegen.resolve.resolver import pending_answers
from codegen.resolve.resolver import resolve_pair as resolve_contracts

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PAIR_1_FRD = SHAPES / "frd" / "f1_pair_1.docx"
PAIR_11_FRD = SHAPES / "frd" / "f1_pair_11_multi_file.docx"
DATE = "2026-01-01"
REC_TYPE_TARGETS = {"HDDR": "REC_TYP_HDR", "DET": "REC_TYP_DTL", "TRLR": "REC_TYP_TRLR"}
BLANK_SHEET = "MAPPING-VC_DISENROLLMENT"


def _resolved(tmp_path: Path, config, workbook):
    sttm = tmp_path / "sttm.xlsx"
    sttm.write_bytes(xlsx_bytes(workbook))
    pair = resolve_pair(sttm, PAIR_1_FRD, config, provider=None, cache_dirs=[],
                        generated_date=DATE)
    frd_json, sttm_json = tmp_path / "frd.json", tmp_path / "sttm.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile, width_answers=pair.width_answers)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config)
    return contract, spec


def _context(spec, config):
    return build_context(spec, config, allow_duplicate_file_name=True,
                         duplicate_rule_text=None, notification_rule_texts=[])


# ----------------------------------------------------- 1. key lookup


def test_rec_type_in_every_segment_resolves_the_detail_key_to_the_detail_column(
        config, tmp_path):
    contract, spec = _resolved(
        tmp_path, config, sttm_fixtures.build_pair1(rec_type_targets=REC_TYPE_TARGETS))
    fields = contract.feeds[0].fields
    # The repro's shape: one source name, three segments, three target columns.
    assert [(f.record_segment, f.stage_column) for f in fields
            if f.source_column == "REC_TYPE"] == [
        ("Header", "REC_TYP_HDR"), ("Detail", "REC_TYP_DTL"), ("Trailer", "REC_TYP_TRLR")]
    assert "REC_TYPE" in contract.feeds[0].load_rules.not_null_columns
    # The Detail key resolves inside the Detail segment, never across segments.
    assert "REC_TYP_DTL" in spec.natural_key_columns
    assert not {"REC_TYP_HDR", "REC_TYP_TRLR"} & set(spec.natural_key_columns)
    assert spec.not_null_columns == spec.natural_key_columns
    # Resolved by its own segment: no key was dropped, so no flag.
    assert not [f for f in spec.provenance_flags if f.startswith("key_column_not_in_table")]
    # The emitter (the first run's KeyError: 'REC_TYP_TRLR') now renders, the
    # sample key column being the Detail record's own source field.
    context = _context(spec, config)
    assert "REC_TYPE" in context["detail_natural_key_source_columns"]
    assert context["natural_key"] == spec.natural_key_columns


def _field(name: str, column: str, segment: str, *, key: bool) -> SttmField:
    return SttmField(source_column=name, description=None, sample_value=None,
                     source_datatype="String", nullable=not key, phi=False, mandatory=key,
                     stage_column=column, stage_datatype="String", standard_column=None,
                     standard_datatype=None, value_spec=None, record_segment=segment,
                     stage_table=f"t_{segment.lower()}")


def _segments(*fields: SttmField) -> list[SegmentSpec]:
    out = []
    for segment in ("Header", "Detail", "Trailer"):
        own = [f for f in fields if f.record_segment == segment]
        if own:
            out.append(SegmentSpec(segment=segment, fields=own, stage_table=ResolvedTable(
                catalog=None, schema_name="s", table=f"t_{segment.lower()}", role="stage")))
    return out


def test_a_key_marked_only_in_the_trailer_is_the_trailers_not_the_details():
    segments = _segments(
        _field("REC_TYPE", "REC_TYP_HDR", "Header", key=False),
        _field("REC_TYPE", "REC_TYP_DTL", "Detail", key=False),
        _field("MEMBER_ID", "MEMBER_ID", "Detail", key=True),
        _field("REC_TYPE", "REC_TYP_TRLR", "Trailer", key=True))
    flags: list[str] = []
    assert resolver._detail_key_columns(segments, ["MEMBER_ID", "REC_TYPE"], flags) == ["MEMBER_ID"]
    assert flags == []             # it HAS a column in its own (trailer) table


def test_a_key_with_no_column_after_per_segment_resolution_is_dropped_and_flagged():
    segments = _segments(_field("MEMBER_ID", "MEMBER_ID", "Detail", key=True))
    flags: list[str] = []
    assert resolver._detail_key_columns(segments, ["MEMBER_ID", "REC_TYPE"], flags) == ["MEMBER_ID"]
    assert len(flags) == 1
    assert flags[0].startswith("key_column_not_in_table:REC_TYPE — ")
    assert "'t_detail'" in flags[0] and "dropped from the natural key" in flags[0]


# ------------------------------------------- 2a. FILE_DETAILS annotation rows


@pytest.mark.parametrize("cells,name_index,expected", [
    (["VENDOR_C", "vc_risk_YYYYMMDD.csv", None, None, "Monthly"], 1, None),
    (["VENDOR_C", "VC RISK FILE.csv", None, None, "Monthly"], 1, None),    # spaces + extension
    (["DataType = String unless stated", None, None, None, None], 1, "assignment"),
    ([None, None, "Files arrive by 6 AM ET", None, None], 1, "no file name"),
    (["VENDOR_C", "see the FRD for the file list", None, None, None], 1, "prose"),
    (["VENDOR_C", "a.csv\nb.csv", None, None, None], 1, "several lines"),
    ([None, None, None, None, None], 1, None),                             # blank: not skipped
])
def test_a_file_details_row_is_an_annotation_only_when_it_does_not_fit(cells, name_index,
                                                                       expected):
    reason = annotation_reason(cells, name_index)
    assert (reason is None) if expected is None else (expected in reason)


def _pair11_with_annotations():
    wb = sttm_fixtures.build_pair11()
    ws = wb["FILE_DETAILS"]
    ws.cell(row=5, column=1, value="DataType = String unless stated")
    ws.cell(row=6, column=3, value="Files arrive by 6 AM ET")
    for col, value in ((1, "VENDOR_C"), (2, "e.g. VC_SAMPLE_20260101.csv"), (5, "Daily")):
        ws.cell(row=7, column=col, value=value).font = Font(italic=True)
    return wb


def _pair11(tmp_path: Path, config, workbook, answers: dict | None = None, frd=PAIR_11_FRD):
    sttm = tmp_path / "p11.xlsx"
    sttm.write_bytes(xlsx_bytes(workbook))
    pair = resolve_pair(sttm, frd, config, provider=MockLayoutProvider([]), cache_dirs=[],
                        use_cache=False, generated_date=DATE, answers=answers)
    frd_json = tmp_path / "frd.contract.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    return sttm, pair, frd_json


def _extract(sttm: Path, pair, frd_json: Path, config, tmp_path: Path):
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    sttm_json = tmp_path / "sttm.contract.json"
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    return contract, sttm_json


def test_file_details_annotation_rows_are_skipped_reported_and_never_files(config, tmp_path):
    sttm, pair, frd_json = _pair11(tmp_path, config, _pair11_with_annotations())
    assert pair.questions == []          # the layout stage did not read them as files either
    contract, sttm_json = _extract(sttm, pair, frd_json, config, tmp_path)
    skipped = [n for n in contract.notes if n.startswith("FILE_DETAILS annotation row skipped")]
    assert [n.split(" — ")[0].rsplit(": ", 1)[1] for n in skipped] == [
        "FILE_DETAILS!row 5", "FILE_DETAILS!row 6", "FILE_DETAILS!row 7"]
    assert "assignment line" in skipped[0] and "a note row" in skipped[1]
    assert "styled as guidance (italic font)" in skipped[2]
    for feed in contract.feeds:
        flags = [f for f in feed.extraction_flags
                 if f.startswith("file_details_annotation_skipped:")]
        assert len(flags) == 3 and all(f.endswith("(not read as a file)") for f in flags)
    specs = resolve_contracts(frd_json, sttm_json, config)
    assert [s.feed_slug for s in specs] == ["vc_enrollment", "vc_disenrollment",
                                            "vc_individual_risk"]
    named = {p for s in specs for p in s.file_name_patterns}
    assert not [p for p in named if "SAMPLE" in p or "=" in p or " " in p]


def test_the_content_driven_reader_skips_annotation_rows_too(config, tmp_path):
    from codegen.extract.generic import read_workbook
    from codegen.layout.discover import discover

    wb = sttm_fixtures.build_pair5()
    ws = wb["File Details"]
    ws.cell(row=3, column=2, value="FileName = the vendor's name, as sent")
    path = tmp_path / "p5.xlsx"
    path.write_bytes(xlsx_bytes(wb))
    found = discover(path, config.extractor)
    ir = read_workbook(found, path.name, config)
    (details,) = [a for a in ir.auxiliary if a.kind == "file_details"]
    assert [r[1] for r in details.rows] == ["FEED_5_*.txt"]
    assert len(ir.annotations) == 1 and ir.annotations[0].startswith("File Details!row 3")


# ------------------------------------------- 2b. blank Schema / TableName


def test_blank_targets_take_the_frd_table_and_hold_the_feed_for_the_schema(config, tmp_path):
    sttm = tmp_path / "sttm.xlsx"
    sttm.write_bytes(xlsx_bytes(sttm_fixtures.build_pair1(blank_targets=True)))
    pair = resolve_pair(sttm, PAIR_1_FRD, config, provider=None, cache_dirs=[],
                        generated_date=DATE)
    assert pair.questions == []
    assert any(f.startswith("frd_unstated:feeds[0].stage_target.schema source_used:none")
               for f in pair.flags)                        # never "the STTM band" when blank
    frd_json = tmp_path / "frd.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    # The table: the FRD's only table. The schema: no FRD statement, no layer
    # convention — NEEDS_ANSWERS with the key, never a mid-run stop.
    with pytest.raises(NeedsAnswersError) as held:
        extract_contract(sttm, frd_json, config, generated_date=DATE, layout=pair.sttm.profile)
    assert [p.key for p in held.value.pending] == ["feeds[0].stage_target.schema",
                                                   "feeds[0].standard_target.schema"]
    assert "every data row of the stage schema column is empty" in held.value.pending[0].reason
    # The keys are answers-file keys: answered, the feed extracts and resolves.
    answers = {"gaps": {
        "feeds[0].stage_target.schema": {"value": "stg_answered", "layer": None,
                                         "source": "user"},
        "feeds[0].standard_target.schema": {"value": "std_answered", "layer": None,
                                            "source": "user"}}}
    pair = resolve_pair(sttm, PAIR_1_FRD, config, provider=None, cache_dirs=[],
                        generated_date=DATE, answers=answers)
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    contract, sttm_json = _extract(sttm, pair, frd_json, config, tmp_path)
    (feed,) = contract.feeds
    assert (feed.stage.schema_name, feed.stage.table) == ("stg_answered", "vnd_p_accum_client")
    assert (feed.standard.schema_name, feed.standard.table) == ("std_answered",
                                                                "vnd_p_accum_client")
    missing = [f.split(" — ")[0] for f in feed.extraction_flags
               if f.startswith("sttm_target_missing:")]
    assert missing == [
        "sttm_target_missing:stage.table source_used:FRD stage target (its only table): "
        "'vnd_p_accum_client'",
        "sttm_target_missing:stage.schema source_used:FRD 'Target Catalog and Schema': "
        "'stg_answered'",
        "sttm_target_missing:standard.table source_used:FRD standard target (its only table): "
        "'vnd_p_accum_client'",
        "sttm_target_missing:standard.schema source_used:FRD 'Target Catalog and Schema': "
        "'std_answered'"]
    (spec,) = resolve_contracts(frd_json, sttm_json, config)
    assert spec.segments[0].stage_table.schema_name == "stg_answered"


def _pair11_one_sheet_blank(target_column: bool = False):
    wb = sttm_fixtures.build_pair11()
    for row in wb[BLANK_SHEET].iter_rows(min_row=3):
        for index in (8, 9, 12, 13):        # stage Schema / TableName, standard Schema / TableName
            row[index].value = None
    if target_column:
        details = wb["FILE_DETAILS"]
        details.cell(row=1, column=6, value="Target Table")
        details.cell(row=3, column=6, value="vc_disenrollment")
    return wb


def test_one_sheet_without_targets_holds_back_its_feed_only(config, tmp_path, monkeypatch,
                                                            capsys):
    sttm, pair, frd_json = _pair11(tmp_path, config, _pair11_one_sheet_blank())
    assert [f.feed_name for f in pair.frd_contract.feeds][-1] == "VC_DISENROLLMENT"
    contract, sttm_json = _extract(sttm, pair, frd_json, config, tmp_path)
    assert [f.feed_id for f in contract.feeds] == ["vc_enrollment", "vc_individual_risk"]
    assert [(p.feed_name, p.key) for p in contract.needs_answers] == [
        ("VC_DISENROLLMENT", "feeds[2].stage_target.schema"),
        ("VC_DISENROLLMENT", "feeds[2].stage_target.tables"),
        ("VC_DISENROLLMENT", "feeds[2].standard_target.schema"),
        ("VC_DISENROLLMENT", "feeds[2].standard_target.tables")]
    assert "MAPPING-VC_DISENROLLMENT!J: every mapping row's stage table cell is empty" in \
        contract.needs_answers[1].reason
    # The resolver skips the held-back feed (it is not "unmatched").
    specs = resolve_contracts(frd_json, sttm_json, config)
    assert [s.feed_id for s in specs] == ["vc_enrollment", "vc_individual_risk"]
    assert pending_answers(sttm_json) == contract.needs_answers
    # The CLI generates the others and names the keys: exit 3, not a FAIL.
    monkeypatch.setenv("CODEGEN_STORAGE_OUTPUTS", f"local:{(tmp_path / 'out').as_posix()}")
    (tmp_path / "out").mkdir()
    monkeypatch.chdir(REPO)
    capsys.readouterr()
    rc = cli.main(["generate", "--frd-contract", str(frd_json), "--sttm-contract",
                   str(sttm_json), "--output-mode", "framework", "--skip-tests", "--dry-run"])
    out = capsys.readouterr().out
    assert ("NEEDS_ANSWERS   VC_DISENROLLMENT (sheet MAPPING-VC_DISENROLLMENT) — answer "
            "`feeds[2].stage_target.tables` under `gaps:`") in out
    assert rc == cli.EXIT_NEEDS_ANSWERS, out
    # Answered, all three feeds extract.
    answers = {"gaps": {
        "feeds[2].stage_target.tables": {"value": "vc_disenrollment", "layer": None,
                                         "source": "user"},
        "feeds[2].stage_target.schema": {"value": "stg_dom_a", "layer": None, "source": "user"},
        "feeds[2].standard_target.tables": {"value": "vc_disenrollment", "layer": None,
                                            "source": "user"},
        "feeds[2].standard_target.schema": {"value": "dom_a", "layer": None, "source": "user"}}}
    sttm, pair, frd_json = _pair11(tmp_path, config, _pair11_one_sheet_blank(), answers=answers)
    contract, sttm_json = _extract(sttm, pair, frd_json, config, tmp_path)
    assert contract.needs_answers == []
    assert sorted(s.feed_id for s in resolve_contracts(frd_json, sttm_json, config)) == [
        "vc_disenrollment", "vc_enrollment", "vc_individual_risk"]


def test_the_file_details_target_and_the_layer_convention_fill_a_blank_sheet(config,
                                                                             tmp_path):
    with_convention = config.model_copy(update={"conventions": config.conventions.model_copy(
        update={"default_schema": {"stage": "stg_dom_a", "standard": "dom_a"}})})
    sttm, pair, frd_json = _pair11(tmp_path, with_convention,
                                   _pair11_one_sheet_blank(target_column=True))
    contract, _sttm_json = _extract(sttm, pair, frd_json, with_convention, tmp_path)
    assert [p.key for p in contract.needs_answers] == ["feeds[2].standard_target.tables"]
    feed_name = contract.needs_answers[0].feed_name
    assert feed_name == "VC_DISENROLLMENT"
    # The stage table came from File Details and the schemas from the layer
    # convention; the standard table has no File Details link (stage only)
    # and the FRD names none — the one answer still needed.
    note = next(n for n in contract.notes if "NEEDS_ANSWERS" in n)
    assert "feeds[2].standard_target.tables" in note


def test_answer_keys_for_targets_travel_through_the_answers_file():
    answers = AnswersFile([], {"feeds[2].stage_target.schema": {"value": "stg_x"},
                               "feeds[0].not_a_target": {"value": "x"}}, {})
    out, notes = apply_answers(answers, [], {})
    assert set(out["gaps"]) == {"feeds[2].stage_target.schema"}
    assert notes == ["gaps['feeds[0].not_a_target'] matches no open question"]


# ------------------------------------------- 3. path normalisation


@pytest.mark.parametrize("raw,lower,expected", [
    ("/inbound/vendor_c/risk.", False, "/inbound/vendor_c/risk"),
    ("inbound//vendor_c//", False, "inbound/vendor_c/"),
    ("\\inbound\\Vendor C\\", False, "/inbound/Vendor C/"),
    ("/x/./y;", False, "/x/y"),
    ("/Inbound/Vendor_C/", True, "/inbound/vendor_c/"),
    ("abfss://landing@acct.example.invalid/in/risk./", False,
     "abfss://landing@acct.example.invalid/in/risk/"),
    ("/inbound/vendor_c/", False, "/inbound/vendor_c/"),                 # clean: untouched
])
def test_normalise_path(raw, lower, expected):
    value, changes = normalise_path(raw, lowercase=lower)
    assert value == expected
    assert bool(changes) == (raw != expected)


@pytest.fixture(scope="module")
def pair11_catalog_specs(config, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("p11_catalog")
    sttm, pair, frd_json = _pair11(tmp, config, sttm_fixtures.build_pair11(),
                                   frd=SHAPES / "frd" / "f1_pair_11_multi_file_catalog.docx")
    _contract, sttm_json = _extract(sttm, pair, frd_json, config, tmp)
    return resolve_contracts(frd_json, sttm_json, config)


def _scoped(config, tmp: Path):
    return config.model_copy(update={"output": config.output.model_copy(update={
        "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")})})


def test_input_path_punctuation_is_normalised_flagged_and_never_fails_the_gate(
        config, pair11_catalog_specs, tmp_path):
    spec = next(s for s in pair11_catalog_specs if s.feed_slug == "vc_individual_risk")
    dirty = spec.model_copy(update={"landing_location": "/inbound//vendor_c/risk."})
    for profile, template in (("acfc_prx", "iig_v2"), (None, None)):
        gate = cli._generate_feed(dirty, _scoped(config, tmp_path / str(profile)), dry_run=True,
                                  skip_tests=True, output_mode="framework",
                                  conventions_profile=profile, iig_template=template)
        failed = {c.name: c.details for c in gate.checks if not c.passed}
        assert "derivations" not in failed and "sql_literals" not in failed, failed
        normalised = [f for f in gate.flags if f.startswith("path_normalised:")]
        assert normalised, (profile, gate.flags)
        assert all("risk./" in f.split(" after ")[0] or "risk." in f.split(" after ")[0]
                   for f in normalised), normalised
        assert all("risk." not in f.split(" after ", 1)[1].split(" (")[0]
                   for f in normalised), normalised


def test_the_family_path_case_lowercases_input_segments_only(config):
    from types import SimpleNamespace

    from codegen.metadata_template import _paths

    tpl = config.metadata.templates["iig_v2"].model_copy(update={"path_case": "lower"})
    feed = SimpleNamespace(landing_location="/Inbound/Vendor_C/Risk.")
    cells = _paths(tpl, "ADLS_DELTA_INGESTION_DETAILS", feed, "VC_RISK", "VC_RISK_RJT")
    value = cells["TGT_ADLS_PATH"]["value"]
    assert value == "/inbound/vendor_c/risk/Processed/vc_risk"     # the shape's literal kept
    note = cells["TGT_ADLS_PATH"]["badge_entry"]["path_note"]
    assert note.startswith("path_normalised:ADLS_DELTA_INGESTION_DETAILS.TGT_ADLS_PATH — before "
                           "'/Inbound/Vendor_C/Risk./Processed/VC_RISK'")
    assert "lowercased (path_case: lower)" in note
    payload = {"tabs": {"ADLS_DELTA_INGESTION_DETAILS": {"rows": [{
        "values": {"TGT_ADLS_PATH": value},
        "badges": {"TGT_ADLS_PATH": cells["TGT_ADLS_PATH"]["badge_entry"]}}]}}}
    assert check_iig_derivations(payload, config).passed


# ------------------------------------------- 4. ruff safe fixes before the gate


_RUFF_TOML = (REPO / "src" / "codegen" / "templates" / "ruff.toml.j2").read_text(encoding="utf-8")
_UNSORTED = '"""Module."""\n\nimport sys\nimport json\n\nVALUES = (json, sys)\n'


def _feed_dir(tmp_path: Path, modules: dict[str, str]) -> Path:
    feed = tmp_path / "feed_x"
    (feed / "pipeline").mkdir(parents=True)
    (feed / "ruff.toml").write_text(_RUFF_TOML, encoding="utf-8", newline="\n")
    for name, text in modules.items():
        (feed / "pipeline" / f"{name}.py").write_text(text, encoding="utf-8", newline="\n")
    return feed


def test_cosmetic_lint_is_fixed_recorded_and_never_a_fail(tmp_path):
    feed = _feed_dir(tmp_path, {"tidy": _UNSORTED})
    check = _ruff_check(feed)
    assert check.passed and not check.not_run, check.details
    assert check.fixes == ["pipeline/tidy.py: I001 x1"]
    assert "import json\nimport sys" in (feed / "pipeline" / "tidy.py").read_text(encoding="utf-8")
    gate = compute_verdict("feed_x", [], [], [check], tests_skipped=False)
    assert gate.verdict == "PASS_WITH_FLAGS"
    assert gate.flags == ["ruff_fixed: 1 kind(s) of ruff finding fixed automatically (safe "
                          "fixes only) before the check — pipeline/tidy.py: I001 x1"]


def test_a_finding_without_a_safe_fix_still_fails_and_says_what_was_fixed(tmp_path):
    feed = _feed_dir(tmp_path, {"tidy": _UNSORTED, "bad": '"""Bad."""\n\nVALUE = undefined\n'})
    check = _ruff_check(feed)
    assert not check.passed
    assert check.fixes == ["pipeline/tidy.py: I001 x1"]
    assert check.details.startswith("1 finding(s) remain after 1 safe fix(es) applied first")
    assert "pipeline/bad.py:3:9: F821" in check.details
    gate = compute_verdict("feed_x", [], [], [check], tests_skipped=False)
    assert gate.verdict == "FAIL" and gate.flags[0].startswith("ruff_fixed:")


def test_a_clean_feed_is_untouched_and_carries_no_flag(tmp_path):
    feed = _feed_dir(tmp_path, {"tidy": '"""Module."""\n\nimport json\nimport sys\n\n'
                                        'VALUES = (json, sys)\n'})
    check = _ruff_check(feed)
    assert check == GateCheck(name="ruff", passed=True, details="ruff clean")


def test_a_fixed_module_resyncs_the_assembled_notebook(config, tmp_path):
    from codegen.emit.emitter import emit_feed
    from codegen.emit.notebook import resync_notebook

    _contract, spec = _resolved(tmp_path, config, sttm_fixtures.build_pair1())
    emit_feed(_context(spec, config), tmp_path / "out")
    feed = tmp_path / "out" / spec.feed_slug
    module = feed / "pipeline" / "mapping.py"
    notebook = feed / f"{spec.feed_slug}.ipynb"
    # The gate formats every emitted .py first (rule 8 of the real-row
    # scorecard): the CLEAN state is the gated one.
    assert _ruff_check(feed).passed
    clean_module, clean_notebook = module.read_bytes(), notebook.read_bytes()
    # The cosmetic defect, in the module AND in the notebook assembled from it.
    swapped = clean_module.replace(
        b"from pyspark.sql import DataFrame\nfrom pyspark.sql import functions as F\n",
        b"from pyspark.sql import functions as F\nfrom pyspark.sql import DataFrame\n")
    assert swapped != clean_module
    module.write_bytes(swapped)
    resync_notebook(notebook, feed)
    assert notebook.read_bytes() != clean_notebook
    check = _ruff_check(feed)
    assert check.passed and check.fixes == ["pipeline/mapping.py: I001 x1"], check.details
    assert module.read_bytes() == clean_module
    assert notebook.read_bytes() == clean_notebook         # never drifts from its modules


def test_the_content_driven_reader_takes_this_feeds_file_details_target(config, tmp_path):
    from acfc_shapes import pair1

    wb = sttm_fixtures.build_pair1(blank_targets=True)
    details = wb.create_sheet("File Details")
    details.append(["Vendor", "FileName", "Frequency", "Target Table"])
    details.append(["VENDOR_A", pair1.FILE_PATTERNS[0], "Daily", "tbl_from_details"])
    details.append(["VENDOR_B", "OTHER_FEED_*.csv", "Daily", "another_feeds_table"])
    sttm = tmp_path / "sttm.xlsx"
    sttm.write_bytes(xlsx_bytes(wb))
    pair = resolve_pair(sttm, PAIR_1_FRD, config, provider=None, cache_dirs=[],
                        generated_date=DATE)
    frd = json.loads(frd_to_json(pair.frd_contract))
    for layer in ("stage_target", "standard_target"):
        frd["feeds"][0][layer]["schema"] = f"{layer[:3]}_frd"
    frd_json = tmp_path / "frd.json"
    frd_json.write_text(json.dumps(frd), encoding="utf-8")
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    (feed,) = contract.feeds
    # The row of THIS feed's file — never another feed's row.
    assert feed.stage.table == "tbl_from_details"
    flag = next(f for f in feed.extraction_flags
                if f.startswith("sttm_target_missing:stage.table"))
    assert "source_used:File Details File Details!D" in flag
