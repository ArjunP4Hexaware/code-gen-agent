"""M0 — the ACFC-shape fixture universe (fixtures/acfc_shapes/).

Asserts that every tracked fixture (a) is reproduced byte for byte by its
builder in tests/acfc_shapes/, (b) has exactly the documented structure —
sheet names, meta-row labels at their rows, band row position/texts/merges,
header row position/texts/count (docs/acfc/SHAPES_FOR_PORT.md), (c) the
pair-1 STTM / FRD / VDD fixtures and the aliased golden all describe the same
21 columns + 3 audit columns with the golden's types, and (d) nothing under
the directory trips the scrub denylist scanner.
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from acfc_shapes import FAMILY, FIXTURE_ROOT, build_all, pair1
from acfc_shapes.frd import DATA_QUALITY, SECTION_TITLES, STRUCTURAL_P1, STRUCTURAL_P2
from acfc_shapes.vdd import FILES_HEADER, V1_HEADER, V2_HEADER, V3_HEADER

REPO = Path(__file__).resolve().parents[1]
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture(scope="module")
def built() -> dict[str, bytes]:
    return build_all()


def _wb(rel: str):
    return load_workbook(FIXTURE_ROOT / rel)


def _row(ws, index: int) -> list:
    return [c for c in next(ws.iter_rows(min_row=index, max_row=index, values_only=True))]


def _texts(ws, index: int) -> list[str]:
    return [str(c) for c in _row(ws, index) if c is not None]


# ---------------------------------------------------------------- registry ---


def test_every_registered_fixture_is_tracked_and_byte_stable(built):
    assert set(built) == set(FAMILY)
    for rel, data in built.items():
        path = FIXTURE_ROOT / rel
        assert path.is_file(), f"missing fixture {rel} — run scripts/build_acfc_shapes.py"
        assert path.read_bytes() == data, f"{rel} drifted from its builder"


def test_gitignore_re_includes_every_binary_fixture():
    ignore = (REPO / ".gitignore").read_text(encoding="utf-8")
    for rel in FAMILY:
        if rel.endswith((".xlsx", ".docx")):
            assert f"!fixtures/acfc_shapes/{rel}" in ignore, rel


def test_readme_maps_every_fixture_to_family_and_pair():
    readme = (FIXTURE_ROOT / "README.md").read_text(encoding="utf-8")
    for rel in FAMILY:
        assert f"`{rel}`" in readme, rel


def test_scrub_scanner_is_clean_over_the_fixture_universe():
    sys.path.insert(0, str(REPO / "scripts"))
    from scrub_check import scan_paths

    assert scan_paths([FIXTURE_ROOT]) == []


# ------------------------------------------------------- STTM families A–E ---

# fixture -> (sheet names, mapping sheets -> (band_row, band_texts, header_row, n_headers))
_STTM_SHAPES = {
    "sttm/pair_1_family_a.xlsx": (
        ["Version", "LOB_CROSSWALK", "FEED_1_MAPPING"],
        {"FEED_1_MAPPING": (14, ["Source Data", "Data Rules and Primary Keys",
                                 "Staging Layer Table", "Standard Layer Table"], 15, 27)},
    ),
    "sttm/pair_8_family_a.xlsx": (
        ["FEED_8_LAYOUT"],
        {"FEED_8_LAYOUT": (7, ["Header Record", "Staging Layer", "Standard Layer"], 8, 21)},
    ),
    "sttm/pair_2_family_b.xlsx": (
        ["FILE_DETAILS", "VERSION_HISTORY", "MAPPING-", "MAPPING-1", "MAPPING-2"],
        {name: (1, ["Source File Layout", "Stage Layer", "Standard Layer"], 2, 16)
         for name in ("MAPPING-", "MAPPING-1", "MAPPING-2")},
    ),
    "sttm/pair_3_family_c.xlsx": (
        ["Version Control", "Layout", "STTM", "FileNames"],
        {"STTM": (4, ["Source - [SRC_SYS_A] File", "Destination",
                      "Target Table - Staging Layer", "Target Table - Standard Layer"], 5, 31)},
    ),
    "sttm/pair_4_family_d.xlsx": (
        ["STTM", "LOB_CROSSWALK", "OVERPUNCH_RULES"],
        {"STTM": (4, ["Source Data", "Data Rules", "Target Table - Staging Layer",
                      "Target Table - Standard Layer"], 5, 25)},
    ),
    "sttm/pair_6_family_d.xlsx": (
        ["Version Control", "Server details", "FEED_6_MAPPING"],
        {"FEED_6_MAPPING": (2, ["Source table", "Stage Layer", "STD layer"], 3, 22)},
    ),
    "sttm/pair_5_family_e.xlsx": (
        ["Version History", "File Details", "Table_Details", "FEED_5_MAPPING",
         "Sheet2", "Sheet3", "Sheet1"],
        {"FEED_5_MAPPING": (10, ["Source Layout", "Stage Layer", "Standard Layer"], 11, 33)},
    ),
    "sttm/pair_7_family_e.xlsx": (
        ["Version History", "File Details", "MAPPING_FEED_7", "Sheet2", "Queries",
         "Analysis_Notes", "Analysis Sheet - Detail", "Analysis Sheet - Compound",
         "Sample Records", "Sheet1"],
        {"MAPPING_FEED_7": (None, [], 1, 19)},   # no band row
    ),
    "sttm/pair_9_family_e.xlsx": (
        ["Summary", "MAPPING_FEED_9", "MAPPING_FEED_9_TBL2", "Sheet2", "Log", "Sheet1"],
        {name: (1, ["Source Tables", "Target - DL Staging Layer",
                    "Target - DL Standard Layer"], 2, 24)
         for name in ("MAPPING_FEED_9", "MAPPING_FEED_9_TBL2")},
    ),
    "sttm/pair_10_family_e.xlsx": (
        ["Version", "Outbound_REGION_A", "Inbound_REGION_A", "Inbound_REGION_B_Adult",
         "Inbound_REGION_B_Child", "Mapping"],
        {"Outbound_REGION_A": (4, ["Source Data", "Stage Layer", "Standard Layer"], 5, 17),
         "Inbound_REGION_A": (4, ["Source Data", "Stage Layer", "Standard Layer"], 5, 20),
         "Inbound_REGION_B_Adult": (4, ["Source Data", "Stage Layer", "Standard Layer"], 5, 18),
         "Inbound_REGION_B_Child": (4, ["Source Data", "Stage Layer", "Standard Layer"], 5, 18)},
    ),
}


@pytest.mark.parametrize("rel", sorted(_STTM_SHAPES))
def test_sttm_fixture_has_documented_sheets_bands_and_headers(rel):
    sheet_names, mapping = _STTM_SHAPES[rel]
    wb = _wb(rel)
    assert wb.sheetnames == sheet_names
    for sheet, (band_row, band_texts, header_row, n_headers) in mapping.items():
        ws = wb[sheet]
        if band_row is not None:
            assert _texts(ws, band_row) == band_texts, (rel, sheet, "band texts")
            if len(ws.merged_cells.ranges) or rel != "sttm/pair_2_family_b.xlsx":
                merged_rows = {r.min_row for r in ws.merged_cells.ranges}
                assert band_row in merged_rows, (rel, sheet, "band merges")
        headers = _row(ws, header_row)
        assert len([h for h in headers if h is not None]) == n_headers, (rel, sheet, "header count")
        # 10–30 plausible rows below the header row on every mapping sheet.
        data_rows = [r for r in ws.iter_rows(min_row=header_row + 1, values_only=True)
                     if any(c is not None for c in r)]
        assert 10 <= len(data_rows) <= 30, (rel, sheet, len(data_rows))


# Documented header texts, verbatim, in order (SHAPES_FOR_PORT §1).
_HEADER_TEXTS = {
    ("sttm/pair_1_family_a.xlsx", "FEED_1_MAPPING", 15): [
        "S.No", "Segment", "Field Name", "Required?", "Format", "Start", "Length", "End",
        "Description", "Data Definition", "PII", "Primary Key", "Critical Data", "Not NULL",
        "Load Rules", "Workspace", "Target Catalog", "Target Schema Name in DL",
        "Target Table Name in DL", "Target_Column_Name_in_DL", "Target Data Type in DL",
        "Workspace", "Target Catalog", "Target Schema Name in DL", "Target Table Name in DL",
        "Target_Column_Name_in_DL", "Target Data Type in DL"],
    ("sttm/pair_2_family_b.xlsx", "MAPPING-1", 2): [
        "Database column Name", "NULL CHECK", "Description", "Sample Value", "DataType",
        "PHI Field", "Mandatory Field", "Comment", "Schema", "TableName", "ColumnName",
        "DataType", "Schema", "TableName", "ColumnName", "DataType"],
    ("sttm/pair_3_family_c.xlsx", "Layout", 4): [
        "Field #", "Field Name", "Req'd?", "Type", "Length", "Start", "End", "Comments",
        "LOB_A", "LOB_B", "LOB_C", "LOB_D", "LOB_E"],
    ("sttm/pair_4_family_d.xlsx", "STTM", 5): [
        "Sl. No.", "Field Name", "Type", "Start", "End", "Len", "Description", "Data Definition",
        "PII", "Primary Key", "Critical Data", "Not NULL", "LOAD Rule", "Workspace",
        "Target Catalog", "Target Schema Name in Lakehouse", "Target Table Name in Lakehouse",
        "Target Column Name in Lakehouse", "Target Data Type in Lakehouse", "Workspace",
        "Target Catalog", "Target Schema Name in Lakehouse", "Target Table Name in Lakehouse",
        "Target Column Name in Lakehouse", "Target Data Type in Lakehouse"],
    ("sttm/pair_6_family_d.xlsx", "FEED_6_MAPPING", 3): [
        "Inscope for Implementation", "SERVER NAME", "TABLE_CATALOG", "TABLE_SCHEMA", "TABLE_NAME",
        "COLUMN_NAME", "Primary Key", "DATA_TYPE", "NULL/NOT NULL", "Load Rules", "Workspace",
        "Catalog", "Schema", "TableName", "ColumnName", "DataType", "Workspace", "Catalog",
        "Schema", "TableName", "ColumnName", "DataType"],
    ("sttm/pair_7_family_e.xlsx", "MAPPING_FEED_7", 1): [
        "Field ID", "Field Name", "Mandatory or Situational", "Column Description", "Comments",
        "Format", "Size", "PHI/PII Field", "Stage Schema", "Stage Table Name",
        "Stage Table - Column Name", "Stage Table - DataType", "Standard Schema",
        "Standard Table Name", "Standard Table - Column Name", "Standard Data Type",
        "Mandatory Fields (Include in DQ Check)", "Recycle Flag", "Comments"],
    ("sttm/pair_8_family_a.xlsx", "FEED_8_LAYOUT", 8): [
        "count", "Sr. No", "Input File", "Output Flat File layout", "Data Definition", "PII",
        "Primary Key", "Critical Data", "NOT NULL", "Workspace", "Catalog", "Schema", "TableName",
        "ColumnName", "DataType", "Workspace", "Catalog", "Schema", "TableName", "ColumnName",
        "DataType"],
    ("sttm/pair_10_family_e.xlsx", "Outbound_REGION_A", 5): [
        "Source Column Name", "Nullable ?", "Description", "PII (Y/N)", "Workspace", "Catalog",
        "Stage Schema", "Stage Table Name", "Stage_Column_Name", "Datatype", "Workspace",
        "Catalog", "Standard Schema", "Standard Table Name", "Standard Column_Name", "Datatype",
        "Comments"],
}


@pytest.mark.parametrize("key", sorted(_HEADER_TEXTS))
def test_sttm_header_texts_are_verbatim(key):
    rel, sheet, row = key
    assert _texts(_wb(rel)[sheet], row) == _HEADER_TEXTS[key]


def test_pair1_meta_rows_labels_and_blanks():
    ws = _wb("sttm/pair_1_family_a.xlsx")["FEED_1_MAPPING"]
    # M9.0: the REAL sheet (HANDOVER_GENIE.md §2) — ten meta rows, no Load
    # Strategy / Notes row; names / example / frequency read TBD, format .dat.
    labels = [_row(ws, i)[0] for i in range(1, 11)]
    assert labels == ["File Names", "File Name Example", "Frequency", "File Format (text, csv)",
                      "File Delimiter", "Last Update Date", "Version", "LOB",
                      "Target table Name Desc", "Feed Type"]
    values = [_row(ws, i)[1] for i in range(1, 11)]
    assert values[:4] == ["TBD", "TBD", "TBD", ".dat"]
    assert {i for i in range(1, 11) if values[i - 1] is None} == {5, 6, 7, 10}
    assert values[7] and values[8]                                   # LOB, table desc
    for gap in (11, 12, 13):
        assert _row(ws, gap) == [None] * 30                          # nothing before the band row


def test_pair1_band_and_header_rows_are_the_captured_geometry():
    """M9.0 — rows 14 / 15 cell by cell as PAIR1_HEADERS.md §1 records them:
    four merges, 27 header texts over 30 columns, J / Q / X empty."""
    ws = _wb("sttm/pair_1_family_a.xlsx")["FEED_1_MAPPING"]
    assert sorted(str(r) for r in ws.merged_cells.ranges) == [
        "A14:I14", "K14:P14", "R14:W14", "Y14:AD14"]
    band = {ws.cell(row=14, column=c).coordinate: ws.cell(row=14, column=c).value
            for c in range(1, 31) if ws.cell(row=14, column=c).value is not None}
    assert band == {"A14": "Source Data", "K14": "Data Rules and Primary Keys",
                    "R14": "Staging Layer Table", "Y14": "Standard Layer Table"}
    header = _row(ws, 15)
    layer = ["Workspace", "Target Catalog", "Target Schema Name in DL",
             "Target Table Name in DL", "Target_Column_Name_in_DL", "Target Data Type in DL"]
    assert header == (["S.No", "Segment", "Field Name", "Required?", "Format", "Start", "Length",
                       "End", "Description", None,
                       "Data Definition", "PII", "Primary Key", "Critical Data", "Not NULL",
                       "Load Rules", None] + layer + [None] + layer)
    assert header[19] == header[26] == "Target Schema Name in DL"    # T / AA
    assert header[20] == header[27] == "Target Table Name in DL"     # U / AB
    rows = [r for r in ws.iter_rows(min_row=16, values_only=True) if any(c is not None for c in r)]
    banners = [r[0] for r in rows if sum(c is not None for c in r) == 1]
    assert banners == ["Header", "Details", "Trailer"]               # one per record block
    data = [r for r in rows if sum(c is not None for c in r) > 1]
    # Schema + table constants repeat on EVERY data row of both bands.
    assert {(r[19], r[20]) for r in data} == {(pair1.STAGE_SCHEMA, pair1.TABLE)}
    assert {(r[26], r[27]) for r in data} == {(pair1.STANDARD_SCHEMA, pair1.TABLE)}
    assert all(r[9] is None and r[16] is None and r[23] is None for r in data)   # J, Q, X
    unmapped = [r for r in data if r[21] == "Do Not Map"]
    assert len(unmapped) == 1
    assert unmapped[0][2] == "Reserved Group\n(01)" and unmapped[0][22] is None
    assert unmapped[0][28] == "Do Not Map" and unmapped[0][29] is None


def test_pair5_meta_rows_and_pair8_meta_rows():
    ws5 = _wb("sttm/pair_5_family_e.xlsx")["FEED_5_MAPPING"]
    assert [_row(ws5, i)[0] for i in range(1, 10)] == [
        "File(s)", "File Generator", "File Location", "LOB", "File frequency", "Domain",
        "Sub-Domain", "NOTE", "File type"]
    assert _row(ws5, 2)[1] is None              # File Generator blank (§4 pair 5)
    ws8 = _wb("sttm/pair_8_family_a.xlsx")["FEED_8_LAYOUT"]
    assert [_row(ws8, i)[0] for i in range(1, 7)] == [
        "File Names", "Last Update Date", "Version", "LOB", "Target table Name Desc", "Feed Type"]
    assert all(_row(ws8, i)[1] is not None for i in range(1, 7))   # all populated (§4 pair 8)
    # Multi-segment: four record blocks, segment labels merged across the source group.
    labels = [r[0] for r in ws8.iter_rows(min_row=9, values_only=True)
              if isinstance(r[0], str) and r[0].endswith(("Record", "Trailer"))]
    assert labels == ["Detail Record", "Trailer Record", "File Trailer"]
    assert len(ws8.merged_cells.ranges) >= 8


def test_pair3_layout_and_sttm_describe_the_same_fields():
    wb = _wb("sttm/pair_3_family_c.xlsx")
    layout = [r[1] for r in wb["Layout"].iter_rows(min_row=5, values_only=True) if r[1]]
    sttm = [r[1] for r in wb["STTM"].iter_rows(min_row=6, values_only=True) if r[1]]
    assert layout == sttm and len(layout) == 13


# --------------------------------------------------------------- FRD F1/F2 ---


def _docx_tables(rel: str) -> list[list[list[str]]]:
    with zipfile.ZipFile(FIXTURE_ROOT / rel) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    tables = []
    for tbl in root.iter(f"{_W}tbl"):
        rows = []
        for tr in tbl.iter(f"{_W}tr"):
            cells = []
            for tc in tr.iter(f"{_W}tc"):
                text = "".join(
                    "\n" if node.tag == f"{_W}br" else (node.text or "")
                    for node in tc.iter() if node.tag in (f"{_W}t", f"{_W}br")
                )
                cells.append(text)
            rows.append(cells)
        tables.append(rows)
    return tables


def _heading_styles(rel: str) -> list[str]:
    with zipfile.ZipFile(FIXTURE_ROOT / rel) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    return [p.find(f"{_W}pPr/{_W}pStyle").get(f"{_W}val")
            for p in root.iter(f"{_W}p") if p.find(f"{_W}pPr/{_W}pStyle") is not None]


def _sections(tables):
    return {t[0][0]: t for t in tables if t[0][0] in SECTION_TITLES}


@pytest.mark.parametrize("rel,structural", [
    ("frd/f1_pair_1.docx", STRUCTURAL_P1),
    ("frd/f1_pair_2_variant.docx", STRUCTURAL_P2),
])
def test_f1_has_six_metadata_sections_with_documented_rows(rel, structural):
    tables = _docx_tables(rel)
    assert _heading_styles(rel) == []                     # F1: no heading styles
    sections = _sections(tables)
    assert set(sections) == set(SECTION_TITLES)
    for title, table in sections.items():
        assert all(len(row) == 3 for row in table)
        assert [row[1] for row in table[1:4]] == ["Name", "Description", "Functional Requirement"]
        assert all(row[0] == title for row in table)     # section prefix in column 0
    assert [row[1] for row in sections["Structural Metadata"][4:]] == structural
    assert [row[1] for row in sections["Data Quality"][4:]] == DATA_QUALITY


def test_f1_pair1_blanks_and_values_match_the_document():
    sections = _sections(_docx_tables("frd/f1_pair_1.docx"))
    values = {(title, row[1]): row[2] for title, table in sections.items() for row in table[1:]}
    assert values[("Descriptive Metadata", "Impact Details")] == ""
    for title in ("Administrative Metadata", "Data Quality", "Vendor Metadata"):
        assert values[(title, "Description")] == ""
    # M9.0: the target cell is an inline layer block; no schema, no catalog.
    assert values[("Structural Metadata", "Target Table Name")] == (
        f"Staging Layer:\nTable: {pair1.TABLE}\nStandard Layer:\nTable: {pair1.TABLE}")
    assert values[("Structural Metadata", "Target Catalog and Schema")] == ""
    assert values[("Structural Metadata", "Load Strategy STD (View)")] == "Append"
    assert values[("Vendor Metadata", "Vendor Abbreviation")] == "VND_P"
    # M9.0 (PAIR1_HEADERS.md §2): the label IS "Object Name"; its cell lists
    # the files, one "<label>: <file name pattern>" line each; the Name row is
    # a sentence about the requirement.
    lines = values[("Descriptive Metadata", "Object Name")].split("\n")
    assert [line.split(": ", 1)[1] for line in lines] == pair1.FILE_PATTERNS
    assert all(": " in line for line in lines)
    assert values[("Descriptive Metadata", "Name")].startswith(
        "Descriptive Metadata for the ingestion of")
    assert values[("Descriptive Metadata", "Tags/Keywords")] == (
        f"Domain: {pair1.DOMAIN}\nSubdomain: {pair1.SUB_DOMAIN}")


def test_f1_pair2_variant_blanks_match_the_document():
    tables = _docx_tables("frd/f1_pair_2_variant.docx")
    sections = [t for t in tables if t[0][0] in SECTION_TITLES]
    assert len(sections) == 12                            # two complete metadata sets
    for table in sections:
        assert table[3][1] == "Functional Requirement" and table[3][2] == ""
    descriptive = next(t for t in sections if t[0][0] == "Descriptive Metadata")
    blank = {row[1] for row in descriptive[4:] if row[2] == ""}
    assert {"Object Name", "Description", "Tags/Keywords"} <= blank
    structural = next(t for t in sections if t[0][0] == "Structural Metadata")
    values = {row[1]: row[2] for row in structural[4:]}
    assert values["Target Schema"] and values["Load Strategy STD"] == "Upsert"
    assert values["ADLS Location"] == "" and values["Archive Schedule"] == ""
    assert values["Inbound/outbound File Folder Path"]


def test_f2_solution_requirement_tables_carry_section_labels_in_row_4():
    tables = _docx_tables("frd/f2_pair_8.docx")
    assert not any(t[0][0] in SECTION_TITLES for t in tables)   # no metadata tables
    styles = _heading_styles("frd/f2_pair_8.docx")
    assert "Heading1" in styles and "Heading4" in styles
    srs = [t for t in tables if t[0][0].startswith("Solution Requirement:")]
    assert len(srs) == 8
    for table in srs:
        assert len(table) == 10 and all(len(row) == 4 for row in table)
        assert table[1][0] == "Name"
        assert table[2][0] == "Business \nRequirement"        # embedded newline
        assert table[3][0] == "Functional Requirement:"       # trailing colon
        assert table[7][0] == "Priority" and table[7][2] == "Source/Reference"
    assert [t[4][0] for t in srs] == [
        "Descriptive Metadata", "Structural Metadata", "Administrative Metadata",
        "Technical Metadata", "Data Quality Considerations/Options", "Vendor Metadata",
        "Reject /Recycle Process", "Email"]


# ------------------------------------------------------------ VDD V1/V2/V3 ---


def test_vdd_files_sheet_and_header_patterns():
    for rel, header, multi in (("vdd/pair_1_v1_segments.xlsx", V1_HEADER, False),
                               ("vdd/pair_2_v2_per_file.xlsx", V2_HEADER, True),
                               ("vdd/pair_9_v3_per_table.xlsx", V3_HEADER, True)):
        wb = _wb(rel)
        assert wb.sheetnames[0] == "FILES"
        assert _texts(wb["FILES"], 1) == FILES_HEADER
        field_sheets = wb.sheetnames[1:]
        assert (len(field_sheets) > 1) is multi
        listed = {r[6] for r in wb["FILES"].iter_rows(min_row=2, values_only=True) if r[6]}
        assert listed == set(field_sheets)
        for name in field_sheets:
            assert _texts(wb[name], 1) == header


def test_vdd_pair1_segments_and_pair2_pair9_blanks():
    v1 = _wb("vdd/pair_1_v1_segments.xlsx")
    segments = [r[9] for r in v1["FEED_1 Fields"].iter_rows(min_row=2, values_only=True)]
    assert set(segments) == {"HDR", "DTL", "TRL"}
    files = list(v1["FILES"].iter_rows(min_row=2, values_only=True))
    assert all(r[3] is None and r[7] == "Y" and r[8] == "SEGMENT_IDENTIFIER" for r in files)
    v2 = _wb("vdd/pair_2_v2_per_file.xlsx")
    for name in v2.sheetnames[1:]:
        rows = list(v2[name].iter_rows(min_row=2, values_only=True))
        assert all(r[3] is None for r in rows)   # Length blank throughout
        assert all(r[8] for r in rows)           # Example Value populated
    v9 = _wb("vdd/pair_9_v3_per_table.xlsx")
    assert all(r[3] is None and r[7] is None and r[8] is None and r[9] is None
               for r in v9["FILES"].iter_rows(min_row=2, values_only=True))


# ------------------------------------------------ pair 1: one structure ---


def test_pair1_sttm_columns_and_types_equal_the_aliased_golden_ddl():
    ddl = (FIXTURE_ROOT / "pair_1/golden/ACCUM_DDL.txt").read_text(encoding="utf-8")
    ddl_columns = re.findall(r"^([A-Z_0-9]+) ([A-Za-z]+(?:\(\d+,\d+\))?),?$", ddl,
                             flags=re.MULTILINE)
    ddl_columns = [c for c in ddl_columns if c[0] != "USING"]
    # two identical blocks (stage + standard)
    assert len(ddl_columns) == 2 * 24
    stage_block = ddl_columns[:24]
    ws = _wb("sttm/pair_1_family_a.xlsx")["FEED_1_MAPPING"]
    rows = [r for r in ws.iter_rows(min_row=16, values_only=True)
            if sum(c is not None for c in r) > 1]                    # banner rows aside
    rows = [r for r in rows if r[21] != "Do Not Map"]                # the unmapped field aside
    field_rows = [r for r in rows if r[2] != "NA"]
    audit_rows = [r for r in rows if r[2] == "NA"]
    # Stage block: (Target_Column_Name_in_DL, Target Data Type in DL) at V / W;
    # standard block at AC / AD — identical to the DDL, distinct columns in order.
    stage = (list(dict.fromkeys((r[21], r[22]) for r in field_rows))
             + [(r[21], r[22]) for r in audit_rows])
    standard = (list(dict.fromkeys((r[28], r[29]) for r in field_rows))
                + [(r[28], r[29]) for r in audit_rows])
    assert stage == stage_block == standard
    assert len(field_rows) == 23 and len(audit_rows) == 3
    assert "prx" not in ddl.lower() and "acfc" not in ddl.lower()
    assert f"{pair1.STAGE_CATALOG}.{pair1.STAGE_SCHEMA}.{pair1.TABLE}" in ddl


def test_pair1_vdd_frd_and_sttm_agree_on_fields_and_positions():
    ws = _wb("sttm/pair_1_family_a.xlsx")["FEED_1_MAPPING"]
    sttm = [(r[1], r[2], r[5], r[6]) for r in ws.iter_rows(min_row=16, values_only=True)
            if r[2] not in (None, "NA") and r[21] != "Do Not Map"]
    v1 = _wb("vdd/pair_1_v1_segments.xlsx")["FEED_1 Fields"]
    vdd = [(r[9], r[1], r[3], r[5]) for r in v1.iter_rows(min_row=2, values_only=True)]
    canon = {"HDDR": "HDR", "DET": "DTL", "TRLR": "TRL"}
    assert [(canon[s], n, st, ln) for s, n, st, ln in sttm] == vdd
    sections = _sections(_docx_tables("frd/f1_pair_1.docx"))
    structural = {row[1]: row[2] for row in sections["Structural Metadata"][4:]}
    assert f"Table: {pair1.TABLE}" in structural["Target Table Name"].split("\n")
    assert structural["Target Catalog and Schema"] == ""     # the STTM band states the schema
    descriptive = {row[1]: row[2] for row in sections["Descriptive Metadata"][4:]}
    for name in ("PROCESS_DATE", "PLAN_YEAR_START_DATE", "INDIVIDUAL_DEDUCTIBLE", "FAMILY_LIMIT"):
        assert name in descriptive["Solution Acceptance Criteria"]


def test_pair1_aliased_golden_iig_keeps_structure_and_drops_raw_tokens():
    raw_golden = REPO / "docs/acfc/rfc_capture/rfc_capture/goldens/pair_1/RFC_ACCUMULATORS_IIG.xlsx"
    source = load_workbook(raw_golden) if raw_golden.is_file() else None
    aliased = _wb("pair_1/golden/RFC_ACCUMULATORS_IIG.xlsx")
    assert aliased.sheetnames == [
        "DATA_FACTORY_PIPELINE_SCHEDULE", "FILE_ADLS_INGESTION_DETAILS",
        "ADLS_DELTA_INGESTION_DETAILS", "STGDELTA_STDDELTA_INGESTION_DET",
        "ADLS_FIXED_WIDTH_HANDLER", "DATABRICKS_NOTEBOOK_DETAILS", "DATA_QUALITY_RULES",
        "EMAIL_TEMPLATE_CONFIG"]
    for ws in aliased.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert not pair1.has_raw_token(value), (ws.title, value[:60])
    if source is not None:   # the raw golden is untracked; compare when present locally
        for ws in aliased.worksheets:
            raw = source[ws.title]
            assert ws.max_row == raw.max_row and ws.max_column == raw.max_column
            assert _texts(ws, 1) == _texts(raw, 1)
    handler = aliased["ADLS_FIXED_WIDTH_HANDLER"]
    rows = {r[4]: (r[6], r[7], r[8]) for r in handler.iter_rows(min_row=2, values_only=True)}
    for segment, fields in pair1.SEGMENTS.items():
        names, lens, starts = rows[segment]
        assert names.split(",") == [f[0] for f in fields]
        assert [int(x) for x in lens.split(",")] == [f[2] for f in fields]
        assert [int(x) for x in starts.split(",")] == [f[1] for f in fields]
