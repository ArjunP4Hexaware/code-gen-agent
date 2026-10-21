"""D1 App demo (2026-10-13): the App ran Generate in notebook mode and the SD STTM's
audit row named its stage column "LOB(Hardcod to MIDS)" — TemplateGapError, no
template populates it. A trailing parenthetical note is stripped before the
column is matched (LOB), the note kept as a hardcode hint (flag
audit_column_annotated, value MIDS, citing the STTM cell). The App's default
output for acfc_prx is framework (output.profile_mode)."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from codegen import cli
from codegen.config import load_config
from codegen.extract import contract_to_json, extract_contract
from codegen.extract.audit_names import split_audit_annotation
from codegen.gate.ddl_names import (
    ddl_column_identifiers,
    ddl_identifier_violations,
    generated_ddl_files,
)
from codegen.resolve.resolver import resolve_pair as resolve_contracts
from test_m4_acceptance import _scoped

ANNOTATED = "LOB(Hardcod to MIDS)"


@pytest.mark.parametrize(("raw", "expected"), [
    (ANNOTATED, ("LOB", "Hardcod to MIDS", "MIDS")),
    ("LOB (hardcode as MIDS)", ("LOB", "hardcode as MIDS", "MIDS")),
    ("SRC_FILE_NAME(file name)", ("SRC_FILE_NAME", "file name", None)),
    ("LOB", ("LOB", None, None)),
    ("REC_CREATION_TIME", ("REC_CREATION_TIME", None, None))])
def test_split_audit_annotation(raw, expected):
    assert split_audit_annotation(raw) == expected


@pytest.fixture(scope="module")
def annotated_spec(pair4_config, tmp_path_factory):
    from acfc_shapes import FIXTURE_ROOT, pair4_nb

    tmp = tmp_path_factory.mktemp("audit_annotation")
    workbook = load_workbook(FIXTURE_ROOT / pair4_nb.STTM_PATH)
    sheet = workbook["NB_COB_REPORT"]
    assert sheet["B22"].value == "NA" and sheet["I22"].value == "REC_UPDATED_TIME"
    sheet["I22"], sheet["J22"] = ANNOTATED, "String"          # the audit row's stage column
    sttm_path = tmp / "sttm.xlsx"
    workbook.save(sttm_path)
    sttm = extract_contract(sttm_path, FIXTURE_ROOT / pair4_nb.FRD_PATH, pair4_config,
                            generated_date="2026-10-13")
    contract = tmp / "sttm.contract.json"
    contract.write_text(contract_to_json(sttm), encoding="utf-8")
    (spec,) = resolve_contracts(FIXTURE_ROOT / pair4_nb.FRD_PATH, contract, pair4_config)
    return sttm, spec


def test_the_annotated_audit_column_is_lob_with_a_cited_hardcode_hint(annotated_spec):
    sttm, _spec = annotated_spec
    names = [a.column for f in sttm.feeds for a in f.audit_columns]
    assert "LOB" in names and ANNOTATED not in names
    flags = [f for feed in sttm.feeds for f in feed.extraction_flags
             if f.startswith("audit_column_annotated:")]
    assert flags and "NB_COB_REPORT!I22" in flags[0]
    assert "hardcode hint 'Hardcod to MIDS' (value 'MIDS')" in flags[0]


@pytest.mark.parametrize("mode", ["notebook", "framework"])
def test_the_annotated_feed_generates_in_both_modes(annotated_spec, pair4_config, tmp_path,
                                                    mode):
    _sttm, spec = annotated_spec
    gate = cli._generate_feed(spec, _scoped(pair4_config, tmp_path), dry_run=True,
                              skip_tests=True, output_mode=mode,
                              conventions_profile="acfc_prx", iig_template="iig_v2")
    assert gate.verdict in ("PASS", "PASS_WITH_FLAGS")
    # No "(" in any column identifier of any generated DDL, and the annotation reaches
    # no artefact (IIG TGT_COLUMN_NAMES / TGT_DATA_TYPE, metadata_inserts.sql) either.
    out = tmp_path / "out"
    files = generated_ddl_files(out)
    assert files and ddl_identifier_violations(files) == []
    assert any("LOB" in ddl_column_identifiers(f.read_text(encoding="utf-8")) for f in files)
    for path in out.rglob("*"):
        if path.suffix in (".txt", ".sql", ".py", ".json", ".md"):
            assert "Hardcod" not in path.read_text(encoding="utf-8", errors="replace"), path
    for path in out.rglob("*_IIG.xlsx"):
        cells = [str(c) for ws in load_workbook(path).worksheets
                 for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        assert not any("Hardcod" in c for c in cells), path


@pytest.mark.parametrize(("text", "names"), [
    ("CREATE OR REPLACE TABLE c.s.t\n(\na String,\nLOB(Hardcod to X) String\n)\n",
     ["a", "LOB(Hardcod"]),
    ("CREATE TABLE IF NOT EXISTS s.t (\n    `a` STRING,\n    `LOB(x)` STRING\n) USING delta",
     ["a", "LOB(x)"])])
def test_ddl_column_identifiers(text, names):
    assert ddl_column_identifiers(text) == names


def test_the_app_default_output_for_acfc_prx_is_framework():
    config = load_config("config/config.yaml")
    assert config.output.mode == ["notebook"]                    # the CLI default is unchanged
    assert config.output.default_parts("acfc_prx") == ["framework"]
    # post91: the App offers framework only (output.app_options), whatever the profile
    assert config.output.app_options == ["framework"]
    assert config.output.default_parts("edo_sfmc") == ["framework"]
    both = config.output.model_copy(update={"app_options": ["notebook", "framework"]})
    assert both.default_parts("edo_sfmc") == ["notebook"]     # notebooks back on
    assert both.default_parts(None) == ["notebook"]
