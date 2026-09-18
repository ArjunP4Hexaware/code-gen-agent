"""M7 §2 f–i — the derivation gate: SQL literals, three-part qualification,
the vendor-only COMMENT fragment, path assembly, the single-line / charset /
cap rules over IIG cells, and the sibling-type check. Pair 11 end to end
under the strict ``acfc_prx`` profile; pair 1 stays flag-free; CV / SFMC
(``edo_sfmc``) are untouched (their baselines are byte-compared elsewhere).
"""

from __future__ import annotations

from pathlib import Path

import pytest

from codegen import cli
from codegen.config import load_config
from codegen.gate.derivations import (
    check_iig_derivations,
    check_sql_literals,
    join_path,
    path_violations,
    short_token,
    sibling_type_flags,
)

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
STTM = SHAPES / "sttm" / "pair_11_family_b.xlsx"
FRD = SHAPES / "frd" / "f1_pair_11_multi_file.docx"
FRD_CATALOG = SHAPES / "frd" / "f1_pair_11_multi_file_catalog.docx"
DATE = "2026-01-01"


def _pair11_specs(frd_path: Path, tmp: Path, config):
    from codegen.extract import contract_to_json as sttm_to_json
    from codegen.extract import extract_contract
    from codegen.extract.frd_docx import contract_to_json as frd_to_json
    from codegen.layout.model import MockLayoutProvider
    from codegen.layout.resolve import resolve_pair as layout_pair
    from codegen.resolve.resolver import resolve_pair

    pair = layout_pair(STTM, frd_path, config, provider=MockLayoutProvider([]), cache_dirs=[],
                       use_cache=False, generated_date=DATE)
    assert pair.questions == []
    frd_json = tmp / "frd.contract.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    contract = extract_contract(STTM, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    sttm_json = tmp / "sttm.contract.json"
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    return resolve_pair(frd_json, sttm_json, config), pair.flags


def _scoped(config, tmp: Path):
    return config.model_copy(update={"output": config.output.model_copy(update={
        "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")})})


@pytest.fixture(scope="module")
def config():
    return load_config(REPO / "config" / "config.yaml")


@pytest.fixture(scope="module")
def pair11(config, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pair11")
    specs, layout_flags = _pair11_specs(FRD, tmp, config)
    return specs, layout_flags


@pytest.fixture(scope="module")
def pair11_catalog(config, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pair11_catalog")
    specs, layout_flags = _pair11_specs(FRD_CATALOG, tmp, config)
    return specs, layout_flags


# -- unit rules ----------------------------------------------------------------- #


def test_join_path_normalizes_segments_and_never_concatenates_raw_values():
    assert join_path("abfss://c@a.dfs.example/", "mftlanding\\inbound\\dom_a\\", "/tbl") == \
        "abfss://c@a.dfs.example/mftlanding/inbound/dom_a/tbl"
    assert join_path("/a//b/", None, "c") == "a/b/c"
    # a segment ending in punctuation is kept for the GATE to fail, not silently fixed
    assert path_violations(join_path("x", "vendor_c."), 1024) == [
        "segment 'vendor_c.' ends in punctuation"]
    assert path_violations("/a b/c", 1024) == ["whitespace"]
    assert "empty segment" in path_violations("/a//c", 1024)
    assert path_violations("/dom_a/public/vendor_c", 1024) == []
    assert path_violations("x" * 20, 10) == ["length 20 > 10"]


def test_short_token_rejects_prose():
    assert short_token("Socially Determined - SD 33732", 32) == "Socially_Determined_SD_33732"
    prose = "Vendor Files = Enrollment, Disenrollment & Individual Risk Reports"
    assert short_token(prose, 32) is None
    assert short_token("two\nlines", 32) is None and short_token("", 32) is None


def test_sql_literal_with_a_newline_fails_citing_the_literal(config):
    text = ("CREATE OR REPLACE TABLE c.s.t (\n  A String COMMENT 'ok')\nUSING delta\n"
            "COMMENT 'Stage target table for feed x\nline two'\nLOCATION 'abfss://c@a/x/ y'")
    check = check_sql_literals([("t.txt", text)], config)
    assert not check.passed
    assert "t.txt: COMMENT literal" in check.details and "line break" in check.details
    assert "LOCATION literal" in check.details and "whitespace" in check.details
    clean = check_sql_literals([("t.txt", "COMMENT 'it''s fine' LOCATION 'abfss://c@a/x/y'")],
                               config)
    assert clean.passed


def test_iig_cells_single_line_charset_and_path_rules(config):
    payload = {"tabs": {"ADLS_DELTA_INGESTION_DETAILS": {"headers": [], "rows": [{
        "values": {"TGT_TABLE_NAME": "vc enrollment", "SRC_ADLS_PATH": "/dom_a/vendor_c.",
                   "DOMAIN": "Domain Alpha\nline two", "PIPELINE_NAME": "WF_OK_NAME"},
        "badges": {"TGT_TABLE_NAME": {"badge": "from FRD"},
                   "SRC_ADLS_PATH": {"badge": "from FRD", "tooltip": "FRD table 5 row 13"},
                   "DOMAIN": {"badge": "from FRD"}, "PIPELINE_NAME": {"badge": "from standards"}},
    }]}}}
    check = check_iig_derivations(payload, config)
    assert not check.passed
    assert "TGT_TABLE_NAME row 2 (from FRD): 'vc enrollment' is not an identifier" in check.details
    assert ("SRC_ADLS_PATH row 2 (from FRD — FRD table 5 row 13): segment 'vendor_c.' ends in "
            "punctuation") in check.details
    assert "DOMAIN row 2 (from FRD): line break" in check.details
    assert "PIPELINE_NAME" not in check.details


# -- pair 11 end to end ---------------------------------------------------------- #


def test_pair11_no_catalog_fails_qualified_names_and_writes_no_ddl(config, pair11, tmp_path):
    specs, _flags = pair11
    scoped = _scoped(config, tmp_path)
    gate = cli._generate_feed(specs[0], scoped, dry_run=True, skip_tests=True,
                              output_mode="framework", conventions_profile="acfc_prx",
                              iig_template="iig_v2")
    assert gate.verdict == "FAIL"
    failed = {c.name: c.details for c in gate.checks if not c.passed}
    assert set(failed) == {"qualified_names"}
    assert "catalog_unstated:stage" in failed["qualified_names"]
    assert "catalog_unstated:standard" in failed["qualified_names"]
    assert "FRD Target Catalog and Schema label, STTM target band catalog column" in \
        failed["qualified_names"]
    framework_dir = tmp_path / "out" / specs[0].feed_slug / "framework"
    assert not list(framework_dir.glob("*_DDL.txt"))            # a two-part name is never written
    assert (framework_dir / "config_rows.xlsx").is_file()      # the review artefacts still land


def test_pair11_with_catalog_emits_three_part_ddl_and_own_paths(config, pair11_catalog, tmp_path):
    specs, _flags = pair11_catalog
    assert [s.feed_slug for s in specs] == ["vc_enrollment", "vc_disenrollment",
                                            "vc_individual_risk"]
    scoped = _scoped(config, tmp_path)
    # combined layout (acfc_prx): every CREATE is catalog.schema.table
    gate = cli._generate_feed(specs[2], scoped, dry_run=True, skip_tests=True,
                              output_mode="framework", conventions_profile="acfc_prx",
                              iig_template="iig_v2")
    assert gate.verdict == "PASS_WITH_FLAGS", [c for c in gate.checks if not c.passed]
    framework_dir = tmp_path / "out" / "vc_individual_risk" / "framework"
    (ddl,) = list(framework_dir.glob("*_DDL.txt"))
    text = ddl.read_text(encoding="utf-8")
    assert "CREATE OR REPLACE TABLE cat_syn.stg_dom_b.vc_individual_risk" in text
    assert "CREATE OR REPLACE TABLE cat_syn.dom_b.vc_individual_risk" in text
    assert "\n\n\n" not in text.replace("\n\n--", "")
    # reference two-file layout on the same feed: LOCATION is this feed's own path, one line
    gate2 = cli._generate_feed(specs[2], _scoped(config, tmp_path / "two"), dry_run=True,
                               skip_tests=True, output_mode="framework",
                               conventions_profile="edo_sfmc", iig_template="iig_v1")
    stage_txt = (tmp_path / "two" / "out" / "vc_individual_risk" / "framework"
                 / "vc_individual_risk_stage_table_creation.txt").read_text(encoding="utf-8")
    location = next(line for line in stage_txt.splitlines() if line.startswith("LOCATION '"))
    assert location.endswith("/mftlanding/inbound/dom_b/dom_a/vendor_c/vc_individual_risk'")
    assert "dom_a/public" not in location                       # not another feed's path
    # COMMENT: vendor label only, no label-prefixed description
    comment = next(line for line in stage_txt.splitlines() if line.startswith("COMMENT '"))
    assert "(source: VENDOR_C" in comment and "=" not in comment
    assert gate2.verdict == "PASS_WITH_FLAGS"


def test_pair11_sibling_type_mismatch_is_exactly_one_flag_citing_the_row(config, pair11_catalog):
    specs, _flags = pair11_catalog
    enrollment = specs[0]
    flags = sibling_type_flags(enrollment, config)
    assert len(flags) == 1
    (flag,) = flags
    assert flag.startswith("sibling_type_mismatch:_pct — vc_enrollment: 3 columns end in _pct, "
                           "2 typed decimal, minority: UNINSURED_PCT String "
                           "(MAPPING-VC_ENROLLMENT row")
    # _pop: two Decimal siblings agree; _dt: one column; no other group
    assert sibling_type_flags(specs[1], config) == []
    assert sibling_type_flags(specs[2], config) == []


def test_pair11_iig_cells_are_single_line_and_names_under_cap(config, pair11_catalog, tmp_path):
    specs, _flags = pair11_catalog
    scoped = _scoped(config, tmp_path)
    for spec in specs:
        gate = cli._generate_feed(spec, scoped, dry_run=True, skip_tests=True,
                                  output_mode="framework", conventions_profile="acfc_prx",
                                  iig_template="iig_v2")
        assert gate.verdict == "PASS_WITH_FLAGS", [c for c in gate.checks if not c.passed]
        assert {c.name for c in gate.checks} >= {"derivations", "sql_literals"}
        from openpyxl import load_workbook

        wb = load_workbook(tmp_path / "out" / spec.feed_slug / "framework" / "config_rows.xlsx")
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for value in row:
                    assert not (isinstance(value, str) and "\n" in value), (ws.title, value)
        sheet = wb["DATA_FACTORY_PIPELINE_SCHEDULE"]
        headers = [c.value for c in sheet[1]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name = row[headers.index("PIPELINE_NAME")]
            assert name is None or len(name) <= 128


def test_pair1_raises_no_derivation_flags(pair1_config, pair1_spec):
    """Pair 1 under the strict profile: no suffix group mixes, every name
    three-part, every literal single-line — zero new flags or failed checks."""
    assert sibling_type_flags(pair1_spec, pair1_config) == []
