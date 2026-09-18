"""M7 §3 — the SQL Server metadata-DB DML deliverable (emit/dml.py):
variables block with assigner + rule, NULL -- ASSIGN + dml_unassigned when
the FAQ carries no id, FAQ-supplied ids fill the DECLAREs, preflight with
RAISERROR, connection insert-or-reuse with SCOPE_IDENTITY(), one INSERT per
IIG row in dependency order, audit columns = @RFC_NUMBER / GETDATE(),
ACTIVE_FLAG 'S', no line break in any literal, environments differ only in
the variables block, every statement parses under sqlglot's tsql dialect,
row counts equal the IIG's, and the runner notebook names secrets only."""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from codegen import cli
from codegen.config import load_config
from codegen.emit.dml import validate_tsql
from codegen.emit.framework import emit_framework
from codegen.faq import FaqAnswer, LoadPatternFaq

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"

pytest.importorskip("sqlglot", reason="sqlglot validates the emitted T-SQL")


def _scoped(config, tmp: Path):
    return config.model_copy(update={"output": config.output.model_copy(update={
        "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")})})


def _variables_block(text: str) -> str:
    start = text.index("-- VARIABLES")
    end = text.index("-- PREFLIGHT")
    return text[start:end]


def _body(text: str) -> str:
    return text[text.index("-- PREFLIGHT"):]


@pytest.fixture(scope="module")
def pair1_dml(pair1_config, pair1_spec, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pair1_dml")
    gate = cli._generate_feed(pair1_spec, _scoped(pair1_config, tmp), dry_run=True,
                              skip_tests=True, output_mode="framework",
                              conventions_profile="acfc_prx", iig_template="iig_v2")
    return gate, tmp / "out" / pair1_spec.feed_slug / "framework"


def test_pair1_dml_files_variables_and_unassigned_ids(pair1_dml, pair1_config):
    gate, framework_dir = pair1_dml
    assert gate.verdict == "PASS_WITH_FLAGS", [c for c in gate.checks if not c.passed]
    for env in ("q1", "a2", "prod"):
        assert (framework_dir / f"config_inserts_{env}.sql").is_file()
        assert (framework_dir / f"Insert_scripts_config_table_{env}.py").is_file()
    text = (framework_dir / "config_inserts_q1.sql").read_text(encoding="utf-8")
    assert text.startswith("-- TARGET SYSTEM = SQL Server metadata DB")
    assert "-- RUN FROM = Databricks notebook" in text
    block = _variables_block(text)
    for name in ("RFC_NUMBER", "PIPELINE_ID", "PARENT_PIPELINE_ID", "GROUP_ID", "OBJECT_ID",
                 "SRC_CONNECTION_ID", "SRC_ADLS_CONNECTION_ID", "METADATA_CONNECTION_ID",
                 "TGT_CONNECTION_ID"):
        line = next(ln for ln in block.splitlines() if ln.startswith(f"DECLARE @{name} "))
        assert "= NULL;" in line and "-- ASSIGN" in line
        assert "assigned by:" in line and "rule:" in line
        assert any(f.startswith(f"dml_unassigned:@{name}") for f in gate.flags), name
    # the FRD's landing path is the connection root (a value, cited)
    assert "DECLARE @SRC_ROOT_PATH NVARCHAR(400) = N'/Pharmacy/Accumulators/';" in block
    # preflight: unused ids, 0-or-1 connection, insert-or-reuse with SCOPE_IDENTITY
    body = _body(text)
    assert ("IF EXISTS (SELECT 1 FROM [dbo].[DATA_FACTORY_PIPELINE_SCHEDULE] WHERE [PIPELINE_ID] "
            "= @PIPELINE_ID) RAISERROR(") in body
    assert "IF @CONNECTION_MATCHES > 1 RAISERROR(" in body
    assert "SET @SRC_CONNECTION_ID = SCOPE_IDENTITY();" in body
    # audit + active conventions on every INSERT that has the columns
    inserts = [ln for ln in body.splitlines() if ln.startswith("INSERT INTO [dbo].[")]
    assert inserts and all("@RFC_NUMBER, GETDATE(), @RFC_NUMBER, GETDATE()" in ln
                           for ln in inserts if "[CREATED_BY]" in ln)
    assert all("N'S'" in ln for ln in inserts if "[ACTIVE_FLAG]" in ln)
    assert not any("N'Y'" in ln and "[ACTIVE_FLAG]" in ln for ln in inserts)
    # dependency order: pipeline schedule first, then file->ADLS, then ADLS->Delta
    order = [m.group(1) for m in re.finditer(r"^-- ([A-Z_]+): \d+ row\(s\)", body, re.M)]
    assert order[:3] == ["DATA_FACTORY_PIPELINE_SCHEDULE", "FILE_ADLS_INGESTION_DETAILS",
                         "ADLS_DELTA_INGESTION_DETAILS"]
    assert any("table not yet described" in ln for ln in body.splitlines())
    assert any(f.startswith("dml_unconfirmed:connection_table") for f in gate.flags)


def test_pair1_dml_row_counts_equal_the_iig_and_statements_parse(pair1_dml):
    gate, framework_dir = pair1_dml
    from openpyxl import load_workbook

    wb = load_workbook(framework_dir / "config_rows.xlsx")
    iig_counts = {ws.title: ws.max_row - 1 for ws in wb.worksheets
                  if not ws.title.startswith("_") and ws.max_row > 1}
    text = (framework_dir / "config_inserts_q1.sql").read_text(encoding="utf-8")
    dml_counts = {m.group(1): int(m.group(2))
                  for m in re.finditer(r"^-- ([A-Z_]+): (\d+) row\(s\)", text, re.M)}
    assert dml_counts == iig_counts
    assert {c.name for c in gate.checks} >= {"dml_parse_q1", "dml_parse_a2", "dml_parse_prod",
                                              "dml_row_counts"}
    assert all(c.passed for c in gate.checks if c.name.startswith("dml_"))
    for env in ("q1", "a2", "prod"):
        assert validate_tsql((framework_dir / f"config_inserts_{env}.sql")
                             .read_text(encoding="utf-8")) == []


def test_environments_differ_only_in_the_variables_block(pair1_dml):
    _gate, framework_dir = pair1_dml
    q1 = (framework_dir / "config_inserts_q1.sql").read_text(encoding="utf-8")
    prod = (framework_dir / "config_inserts_prod.sql").read_text(encoding="utf-8")
    assert _body(q1) == _body(prod)
    assert "DECLARE @ENV NVARCHAR(16) = N'q1';" in q1 and "N'prod';" in prod
    # path columns go through the environment prefix, so the body can stay identical
    assert "CONCAT(@PATH_PREFIX, N'" in _body(q1)


def test_no_literal_spans_lines_and_none_carries_pointer_text(pair1_dml):
    _gate, framework_dir = pair1_dml
    text = (framework_dir / "config_inserts_q1.sql").read_text(encoding="utf-8")
    for match in re.finditer(r"N'((?:[^']|'')*)'", text):
        assert "\n" not in match.group(1)
        assert "refer to" not in match.group(1).lower()


def test_faq_supplied_ids_fill_the_declares_and_raise_no_unassigned_flags(pair1_config,
                                                                          pair1_spec, tmp_path):
    faq = LoadPatternFaq(
        rfc_number=FaqAnswer(value="SYN001", source="engineer", evidence="ticket"),
        feed_abbreviation=FaqAnswer(value="ACCUM", source="engineer"),
        pipeline_id=FaqAnswer(value="9101", source="engineer", evidence="assigned 2026-09"),
        parent_pipeline_id=FaqAnswer(value="0", source="engineer"),
        group_id=FaqAnswer(value="9201", source="engineer"),
        object_id=FaqAnswer(value="1", source="engineer"),
        source_host=FaqAnswer(value="mft.synthetic.example", source="engineer",
                              evidence="platform team"),
        connection_ids={"SRC_CONNECTION_ID": "31", "SRC_ADLS_CONNECTION_ID": "32",
                        "METADATA_CONNECTION_ID": "33", "TGT_CONNECTION_ID": "34"},
    )
    framework = emit_framework(pair1_spec, faq, [], pair1_config, tmp_path,
                               conventions_profile="acfc_prx", iig_template="iig_v2")
    assert not any(f.startswith("dml_unassigned:") for f in framework.flags)
    sql = next(p for p in framework.files if p.name == "config_inserts_q1.sql")
    block = _variables_block(sql.read_text(encoding="utf-8"))
    assert "DECLARE @RFC_NUMBER NVARCHAR(50) = N'SYN001';" in block
    assert "DECLARE @PIPELINE_ID INT = 9101;" in block and "-- ASSIGN" not in block.split(
        "DECLARE @PIPELINE_ID")[1].split("\n")[0]
    assert "DECLARE @GROUP_ID INT = 9201;" in block
    assert "DECLARE @SRC_CONNECTION_ID INT = 31;" in block
    assert "DECLARE @SRC_HOST_NAME NVARCHAR(200) = N'mft.synthetic.example';" in block
    assert all(c.passed for c in framework.checks if c.name.startswith("dml_"))


def test_notebook_source_names_secrets_and_never_holds_a_credential(pair1_dml, pair1_config):
    _gate, framework_dir = pair1_dml
    text = (framework_dir / "Insert_scripts_config_table_q1.py").read_text(encoding="utf-8")
    assert text.startswith("# Databricks notebook source\n# TARGET SYSTEM = SQL Server metadata "
                           "DB\n# RUN FROM = Databricks notebook\n")
    dml = pair1_config.dml
    assert f'SECRET_SCOPE = "{dml.secret_scope}"' in text
    assert "dbutils.secrets.get(SECRET_SCOPE, KEY_NAME_JDBC_URL)" in text
    assert "connection.setAutoCommit(False)" in text and "connection.rollback()" in text
    assert "SQL_NAME = \"config_inserts_q1.sql\"" in text
    assert not re.search(r"password\s*=\s*['\"][^'\"]+['\"]", text, re.IGNORECASE)


def test_reference_profile_writes_no_dml(pair1_config, pair1_spec, tmp_path):
    framework = emit_framework(pair1_spec, LoadPatternFaq(), [], pair1_config, tmp_path,
                               conventions_profile="edo_sfmc", iig_template="iig_v1")
    assert not any(p.suffix == ".sql" for p in framework.files)
    assert not any(f.startswith("dml_") for f in framework.flags)


def test_multi_file_frd_dml_takes_the_sttm_frequency_and_no_pointer_text(tmp_path):
    from test_derivations import FRD_CATALOG, _pair11_specs

    config = load_config(REPO / "config" / "config.yaml")
    (tmp_path / "contracts").mkdir()
    specs, _flags = _pair11_specs(FRD_CATALOG, tmp_path / "contracts", config)
    risk = next(s for s in specs if s.feed_slug == "vc_individual_risk")
    gate = cli._generate_feed(risk, _scoped(config, tmp_path), dry_run=True, skip_tests=True,
                              output_mode="framework", conventions_profile="acfc_prx",
                              iig_template="iig_v2")
    assert gate.verdict == "PASS_WITH_FLAGS", [c for c in gate.checks if not c.passed]
    text = (tmp_path / "out" / "vc_individual_risk" / "framework" / "config_inserts_q1.sql"
            ).read_text(encoding="utf-8")
    literals = [m.group(1) for m in re.finditer(r"N'((?:[^']|'')*)'", text)]
    assert "Monthly" in literals                       # FREQUENCY from the STTM File Details
    assert not any("refer to" in lit.lower() or "\n" in lit for lit in literals)
    assert "Vendor Files =" not in text                # the label-prefixed description never lands
