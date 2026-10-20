"""The EDH E2E gold-mapping reader (codegen.layouts.edh_e2e) on the synthetic fixtures.

Every value here is fabricated (tests/fixtures/edh_gold_synthetic/, SYN_ names only). The rule
under test is the ingestion layouts' one: stated -> filled, derived -> proposed, unstated or
free text -> QUESTION — and never silence.
"""

from __future__ import annotations

import importlib.util
import json
import subprocess
import sys
from pathlib import Path

import pytest

from codegen.config import EdhE2eConfig
from codegen.contracts.gold import ROLES, GoldContract
from codegen.layouts import edh_e2e
from codegen.layouts.edh_e2e import Untranslatable, read_dir, role_counts, translate_edit

REPO = Path(__file__).resolve().parent.parent
FIXTURES = REPO / "tests" / "fixtures" / "edh_gold_synthetic"
QUALIFY = {"syn_a": "`SYN_T`.`SYN_A`"}


@pytest.fixture(scope="module")
def contract() -> GoldContract:
    return read_dir(FIXTURES, EdhE2eConfig())


def _feed(contract, json_prefix):
    return next(f for f in contract.feeds if f.json_file.startswith(json_prefix))


def _keys(feed, role=None):
    return {q.key for q in feed.questions if role is None or q.role == role}


def _proposed(feed, role=None):
    return {p.key for p in feed.proposals if role is None or p.role == role}


# ---- every value is filled, proposed or asked ---------------------------------------------


def test_every_mapping_row_gets_an_expression_or_a_question(contract):
    for feed in contract.feeds:
        for j, col in enumerate(feed.columns):
            key = f"feeds[{feed.index}].columns[{j}]"
            asked = {f"{key}.expression", f"{key}.source_table"} & _keys(feed)
            assert (col.expression is not None) != bool(asked), key
            if col.expression is not None:
                assert col.expression.startswith(
                    f"/* {feed.json_file} block_index={col.block_index} "
                    f"source_row={col.source_row} */ "
                )


def test_every_mapping_row_gets_a_type_or_a_question(contract):
    for feed in contract.feeds:
        for j, col in enumerate(feed.columns):
            key = f"feeds[{feed.index}].columns[{j}].target_type"
            assert (col.target_type is not None) != (key in _keys(feed)), key
            if col.target_type_status == "proposed":
                assert key in _proposed(feed)


def test_every_selection_criteria_cell_becomes_a_join_filter_or_question(contract):
    for feed in contract.feeds:
        cells = {}
        for col in feed.columns:
            text = next(
                r["selection_criteria"]
                for r in _rows(feed)
                if r["source_row"] == col.source_row and r["block_index"] == col.block_index
            )
            if text:
                cells.setdefault((col.block_index, text), []).append(col.source_row)
        for (block, text), rows in cells.items():
            sset = next(s for s in feed.source_sets if s.block_index == block)
            used = {r for j in sset.joins for r in j.rows} | {
                r for f in sset.filters for r in f.rows
            }
            asked = {
                q.where.source_row
                for q in feed.questions
                if q.quoted == text and q.where is not None
            }
            assert rows[0] in used | asked, (feed.json_file, block, rows[0])


def _rows(feed):
    doc = json.loads((FIXTURES / feed.json_file).read_text(encoding="utf-8"))
    return doc["mappings"]


def test_feed_level_roles_are_filled_or_asked(contract):
    for feed in contract.feeds:
        k = f"feeds[{feed.index}]"
        asked = _keys(feed)
        assert (feed.load_mode is not None) != (f"{k}.load_mode" in asked)
        assert (feed.target.schema_name is None) == (f"{k}.target_schema" in asked)
        assert f"{k}.scd2_mode" in asked and f"{k}.sk_strategy" in asked
        assert bool(feed.primary_key) != (f"{k}.primary_key" in asked)
        for s in feed.source_sets:
            for n, src in enumerate(s.sources):
                key = f"{k}.source_sets[{s.block_index}].sources[{n}].catalog_schema"
                assert (src.catalog_schema is None) == (key in asked)


def test_document_notes_and_rules_are_one_question_each(contract):
    for name in {f.json_file for f in contract.feeds}:
        feeds = [f for f in contract.feeds if f.json_file == name]
        doc = json.loads((FIXTURES / name).read_text(encoding="utf-8"))
        # asked once per FILE, however many target feeds the file yields
        assert sum(len(_keys(f, "document_note")) for f in feeds) == len(
            doc["document_notes"] or []
        )
        assert len(_keys(feeds[0], "rule")) == len(doc["rules"] or [])
        assert not any(_keys(f, "rule") for f in feeds[1:])
        for q in feeds[0].questions:
            if q.role == "rule":
                assert q.quoted in doc["rules"]


def test_every_counted_role_is_in_the_vocabulary(contract):
    for feed in contract.feeds:
        assert set(role_counts(feed)) <= set(ROLES)


# ---- the profile fills what the documents never state -------------------------------------


def test_profile_values_fill_instead_of_asking():
    cfg = EdhE2eConfig(
        target_schema="syn_gold",
        source_catalog_map={"SYNFCT": "syn_cat.syn_src"},
        scd2_mode="syn_type2",
        sk_strategy="syn_identity",
        function_map={"SynBlank": "''"},
    )
    feed = read_dir(FIXTURES, cfg).feeds[0]
    asked = _keys(feed)
    assert not {
        r for r in ("target_schema", "scd2_mode", "sk_strategy") if f"feeds[0].{r}" in asked
    }
    assert not _keys(feed, "source_catalog_schema")
    assert feed.target.schema_name == "syn_gold"
    assert {s.catalog_schema for ss in feed.source_sets for s in ss.sources} == {"syn_cat.syn_src"}
    assert role_counts(feed)["scd2_mode"]["filled"] == 1
    # the mapped zero-argument function lets both empty / null tests translate
    exprs = [c.expression for c in feed.columns if c.edit_class == "COALESCE-equivalent"]
    assert exprs and all(e is not None for e in exprs)


def test_without_a_function_mapping_the_expression_is_asked_quoting_the_cell(contract):
    feed = _feed(contract, "syn_mbr_dim")
    q = [q for q in feed.questions if q.role == "expression" and "SynBlank" in (q.quoted or "")]
    assert q and all("edh_e2e.function_map" in x.reason for x in q)


# ---- the multi-block case -----------------------------------------------------------------


def test_two_blocks_yield_two_source_sets_for_one_target(contract):
    feed = _feed(contract, "syn_mbr_dim")
    assert [s.block_index for s in feed.source_sets] == [0, 1]
    first, second = ({src.table for src in s.sources} for s in feed.source_sets)
    assert first and second and not first & second
    assert {c.target_attribute for c in feed.columns if c.block_index == 1} < {
        c.target_attribute for c in feed.columns if c.block_index == 0
    }


# ---- joins / filters / keys / regions ------------------------------------------------------


def test_join_grammar(contract):
    feed = _feed(contract, "syn_elig_fact")
    joins = feed.source_sets[0].joins
    same_name = [j for j in joins if j.on_assumption == "same_name"]
    assert same_name and all(j.status == "proposed" and j.type == "INNER" for j in same_name)
    assert len(_proposed(feed, "join")) == len(same_name)
    # the second joined table of one cell is a second join with the same column
    by_rows = {}
    for j in same_name:
        by_rows.setdefault(tuple(j.rows), []).append(j.joined)
    assert any(len(v) == 2 for v in by_rows.values())
    assert any(len(j.on) == 2 for j in same_name)  # "using <col> & <col2>"
    explicit = [j for j in joins if j.on_assumption == "explicit"]
    assert {j.type for j in explicit} == {"INNER", "LEFT"}  # "*=" is the outer side
    assert all(j.pairs and j.base is None for j in explicit)
    assert len([q for q in feed.questions if q.key.endswith(".tables")]) == len(explicit)


def test_filter_grammar(contract):
    feed = _feed(contract, "syn_elig_fact")
    by_pattern = {}
    for f in feed.source_sets[0].filters:
        by_pattern.setdefault(f.pattern, []).append(f)
    assert set(by_pattern) == {"F1", "F2", "F3", "F4", "F5"}
    assert "date_sub(current_date(), 30)" in by_pattern["F3"][0].predicate
    assert by_pattern["F4"][0].status == "proposed"
    assert " AND " in by_pattern["F5"][0].predicate
    assert len(_keys(feed, "filter")) == 3  # the three free-text cells


def test_keys(contract):
    feeds = [f for f in contract.feeds if f.json_file.startswith("syn_memo")]
    assert feeds[0].primary_key
    other = [q for f in feeds for q in f.questions if q.role == "primary_key" and q.quoted]
    assert len(other) == 2  # the two non-PK indexes cells, quoted
    no_pk = [f for f in feeds if not f.primary_key]
    assert all(f"feeds[{f.index}].primary_key" in _keys(f) for f in no_pk)


def test_regions_drift_is_a_proposal_and_sentinels_never_fill(contract):
    feed = _feed(contract, "syn_elig_fact")
    assert _proposed(feed, "region_drift")
    assert len({r.label for r in feed.regions}) == len(feed.regions)
    assert "edh_region_column_not_identifier:1" in feed.flags
    empty = _feed(contract, "syn_plan_ref")
    assert not empty.regions and len(_keys(empty, "target_type")) == len(empty.columns)


def test_drag_fill_is_flagged_and_kept(contract):
    feed = _feed(contract, "syn_memo")
    flagged = [x for x in feed.flags if x.startswith("edh_drag_fill:")]
    assert len(flagged) == 3
    j = int(flagged[0].split("columns[")[1].split("]")[0])
    assert feed.columns[j].source_format.upper().startswith("VARCHAR")


# ---- the expression translator (fabricated cells) ------------------------------------------


@pytest.mark.parametrize(
    "cell, sql",
    [
        ("set to NULL", "NULL"),
        ("set NULL", "NULL"),
        ("current timestamp", "current_timestamp()"),
        ("SYN_A[2,5]", "substr(`SYN_T`.`SYN_A`, 2, 4)"),
        ("TRIM(SYN_B)[1,3]", "substr(trim(`SYN_B`), 1, 3)"),
        ("if SYN_A IS NULL THEN 'X' ELSE SYN_A", "coalesce(`SYN_T`.`SYN_A`, 'X')"),
        ("if ISNULL(SYN_A) THEN 'X' ELSE SYN_A", "coalesce(`SYN_T`.`SYN_A`, 'X')"),
        (
            "if SYN_A IS a empty string OR NULL THEN 'X' ELSE SYN_A",
            "coalesce(nullif(`SYN_T`.`SYN_A`, ''), 'X')",
        ),
        (
            "if SYN_A IS an empty string THEN 'X' ELSE SYN_A",
            "CASE WHEN `SYN_T`.`SYN_A` = '' THEN 'X' ELSE `SYN_T`.`SYN_A` END",
        ),
        (
            "if SYN_B = 'P' THEN 'Q' ELSE if SYN_B = 'R' THEN 'S' ELSE 'T'",
            "CASE WHEN `SYN_B` = 'P' THEN 'Q' WHEN `SYN_B` = 'R' THEN 'S' ELSE 'T' END",
        ),
        ("SYN_B : '-' : SYN_C[1,2]", "concat(`SYN_B`, '-', substr(`SYN_C`, 1, 2))"),
        ("CAST(SYN_B AS DECIMAL(10,2))", "CAST(`SYN_B` AS DECIMAL(10,2))"),
        (
            "if SYN_B = < SYN_C THEN SYN_B ELSE SYN_C",
            "CASE WHEN `SYN_B` <= `SYN_C` THEN `SYN_B` ELSE `SYN_C` END",
        ),
    ],
)
def test_translated_classes(cell, sql):
    assert translate_edit(cell, qualify=QUALIFY, functions={}) == sql


@pytest.mark.parametrize(
    "cell",
    [
        "TRIM(CONVERT('SYN_X', 'SYN_Y', s.SYN_B))[1,9]",  # character-set arguments
        "if SYN_B IS NULL THEN SynFn() ELSE SYN_B",  # an unmapped function
        "only for SYN_REGION .",  # free text
        "SYN_B > 0",  # a predicate, not an expression
        "if (SYN_B >= SYN_C) AND (SYN_B = < SYN_D)",  # an if without THEN
    ],
)
def test_untranslatable_cells_raise(cell):
    with pytest.raises(Untranslatable):
        translate_edit(cell, qualify=QUALIFY, functions={})


# ---- the CLI and the shared grammar --------------------------------------------------------


def test_cli_prints_answer_lines_and_exits_3(tmp_path):
    out = tmp_path / "gold.json"
    done = subprocess.run(
        [
            sys.executable,
            "-m",
            "codegen.cli",
            "read-edh-e2e",
            "--input",
            str(FIXTURES),
            "--out",
            str(out),
            "--table",
            "1",
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=REPO,
    )
    assert done.returncode == 3, done.stderr
    contract = GoldContract.model_validate_json(out.read_text(encoding="utf-8"))
    assert [f.index for f in contract.feeds] == [1]
    lines = [ln for ln in done.stdout.splitlines() if ln.startswith("QUESTION")]
    assert len(lines) == len(contract.feeds[0].questions)
    assert all(ln[15:].startswith("feeds[1].") and " — " in ln for ln in lines)


def test_reader_and_profiler_share_the_grammar():
    spec = importlib.util.spec_from_file_location(
        "profile_edh", REPO / "scripts" / "gold" / "profile_edh.py"
    )
    profiler = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(profiler)
    for name in ("J_USING", "J_COLUMN_NO_USING", "J_EQUALITIES", "J_TABLE_ONLY", "PREAMBLE"):
        assert getattr(profiler, name).pattern == getattr(edh_e2e, name).pattern, name
    assert [(n, rx.pattern) for n, rx in profiler.F_PATTERNS] == [
        (n, rx.pattern) for n, rx in edh_e2e.F_PATTERNS
    ]


# ---- v2.2: one feed per target table ------------------------------------------------------


def test_a_file_whose_blocks_load_two_tables_yields_two_feeds(contract):
    split = [f for f in contract.feeds if f.json_file.startswith("syn_split")]
    assert len(split) == 2
    assert len({f.manifest_index for f in split}) == 1
    assert len({f.target.table for f in split}) == 2
    assert [f.blocks for f in split] == [[0], [1]]
    assert [len(f.source_sets) for f in split] == [1, 1]
    # the mixed_target note is asked once, by the feed of the lowest block it names
    assert [len(_keys(f, "document_note")) for f in split] == [2, 0]
    indexes = [f.index for f in contract.feeds]
    assert indexes == list(range(len(contract.feeds)))


def test_feed_numbering_is_the_same_for_one_table(contract):
    split = [f for f in contract.feeds if f.json_file.startswith("syn_split")]
    alone = read_dir(FIXTURES, EdhE2eConfig(), only=split[0].manifest_index)
    assert [f.index for f in alone.feeds] == [f.index for f in split]
    assert [q.key for f in alone.feeds for q in f.questions] == [
        q.key for f in split for q in f.questions
    ]


def test_lengths_are_proposed_from_target_format(contract):
    feed = next(f for f in contract.feeds if f.json_file.startswith("syn_split"))
    lengths = [(c.length, c.length_status) for c in feed.columns]
    assert lengths == [("12", "proposed"), ("40", "proposed"), (None, None)]  # DATE: none
    assert len(_proposed(feed, "length")) == 2
    stated = _feed(contract, "syn_plan_ref")
    assert all(c.length_status == "filled" for c in stated.columns)
    assert role_counts(stated)["length"]["filled"] == len(stated.columns)


# ---- the questions workbook and the answers that come back ---------------------------------


def test_questions_workbook_has_one_sheet_per_feed_and_one_row_per_question(contract):
    from openpyxl import load_workbook

    data, rows = edh_e2e.questions_workbook(contract)
    assert rows == sum(len(f.questions) for f in contract.feeds)
    import io

    book = load_workbook(io.BytesIO(data), read_only=True)
    assert book.sheetnames == [f"feed_{f.index}" for f in contract.feeds]
    for feed, sheet in zip(contract.feeds, book.worksheets, strict=True):
        lines = list(sheet.iter_rows(values_only=True))
        assert lines[0] == edh_e2e.QUESTION_COLUMNS
        assert [ln[0] for ln in lines[1:]] == [q.key for q in feed.questions]
        assert all(ln[-1] is None for ln in lines[1:])  # answers left empty


def _fill(data: bytes, answers: dict[str, str]) -> bytes:
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data))
    for sheet in book.worksheets:
        for row in sheet.iter_rows(min_row=2):
            if row[0].value in answers:
                row[6].value = answers[row[0].value]
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def test_a_filled_workbook_closes_its_questions(contract):
    feed = _feed(contract, "syn_mbr_dim")
    k = f"feeds[{feed.index}]"
    expr_key = next(q.key for q in feed.questions if q.role == "expression")
    given = {
        f"{k}.target_schema": "syn_gold",
        f"{k}.scd2_mode": "syn_type2",
        expr_key: "upper(`SYN_X`)",
    }
    data, _ = edh_e2e.questions_workbook(contract)
    answers = edh_e2e.answers_from_workbook(_fill(data, given))
    assert answers == given
    after = read_dir(FIXTURES, EdhE2eConfig(), answers=answers)
    again = next(f for f in after.feeds if f.index == feed.index)
    assert not set(given) & _keys(again)
    assert {a.key for a in again.answered} == set(given)
    assert again.target.schema_name == "syn_gold" and again.scd2_mode == "syn_type2"
    j = int(expr_key.split("columns[")[1].split("]")[0])
    assert again.columns[j].expression.endswith("*/ upper(`SYN_X`)")
    assert sum(len(f.questions) for f in after.feeds) == sum(
        len(f.questions) for f in contract.feeds
    ) - len(given)


def test_answers_yaml_is_read_by_the_ingestion_loader(tmp_path, contract):
    feed = _feed(contract, "syn_plan_ref")
    path = tmp_path / "answers.yaml"
    path.write_text(
        f'gaps:\n  "feeds[{feed.index}].load_mode": {{value: "full"}}\n', encoding="utf-8"
    )
    answers = edh_e2e.answers_from_yaml(path)
    again = next(
        f
        for f in read_dir(FIXTURES, EdhE2eConfig(), answers=answers).feeds
        if f.index == feed.index
    )
    assert again.load_mode == "full"
    assert f"feeds[{feed.index}].load_mode" not in _keys(again)


def test_cli_questions_workbook_is_refused_inside_the_repo_and_read_back(tmp_path):
    base = [
        sys.executable,
        "-m",
        "codegen.cli",
        "read-edh-e2e",
        "--input",
        str(FIXTURES),
        "--out",
        str(tmp_path / "gold.json"),
    ]
    inside = subprocess.run(
        base + ["--questions-xlsx", str(REPO / "q.xlsx")],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=REPO,
    )
    assert inside.returncode == 1 and not (REPO / "q.xlsx").exists()
    book = tmp_path / "questions_syn.xlsx"
    done = subprocess.run(
        base + ["--questions-xlsx", str(book)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=REPO,
    )
    assert done.returncode == 3, done.stderr
    contract = GoldContract.model_validate_json((tmp_path / "gold.json").read_text("utf-8"))
    total = sum(len(f.questions) for f in contract.feeds)
    assert f"questions workbook: {book} ({total} rows)" in done.stdout
    first = contract.feeds[0]
    book.write_bytes(_fill(book.read_bytes(), {first.questions[0].key: "syn answer"}))
    back = subprocess.run(
        base + ["--answers", str(book)], capture_output=True, text=True, encoding="utf-8", cwd=REPO
    )
    assert back.returncode == 3
    assert "1 answered" in back.stdout
