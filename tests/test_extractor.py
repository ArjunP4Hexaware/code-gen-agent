"""Workbook->STTM extractor: header-dialect tolerance, audit-row peeling,
verbatim recycle text, segmented/pairing diagnostics, golden-output byte
comparison, and end-to-end resolution of the emitted contract.

The CV/golden pair (fixtures/workbooks + FRD_demo_cv_golden) is a different
anonymization universe from the MIDS fixture contract (cv_* vs sd_*) — the
two are structural twins but never byte-comparable; see
docs/EXTRACTOR_RECON.md §4d.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from codegen.cli import main
from codegen.contracts.frd import FrdContract, GatedAmbiguity
from codegen.extract import (
    ExtractionError,
    SegmentedWorkbookError,
    WorkbookParseError,
    contract_to_json,
    extract_contract,
    parse_workbook,
)
from codegen.extract.workbook import parse_recycle_text
from codegen.resolve.resolver import resolve_pair

REPO = Path(__file__).resolve().parents[1]
WORKBOOK = REPO / "fixtures" / "workbooks" / "demo_sttm_cv_golden.xlsx"
FRD = REPO / "fixtures" / "contracts" / "FRD_demo_cv_golden.contract.json"
EXPECTED = REPO / "fixtures" / "contracts" / "sttm_mapping_contracts_cv_golden.json"
CONFIG = REPO / "config" / "config.yaml"
# Injected so extraction is byte-reproducible; must match the committed
# expected-output fixture's generated_date.
GENERATED_DATE = "2026-08-07"

SHEET_DEMOGRAPHIC = "MAPPING-CV_COMMUNITY_DEMOGRAPHI"
SHEET_COMMUNITY = "MAPPING-CV_COMMUNITY_RISK"
SHEET_INDIV = "MAPPING-CV_INDIV_RISK"


@pytest.fixture(scope="module")
def workbook_ir(config):
    return parse_workbook(WORKBOOK, config.extractor)


@pytest.fixture(scope="module")
def sheets_by_name(workbook_ir):
    return {sheet.sheet_name: sheet for sheet in workbook_ir.sheets}


@pytest.fixture(scope="module")
def cv_contract(config):
    return extract_contract(WORKBOOK, FRD, config, generated_date=GENERATED_DATE)


# ------------------------------------------------------------ FRD side


def test_cv_frd_fixture_parses_with_structured_ambiguities():
    frd = FrdContract.model_validate(json.loads(FRD.read_text(encoding="utf-8")))
    assert [f.feed_name for f in frd.feeds] == [
        "cv_community_demographic_risk",
        "cv_community_risk",
        "cv_individual_risk",
    ]
    assert frd.provenance is not None
    assert any(isinstance(a, GatedAmbiguity) for a in frd.provenance.ambiguities)


# ------------------------------------------------------- workbook parsing


def test_all_three_header_dialects_resolve(sheets_by_name):
    # One workbook, three source-block header dialects (recon §3a); all
    # three sheets must parse into the same IR shape.
    assert set(sheets_by_name) == {SHEET_DEMOGRAPHIC, SHEET_COMMUNITY, SHEET_INDIV}
    for sheet in sheets_by_name.values():
        assert sheet.rows and sheet.audit
        assert sheet.stage_table and sheet.standard_table


def test_audit_rows_peeled_not_emitted_as_fields(sheets_by_name):
    # 4-row delta per sheet: 90/271/50 sheet data rows -> 86/267/46 fields.
    expected = {SHEET_DEMOGRAPHIC: 86, SHEET_COMMUNITY: 267, SHEET_INDIV: 46}
    for name, count in expected.items():
        sheet = sheets_by_name[name]
        assert len(sheet.rows) == count, name
        assert [a.column for a in sheet.audit] == [
            "LOB",
            "SRC_FILE_NAME",
            "REC_CREATION_TIME",
            "REC_UPDATED_TIME",
        ]
        # Workbook says lowercase 'string'/'timestamp'; the contract Literal
        # requires normalized casing.
        assert [a.datatype for a in sheet.audit] == [
            "String",
            "String",
            "Timestamp",
            "Timestamp",
        ]
        assert not any(r.source_column == "NA" for r in sheet.rows)


def test_value_spec_comes_only_from_the_comment_dialect(sheets_by_name):
    indiv = sheets_by_name[SHEET_INDIV]
    assert all(r.value_spec is not None for r in indiv.rows)
    for name in (SHEET_DEMOGRAPHIC, SHEET_COMMUNITY):
        assert all(r.value_spec is None for r in sheets_by_name[name].rows)


def test_version_and_file_details(workbook_ir):
    # VERSION_HISTORY floats at E6 in the golden workbook; latest row wins.
    assert workbook_ir.version == "1.1"
    assert [row.vendor for row in workbook_ir.file_details] == ["Civic Vantage"] * 3


def test_recycle_text_is_verbatim_client_phrasing(sheets_by_name):
    recycle = parse_recycle_text(sheets_by_name[SHEET_INDIV])
    assert recycle is not None
    assert recycle.applies_to == "member_id"
    assert recycle.window_days == 7
    assert recycle.validation.startswith("Check with SUBS_ID from CoreMember")
    # Decision D3: no canonical rewriting.
    assert "against" not in recycle.validation.lower()
    for name in (SHEET_DEMOGRAPHIC, SHEET_COMMUNITY):
        assert parse_recycle_text(sheets_by_name[name]) is None


# ------------------------------------------------- loud-failure diagnostics


def _write_minimal_workbook(
    path: Path,
    *,
    second_stage_table: str | None = None,
    extra_source_header: str | None = None,
) -> Path:
    workbook = Workbook()
    details = workbook.active
    details.title = "FILE_DETAILS"
    details.append(["Vendor", "FileName", "File Description", "Location", "Frequency"])
    details.append(["Acme", "acme_YYYYMMDD.csv", None, None, "Monthly"])

    history = workbook.create_sheet("VERSION_HISTORY")
    history.append(["Date", "Version", "Author(s)", "Description of Version/Change"])
    history.append(["2026-01-01", "1.0", "T. Author", "Initial"])

    ws = workbook.create_sheet("MAPPING-ACME")
    source_headers = [
        "Database column Name",
        "NULL CHECK",
        "Description",
        "Sample Value",
        "DataType",
        "PHI Field",
        "Mandatory Field",
    ]
    if extra_source_header is not None:
        source_headers[2] = extra_source_header
    ws.append(
        ["Source File Layout"] + [None] * 6 + ["Stage Layer"] + [None] * 3
        + ["Standard Layer"] + [None] * 3
    )
    ws.append(
        source_headers
        + ["Schema", "TableName", "ColumnName", "DataType"]
        + ["Schema", "TableName", "ColumnName", "DataType"]
    )

    def _row(source_column: str, stage_table: str) -> list:
        return [
            source_column, "Not NULL", "desc", "1", "String", "No", "Yes",
            "stg_acme", stage_table, source_column, "String",
            "acme", "acme_table", source_column, "String",
        ]

    ws.append(_row("col_a", "acme_table"))
    ws.append(_row("col_b", second_stage_table or "acme_table"))
    ws.append(
        ["NA", "NULL", "NA", "NA", "NA", "No", "No",
         "stg_acme", "acme_table", "LOB", "string",
         "acme", "acme_table", "LOB", "string"]
    )
    workbook.save(path)
    return path


def test_segmented_layout_multiple_stage_tables_is_loud(tmp_path, config):
    path = _write_minimal_workbook(tmp_path / "seg.xlsx", second_stage_table="acme_table_hdr")
    with pytest.raises(SegmentedWorkbookError, match="segmented"):
        parse_workbook(path, config.extractor)


def test_segment_header_is_loud(tmp_path, config):
    path = _write_minimal_workbook(tmp_path / "seg2.xlsx", extra_source_header="Record Segment")
    with pytest.raises(SegmentedWorkbookError, match="segment"):
        parse_workbook(path, config.extractor)


def test_unrecognized_source_header_is_loud(tmp_path, config):
    path = _write_minimal_workbook(tmp_path / "bad.xlsx", extra_source_header="Mystery Column")
    with pytest.raises(WorkbookParseError, match="MAPPING-ACME.*Mystery Column"):
        parse_workbook(path, config.extractor)


def _doctored_frd(tmp_path: Path, mutate) -> Path:
    data = json.loads(FRD.read_text(encoding="utf-8"))
    mutate(data)
    path = tmp_path / "frd.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def test_sheet_with_no_frd_feed_is_loud(tmp_path, config):
    def _mutate(data):
        data["feeds"][0]["stage_target"]["tables"] = ["somewhere_else"]

    with pytest.raises(ExtractionError, match="cv_community_demographic_risk.*0 FRD feeds"):
        extract_contract(WORKBOOK, _doctored_frd(tmp_path, _mutate), config)


def test_frd_feed_with_no_sheet_is_loud(tmp_path, config):
    def _mutate(data):
        extra = json.loads(json.dumps(data["feeds"][2]))
        extra["feed_name"] = "cv_phantom_feed"
        extra["stage_target"]["tables"] = ["cv_phantom"]
        data["feeds"].append(extra)

    with pytest.raises(ExtractionError, match="cv_phantom_feed.*no mapping sheet"):
        extract_contract(WORKBOOK, _doctored_frd(tmp_path, _mutate), config)


# ------------------------------------- FILE_DETAILS date-placeholder pairing


def test_ccyy_patterns_pair_via_canonicalization(tmp_path, config):
    # Models the real-pair finding: FRD patterns say CCYY/CCYYMMDD where the
    # workbook FILE_DETAILS names say YYYY/YYYYMMDD. Exact membership fails;
    # canonicalized membership must pair all three feeds.
    def _mutate(data):
        only_ccyy = {
            "cv_community_demographic_risk": ["demographic_extract_CCYY_MM.csv"],
            "cv_community_risk": ["community_metrics_CCYY_MM.csv"],
            "cv_individual_risk": ["cv_ind_risk_data_package_mrdn_OH_CCYYMMDD_HHMM.psv"],
        }
        for feed in data["feeds"]:
            feed["file_name_patterns"] = only_ccyy[feed["feed_name"]]

    contract = extract_contract(
        WORKBOOK, _doctored_frd(tmp_path, _mutate), config, generated_date=GENERATED_DATE
    )
    # Pairing succeeded, and the emitted name_pattern stays the VERBATIM
    # workbook file name (canonicalization is comparison-only).
    assert [f.source_file.name_pattern for f in contract.feeds] == [
        "demographic_extract_YYYY_MM.csv",
        "community_metrics_YYYY_MM.csv",
        "cv_ind_risk_data_package_mrdn_OH_YYYYMMDD_HHMM.psv",
    ]


def test_unrelated_filename_still_fails_pairing(tmp_path, config):
    def _mutate(data):
        data["feeds"][2]["file_name_patterns"] = ["totally_unrelated_file.psv"]

    with pytest.raises(ExtractionError, match="cv_individual_risk.*0 FILE_DETAILS rows"):
        extract_contract(WORKBOOK, _doctored_frd(tmp_path, _mutate), config)


# --------------------------------------- segmented layout-family diagnostic


def _write_segmented_family_workbook(path: Path) -> Path:
    """Minimal synthetic sheet with the family signature: key:value
    metadata block, band-label row (with the 'Source Layout' variant),
    and a per-row Segment column in the header row. Generic names only."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "feed_x"
    ws.append(["File(s)", "feed_x_*.txt"])
    ws.append(["File type", "txt (pipe delimited)"])
    ws.append(["Source Layout"] + [None] * 4 + ["Stage Layer"] + [None] * 3
              + ["Standard Layer"])
    ws.append(["#", "Field Name", "Data Type", "Segment", "PII",
               "Schema", "TableName", "ColumnName", "DataType", "Schema"])
    workbook.save(path)
    return path


def test_segmented_family_gets_the_accurate_diagnostic(tmp_path, config):
    path = _write_segmented_family_workbook(tmp_path / "family.xlsx")
    with pytest.raises(SegmentedWorkbookError, match="SEGMENTED_MODE_DESIGN"):
        parse_workbook(path, config.extractor)


def test_alien_workbook_still_gets_the_original_error(tmp_path, config):
    workbook = Workbook()
    ws = workbook.active
    ws.title = "notes"
    ws.append(["hello", "world"])
    path = tmp_path / "alien.xlsx"
    workbook.save(path)
    with pytest.raises(WorkbookParseError, match="no mapping sheets found"):
        parse_workbook(path, config.extractor)


# --------------------------------------------------------- emitted contract


def test_extracted_contract_shape(cv_contract):
    assert [f.feed_id for f in cv_contract.feeds] == [
        "cv_community_demographic_risk",
        "cv_community_risk",
        "cv_individual_risk",
    ]
    assert [f.field_count for f in cv_contract.feeds] == [86, 267, 46]
    assert cv_contract.sttm_version == "1.1"
    assert not cv_contract.synthetic
    for feed in cv_contract.feeds:
        assert feed.standard is not None
        assert not feed.is_segmented
        assert len(feed.audit_columns) == 4
    indiv = cv_contract.feeds[2]
    assert indiv.source_file.delimiter == "|"  # implied from FRD psv format
    assert indiv.load_rules.phi_columns == ["member_id"]
    assert indiv.load_rules.recycle is not None
    assert indiv.load_rules.recycle.recycle_window_days == 7


def test_golden_output_is_byte_identical(cv_contract):
    assert contract_to_json(cv_contract) == EXPECTED.read_text(encoding="utf-8")


def test_extracted_contract_resolves_through_resolver(config):
    specs = resolve_pair(FRD, EXPECTED, config)
    assert [s.feed_id for s in specs] == [
        "cv_community_demographic_risk",
        "cv_community_risk",
        "cv_individual_risk",
    ]
    assert [s.delimiter for s in specs] == [",", ",", "|"]
    recycle = specs[2].recycle
    assert recycle is not None
    # The VERBATIM client phrasing resolved through the extended reference
    # parser (decision D3).
    assert recycle.reference_table == "PR_STD.COREMEMBER.CM_SUBS_MASTER"
    assert recycle.reference_id_column == "SUBS_ID"
    assert recycle.reference_filter == "GRP_CK = 47"
    assert recycle.recycle_table.table == "cv_individual_risk_recycle"


def test_cli_extract_sttm_reproduces_the_golden_fixture(tmp_path):
    out = tmp_path / "extracted.json"
    exit_code = main(
        [
            "extract-sttm",
            "--config", str(CONFIG),
            "--workbook", str(WORKBOOK),
            "--frd-contract", str(FRD),
            "--out", str(out),
            "--generated-date", GENERATED_DATE,
        ]
    )
    assert exit_code == 0
    assert out.read_text(encoding="utf-8") == EXPECTED.read_text(encoding="utf-8")
