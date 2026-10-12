"""The file of a split feed (layout.resolve.split_frd_feeds_by_sttm), 2026-10-12.

Regression from 1faa83d: once the FRD reader lifted a block's file names into
its one feed's ``file_name_patterns``, the split handed that WHOLE list to every
feed its name matching could not pair — an unanswered run then failed in
extract-sttm on "N FILE_DETAILS rows match" instead of naming the question.
Now:

* a feed whose file is asked owns NO pattern; layout --require-complete prints
  ``QUESTION feeds[i].file_name_patterns`` with the candidates and exits 3, and
  extract-sttm names the question (never "rows match");
* when exactly ONE table and ONE file are left after the unique name matches,
  the file is taken BY ELIMINATION (owner decision: fewer questions), said in a
  NOTE line — never a QUESTION — and the run completes unanswered.

Fixtures: the synthetic cold_4 pair (two mapping sheets, two files; the
reversal file pairs by name, the claim file does not), and the same pair with
a third FILE_DETAILS row so the elimination is no longer unique.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest
import yaml

import test_cold_pairs
from codegen import cli
from codegen.extract.frd_docx import read_frd
from codegen.layout.resolve import split_frd_feeds_by_sttm
from test_cold_pairs import DATE, _answers, _paths, _run
from test_harness_answer_lines import _NEEDED_KEY_RE

cold_env = test_cold_pairs.cold_env          # the harness's environment (fixture)
KEY = "feeds[0].file_name_patterns"


def _labelled(out: str, label: str) -> list[str]:
    return [line[cli.ANSWER_LABEL_WIDTH:].split(" — ", 1)[0] for line in out.splitlines()
            if line.startswith(label + " ")]


def _role_answers(tmp: Path) -> Path:
    """cold_4's answers.yaml WITHOUT the file answer: only the column placements."""
    answers = _answers("cold_4")
    answers.pop("gaps", None)
    path = tmp / "roles_only.yaml"
    path.write_text(yaml.safe_dump(answers), encoding="utf-8")
    return path


def _with_extra_file(tmp: Path) -> Path:
    """cold_4's STTM with a third FILE_DETAILS row: two files left for one table."""
    sttm, _frd = _paths("cold_4")
    copy = tmp / sttm.name
    shutil.copyfile(sttm, copy)
    book = openpyxl.load_workbook(copy)
    sheet = next(book[n] for n in book.sheetnames if "FILE_DETAIL" in n.upper())
    last = sheet.max_row
    sheet.append([sheet.cell(last, c).value for c in range(1, sheet.max_column + 1)])
    sheet.cell(last + 1, 2).value = "CH_RX_ADJ_YYYYMMDD.csv"
    book.save(copy)
    return copy


def _layout(capsys, tmp: Path, sttm: Path, answers: Path) -> tuple[int, str, Path]:
    _s, frd = _paths("cold_4")
    frd_out = tmp / "frd.contract.json"
    code, out = _run(capsys, ["layout", "--workbook", str(sttm), "--frd", str(frd),
                              "--dry-run", "--no-cache", "--require-complete",
                              "--frd-contract-out", str(frd_out), "--answers", str(answers)])
    return code, out, frd_out


def test_a_unique_elimination_is_a_note_and_the_run_completes_unanswered(cold_env, capsys):
    sttm, _frd = _paths("cold_4")
    answers = _role_answers(cold_env)
    code, out, frd_out = _layout(capsys, cold_env, sttm, answers)
    assert code == 0, out
    notes = [k for k in _labelled(out, "NOTE") if k.startswith("feeds[")]
    assert notes == [KEY]                     # (other NOTEs: the answers file's confirmations)
    assert "(by elimination" in out and "informational, not required to run" in out
    assert not [line for line in out.splitlines() if _NEEDED_KEY_RE.match(line)]
    code, out = _run(capsys, ["extract-sttm", "--workbook", str(sttm), "--frd-contract",
                              str(frd_out), "--out", str(cold_env / "sttm.json"),
                              "--generated-date", DATE, "--answers", str(answers)])
    assert code == 0, out


def test_a_non_unique_leftover_is_asked_and_the_feed_owns_no_pattern(cold_env, capsys):
    sttm = _with_extra_file(cold_env)
    answers = _role_answers(cold_env)
    code, out, frd_out = _layout(capsys, cold_env, sttm, answers)
    assert code == cli.EXIT_NEEDS_ANSWERS, out
    assert _labelled(out, "QUESTION") == [KEY]
    question = next(line for line in out.splitlines() if line.startswith("QUESTION"))
    for name in ("CH_RX_CLM_<LOB>_YYYYMMDD.csv", "CH_RX_REV_YYYYMMDD.csv",
                 "CH_RX_ADJ_YYYYMMDD.csv"):
        assert name in question                       # the candidate list, names only
    contract = yaml.safe_load(frd_out.read_text(encoding="utf-8"))
    assert contract["feeds"][0]["file_name_patterns"] == []          # never the shared list
    assert contract["feeds"][1]["file_name_patterns"] == ["CH_RX_REV_YYYYMMDD.csv"]
    # Run past the question: extract names it — never "N FILE_DETAILS rows match".
    code, out = _run(capsys, ["extract-sttm", "--workbook", str(sttm), "--frd-contract",
                              str(frd_out), "--out", str(cold_env / "sttm.json"),
                              "--generated-date", DATE, "--answers", str(answers)])
    assert "FILE_DETAILS rows match" not in out
    assert "File for table ch_rx_claim" in out


def test_the_answer_still_wins_over_the_question(cold_env, capsys):
    sttm = _with_extra_file(cold_env)
    answers = _answers("cold_4")                      # roles + the claim file answered
    path = cold_env / "full.yaml"
    path.write_text(yaml.safe_dump(answers), encoding="utf-8")
    code, out, frd_out = _layout(capsys, cold_env, sttm, path)
    assert code == 0, out
    contract = yaml.safe_load(frd_out.read_text(encoding="utf-8"))
    assert contract["feeds"][0]["file_name_patterns"] == ["CH_RX_CLM_<LOB>_YYYYMMDD.csv"]


def test_the_split_never_hands_the_shared_list_to_an_asked_feed(config, tmp_path):
    """Unit form of the 1faa83d shape: the FRD feed carries EVERY file of its
    block; one table pairs by name, two do not -> both asked, both empty, the
    paired one keeps its file. (Synthetic names.)"""
    from test_frd_multi_object import _multi_object_docx, discover_frd, read_docx

    path = tmp_path / "frd.docx"
    path.write_bytes(_multi_object_docx("alpha_claims_YYYYMMDD.csv"))
    content = read_docx(path)
    contract = read_frd(content, discover_frd(content, config.extractor.frd), config,
                        document_name="frd.docx", generated_date="2026-01-01")
    files = ["alpha_claims_YYYYMMDD.csv", "vendor_feed_one_YYYYMMDD.csv",
             "vendor_feed_two_YYYYMMDD.csv"]
    lifted = contract.feeds[0].model_copy(update={"file_name_patterns": files})
    contract = contract.model_copy(update={"feeds": [lifted]})
    facts = {"meta": {}, "files": [(f, f"FILE_DETAILS!B{i}") for i, f in enumerate(files, 2)],
             "sheet_tables": [
                 {"sheet": f"MAPPING-{t.upper()}", "stage_table": t, "standard_table": None,
                  "stage_schema": "stg_syn"}
                 for t in ("alpha_claims", "beta_members", "gamma_scores")]}
    split, flags, questions = split_frd_feeds_by_sttm(contract, facts, {}, config)
    assert [f.file_name_patterns for f in split.feeds] == [["alpha_claims_YYYYMMDD.csv"], [], []]
    assert [q.key for q in questions] == ["feeds[1].file_name_patterns",
                                          "feeds[2].file_name_patterns"]
    assert all(len(q.candidates) == 3 for q in questions)
    assert not any("(by elimination" in f for f in flags)


@pytest.mark.parametrize("answered", [False, True])
def test_owed_key_is_a_question_for_the_harness(answered):
    """The file question is a choice under gaps: - owed (QUESTION), the line
    the harness's regex reads; a by-elimination fill is never one."""
    assert cli._owed_key("frd", KEY, "choice") is True
    line = cli.answer_line("QUESTION" if not answered else "NOTE", KEY, "reason")
    assert bool(_NEEDED_KEY_RE.match(line)) is (not answered)
