"""Option B (framework output): snapshot-guarded Option A + framework mode.

The notebook-mode snapshot (``tests/snapshots/notebook_mode.json``) was
generated from the pre-change emitter at the start of the Option B session;
the suite regenerates the three CV feeds in notebook mode and asserts every
emitted file's sha256 against it — Option A stays byte-identical or this
fails. Framework-mode tests assert the artefact contract: workbooks open,
sheets match the configured layout, IDs are placeholders, banners present,
inserts are dialect-sane, verdict identical across modes, and a second
render is byte-identical.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from codegen import cli
from codegen.resolve.resolver import resolve_pair

REPO = Path(__file__).resolve().parents[1]
_FRD = REPO / "fixtures" / "contracts" / "FRD_demo_cv_golden.contract.json"
_STTM = REPO / "fixtures" / "contracts" / "sttm_mapping_contracts_cv_golden.json"
_SNAPSHOT = REPO / "tests" / "snapshots" / "notebook_mode.json"

needs_demo_pair = pytest.mark.skipif(
    not (_FRD.is_file() and _STTM.is_file()),
    reason="demo fixture pair not restored (removed 2026-08-22)",
)


def _config_into(config, tmp: Path, mode: str):
    return config.model_copy(update={
        "output": config.output.model_copy(update={
            "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports"),
            "mode": mode,
        })
    })


def _generate_all(config, tmp: Path, mode: str, output_mode: str | None = None):
    scoped = _config_into(config, tmp, mode)
    specs = resolve_pair(_FRD, _STTM, scoped)
    for spec in specs:
        gate = cli._generate_feed(spec, scoped, dry_run=True, skip_tests=True,
                                  output_mode=output_mode)
        yield spec, gate


def _manifest(tmp: Path) -> dict[str, str]:
    manifest = {}
    for base, label in ((tmp / "out", "out"), (tmp / "reports", "reports")):
        for p in sorted(base.rglob("*")):
            if p.is_file():
                rel = f"{label}/{p.relative_to(base).as_posix()}"
                manifest[rel] = hashlib.sha256(p.read_bytes()).hexdigest()
    return manifest


@needs_demo_pair
def test_notebook_mode_matches_pre_change_snapshot(config, tmp_path):
    for _spec, _gate in _generate_all(config, tmp_path, "notebook"):
        pass
    snapshot = json.loads(_SNAPSHOT.read_text(encoding="utf-8"))
    assert _manifest(tmp_path) == snapshot


@pytest.fixture(scope="module")
def framework_run(config, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("framework_mode")
    results = list(_generate_all(config, tmp, "framework"))
    return tmp, results


@needs_demo_pair
def test_framework_mode_tree_and_verdicts(config, framework_run, tmp_path):
    tmp, results = framework_run
    notebook_gates = {
        spec.feed_slug: gate.verdict
        for spec, gate in _generate_all(config, tmp_path, "notebook")
    }
    for spec, gate in results:
        feed_dir = tmp / "out" / spec.feed_slug
        # Verdict identical across modes — the gate ran on the full render.
        assert gate.verdict == notebook_gates[spec.feed_slug] == "PASS_WITH_FLAGS"
        # No notebook, no module tree; ddl + framework + candidates only.
        assert not (feed_dir / f"{spec.feed_slug}.ipynb").exists()
        assert not (feed_dir / "pipeline").exists()
        assert (feed_dir / "ddl").is_dir()
        names = sorted(p.name for p in (feed_dir / "framework").iterdir())
        assert names == ["ADDITION.md", "config_inserts.xlsx",
                         "config_rows.xlsx",
                         f"{spec.feed_slug}_stage_table_creation.txt",
                         f"{spec.feed_slug}_standard_table_creation.txt"]


@needs_demo_pair
def test_framework_workbooks_layout_and_provenance(config, framework_run):
    tmp, results = framework_run
    layout = config.demo.metadata_sheet
    for spec, _gate in results:
        framework_dir = tmp / "out" / spec.feed_slug / "framework"

        rows_workbook = load_workbook(framework_dir / "config_rows.xlsx")
        assert rows_workbook.sheetnames == [*layout.tabs, "_provenance", "_inputs"]
        for name, tab in layout.tabs.items():
            assert [c.value for c in rows_workbook[name][1]] == tab.headers
        inputs = {row[0].value: row[1].value
                  for row in rows_workbook["_inputs"].iter_rows(min_row=2)}
        assert spec.frd_contract_sha256 in inputs["FRD contract"]
        assert spec.sttm_contract_sha256 in inputs["STTM contract"]
        assert "Load-pattern FAQ" in inputs
        assert "client IIG template" in inputs["Layout"]



_REF = REPO / "fixtures" / "reference"


def test_config_layout_matches_scrubbed_iig_workbook(config):
    """The IIG layout in config/config.yaml was transcribed from the scrubbed
    reference workbook — pin them together so neither can drift silently.
    Tab order, tab names, and every header must match the workbook verbatim
    (including the client's own spellings, e.g. CRETAED_BY)."""
    workbook = load_workbook(_REF / "SFMC_IIG.xlsx", read_only=True,
                             data_only=True)
    try:
        workbook_layout = {
            sheet.title: [
                str(value).strip()
                for value in next(sheet.iter_rows(max_row=1, values_only=True))
                if value is not None
            ]
            for sheet in workbook.worksheets
        }
    finally:
        workbook.close()
    config_layout = {name: list(tab.headers)
                     for name, tab in config.demo.metadata_sheet.tabs.items()}
    assert list(config_layout) == list(workbook_layout)  # names AND order
    assert config_layout == workbook_layout


def _txt_paths(tmp: Path, spec) -> tuple[Path, Path]:
    framework_dir = tmp / "out" / spec.feed_slug / "framework"
    return (framework_dir / f"{spec.feed_slug}_stage_table_creation.txt",
            framework_dir / f"{spec.feed_slug}_standard_table_creation.txt")


@needs_demo_pair
def test_ddl_txt_files_exist_and_are_plain_sql(framework_run):
    tmp, results = framework_run
    for spec, _gate in results:
        for path in _txt_paths(tmp, spec):
            assert path.is_file(), path
            text = path.read_text(encoding="utf-8")
            assert text.strip(), path
            assert text.splitlines()[0].startswith("--"), path
            assert "```" not in text, path
            # provenance banner as leading SQL comments
            assert spec.frd_contract_sha256 in text
            assert "never creates target tables" in text


def _tblproperties_keys(text: str) -> set[str]:
    import re
    return set(re.findall(r"'(delta\.[A-Za-z.]+)'\s*=", text))


@needs_demo_pair
def test_ddl_txt_style_conforms_to_reference_goldens(framework_run):
    """Structural conformance to the client's scrubbed reference goldens:
    clause presence/order, LOCATION on stage only, CLUSTER BY AUTO on
    standard only, TBLPROPERTIES key sets, trailing SET TAGS — never value
    equality (different feed, different columns)."""
    golden_stage = (_REF / "SFMC_stage_table_creation.txt").read_text(encoding="utf-8")
    golden_standard = (_REF / "SFMC_standard_table_creation.txt").read_text(encoding="utf-8")
    tmp, results = framework_run
    for spec, _gate in results:
        stage_path, standard_path = _txt_paths(tmp, spec)
        stage = stage_path.read_text(encoding="utf-8")
        standard = standard_path.read_text(encoding="utf-8")

        for text in (stage, standard):
            assert "CREATE OR REPLACE TABLE " in text
            assert " COMMENT '" in text          # per-column COMMENT clauses
            assert "USING delta" in text
            assert "SET TAGS ('DOMAIN' = " in text
            assert "CREATE TABLE IF NOT EXISTS" not in text

        # LOCATION on stage only; CLUSTER BY AUTO on standard only.
        assert "LOCATION '" in stage and "LOCATION '" not in standard
        assert "CLUSTER BY AUTO" in standard and "CLUSTER BY AUTO" not in stage
        # The synthetic location is labeled and never a real ADLS host.
        assert "SYNTHETIC" in stage
        assert "dfs.core.windows.net" not in stage

        # TBLPROPERTIES: exactly the goldens' key sets, per layer.
        assert _tblproperties_keys(stage) == _tblproperties_keys(golden_stage)
        assert _tblproperties_keys(standard) == _tblproperties_keys(golden_standard)

        # Clause order within the first table section mirrors the goldens.
        for text, has_location in ((stage, True), (standard, False)):
            create = text.index("CREATE OR REPLACE TABLE ")
            using = text.index("USING delta")
            props = text.index("TBLPROPERTIES (")
            tags = text.index("SET TAGS (")
            assert text.index(f"--{text[create:].split()[4]}") < create
            assert create < using < props < tags
            if has_location:
                assert using < text.index("LOCATION '") < props
            else:
                assert using < text.index("CLUSTER BY AUTO") < props


def _inserts_workbook(tmp: Path, spec):
    return load_workbook(tmp / "out" / spec.feed_slug / "framework" /
                         "config_inserts.xlsx")


@needs_demo_pair
def test_config_inserts_workbook_layout_and_statements(config, framework_run):
    tmp, results = framework_run
    layout = config.demo.metadata_sheet
    always_blank = set(layout.always_blank)
    for spec, _gate in results:
        workbook = _inserts_workbook(tmp, spec)
        # One sheet per POPULATED tab + _provenance; every populated sheet is
        # a real IIG tab name.
        assert workbook.sheetnames[-1] == "_provenance"
        data_sheets = workbook.sheetnames[:-1]
        assert data_sheets
        assert set(data_sheets) <= {name[:31] for name in layout.tabs}
        for name in data_sheets:
            sheet = workbook[name]
            tab_headers = next(
                tab.headers for tab_name, tab in layout.tabs.items()
                if tab_name[:31] == name)
            headers = [c.value for c in sheet[1]]
            # Value columns on the real layout, then the statement (oversize
            # statements continue in _PART<n> columns, never truncated).
            assert headers[:len(tab_headers) + 1] == [*tab_headers,
                                                      "INSERT_STATEMENT"]
            assert all(h.startswith("INSERT_STATEMENT_PART")
                       for h in headers[len(tab_headers) + 1:])
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            assert rows  # populated sheets only
            # INSERT count per sheet == row count; IDs are placeholders.
            for row in rows:
                statement = "".join(str(part) for part in
                                    row[len(tab_headers):] if part)
                assert statement.startswith("INSERT INTO ")
                assert statement.rstrip().endswith(";")
                for header, value in zip(tab_headers, row, strict=False):
                    if header in always_blank:
                        assert value == config.framework.id_placeholder
        provenance = {r[0]: r[1] for r in
                      workbook["_provenance"].iter_rows(min_row=2,
                                                        values_only=True)}
        assert spec.frd_contract_sha256 in provenance["FRD contract"]
        assert provenance["dialect"] == "sqlserver"
        assert "only after" in provenance["note"]


@needs_demo_pair
def test_config_inserts_deterministic_across_builds(config, tmp_path):
    first = tmp_path / "a"
    second = tmp_path / "b"
    digests = []
    for tmp in (first, second):
        run_digests = {}
        for spec, _gate in _generate_all(config, tmp, "framework"):
            path = (tmp / "out" / spec.feed_slug / "framework" /
                    "config_inserts.xlsx")
            run_digests[spec.feed_slug] = hashlib.sha256(
                path.read_bytes()).hexdigest()
        digests.append(run_digests)
    assert digests[0] == digests[1]


@needs_demo_pair
def test_lakebase_dialect(config, tmp_path):
    lakebase = config.model_copy(update={
        "framework": config.framework.model_copy(update={
            "sql_dialect": "lakebase",
            "tables": {"ADLS_DELTA_INGESTION_DETAILS": "meta.adls_ingestion"},
        })
    })
    results = list(_generate_all(lakebase, tmp_path, "framework"))
    spec, _gate = results[0]
    workbook = _inserts_workbook(tmp_path, spec)
    sheet = workbook["ADLS_DELTA_INGESTION_DETAILS"]
    n_values = len(config.demo.metadata_sheet.tabs[
        "ADLS_DELTA_INGESTION_DETAILS"].headers)
    statement = "".join(
        str(c) for c in next(sheet.iter_rows(min_row=2, values_only=True))
        [n_values:] if c)
    assert statement.startswith('INSERT INTO "meta"."adls_ingestion" ("GROUP_ID"')
    # No sqlserver unicode string literals in lakebase (a plain 'N' flag
    # value is fine; the N-prefix form would follow a separator).
    assert ", N'" not in statement and "(N'" not in statement


@needs_demo_pair
def test_both_mode_is_the_union(config, tmp_path):
    for spec, gate in _generate_all(config, tmp_path, "both"):
        feed_dir = tmp_path / "out" / spec.feed_slug
        assert gate.verdict == "PASS_WITH_FLAGS"
        assert (feed_dir / f"{spec.feed_slug}.ipynb").is_file()
        assert (feed_dir / "pipeline").is_dir()
        assert (feed_dir / "framework" / "config_rows.xlsx").is_file()
        report = (tmp_path / "reports" / f"{spec.feed_slug}.md").read_text(
            encoding="utf-8")
        assert "## Framework output (Option B)" in report


@needs_demo_pair
def test_no_raw_client_values_in_reference_or_emitted_artefacts(
        config, framework_run, tmp_path):
    """Denylist scan (scripts/scrub_check.py): the tracked reference fixtures
    and EVERY emitted artefact — framework AND notebook mode — must carry
    zero raw client values (generic infra patterns always; the concrete
    harvested values too when inputs/reference_raw/ is present locally)."""
    import sys
    sys.path.insert(0, str(REPO / "scripts"))
    from scrub_check import scan_paths

    for _ in _generate_all(config, tmp_path, "notebook"):
        pass
    tmp, _results = framework_run
    targets = [REPO / "fixtures" / "reference",
               tmp / "out", tmp / "reports",
               tmp_path / "out", tmp_path / "reports"]
    hits = scan_paths([t for t in targets if t.exists()])
    assert hits == [], hits


@needs_demo_pair
def test_framework_mode_rendered_twice_is_byte_identical(config, tmp_path):
    first = tmp_path / "a"
    second = tmp_path / "b"
    for tmp in (first, second):
        for _ in _generate_all(config, tmp, "framework"):
            pass
    assert _manifest(first) == _manifest(second)
