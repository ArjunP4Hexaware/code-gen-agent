"""FRD gap handling (resolve/gapfill.py + layout fills + resolver fallbacks).

* a docx FRD missing the file format and both load strategies, with an STTM
  stating them per layer → no pause, provenance flags, the run resolves;
* FRD and STTM disagreeing on the format → one "choice" pause with both
  candidates; the person's choice lands with source=user;
* an STTM with one layer-less "Load Strategy" → one "layer" pause;
* the STTM stage band is never asked about; domain stays FRD-only.
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from acfc_shapes import frd as frd_fixtures
from acfc_shapes import pair1
from acfc_shapes.common import docx_bytes
from codegen.extract import contract_to_json as sttm_to_json
from codegen.extract import extract_contract
from codegen.extract.frd_docx import contract_to_json as frd_to_json
from codegen.layout.fingerprint import fingerprint
from codegen.layout.model import MockLayoutProvider
from codegen.layout.resolve import parse_answers, resolve_pair
from codegen.resolve.gapfill import canonical_strategy, parse_load_strategy_text
from codegen.resolve.resolver import resolve_pair as resolve_contracts

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PROFILES = REPO / "fixtures" / "layout_profiles"
DATE = "2026-01-01"
MARKERS = {"stage": ["STG", "Stage", "Staging"], "standard": ["STD", "Standard"]}


def test_load_strategy_text_parsing():
    assert parse_load_strategy_text("Append", MARKERS) == {"any": "Append"}
    assert parse_load_strategy_text("STG: Append; STD: Upsert", MARKERS) == {
        "stage": "Append", "standard": "Upsert"}
    assert parse_load_strategy_text("Stage - Truncate and Load\nStandard - merge", MARKERS) == {
        "stage": "Truncate and Load", "standard": "Upsert"}
    assert parse_load_strategy_text("", MARKERS) == {}
    assert parse_load_strategy_text(None, MARKERS) == {}
    assert parse_load_strategy_text("whatever", MARKERS) == {}
    assert canonical_strategy("truncate & load") == "Truncate and Load"
    assert canonical_strategy("OVERWRITE") == "Truncate and Load"
    assert canonical_strategy("nonsense") is None


# -- fixtures: pair 1 with holes -------------------------------------------------- #


def _pair1_docx_without(missing: set[str], drop_labels: set[str] = frozenset()) -> bytes:
    """The pair-1 F1 FRD with the given Structural labels blanked (or set);
    ``drop_labels`` removes the label rows altogether (a label present with a
    blank value is "stated blank" and raises no layout question)."""
    feed = pair1.FEED_NAME
    stage = f"{pair1.STAGE_CATALOG}.{pair1.STAGE_SCHEMA}"
    standard = f"{pair1.STANDARD_CATALOG}.{pair1.STANDARD_SCHEMA}"
    functional = "Ingest."
    structural = {
        "Object/data Format": pair1.FILE_FORMAT,
        "Target Catalog and Schema": f"STG: {stage}; STD: {standard}",
        "Target Table Name": pair1.TABLE,
        "Domain and Subdomain": f"{pair1.DOMAIN} / {pair1.SUB_DOMAIN}",
        "Load Strategy STG": "Append",
        "Load Strategy STD (View)": "Append",
        "Load Strategy Consumption (EDH, BSL)": "N/A",
        "Archive Schedule": "Archive after load",
        "Source Data Dictionary": f"{feed} VDD",
        "ADLS Location": f"/{pair1.DOMAIN}/{pair1.SUB_DOMAIN}/",
        "Inbound File Folder Path": f"inbound\\{pair1.DOMAIN.lower()}",
    }
    overrides = missing.items() if isinstance(missing, dict) else ((m, "") for m in missing)
    for label, value in overrides:
        structural[label] = value
    sections = [
        frd_fixtures._section("Descriptive Metadata", feed, "desc", functional,
                              frd_fixtures.DESCRIPTIVE, {
                                  "Data Source": pair1.VENDOR_NAME, "Object Name": feed,
                                  "Frequency": pair1.FREQUENCY, "LOBs": pair1.LOB}),
        frd_fixtures._section("Structural Metadata", feed, "struct", functional,
                              [lab for lab in frd_fixtures.STRUCTURAL_P1
                               if lab not in drop_labels], structural),
        frd_fixtures._section("Administrative Metadata", feed, "", functional,
                              frd_fixtures.ADMINISTRATIVE, {}),
        frd_fixtures._section("Technical Metadata", feed, "", functional,
                              frd_fixtures.TECHNICAL, {}),
        frd_fixtures._section("Data Quality", feed, "", functional, frd_fixtures.DATA_QUALITY, {}),
        frd_fixtures._section("Vendor Metadata", feed, "", functional, frd_fixtures.VENDOR,
                              {"Vendor Name": pair1.VENDOR_NAME}),
    ]
    leading, trailing = frd_fixtures._common_tables("PROJECT_ALPHA", feed)
    return docx_bytes(leading + sections + trailing)


def _sttm_with_load_strategy(tmp: Path, text: str) -> tuple[Path, Path]:
    """Pair-1 STTM with its 'Load Strategy' meta value replaced, plus a cache
    dir holding the tracked layout profile re-keyed to the new fingerprint
    (the meta value sits inside the fingerprinted header region)."""
    path = tmp / "pair_1_family_a.xlsx"
    shutil.copyfile(SHAPES / "sttm" / "pair_1_family_a.xlsx", path)
    wb = load_workbook(path)
    ws = wb["FEED_1_MAPPING"]
    for row in ws.iter_rows(min_row=1, max_row=15):
        if row[0].value == "Load Strategy":
            row[1].value = text
    wb.save(path)
    cache = tmp / "profiles"
    cache.mkdir()
    tracked = PROFILES / "sttm_pair_1_family_a.layout.json"
    profile = json.loads(tracked.read_text(encoding="utf-8"))
    profile["fingerprint"] = fingerprint(load_workbook(path))
    (cache / "sttm_pair_1_variant.layout.json").write_text(json.dumps(profile), encoding="utf-8")
    return path, cache


def _resolve(tmp: Path, docx: bytes, sttm: Path, cache: Path, answers=None):
    frd_path = tmp / "frd.docx"
    frd_path.write_bytes(docx)
    return resolve_pair(sttm, frd_path, load_config_for_tests(), provider=MockLayoutProvider(
        [PROFILES / "mock", cache, PROFILES]), cache_dirs=[cache, PROFILES],
        generated_date=DATE, answers=answers, use_cache=True)


def load_config_for_tests():
    from codegen.config import load_config

    return load_config(REPO / "config" / "config.yaml",
                       overlays=[SHAPES / "pair_1" / "config_overlay.yaml"])


# -- 1. silent FRD, STTM states per layer -> no pause ------------------------------- #


def test_silent_frd_filled_from_sttm_without_a_pause(tmp_path):
    sttm, cache = _sttm_with_load_strategy(tmp_path, "STG: Append; STD: Append")
    pair = _resolve(tmp_path, _pair1_docx_without(
        {"Object/data Format", "Load Strategy STG", "Load Strategy STD (View)"}), sttm, cache)
    assert pair.questions == [], [q.key for q in pair.questions]
    fills = {f["field"]: f for f in pair.gap_fills}
    assert fills["feeds[0].file_format"]["value"] == pair1.FILE_FORMAT
    assert fills["feeds[0].file_format"]["source"] == "STTM"
    assert "File Format" in fills["feeds[0].file_format"]["cell"]
    assert fills["feeds[0].stage_target.load_strategy"]["value"] == "Append"
    assert fills["feeds[0].standard_target.load_strategy"]["value"] == "Append"
    flagged = [f for f in pair.flags if f.startswith("frd_unstated:")]
    assert len([f for f in flagged if "source_used:STTM" in f]) == 3
    assert not any(f.startswith("layout_unresolved:frd") for f in pair.flags)
    feed = pair.frd_contract.feeds[0]
    assert (feed.file_format, feed.stage_target.load_strategy,
            feed.standard_target.load_strategy) == (pair1.FILE_FORMAT, "Append", "Append")
    # The run resolves end to end with the filled contract.
    config = load_config_for_tests()
    frd_json = tmp_path / "frd.contract.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    sttm_json = tmp_path / "sttm.contract.json"
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config)
    assert spec.file_format == pair1.FILE_FORMAT and spec.stage_load_strategy == "Append"


# -- 2. FRD vs STTM disagree on the format -> one choice pause ---------------------- #


def test_format_disagreement_pauses_with_both_candidates_and_user_choice_wins(tmp_path):
    sttm, cache = _sttm_with_load_strategy(tmp_path, "STG: Append; STD: Append")
    docx = _pair1_docx_without({"Object/data Format": "Pipe Delimited Text"})
    pair = _resolve(tmp_path, docx, sttm, cache)
    keys = [q.key for q in pair.questions]
    assert keys == ["feeds[0].file_format"]
    (question,) = pair.questions
    assert question.kind == "choice" and question.title == "File format"
    assert [(c["source"], c["value"]) for c in question.candidates] == [
        ("FRD", "Pipe Delimited Text"), ("STTM", pair1.FILE_FORMAT)]
    assert all(c["cell"] for c in question.candidates)
    answers = parse_answers({"gaps": {"feeds[0].file_format": {
        "value": pair1.FILE_FORMAT, "source": "STTM"}}})
    resolved = _resolve(tmp_path, docx, sttm, cache, answers=answers)
    assert resolved.questions == []
    assert resolved.frd_contract.feeds[0].file_format == pair1.FILE_FORMAT
    fill = next(f for f in resolved.gap_fills if f["field"] == "feeds[0].file_format")
    assert fill["source"] == "user" and "STTM" in fill["cell"]
    assert any(f.startswith("frd_unstated:feeds[0].file_format source_used:user")
               for f in resolved.flags)


# -- 3. one layer-less STTM load strategy -> one layer pause ------------------------- #


def test_layerless_sttm_strategy_pauses_with_the_layer_question(tmp_path):
    sttm = SHAPES / "sttm" / "pair_1_family_a.xlsx"   # meta 'Load Strategy' = "Append"
    docx = _pair1_docx_without({"Load Strategy STG", "Load Strategy STD (View)"})
    pair = _resolve(tmp_path, docx, sttm, PROFILES / "mock")
    (question,) = pair.questions
    assert question.key == "feeds[0].load_strategy" and question.kind == "layer"
    assert [c["layer"] for c in question.candidates] == ["stage", "standard", "both"]
    assert {c["value"] for c in question.candidates} == {"Append"}
    assert "without saying which layer" in question.reason
    answers = parse_answers({"gaps": {"feeds[0].load_strategy": {
        "value": "Append", "layer": "both", "source": "STTM"}}})
    resolved = _resolve(tmp_path, docx, sttm, PROFILES / "mock", answers=answers)
    assert resolved.questions == []
    feed = resolved.frd_contract.feeds[0]
    assert (feed.stage_target.load_strategy, feed.standard_target.load_strategy) == (
        "Append", "Append")
    assert sum(1 for f in resolved.gap_fills if f["source"] == "user") == 2
    with pytest.raises(ValueError):
        parse_answers({"gaps": {"feeds[0].load_strategy": {"value": "Append", "layer": "all"}}})


# -- 4. stage band authoritative; domain FRD-only ------------------------------------ #


def test_stage_band_is_never_asked_and_domain_stays_an_frd_question(tmp_path):
    sttm, cache = _sttm_with_load_strategy(tmp_path, "STG: Append; STD: Append")
    # Labels ABSENT from the document (not merely blank): the recognizer has
    # nothing to place, so each would be a question — unless the STTM answers.
    docx = _pair1_docx_without(set(), drop_labels={"Target Catalog and Schema",
                                                   "Target Table Name", "Domain and Subdomain"})
    pair = _resolve(tmp_path, docx, sttm, cache)
    keys = [q.key for q in pair.questions]
    assert "feeds[0].stage_target.schema" not in keys and "feeds[0].stage_target.tables" not in keys
    assert keys == ["feeds[0].domain"]
    assert any(f.startswith("frd_unstated:feeds[0].stage_target.schema source_used:STTM stage band")
               for f in pair.flags)


# -- 5. the CLI path: resolver fallbacks + remaining hard stops ---------------------- #


def test_resolver_fallbacks_and_remaining_hard_stops(tmp_path):
    from codegen.resolve.resolver import ContractMismatchError

    config = load_config_for_tests()
    sttm, cache = _sttm_with_load_strategy(tmp_path, "STG: Append; STD: Upsert")
    pair = _resolve(tmp_path, _pair1_docx_without(set()), sttm, cache)
    frd_json = tmp_path / "frd.contract.json"
    sttm_json = tmp_path / "sttm.contract.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    # Blank the FRD contract's format and strategies: the resolver fills them.
    data = json.loads(frd_json.read_text(encoding="utf-8"))
    feed = data["feeds"][0]
    feed["file_format"] = None
    feed["stage_target"]["load_strategy"] = None
    feed["standard_target"]["load_strategy"] = None
    frd_json.write_text(json.dumps(data), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config)
    assert spec.file_format == pair1.FILE_FORMAT
    assert (spec.stage_load_strategy, spec.standard_load_strategy) == ("Append", "Upsert")
    kinds = sorted(f.split(" ")[0] for f in spec.provenance_flags if f.startswith("frd_unstated"))
    assert kinds == ["frd_unstated:file_format", "frd_unstated:stage_target.load_strategy",
                     "frd_unstated:standard_target.load_strategy"]
    # Every source silent on the format -> the one remaining hard stop names all three.
    data2 = json.loads(sttm_json.read_text(encoding="utf-8"))
    data2["feeds"][0]["meta_rows"].pop("file_format", None)
    (tmp_path / "sttm2.json").write_text(json.dumps(data2), encoding="utf-8")
    with pytest.raises(ContractMismatchError, match="no source states the file format"):
        resolve_contracts(frd_json, tmp_path / "sttm2.json", config)
