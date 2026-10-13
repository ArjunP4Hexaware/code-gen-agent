"""tools/iig_diff.py --export-scrubbed and --golden-structure (2026-10-12).

The golden is the scrubbed SFMC reference IIG (a committed fixture) with
synthetic "client" values injected in a tmp copy; nothing here reads a real
client file."""

from __future__ import annotations

import functools
import importlib.util
import json
import re
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("iig_diff", REPO / "tools" / "iig_diff.py")
iig_diff = importlib.util.module_from_spec(_spec)
sys.modules["iig_diff"] = iig_diff  # dataclasses look their module up
_spec.loader.exec_module(iig_diff)

KEYS = REPO / "tools" / "iig_keys.yaml"
SFMC = REPO / "fixtures" / "reference" / "SFMC_IIG.xlsx"
TOKEN = "ZqClientSecretName42"        # a "client" token: in the golden, not in the fixtures
AGENT_ONLY = "QxAgentOnlyThing77"     # in the agent only, not in the fixtures
LONG_ID = "8675309"                   # a 7-digit id-like number
DENY_TERM = "Blorptastic"             # a local-denylist term
CLIENT_DOMAIN = "clientcorp" + ".com"   # split: the file itself must stay scrub-clean
EMAIL = "jane.doe@" + CLIENT_DOMAIN
WORKSPACE_USERS = "/Workspace" + "/Users/"
SCHED, ADLS = "DATA_FACTORY_PIPELINE_SCHEDULE", "ADLS_DELTA_INGESTION_DETAILS"
STG, NB = "STGDELTA_STDDELTA_INGESTION_DET", "DATABRICKS_NOTEBOOK_DETAILS"
LEAKS = (TOKEN, AGENT_ONLY, LONG_ID, DENY_TERM, "clientcorp", "jane.doe", "jdoe", "someone.else")

pytestmark = pytest.mark.skipif(not SFMC.exists(), reason=f"{SFMC} absent")


@pytest.fixture(autouse=True)
def _fresh_universe():
    iig_diff.fixture_universe.cache_clear()
    yield
    iig_diff.fixture_universe.cache_clear()


def _set(ws, row: int, column: str, value) -> None:
    header = [c.value for c in ws[1]]
    ws.cell(row, header.index(column) + 1).value = value


def _build_pair(tmp_path: Path) -> tuple[Path, Path]:
    golden, agent = tmp_path / "golden.xlsx", tmp_path / "run" / "agent.xlsx"
    wb = openpyxl.load_workbook(SFMC)
    _set(wb[SCHED], 2, "PIPELINE_NAME", f"PL_{TOKEN}")            # a key: still pairs
    _set(wb[ADLS], 2, "SRC_ADLS_PATH", f"{TOKEN}/{LONG_ID}")
    _set(wb[ADLS], 2, "SOURCE", f"mail {EMAIL} or {WORKSPACE_USERS}jdoe@{CLIENT_DOMAIN}/nb")
    _set(wb[NB], 2, "PROCESS_NAME", f"{DENY_TERM}_proc")
    wb.save(golden)
    wb = openpyxl.load_workbook(golden)
    _set(wb[ADLS], 2, "SRC_ADLS_PATH", AGENT_ONLY)
    _set(wb[ADLS], 2, "SOURCE", "SFMC")
    _set(wb[STG], 2, "DOMAIN", TOKEN)                               # the golden token, agent side
    _set(wb[NB], 2, "PROCESS_NAME", "other_proc")
    _set(wb[SCHED], 2, "CREATED_BY", "someone.else")               # needs_assignment
    wb[SCHED].delete_rows(5)                                        # one row dropped
    agent.parent.mkdir()
    wb.save(agent)
    return golden, agent


def _deny(tmp_path: Path, *terms: str) -> Path:
    path = tmp_path / "denylist_local.txt"
    path.write_text("# synthetic terms only\n" + " | ".join(terms) + "\n", encoding="utf-8")
    return path


def _export(tmp_path: Path, deny_terms=(DENY_TERM,)):
    golden, agent = _build_pair(tmp_path)
    # the scrubbing contract on natural-key row labels (identity labels: in
    # tests/test_iig_diff_identity.py)
    payload, diffs, keys = iig_diff._run_pair(golden, agent, tmp_path / "out", KEYS, "key")
    export = tmp_path / "export"
    check = iig_diff.export_scrubbed(diffs, keys, export,
                                     local_denylist=_deny(tmp_path, *deny_terms), harvest={})
    return golden, payload, export, check


def _table(md: str) -> list[list[str]]:
    rows = []
    for line in md.splitlines():
        if line.startswith("| ") and not line.startswith("| ---"):
            rows.append([c.strip() for c in re.split(r"(?<!\\)\|", line)[1:-1]])
    return rows[1:] if rows else rows


def _texts(folder: Path) -> str:
    return "\n".join(p.read_text(encoding="utf-8") for p in sorted(folder.iterdir()))


def test_export_lists_every_introduced_mismatch(tmp_path: Path) -> None:
    golden, payload, export, check = _export(tmp_path)
    assert {p.name for p in export.iterdir()} == {"mismatches.md", "row_keys.md",
                                                  "SELF_CHECK.md"}
    rows = _table((export / "mismatches.md").read_text(encoding="utf-8"))
    t = payload["totals"]
    exported = t["mismatch"] + t["missing-in-agent"] + t["extra-in-agent"]
    assert len(rows) == check["cells"] == exported
    found = {(r[0], r[1], r[2]) for r in rows}
    assert {(ADLS, "SRC_ADLS_PATH", "mismatch"), (ADLS, "SOURCE", "mismatch"),
            (NB, "PROCESS_NAME", "mismatch")} <= found
    assert any(r[0] == STG and r[1] == "DOMAIN" for r in rows)
    # the dropped schedule row: every scored cell of it is missing-in-agent (the
    # scorer counts a one-sided row's every cell), and it is the only one
    keys = iig_diff.load_keys(KEYS)
    wb = openpyxl.load_workbook(golden, read_only=True)
    header = next(wb[SCHED].iter_rows(values_only=True))
    wb.close()
    expected = {h for h in header if h and not keys.is_assignment(h)
                and not keys.is_convention(h)}               # never scored
    missing = [r for r in rows if r[0] == SCHED and r[2] == "missing-in-agent"]
    assert {r[1] for r in missing} == expected and len(missing) == len(expected)
    assert all("golden row 5, agent row -" in r[5] for r in missing)
    # needs_assignment / withheld columns are never exported
    assert not any(r[1] in ("CREATED_BY", "UPDATED_BY", "PIPELINE_ID", "GROUP_ID") for r in rows)
    assert check["rows_listed"] == {"golden": sum(s["rows"]["golden"]
                                                  for s in payload["sheets"].values()),
                                    "agent": sum(s["rows"]["agent"]
                                                 for s in payload["sheets"].values())}


def test_client_tokens_become_consistent_placeholders(tmp_path: Path) -> None:
    _, _, export, check = _export(tmp_path)
    text = _texts(export)
    for leak in LEAKS:
        assert leak.lower() not in text.lower(), leak
    rows = _table((export / "mismatches.md").read_text(encoding="utf-8"))
    by_cell = {(r[0], r[1]): r for r in rows}
    m = re.fullmatch(r"(G\d+)/(G\d+)", by_cell[(ADLS, "SRC_ADLS_PATH")][3])
    assert m, by_cell[(ADLS, "SRC_ADLS_PATH")]
    g_token, g_id = m.groups()
    assert g_token != g_id                                          # two tokens, two numbers
    assert re.fullmatch(r"A\d+", by_cell[(ADLS, "SRC_ADLS_PATH")][4])   # agent-only token
    assert by_cell[(STG, "DOMAIN")][4] == g_token                   # agent side, same G
    assert re.fullmatch(r"\S+ <EMAIL> or <WORKSPACE_USER_PATH>/nb", by_cell[(ADLS, "SOURCE")][3])
    assert by_cell[(NB, "PROCESS_NAME")][3].startswith("<LOCAL_DENYLIST>_")
    # row_keys.md: the same token in a key maps to the same G<n>
    keys_md = (export / "row_keys.md").read_text(encoding="utf-8")
    assert f"| PL_{g_token} | paired |" in keys_md
    assert "unpaired" in keys_md                                    # the dropped row
    assert check["placeholders"]["G"] >= 2 and check["placeholders"]["A"] >= 1


def test_self_check_counts_and_rescan_is_clean(tmp_path: Path) -> None:
    _, _, export, check = _export(tmp_path)
    assert check["rescan"] == {}
    assert check["before"]["email_address"] >= 1
    assert check["before"]["workspace_user_path"] >= 1
    assert check["before"]["local_denylist"] >= 1
    assert check["universe_source"] == "git ls-files fixtures/" or not (REPO / ".git").exists()
    assert check["line"].startswith("SELF_CHECK: ")
    assert check["line"].endswith("re-scan: 0 hits")
    self_check = (export / "SELF_CHECK.md").read_text(encoding="utf-8")
    assert "scrub_check: 0 hits" in self_check
    sc = iig_diff.load_scrub_check()
    assert sc.scan_paths([export], deny={}) == []


def test_a_dirty_rescan_fails_names_the_class_and_deletes_the_evidence(tmp_path: Path) -> None:
    # "pairing" is in row_keys.md's own static text: the re-scan must catch it
    _, _, export, check = _export(tmp_path, deny_terms=(DENY_TERM, "pairing"))
    assert check["rescan"].get("local_denylist", 0) >= 1
    assert not (export / "mismatches.md").exists()
    assert not (export / "row_keys.md").exists()
    assert "FAILED - local_denylist" in check["line"]
    assert "pairing" not in check["line"]
    assert "FAILED - local_denylist" in (export / "SELF_CHECK.md").read_text(encoding="utf-8")


def test_cli_export_scrubbed(tmp_path: Path, capsys, monkeypatch) -> None:
    golden, agent = _build_pair(tmp_path)
    (tmp_path / "run" / "CODEGEN_BUILD.json").write_text(json.dumps(
        {"sha": "abc1234", "line": "codegen 0.0.test · build abc1234 on b"}), encoding="utf-8")
    monkeypatch.setattr(iig_diff, "export_scrubbed", functools.partial(
        iig_diff.export_scrubbed, local_denylist=_deny(tmp_path, DENY_TERM), harvest={}))
    out, export = tmp_path / "out", tmp_path / "export"
    assert iig_diff.main(["--golden", str(golden), "--agent", str(agent), "--out", str(out),
                          "--export-scrubbed", str(export)]) == 0
    printed = capsys.readouterr().out
    line = next(x for x in printed.splitlines() if x.startswith("SELF_CHECK: "))
    assert line.endswith("re-scan: 0 hits")
    assert {p.name for p in out.iterdir()} == {"score.json", "diff.xlsx", "routing.csv"}
    assert (export / "mismatches.md").is_file()


def test_out_is_still_required_without_the_new_flags(tmp_path: Path, capsys) -> None:
    with pytest.raises(SystemExit) as exc:
        iig_diff.main(["--golden", str(SFMC), "--agent", str(SFMC)])
    assert exc.value.code == 2
    assert "the following arguments are required: --out" in capsys.readouterr().err


# ------------------------------------------------------------- the universe


def test_without_git_the_universe_is_a_gitignore_aware_disk_walk(tmp_path, monkeypatch) -> None:
    def no_git(root):
        raise FileNotFoundError("git unavailable")

    monkeypatch.setattr(iig_diff, "_git_ls_fixtures", no_git)
    _, _, export, check = _export(tmp_path)
    assert check["universe_source"].startswith("disk walk of fixtures/ (no git)")
    assert check["rescan"] == {}
    text = _texts(export)
    for leak in LEAKS:
        assert leak.lower() not in text.lower(), leak
    assert "disk walk of fixtures/ (no git)" in (export / "SELF_CHECK.md").read_text(
        encoding="utf-8")


def test_the_disk_walk_skips_gitignored_files(tmp_path: Path) -> None:
    root = tmp_path / "repo"                                    # no .git here
    (root / "fixtures" / "contracts").mkdir(parents=True)
    (root / ".gitignore").write_text("*.xlsx\nfixtures/contracts/\n*.secret\n"
                                     "!fixtures/keep.xlsx\n", encoding="utf-8")
    (root / "fixtures" / "a.txt").write_text("alpha bravo", encoding="utf-8")
    (root / "fixtures" / "contracts" / "c.json").write_text('{"x": "ZqLeaked"}', encoding="utf-8")
    (root / "fixtures" / "x.secret").write_text("ZqOther", encoding="utf-8")
    for name, value in (("keep.xlsx", "Delta9"), ("drop.xlsx", "ZqDropped")):
        wb = openpyxl.Workbook()
        wb.active["A1"] = value
        wb.save(root / "fixtures" / name)
    with zipfile.ZipFile(root / "fixtures" / "f.docx", "w") as z:
        z.writestr("word/document.xml",
                   "<w:document><w:p><w:r><w:t>Char</w:t></w:r><w:r><w:t>lie</w:t></w:r>"
                   "</w:p><w:p><w:r><w:t>echo</w:t></w:r></w:p></w:document>")
    tokens, source, n = iig_diff.fixture_universe(root)
    assert source == "disk walk of fixtures/ (no git), .gitignore honoured"
    assert n == 3
    assert {"alpha", "bravo", "delta9", "charlie", "echo"} <= tokens
    assert not {"zqleaked", "zqother", "zqdropped"} & tokens


@pytest.mark.skipif(not (REPO / ".git").exists(), reason="no git checkout")
def test_the_disk_walk_never_counts_more_than_git_tracks() -> None:
    try:
        tracked = {p.resolve() for p in iig_diff._git_ls_fixtures(REPO)}
    except (OSError, subprocess.SubprocessError):
        pytest.skip("git unavailable")
    walked = {p.resolve() for p in iig_diff._walk_fixtures(REPO)}
    assert walked and walked <= tracked


# --------------------------------------------------------- golden structure


def _direct_counts(path: Path) -> dict[str, dict[str, Counter]]:
    """sheet -> (grouping column, Counter of id -> rows), straight from openpyxl."""
    out = {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        header, *rows = list(ws.iter_rows(values_only=True))
        rows = [r for r in rows if any(v not in (None, "") for v in r)]
        col = next((c for c in ("GROUP_ID", "PIPELINE_ID") if c in header), None)
        if col:
            i = header.index(col)
            out[ws.title] = (col, Counter(str(r[i]) for r in rows if r[i] not in (None, "")))
    wb.close()
    return out


def test_golden_structure_on_sfmc(tmp_path: Path, capsys) -> None:
    out = tmp_path / "out"
    assert iig_diff.main(["--golden", str(SFMC), "--golden-structure", "--out", str(out)]) == 0
    printed = capsys.readouterr().out
    data = json.loads((out / "golden_structure.json").read_text(encoding="utf-8"))
    md = (out / "golden_structure.md").read_text(encoding="utf-8")
    direct = _direct_counts(SFMC)
    prefix = {"GROUP_ID": "group", "PIPELINE_ID": "pipeline"}
    ids: dict[str, set[str]] = {"GROUP_ID": set(), "PIPELINE_ID": set()}
    wb = openpyxl.load_workbook(SFMC, read_only=True, data_only=True)
    for ws in wb.worksheets:
        header, *rows = list(ws.iter_rows(values_only=True))
        for col in ids:
            if col in header:
                ids[col] |= {str(r[header.index(col)]) for r in rows
                             if r[header.index(col)] not in (None, "")}
    wb.close()
    assert data["group_ids"] == len(ids["GROUP_ID"]) == 2
    assert data["pipeline_ids"] == len(ids["PIPELINE_ID"]) == 4
    for sheet, (col, counts) in direct.items():
        assert data["sheets"][sheet]["grouped_by"] == col
        order = sorted(ids[col], key=lambda v: v.casefold())
        for value, n in counts.items():
            label = f"{prefix[col]}#{order.index(value) + 1}"
            assert data["groups"][label]["sheets"][sheet]["rows"] == n, (sheet, label)
    assert data["sheets"]["EMAIL_TEMPLATE_CONFIG"]["grouped_by"] is None
    assert data["sheets"][SCHED]["grouped_by"] == "PIPELINE_ID"
    g1, g2 = data["groups"]["group#1"], data["groups"]["group#2"]
    assert g1["sheets"][ADLS]["distinct"] == {"OBJECT_NAME": 1, "TGT_TABLE_NAME": 1}
    assert g1["sheets"][ADLS]["shared_rows"] == {"TGT_TABLE_NAME": 1, "OBJECT_NAME": 1}
    assert g1["pairs_with"]["TGT_TABLE_NAME"] == ["group#2"]
    assert g2["reads_tables_of"] == ["group#1"]
    assert g1["linked"] == ["pipeline#1"]
    assert "pipeline#1 (grouped by PIPELINE_ID)" in md
    assert "PIPELINE_ID (sheet has no GROUP_ID)" in md
    for text in (printed, md, json.dumps(data)):
        assert "real GROUP_IDs withheld" in text
        for value in ids["GROUP_ID"] | ids["PIPELINE_ID"]:
            assert value.casefold() not in text.casefold()
        assert "sfmc_member_outreach_email_tracking" not in text.casefold()


def test_golden_structure_alone_prints_and_writes_nothing(tmp_path: Path, capsys) -> None:
    book = tmp_path / "g.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ADLS
    ws.append(["GROUP_ID", "OBJECT_ID", "OBJECT_NAME", "TGT_TABLE_NAME", "PIPELINE_ID"])
    for row in ([730063, 1, "obj_a", "tbl_x", 730263], [730064, 1, "obj_b", "tbl_y", 730264],
                [730099, 1, "obj_c", "tbl_x", 730299], [730099, 2, "obj_d", "tbl_z", 730299]):
        ws.append(row)
    wb.save(book)
    before = sorted(tmp_path.iterdir())
    assert iig_diff.main(["--golden", str(book), "--golden-structure"]) == 0
    assert sorted(tmp_path.iterdir()) == before
    printed = capsys.readouterr().out
    for value in ("730063", "730064", "730099", "730263", "tbl_x", "obj_a"):
        assert value not in printed
    data = iig_diff.golden_structure(book, iig_diff.load_keys(KEYS))
    assert [g["sheets"][ADLS]["rows"] for g in data["groups"].values()
            if g["column"] == "GROUP_ID"] == [1, 1, 2]
    g3 = data["groups"]["group#3"]["sheets"][ADLS]
    assert g3["distinct"]["TGT_TABLE_NAME"] == 2
    assert g3["shared_rows"]["TGT_TABLE_NAME"] == 1
    assert g3["pairs_with"]["TGT_TABLE_NAME"] == ["group#1"]
    assert data["groups"]["group#2"]["pairs_with"]["TGT_TABLE_NAME"] == []
    assert "pairs with group#k on TGT_TABLE_NAME hash" in printed
