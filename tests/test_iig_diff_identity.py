"""iig_diff: pair rows by what they describe, not by generated names (2026-10-13).

D1 pair_2 paired zero rows: the prod goldens name pipelines / notebooks in the
legacy convention (``PL_File_<DOMAIN>_<PKG>_<FLOW>_Incr``,
``FileInvokeDQDeltaTable_V0.1``), the agent in the EDO standard
(``WF_DLK_NSP_<TABLE>_<DOMAIN>_<SUBDOM>_<REGION>_<FREQ>``, ``NB_..._INGEST``).
Now: each golden row's feed + flow from the group index (GROUP_ID /
PIPELINE_ID -> flow_type + matched_table), each agent row's from its workbook
and semantic cells; per sheet an IDENTITY (iig_keys.yaml ``identity``); the
convention columns never pair and never score; a sheet without identity pairs
by position only when the counts match; every sheet says how its rows paired.
"""

from __future__ import annotations

import csv
import importlib.util
import json
import sys
import time
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("iig_diff", REPO / "tools" / "iig_diff.py")
iig_diff = importlib.util.module_from_spec(_spec)
sys.modules["iig_diff"] = iig_diff
_spec.loader.exec_module(iig_diff)

KEYS = REPO / "tools" / "iig_keys.yaml"
PIPE = "DATA_FACTORY_PIPELINE_SCHEDULE"
ADLS = "ADLS_DELTA_INGESTION_DETAILS"
STG = "STGDELTA_STDDELTA_INGESTION_DET"
NB = "DATABRICKS_NOTEBOOK_DETAILS"
DQ = "DATA_QUALITY_RULES"
EMAIL = "EMAIL_TEMPLATE_CONFIG"

# Synthetic "Northwind" feeds in the REAL naming shapes.
FEEDS = {  # feed (= stage table) -> (domain, package)
    "nw_member_risk": ("MEMBER", "RISK_PKG"),
    "nw_claim_events": ("CLAIM", "EVENTS_PKG"),
}
FLOWS = (("file_to_stg", "ADLS_To_Delta", "STAGE", "Stage Load for"),
         ("stg_to_std", "Delta_To_STD", "STANDARD", "Standard Load for"))
PIPE_HEAD = ["PIPELINE_ID", "PIPELINE_NAME", "PIPELINE_FREQUENCY", "ACTIVE_FLAG", "LAYER_NAME",
             "CREATED_BY"]
ADLS_HEAD = ["GROUP_ID", "OBJECT_ID", "OBJECT_NAME", "PIPELINE_ID", "TGT_TABLE_NAME",
             "TGT_DATABASE_NAME", "TGT_LOAD_OPTION", "SRC_FORMAT"]
STG_HEAD = ["GROUP_ID", "OBJECT_ID", "PIPELINE_ID", "TGT_TABLE_NAME", "TGT_SCHEMA_NAME",
            "TGT_LOAD_OPTION"]
NB_HEAD = ["PIPELINE_ID", "PIPELINE_NAME", "GROUP_ID", "SEQ_NM", "PROCESS_NAME",
           "TGT_REFRESH_TYPE", "DATABRICKS_NOTEBOOK_NAME"]
DQ_HEAD = ["GROUP_ID", "OBJECT_ID", "RULE_TYPE", "RULE_CLASS", "SOURCE_COLUMN", "INPUT_PARAM"]


def _book(path: Path, sheets: dict[str, list[list]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _ids(feed_no: int, flow_no: int) -> tuple[int, int]:
    """(GROUP_ID, PIPELINE_ID) of a feed x flow, as the prod export numbers them."""
    group = 180250 + 2 * feed_no + flow_no
    return group, group + 200


def _index(path: Path, feeds: dict = FEEDS) -> Path:
    with path.open("w", encoding="utf-8", newline="") as fh:
        out = csv.writer(fh)
        out.writerow(["pair_id", "group_id", "pipeline_id", "object_id", "flow_type",
                      "match_route", "confidence", "matched_table"])
        for f, feed in enumerate(feeds):
            for k, (flow, *_rest) in enumerate(FLOWS):
                group, pipeline = _ids(f, k)
                out.writerow(["pair_x", group, pipeline, 1, flow, "adls", "HIGH", feed])
    return path


def _golden(path: Path, feeds: dict = FEEDS) -> Path:
    """The prod shape: legacy names, one GROUP_ID per feed x flow, ids filled."""
    pipe, adls, stg, nb, dq = [PIPE_HEAD], [ADLS_HEAD], [STG_HEAD], [NB_HEAD], [DQ_HEAD]
    for f, (feed, (domain, pkg)) in enumerate(feeds.items()):
        for k, (_flow, leg, layer, process) in enumerate(FLOWS):
            group, pipeline = _ids(f, k)
            pipe.append([pipeline, f"PL_File_{domain}_{pkg}_{leg}_Incr", "Monthly", "Y", layer,
                         "prod.user"])
            notebook = ("FileInvokeDQDeltaTable_V0.1" if k == 0
                        else "DeltaInvokeDQStandardDeltaTable_V2")
            nb.append([pipeline, f"PL_File_{domain}_{pkg}_{leg}_Incr", group, k + 1,
                       f"{process} {feed}", "Append" if k == 0 else "Upsert", notebook])
            if k == 0:
                adls.append([group, 1, f"{feed.upper()}_YYYYMMDD.psv", pipeline,
                             f"stg_nw.{feed}", "stg_nw", "Append", "csv"])
                dq.append([group, 1, "Predefined", "CheckRule", "member_id", "NOT NULL"])
                dq.append([group, 1, "Custom", "DateFormat", "enroll_date", "yyyy-MM-dd"])
            else:
                stg.append([group, 1, pipeline, feed, "nw", "Upsert"])
    return _book(path, {PIPE: pipe, ADLS: adls, STG: stg, NB: nb, DQ: dq})


def _agent(folder: Path, feed: str, domain: str) -> Path:
    """The agent's iig_v1 shape: EDO names, every id blank (assigned later)."""
    wf = f"WF_DLK_NSP_{feed.upper()}_{domain}_CARE_MANAGEMENT_MI_MTH"
    nb = f"NB_DLK_NSP_{domain}_CARE_MANAGEMENT_INGEST"
    return _book(folder / feed / f"{feed}_IIG.xlsx", {
        PIPE: [PIPE_HEAD] + [[None, wf, "Monthly", "Y", layer, None]
                             for _f, _l, layer, _p in FLOWS],
        ADLS: [ADLS_HEAD, [None, None, f"{feed.upper()}_YYYYMMDD.psv", None,
                           f"stg_nw.{feed}", "stg_nw", "Append", "csv"]],
        STG: [STG_HEAD, [None, None, None, feed, "nw", "Upsert"]],
        NB: [NB_HEAD] + [[None, wf, None, k + 1, f"{process} {feed}",
                          "Append" if k == 0 else "Upsert", nb]
                         for k, (_f, _l, _layer, process) in enumerate(FLOWS)],
        DQ: [DQ_HEAD, [None, None, "Predefined", "CheckRule", "member_id", "NOT NULL"],
             [None, None, "Custom", "DateFormat", "enroll_date", "yyyy-MM-dd"]],
    })


def _stamp(folder: Path) -> None:
    (folder / "CODEGEN_BUILD.json").write_text(json.dumps(
        {"sha": "abc1234", "line": "codegen 0.0.test · build abc1234 on b"}), encoding="utf-8")


def _run(tmp_path: Path, golden: Path, agents: list[Path], **kw) -> dict:
    index = _index(tmp_path / "index.csv")
    return iig_diff.run_pair(golden, agents if len(agents) > 1 else agents[0],
                             tmp_path / "out", KEYS, group_index=[index], **kw)


def test_real_shapes_pair_by_identity_comparable_100(tmp_path: Path) -> None:
    """The brief's test: golden PIPELINE_NAME PL_File_<DOMAIN>_<PKG>_<FLOW>_Incr,
    agent WF_DLK_NSP_<TABLE>_<DOMAIN>_<SUBDOM>_<REGION>_<FREQ>; identity columns
    and every non-key cell identical -> every row paired by identity, comparable
    100, the convention cells counted as differing, zero mismatches."""
    golden = _golden(tmp_path / "golden.xlsx")
    run = tmp_path / "run"
    agents = [_agent(run, feed, domain) for feed, (domain, _pkg) in FEEDS.items()]
    payload = _run(tmp_path, golden, agents)
    for sheet, rows in ((PIPE, 4), (ADLS, 2), (STG, 2), (NB, 4), (DQ, 4)):
        s = payload["sheets"][sheet]
        assert s["pairing"] == {"identity": rows, "position": 0, "key": 0}, (sheet, s)
        assert s["unpaired"] == {"golden": 0, "agent": 0}, (sheet, s["warnings"])
        assert s["counts"]["mismatch"] == 0, sheet
        assert s["counts"]["missing-in-agent"] == s["counts"]["extra-in-agent"] == 0, sheet
        assert s["comparable_score"] == 100.0, sheet
    t = payload["totals"]
    assert t["comparable_score"] == 100.0 and t["mismatch"] == 0
    assert t["pairing"]["identity"] == 16 and t["unpaired"] == {"golden": 0, "agent": 0}
    # PIPELINE_NAME (schedule + notebook) and DATABRICKS_NOTEBOOK_NAME: 12 cells, all differ
    assert t["convention_cells"] == {"cells": 12, "differing": 12}
    assert t["convention"] == 12                       # a status of its own, never scored
    # the identities that paired: feed x flow, never a name
    assert payload["sheets"][PIPE]["key"] == ["@feed", "@flow"]
    assert payload["sheets"][ADLS]["key"] == ["@feed", "@flow", "TGT_TABLE_NAME"]


def test_a_golden_row_of_a_feed_the_agent_did_not_write_is_out_of_scope(tmp_path: Path) -> None:
    every = {**FEEDS, "nw_provider_roster": ("PROVIDER", "ROSTER_PKG")}
    golden = _golden(tmp_path / "golden.xlsx", every)
    _index(tmp_path / "index.csv", every)
    run = tmp_path / "run"
    agents = [_agent(run, feed, domain) for feed, (domain, _pkg) in FEEDS.items()]
    payload = iig_diff.run_pair(golden, agents, tmp_path / "out", KEYS,
                                group_index=[tmp_path / "index.csv"])
    pipe = payload["sheets"][PIPE]
    assert pipe["out_of_scope_rows"] == 2                    # the roster feed's two pipelines
    assert pipe["unpaired"] == {"golden": 0, "agent": 0}
    assert pipe["comparable_score"] == 100.0


def test_a_flow_the_agent_never_states_is_dropped_and_said(tmp_path: Path) -> None:
    """EMAIL_TEMPLATE_CONFIG: the golden has a template per pipeline (feed x
    flow x status), the agent one per feed (no flow column) — the flow is
    dropped from that sheet's identity, said, and the extra golden rows stay
    unpaired, counted."""
    golden = _book(tmp_path / "g.xlsx", {EMAIL: [
        ["PIPELINE_ID", "TEMPLATE_NAME", "STATUS", "SUBJECT"],
        [180450, "Member Risk", "Success", "done"], [180450, "Member Risk", "Failed", "fail"],
        [180451, "Member Risk", "Success", "done"], [180451, "Member Risk", "Failed", "fail"]]})
    agent = _book(tmp_path / "run" / "nw_member_risk" / "nw_member_risk_IIG.xlsx", {EMAIL: [
        ["TEMPLATE_NAME", "STATUS", "SUBJECT"],
        ["nw_member_risk", "Success", "done"], ["nw_member_risk", "Failed", "fail"]]})
    s = _run(tmp_path, golden, [agent])["sheets"][EMAIL]
    assert s["key"] == ["@feed", "STATUS"]
    assert any("@flow" in d for d in s["identity_dropped"])
    assert s["pairing"]["identity"] == 2 and s["unpaired"] == {"golden": 2, "agent": 0}
    assert s["counts"]["mismatch"] == 0
    assert s["convention_cells"] == {"cells": 2, "differing": 2}     # TEMPLATE_NAME


def test_no_identity_pairs_by_position_only_when_counts_match(tmp_path: Path) -> None:
    head = ["RUN_ID", "STATUS"]
    golden = _book(tmp_path / "g.xlsx", {"FILE_LOG_INFORMATION": [head, ["r1", "ok"],
                                                                  ["r2", "ok"]]})
    equal = _book(tmp_path / "run" / "a" / "a_IIG.xlsx", {"FILE_LOG_INFORMATION": [
        head, ["x1", "ok"], ["x2", "ok"]]})
    s = _run(tmp_path, golden, [equal])["sheets"]["FILE_LOG_INFORMATION"]
    assert s["pairing"] == {"identity": 0, "position": 2, "key": 0}
    assert any("BY POSITION" in w for w in s["warnings"])
    fewer = _book(tmp_path / "run" / "b" / "b_IIG.xlsx", {"FILE_LOG_INFORMATION": [
        head, ["x1", "ok"]]})
    s = _run(tmp_path, golden, [fewer])["sheets"]["FILE_LOG_INFORMATION"]
    assert s["pairing"] == {"identity": 0, "position": 0, "key": 0}
    assert s["unpaired"] == {"golden": 2, "agent": 1}
    assert any("row counts differ" in w and "NOT paired" in w for w in s["warnings"])


def test_an_incomplete_identity_leaves_the_row_unpaired_and_says_so(tmp_path: Path) -> None:
    golden = _golden(tmp_path / "golden.xlsx")
    run = tmp_path / "run"
    agents = [_agent(run, feed, domain) for feed, (domain, _pkg) in FEEDS.items()]
    wb = openpyxl.load_workbook(agents[0])
    wb[ADLS]["E2"] = None                                   # TGT_TABLE_NAME blank
    wb.save(agents[0])
    s = _run(tmp_path, golden, agents)["sheets"][ADLS]
    assert s["pairing"]["identity"] == 1
    assert s["unpaired"] == {"golden": 1, "agent": 1}
    assert any("identity incomplete" in w for w in s["warnings"])


def test_3000_row_all_files_sheet_in_scope_diffs_in_seconds(tmp_path: Path) -> None:
    """D1 stripped ALL_FILES_STATIC_INFORMATION (2,995 rows) to dodge a 600 s
    timeout: diff.xlsx called ws.max_row per cell (quadratic) and the scope test
    compared every row with every stem. Both sides 3,000 rows + a 3,000-row
    golden against 23 agent rows (2,977 out of scope), scope ON."""
    head = ["INVENTORY_ID", "FILE_NAME", "DESCRIPTION", "LOB", "FREQUENCY", "SUPPLIER"]
    rows = [[i, f"VENDOR{i % 97}_FILE_{i}_YYYYMMDD.txt", f"desc {i}", "ALL", "Monthly",
             f"v{i % 97}"] for i in range(3000)]
    golden = _book(tmp_path / "g.xlsx", {"ALL_FILES_STATIC_INFORMATION": [head, *rows]})
    full = _book(tmp_path / "run" / "full" / "full_IIG.xlsx", {
        "ALL_FILES_STATIC_INFORMATION": [head, *[[None, *r[1:]] for r in rows]]})
    start = time.perf_counter()
    s = _run(tmp_path, golden, [full])["sheets"]["ALL_FILES_STATIC_INFORMATION"]
    both = time.perf_counter() - start
    assert s["pairing"]["identity"] == 3000 and s["comparable_score"] == 100.0
    few = _book(tmp_path / "run" / "few" / "few_IIG.xlsx", {
        "ALL_FILES_STATIC_INFORMATION": [head, *[[None, *r[1:]] for r in rows[:23]]]})
    start = time.perf_counter()
    s = _run(tmp_path, golden, [few])["sheets"]["ALL_FILES_STATIC_INFORMATION"]
    scoped = time.perf_counter() - start
    assert s["pairing"]["identity"] == 23
    assert both < 30 and scoped < 30, (both, scoped)


def test_key_samples_and_the_export_show_identities(tmp_path: Path, capsys) -> None:
    golden = _golden(tmp_path / "golden.xlsx")
    run = tmp_path / "run"
    agents = [_agent(run, feed, domain) for feed, (domain, _pkg) in FEEDS.items()]
    _stamp(run)
    out, export = tmp_path / "out", tmp_path / "export"
    assert iig_diff.main(["--golden", str(golden), "--agent", *map(str, agents),
                          "--out", str(out), "--group-index", str(_index(tmp_path / "i.csv")),
                          "--export-scrubbed", str(export)]) == 0
    printed = capsys.readouterr().out
    sheet = json.loads((out / "score.json").read_text(encoding="utf-8"))["sheets"][PIPE]
    assert sheet["key_samples"]["golden"][:2] == ["feed#1 | file_to_stg", "feed#1 | stg_to_std"]
    assert sheet["key_samples"]["agent"][:2] == ["feed#1 | file_to_stg", "feed#1 | stg_to_std"]
    assert f"{PIPE}: score" in printed and "paired by identity 4, position 0" in printed
    assert "PAIRING: identity 16, position 0, key 0 row(s)" in printed
    assert "KEY CONVENTIONS (D1 only - unscrubbed" in printed
    row_keys = (export / "row_keys.md").read_text(encoding="utf-8")
    assert "key columns used for pairing: @feed, @flow" in row_keys
    assert "| feed#1 \\| file_to_stg | paired |" in row_keys
    assert "PL_File_" not in row_keys                       # names are not the label


def test_identity_values_drop_case_separators_and_the_schema_prefix(tmp_path: Path) -> None:
    """``stg_nw.nw_member_risk`` (golden) and ``NW-MEMBER-RISK`` (agent) are one
    target table for pairing; the cell itself still compares as written."""
    golden_table = iig_diff.identity_value("stg_nw.nw_member_risk", "TGT_TABLE_NAME")
    agent_table = iig_diff.identity_value("NW-MEMBER-RISK", "TGT_TABLE_NAME")
    assert golden_table == agent_table == "nwmemberrisk"
    assert iig_diff.identity_value("Vendor.File_1.txt", "FILE_NAME") == "vendorfile1txt"
    golden = _golden(tmp_path / "golden.xlsx")
    run = tmp_path / "run"
    agents = [_agent(run, feed, domain) for feed, (domain, _pkg) in FEEDS.items()]
    wb = openpyxl.load_workbook(agents[0])
    wb[ADLS]["E2"] = "NW-MEMBER-RISK"
    wb.save(agents[0])
    s = _run(tmp_path, golden, agents)["sheets"][ADLS]
    assert s["pairing"]["identity"] == 2 and s["unpaired"] == {"golden": 0, "agent": 0}
    assert s["counts"]["mismatch"] == 1                     # the table cell as written


def test_the_identity_check_tool_counts_a_corpus(tmp_path: Path, capsys) -> None:
    spec = importlib.util.spec_from_file_location(
        "iig_identity_check", REPO / "tools" / "iig_identity_check.py")
    tool = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(tool)
    corpus = tmp_path / "corpus"
    _golden(corpus / "g1.xlsx")
    _golden(corpus / "g2.xlsx")
    index = _index(tmp_path / "index.csv")
    assert tool.main([str(corpus), "--group-index", str(index)]) == 0
    out = capsys.readouterr().out
    assert "IDENTITY CHECK over 2 golden(s)" in out
    assert (f"  {PIPE}: rows 8, identified 8, by position 0, unidentified 0, "
            "repeated identities 0") in out
    assert "PL_File_" not in out                            # counts only
