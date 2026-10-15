"""iig_diff across templates (2026-10-13): an iig_v1 golden's
ALL_FILES_STATIC_INFORMATION against an iig_v2 agent's FILE_ADLS_INGESTION_DETAILS.

The two sheets hold different grains: ALL_FILES is one row per FILE of the
inventory, FILE_ADLS one row per FEED, naming no file (the pair_1 golden: one
row for a four-file feed). A golden file row pairs with the FILE_ADLS row of the
agent feed whose ADLS_DELTA_INGESTION_DETAILS SRC_FILE_NAME is that file (case,
separators, extension and CCYY_MM vs YYYY_MM normalized), scoped like every
sheet; only the headers both templates have are compared, the rest are "not
comparable". ADLS_FIXED_WIDTH_HANDLER stays "sheet not in golden".
"""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import openpyxl
import yaml

REPO = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("iig_diff", REPO / "tools" / "iig_diff.py")
iig_diff = importlib.util.module_from_spec(_spec)
sys.modules["iig_diff"] = iig_diff
_spec.loader.exec_module(iig_diff)

KEYS = REPO / "tools" / "iig_keys.yaml"
V1 = "ALL_FILES_STATIC_INFORMATION"
V2 = "FILE_ADLS_INGESTION_DETAILS"
ADLS = "ADLS_DELTA_INGESTION_DETAILS"
FIXED = "ADLS_FIXED_WIDTH_HANDLER"
LABEL = f"{V1} → {V2}"


def _template_headers() -> tuple[list[str], list[str]]:
    config = yaml.safe_load((REPO / "config" / "config.yaml").read_text(encoding="utf-8"))
    v1 = config["demo"]["metadata_sheet"]["tabs"][V1]["headers"]
    v2 = config["metadata"]["templates"]["iig_v2"]["tabs"][V2]["headers"]
    return v1, v2


V1_HEAD, V2_HEAD = _template_headers()
ADLS_HEAD = ["GROUP_ID", "OBJECT_ID", "OBJECT_NAME", "DOMAIN", "SUBDOMAIN", "SRC_FILE_NAME",
             "TGT_TABLE_NAME"]

# Synthetic "Northwind" feeds: feed -> (domain, subdomain, its files as the golden
# inventory writes them, as the agent's SRC_FILE_NAME writes them)
FEEDS = {
    "nw_member_risk": ("Member", "Risk", ["NW_MEMBER_RISK_CCYY_MM.csv"],
                       ["NW_MEMBER_RISK_YYYY_MM.csv"]),
    "nw_claim_events": ("Claim", "Events",
                        ["NW_CLAIM_EVENTS_HDR_CCYYMMDD.txt", "nw-claim-events-dtl_ccyymmdd.TXT"],
                        ["NW_CLAIM_EVENTS_HDR_*.txt", "NW_CLAIM_EVENTS_DTL*"]),
}
# The rest of the inventory: other feeds' files (out of scope, never missing)
OTHER_FILES = ["NW_PROVIDER_ROSTER_CCYYMMDD.csv", "NW_PHARMACY_FILL_CCYYMM.dat"]


def _book(path: Path, sheets: dict[str, list[list]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _row(head: list[str], **cells) -> list:
    return [cells.get(h) for h in head]


def _golden(path: Path, others: bool = True) -> Path:
    rows = []
    for i, (_feed, (domain, sub, files, _agent)) in enumerate(FEEDS.items()):
        rows += [_row(V1_HEAD, INVENTORY_ID=900 + 10 * i + j, OBJECT_ID=j + 1, FILE_NAME=f,
                      DESCRIPTION=f"inventory row {f}", FREQUENCY="Monthly", domain=domain,
                      subdomain=sub, vendor_name="Northwind")
                 for j, f in enumerate(files)]
    rows += [_row(V1_HEAD, INVENTORY_ID=990 + k, FILE_NAME=f, domain="Other", subdomain="Other")
             for k, f in enumerate(OTHER_FILES if others else [])]
    return _book(path, {V1: [V1_HEAD, *rows]})


def _agent(run: Path) -> list[Path]:
    books = []
    for feed, (domain, sub, _golden_files, files) in FEEDS.items():
        books.append(_book(run / feed / f"{feed}_IIG.xlsx", {
            V2: [V2_HEAD, _row(V2_HEAD, DOMAIN=domain, SUBDOMAIN=sub, SRC_ROOT_DIR="inbound",
                               TGT_ADLS_PATH=f"landing/{feed}", SOURCE_TYPE="File")],
            ADLS: [ADLS_HEAD, *[[None, j + 1, f"obj_{j}", domain, sub, f, f"stg_{feed}"]
                                for j, f in enumerate(files)]],
            FIXED: [["SEGMENT", "TGT_TABLE", "LEN"], ["D", f"stg_{feed}", 10]],
        }))
    (run / "CODEGEN_BUILD.json").write_text(json.dumps(
        {"sha": "abc1234", "line": "codegen 0.0.test · build abc1234 on b"}), encoding="utf-8")
    return books


def test_the_column_mapping_is_the_headers_both_templates_have() -> None:
    """Derived, not guessed: sheet_aliases.columns is exactly the intersection
    of the two templates' headers (case / separators normalized)."""
    alias = next(a for a in iig_diff.load_keys(KEYS).sheet_aliases if a.golden_sheet == V1)
    assert alias.agent_sheet == V2
    hid = iig_diff.header_id
    shared = {hid(h) for h in V1_HEAD} & {hid(h) for h in V2_HEAD}
    assert {hid(g) for g, _a in alias.columns} == shared
    assert all(hid(g) == hid(a) for g, a in alias.columns)
    assert all(g in V1_HEAD for g, _a in alias.columns)
    assert all(a in V2_HEAD for _g, a in alias.columns)


def test_file_identity_normalizes_case_separators_and_date_tokens() -> None:
    tokens = iig_diff.load_keys(KEYS).file_name_tokens
    same = ["NW_X_CCYY_MM.csv", "nw-x-yyyy-mm.CSV", "NW_X_*.csv", "NW_X*", "in/NW_X_202410.csv"]
    assert {iig_diff.file_identity(v, tokens) for v in same} == {"nwx"}
    # a date / wildcard is dropped wherever it sits: end, inside, start
    assert iig_diff.file_identity("NW_X_HDR_CCYYMMDD.txt", tokens) == "nwxhdr"
    assert iig_diff.file_identity("NW_X_HDR_*.txt", tokens) == "nwxhdr"
    assert iig_diff.file_identity("NW_X_CCYYMMDD_HDR.txt", tokens) == "nwxhdr"
    assert iig_diff.file_identity("NW_X_*_HDR.txt", tokens) == "nwxhdr"
    # pair_12 (CAQH-shaped, synthetic): the inventory puts the date FIRST and a
    # wildcard INSIDE; the agent's SRC_FILE_NAME drops both and ends in '*';
    # the prod golden collapses the LOB suffix into its wildcard (2026-10-14)
    caqh = {iig_diff.file_identity(v, tokens) for v in (
        "YYYYMMDD_2041_7310_NW_*_P_I_1_T.txt", "2041_7310_NW_P_I_1_T*",
        "20261013_2041_7310_NW_P_I_1_T.txt", "*2041_7310_NW*.txt", "2041_7310_NW_P_I_1_T")}
    assert caqh == {"20417310nw"}
    # one one-letter segment is a name, not a suffix
    assert iig_diff.file_identity("NW_X_CCYYMMDD.txt", tokens) == "nwx"
    assert iig_diff.file_identity("2041_7311_NW_P_I_1_T*", tokens) != caqh.pop()
    assert iig_diff.file_identity("NW_Y_CCYYMMDD.txt", tokens) != "nwx"
    assert iig_diff.file_identity("", tokens) is None


def test_v1_all_files_against_v2_file_adls_all_paired_score_100(tmp_path: Path, capsys) -> None:
    """The brief's test: identical shared columns -> every row paired, score 100."""
    golden = _golden(tmp_path / "golden.xlsx", others=False)
    books = _agent(tmp_path / "run")
    out = tmp_path / "out"
    assert iig_diff.main(["--golden", str(golden), "--agent", *map(str, books),
                          "--out", str(out)]) == 0
    printed = capsys.readouterr().out
    payload = json.loads((out / "score.json").read_text(encoding="utf-8"))
    sheets = payload["sheets"]

    s = sheets[LABEL]
    assert V1 not in sheets and V2 not in sheets           # one mapped pair, not two halves
    assert s["presence"] == "compared"
    assert s["pairing"] == {"identity": 3, "position": 0, "key": 0}
    assert s["unpaired"] == {"golden": 0, "agent": 0}
    assert s["out_of_scope_rows"] == 0
    assert s["score"] == 100.0 and s["comparable_score"] == 100.0
    counts = s["counts"]
    assert counts["mismatch"] == counts["missing-in-agent"] == counts["extra-in-agent"] == 0
    assert counts["exact"] == 3 * 2                         # domain + subdomain per file
    # only the shared headers are compared; the rest are not comparable, per side
    assert s["mapped_columns"] == {"OBJECT_ID": "OBJECT_ID", "domain": "DOMAIN",
                                   "subdomain": "SUBDOMAIN", "CREATED_DATE": "CREATED_DATE",
                                   "UPDATED_DATE": "UPDATED_DATE"}
    assert "FILE_NAME" in s["not_comparable"]["golden"]
    assert "DESCRIPTION" in s["not_comparable"]["golden"]
    assert "TGT_ADLS_PATH" in s["not_comparable"]["agent"]
    assert "DOMAIN" not in s["not_comparable"]["agent"]
    compared = {c for c in counts if counts[c]}
    assert compared <= {"exact", "needs_assignment"}

    # the fixed-width handler has no v1 counterpart: reported, not scored
    assert sheets[FIXED]["note"] == "sheet not in golden" and sheets[FIXED]["score"] is None
    assert payload["totals"]["comparable_score"] == 100.0

    line = next(x for x in printed.splitlines() if x.startswith(LABEL + ":"))
    assert "paired by identity 3, position 0" in line
    assert "unpaired golden 0, agent 0" in line
    assert "not comparable: golden" in line
    assert openpyxl.load_workbook(out / "diff.xlsx", read_only=True).sheetnames.count(V1) == 1


def test_a_differing_cell_and_a_file_the_agent_did_not_name(
        tmp_path: Path) -> None:
    golden = _golden(tmp_path / "golden.xlsx")
    run = tmp_path / "run"
    books = _agent(run)
    # the claim feed's agent workbook loses its detail file pattern
    feed = "nw_claim_events"
    domain, sub, _g, files = FEEDS[feed]
    _book(books[1], {
        V2: [V2_HEAD, _row(V2_HEAD, DOMAIN=domain, SUBDOMAIN="Wrong")],
        ADLS: [ADLS_HEAD, [None, 1, "obj_0", domain, sub, files[0], f"stg_{feed}"]],
    })
    payload = iig_diff.run_pair(golden, books, tmp_path / "out", KEYS)
    s = payload["sheets"][LABEL]
    assert s["pairing"]["identity"] == 2
    assert s["counts"]["mismatch"] == 1                     # SUBDOMAIN of the header file
    assert s["comparable_score"] < 100.0
    # scoped like every sheet: a file no agent feed names, naming only itself, is
    # out of scope; --scope all scores it (and the other feeds' files) as missing
    assert s["out_of_scope_rows"] == len(OTHER_FILES) + 1 and s["unpaired"]["golden"] == 0
    assert s["score"] < s["comparable_score"]               # the overall score counts them
    every = iig_diff.run_pair(golden, books, tmp_path / "all", KEYS, scope_mode="all")
    s = every["sheets"][LABEL]
    assert s["out_of_scope_rows"] == 0 and s["unpaired"]["golden"] == len(OTHER_FILES) + 1


def test_no_alias_when_both_sides_have_the_golden_sheet(tmp_path: Path) -> None:
    """An iig_v1 agent (it writes ALL_FILES itself) compares sheet with sheet."""
    golden = _golden(tmp_path / "golden.xlsx")
    run = tmp_path / "run"
    book = _book(run / "a" / "a_IIG.xlsx", {V1: [V1_HEAD]})
    (run / "CODEGEN_BUILD.json").write_text(json.dumps({"sha": "abc1234"}), encoding="utf-8")
    payload = iig_diff.run_pair(golden, book, tmp_path / "out", KEYS)
    assert V1 in payload["sheets"] and LABEL not in payload["sheets"]


def test_caqh_shaped_names_pair_and_the_file_names_print_d1_only(tmp_path: Path,
                                                                  capsys) -> None:
    """pair_12 (2026-10-13): ALL_FILES -> FILE_ADLS paired 0 rows. The inventory
    writes a date FIRST and a wildcard INSIDE the name; the agent's
    SRC_FILE_NAME drops both. Synthetic names, the same shape."""
    golden_files = ["YYYYMMDD_2041_7310_NW_*_P_I_1_T.txt", "YYYYMMDD_2041_7311_NW_*_P_I_1_T.txt"]
    agent_files = ["2041_7310_NW_P_I_1_T*", "2041_7311_NW_P_I_1_T*"]
    golden = _book(tmp_path / "golden.xlsx", {V1: [V1_HEAD, *[
        _row(V1_HEAD, FILE_NAME=f, domain="Member", subdomain="TPL") for f in golden_files]]})
    run = tmp_path / "run"
    book = _book(run / "nw_tpl_dtl" / "nw_tpl_dtl_IIG.xlsx", {
        V2: [V2_HEAD, _row(V2_HEAD, DOMAIN="Member", SUBDOMAIN="TPL")],
        ADLS: [ADLS_HEAD, *[[None, j + 1, f"obj_{j}", "Member", "TPL", f, "stg_nw_tpl_dtl"]
                            for j, f in enumerate(agent_files)]]})
    (run / "CODEGEN_BUILD.json").write_text(json.dumps(
        {"sha": "abc1234", "line": "codegen 0.0.test · build abc1234 on b"}), encoding="utf-8")
    out = tmp_path / "out"
    assert iig_diff.main(["--golden", str(golden), "--agent", str(book),
                          "--out", str(out)]) == 0
    printed = capsys.readouterr().out
    s = json.loads((out / "score.json").read_text(encoding="utf-8"))["sheets"][LABEL]
    assert s["pairing"]["identity"] == 2 and s["comparable_score"] == 100.0
    samples = s["file_samples"]
    assert [g["name"] for g in samples["golden"]] == golden_files
    assert [a["name"] for a in samples["agent"]] == agent_files
    assert [g["identity"] for g in samples["golden"]] == [a["identity"]
                                                          for a in samples["agent"]]
    block = printed.split(f"FILE NAMES (D1 only - unscrubbed) {LABEL}")[1]
    assert golden_files[0] in block and agent_files[0] in block
    assert "golden 2 row(s), agent 1 row(s)" in block


def test_rows_sharing_feed_flow_table_pair_by_file_despite_a_lob_suffix(tmp_path: Path,
                                                                         capsys) -> None:
    """D1 pair_12 (2026-10-14): 18 CAQH ADLS rows share feed x flow x target
    table — the identity could not tell them apart; the golden names each file
    ``*<n>_<n>_X*.txt`` (the LOB suffix collapsed into the wildcard), the agent
    ``<n>_<n>_X_P_I_1_T*``. Identity @file (the normalized SRC_FILE_NAME) splits
    the 18, the suffix normalizes away: all 18 pair by identity, in any order,
    and ALL_FILES -> FILE_ADLS (16 golden files, one agent feed row) pairs too.
    Synthetic names, the same shape."""
    numbers = [7300 + n for n in range(18)]
    lob = {n: f"{(i * 37) % 9000 + 100:04d}" for i, n in enumerate(numbers)}
    head = [*ADLS_HEAD, "LOB"]
    table = "stg_nw_tpl_dtl"
    golden_rows = [[None, i + 1, f"*2041_{n}_NW", "Member", "TPL", f"*2041_{n}_NW*.txt",
                    table, lob[n]] for i, n in enumerate(numbers)]
    agent_rows = [[None, i + 1, f"2041_{n}_NW_P_I_1_T", "Member", "TPL",
                   f"2041_{n}_NW_P_I_1_T*", table, lob[n]]
                  for i, n in enumerate(reversed(numbers))]
    golden = _book(tmp_path / "golden.xlsx", {
        ADLS: [head, *golden_rows],
        V1: [V1_HEAD, *[_row(V1_HEAD, FILE_NAME=f"*2041_{n}_NW*.txt", domain="Member",
                             subdomain="TPL") for n in numbers[:16]]]})
    run = tmp_path / "run"
    book = _book(run / "nw_tpl_dtl" / "nw_tpl_dtl_IIG.xlsx", {
        V2: [V2_HEAD, _row(V2_HEAD, DOMAIN="Member", SUBDOMAIN="TPL")],
        ADLS: [head, *agent_rows]})
    (run / "CODEGEN_BUILD.json").write_text(json.dumps(
        {"sha": "abc1234", "line": "codegen 0.0.test · build abc1234 on b"}), encoding="utf-8")
    out = tmp_path / "out"
    assert iig_diff.main(["--golden", str(golden), "--agent", str(book),
                          "--out", str(out)]) == 0
    capsys.readouterr()
    sheets = json.loads((out / "score.json").read_text(encoding="utf-8"))["sheets"]
    adls = sheets[ADLS]
    assert adls["pairing"]["identity"] == 18 and adls["key"][-1] == "@file"
    # each row pairs with its own file: the LOB beside it agrees on all 18
    assert adls["column_counts"]["LOB"]["exact"] == 18
    assert adls["column_counts"]["TGT_TABLE_NAME"]["exact"] == 18
    assert sheets[LABEL]["pairing"]["identity"] == 16


MFT_ADLS_HEAD = [*ADLS_HEAD, "SRC_CONTAINER_NAME", "SRC_ADLS_PATH"]


def _mft_agent(run: Path) -> list[Path]:
    """The agent of an MFT-landed feed (framework owner's rule 3): FILE_ADLS
    written with its headers only; the ADLS rows land under mftlanding."""
    books = []
    for feed, (domain, sub, _golden_files, files) in FEEDS.items():
        books.append(_book(run / feed / f"{feed}_IIG.xlsx", {
            V2: [V2_HEAD],
            ADLS: [MFT_ADLS_HEAD, *[[None, j + 1, f"obj_{j}", domain, sub, f, f"stg_{feed}",
                                     "mftlanding", f"/mftlanding/inbound/{feed}/"]
                                    for j, f in enumerate(files)]],
        }))
    (run / "CODEGEN_BUILD.json").write_text(json.dumps(
        {"sha": "abc1234", "line": "codegen 0.0.test · build abc1234 on b"}), encoding="utf-8")
    return books


def test_an_mft_feed_pairs_all_files_with_adls_delta_and_file_adls_is_empty_by_rule(
        tmp_path: Path, capsys) -> None:
    """Rule 3: the golden's ALL_FILES file rows pair with the agent's ADLS_DELTA
    rows on the file identity; FILE_ADLS (the golden's rows too) is reported
    "empty by rule (MFT)" — never missing, no score."""
    golden = _golden(tmp_path / "golden.xlsx", others=False)
    wb = openpyxl.load_workbook(golden)                    # a v2-style FILE_ADLS row too
    ws = wb.create_sheet(V2)
    ws.append(V2_HEAD)
    ws.append(_row(V2_HEAD, DOMAIN="Member", SUBDOMAIN="Risk", SOURCE_TYPE="File"))
    wb.save(golden)
    books = _mft_agent(tmp_path / "run")
    keys = iig_diff.load_keys(KEYS)
    agent = iig_diff.read_agent_workbooks(books, keys)
    assert iig_diff.agent_mft_marker(agent, keys) == "mftlanding"
    out = tmp_path / "out"
    assert iig_diff.main(["--golden", str(golden), "--agent", *map(str, books),
                          "--out", str(out)]) == 0
    capsys.readouterr()
    sheets = json.loads((out / "score.json").read_text(encoding="utf-8"))["sheets"]
    s = sheets[f"{V1} → {ADLS}"]
    assert s["pairing"] == {"identity": 3, "position": 0, "key": 0}
    assert s["unpaired"] == {"golden": 0, "agent": 0}
    assert s["counts"]["exact"] == 3 * 2 and s["counts"]["missing-in-agent"] == 0
    assert "MFT feed (landing 'mftlanding')" in s["template_mismatch"]["note"]
    empty = sheets[V2]
    assert empty["presence"] == "empty-by-rule" and empty["note"] == "empty by rule (MFT)"
    assert empty["score"] is None and empty["counts"]["missing-in-agent"] == 0
    assert empty["rows"]["golden"] == 1
