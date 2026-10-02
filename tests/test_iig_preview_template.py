"""IIG-first M2: the metadata-sheet preview renders in the IIG template
version the served run actually used, not the config default.

Before the fix both preview routes called ``metadata_sheet_payload`` without
a template, so an ``iig_v2`` run was previewed in the ``iig_v1`` layout.
"""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402
from ui.backend import main as ui_main  # noqa: E402


def _client(monkeypatch, config, spec, iig_template):
    store = SimpleNamespace(
        runs={spec.feed_slug: SimpleNamespace(spec=spec, outcomes=[])},
        config=config, label="demo_preview", iig_template=iig_template)
    monkeypatch.setattr(ui_main, "_require_store", lambda: store)
    return TestClient(ui_main.app)


def test_preview_uses_the_runs_iig_v2_template(monkeypatch, pair1_config, pair1_spec):
    client = _client(monkeypatch, pair1_config, pair1_spec, "iig_v2")
    payload = client.get("/api/demo/metadata-sheet").json()
    expected = list(pair1_config.metadata.templates["iig_v2"].tabs)
    assert payload["template"] == "iig_v2"
    assert list(payload["tabs"]) == expected
    workbook = load_workbook(io.BytesIO(client.get("/api/demo/metadata-sheet.xlsx").content))
    assert [s for s in workbook.sheetnames if not s.startswith("_")] == expected


def test_preview_without_a_recorded_template_is_the_default(monkeypatch, pair1_config,
                                                            pair1_spec):
    client = _client(monkeypatch, pair1_config, pair1_spec, None)
    payload = client.get("/api/demo/metadata-sheet").json()
    assert pair1_config.metadata.template == "iig_v1"
    assert "template" not in payload                    # the iig_v1 payload is unchanged
    assert list(payload["tabs"]) == list(pair1_config.demo.metadata_sheet.tabs)


def test_the_store_records_and_forgets_the_runs_template(tmp_path):
    from ui.backend.service import REPO_ROOT, GenerationStore

    store = GenerationStore(str(REPO_ROOT / "config" / "config.yaml"))
    assert store.iig_template is None
    store.adopt({}, [], mode="live", label="demo_x", out_root=tmp_path, reports_root=tmp_path,
                iig_template="iig_v2")
    assert store.iig_template == "iig_v2"
    assert store.forget_runs(["demo_x"]) is True
    assert store.iig_template is None
