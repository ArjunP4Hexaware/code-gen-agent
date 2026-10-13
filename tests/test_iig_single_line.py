"""No IIG cell may hold a line break (2026-10-13, pair_12 CAQH on D1).

The CAQH FRD states BOTH labels in one cell on two lines — "Domain: Member\\n
Sub Domain: Third Party Liability" — and the reader carried the whole cell
into SUBDOMAIN: 22 cells, the derivations check FAILed the feed. Now (1) the
FRD reader parses a cell that carries a ``Domain:`` AND a ``Sub Domain:``
label into the two values; (2) every IIG payload cell is single-line before
any workbook is written (normalised at the emitter boundary, each change a
flag naming the cell); (3) a check over the WRITTEN workbooks fails with the
cell address if a line break still slips through.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from codegen.extract.frd_docx import split_domain_labels
from codegen.gate.derivations import check_iig_workbooks_single_line, normalise_iig_line_breaks

CAQH_CELL = "Domain: Member\nSub Domain: Third Party Liability"


def test_the_exact_two_line_cell_splits_into_domain_and_sub_domain():
    assert split_domain_labels(CAQH_CELL) == ("Member", "Third Party Liability")


@pytest.mark.parametrize("text", [
    "domain: Member; sub domain: Third Party Liability",
    "DOMAIN: Member\r\nSUB-DOMAIN: Third Party Liability",
    "Domain = Member | Subdomain = Third Party Liability",
    "Domain: Member, Sub_Domain: Third Party Liability",
    "Sub Domain: Third Party Liability\nDomain: Member",
])
def test_any_case_any_separator(text):
    assert split_domain_labels(text) == ("Member", "Third Party Liability")


@pytest.mark.parametrize("text", [
    "", None, "Member / Third Party Liability", "Domain: Member",
    "Sub Domain: Third Party Liability", "Member Domain and more",
])
def test_a_cell_without_both_labels_is_not_split(text):
    assert split_domain_labels(text) is None


def _frd_with_domain_cell(tmp_path, cell: str):
    """The synthetic family-E FRD with its "Domain and Subdomain" row holding
    ``cell`` (the CAQH label shape)."""
    from acfc_shapes import family_e_multi as fm

    frd = tmp_path / "FRD.docx"
    frd.write_bytes(fm.build_frd(domain_cell=cell))
    return frd


def test_the_frd_reader_reads_both_labels(tmp_path, config):
    from codegen.extract.frd_docx import extract_frd_contract

    frd = _frd_with_domain_cell(tmp_path, CAQH_CELL)
    contract, _ = extract_frd_contract(frd, config, generated_date="2026-10-13")
    (feed,) = contract.feeds
    assert (feed.domain, feed.sub_domain) == ("Member", "Third Party Liability")
    assert any(f.startswith("frd_domain_labels_split:feeds[0]") and "'Third Party Liability'" in f
               for f in contract.extraction_flags), contract.extraction_flags


def test_a_plain_domain_cell_reads_as_before(tmp_path, config):
    from codegen.extract.frd_docx import extract_frd_contract

    frd = _frd_with_domain_cell(tmp_path, "Member / Care Management")
    contract, _ = extract_frd_contract(frd, config, generated_date="2026-10-13")
    (feed,) = contract.feeds
    assert (feed.domain, feed.sub_domain) == ("Member", "Care Management")
    assert not any(f.startswith("frd_domain_labels_split") for f in contract.extraction_flags)


def _payload(value: str) -> dict:
    return {"tabs": {"ADLS_DELTA_INGESTION_DETAILS": {
        "headers": ["GROUP_ID", "DOMAIN", "SUBDOMAIN"],
        "rows": [{"values": {"GROUP_ID": "", "DOMAIN": "Member", "SUBDOMAIN": value},
                  "badges": {h: {"badge": "from_frd"} for h in ("GROUP_ID", "DOMAIN",
                                                                "SUBDOMAIN")}}]}}}


def test_the_emitter_boundary_makes_every_cell_single_line():
    payload = _payload(CAQH_CELL)
    flags = normalise_iig_line_breaks(payload)
    row = payload["tabs"]["ADLS_DELTA_INGESTION_DETAILS"]["rows"][0]
    assert row["values"]["SUBDOMAIN"] == "Domain: Member Sub Domain: Third Party Liability"
    assert flags == [
        "iig_line_break_normalised:ADLS_DELTA_INGESTION_DETAILS!C2 (SUBDOMAIN) — "
        "'Domain: Member\\nSub Domain: Third Party Liability' written as one line"]
    assert normalise_iig_line_breaks(payload) == []          # idempotent


def test_a_line_break_in_a_written_workbook_fails_with_its_address(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ADLS_DELTA_INGESTION_DETAILS"
    sheet.append(["GROUP_ID", "DOMAIN", "SUBDOMAIN"])
    sheet.append(["", "Member", CAQH_CELL])
    path = tmp_path / "x_IIG.xlsx"
    workbook.save(path)
    check = check_iig_workbooks_single_line([path])
    assert not check.passed
    assert check.name == "iig_single_line"
    assert "x_IIG.xlsx ADLS_DELTA_INGESTION_DETAILS!C2 (SUBDOMAIN)" in check.details
    sheet["C2"] = "Third Party Liability"
    workbook.save(path)
    assert check_iig_workbooks_single_line([path]).passed


def test_caqh_shape_end_to_end_writes_single_line_iig(tmp_path, capsys, monkeypatch):
    """layout -> extract-sttm -> generate (mock) with the CAQH domain cell:
    DOMAIN / SUBDOMAIN carry the two values, the gate passes the derivations
    and the written-workbook check."""
    from openpyxl import load_workbook

    from acfc_shapes import family_e_multi as fm
    from codegen import cli

    monkeypatch.setenv("CODEGEN_NOTIFICATION_EMAILS", "syn.dl@synthetic.example")
    # acfc_prx needs its catalogs: the ACFC overlay the suite otherwise skips.
    overlay = Path(__file__).resolve().parents[1] / "config" / "overlays" / "acfc_env.yaml"
    monkeypatch.setenv("CODEGEN_CONFIG_OVERLAYS", str(overlay))
    for name in ("state", "out"):
        (tmp_path / name).mkdir()
    monkeypatch.setenv("CODEGEN_STORAGE_STATE", f"local:{(tmp_path / 'state').as_posix()}")
    monkeypatch.setenv("CODEGEN_STORAGE_OUTPUTS", f"local:{(tmp_path / 'out').as_posix()}")
    sttm = tmp_path / "STTM.xlsx"
    fm.build_sttm().save(sttm)
    frd = _frd_with_domain_cell(tmp_path, CAQH_CELL)
    frd_out, sttm_out = tmp_path / "frd.json", tmp_path / "sttm.json"
    assert cli.main(["layout", "--workbook", str(sttm), "--frd", str(frd), "--no-cache",
                     "--frd-contract-out", str(frd_out)]) == 0
    assert cli.main(["extract-sttm", "--workbook", str(sttm), "--frd-contract", str(frd_out),
                     "--out", str(sttm_out), "--generated-date", "2026-10-13"]) == 0
    capsys.readouterr()
    assert cli.main(["generate", "--frd-contract", str(frd_out), "--sttm-contract",
                     str(sttm_out), "--feed", "nb_cm_north", "--skip-tests",
                     "--output-mode", "framework", "--profile", "acfc_prx",
                     "--iig-template", "iig_v2"]) == 0, capsys.readouterr().out
    (clean,) = (tmp_path / "out").rglob("nb_cm_north_IIG.xlsx")
    workbook = load_workbook(clean, read_only=True)
    seen = set()
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        headers = list(rows[0])
        for row in rows[1:]:
            for header, value in zip(headers, row, strict=False):
                assert not (isinstance(value, str) and ("\n" in value or "\r" in value))
                if header in ("DOMAIN", "SUBDOMAIN") and value:
                    seen.add((header, value))
    assert ("SUBDOMAIN", "Third Party Liability") in seen
    assert ("DOMAIN", "Member") in seen
