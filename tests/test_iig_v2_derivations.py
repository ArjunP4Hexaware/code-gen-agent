"""IIG-first P1 (2026-10-07): iig_v2 derivation fixes, scored against the
real ACFC rows (the real values themselves are not in the repo).

a. STGDELTA SRC_ / TGT_CATALOG_NAME = the DDL's catalog chain.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from codegen import cli
from codegen.config import load_config
from test_derivations import FRD, FRD_CATALOG, REPO, _pair11_specs, _scoped

ACFC_OVERLAY = REPO / "config" / "overlays" / "acfc_env.yaml"


def _write_overlay(tmp: Path, text: str) -> Path:
    path = tmp / "overlay.yaml"
    path.write_text(text, encoding="utf-8")
    return path


def _run(frd: Path, config, tmp: Path, index: int = 0):
    (tmp / "contracts").mkdir(parents=True, exist_ok=True)
    specs, _ = _pair11_specs(frd, tmp / "contracts", config)
    spec = specs[index]
    gate = cli._generate_feed(spec, _scoped(config, tmp), dry_run=True, skip_tests=True,
                              output_mode="framework", conventions_profile="acfc_prx",
                              iig_template="iig_v2")
    return spec, gate, tmp / "out" / spec.feed_slug / "framework" / "config_rows.xlsx"


def _rows(xlsx: Path, sheet: str) -> list[dict[str, tuple]]:
    """[{header: (value, badge, tooltip)}] for one sheet of config_rows.xlsx."""
    workbook = load_workbook(xlsx)
    ws = workbook[sheet]
    headers = [c.value for c in ws[1]]
    prov: dict[tuple[int, str], tuple] = {}
    for tab, row, header, badge, tooltip in workbook["_provenance"].iter_rows(
            min_row=2, values_only=True):
        if tab == sheet:
            prov[(row, header)] = (badge, tooltip)
    out = []
    for index, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        out.append({h: (v, *prov.get((index, h), (None, None)))
                    for h, v in zip(headers, values, strict=False)})
    return out


# -- a. catalogs ----------------------------------------------------------------- #


def test_stg_std_catalogs_follow_the_ddl_chain_default_catalog(tmp_path):
    config = load_config(REPO / "config" / "config.yaml", overlays=[ACFC_OVERLAY])
    _spec, gate, xlsx = _run(FRD, config, tmp_path)            # FRD + STTM state no catalog
    assert gate.verdict != "FAIL"
    (row,) = _rows(xlsx, "STGDELTA_STDDELTA_INGESTION_DET")
    assert row["SRC_CATALOG_NAME"][:2] == ("d1_dlk", "SYNTHETIC")
    assert row["TGT_CATALOG_NAME"][:2] == ("d1_std", "SYNTHETIC")
    assert "config_default" in row["SRC_CATALOG_NAME"][2]
    assert "default_catalog[standard]" in row["TGT_CATALOG_NAME"][2]


def test_stg_std_catalogs_stated_by_the_frd_win(tmp_path):
    config = load_config(REPO / "config" / "config.yaml", overlays=[ACFC_OVERLAY])
    spec, _gate, xlsx = _run(FRD_CATALOG, config, tmp_path)
    (row,) = _rows(xlsx, "STGDELTA_STDDELTA_INGESTION_DET")
    assert row["SRC_CATALOG_NAME"][0] == spec.detail_segment.stage_table.catalog != "d1_dlk"
    assert row["TGT_CATALOG_NAME"][0] == spec.standard_table.catalog
    assert row["SRC_CATALOG_NAME"][1] == "from FRD"


def test_unresolved_catalog_leaves_the_cell_to_a_constant(tmp_path):
    overlay = _write_overlay(tmp_path, (
        "metadata: {templates: {iig_v2: {constants: {STGDELTA_STDDELTA_INGESTION_DET: "
        "{SRC_CATALOG_NAME: const_stg}}}}}\n"))
    config = load_config(REPO / "config" / "config.yaml", overlays=[overlay])
    _spec, _gate, xlsx = _run(FRD, config, tmp_path)
    (row,) = _rows(xlsx, "STGDELTA_STDDELTA_INGESTION_DET")
    assert row["SRC_CATALOG_NAME"][0] == "const_stg"           # was always blank before
    assert row["TGT_CATALOG_NAME"][0] in (None, "")             # still open, never guessed


# -- b. TGT_LOAD_OPTION ------------------------------------------------------------ #


def test_load_option_speaks_the_framework_vocabulary():
    from codegen.metadata_template import _load_option_cell

    cell = _load_option_cell("Truncate and Load", "STG")
    assert cell["value"] == "Overwrite"
    assert cell["badge_entry"]["badge"] == "from_frd"
    assert "'Truncate and Load'" in cell["badge_entry"]["tooltip"]      # the FRD's own text
    assert "Load Strategy STG" in cell["badge_entry"]["tooltip"]
    same = _load_option_cell("Append", "STD")                             # already framework
    assert same["value"] == "Append"
    assert same["badge_entry"]["tooltip"] == "FRD Structural Metadata → Load Strategy STD"


def test_load_option_both_sheets_on_a_truncate_and_load_feed(tmp_path):
    config = load_config(REPO / "config" / "config.yaml", overlays=[ACFC_OVERLAY])
    spec, _gate, xlsx = _run(FRD, config, tmp_path)
    expected = {"ADLS_DELTA_INGESTION_DETAILS": spec.stage_load_strategy,
                "STGDELTA_STDDELTA_INGESTION_DET": spec.standard_load_strategy}
    for sheet, strategy in expected.items():
        for row in _rows(xlsx, sheet):
            want = "Overwrite" if strategy == "Truncate and Load" else strategy
            assert row["TGT_LOAD_OPTION"][0] == want


# -- c. FREQUENCY / PIPELINE_FREQUENCY ------------------------------------------- #


def test_frequency_token_reading():
    from codegen.metadata_template import _frequency_token

    assert _frequency_token("Monthly Run; File will be received yearly twice.") == "Monthly"
    assert _frequency_token("daily") == "Daily"
    assert _frequency_token("Files arrive weekly on Monday") == "Weekly"
    assert _frequency_token("Received annually") == "Yearly"
    assert _frequency_token("twice a month") is None                   # no token
    assert _frequency_token("Files arrive daily and monthly") is None  # several, none leads


def _feed(frequency):
    from types import SimpleNamespace

    return SimpleNamespace(frequency=frequency)


def test_frequency_cell_normalises_and_keeps_the_text():
    from codegen.metadata_template import _frequency

    cell = _frequency(_feed("Monthly Run; File reception: 11-15th of every month"), None)
    assert cell["value"] == "Monthly"
    assert cell["badge_entry"]["badge"] == "from_frd"
    assert "'Monthly Run; File reception: 11-15th of every month'" in \
        cell["badge_entry"]["tooltip"]
    plain = _frequency(_feed("Daily"), None)                     # already a token: unchanged
    assert plain["value"] == "Daily" and "tooltip" not in plain["badge_entry"]
    open_cell = _frequency(_feed("twice a month"), None)
    assert open_cell["value"] == "" and open_cell["badge_entry"]["badge"] == "needs_template"
    assert "'twice a month'" in open_cell["badge_entry"]["tooltip"]


# -- d. path placeholders ------------------------------------------------------- #

ACFC_PATHS = """
metadata:
  templates:
    iig_v2:
      path_patterns:
        ADLS_DELTA_INGESTION_DETAILS:
          SRC_ADLS_PATH: "{landing_rel}"
          SRC_ADLS_ARCHVL_PATH: "{landing_rel}Archive/"
          TGT_ADLS_PATH: "/{domain_path}Processed/{stage_table}"
          TGT_RJT_ADLS_PATH: "/{domain_path}Reject/{reject_table}"
"""


def test_landing_rel_and_domain_path():
    from types import SimpleNamespace

    from codegen.metadata_template import _domain_path, _landing_rel

    tpl = SimpleNamespace(constants={"ADLS_DELTA_INGESTION_DETAILS": {
        "SRC_CONTAINER_NAME": "mftlanding"}})
    rel = _landing_rel(tpl, "STGDELTA_STDDELTA_INGESTION_DET",
                       "/mftlanding/inbound/sdh/public/vendor_x/")
    assert rel == "/inbound/sdh/public/vendor_x/"
    assert _domain_path(rel) == "sdh/public/vendor_x/"
    assert _landing_rel(tpl, "X", "/other/inbound/a/") == "/other/inbound/a/"   # not the container
    assert _domain_path("/a/b/") == "a/b/"                                       # no inbound/
    no_container = SimpleNamespace(constants={})
    assert _landing_rel(no_container, "X", "/mftlanding/a/") == "/mftlanding/a/"


def test_overlay_path_shapes_render_the_real_target_shape(tmp_path):
    overlay = _write_overlay(tmp_path, ACFC_PATHS)
    config = load_config(REPO / "config" / "config.yaml", overlays=[ACFC_OVERLAY, overlay])
    spec, _gate, xlsx = _run(FRD, config, tmp_path)
    (row, *_more) = _rows(xlsx, "ADLS_DELTA_INGESTION_DETAILS")
    from codegen.metadata_template import _domain_path, _landing, _landing_rel

    rel = _landing_rel(config.metadata.templates["iig_v2"], "ADLS_DELTA_INGESTION_DETAILS",
                       _landing(_frd_feed(spec)))
    stage = spec.detail_segment.stage_table.table
    assert row["SRC_ADLS_PATH"][0] == rel                             # the shape wins
    assert row["SRC_ADLS_ARCHVL_PATH"][0] == f"{rel}Archive/"
    assert row["TGT_ADLS_PATH"][0] == f"/{_domain_path(rel)}Processed/{stage}"
    assert row["TGT_RJT_ADLS_PATH"][0] == f"/{_domain_path(rel)}Reject/{stage}_reject"


def test_shipped_iig_v2_path_shapes_unchanged():
    shapes = load_config(REPO / "config" / "config.yaml").metadata.templates["iig_v2"] \
        .path_patterns["ADLS_DELTA_INGESTION_DETAILS"]
    assert shapes == {"SRC_ADLS_ARCHVL_PATH": "{landing}Archive/",
                      "TGT_ADLS_PATH": "{landing}Processed/{stage_table}",
                      "TGT_RJT_ADLS_PATH": "{landing}Processed/{reject_table}"}


def _frd_feed(spec):
    from codegen.metadata_sheet import _feed_from_spec

    return _feed_from_spec(spec)


# -- e. src_columns_style: named ------------------------------------------------- #


def _field(source, stage):
    from types import SimpleNamespace

    return SimpleNamespace(source_column=source, stage_column=stage)


def test_named_sources_rule():
    from codegen.metadata_template import _named_sources

    assert _named_sources([_field("zip_code", "zip_code"), _field("pop", "pop")])
    assert not _named_sources([_field("col1", "a"), _field("pop", "pop")])
    assert not _named_sources([_field("COL_2", "a")])
    assert not _named_sources([_field(None, "a")])
    assert not _named_sources([])


def test_named_style_is_overlay_selectable(tmp_path):
    overlay = _write_overlay(tmp_path, "metadata: {templates: {iig_v2: "
                                       "{src_columns_style: named}}}\n")
    config = load_config(REPO / "config" / "config.yaml", overlays=[ACFC_OVERLAY, overlay])
    spec, _gate, xlsx = _run(FRD, config, tmp_path)
    from codegen.metadata_template import _distinct_fields, _named_sources

    fields = _distinct_fields(spec)
    (row, *_more) = _rows(xlsx, "ADLS_DELTA_INGESTION_DETAILS")
    if _named_sources(fields):
        expected = ",".join(f"{f.source_column}:{f.source_column}" for f in fields)
    else:
        expected = ",".join(f"col{i}:{f.stage_column}" for i, f in enumerate(fields, start=1))
    assert row["SRC_COLUMNS"][0] == expected
    shipped = load_config(REPO / "config" / "config.yaml").metadata.templates["iig_v2"]
    assert shipped.src_columns_style == "positional"                   # the shipped default


# -- f. HEADER_FLAG / FILE_HEADER_FLAG / FILE_FOOTER_FLAG ------------------------ #


def _faq(header=None, trailer=None):
    from types import SimpleNamespace

    from codegen.faq import FaqAnswer

    def answer(value):
        return (FaqAnswer(value=value, source="engineer", evidence="FRD says so")
                if value is not None else FaqAnswer(value="unknown"))

    return SimpleNamespace(has_header=answer(header), has_trailer=answer(trailer))


def _spec(delimiter):
    from types import SimpleNamespace

    return SimpleNamespace(delimiter=delimiter)


def test_delimited_flags_only_from_stated_answers():
    from codegen.metadata_template import _header_flag_cells

    cells = _header_flag_cells(_spec(","), _faq(header="yes", trailer="no"))
    assert {k: v["value"] for k, v in cells.items()} == {
        "HEADER_FLAG": "Y", "FILE_HEADER_FLAG": "Y", "FILE_FOOTER_FLAG": "N"}
    assert all(v["badge_entry"]["badge"] == "from_faq" for v in cells.values())
    assert "has_header 'yes' → Y" in cells["HEADER_FLAG"]["badge_entry"]["tooltip"]
    only_header = _header_flag_cells(_spec("|"), _faq(header="no"))
    assert {k: v["value"] for k, v in only_header.items()} == {
        "HEADER_FLAG": "N", "FILE_HEADER_FLAG": "N"}                   # trailer stays open
    assert _header_flag_cells(_spec(","), _faq()) == {}               # nothing stated: open
    assert _header_flag_cells(_spec(","), None) == {}


def test_fixed_width_keeps_the_pre_existing_reading():
    """No delimiter: the golden's flags do not follow header / trailer
    SEGMENTS (open framework question) — HEADER_FLAG as the FAQ states it,
    the FILE_* flags open."""
    from codegen.metadata_template import _header_flag_cells

    cells = _header_flag_cells(_spec(None), _faq(header="no", trailer="yes"))
    assert list(cells) == ["HEADER_FLAG"]
    assert cells["HEADER_FLAG"]["value"] == "no"
