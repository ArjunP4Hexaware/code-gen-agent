"""post74 (D1 post73 item 6): pair_11 asked 15 notebook-only keys
(``feeds.feed_<n>.notebook.landing_path / archive_path / reject_path``) while
pair_02 asked none. The cause: the notebook (v2 and v1's feed_spec) takes its
landing / archive / reject paths from the IIG ADLS_DELTA row's SRC_ADLS_PATH /
SRC_ADLS_ARCHVL_PATH / TGT_RJT_ADLS_PATH, and those derive from the FRD's
landing location — pair_02's FRD states it, the Pyx FRD (pair_11) states no
inbound path, so every path cell is open.

Now an open landing is asked ONCE per feed, by the IIG cell's key
``feeds[<i>].src_adls_path`` — the same key for the workbook (INPUT_NEEDED) and
the notebook (blocking on a notebook run) — with the template's landing
convention shown as the agent's proposal, never written unanswered. The
answer fills the workbook's three path cells and the notebook's paths."""

from __future__ import annotations

import json

import yaml
from openpyxl import load_workbook

from codegen import cli, run_stages
from codegen.metadata_template import landing_proposal
from test_iig_corpus_cells import _pair11

ADLS = "ADLS_DELTA_INGESTION_DETAILS"


def _without_landing(pair, index: int = 0):
    frd_path = pair[0]
    data = json.loads(frd_path.read_text(encoding="utf-8"))
    data["feeds"][index]["landing_location"] = None
    data.get("field_provenance", {}).pop(f"feeds[{index}].landing_location", None)
    frd_path.write_text(json.dumps(data), encoding="utf-8")


def _run(config, pair, parts, **kw) -> int:
    return cli._run_pairs([pair], config, dry_run=True, skip_tests=True, output_mode=parts,
                          conventions_profile="acfc_prx", iig_template="iig_v2", **kw)


def _adls_paths(out, feed_slug) -> dict:
    (xlsx,) = (out / feed_slug).rglob("*_IIG.xlsx")
    ws = load_workbook(xlsx, read_only=True)[ADLS]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    return {h: rows[1][header.index(h)] for h in
            ("SRC_ADLS_PATH", "SRC_ADLS_ARCHVL_PATH", "TGT_RJT_ADLS_PATH")}


def _lines(out: str, label: str) -> list[str]:
    return [ln for ln in out.splitlines() if ln.startswith(label)]


def test_a_framework_run_asks_the_iig_cell_once_with_a_proposal(tmp_path, capsys,
                                                                landing_autoset_off):
    config, specs, pair = _pair11(tmp_path)
    _without_landing(pair)
    feed = specs[0]
    code = _run(config, pair, ["framework"], only_feed=feed.feed_id)
    out = capsys.readouterr().out
    assert code == cli.EXIT_OK                      # non-blocking on a framework run
    (line,) = [ln for ln in _lines(out, "INPUT_NEEDED") if "src_adls_path" in ln]
    assert line.startswith("INPUT_NEEDED   feeds[0].src_adls_path — no document states the "
                           "inbound path")
    proposal = landing_proposal(config.metadata.templates["iig_v2"], feed)
    assert proposal and f"proposed (agent_proposal, not applied): {proposal!r}" in line
    assert all(v in (None, "") for v in _adls_paths(tmp_path / "out", feed.feed_slug).values())
    assert not [ln for ln in out.splitlines() if ".notebook." in ln.split(" — ")[0]]


def test_a_notebook_run_asks_the_same_key_blocking(tmp_path, capsys, landing_autoset_off):
    config, specs, pair = _pair11(tmp_path)
    _without_landing(pair)
    feed = specs[0]
    code = _run(config, pair, ["notebook", "framework"], only_feed=feed.feed_id)
    out = capsys.readouterr().out
    assert code == cli.EXIT_NEEDS_ANSWERS
    keys = [ln[15:].split(" — ", 1)[0] for ln in _lines(out, "QUESTION")]
    assert keys == ["feeds[0].src_adls_path"]       # one key, never notebook.<path>
    manifest = run_stages.read_stage_manifest(tmp_path / "out")
    assert [q["key"] for q in manifest["questions"] if q["blocking"]] == [
        "feeds[0].src_adls_path"]


def test_the_answer_fills_the_workbook_and_the_notebook(tmp_path, capsys):
    config, specs, pair = _pair11(tmp_path)
    _without_landing(pair)
    feed = specs[0]
    answers = tmp_path / "answers.yaml"
    answers.write_text(yaml.safe_dump({"inputs": {
        "feeds[0].src_adls_path": {"value": "/inbound/claims/vendor_k/"}}}), encoding="utf-8")
    code = _run(config, pair, ["notebook", "framework"], only_feed=feed.feed_id,
                answers_path=answers)
    out = capsys.readouterr().out
    assert "src_adls_path" not in out.replace("feeds[0].src_adls_path matches", "")
    paths = _adls_paths(tmp_path / "out", feed.feed_slug)
    assert paths["SRC_ADLS_PATH"] == "/inbound/claims/vendor_k/"
    assert paths["SRC_ADLS_ARCHVL_PATH"] == "/inbound/claims/vendor_k/Archive/"
    assert str(paths["TGT_RJT_ADLS_PATH"]).startswith("/claims/vendor_k/Reject/")
    assert code in (cli.EXIT_OK, cli.EXIT_NEEDS_ANSWERS)
    assert not [ln for ln in _lines(out, "QUESTION") if "adls_path" in ln]


def test_no_proposal_from_an_unstated_source(tmp_path):
    config, specs, _pair = _pair11(tmp_path)
    tpl = config.metadata.templates["iig_v2"]
    spec = specs[0].model_copy(update={"source_system": "unstated", "landing_location": None})
    assert landing_proposal(tpl, spec) == ""


# --------------------------------------------- post84: landing_autoset (written)


def test_autoset_writes_the_convention_unasked_and_the_notebook_runs(tmp_path, capsys):
    """post84: acfc_env.yaml's landing_autoset — no document states the landing,
    so the convention is WRITTEN (the workbook's three path cells, the notebook's
    paths), one NOTE line, no question, and the notebook is not withheld."""
    config, specs, pair = _pair11(tmp_path)
    tpl = config.metadata.templates["iig_v2"]
    assert tpl.landing_autoset is True
    _without_landing(pair)
    feed = specs[0]
    proposal = landing_proposal(tpl, feed.model_copy(update={"landing_location": None}))
    code = _run(config, pair, ["notebook", "framework"], only_feed=feed.feed_id)
    out = capsys.readouterr().out
    assert not [ln for ln in out.splitlines()
                if ln.startswith(("QUESTION", "INPUT_NEEDED")) and "src_adls_path" in ln]
    (note,) = [ln for ln in _lines(out, "NOTE") if "feeds[0].src_adls_path" in ln]
    assert f"'{proposal}'" in note and "profile default" in note
    paths = _adls_paths(tmp_path / "out", feed.feed_slug)
    assert paths["SRC_ADLS_PATH"] == proposal
    assert paths["SRC_ADLS_ARCHVL_PATH"] == f"{proposal}Archive/"
    manifest = run_stages.read_stage_manifest(tmp_path / "out")
    assert not [q for q in manifest["questions"] if "src_adls_path" in q["key"]]
    (entry,) = manifest["feeds"]
    assert entry["written"] is True
    assert list((tmp_path / "out" / feed.feed_slug).rglob("*.ipynb"))       # not withheld
    assert code in (cli.EXIT_OK, cli.EXIT_NEEDS_ANSWERS)
    report = (tmp_path / "reports" / f"{feed.feed_slug}.md").read_text(encoding="utf-8")
    assert "landing_from_convention:" in report


def test_a_stated_landing_is_never_replaced_by_the_convention(tmp_path):
    from codegen.metadata_template import landing_default, with_landing_default

    config, specs, _pair = _pair11(tmp_path)
    tpl = config.metadata.templates["iig_v2"]
    stated = specs[0].model_copy(update={"landing_location": "/inbound/syn/stated/"})
    assert landing_default(tpl, stated) == ""
    assert with_landing_default(tpl, stated) is stated
    unset = specs[0].model_copy(update={"landing_location": None})
    applied = with_landing_default(tpl, unset)
    assert applied.landing_location == landing_proposal(tpl, unset)
    assert landing_default(tpl, applied) == applied.landing_location   # recognised as ours
    off = tpl.model_copy(update={"landing_autoset": False})
    assert landing_default(off, unset) == ""


def test_the_app_lists_the_default_as_changeable_and_a_change_reaches_the_resolver(tmp_path):
    """post84: the App shows the written landing as a profile default (the
    "Change defaults" panel); a confirmed change is the next run's gaps: answer."""
    import pytest

    pytest.importorskip("fastapi")
    from types import SimpleNamespace

    from ui.backend.demo import DemoRunner

    from codegen.input_choices import Choice

    config, specs, (frd, _sttm) = _pair11(tmp_path)
    spec = specs[0].model_copy(update={"landing_location": None})
    runner = SimpleNamespace(effective_output_parts=lambda: ["notebook", "framework"],
                             iig_template="iig_v2", conventions_profile="acfc_prx",
                             input_choices={})
    items = [i for i in DemoRunner._owed_inputs(runner, spec, config, frd)
             if i["key"].endswith(".src_adls_path")]
    (item,) = items
    assert item["kind"] == "profile_default" and item["applied"].startswith("/inbound/")
    runner.input_choices = {item["key"]: Choice("naming_standard", "/inbound/syn/changed/"),
                            "run.email_to": Choice("use_default")}
    assert DemoRunner._landing_choices(runner) == {item["key"]: "/inbound/syn/changed/"}
