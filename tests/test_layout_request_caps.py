"""P1-7 (red team LM-5 / PPL-3, 2026-10-30): a layout model request is
bounded. A 296 KB STTM produced a 12.2 MB request; the retry resent it.

- every header-region cell (and every partial-profile / label string) the model
  sees is clipped to ``layout.max_cell_chars``;
- cells go only for the sheets with an open role, at most
  ``layout.max_sheets_sent`` of them - every other sheet is its summary line;
- a request whose JSON is over ``layout.max_prompt_chars`` is refused BEFORE the
  provider is called (REASONS ``layout.request_too_large``, integers only) and so
  never retried;
- none of it moves the fingerprint - the layout-profile cache key.
"""

from __future__ import annotations

import json

from openpyxl import load_workbook

from codegen.layout.fingerprint import fingerprint, sheet_region
from codegen.layout.resolve import resolve_frd, resolve_workbook
from codegen.run_errors import REASONS
from layout_attack import (
    PAIR4_SHEET,
    STAGE_SCHEMA,
    StubProvider,
    frd_with_label,
    pair4_copy,
    with_layout,
)

LONG = "Z" * 32000


def _request_chars(request: dict) -> int:
    return len(json.dumps(request, indent=2))


def _empty(tmp_path):
    empty = tmp_path / "empty_cache"
    empty.mkdir(exist_ok=True)
    return [empty]


def _probe_workbook(tmp_path, sheets: int = 60, cols: int = 200, chars: int = 1000):
    """The red team's probe: pair 4 with one open role (stage schema), a 32 000
    character meta value above its header row and ``sheets`` extra sheets, each a
    header row of ``cols`` cells of ``chars`` characters over one data row."""
    path = pair4_copy(tmp_path / "probe.xlsx", headers={STAGE_SCHEMA: "Zone"},
                      cells={(1, 2): LONG})
    wb = load_workbook(path)
    for index in range(sheets):
        ws = wb.create_sheet(f"S{index}")
        ws.append([f"H{c}" + "Q" * chars for c in range(cols)])
        ws.append(["d"] * cols)
    wb.save(path)
    return path


def test_the_probe_workbook_request_is_bounded_and_the_cache_key_unchanged(config, tmp_path):
    path = _probe_workbook(tmp_path)
    before = fingerprint(load_workbook(path))
    stub = StubProvider()
    doc, workbook = resolve_workbook(path, config, provider=stub, cache_dirs=_empty(tmp_path))
    assert doc.fingerprint == before == fingerprint(workbook)     # the cache key never moves
    (request,) = stub.requests
    assert _request_chars(request) <= config.layout.max_prompt_chars
    regions = request["header_regions"]
    assert len(regions) == len(workbook.sheetnames)
    # only the sheet with the open role sends cells; the 60 others are summary lines
    sent = [r for r in regions if "cells withheld" not in r.splitlines()[0]]
    assert len(sent) == 1 and sent[0].startswith(f"sheet {PAIR4_SHEET!r}")
    assert "Q" * 100 not in json.dumps(request)
    # the long meta value is clipped in the model's copy, kept whole in the hash input
    assert LONG not in sent[0] and "[clipped, 32000 chars]" in sent[0]
    assert LONG in dict(sheet_region(workbook[PAIR4_SHEET]).cells).values()
    assert all(len(line) <= config.layout.max_cell_chars + 40
               for region in regions for line in region.splitlines()[2:])


def test_an_oversized_request_is_refused_before_the_call_and_never_retried(config, tmp_path):
    path = pair4_copy(tmp_path / "a.xlsx", headers={STAGE_SCHEMA: "Zone"})
    small = with_layout(config, max_prompt_chars=2000)
    stub = StubProvider()
    doc, _ = resolve_workbook(path, small, provider=stub, cache_dirs=_empty(tmp_path))
    assert stub.requests == [] and doc.provider_calls == 0
    (refusal,) = [r for r in doc.rejections if r.reason.startswith("layout.request_too_large")]
    assert "layout.request_too_large" in REASONS
    # integers and vocabulary only: no header text, no sheet name
    assert "Zone" not in refusal.render() and PAIR4_SHEET not in refusal.render()
    assert [(u.layer, u.role) for u in doc.profile.unresolved] == [("stage", "schema")]
    assert doc.questions                       # the open role is a question for a person


def test_sheets_over_max_sheets_sent_go_as_summary_lines(config, tmp_path):
    path = pair4_copy(tmp_path / "a.xlsx", headers={STAGE_SCHEMA: "Zone"})
    wb = load_workbook(path)
    second = wb.copy_worksheet(wb[PAIR4_SHEET])
    second.title = "STTM_B"
    wb.save(path)
    stub = StubProvider()
    doc, _ = resolve_workbook(path, with_layout(config, max_sheets_sent=1), provider=stub,
                              cache_dirs=_empty(tmp_path))
    (request,) = stub.requests
    heads = [r.splitlines()[0] for r in request["header_regions"]]
    assert sum("cells withheld" not in h for h in heads) == 1
    assert any("over layout.max_sheets_sent 1" in r.reason for r in doc.rejections)


def test_frd_label_and_profile_strings_are_clipped_and_the_cap_refuses(config, tmp_path):
    path = frd_with_label(tmp_path / "frd.docx", "Target Catalog and Schema",
                          "Zone " + "W" * 5000)
    stub = StubProvider()
    doc, _ = resolve_frd(path, config, provider=stub, cache_dirs=_empty(tmp_path))
    (request,) = stub.requests
    assert "W" * 600 not in json.dumps(request)
    assert any(label.endswith(" chars]") and "Zone W" in label for label in request["table_labels"])
    stub = StubProvider()
    doc, _ = resolve_frd(path, with_layout(config, max_prompt_chars=2000), provider=stub,
                         cache_dirs=_empty(tmp_path))
    assert stub.requests == [] and doc.provider_calls == 0
    assert any(r.reason.startswith("layout.request_too_large") for r in doc.rejections)
