"""M10.1 — the two v0.6.0 real-run bugs (docs/acfc/RUN_v060_pair1.md on the
unscrubbed ``acfc-runs`` branch; nothing from it is copied here).

1. A natural key that includes a Trailer-segment column crashed
   ``build_context`` (``KeyError``): the key was resolved against the Detail
   segment's fields only. Now every segment is searched; a column no segment
   carries is the gate check ``natural_key_columns`` = FAIL naming it.
2. The FRD's ADLS Location cell read ``/Path : <storage>/<landing>/…`` — a
   label token before the path — and the derivations check FAILed every
   path that embedded it. The label is stripped (``extractor.frd.
   value_label_prefixes``) and recorded in the field's evidence.
"""

from __future__ import annotations

import time
from pathlib import Path

import pytest

from acfc_shapes import frd as frd_fixtures
from acfc_shapes import pair1
from acfc_shapes import sttm as sttm_fixtures
from acfc_shapes.common import xlsx_bytes
from codegen import cli
from codegen.emit.context import build_context
from codegen.extract import contract_to_json as sttm_to_json
from codegen.extract import extract_contract
from codegen.extract.frd_docx import contract_to_json as frd_to_json
from codegen.extract.frd_docx import strip_value_label
from codegen.gate import natural_key_check
from codegen.layout.resolve import resolve_pair
from codegen.resolve.resolver import resolve_pair as resolve_contracts

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PAIR_1_FRD = SHAPES / "frd" / "f1_pair_1.docx"
PAIR_1_VDD = SHAPES / "vdd" / "pair_1_v1_segments.xlsx"
DATE = "2026-01-01"
# The v0.6.0 cell: a label token, then the storage-rooted folder path.
STORAGE = "/synstorage/mftlanding"
LABELLED_LANDING = f"/Path : {STORAGE}/{pair1.DOMAIN}/{pair1.SUB_DOMAIN}/"
BARE_LANDING = f"{STORAGE}/{pair1.DOMAIN}/{pair1.SUB_DOMAIN}/"


def _spec(tmp_path: Path, config, sttm: Path, frd: Path):
    pair = resolve_pair(sttm, frd, config, provider=None, cache_dirs=[], vdd_path=PAIR_1_VDD,
                        generated_date=DATE)
    assert pair.questions == [], [q.key for q in pair.questions]
    frd_json, sttm_json = tmp_path / "frd.json", tmp_path / "sttm.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config)
    return pair, spec


def _scoped(config, tmp: Path):
    return config.model_copy(update={"output": config.output.model_copy(update={
        "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")})})


# ------------------------------------------------ 1. natural key across segments


@pytest.fixture(scope="module")
def trailer_key_pair(pair1_config, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("m101_trailer_key")
    sttm = tmp / "sttm_trailer_key.xlsx"
    sttm.write_bytes(xlsx_bytes(sttm_fixtures.build_pair1(trailer_key=True)))
    return tmp, sttm


def test_the_variant_has_the_captured_shape(pair1_config, trailer_key_pair, tmp_path):
    tmp, sttm = trailer_key_pair
    _pair, spec = _spec(tmp_path, pair1_config, sttm, PAIR_1_FRD)
    assert spec.is_segmented
    trailer = next(s for s in spec.segments if s.segment == "Trailer")
    by_source = {f.source_column: f.stage_column for f in trailer.fields}
    assert by_source[pair1.TRAILER_KEY_SOURCE] == pair1.TRAILER_KEY_STAGE
    detail = {f.source_column: f for f in spec.detail_segment.fields}
    assert detail[pair1.TRAILER_KEY_SOURCE].stage_column == pair1.TRAILER_KEY_SOURCE
    assert not detail[pair1.TRAILER_KEY_SOURCE].nullable


def test_a_shared_source_name_keys_on_the_detail_segments_stage_column(
        pair1_config, trailer_key_pair, tmp_path):
    """The resolver used to let the LAST segment win the source-name -> stage
    map: the detail MERGE was keyed on the trailer's stage column."""
    tmp, sttm = trailer_key_pair
    _pair, spec = _spec(tmp_path, pair1_config, sttm, PAIR_1_FRD)
    assert pair1.TRAILER_KEY_SOURCE in spec.natural_key_columns
    assert pair1.TRAILER_KEY_STAGE not in spec.natural_key_columns
    assert spec.natural_key_columns == spec.not_null_columns


def test_build_context_resolves_the_key_against_every_segment(pair1_config, trailer_key_pair,
                                                               tmp_path):
    """Defence behind the resolver fix: a key column that IS a stage column of
    another segment resolves there — never a KeyError."""
    tmp, sttm = trailer_key_pair
    _pair, spec = _spec(tmp_path, pair1_config, sttm, PAIR_1_FRD)
    as_v060 = spec.model_copy(update={"natural_key_columns": [
        pair1.TRAILER_KEY_STAGE if c == pair1.TRAILER_KEY_SOURCE else c
        for c in spec.natural_key_columns]})
    context = build_context(as_v060, pair1_config, allow_duplicate_file_name=True,
                            duplicate_rule_text=None, notification_rule_texts=[])
    assert context["natural_key_missing"] == []
    assert context["natural_key_segments"][pair1.TRAILER_KEY_STAGE] == "Trailer"
    assert context["natural_key_segments"]["CARDHOLDER_ID"] == "Detail"
    assert len(context["detail_natural_key_source_columns"]) == len(as_v060.natural_key_columns)
    assert natural_key_check(context).passed


def test_a_key_column_no_segment_carries_is_a_named_fail_never_a_key_error(
        pair1_config, trailer_key_pair, tmp_path):
    tmp, sttm = trailer_key_pair
    _pair, spec = _spec(tmp_path, pair1_config, sttm, PAIR_1_FRD)
    broken = spec.model_copy(update={"natural_key_columns": [*spec.natural_key_columns,
                                                             "REC_TYP_NOWHERE"]})
    context = build_context(broken, pair1_config, allow_duplicate_file_name=True,
                            duplicate_rule_text=None, notification_rule_texts=[])
    assert context["natural_key_missing"] == ["REC_TYP_NOWHERE"]
    check = natural_key_check(context)
    assert not check.passed and check.name == "natural_key_columns"
    assert "REC_TYP_NOWHERE" in check.details and "Trailer" in check.details
    gate = cli._generate_feed(broken, _scoped(pair1_config, tmp_path), dry_run=True,
                              skip_tests=True, output_mode="framework",
                              conventions_profile="acfc_prx", iig_template="iig_v2")
    assert gate.verdict == "FAIL"
    assert any(c.name == "natural_key_columns" and not c.passed for c in gate.checks)


def test_the_cli_completes_on_the_trailer_key_variant(pair1_config, trailer_key_pair,
                                                      tmp_path):
    tmp, sttm = trailer_key_pair
    _pair, spec = _spec(tmp_path, pair1_config, sttm, PAIR_1_FRD)
    gate = cli._generate_feed(spec, _scoped(pair1_config, tmp_path), dry_run=True,
                              skip_tests=True, output_mode="framework",
                              conventions_profile="acfc_prx", iig_template="iig_v2")
    assert gate.verdict != "FAIL", [c for c in gate.checks if not c.passed]
    # a passing key check is not added — it would move every baseline report
    assert not any(c.name == "natural_key_columns" for c in gate.checks)


def test_the_app_runner_completes_on_the_trailer_key_variant(monkeypatch, tmp_path,
                                                             trailer_key_pair):
    """The App path (DemoRunner._execute): the pair-1 variant in an extra
    input folder, chosen through the chooser, generated end to end."""
    pytest.importorskip("fastapi")
    from ui.backend import stores as ui_stores
    from ui.backend.demo import DemoRunner
    from ui.backend.service import GenerationStore

    _tmp, sttm = trailer_key_pair
    folder = tmp_path / "pairs" / "pair_1"
    folder.mkdir(parents=True)
    (folder / "STTM_alpha.xlsx").write_bytes(sttm.read_bytes())
    (folder / "FRD_bravo.docx").write_bytes(PAIR_1_FRD.read_bytes())
    (folder / "VDD_charlie.xlsx").write_bytes(PAIR_1_VDD.read_bytes())
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_PROVIDER", "1")
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    monkeypatch.setenv("CODEGEN_EXTRA_INPUT_DIRS", f"local:{(tmp_path / 'pairs').as_posix()}")
    monkeypatch.setenv("CODEGEN_STORAGE_STATE", f"local:{(tmp_path / 'state').as_posix()}")
    monkeypatch.setenv("CODEGEN_STORAGE_OUTPUTS", f"local:{(tmp_path / 'out').as_posix()}")
    (tmp_path / "state").mkdir()
    (tmp_path / "out").mkdir()
    ui_stores.reset_stores()
    try:
        store = GenerationStore(str(REPO / "config" / "config.yaml"))
        config = store.config
        monkeypatch.setattr(store, "config", config.model_copy(update={
            "conventions": config.conventions.model_copy(update={"profile": "acfc_prx"}),
            "metadata": config.metadata.model_copy(update={"template": "iig_v2"}),
            "inputs": config.inputs.model_copy(update={"listing_ttl_seconds": 0.0})}))
        runner = DemoRunner(store)
        runner.select_output_mode("framework")
        runner.select_workbook("STTM_alpha.xlsx")
        runner.start_live()
        deadline = time.monotonic() + 600
        while runner.state in ("running", "needs_layout"):
            assert time.monotonic() < deadline, "the live run did not finish"
            time.sleep(0.5)
        assert runner.state == "done", f"{runner.error}\n{runner.stages}"
        assert store.runs, "the run produced no feeds"
        (run,) = store.runs.values()
        assert run.gate.verdict != "FAIL", [c for c in run.gate.checks if not c.passed]
        assert pair1.TRAILER_KEY_SOURCE in run.spec.natural_key_columns
        assert pair1.TRAILER_KEY_STAGE not in run.spec.natural_key_columns
    finally:
        ui_stores.reset_stores()


# ---------------------------------------------- 2. label-prefixed path values


@pytest.mark.parametrize("cell, remainder, label", [
    ("/Path : abfss://c@s.dfs.core.windows.net/dom/sub/",
     "abfss://c@s.dfs.core.windows.net/dom/sub/", "Path"),
    ("Path: /dom/sub/", "/dom/sub/", "Path"),
    ("ADLS Location :  /dom/sub", "/dom/sub", "ADLS Location"),
    ("adls path:/x/y", "/x/y", "adls path"),
    ("/dom/sub/", "/dom/sub/", None),                    # a bare path: untouched
    ("Pathway: /x", "Pathway: /x", None),               # not a label token
    ("Path :", "Path :", None),                         # a label with no value: untouched
])
def test_strip_value_label(cell, remainder, label):
    prefixes = ["Path", "ADLS Path", "Location", "ADLS Location", "Landing Path"]
    assert strip_value_label(cell, prefixes) == (remainder, label)


def test_the_tracked_config_lists_the_labels(config):
    assert config.extractor.frd.value_label_prefixes == [
        "Path", "ADLS Path", "Location", "ADLS Location", "Landing Path"]


def test_a_labelled_adls_location_is_the_path_without_the_label(pair1_config, tmp_path):
    frd = tmp_path / "frd_labelled.docx"
    frd.write_bytes(frd_fixtures.build_f1_pair1(landing=LABELLED_LANDING))
    sttm = SHAPES / "sttm" / "pair_1_family_a.xlsx"
    pair, spec = _spec(tmp_path, pair1_config, sttm, frd)
    feed = pair.frd_contract.feeds[0]
    assert feed.landing_location == BARE_LANDING
    evidence = pair.frd_contract.field_provenance["feeds[0].landing_location"]
    assert evidence.stripped_label == "Path"
    assert spec.landing_location == BARE_LANDING
    # the bare cell records no stripped label (and its JSON is unchanged)
    (tmp_path / "bare").mkdir()
    bare_pair, _ = _spec(tmp_path / "bare", pair1_config, sttm, PAIR_1_FRD)
    assert bare_pair.frd_contract.field_provenance[
        "feeds[0].landing_location"].stripped_label is None
    assert "stripped_label" not in frd_to_json(bare_pair.frd_contract)


def test_the_derivations_check_passes_and_the_iig_dml_carry_the_bare_path(pair1_config,
                                                                          tmp_path):
    frd = tmp_path / "frd_labelled.docx"
    frd.write_bytes(frd_fixtures.build_f1_pair1(landing=LABELLED_LANDING))
    sttm = SHAPES / "sttm" / "pair_1_family_a.xlsx"
    _pair, spec = _spec(tmp_path, pair1_config, sttm, frd)
    gate = cli._generate_feed(spec, _scoped(pair1_config, tmp_path), dry_run=True,
                              skip_tests=True, output_mode="framework",
                              conventions_profile="acfc_prx", iig_template="iig_v2")
    derivations = next(c for c in gate.checks if c.name == "derivations")
    assert derivations.passed, derivations.details
    framework = tmp_path / "out" / spec.feed_slug / "framework"
    q1 = (framework / "config_inserts_q1.sql").read_text(encoding="utf-8")
    assert BARE_LANDING.rstrip("/") in q1
    assert "Path :" not in q1 and "/Path" not in q1
    from openpyxl import load_workbook

    rows = load_workbook(framework / "config_rows.xlsx", read_only=True)
    cells = [str(c) for ws in rows.worksheets for row in ws.iter_rows(values_only=True)
             for c in row if c is not None]
    assert any(BARE_LANDING.rstrip("/") in c for c in cells)
    assert not any("Path :" in c or c.startswith("/Path") for c in cells)
