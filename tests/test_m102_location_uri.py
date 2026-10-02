"""M10.2 (ported from staging 435e4a5): a location URI landing is carried as
written through the IIG and the DML.

Pair 1 with the FRD's ADLS Location set to an ``abfss://`` URI, run end to
end under ``acfc_prx`` + ``iig_v2`` in framework mode. The variant is built at
runtime: the tracked ``fixtures/acfc_shapes/frd/f1_pair_1.docx`` is copied into
``tmp_path`` and its ADLS Location cell is set to a synthetic URI on the
reserved ``.invalid`` TLD — no extra tracked fixture. Before the port the URI went
through the folder-path rules: every path cell became ``/abfss:/…`` or
``/Archive/abfss:/…`` and the derivations check FAILed the run.

Adapted from staging's ``tests/test_m101_real_run_bugs.py`` (absent on this
branch): the URI here is bare, not ``/Path : <uri>`` — the label strip is the
M10.1 reader change, which was not ported — and the reader's
``value_kind`` provenance is not asserted for the same reason.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from acfc_shapes import pair1
from codegen import cli
from conftest import with_dml

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
LAYOUT_PROFILES = REPO / "fixtures" / "layout_profiles"
FRD_PAIR1 = SHAPES / "frd" / "f1_pair_1.docx"
DATE = "2026-01-01"
FOLDER_LANDING = f"/{pair1.DOMAIN}/{pair1.SUB_DOMAIN}/"     # the tracked FRD's ADLS Location
URI_LANDING = f"abfss://syn-container@synstorage.example.invalid{FOLDER_LANDING}"


def _uri_variant(tmp: Path) -> Path:
    """A copy of the pair-1 FRD whose ADLS Location cell is URI_LANDING."""
    out = tmp / "f1_pair_1_uri.docx"
    with zipfile.ZipFile(FRD_PAIR1) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                assert xml.count(FOLDER_LANDING) == 1, "the ADLS Location cell moved"
                data = xml.replace(FOLDER_LANDING, URI_LANDING).encode("utf-8")
            dst.writestr(item, data)
    return out


def _pair1_spec(config, tmp: Path, frd: Path):
    from codegen.extract import contract_to_json as sttm_to_json
    from codegen.extract import extract_contract
    from codegen.extract.frd_docx import contract_to_json as frd_to_json
    from codegen.extract.vdd import contract_to_json as vdd_to_json
    from codegen.extract.vdd import extract_vdd_contract
    from codegen.layout.model import MockLayoutProvider
    from codegen.layout.resolve import resolve_pair as resolve_layout_pair
    from codegen.resolve.resolver import resolve_pair as resolve_contracts

    sttm = SHAPES / "sttm" / "pair_1_family_a.xlsx"
    vdd_path = SHAPES / "vdd" / "pair_1_v1_segments.xlsx"
    pair = resolve_layout_pair(
        sttm, frd, config,
        provider=MockLayoutProvider([LAYOUT_PROFILES / "mock", LAYOUT_PROFILES]),
        cache_dirs=[tmp / "cache"], generated_date=DATE, vdd_path=vdd_path)
    assert pair.questions == []
    frd_json = tmp / "frd.contract.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    sttm_json = tmp / "sttm.contract.json"
    sttm_json.write_text(sttm_to_json(extract_contract(
        sttm, frd_json, config, generated_date=DATE, layout=pair.sttm.profile)),
        encoding="utf-8")
    vdd_json = tmp / "vdd.contract.json"
    vdd, _ = extract_vdd_contract(vdd_path, config, generated_date=DATE)
    vdd_json.write_text(vdd_to_json(vdd), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config, vdd_path=vdd_json)
    return pair, spec


@pytest.fixture(scope="module")
def uri_run(pair1_config, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pair1_uri")
    pair, spec = _pair1_spec(pair1_config, tmp, _uri_variant(tmp))
    scoped = with_dml(pair1_config).model_copy(update={"output": pair1_config.output.model_copy(
        update={"dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")})})
    gate = cli._generate_feed(spec, scoped, dry_run=True, skip_tests=True,
                              output_mode="framework", conventions_profile="acfc_prx",
                              iig_template="iig_v2")
    framework = tmp / "out" / spec.feed_slug / "framework"
    rows = load_workbook(framework / "config_rows.xlsx", read_only=True)
    values = [str(c) for ws in rows.worksheets for row in ws.iter_rows(values_only=True)
              for c in row if c is not None]
    return pair, spec, gate, framework, values


def test_the_uri_reaches_the_spec_unchanged(uri_run):
    pair, spec, _gate, _framework, _values = uri_run
    assert pair.frd_contract.feeds[0].landing_location == URI_LANDING
    assert spec.landing_location == URI_LANDING


def test_the_uri_is_kept_intact_in_the_iig_and_the_dml(uri_run):
    _pair, _spec, _gate, framework, values = uri_run
    assert URI_LANDING in values                                    # TGT_ADLS_PATH
    assert not any(v.startswith(("/abfss", "/Archive/abfss", "abfss:/syn")) for v in values)
    assert all(v.count("://") <= 1 for v in values)
    for env in ("q1", "a2", "prod"):
        dml = (framework / f"config_inserts_{env}.sql").read_text(encoding="utf-8")
        assert f"N'{URI_LANDING}'" in dml
        assert "/abfss" not in dml and "abfss:/syn" not in dml


def test_archive_and_processed_are_appended_inside_the_uri(uri_run):
    _pair, spec, _gate, framework, values = uri_run
    assert URI_LANDING + "Archive/" in values      # both {landing}Archive/ and /Archive{landing}
    assert URI_LANDING + "Processed/" in values
    assert URI_LANDING + f"Processed/{spec.detail_segment.stage_table.table}" in values
    q1 = (framework / "config_inserts_q1.sql").read_text(encoding="utf-8")
    assert f"N'{URI_LANDING}Archive/'" in q1
    assert f"N'{URI_LANDING}Processed/{spec.detail_segment.stage_table.table}'" in q1


def test_a_template_shape_that_prefixes_the_landing_is_flagged(uri_run):
    _pair, _spec, gate, _framework, _values = uri_run
    shape_flags = [f for f in gate.flags if f.startswith("iig_path_shape_on_uri:")]
    assert shape_flags, gate.flags
    assert all("'/Archive{landing}'" in f for f in shape_flags)
    assert any(f.startswith("iig_path_shape_on_uri:STGDELTA_STDDELTA_INGESTION_DET.")
               for f in shape_flags)


def test_the_run_passes_with_flags_under_the_uri_rule_instead_of_failing(uri_run):
    _pair, _spec, gate, _framework, _values = uri_run
    derivations = next(c for c in gate.checks if c.name == "derivations")
    assert derivations.passed, derivations.details
    assert "location URI cell(s) checked by the location URI rule" in derivations.details
    assert next(c for c in gate.checks if c.name == "sql_literals").passed
    assert gate.verdict == "PASS_WITH_FLAGS", [c for c in gate.checks if not c.passed]
