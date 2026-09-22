"""M11 addendum (SHAPES_ROUND2) — items 7 and 11, the two that move the most
pairs: the FRD family never blocks (pairs 8/9/10), and a sheet's REAL used
range replaces a trusted max_row (pair 3)."""

from __future__ import annotations

import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from acfc_shapes.common import docx_bytes, paragraph, table
from codegen.layout.extent import load_document, trim_worksheet, used_max_row
from codegen.layout.resolve import resolve_pair
from codegen.layout.size import check_workbook_size, workbook_cells

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PAIR_1_STTM = SHAPES / "sttm" / "pair_1_family_a.xlsx"
PAIR_1_FRD = SHAPES / "frd" / "f1_pair_1.docx"
PAIR_1_VDD = SHAPES / "vdd" / "pair_1_v1_segments.xlsx"
DATE = "2026-01-01"


# ======================================================= item 11: used range


def test_used_max_row_stops_at_the_first_long_gap():
    assert used_max_row([], 500) == (0, 0)
    assert used_max_row([1, 2, 3, 300], 500) == (300, 0)          # gaps shorter than 500
    assert used_max_row([1, 2, 3, 2000], 500) == (3, 1)            # 1996 empty rows: stop
    assert used_max_row([1, 5000], 0) == (5000, 0)                 # 0 = trust everything


def test_a_dead_tail_is_trimmed_and_valued_rows_beyond_a_gap_are_reported():
    workbook = Workbook()
    ws = workbook.active
    for row in range(1, 11):
        ws.cell(row=row, column=1).value = f"v{row}"
    ws.cell(row=2000, column=1).value = "after the gap"
    ws.cell(row=900_000, column=3).font = Font(bold=True)          # formatting only
    declared, used, beyond = trim_worksheet(ws, 500)
    assert (declared, used, beyond) == (900_000, 10, 1)
    assert ws.max_row == 10


def test_a_few_formatted_empty_rows_keep_max_row_exact():
    """max_row is fingerprint input: a normal sheet must not move it."""
    workbook = Workbook()
    ws = workbook.active
    for row in range(1, 11):
        ws.cell(row=row, column=1).value = row
    ws.cell(row=40, column=2).font = Font(bold=True)               # 30 empty rows: a tail < 500
    assert trim_worksheet(ws, 500) == (40, 40, 0)
    assert ws.max_row == 40


def _bloated_pair1(tmp_path: Path) -> Path:
    """The pair-1 STTM with ONE formatted empty cell a million rows down —
    the SHAPES_ROUND2 §3 mechanism (max_row 1,048,538 over 336 real rows)."""
    workbook = load_workbook(PAIR_1_STTM)
    workbook["FEED_1_MAPPING"].cell(row=1_000_000, column=30).font = Font(bold=True)
    path = tmp_path / "bloated.xlsx"
    workbook.save(path)
    assert load_workbook(path, data_only=True)["FEED_1_MAPPING"].max_row == 1_000_000
    return path


def test_the_loader_reads_the_used_range_not_max_row(tmp_path, config):
    workbook = load_document(_bloated_pair1(tmp_path), config.extractor.used_range_empty_rows)
    assert workbook["FEED_1_MAPPING"].max_row < 200
    trimmed = workbook.trim_report.trimmed
    assert "FEED_1_MAPPING" in trimmed and trimmed["FEED_1_MAPPING"][0] == 1_000_000
    assert "formatting only" in workbook.trim_report.notes()[0]


def test_a_bloated_sheet_resolves_fast_and_exactly_like_the_clean_one(tmp_path, pair1_config):
    """Without the used range this resolve took 151 s here — past the App's
    120 s budget, the v0.6.2 pair-3 failure. With it: well under a second."""
    from codegen.extract import contract_to_json as sttm_to_json
    from codegen.extract import extract_contract
    from codegen.extract.frd_docx import contract_to_json as frd_to_json

    def contract(sttm: Path, work: Path) -> str:
        pair = resolve_pair(sttm, PAIR_1_FRD, pair1_config, provider=None, cache_dirs=[],
                            generated_date=DATE)
        assert pair.questions == []
        frd = work / "frd.json"
        frd.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
        text = sttm_to_json(extract_contract(sttm, frd, pair1_config, generated_date=DATE,
                                             layout=pair.sttm.profile))
        return text.replace(sttm.name, "<workbook>")

    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    started = time.monotonic()
    bloated = contract(_bloated_pair1(tmp_path / "a"), tmp_path / "a")
    elapsed = time.monotonic() - started
    clean = contract(PAIR_1_STTM, tmp_path / "b")
    assert elapsed < 30, f"the bloated sheet took {elapsed:.1f}s"
    assert bloated == clean


def _dimension_lie(tmp_path: Path, rows: int, cols: int, claimed: str) -> Path:
    """A workbook holding ``rows`` x ``cols`` whose <dimension> CLAIMS more —
    what read-only mode reports as max_row."""
    workbook = Workbook()
    ws = workbook.active
    for r in range(1, rows + 1):
        ws.append([f"c{r}_{c}" for c in range(cols)])
    honest = tmp_path / "honest.xlsx"
    workbook.save(honest)
    lying = tmp_path / "lying.xlsx"
    with zipfile.ZipFile(honest) as src, zipfile.ZipFile(lying, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                text = data.decode("utf-8")
                start = text.index('<dimension ref="') + len('<dimension ref="')
                end = text.index('"', start)
                data = (text[:start] + claimed + text[end:]).encode("utf-8")
            dst.writestr(item, data)
    return lying


def test_the_cap_measures_the_used_range_not_the_declared_one(tmp_path, config):
    """A dimension that claims 1M rows over 300 real ones (SHAPES_ROUND2 §3:
    1,048,538 x 34 declared, ~336 rows held) is READ — refusing it on the
    declaration would make a readable document permanently unreadable."""
    lying = _dimension_lie(tmp_path, rows=300, cols=34, claimed="A1:AH1048538")
    read_only = load_workbook(lying, read_only=True)
    assert read_only.active.max_row == 1_048_538                    # the claim
    read_only.close()
    total, per_sheet, streamed = workbook_cells(lying, config.extractor.used_range_empty_rows,
                                                config.inputs.max_workbook_cells)
    assert streamed and total == 300 * 34                           # what is there
    check_workbook_size(lying, config.inputs.max_workbook_cells,
                        config.extractor.used_range_empty_rows)     # passes
    del per_sheet


def test_the_cap_still_refuses_a_workbook_that_is_really_that_big(tmp_path, config):
    big = tmp_path / "big.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("BIG")
    for r in range(2600):
        sheet.append([f"c{r}_{c}" for c in range(100)])             # 260k used cells
    workbook.save(big)
    from codegen.layout.size import WorkbookTooLarge

    with pytest.raises(WorkbookTooLarge, match="used cells, > cap"):
        check_workbook_size(big, config.inputs.max_workbook_cells,
                            config.extractor.used_range_empty_rows)


# ============================================== item 7: family never blocks


def _unrecognized_frd(tmp_path: Path) -> Path:
    """An FRD with no F1 metadata table and no Solution Requirement table —
    only boilerplate (the v0.6.2 pairs 8/9/10 error, before item 8 / 9)."""
    path = tmp_path / "topic_frd.docx"
    path.write_bytes(docx_bytes([
        paragraph("Revision History", style="Title"),
        table([["Date", "Version", "Author(s)", "Description of Version/Changes"],
               ["12/15/25", "1.0", "SYN Author", "Initial"]]),
        paragraph("Definitions and Acronyms", style="Heading2"),
        table([["Acronym", "Definition"], ["DL", "Data Lake"]]),
    ]))
    return path


def test_an_unrecognized_frd_is_filled_by_the_chain_and_asks_the_rest(tmp_path, pair1_config):
    pair = resolve_pair(PAIR_1_STTM, _unrecognized_frd(tmp_path), pair1_config, provider=None,
                        cache_dirs=[], generated_date=DATE, vdd_path=PAIR_1_VDD)
    contract = pair.frd_contract
    assert any(f.startswith("frd_family_unrecognized:") and "2 table(s)" in f
               for f in contract.extraction_flags)
    feed = contract.feeds[0]
    # the chain: the STTM bands name the feed and its tables, the VDD the format
    assert feed.feed_name == "vnd_p_accum_client"
    assert feed.stage_target.tables == ["vnd_p_accum_client"]
    assert feed.file_format and feed.file_name_patterns
    keys = {q.key: q for q in pair.questions}
    assert "feeds[0].feed_name" not in keys                 # the chain answered it
    for field in ("source_system", "lobs", "domain"):
        question = keys[f"feeds[0].{field}"]
        assert question.kind == "text" and question.candidates == []   # TYPED, answerable
        assert "frd_family_unrecognized" in question.reason


def test_typed_answers_fill_an_unrecognized_frd(tmp_path, pair1_config):
    answers = {"sttm": {}, "frd": {}, "vdd": {}, "gaps": {
        "feeds[0].source_system": {"value": "VENDOR_A", "source": "user"},
        "feeds[0].lobs": {"value": "MEDICAID; CHIP", "source": "user"},
        "feeds[0].domain": {"value": "Pharmacy", "source": "user"},
        "feeds[0].frequency": {"value": "Daily", "source": "user"},
        "feeds[0].load_strategy": {"value": "Append", "layer": "both", "source": "user"},
    }}
    pair = resolve_pair(PAIR_1_STTM, _unrecognized_frd(tmp_path), pair1_config, provider=None,
                        cache_dirs=[], generated_date=DATE, vdd_path=PAIR_1_VDD,
                        answers=answers)
    feed = pair.frd_contract.feeds[0]
    assert feed.source_system == "VENDOR_A"
    assert feed.lobs == ["MEDICAID", "CHIP"]                 # a list, split
    assert feed.domain == "Pharmacy"
    assert not [q for q in pair.questions if q.key in answers["gaps"]]


def test_the_cli_extract_frd_never_stops_on_the_family(tmp_path):
    import json

    from codegen.cli import main

    out = tmp_path / "frd.json"
    assert main(["extract-frd", "--config", str(REPO / "config" / "config.yaml"),
                 "--docx", str(_unrecognized_frd(tmp_path)), "--out", str(out)]) == 0
    contract = json.loads(out.read_text(encoding="utf-8"))
    assert any(f.startswith("frd_family_unrecognized:") for f in contract["extraction_flags"])
