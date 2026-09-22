"""M11 addendum items 10, 8, 12, 13, 9 (v0.7.2) — built from the round-2
shape capture, docs/acfc/SHAPES_ROUND2.md on ``origin/acfc-runs``. Every
fixture here is built in-test from the capture's STRUCTURE (header texts,
row anatomy, merged ranges); vendor and business terms are synthetic
stand-ins for the capture's ``<TOKEN_n>`` placeholders."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from codegen.layout.resolve import resolve_workbook

REPO = Path(__file__).resolve().parents[1]


# =========================================== item 10: layer-prefixed headers


# SHAPES_ROUND2 §2.1, the pair-7 mapping sheet's header row (row 1), column
# for column. No band row, zero merged ranges. "<TOKEN_19>" is the vendor.
PAIR7_HEADER = [
    "Field ID", "Field Name", "Mandatory or Situational", "Column Description",
    "VENDOR_G Data Fields - Comments", "Format", "Size", "PHI\\PII Field",
    "Stage Schema", "Stage Table Name", "Stage Table - Column Name", "Stage Table - DataType",
    "Standard Schema", "Standard Table Name", "Standard Table - Column Name",
    "Standard Data Type",
    "Mandatory Fields (Include in DQ Check)", "Recycle Flag ( Enabled for 7 Days)", "Comments",
]


def _pair7_round2(tmp_path: Path) -> Path:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "MAPPING_G"
    ws.append(PAIR7_HEADER)
    for n in range(1, 6):
        ws.append([n, f"FIELD_{n}", "M", f"Field {n}", "", "A/N", 10, "N",
                   "stg_g", "g_claims", f"FIELD_{n}", "string",
                   "std_g", "g_claims", f"FIELD_{n}", "string", "Y", "", ""])
    for audit in ("SRC_FILE_NAME", "REC_CREATION_TIME"):
        ws.append(["", "NA", "", "", "", "", "", "",
                   "stg_g", "g_claims", audit, "string",
                   "std_g", "g_claims", audit, "string", "", "", ""])
    details = workbook.create_sheet("File Details")
    details.append(["Vendor", "FileName", "File Description", "Location",
                    "Frequency (Historical Drops)"])
    details.append(["VENDOR_G", "G_CLAIMS_YYYYMMDD.txt", "claims", "", "Daily"])
    path = tmp_path / "pair7_round2.xlsx"
    workbook.save(path)
    return path


def test_layer_prefixed_headers_resolve_with_zero_questions(config, tmp_path):
    """SHAPES_ROUND2 §2.2 Q4/Q5: stage = I-L, standard = M-P."""
    empty = tmp_path / "cache"
    empty.mkdir()
    doc, _ = resolve_workbook(_pair7_round2(tmp_path), config, provider=None,
                              cache_dirs=[empty])
    assert doc.questions == [], [q.key for q in doc.questions]
    assert doc.complete
    (sheet,) = doc.profile.mapping_sheets
    assert sheet.band_row is None and sheet.header_row == 1
    stage = sheet.band("stage")
    standard = sheet.band("standard")
    assert {r: stage.roles[r] for r in ("schema", "table", "column", "target_type")} == {
        "schema": 9, "table": 10, "column": 11, "target_type": 12}
    assert {r: standard.roles[r] for r in ("schema", "table", "column", "target_type")} == {
        "schema": 13, "table": 14, "column": 15, "target_type": 16}
    assert sheet.band("source").roles["field_name"] == 2
    # the trailing group's qualified header ("… ( Enabled for 7 Days)")
    assert sheet.band("rules").roles.get("recycle_flag") == 18


def test_the_prefix_is_a_fallback_never_an_override(config):
    """A header the base synonyms already match is untouched by the prefix
    rule; a remainder that matches nothing stays unresolved."""
    from codegen.layout.discover import _fallback_matches, normalize

    disc = config.extractor.discovery
    table = {role: {normalize(s) for s in spellings}
             for role, spellings in disc.roles["target"].items()}
    assert _fallback_matches(normalize("Stage Table - Column Name"), "stage", "target",
                             table, disc) == ["column"]
    assert _fallback_matches(normalize("Standard Table - DataType"), "standard", "target",
                             table, disc) == ["target_type"]
    assert _fallback_matches(normalize("Stage Zqx 11"), "stage", "target", table, disc) == []
    # a SOURCE-band header is never prefix-stripped
    assert _fallback_matches(normalize("Stage Table - Column Name"), "source", "source",
                             table, disc) == []


# ======================================================= item 8: F2 detection


def _frd_contract(path: Path, config):
    from codegen.extract.frd_docx import discover_frd, read_docx, read_frd

    content = read_docx(path)
    profile = discover_frd(content, config.extractor.frd)
    return profile, read_frd(content, profile, config, document_name=path.name,
                             generated_date="2026-01-01")


FEED_FACTS = ("file_format", "frequency", "lobs", "domain", "sub_domain", "source_system")


def test_the_round2_pair8_geometry_is_f2_and_reads_like_round1(config, tmp_path):
    """SHAPES_ROUND2 §1.1: "Solution Requirement: N" in a cell merged across
    B-D after a leading column; row 4's LABEL names the section. Both
    plausible Word structures of that leading column read as F2, with the
    same feed facts as the round-1 (documented) fixture."""
    from acfc_shapes import frd as frd_fixtures

    _p, round1 = _frd_contract(REPO / "fixtures" / "acfc_shapes" / "frd" / "f2_pair_8.docx",
                               config)
    expected = {f: getattr(round1.feeds[0], f) for f in FEED_FACTS}
    expected_tables = round1.feeds[0].stage_target.tables
    for geometry in ("vmerge", "lead"):
        path = tmp_path / f"pair8_{geometry}.docx"
        path.write_bytes(frd_fixtures.build_f2_round2_pair8(geometry))
        profile, contract = _frd_contract(path, config)
        assert profile.family == "F2", geometry
        assert [r.section for r in profile.sections] == [
            "descriptive", "structural", "administrative", "technical", "data_quality",
            "vendor", "email"], geometry
        feed = contract.feeds[0]
        assert {f: getattr(feed, f) for f in FEED_FACTS} == expected, geometry
        assert feed.stage_target.tables == expected_tables, geometry
        assert not any(f.startswith("frd_family_unrecognized") for f in
                       contract.extraction_flags)
        if geometry == "vmerge":
            assert any("leading merged column" in n for n in profile.notes)


def test_a_solution_requirement_title_with_a_sub_id_and_a_title_is_detected(config):
    from codegen.extract.frd_docx import _sr_title, normalize_label

    prefix = normalize_label(config.extractor.frd.solution_requirement_prefix)
    assert _sr_title(["", "Solution Requirement: 2.2 – Inbound File Data Ingestion"],
                     prefix) == "Solution Requirement: 2.2 – Inbound File Data Ingestion"
    assert _sr_title(["Solution Requirement: 1", "", ""], prefix)          # round 1: col 0
    assert _sr_title(["Non-Functional Requirement ID: 3"], prefix) is None


# ====================================== item 12: a blank band label; NOTE rows


# SHAPES_ROUND2 §4.1 (pair 5): meta rows r1-r9, band row r10 with three
# merged groups — A10:K10 "Source Layout", L10:V10 "Stage Layer", W10:AG10
# BLANK (the standard layer implied) — and the header row r11.
PAIR5_META = [("File(s)", "VENDOR_E_OUTREACH_*.txt"), ("File Generator", None),
              ("File Location", "/landing/vendor_e/"), ("LOB", "ALL"),
              ("File frequency", "Monthly"), ("Domain", "Care Management"),
              ("Sub-Domain", "Outreach"),
              ("NOTE", 'Fields Id to be separated by "|" delimiter and field values to be in '
                       'double (") quotes.'),
              ("File type", "Text")]
PAIR5_SOURCE = ["Field ID", "Field Name", "Description", "Data Type", "Length", "Start", "End",
                "PII", "Sample Value", "Comments", "Required"]
PAIR5_TARGET = ["Workspace", "Catalog", "Schema", "Table Name", "Column Name", "Data Type",
                "Primary Key", "Mandatory Column", "Constraints", "Field Description",
                "Transformation"]


def _pair5_round2(tmp_path: Path) -> Path:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "MAPPING_E"
    for label, value in PAIR5_META:
        ws.append([label, value])
    ws.append(["Source Layout"] + [None] * 10 + ["Stage Layer"] + [None] * 21)
    ws.merge_cells("A10:K10")
    ws.merge_cells("L10:V10")
    ws.merge_cells("W10:AG10")                                     # blank label
    ws.append(PAIR5_SOURCE + PAIR5_TARGET + PAIR5_TARGET)
    for n in range(1, 5):
        target = ["DLK", "cat_e", "stg_e", "e_outreach", f"FIELD_{n}", "string",
                  "N", "N", "", f"Field {n}", "Load as is"]
        ws.append([n, f"FIELD_{n}", f"Field {n}", "A/N", 10, None, None, "N", "x", "", "N",
                   *target, *[("std_e" if v == "stg_e" else v) for v in target]])
    for audit in ("SRC_FILE_NAME", "REC_CREATION_TIME"):
        target = ["DLK", "cat_e", "stg_e", "e_outreach", audit, "string", "", "", "", "", ""]
        ws.append([None, "NA", None, None, None, None, None, None, None, None, None,
                   *target, *[("std_e" if v == "stg_e" else v) for v in target]])
    path = tmp_path / "pair5_round2.xlsx"
    workbook.save(path)
    return path


def test_a_blank_band_label_takes_its_layer_from_the_headers(config, tmp_path):
    empty = tmp_path / "cache"
    empty.mkdir()
    doc, _ = resolve_workbook(_pair5_round2(tmp_path), config, provider=None,
                              cache_dirs=[empty])
    (sheet,) = doc.profile.mapping_sheets
    assert sheet.band_row == 10 and sheet.header_row == 11
    stage, standard = sheet.band("stage"), sheet.band("standard")
    assert (stage.col_start, stage.col_end) == (12, 22)
    assert (standard.col_start, standard.col_end) == (23, 33)       # W:AG, no label
    assert standard.label is None
    for role, col in (("schema", 25), ("table", 26), ("column", 27), ("target_type", 28)):
        assert standard.roles[role] == col
    assert doc.questions == [], [q.key for q in doc.questions]


def test_a_note_meta_row_is_never_a_value_source(config, tmp_path):
    empty = tmp_path / "cache"
    empty.mkdir()
    doc, _ = resolve_workbook(_pair5_round2(tmp_path), config, provider=None,
                              cache_dirs=[empty])
    (sheet,) = doc.profile.mapping_sheets
    note = next(m for m in sheet.meta_rows if m.label == "NOTE")
    assert note.key == "notes" and note.value_col is None            # recorded, never read
    # the '|' in the NOTE's prose never becomes the file's delimiter
    stated = {m.key for m in sheet.meta_rows if m.value_col is not None}
    assert "delimiter" not in stated


# ============================== item 13: scope, audit rows, types, RDBMS source


# SHAPES_ROUND2 §4.2 (pair 6): band row r2 — B2:I2 "Source table", L2:Q2
# "Stage Layer", S2:X2 "STD layer"; header row r3 (22 columns, gaps at K and
# R); column A holds In Scope / Out of scope; the audit rows are marked in
# the Load Rules column J ("Audit Column"), two of them in ONE layer only, and
# their types include tinyint and timestamp(YYYY-MM-DD HH:MM:SS).
PAIR6_TARGET = ["Workspace", "Catalog", "Schema", "TableName", "ColumnName", "DataType"]
PAIR6_HEADER = (["Inscope for VENDOR_F Implementation", "SERVER NAME", "TABLE_CATALOG",
                 "TABLE_SCHEMA", "TABLE_NAME", "COLUMN_NAME", "Primary Key", "DATA_TYPE",
                 "NULL/NOT NULL", "Load Rules", None] + PAIR6_TARGET + [None] + PAIR6_TARGET)


def _pair6_round2(tmp_path: Path) -> Path:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "EXT_F_XWLK"
    ws.append(["Source to target mapping"])
    band = [None] * 24
    band[1], band[11], band[18] = "Source table", "Stage Layer", "STD layer"
    ws.append(band)
    ws.merge_cells("B2:I2")
    ws.merge_cells("L2:Q2")
    ws.merge_cells("S2:X2")
    ws.append(PAIR6_HEADER)

    def target(schema, column, dtype):
        return ["DLK", "cat_f", schema, "f_xwlk", column, dtype]

    for n, scope in enumerate(["In Scope", "In Scope", "Out of scope", "In Scope",
                               "Out of scope"], start=1):
        ws.append([scope, "SRV_REG_1", "DB_F", "dbo", "XWLK", f"COL_{n}", "N", "varchar",
                   "NULL", "Straight move", None, *target("stg_f", f"COL_{n}", "string"), None,
                   *target("std_f", f"COL_{n}", "string")])
    empty_source = [None] * 9
    audit = [
        ("Audit Column", ("DELETE_FLAG", "tinyint"), None),
        ("Audit Column", None, ("SRC_FILE_NAME", "string")),
        ("Audit Column", ("REC_CREATION_TIME", "timestamp(YYYY-MM-DD HH:MM:SS)"),
         ("REC_CREATION_TIME", "timestamp(YYYY-MM-DD HH:MM:SS)")),
        ("Audit Column + Hard code note", ("REGION_NAME", "string"), ("REGION_NAME", "string")),
    ]
    for rule, stage, std in audit:
        ws.append(["In Scope", *empty_source[1:], rule, None,
                   *(target("stg_f", *stage) if stage else [None] * 6), None,
                   *(target("std_f", *std) if std else [None] * 6)])
    path = tmp_path / "pair6_round2.xlsx"
    workbook.save(path)
    return path


def _pair6_contract(tmp_path: Path, config):
    import json

    from codegen.extract import extract_contract
    from segmented_fixture import frd_contract_dict

    sttm = _pair6_round2(tmp_path)
    empty = tmp_path / "cache"
    empty.mkdir()
    doc, _ = resolve_workbook(sttm, config, provider=None, cache_dirs=[empty])
    data = frd_contract_dict()
    data["feeds"][0].update({
        "feed_name": "f_xwlk", "file_name_patterns": ["F_XWLK_*.txt"], "file_format": "csv",
        "delimiter": ",", "record_segments": [],
        "stage_target": {"catalog": "cat_f", "schema": "stg_f", "tables": ["f_xwlk"],
                         "load_strategy": "Truncate and Load"},
        "standard_target": {"catalog": "cat_f", "schema": "std_f", "tables": ["f_xwlk"],
                            "load_strategy": "Append"}})
    frd = tmp_path / "frd.json"
    frd.write_text(json.dumps(data), encoding="utf-8")
    return doc, frd, extract_contract(sttm, frd, config, generated_date="2026-01-01",
                                      layout=doc.profile)


def test_scope_markers_filter_rows_with_one_grouped_flag(config, tmp_path):
    doc, _frd, contract = _pair6_contract(tmp_path, config)
    assert doc.questions == [], [q.key for q in doc.questions]
    feed = contract.feeds[0]
    assert [f.source_column for f in feed.fields] == ["COL_1", "COL_2", "COL_4"]
    grouped = [f for f in feed.extraction_flags if f.startswith("rows_out_of_scope:")]
    assert len(grouped) == 1 and "2 row(s)" in grouped[0] and "column A" in grouped[0]


def test_load_rule_audit_rows_one_layer_and_type_formats(config, tmp_path):
    _doc, _frd, contract = _pair6_contract(tmp_path, config)
    feed = contract.feeds[0]
    audit = {a.column: a.datatype for a in feed.audit_columns}
    assert audit == {"DELETE_FLAG": "tinyint", "SRC_FILE_NAME": "String",
                     "REC_CREATION_TIME": "Timestamp", "REGION_NAME": "String"}
    flags = feed.extraction_flags
    assert any(f.startswith("audit_column_one_layer:DELETE_FLAG") and "stage layer" in f
               for f in flags)
    assert any(f.startswith("audit_column_one_layer:SRC_FILE_NAME") and "standard layer" in f
               for f in flags)
    assert any(f.startswith("type_format:REC_CREATION_TIME") and "YYYY-MM-DD HH:MM:SS" in f
               for f in flags)


def test_split_type_format_keeps_a_precision():
    from codegen.formats import split_type_format

    assert split_type_format("timestamp(YYYY-MM-DD HH:MM:SS)") == (
        "timestamp", "YYYY-MM-DD HH:MM:SS")
    assert split_type_format("decimal(10,2)") == ("decimal(10,2)", None)
    assert split_type_format("varchar(50)") == ("varchar(50)", None)
    assert split_type_format("string") == ("string", None)


def test_an_rdbms_source_writes_ddl_and_a_dml_review_block(pair1_config, tmp_path):
    from codegen import cli
    from codegen.emit.dml import validate_tsql
    from codegen.resolve.resolver import resolve_pair as resolve_contracts

    _doc, frd, contract = _pair6_contract(tmp_path, pair1_config)
    feed = contract.feeds[0]
    assert feed.source_kind == "rdbms"
    assert any(f.startswith("source_kind_rdbms:") for f in feed.extraction_flags)
    from codegen.extract import contract_to_json

    sttm_json = tmp_path / "sttm.json"
    sttm_json.write_text(contract_to_json(contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd, sttm_json, pair1_config)
    assert spec.source_kind == "rdbms"
    config = pair1_config.model_copy(update={"output": pair1_config.output.model_copy(
        update={"dir": str(tmp_path / "out"), "reports_dir": str(tmp_path / "reports")})})
    gate = cli._generate_feed(spec, config, dry_run=True, skip_tests=True,
                              output_mode="framework", conventions_profile="acfc_prx",
                              iig_template="iig_v2")
    framework = tmp_path / "out" / spec.feed_slug / "framework"
    assert list(framework.glob("*_DDL.txt"))                     # the DDL is generated
    q1 = (framework / "config_inserts_q1.sql").read_text(encoding="utf-8")
    assert "-- REVIEW FILE_ADLS_INGESTION_DETAILS: the source is an RDBMS table" in q1
    assert "INSERT INTO [dbo].[FILE_ADLS_INGESTION_DETAILS]" not in q1
    assert "FILE_CONNECTION_DETAILS" not in q1                   # no file-connection lookup
    assert validate_tsql(q1) == []
    assert any(f.startswith("dml_rdbms_review:FILE_ADLS_INGESTION_DETAILS") for f in gate.flags)
    assert all(c.passed for c in gate.checks if c.name.startswith("dml_"))


# ================================ item 9: F3 — topic-organized requirements


def test_pair9_is_f3_classified_with_no_field_read(config, tmp_path):
    """SHAPES_ROUND2 §1.2: nine topic requirement tables (none names a
    metadata section), twelve NFR tables, boilerplate — recognized and
    classified; no value is read (there is no domain table)."""
    from acfc_shapes import frd as frd_fixtures

    path = tmp_path / "pair9.docx"
    path.write_bytes(frd_fixtures.build_f3_pair9())
    profile, contract = _frd_contract(path, config)
    assert profile.family == "F3"
    note = profile.notes[0]
    assert "9 requirement" in note and "12 nfr" in note and "boilerplate" in note
    assert profile.fields == {}
    assert any(f.startswith("frd_family_f3:") for f in contract.extraction_flags)
    assert not any(f.startswith("frd_family_unrecognized") for f in contract.extraction_flags)


def test_pair10_is_f3_and_reads_its_domain_table(config, tmp_path):
    """SHAPES_ROUND2 §1.3: the 2-column Domain / SubDomain table is read;
    sub-numbered requirement titles ("2.2 – Inbound File Data Ingestion")
    are requirement tables; 13 NFR tables."""
    from acfc_shapes import frd as frd_fixtures

    path = tmp_path / "pair10.docx"
    path.write_bytes(frd_fixtures.build_f3_pair10())
    profile, contract = _frd_contract(path, config)
    assert profile.family == "F3"
    assert "4 requirement" in profile.notes[0] and "13 nfr" in profile.notes[0]
    assert "1 domain" in profile.notes[0]
    feed = contract.feeds[0]
    assert (feed.domain, feed.sub_domain) == ("Care Management", "Assessments")
    assert contract.field_provenance["feeds[0].domain"].label == "Domain"


def test_an_f3_document_is_filled_by_the_chain_and_asks_typed_questions(pair1_config,
                                                                         tmp_path):
    from acfc_shapes import frd as frd_fixtures
    from codegen.layout.resolve import resolve_pair

    shapes = REPO / "fixtures" / "acfc_shapes"
    path = tmp_path / "pair10.docx"
    path.write_bytes(frd_fixtures.build_f3_pair10())
    pair = resolve_pair(shapes / "sttm" / "pair_1_family_a.xlsx", path, pair1_config,
                        provider=None, cache_dirs=[], generated_date="2026-01-01",
                        vdd_path=shapes / "vdd" / "pair_1_v1_segments.xlsx")
    feed = pair.frd_contract.feeds[0]
    assert feed.stage_target.tables == ["vnd_p_accum_client"]            # the STTM band
    assert feed.domain == "Care Management"                              # the F3 domain table
    keys = {q.key: q for q in pair.questions}
    assert "feeds[0].domain" not in keys                                 # read, not asked
    assert keys["feeds[0].source_system"].kind == "text"                 # typed, answerable


def test_the_round1_f2_fixture_is_still_f2(config):
    profile, _contract = _frd_contract(
        REPO / "fixtures" / "acfc_shapes" / "frd" / "f2_pair_8.docx", config)
    assert profile.family == "F2" and len(profile.sections) == 8
