"""M12 — the four causes read out of the v0.7.2 all-pairs run
(``docs/acfc/RUN_v072_details.md`` on ``origin/acfc-runs``; read in place,
never copied).

1. A layout model's answer is judged on its PLACEMENT only: the response
   schema carries no provenance, extra keys are dropped, and the resolver
   stamps ``source="model"`` itself. Pairs 5, 7, 8, 9 and 10 lost every
   model-placed role to a literal_error on ``source`` / ``role_sources``.
2. The gate's ruff runs with the PINNED rule set, wherever the tree lives,
   and ignores EXE002 (a file mode); the runner notebook's unused noqa is gone.
3. When the document parser child cannot start, documents are parsed
   in-process — a selection never waits on a parser that did not start.
4. A VDD tie is a choice question, not "no VDD".
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest
from pydantic import ValidationError

from codegen.layout.frd_profile import FrdLayoutProfile
from codegen.layout.model import MockLayoutProvider, build_frd_request
from codegen.layout.profile import LayoutProfile
from codegen.layout.resolve import resolve_frd, resolve_workbook
from codegen.layout.response import (
    FrdLayoutResponse,
    LayoutResponse,
    frd_partial,
    frd_profile_from_response,
    layout_profile_from_response,
)

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PROFILES = REPO / "fixtures" / "layout_profiles"
MOCK = PROFILES / "mock"
PROVENANCE = {"source", "role_sources", "confidence", "fingerprint", "strategy",
              "field_sources", "family", "unresolved"}


# ============================================ item 1: the response schema


def _placement(profile_file: str) -> dict:
    return json.loads((PROFILES / profile_file).read_text(encoding="utf-8"))


def _role_keys(payload: dict) -> list[str]:
    return [f"{s['name']}/{b['layer']}/{r}" for s in payload["sheets"] for b in s["bands"]
            for r in b["roles"]]


@pytest.mark.parametrize("resolved", [1, 4, 8, 15])
def test_the_four_real_rejection_shapes_now_validate(resolved):
    """RUN_v072_details §4: one literal_error on ``source`` plus one per
    ``role_sources`` key, for 1 / 4 / 8 / 15 resolved roles. The same payload
    is (still) invalid as a LayoutProfile — the shape is the real one — and is
    ACCEPTED as a response, provenance stamped by us."""
    payload = _placement("sttm_pair_1_family_a.layout.json")
    keys = _role_keys(payload)
    assert len(keys) >= 15
    payload["source"] = "model_inference"
    payload["role_sources"] = dict.fromkeys(keys[:resolved], "header_match")

    with pytest.raises(ValidationError) as caught:
        LayoutProfile.model_validate(payload)
    errors = caught.value.errors()
    assert len(errors) == 1 + resolved
    assert {e["type"] for e in errors} == {"literal_error"}
    assert sorted(e["loc"][0] for e in errors) == ["role_sources"] * resolved + ["source"]

    profile = layout_profile_from_response(payload, "the-document-fingerprint")
    assert profile.source == "model" and profile.strategy == "model"
    assert profile.fingerprint == "the-document-fingerprint"     # never the answer's
    assert profile.role_sources == dict.fromkeys(keys, "model")
    assert profile.confidence == {}


class _Chatty:
    """A provider whose answers are correct placements that ALSO volunteer
    provenance — invented words for ``source`` / ``role_sources``, a
    confidence, a strategy, a wrong fingerprint."""

    name = "chatty"

    def __init__(self, inner: MockLayoutProvider) -> None:
        self._inner = inner
        self.requests = inner.requests

    def complete_layout(self, request: dict) -> dict:
        answer = dict(self._inner.complete_layout(request))
        answer.update(source="inferred_from_headers", strategy="llm",
                      fingerprint="0" * 64, confidence={"x": 7.0},
                      role_sources=dict.fromkeys(_role_keys(answer), "header match"))
        return answer

    def advise_layout(self, request: dict) -> dict:
        return self._inner.advise_layout(request)


def test_a_model_answer_that_volunteers_provenance_is_accepted_and_stamped(config, tmp_path):
    """Pair 8 needs the model. The chatty answer resolves EXACTLY like the
    plain one: no schema rejection, source="model", every role source one of
    ours."""
    empty = tmp_path / "empty"
    empty.mkdir()
    workbook = SHAPES / "sttm" / "pair_8_family_a.xlsx"
    plain, _ = resolve_workbook(workbook, config, provider=MockLayoutProvider([MOCK, PROFILES]),
                                cache_dirs=[empty], runtime_cache_dir=tmp_path / "rt_a")
    chatty, _ = resolve_workbook(workbook, config,
                                 provider=_Chatty(MockLayoutProvider([MOCK, PROFILES])),
                                 cache_dirs=[empty], runtime_cache_dir=tmp_path / "rt_b")
    assert chatty.provider_calls == 1 and chatty.rejections == []
    assert chatty.complete and chatty.profile.source == "model"
    assert set(chatty.profile.role_sources.values()) <= {"synonyms", "model"}
    assert chatty.profile == plain.profile
    assert not (tmp_path / "rt_b" / "rejections").exists()


def test_the_prompt_shows_and_asks_for_the_response_shape_only(config, tmp_path):
    """The partial profile the model completes and the schema it is judged
    against are the RESPONSE model — no provenance field in either."""
    provider = MockLayoutProvider([MOCK, PROFILES])
    empty = tmp_path / "empty"
    empty.mkdir()
    resolve_workbook(SHAPES / "sttm" / "pair_8_family_a.xlsx", config, provider=provider,
                     cache_dirs=[empty], runtime_cache_dir=tmp_path / "rt")
    (request,) = provider.requests
    assert request["response_schema"] == LayoutResponse.model_json_schema()
    assert not PROVENANCE & set(request["response_schema"]["properties"])
    assert not PROVENANCE & set(request["partial_profile"])
    assert request["partial_profile"]["sheets"]

    frd = FrdLayoutProfile.model_validate(_placement("frd_f1_pair_1.layout.json"))
    frd_request = build_frd_request("fp", [], frd_partial(frd), [], config)
    assert frd_request["response_schema"] == FrdLayoutResponse.model_json_schema()
    assert not PROVENANCE & set(frd_request["partial_profile"])
    assert not PROVENANCE & set(frd_request["response_schema"]["properties"])


def test_the_frd_response_has_the_same_fix(config, tmp_path):
    """An FRD answer with an invented source, field sources and family: invalid
    as an FrdLayoutProfile, accepted as a response — the family is what
    DISCOVERY found, the provenance ours."""
    payload = _placement("frd_f1_pair_1.layout.json")
    payload.update(source="inferred", family="F9",
                   field_sources=dict.fromkeys(payload["fields"], "label match"))
    with pytest.raises(ValidationError):
        FrdLayoutProfile.model_validate(payload)
    profile = frd_profile_from_response(payload, "fp", "F1")
    assert profile.family == "F1" and profile.source == "model"
    assert profile.field_sources == dict.fromkeys(payload["fields"], "model")
    assert set(profile.fields) == set(payload["fields"])

    from acfc_shapes import frd as frd_fixtures

    class _ChattyFrd:
        name = "chatty"
        requests: list = []

        def complete_layout(self, request):
            self.requests.append(request)
            return {"fields": {}, "source": "inferred", "family": "F9",
                    "field_sources": {"feeds[0].frequency": "label match"}}

    docx = tmp_path / "pair9.docx"
    docx.write_bytes(frd_fixtures.build_f3_pair9())
    provider = _ChattyFrd()
    empty = tmp_path / "empty"
    empty.mkdir()
    doc, _ = resolve_frd(docx, config, provider=provider, cache_dirs=[empty],
                         runtime_cache_dir=tmp_path / "rt")
    assert doc.provider_calls == 1 and provider.requests
    assert not any("schema validation" in r.reason for r in doc.rejections), doc.rejections


def test_a_placement_that_does_not_parse_is_still_rejected(config):
    """Only provenance became lenient: a wrong PLACEMENT still fails."""
    with pytest.raises(ValidationError):
        layout_profile_from_response({"sheets": [{"name": "S", "kind": "mapping",
                                                  "bands": [{"layer": "stage",
                                                             "col_start": 0,
                                                             "col_end": 3}]}]}, "fp")
    with pytest.raises(ValidationError):
        layout_profile_from_response({"sheets": [{"name": "S", "kind": "novel"}]}, "fp")


# ============================================ item 2: the gate's ruff


EMITTED_RUFF_TOML = """# GENERATED by codegen-data-engineer-agent — lint config for the generated
# pipeline, mirrored from the generator repo's own settings.
line-length = 100
target-version = "py311"

# The assembled notebook duplicates already-linted module code cell-by-cell;
# ruff would re-flag the notebook's by-design shared namespace (F821 spark /
# dbutils, cross-cell redefinitions), so it is excluded rather than suppressed.
extend-exclude = ["*.ipynb"]

[lint]
select = ["E", "F", "I", "UP", "B", "SIM"]

[lint.isort]
# Local modules of the generated package; conftest is imported by tests.
known-first-party = ["pipeline", "conftest"]
"""


def test_the_emitted_ruff_toml_is_unchanged_and_the_gate_pins_the_same_rules():
    """One definition (codegen.lint_rules) renders the emitted file — byte for
    byte what it was — and the gate's arguments."""
    from codegen import lint_rules
    from codegen.emit.emitter import _environment

    rendered = _environment().get_template("ruff.toml.j2").render(lint_rules.template_context())
    assert rendered == EMITTED_RUFF_TOML
    args = lint_rules.ruff_config_args()
    assert args[0] == "--isolated"
    assert 'lint.ignore = ["EXE002"]' in args
    assert 'lint.select = ["E", "F", "I", "UP", "B", "SIM"]' in args


@pytest.fixture(scope="module")
def runner_notebooks(pair1_config, pair1_spec, tmp_path_factory):
    """The pair-1 DML runner notebooks, generated OUTSIDE the checkout (a
    temp root, like a remote outputs role): no ruff config above them."""
    pytest.importorskip("sqlglot", reason="the acfc_prx framework path emits DML")
    from codegen import cli

    tmp = tmp_path_factory.mktemp("m12_runner")
    config = pair1_config.model_copy(update={"output": pair1_config.output.model_copy(update={
        "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")})})
    cli._generate_feed(pair1_spec, config, dry_run=True, skip_tests=True,
                       output_mode="framework", conventions_profile="acfc_prx",
                       iig_template="iig_v2")
    notebooks = sorted((tmp / "out" / pair1_spec.feed_slug / "framework")
                       .glob("Insert_scripts_config_table_*.py"))
    assert len(notebooks) == 3
    lint_dir = tmp / "lint_only"
    lint_dir.mkdir()
    copies = []
    for notebook in notebooks:
        copy = lint_dir / notebook.name
        copy.write_bytes(notebook.read_bytes())
        os.chmod(copy, 0o755)
        copies.append(copy)
    return lint_dir, copies


def test_runner_notebooks_pass_the_gate_ruff_at_mode_0o755(runner_notebooks):
    """The ACFC failure's own shape: the runner notebooks, executable, in a
    directory with no ruff config anywhere above it. Before M12 ruff fell back
    to its defaults there (RUF100 on the SLF001 noqa)."""
    from codegen.gate.preflight import _ruff_check

    lint_dir, copies = runner_notebooks
    assert not any((parent / name).exists() for parent in [lint_dir, *lint_dir.parents]
                   for name in ("ruff.toml", ".ruff.toml", "pyproject.toml"))
    if os.name != "nt":
        assert all(os.stat(c).st_mode & 0o777 == 0o755 for c in copies)
    check = _ruff_check(lint_dir)
    assert check.passed and not check.not_run, check.details


def test_runner_notebooks_carry_no_unused_noqa(runner_notebooks):
    """RUF100 on top of the pinned rules: every noqa names a rule that fires."""
    from codegen.lint_rules import ruff_config_args

    lint_dir, _copies = runner_notebooks
    result = subprocess.run(
        [sys.executable, "-m", "ruff", "check", "--no-cache", "--output-format", "json",
         *ruff_config_args(), "--extend-select", "RUF100", str(lint_dir)],
        capture_output=True, text=True, check=False)
    assert json.loads(result.stdout) == [], result.stdout


def test_the_gate_does_not_consult_a_config_it_finds(tmp_path):
    """A config lying around the tree (a user's, a parent directory's) no
    longer changes the verdict: the rules are the pinned ones."""
    from codegen.gate.preflight import _ruff_check

    (tmp_path / "ruff.toml").write_text('[lint]\nselect = ["ALL"]\n', encoding="utf-8")
    (tmp_path / "module.py").write_text('"""A clean module."""\n\nVALUE = 1\n',
                                        encoding="utf-8")
    assert _ruff_check(tmp_path).passed


# ============================================ item 3: the parser mode


def _never_ready(tag: str) -> list[str]:
    """A parser command that starts and never says ready (the serverless / App
    shape of RUN_v072_details §2). ``tag`` keeps each test's probe its own."""
    return [sys.executable, "-c", f"import time; time.sleep(3600)  # {tag}"]


@pytest.fixture
def ws(monkeypatch, tmp_path):
    pytest.importorskip("fastapi")
    from ui.backend import stores as ui_stores

    from test_m93_chooser import _Workspace

    workspace = _Workspace(monkeypatch, tmp_path)
    yield workspace
    workspace.release()
    ui_stores.reset_stores()


def _probe_budget(ws, monkeypatch, probe: float = 0.5, start: float = 60.0) -> None:
    config = ws.store.config
    monkeypatch.setattr(ws.store, "config", config.model_copy(update={
        "inputs": config.inputs.model_copy(update={"parser_probe_seconds": probe,
                                                   "parser_start_timeout_seconds": start})}))


def test_a_parser_that_never_starts_selects_in_process_within_budget(monkeypatch, ws, tmp_path):
    """Pair 3's trace: 'start parser' warned, classify waited out the budget.
    Now the probe answers `inprocess` within its short budget, the parse runs in
    a worker thread, and the selection completes — far inside the 60 s a child
    start would have been waited for."""
    from ui.backend import docindex

    command = _never_ready(tmp_path.name)
    monkeypatch.setattr(docindex, "worker_command", lambda config, base: command)
    _probe_budget(ws, monkeypatch)
    ws.pairs("pair_1")
    runner = ws.runner()
    started = time.monotonic()
    runner.start_selection("STTM_alpha.xlsx")
    job = runner.wait_selection(45.0)
    took = time.monotonic() - started
    assert job["state"] == "done", job
    assert took < 20, f"the selection took {took:.1f}s"
    steps = {s["step"]: s for s in job["steps"]}
    assert steps["start parser"]["state"] == "done"
    assert "in-process" in steps["start parser"]["detail"]
    assert steps["classify"]["state"] == "done"
    assert steps["classify"]["detail"].startswith("sttm:")
    status = runner.status()
    assert status["selection"]["sttm"] == "STTM_alpha.xlsx"
    assert status["parser_mode"]["mode"] == "inprocess"
    assert "did not start" in status["parser_mode"]["reason"]


def _bloated_pair1(tmp_path: Path) -> Path:
    """Pair 1's STTM with one formatted cell a million rows down: the sheet's
    dimension claims 1,000,000 rows over ~100 real ones (SHAPES_ROUND2 §3)."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    workbook = load_workbook(SHAPES / "sttm" / "pair_1_family_a.xlsx")
    workbook["FEED_1_MAPPING"].cell(row=1_000_000, column=30).font = Font(bold=True)
    path = tmp_path / "bloated.xlsx"
    workbook.save(path)
    return path


class _Doc:
    def __init__(self, path: Path) -> None:
        self.name, self.uri, self.version = path.name, str(path), f"{path}#1:1"


def test_a_workbook_claiming_1m_rows_classifies_in_process_under_2s(monkeypatch, config,
                                                                    tmp_path):
    """In-process parsing reads through the used-range loader and the streamed
    cell cap — a declared million rows costs nothing."""
    from ui.backend import docindex

    import codegen.layout.docworker  # noqa: F401 — warm imports; the read is timed
    import codegen.pairing  # noqa: F401

    command = _never_ready(tmp_path.name)
    monkeypatch.setattr(docindex, "worker_command", lambda cfg, base: command)
    probe_config = config.model_copy(update={"inputs": config.inputs.model_copy(update={
        "parser_probe_seconds": 0.5})})
    index = docindex.DocumentIndex(lambda: probe_config, lambda: tmp_path / "index.json",
                                   lambda: None, REPO)
    assert index.parser_mode(wait=10.0)["mode"] == docindex.PARSER_INPROCESS
    workbook = _bloated_pair1(tmp_path)
    started = time.monotonic()
    verdict = index.read_now(_Doc(workbook), workbook, 30.0)
    took = time.monotonic() - started
    assert verdict["state"] == "sttm", verdict
    assert verdict["facts"] is not None
    assert took < 2.0, f"in-process classify took {took:.2f}s"
    assert index.request_parses == ["bloated.xlsx"]


def test_a_late_in_process_read_is_abandoned_and_holds_its_lane(monkeypatch, config, tmp_path):
    """A thread cannot be killed: a read past its budget is `timed_out`, and
    the next read of that lane waits for it (within its own budget) instead of
    stacking a second runaway parse."""
    import threading

    from ui.backend import docindex

    from codegen.layout import docworker

    release = threading.Event()

    def slow(local, name, cfg, base):
        release.wait(30)
        return {"state": "sttm", "reason": "slow", "facts": None}

    monkeypatch.setattr(docworker, "read_document", slow)
    index = docindex.DocumentIndex(lambda: config, lambda: tmp_path / "index.json",
                                   lambda: None, REPO)
    doc = _Doc(tmp_path / "a.xlsx")
    try:
        first = index._parse_inprocess(doc, tmp_path / "a.xlsx", 0.3, "request")
        assert first["timed_out"] and "abandoned" in first["reason"]
        second = index._parse_inprocess(doc, tmp_path / "a.xlsx", 0.3, "request")
        assert second["timed_out"] and "still running" in second["reason"]
    finally:
        release.set()
    time.sleep(0.2)
    monkeypatch.setattr(docworker, "read_document",
                        lambda *a: {"state": "vdd", "reason": "ok", "facts": None})
    assert index._parse_inprocess(doc, tmp_path / "a.xlsx", 5.0, "request")["state"] == "vdd"


def test_a_child_that_stops_starting_demotes_the_process(monkeypatch, config, tmp_path):
    """Subprocess mode, then a start that fails: that read finishes in-process
    and the mode becomes inprocess — never a second wait on the child."""
    from ui.backend import docindex

    from codegen.layout import docworker

    index = docindex.DocumentIndex(lambda: config, lambda: tmp_path / "index.json",
                                   lambda: None, REPO)
    command = [sys.executable, "-c", f"print('stub')  # {tmp_path.name}"]
    monkeypatch.setattr(docindex, "worker_command", lambda cfg, base: command)
    key = (tuple(command), str(REPO))
    import threading

    done = threading.Event()
    done.set()
    monkeypatch.setitem(docindex._modes, key, {"mode": docindex.PARSER_SUBPROCESS,
                                               "reason": "", "done": done})

    class _Refuses:
        def parse(self, *a, **k):
            raise docindex.ParserStartError("the document parser process did not start")

    monkeypatch.setattr(index, "_parser", lambda lane: _Refuses())
    monkeypatch.setattr(docworker, "read_document",
                        lambda *a: {"state": "sttm", "reason": "in-process", "facts": None})
    result = index._parse(_Doc(tmp_path / "a.xlsx"), tmp_path / "a.xlsx", 5.0, "request")
    assert result["state"] == "sttm"
    assert index.parser_mode()["mode"] == docindex.PARSER_INPROCESS
    assert "from now on" in index.parser_mode()["reason"]


# ============================================ item 4: a VDD tie is a question


def _tied(config, *, scores=(1.0, 1.0, 1.0)):
    from codegen.pairing import PairCandidate, PairSignal, decide

    signal = PairSignal("meta_frequency", 1.0, "'Monthly' at meta row 'Frequency' ↔ FILES")
    candidates = [PairCandidate(f"VDD_{chr(97 + i)}.xlsx", score, (signal,) if score else ())
                  for i, score in enumerate(scores)]
    return decide("vdd", "STTM_pair_5.xlsx", candidates, config)


def test_a_vdd_tie_is_a_choice_question_not_no_vdd(config):
    """Pair 5: three dictionaries at 1.0 each, below min_score — nothing can
    separate them, so a person is asked (it used to read "no VDD")."""
    decision = _tied(config)
    assert decision.chosen is None and decision.ambiguous
    question = decision.question()
    assert question["key"] == "pair.vdd" and question["kind"] == "choice"
    assert [c["value"] for c in question["candidates"]] == [
        "VDD_a.xlsx", "VDD_b.xlsx", "VDD_c.xlsx"]


def test_a_weak_vdd_without_a_tie_still_pairs_nothing_and_asks_nothing(config):
    """Unchanged: a VDD is optional, one weak candidate is not a question."""
    decision = _tied(config, scores=(1.0, 0.0, 0.0))
    assert decision.chosen is None and not decision.ambiguous


def test_a_vdd_tie_reaches_the_status_as_a_question(monkeypatch, ws, config):
    """The runner surfaces the tie (pairing.vdd.question, pair_candidates)."""
    from codegen import pairing

    real = pairing.pair_by_content

    def tie_for_vdd(kind, sttm_path, candidates, cfg, base, **kw):
        if kind == "vdd":
            return _tied(cfg)
        return real(kind, sttm_path, candidates, cfg, base, **kw)

    monkeypatch.setattr(pairing, "pair_by_content", tie_for_vdd)
    ws.pairs("pair_1")
    runner = ws.runner()
    runner.start_selection("STTM_alpha.xlsx")
    job = runner.wait_selection(60.0)
    assert job["state"] == "done", job
    status = runner.status()
    assert status["pairing"]["vdd"]["question"]["key"] == "pair.vdd"
    assert len(status["pair_candidates"]["vdd"]["candidates"]) == 3
    assert status["selection"]["vdd"] is None
