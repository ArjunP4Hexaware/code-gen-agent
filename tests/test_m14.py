"""M14 — question quality and the last two blockers, from the v0.7.3 all-pairs
run (``docs/acfc/RUN_v073_all_pairs.md`` on ``origin/acfc-runs`` — read in
place; the candidate TEXTS below are the structural values of its question
payloads, not client terms).

1. Candidate texts are normalized (config ``value_vocabulary``) before the gap
   chain decides the documents disagree: only canonical values that genuinely
   differ are a question; transport phrases, date schedules and vague cadences
   are dropped / yield, flagged.
2. F2 / F3 content questions are TEXT questions with the document's own
   sentences as evidence — never a picker over row labels.
3. The same question for several feeds is asked once (``feeds[*]``), a
   per-feed answer overrides it.
4. Role questions always carry candidates (column letter + header).
5. Identical flags appear once, with a count.
6. The verdict line names every failed / not-run check with its first finding.
7. The parser SUBPROCESS reads through the used-range loader, on the same
   per-file budget as in-process.
"""

from __future__ import annotations

import re
import shutil
import sys
import time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from codegen.config import load_config
from codegen.gate.preflight import GateCheck
from codegen.gate.verdict import GateResult, dedupe_flags
from codegen.layout.profile import BandProfile, LayoutProfile, SheetProfile, UnresolvedRole
from codegen.layout.resolve import (
    LayoutQuestion,
    _one_question_for_all_feeds,
    _questions_for,
    _typed,
    expand_all_feeds,
    fill_frd_gaps,
)
from codegen.report.generation_report import verdict_line
from codegen.resolve.gapfill import Statement, reconcile

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PROFILES = REPO / "fixtures" / "layout_profiles"
DATE = "2026-01-01"
SCHEDULE = "1) 03/20/2026 \n2) 03/31/2026 \n3) 04/02/2026 \n4)  04/08/2026 \n5) 04/15/2026"


@pytest.fixture(scope="module")
def vocabulary():
    return load_config(REPO / "config" / "config.yaml").value_vocabulary


# ============================================ 1. normalize, then compare


# (pair, field, the two candidates as the run asked them, what is left, dropped)
RUN_V073_CANDIDATES = [
    ("pair 2", "file_format", ["File Data Ingestion", "CSV"], ["CSV"], ["File Data Ingestion"]),
    ("pair 4", "file_format", [".xlsx", "Fixed-width"], [".xlsx", "Fixed-width"], []),
    ("pair 5", "file_format", [".txt files", "Delimited"], ["Delimited"], []),
    ("pair 6", "file_format", ["SQL Server Table to  Databricks Table", "Database Table"],
     ["Database Table"], ["SQL Server Table to  Databricks Table"]),
    ("pair 7", "file_format", ["Xlsx files", "Fixed-width"], ["Xlsx files", "Fixed-width"], []),
    ("pair 7", "frequency", [SCHEDULE, "Historical (one-time)"], ["Historical (one-time)"],
     [SCHEDULE]),
    ("pair 10", "file_format", ["CSV", "Delimited"], ["CSV"], []),
    ("pair 10", "delimiter", ["Comma", "Pipe (|)"], ["Comma", "Pipe (|)"], []),
    ("pair 10", "frequency", ["Periodic", "Weekly"], ["Weekly"], ["Periodic"]),
]


@pytest.mark.parametrize(("pair", "field", "candidates", "kept", "dropped"), RUN_V073_CANDIDATES,
                         ids=[f"{c[0]}-{c[1]}-{i}" for i, c in enumerate(RUN_V073_CANDIDATES)])
def test_the_real_candidate_pairs_reduce_to_genuine_disagreements(vocabulary, pair, field,
                                                                   candidates, kept, dropped):
    statements = [Statement(value, source, f"{source} cell")
                  for value, source in zip(candidates, ("STTM", "VDD"), strict=True)]
    outcome = reconcile(field, statements, vocabulary)
    assert [s.value for s in outcome.kept] == kept, pair
    assert [s.value for s, _why in outcome.dropped] == dropped, pair
    for _statement, why in outcome.dropped:
        assert why                                     # every drop says why


def test_same_meaning_never_asks_and_the_frd_value_stays(vocabulary):
    for field, frd, other in (("file_format", "Fixed Width", ".dat"),
                              ("file_format", "Delimited", "CSV"),
                              ("file_format", "Fixed Width Text File", "Fixed-width"),
                              ("delimiter", "|", "Pipe (|)"),
                              ("delimiter", ",", "Comma"),
                              ("frequency", "Weekly", "Periodic")):
        outcome = reconcile(field, [Statement(frd, "FRD", "c"), Statement(other, "STTM", "c")],
                            vocabulary)
        assert [s.value for s in outcome.kept] == [frd], (field, frd, other)


def test_unknown_texts_still_compare_as_before(vocabulary):
    """A text the vocabulary does not know falls back to the old agreement
    rule (equal, or one contains the other)."""
    same = reconcile("file_format", [Statement("Proprietary blob", "FRD", "c"),
                                     Statement("proprietary blob", "STTM", "c")], vocabulary)
    differ = reconcile("file_format", [Statement("Proprietary blob", "FRD", "c"),
                                       Statement("Avro", "STTM", "c")], vocabulary)
    assert len(same.kept) == 1 and len(differ.kept) == 2


def test_a_transport_phrase_in_the_frd_is_dropped_and_the_sttm_fills(tmp_path):
    """Pair 2's shape end to end through the layout stage: the FRD's format
    cell says how the data moves; the STTM states the format — no question."""
    from acfc_shapes import pair1
    from test_frd_gapfill import _pair1_docx_without, _resolve, _sttm_with_load_strategy

    sttm, cache = _sttm_with_load_strategy(tmp_path, "STG: Append; STD: Append")
    pair = _resolve(tmp_path, _pair1_docx_without({"Object/data Format": "File Data Ingestion"}),
                    sttm, cache)
    assert "feeds[0].file_format" not in [q.key for q in pair.questions]
    assert pair.frd_contract.feeds[0].file_format == pair1.FILE_FORMAT
    (flag,) = [f for f in pair.flags if f.startswith("candidate_dropped:feeds[0].file_format")]
    assert "'File Data Ingestion' (FRD table" in flag and "not a file format" in flag


# ============================================ 2. F2 / F3: text + evidence


@pytest.fixture(scope="module")
def pair9_content(tmp_path_factory):
    from acfc_shapes import frd as frd_fixtures
    from codegen.extract.frd_docx import read_docx

    path = tmp_path_factory.mktemp("m14_f3") / "pair9.docx"
    path.write_bytes(frd_fixtures.build_f3_pair9())
    return path, read_docx(path)


def test_an_f2_row_label_picker_becomes_a_text_question_with_evidence(pair9_content):
    config = load_config(REPO / "config" / "config.yaml")
    _path, content = pair9_content
    picker = LayoutQuestion(
        document="frd", sheet=None, layer=None, role="feeds[0].source_system", kind="role",
        reason="no label matched", title="Source system / vendor",
        candidates=[{"table": 1, "row": 1, "col": 0, "label": "Business Requirement"},
                    {"table": 2, "row": 1, "col": 0, "label": "Impact Details"}])
    typed = _typed(picker, content, config, "F2")
    assert typed.kind == "text" and typed.candidates == []
    assert "frd_family_f2" in typed.reason and typed.evidence
    cells = {(t, r, c): " ".join((content.tables[t][r][c] or "").split())
             for t, rows in enumerate(content.tables) for r, row in enumerate(rows)
             for c, _cell in enumerate(row)}
    for item in typed.evidence:                         # verbatim, with a real cell
        t, r, c = map(int, re.findall(r"\d+", item["cell"]))
        assert item["text"].rstrip("…") in cells[(t, r, c)]
        assert "source system" in item["text"].lower() or "vendor" in item["text"].lower()
    texts = [e["text"].lower().rstrip(" .") for e in typed.evidence]
    assert len(texts) == len(set(texts))                # no near-duplicates
    assert typed.as_dict()["evidence"] == typed.evidence


def test_a_prose_frd_asks_no_role_question(pair9_content, tmp_path):
    from codegen.layout.model import MockLayoutProvider
    from codegen.layout.resolve import resolve_pair

    path, _content = pair9_content
    pair = resolve_pair(SHAPES / "sttm" / "pair_9_family_e.xlsx", path,
                        load_config(REPO / "config" / "config.yaml"),
                        provider=MockLayoutProvider([PROFILES / "mock", PROFILES]),
                        cache_dirs=[tmp_path / "cache"], generated_date=DATE)
    frd_questions = [q for q in pair.questions if q.document == "frd"]
    assert frd_questions and all(q.kind != "role" for q in frd_questions)
    source = next(q for q in frd_questions if q.key == "feeds[0].source_system")
    assert source.kind == "text" and source.evidence


# ============================================ 3. once for all feeds


def _choice(index: int, field: str = "delimiter") -> LayoutQuestion:
    return LayoutQuestion(document="frd", sheet=None, layer=None, role=f"feeds[{index}].{field}",
                          kind="choice", reason="the documents disagree", title="Delimiter",
                          candidates=[{"value": "Comma", "source": "STTM", "cell": f"r{index}"},
                                      {"value": "Pipe (|)", "source": "VDD", "cell": f"v{index}"}])


def test_the_same_question_for_four_feeds_is_asked_once():
    """Pair 10: Comma vs Pipe for feeds 0-3 — one question, keyed feeds[*]."""
    questions = [_choice(i) for i in range(4)] + [
        LayoutQuestion(document="frd", sheet=None, layer=None, role=f"feeds[{i}].file_patterns",
                       kind="text", reason="none", title="File name patterns")
        for i in range(2)]
    collapsed = _one_question_for_all_feeds(questions)
    keys = [q.key for q in collapsed]
    assert keys == ["feeds[*].delimiter", "feeds[0].file_patterns", "feeds[1].file_patterns"]
    (once,) = [q for q in collapsed if q.key == "feeds[*].delimiter"]
    assert once.feeds == [0, 1, 2, 3] and "overrides" in once.reason
    assert once.as_dict()["feeds"] == [0, 1, 2, 3]
    # different candidates are different questions
    odd = _choice(9).candidates[:1]
    mixed = _one_question_for_all_feeds([_choice(0), LayoutQuestion(
        document="frd", sheet=None, layer=None, role="feeds[1].delimiter", kind="choice",
        reason="x", title="Delimiter", candidates=odd)])
    assert [q.key for q in mixed] == ["feeds[0].delimiter", "feeds[1].delimiter"]


def test_a_feeds_star_answer_applies_to_all_and_a_per_feed_answer_wins():
    gaps = {"feeds[*].delimiter": {"value": "Comma", "source": "STTM"},
            "feeds[2].delimiter": {"value": "Pipe (|)", "source": "VDD"}}
    out = expand_all_feeds(gaps, 4)
    assert [out[f"feeds[{i}].delimiter"]["value"] for i in range(4)] == [
        "Comma", "Comma", "Pipe (|)", "Comma"]


def test_the_answers_file_accepts_feeds_star_through_the_gap_chain(tmp_path):
    """End to end on a two-feed contract (pair 9's F3 FRD splits in two): a
    feeds[*] answer fills every feed, the per-feed one overrides."""
    from acfc_shapes import frd as frd_fixtures
    from codegen.layout.model import MockLayoutProvider
    from codegen.layout.resolve import parse_answers, resolve_pair

    docx = tmp_path / "pair9.docx"
    docx.write_bytes(frd_fixtures.build_f3_pair9())
    config = load_config(REPO / "config" / "config.yaml")
    base = resolve_pair(SHAPES / "sttm" / "pair_9_family_e.xlsx", docx, config,
                        provider=MockLayoutProvider([PROFILES / "mock", PROFILES]),
                        cache_dirs=[tmp_path / "c"], generated_date=DATE).frd_contract
    count = len(base.feeds)
    assert count >= 2, "the override needs a feed besides the one it overrides"
    gaps = parse_answers({"gaps": {
        "feeds[*].stage_target.load_strategy": {"value": "Append"},
        f"feeds[{count - 1}].stage_target.load_strategy": {"value": "Truncate and Load"}}})
    result = fill_frd_gaps(base, None, {}, {}, gaps["gaps"], config, REPO)
    got = [f.stage_target.load_strategy for f in result.contract.feeds]
    assert got == ["Append"] * (count - 1) + ["Truncate and Load"]


# ============================================ 4. role candidates, always


def test_a_role_question_never_ships_an_empty_candidate_list():
    """Pair 5's shape: the band's own columns are all claimed (or blank) —
    the question still offers every headed column, each with its letter."""
    wb = Workbook()
    ws = wb.active
    ws.title = "MAP"
    ws.append(["", "", "", ""])
    ws.append(["Stage Table", "Stage Column", "Target Schema", "Notes"])
    profile = LayoutProfile(
        fingerprint="fp", source="synonyms", strategy="content",
        sheets=[SheetProfile(name="MAP", kind="mapping", header_row=2, band_row=1, bands=[
            BandProfile(layer="stage", col_start=1, col_end=2,
                        roles={"table": 1, "column": 2})])],
        unresolved=[UnresolvedRole(sheet="MAP", layer="stage", role="schema",
                                   reason="required role has no header")])
    (question,) = _questions_for(profile, wb)
    assert question.candidates == [{"col": 3, "letter": "C", "header": "Target Schema"},
                                   {"col": 4, "letter": "D", "header": "Notes"}]
    assert "every headed column" in question.reason


# ============================================ 5. + 6. flags and the verdict line


def test_identical_flags_appear_once_with_a_count():
    assert dedupe_flags(["a", "b", "a", "c", "a", "b"]) == ["a (×3)", "b (×2)", "c"]
    assert dedupe_flags(["x", "y"]) == ["x", "y"]         # distinct: untouched


def test_the_verdict_line_names_every_failed_check_with_its_first_finding():
    ruff = "2 finding(s)\npipeline/a.py:1:1: F401 unused\npipeline/b.py:2:1: E501"
    checks = [GateCheck(name="ruff", passed=False, details=ruff),
              GateCheck(name="secrets", passed=True, details="no hardcoded secrets"),
              GateCheck(name="tests", passed=True, not_run=True, details="pytest not found")]
    line = verdict_line(GateResult(feed_id="f", verdict="FAIL", checks=checks, flags=[]))
    assert line == ("**Verdict: FAIL** — ruff: pipeline/a.py:1:1: F401 unused; "
                    "tests (not run): pytest not found")
    passing = GateResult(feed_id="f", verdict="PASS_WITH_FLAGS", checks=checks[1:2], flags=["x"])
    assert verdict_line(passing) == "**Verdict: PASS_WITH_FLAGS**"      # as before M14


# ============================================ 7. the subprocess on the used range


def _bloated(tmp_path: Path) -> Path:
    path = tmp_path / "bloated.xlsx"
    shutil.copyfile(SHAPES / "sttm" / "pair_1_family_a.xlsx", path)
    wb = load_workbook(path)
    wb["FEED_1_MAPPING"].cell(row=1_000_000, column=30).font = Font(bold=True)
    wb.save(path)
    return path


def test_the_parser_subprocess_reads_a_1m_row_dimension_fast(config, tmp_path):
    pytest.importorskip("fastapi")
    from ui.backend import docindex

    parser = docindex.ParserProcess(docindex.worker_command(config, REPO), REPO, 90.0)
    try:
        parser.ensure()                                 # the start is not the file's cost
        workbook = _bloated(tmp_path)
        started = time.monotonic()
        verdict = parser.parse(workbook, workbook.name, 60.0)
        took = time.monotonic() - started
    finally:
        parser.stop()
    assert verdict["state"] == "sttm" and verdict["facts"] is not None, verdict
    assert took < 10.0, f"the subprocess took {took:.1f}s on a 1M-row dimension"


def test_subprocess_and_in_process_reads_get_the_same_budget(monkeypatch, config, tmp_path):
    pytest.importorskip("fastapi")
    import threading

    from ui.backend import docindex

    index = docindex.DocumentIndex(lambda: config, lambda: tmp_path / "i.json", lambda: None,
                                   REPO)
    budgets: dict[str, float] = {}
    done = threading.Event()
    done.set()

    class _Recorder:
        def parse(self, local, name, timeout):
            budgets["subprocess"] = timeout
            return {"state": "sttm", "reason": "", "facts": None}

    command = [sys.executable, "-c", f"pass  # {tmp_path.name}"]
    monkeypatch.setattr(docindex, "worker_command", lambda cfg, base: command)
    monkeypatch.setitem(docindex._modes, (tuple(command), str(REPO)),
                        {"mode": docindex.PARSER_SUBPROCESS, "reason": "", "done": done})
    monkeypatch.setattr(index, "_parser", lambda lane: _Recorder())

    class _Doc:
        name = uri = "a.xlsx"

    index._parse(_Doc(), tmp_path / "a.xlsx", 7.0, "request")
    monkeypatch.setitem(docindex._modes, (tuple(command), str(REPO)),
                        {"mode": docindex.PARSER_INPROCESS, "reason": "", "done": done})
    seen: dict[str, float] = {}

    def spy(doc, local, timeout, lane):
        seen["inprocess"] = timeout
        return {"state": "sttm", "reason": "", "facts": None}

    monkeypatch.setattr(index, "_parse_inprocess", spy)
    index._parse(_Doc(), tmp_path / "a.xlsx", 7.0, "request")
    assert 6.5 < budgets["subprocess"] <= 7.0 and 6.5 < seen["inprocess"] <= 7.0
