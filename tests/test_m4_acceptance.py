"""M4 acceptance — conventions profiles, IIG template versions, drag-fill.

Pinned against the aliased pair-1 golden (docs/acfc/rfc_capture goldens,
tracked copy at fixtures/acfc_shapes/pair_1/golden/):

* the pair-1 run under ``acfc_prx`` + ``iig_v2`` writes ``ACCUM_DDL.txt``
  byte-identical to the golden — INCLUDING the Decimal(17,2)→(22,2) run,
  which the gate flags (``drag_fill_suspect``) but never alters;
* the IIG workbook matches the golden's sheet names, order, headers, column
  order and row counts, the cells the fixture universe determines match the
  golden's values, and every other cell is blank AND flagged (pinned list);
* a docx FRD that names no file pattern / segments takes both from the STTM
  with provenance flags;
* the default profile/template (``edo_sfmc`` / ``iig_v1``) is untouched —
  the CV / SFMC baselines are guarded by the notebook-mode snapshot and the
  Option B suite; here the default framework emit still writes the two
  reference-style files and no combined file.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from acfc_shapes.pair1 import alias_golden_iig
from codegen import cli
from codegen.extract import contract_to_json as sttm_to_json
from codegen.extract import extract_contract
from codegen.extract.frd_docx import contract_to_json as frd_to_json
from codegen.extract.vdd import contract_to_json as vdd_to_json
from codegen.extract.vdd import extract_vdd_contract
from codegen.gate.drag_fill import drag_fill_flags
from codegen.layout.model import MockLayoutProvider
from codegen.layout.resolve import resolve_pair
from codegen.resolve.resolver import resolve_pair as resolve_contracts

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PROFILES = REPO / "fixtures" / "layout_profiles"
GOLDEN_DDL = SHAPES / "pair_1" / "golden" / "ACCUM_DDL.txt"
GOLDEN_IIG = SHAPES / "pair_1" / "golden" / "RFC_ACCUMULATORS_IIG.xlsx"
DATE = "2026-01-01"

needs_golden = pytest.mark.skipif(
    not (GOLDEN_DDL.is_file() and GOLDEN_IIG.is_file()),
    reason="pair-1 golden not present (the IIG is untracked on staging)",
)


# -- pair-1 run ----------------------------------------------------------------- #


@pytest.fixture(scope="module")
def pair1_spec(config, tmp_path_factory):
    """Pair 1 end to end the M4 way: mock layout, STTM + FRD(docx) + VDD
    contracts, resolver takes file pattern + segments from the STTM."""
    tmp = tmp_path_factory.mktemp("pair1_contracts")
    pair = resolve_pair(SHAPES / "sttm" / "pair_1_family_a.xlsx",
                        SHAPES / "frd" / "f1_pair_1.docx", config,
                        provider=MockLayoutProvider([PROFILES / "mock", PROFILES]),
                        cache_dirs=[tmp / "cache"], generated_date=DATE,
                        vdd_path=SHAPES / "vdd" / "pair_1_v1_segments.xlsx")
    frd_json = tmp / "frd.contract.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    assert json.loads(frd_json.read_text(encoding="utf-8"))["feeds"][0]["file_name_patterns"] == []
    sttm_json = tmp / "sttm.contract.json"
    contract = extract_contract(SHAPES / "sttm" / "pair_1_family_a.xlsx", frd_json, config,
                                generated_date=DATE, layout=pair.sttm.profile)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    vdd_json = tmp / "vdd.contract.json"
    vdd_contract, _ = extract_vdd_contract(SHAPES / "vdd" / "pair_1_v1_segments.xlsx", config,
                                           generated_date=DATE)
    vdd_json.write_text(vdd_to_json(vdd_contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config, vdd_path=vdd_json)
    return spec


def _scoped(config, tmp: Path):
    return config.model_copy(update={
        "output": config.output.model_copy(update={
            "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")}),
    })


@pytest.fixture(scope="module")
def pair1_run(config, pair1_spec, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pair1_run")
    scoped = _scoped(config, tmp)
    gate = cli._generate_feed(pair1_spec, scoped, dry_run=True, skip_tests=True,
                              output_mode="framework", conventions_profile="acfc_prx",
                              iig_template="iig_v2")
    return gate, tmp / "out" / pair1_spec.feed_slug / "framework"


def _golden_iig():
    return load_workbook(io.BytesIO(alias_golden_iig(GOLDEN_IIG)))


def _sheet_rows(sheet) -> tuple[list[str], list[list[str]]]:
    rows = list(sheet.iter_rows(values_only=True))
    headers = [c for c in rows[0] if c is not None]
    body = [["" if c is None else str(c).strip() for c in r[:len(headers)]] for r in rows[1:]]
    return headers, body


# -- resolver provenance -------------------------------------------------------- #


def test_pair1_resolver_takes_pattern_and_segments_from_the_sttm_flagged(pair1_spec):
    assert pair1_spec.file_name_patterns == [
        "I_ACCUM_*_TO_CLIENT_*.csv", "I_ACCUM_*_FROM_CLIENT_*.csv",
        "F_ACCUM_*_TO_CLIENT_*.csv", "F_ACCUM_*_FROM_CLIENT_*.csv"]
    assert [s.segment for s in pair1_spec.segments] == ["Header", "Detail", "Trailer"]
    assert pair1_spec.delimiter == ""  # fixed width: no implied delimiter
    kinds = sorted(f.split(":")[0] for f in pair1_spec.provenance_flags)
    assert kinds == ["file_pattern_from_sttm", "segments_from_sttm"]
    # Facts, not guesses: the flags quote where the values came from.
    assert "meta row 'File Names'" in pair1_spec.provenance_flags[1]
    assert "sheet 'FEED_1_MAPPING'" in pair1_spec.provenance_flags[0]
    # Not-null / natural key are the DETAIL segment's, deduped.
    assert pair1_spec.natural_key_columns == ["SEGMENT_IDENTIFIER", "CARDHOLDER_ID", "PLAN_ID"]


# -- DDL ------------------------------------------------------------------------- #


@needs_golden
def test_pair1_combined_ddl_is_byte_identical_to_the_golden(pair1_run):
    gate, framework_dir = pair1_run
    ours = (framework_dir / "ACCUM_DDL.txt").read_bytes()
    golden = GOLDEN_DDL.read_bytes()
    assert ours == golden
    # The golden's whitespace quirks are reproduced, not normalized.
    assert golden.startswith(b"\n--stage table\n\n")
    assert b")\n \nUSING delta" in golden and not golden.endswith(b"\n")
    # The 17,2 -> 22,2 run is transcribed as the STTM states it ...
    assert b"INDIVIDUAL_DEDUCTIBLE Decimal(17,2)" in ours
    assert b"FAMILY_LIMIT Decimal(22,2)" in ours
    # ... and no two-file output exists under this profile.
    assert not list(framework_dir.glob("*_table_creation.txt"))


@needs_golden
def test_pair1_gate_flags_drag_fill_and_provenance(pair1_run):
    gate, _ = pair1_run
    assert gate.verdict == "PASS_WITH_FLAGS"
    drag = [f for f in gate.flags if f.startswith("drag_fill_suspect:")]
    assert len(drag) == 1
    assert drag[0].startswith(
        "drag_fill_suspect:stage — vnd_p_accum_client columns INDIVIDUAL_DEDUCTIBLE, "
        "FAMILY_DEDUCTIBLE, COPAY, COINSURANCE, INDIVIDUAL_LIMIT, FAMILY_LIMIT: "
        "Decimal(17,2) → Decimal(22,2)")
    assert "(STTM rows 27–32)" in drag[0]
    assert any(f.startswith("file_pattern_from_sttm:") for f in gate.flags)
    assert any(f.startswith("segments_from_sttm:") for f in gate.flags)
    assert not any(f.startswith("ddl_file_name_from_slug") for f in gate.flags)
    assert not any(f.startswith("vdd_") for f in gate.flags)


def test_drag_fill_needs_three_adjacent_steps(pair1_spec):
    """Breaking the run at COPAY leaves [17,18] (no flag) and [20,21,22]."""
    detail = pair1_spec.detail_segment
    fields = [f.model_copy(update={"stage_datatype": "Decimal(19,4)",
                                   "standard_datatype": "Decimal(19,4)"})
              if f.stage_column == "COPAY" else f for f in detail.fields]
    segments = [detail.model_copy(update={"fields": fields}) if s is detail else s
                for s in pair1_spec.segments]
    spec = pair1_spec.model_copy(update={"segments": segments})
    (flag,) = drag_fill_flags(spec)
    assert "columns COINSURANCE, INDIVIDUAL_LIMIT, FAMILY_LIMIT: Decimal(20,2) → Decimal(22,2)" \
        in flag
    # Standard layer identical to stage -> reported once, for the stage.
    assert flag.startswith("drag_fill_suspect:stage")


# -- IIG workbook ---------------------------------------------------------------- #

# Cells the fixture universe determines: compared to the golden verbatim.
PINNED_COLUMNS = {
    "DATA_FACTORY_PIPELINE_SCHEDULE": [
        "PIPELINE_NAME", "PIPELINE_DESCRIPTION", "APPLICATION_NAME", "NO_OF_CYCLE_PER_DAY",
        "DAY_OF_SCHEDULE", "ACTIVE_FLAG", "ACTIVE_END_DATE"],
    "FILE_ADLS_INGESTION_DETAILS": ["DOMAIN", "SUBDOMAIN", "TGT_ADLS_PATH", "SOURCE_TYPE"],
    "ADLS_DELTA_INGESTION_DETAILS": [
        "OBJECT_NAME", "DOMAIN", "SUBDOMAIN", "SRC_ADLS_PATH", "SRC_FILE_NAME", "SRC_FORMAT",
        "SRC_COLUMNS", "SRC_DATA_TYPE", "SRC_ADLS_ARCHVL_PATH", "TGT_DATABASE_NAME",
        "TGT_TABLE_NAME", "TGT_ADLS_PATH", "TGT_COLUMN_NAMES", "TGT_FORMAT", "TGT_DATA_TYPE",
        "TGT_RJT_TABLE_NAME", "TGT_RJT_ADLS_PATH", "RECYCL_ENBL_FLG", "ACTIVE_FLAG"],
    "STGDELTA_STDDELTA_INGESTION_DET": [
        "DOMAIN", "SUBDOMAIN", "FREQUENCY", "SRC_ADLS_PATH", "SRC_TABLE_NAME",
        "SRC_CATALOG_NAME", "SRC_SCHEMA_NAME", "SRC_FORMAT", "SRC_COLUMNS", "SRC_DATA_TYPE",
        "SRC_ADLS_ARCHVL_PATH", "TGT_CATALOG_NAME", "TGT_SCHEMA_NAME", "TGT_TABLE_NAME",
        "TGT_ADLS_PATH", "TGT_COLUMN_NAMES", "TGT_FORMAT", "TGT_DATA_TYPE", "TGT_LOAD_OPTION",
        "TGT_RJT_TABLE_NAME", "TGT_RJT_ADLS_PATH"],
    "ADLS_FIXED_WIDTH_HANDLER": [
        "PROCESS_NAME", "SEGMENT", "SEGMENT_FILTER", "COL", "LEN", "start_ind", "TGT_TABLE",
        "TGT_AUDIT_CLMS", "SRC_ADLS_ARCHVL_PATH", "ACTIVE_FLAG", "IS_MANDATORY",
        "COLUMN_VALIDATIONS"],
    "DATABRICKS_NOTEBOOK_DETAILS": [
        "PIPELINE_NAME", "SEQ_NM", "PROCESS_NAME", "ACTIVE_FLAG", "DATABRICKS_NOTEBOOK_NAME"],
    "DATA_QUALITY_RULES": [
        "SEQUENCE_NO", "RULE_TYPE", "RULE_CLASS", "ACTIVE_RULE_FLG", "SOURCE_COLUMN",
        "INPUT_PARAM", "TARGET_COLUMN"],
    "EMAIL_TEMPLATE_CONFIG": ["TEMPLATE_NAME", "PROCESS_NAME", "STATUS", "ACTIVE_FLAG"],
}

# Blank-and-flag: framework-assigned IDs / connections / containers / audit
# stamps, client wording (email bodies), notebook paths, and the
# fixed-width-handler VERSION / SEGMNT_TYP / FILE_TYPE / EXTENSION that no
# input document states. Pinned exactly.
EXPECTED_BLANK = {
    "ADLS_DELTA_INGESTION_DETAILS": [
        "CLAIM_TYPE_ID", "CREATED_BY", "CREATED_DATE", "FILE_FOOTER_FLAG", "FILE_HEADER_FLAG",
        "GROUP_ID", "MAPPING_EXPRESSION", "METADATA_CONNECTION_ID", "OBJECT_ID", "PIPELINE_ID",
        "RECYCL_RETN_DAYS", "SRC_ADLS_CONNECTION_ID", "SRC_COL_LNGTH", "SRC_COL_STRT_END_INDX",
        "SRC_CONTAINER_NAME", "SRC_FILE_DELIMITER", "SRC_REC_LNGTH", "TGT_CONNECTION_ID",
        "TGT_CONTAINER_NAME", "UPDATED_BY", "UPDATED_DATE"],
    "ADLS_FIXED_WIDTH_HANDLER": [
        "CREATED_BY", "CREATED_DATE", "EXTENSION", "FILE_REJECTION", "FILE_TYPE", "PIPELINE_ID",
        "RJCT_RSN_COLUMN_NM", "SEGMNT_TYP", "UPDATED_BY", "UPDATED_DATE", "VERSION"],
    "DATABRICKS_NOTEBOOK_DETAILS": [
        "CLUSTER_DETAILS_ID", "CREATED_BY", "CREATED_DATE", "DATABRICKS_CLUSTERID",
        "DATABRICKS_NOTEBOOK_PATH", "DATABRICKS_WORKSPACE_SECRET", "DATABRICKS_WORKSPACE_URL",
        "DELETE_DATABRICKS_NOTEBOOK_NAME", "DQ_NOTEBOOK_PATH", "GROUP_ID", "PIPELINE_ID",
        "UPDATED_BY", "UPDATED_DATE"],
    "DATA_FACTORY_PIPELINE_SCHEDULE": [
        "ACTIVE_START_DATE", "CREATED_BY", "CREATED_DATE", "ESTIMATED_START_TIME",
        "PARENT_PIPELINE_ID", "PIPELINE_ID", "UPDATED_BY", "UPDATED_DATE"],
    "DATA_QUALITY_RULES": [
        "CREATED_BY", "CREATED_DATE", "GROUP_ID", "OBJECT_ID", "UPDATED_BY", "UPDATED_DATE"],
    "EMAIL_TEMPLATE_CONFIG": [
        "BODY", "BODY_QUERY", "CREATED_BY", "CREATED_DATE", "EMAIL_CC", "SENDER_EMAIL",
        "SENDER_NAME", "SUBJECT", "TEMPLATE_ID", "UPDATED_BY", "UPDATED_DATE"],
    "FILE_ADLS_INGESTION_DETAILS": [
        "CREATED_BY", "CREATED_DATE", "GROUP_ID", "OBJECT_ID", "PIPELINE_ID", "SRC_CONNECTION_ID",
        "SRC_EXTRACT_START_TIME", "SRC_ROOT_DIR", "TGT_CONTAINER_NAME", "TGT_STORAGE_ACCOUNT_NAME",
        "UPDATED_BY", "UPDATED_DATE"],
    "STGDELTA_STDDELTA_INGESTION_DET": [
        "CREATED_BY", "CREATED_DATE", "GROUP_ID", "METADATA_CONNECTION_ID", "OBJECT_ID",
        "PIPELINE_ID", "SRC_ADLS_CONNECTION_ID", "SRC_CONTAINER_NAME", "TGT_CONNECTION_ID",
        "TGT_CONTAINER_NAME", "UPDATED_BY", "UPDATED_DATE"],
}


@needs_golden
def test_pair1_iig_matches_golden_sheets_headers_and_row_counts(pair1_run):
    _gate, framework_dir = pair1_run
    ours = load_workbook(framework_dir / "config_rows.xlsx")
    golden = _golden_iig()
    assert [s for s in ours.sheetnames if not s.startswith("_")] == golden.sheetnames
    for name in golden.sheetnames:
        g_headers, g_rows = _sheet_rows(golden[name])
        o_headers, o_rows = _sheet_rows(ours[name])
        assert o_headers == g_headers, name
        assert len(o_rows) == len(g_rows), name


@needs_golden
def test_pair1_iig_pinned_cells_match_the_golden(pair1_run):
    _gate, framework_dir = pair1_run
    ours = load_workbook(framework_dir / "config_rows.xlsx")
    golden = _golden_iig()
    mismatches = []
    for name, columns in PINNED_COLUMNS.items():
        g_headers, g_rows = _sheet_rows(golden[name])
        _o_headers, o_rows = _sheet_rows(ours[name])
        for column in columns:
            index = g_headers.index(column)
            for row_number, (g_row, o_row) in enumerate(zip(g_rows, o_rows, strict=True), 1):
                if g_row[index] != o_row[index]:
                    mismatches.append((name, row_number, column, g_row[index], o_row[index]))
    assert mismatches == []
    # A few of the pins, spelled out (the golden's own values):
    _h, rows = _sheet_rows(ours["ADLS_DELTA_INGESTION_DETAILS"])
    row = dict(zip(_h, rows[0], strict=True))
    assert row["SRC_COLUMNS"].startswith("col1:SEGMENT_IDENTIFIER,col2:FILE_TYPE,")
    assert row["SRC_DATA_TYPE"].startswith("String:String,String:String,String:String,String:Date,")
    assert row["TGT_DATA_TYPE"].endswith(
        "Decimal(17,2),Decimal(18,2),Decimal(19,2),Decimal(20,2),Decimal(21,2),Decimal(22,2),"
        "Date,Date,String,String,String,String,TIMESTAMP,TIMESTAMP")
    assert row["TGT_COLUMN_NAMES"].endswith(",SRC_FILE_NAME,REC_CREATION_TIME,REC_UPDATED_TIME")
    _h, rows = _sheet_rows(ours["ADLS_FIXED_WIDTH_HANDLER"])
    handler = [dict(zip(_h, r, strict=True)) for r in rows]
    assert [r["SEGMENT"] for r in handler] == ["HDDR", "DET", "TRLR"]
    assert [r["SEGMENT_FILTER"] for r in handler] == [
        "value like '00%'", "value like '10%'", "value like '99%'"]
    assert handler[1]["LEN"] == "2,15,16,17,18,19,20,21,22,23,24,25"
    assert handler[1]["start_ind"] == "1,3,18,34,51,69,88,108,129,151,174,198"
    assert handler[0]["TGT_TABLE"] == "pr_dlk_vnd_p.stg_vnd_p_accum.vnd_p_accum_client"
    _h, rows = _sheet_rows(ours["DATA_QUALITY_RULES"])
    dq = [dict(zip(_h, r, strict=True)) for r in rows]
    assert dq[0]["INPUT_PARAM"] == "yyyyMMdd--yyyy-MM-dd"
    assert dq[1]["SOURCE_COLUMN"] == (
        "INDIVIDUAL_DEDUCTIBLE,FAMILY_DEDUCTIBLE,COPAY,COINSURANCE,INDIVIDUAL_LIMIT,"
        "FAMILY_LIMIT,PROCESS_DATE,PLAN_YEAR_START_DATE,PLAN_YEAR_END_DATE")


@needs_golden
def test_pair1_iig_everything_else_is_blank_and_flagged(pair1_run):
    gate, framework_dir = pair1_run
    flags = sorted(f for f in gate.flags if f.startswith("iig_blank:"))
    expected = sorted(f"iig_blank:{tab}.{header}"
                      for tab, headers in EXPECTED_BLANK.items() for header in headers)
    assert flags == expected
    # Every flagged column really is blank in the workbook, in every row.
    ours = load_workbook(framework_dir / "config_rows.xlsx")
    for tab, headers in EXPECTED_BLANK.items():
        o_headers, rows = _sheet_rows(ours[tab])
        for header in headers:
            index = o_headers.index(header)
            assert all(r[index] == "" for r in rows), (tab, header)
    # And nothing the golden fills with a real client-assigned value was
    # invented: no cell carries a SYN- identifier or RFC stamp.
    for sheet in ours.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                assert not (isinstance(value, str) and (value.startswith("SYN-")
                                                        or value.startswith("RFC#")))


@needs_golden
def test_iig_v2_config_headers_are_the_golden_headers(config):
    template = config.metadata.templates["iig_v2"]
    golden = _golden_iig()
    assert list(template.tabs) == golden.sheetnames
    for name in golden.sheetnames:
        headers, _rows = _sheet_rows(golden[name])
        assert template.tabs[name].headers == headers, name


# -- defaults untouched ------------------------------------------------------------ #


def test_default_profile_and_template_are_the_edo_sfmc_two_file_layout(config):
    assert config.conventions.profile == "edo_sfmc"
    assert config.conventions.get(None).ddl_layout == "two_files"
    assert config.metadata.template == "iig_v1"
    assert config.metadata.resolve(None) == ("iig_v1", None)
    with pytest.raises(ValueError):
        config.conventions.get("no_such_profile")
    with pytest.raises(ValueError):
        config.metadata.resolve("iig_v9")


def test_pair1_default_profile_writes_two_files_and_no_combined_ddl(config, pair1_spec,
                                                                    tmp_path):
    scoped = _scoped(config, tmp_path)
    gate = cli._generate_feed(pair1_spec, scoped, dry_run=True, skip_tests=True,
                              output_mode="framework")
    framework_dir = tmp_path / "out" / pair1_spec.feed_slug / "framework"
    names = sorted(p.name for p in framework_dir.iterdir())
    assert names == ["ADDITION.md", "accumulators_stage_table_creation.txt",
                     "accumulators_standard_table_creation.txt", "config_inserts.xlsx",
                     "config_rows.xlsx"]
    # iig_v1: the 7-tab reference layout, no blank-and-flag flags.
    assert not any(f.startswith("iig_blank:") for f in gate.flags)
    ours = load_workbook(framework_dir / "config_rows.xlsx")
    assert [s for s in ours.sheetnames if not s.startswith("_")] == list(
        config.demo.metadata_sheet.tabs)


def test_combined_ddl_without_feed_abbreviation_falls_back_to_the_slug_and_flags(
        config, pair1_spec, tmp_path):
    from codegen.emit.framework import emit_framework
    from codegen.faq import LoadPatternFaq

    artefacts = emit_framework(pair1_spec, LoadPatternFaq(), [], config, tmp_path,
                               conventions_profile="acfc_prx", iig_template="iig_v2")
    names = [p.name for p in artefacts.files]
    assert "ACCUMULATORS_DDL.txt" in names
    assert any(f.startswith("ddl_file_name_from_slug:") for f in artefacts.flags)
    # No FAQ -> no process name / discriminators: those cells join the blank list.
    assert "iig_blank:ADLS_FIXED_WIDTH_HANDLER.PROCESS_NAME" in artefacts.flags
    assert "iig_blank:ADLS_FIXED_WIDTH_HANDLER.SEGMENT_FILTER" in artefacts.flags
