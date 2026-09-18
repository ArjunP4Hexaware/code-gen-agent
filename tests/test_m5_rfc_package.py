"""M5 — the `rfc` output mode: RFC deployment packages.

* pair 1 under ``acfc_prx`` + ``iig_v2`` + ``main_single`` writes the
  documented PRX package file set (RFC_PACKAGE_SHAPES §1) with the
  documented names; the DDL is the byte-identical golden; the IIG carries
  the template's eight sheets only; the single-sheet playbook has the §1
  header pattern, config task rows, the RFC# column from the run input and
  every date / time / owner / team / status cell blank AND flagged (pinned);
* the SFMC pair under ``edo_sfmc`` + ``iig_v1`` + ``sfmc_7sheet`` (file log
  on) writes the INGESTION_A file set; the playbook's sheets and header rows
  match the scrubbed reference template, its values are blank-and-flag;
  FILE_LOG_INFORMATION.txt is the documented T-SQL shape;
* MANIFEST.md lists every file written; Runbook / TDD are out of scope;
* the RFC number is a run input (FAQ ``rfc_number``): unanswered → the
  documented placeholder + flag; answered → the package is named after it;
* config overlays deep-merge and stay loud on unknown sections; the shipped
  config carries no client-shaped template rows (they live in the pair-1
  fixture overlay).
"""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from codegen import cli
from codegen.config import load_config
from codegen.emit.rfc import emit_rfc_package
from codegen.faq import FaqAnswer, LoadPatternFaq
from codegen.resolve.resolver import resolve_pair as resolve_contracts
from test_m4_acceptance import GOLDEN_DDL, _scoped, needs_golden

REPO = Path(__file__).resolve().parents[1]
REFERENCE_PLAYBOOK = REPO / "fixtures" / "reference" / "SFMC_Deployment_Playbook.xlsx"
SFMC_FRD = REPO / "fixtures" / "contracts" / "FRD_sfmc_email_campaign.contract.json"
SFMC_STTM = REPO / "fixtures" / "contracts" / "sttm_mapping_contracts_sfmc.json"

needs_sfmc = pytest.mark.skipif(
    not (SFMC_FRD.is_file() and SFMC_STTM.is_file()),
    reason="SFMC contract pair not restored (removed 2026-08-22)",
)
needs_reference = pytest.mark.skipif(not REFERENCE_PLAYBOOK.is_file(),
                                     reason="scrubbed SFMC playbook reference absent")

MAIN_HEADERS = ["Task #", "Task", "Detail", "", "RFC#s", "Deployment Team",
                "Planned Start Date", "Planned End Date", "Actual Start Date", "Actual End Date",
                "Planned Start Time", "Planned End Time", "Actual Start Time", "Actual End Time",
                "Task Owner", "Status"]
IIG_V2_SHEETS = ["DATA_FACTORY_PIPELINE_SCHEDULE", "FILE_ADLS_INGESTION_DETAILS",
                 "ADLS_DELTA_INGESTION_DETAILS", "STGDELTA_STDDELTA_INGESTION_DET",
                 "ADLS_FIXED_WIDTH_HANDLER", "DATABRICKS_NOTEBOOK_DETAILS", "DATA_QUALITY_RULES",
                 "EMAIL_TEMPLATE_CONFIG"]
MAIN_BLANK = ["Deployment Team", "Planned Start Date", "Planned End Date", "Actual Start Date",
              "Actual End Date", "Planned Start Time", "Planned End Time", "Actual Start Time",
              "Actual End Time", "Task Owner", "Status"]

# The 7-sheet template's blank-and-flag list, pinned (every value the
# reference deployment carried: schedule grid, cover values, team / dates /
# times / owner / status / dependency / instructions / comments, contacts).
SFMC_BLANK = {
    "Overview": ["schedule grid"],
    "Cover Page": ["Project", "Revision history Date", "Revision history Version",
                   "Revision history Author(s)",
                   "Revision history Description of Version/Changes"],
    "PrePost-Prod & Prod Execution": [
        "Deployment Team", "Planned Start Date", "Planned End Date", "Actual Start Date",
        "Actual End Date", "Planned Start Time", "Planned End Time", "Actual Start Time",
        "Actual End Time", "Task Owner", "Status", "Task Dependency", "Task Instructions",
        "Comments: Touch Point Calls/Emails"],
    "Rollback Execution": [
        "Rollback Execution Team", "Rollback Execution Planned Start Date",
        "Rollback Execution Planned End Date", "Rollback Execution Actual Start Date",
        "Rollback Execution Actual End Date", "Rollback Execution Planned Start Time",
        "Rollback Execution Planned End Time", "Rollback Execution Actual Start Time",
        "Rollback Execution Actual End Time", "Rollback Execution Task Owner",
        "Rollback Execution Status", "Task Dependency", "Task Instructions", "Communication",
        "Comments: Touch Point Calls/Emails"],
    "Rollback Validation": [
        "Rollback Validation Team", "Rollback Validation Planned Start Date",
        "Rollback Validation Planned End Date", "Rollback Validation Actual Start Date",
        "Rollback Validation Actual End Date", "Rollback Validation Planned Start Time",
        "Rollback Validation Planned End Time", "Rollback Validation Actual Start Time",
        "Rollback Validation Actual End Time", "Task / App Owner or App Support Resource",
        "Production RFC", "Rollback Validation Conditions", "Rollback Validation Status",
        "Rollback Validation Done By", "Task Dependency", "Task Instructions", "Communication",
        "Comments: Touch Point Calls/Emails"],
    "Contact List": ["Contact Name", "Deployment Team", "Cell Phone", "Office",
                     "Escalation Point (name/number)"],
}


def _package_dir(tmp: Path, slug: str) -> Path:
    (package,) = [p for p in (tmp / "out" / slug).iterdir() if p.name.startswith("RFC")]
    return package


# M7: the SQL Server DML deliverable rides the acfc_prx profile into the package.
DML_FILES = sorted([f"Insert_scripts_config_table_{env}.py" for env in ("q1", "a2", "prod")]
                   + [f"config_inserts_{env}.sql" for env in ("q1", "a2", "prod")])


def _manifest_files(package: Path) -> list[str]:
    import re

    text = (package / "MANIFEST.md").read_text(encoding="utf-8")
    return re.findall(r"^\| `([^`]+)` \|", text, flags=re.MULTILINE)


# -- pair 1: the PRX package ------------------------------------------------------ #


@pytest.fixture(scope="module")
def pair1_rfc(pair1_config, pair1_spec, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pair1_rfc")
    gate = cli._generate_feed(pair1_spec, _scoped(pair1_config, tmp), dry_run=True,
                              skip_tests=True, output_mode="rfc",
                              conventions_profile="acfc_prx", iig_template="iig_v2",
                              playbook_template="main_single")
    return gate, tmp, _package_dir(tmp, pair1_spec.feed_slug)


@needs_golden
def test_pair1_prx_package_file_set_and_names(pair1_rfc, pair1_spec):
    gate, tmp, package = pair1_rfc
    assert gate.verdict == "PASS_WITH_FLAGS"
    assert package.name == "RFC######_ACCUM"
    assert sorted(p.name for p in package.iterdir()) == sorted([
        "ACCUM_DDL.txt", "ACCUM_IIG.xlsx", *DML_FILES, "MANIFEST.md",
        "RFC######_ACCUM_Deployment_Playbook.xlsx"])
    # The DDL is the golden, byte for byte; framework/ still holds the sources.
    assert (package / "ACCUM_DDL.txt").read_bytes() == GOLDEN_DDL.read_bytes()
    framework_dir = tmp / "out" / pair1_spec.feed_slug / "framework"
    assert (framework_dir / "ACCUM_DDL.txt").is_file()
    assert (framework_dir / "config_rows.xlsx").is_file()
    # The IIG carries the template's sheets only — provenance stays in framework/.
    iig = load_workbook(package / "ACCUM_IIG.xlsx")
    assert iig.sheetnames == IIG_V2_SHEETS
    assert not any(s.startswith("_") for s in iig.sheetnames)
    assert "_provenance" in load_workbook(framework_dir / "config_rows.xlsx").sheetnames


@needs_golden
def test_pair1_single_sheet_playbook(pair1_rfc, pair1_config):
    gate, _tmp, package = pair1_rfc
    wb = load_workbook(package / "RFC######_ACCUM_Deployment_Playbook.xlsx")
    assert wb.sheetnames == ["Main"]
    assert [c.value or "" for c in wb["Main"][1]] == MAIN_HEADERS
    rows = [[("" if v is None else str(v)) for v in r]
            for r in wb["Main"].iter_rows(min_row=2, values_only=True)]
    tasks = pair1_config.playbook.tasks_for("acfc_prx")
    expected = (len(tasks.pre_production) + len(tasks.production) + len(tasks.post_production)
                + len(tasks.rollback_execution) + len(tasks.rollback_validation))
    rows = [r for r in rows if any(r)]
    assert len(rows) == expected == 7
    col = {h: i for i, h in enumerate(MAIN_HEADERS)}
    assert [r[col["Task #"]] for r in rows] == [str(i) for i in range(1, 8)]
    assert all(r[col["RFC#s"]] == "RFC######" for r in rows)
    assert rows[0][col["Task"]] == "Execute the deployment DDL"
    assert "ACCUM_DDL.txt" in rows[0][col["Detail"]]
    assert "ACCUM_IIG.xlsx" in rows[1][col["Detail"]]
    for header in MAIN_BLANK:
        assert all(r[col[header]] == "" for r in rows), header
    blank = sorted(f for f in gate.flags if f.startswith("playbook_blank:"))
    assert blank == sorted(f"playbook_blank:Main.{h}" for h in MAIN_BLANK)
    assert any(f.startswith("rfc_number_unanswered:") for f in gate.flags)
    assert not any(f.startswith("rfc_feed_name_from_slug") for f in gate.flags)
    assert not any(f.startswith("file_log_") for f in gate.flags)


@needs_golden
def test_pair1_manifest_lists_every_file_and_the_out_of_scope_documents(pair1_rfc):
    _gate, _tmp, package = pair1_rfc
    listed = _manifest_files(package)
    written = sorted(p.name for p in package.iterdir() if p.name != "MANIFEST.md")
    assert sorted(listed) == written
    text = (package / "MANIFEST.md").read_text(encoding="utf-8")
    assert "`drag_fill_suspect`" in text and "`iig_blank` ×" in text
    assert "RFC###### (placeholder — unanswered)" in text
    assert "Operational_RunBook_<Feed>.docx" in text and "TDD_<description>.docx" in text
    assert "Out of scope" in text


@needs_golden
def test_answered_rfc_number_names_the_package(pair1_rfc, pair1_config, pair1_spec, tmp_path):
    from codegen.emit.framework import emit_framework

    faq = LoadPatternFaq(
        rfc_number=FaqAnswer(value="SYN001", source="engineer", evidence="ticket"),
        feed_abbreviation=FaqAnswer(value="ACCUM", source="engineer", evidence="package"),
    )
    framework = emit_framework(pair1_spec, faq, [], pair1_config, tmp_path,
                               conventions_profile="acfc_prx", iig_template="iig_v2")
    rfc = emit_rfc_package(pair1_spec, faq, pair1_config, tmp_path, framework, flags_so_far=[],
                           conventions_profile="acfc_prx", iig_template="iig_v2",
                           playbook_template="main_single")
    assert rfc.package_dir.name == "RFCSYN001_ACCUM"
    assert sorted(p.name for p in rfc.files) == sorted([
        "ACCUM_DDL.txt", "ACCUM_IIG.xlsx", *DML_FILES, "MANIFEST.md",
        "RFCSYN001_ACCUM_Deployment_Playbook.xlsx"])
    assert not any(f.startswith("rfc_") for f in rfc.flags)
    wb = load_workbook(rfc.package_dir / "RFCSYN001_ACCUM_Deployment_Playbook.xlsx")
    rows = [r for r in wb["Main"].iter_rows(min_row=2, values_only=True) if any(r)]
    assert {r[MAIN_HEADERS.index("RFC#s")] for r in rows} == {"RFCSYN001"}


# -- SFMC: the INGESTION_A package ------------------------------------------------- #


@pytest.fixture(scope="module")
def sfmc_rfc(config, tmp_path_factory):
    if not (SFMC_FRD.is_file() and SFMC_STTM.is_file()):
        pytest.skip("SFMC contract pair not restored (removed 2026-08-22)")
    tmp = tmp_path_factory.mktemp("sfmc_rfc")
    scoped = _scoped(config, tmp).model_copy(update={
        "rfc": config.rfc.model_copy(update={"file_log_information": True})})
    (spec,) = resolve_contracts(SFMC_FRD, SFMC_STTM, scoped)
    gate = cli._generate_feed(spec, scoped, dry_run=True, skip_tests=True, output_mode="rfc")
    return gate, tmp, spec, _package_dir(tmp, spec.feed_slug)


@needs_sfmc
@needs_reference
def test_sfmc_ingestion_a_package_file_set(sfmc_rfc, config):
    gate, tmp, spec, package = sfmc_rfc
    token = "SFMC_EMAIL_CAMPAIGN_TRACKING"
    assert package.name == f"RFC######_{token}"
    assert sorted(p.name for p in package.iterdir()) == sorted([
        f"{token}_stage_table_creation.txt", f"{token}_standard_table_creation.txt",
        f"{token}_IIG.xlsx", f"RFC######_{token}_Deployment_Playbook.xlsx",
        "FILE_LOG_INFORMATION.txt", "MANIFEST.md"])
    # DDL exactly as the edo_sfmc profile emitted it (two files, renamed).
    framework_dir = tmp / "out" / spec.feed_slug / "framework"
    for layer in ("stage", "standard"):
        assert ((package / f"{token}_{layer}_table_creation.txt").read_bytes()
                == (framework_dir / f"{spec.feed_slug}_{layer}_table_creation.txt").read_bytes())
    # IIG: the iig_v1 (7-tab reference) sheets only.
    assert load_workbook(package / f"{token}_IIG.xlsx").sheetnames == list(
        config.demo.metadata_sheet.tabs)
    assert sorted(_manifest_files(package)) == sorted(
        p.name for p in package.iterdir() if p.name != "MANIFEST.md")
    assert any(f.startswith("rfc_feed_name_from_slug:") for f in gate.flags)
    assert any(f.startswith("rfc_number_unanswered:") for f in gate.flags)


@needs_sfmc
@needs_reference
def test_sfmc_playbook_matches_the_reference_template_and_is_blank_and_flag(sfmc_rfc, config):
    gate, _tmp, _spec, package = sfmc_rfc
    ours = load_workbook(next(package.glob("*_Deployment_Playbook.xlsx")))
    reference = load_workbook(REFERENCE_PLAYBOOK)
    assert ours.sheetnames == reference.sheetnames
    tpl = config.playbook.templates["sfmc_7sheet"]
    for name in (tpl.task_sheet, tpl.rollback_execution_sheet, tpl.rollback_validation_sheet):
        for row in (1, 2):
            ours_row = [c.value for c in ours[name][row]]
            assert ours_row == [c.value for c in reference[name][row]], (name, row)
    assert [c.value for c in ours[tpl.contact_sheet][1]] == [
        c.value for c in reference[tpl.contact_sheet][1]]
    assert ours[tpl.contact_sheet].max_row == 1
    # Overview: nothing survives (dates, times, people).
    assert all(c.value is None for row in ours[tpl.overview_sheet].iter_rows() for c in row)
    # Cover page: labels kept, values blank.
    cover = ours[tpl.cover_sheet]
    assert cover["B6"].value == "Project:" and cover["B7"].value == "Project Manager:"
    assert cover["B14"].value == "Date" and cover["B15"].value is None
    assert cover["E15"].value is None
    # Task rows: config tasks under the section bands, RFC# from the run input.
    tasks = config.playbook.tasks_for("edo_sfmc")
    task_sheet = ours[tpl.task_sheet]
    values = [[c.value for c in row] for row in task_sheet.iter_rows(min_row=3)]
    labels = [r[0] for r in values if r[0] in ("PRE-PRODUCTION", "PRODUCTION", "POST PRODUCTION")]
    assert labels == ["PRE-PRODUCTION", "PRODUCTION", "POST PRODUCTION"]
    task_rows = [r for r in values if isinstance(r[0], int)]
    assert [r[1] for r in task_rows] == [t.task for t in tasks.production] + [
        t.task for t in tasks.post_production]
    assert {r[3] for r in task_rows} == {"RFC######"}
    assert all(all(v in (None, "") for v in r[4:]) for r in task_rows)
    blank = sorted(f for f in gate.flags if f.startswith("playbook_blank:"))
    assert blank == sorted(f"playbook_blank:{sheet}.{h}"
                           for sheet, headers in SFMC_BLANK.items() for h in headers)


@needs_sfmc
def test_sfmc_file_log_information_is_the_documented_t_sql_shape(sfmc_rfc, config):
    gate, _tmp, _spec, package = sfmc_rfc
    text = (package / "FILE_LOG_INFORMATION.txt").read_text(encoding="utf-8")
    lines = text.splitlines()
    assert lines[0] == "CREATE TABLE [dbo].[FILE_LOG_INFORMATION]("
    assert lines[-1] == ")"
    body = lines[1:-1]
    assert len(body) == len(config.rfc.file_log_columns) == 13
    assert body[0] == "  [ID] [int] IDENTITY(1,1) NOT NULL,"
    assert body[-1] == "  [UPDATED_DATE] [datetime] NOT NULL"
    assert all(line.endswith(",") for line in body[:-1])
    assert "[CREATED_BY] [varchar](<length>) NOT NULL" in text
    assert any(f.startswith("file_log_placeholder_unstated:") for f in gate.flags)


# -- config: mode, overlays, shipped vocabulary ------------------------------------- #


def test_rfc_mode_and_playbook_knobs_exist(config):
    assert config.output.mode == "notebook"
    assert config.playbook.template == "sfmc_7sheet"
    assert sorted(config.playbook.templates) == ["main_single", "sfmc_7sheet"]
    assert config.playbook.templates["main_single"].headers == MAIN_HEADERS
    assert config.rfc.file_log_information is False
    assert config.rfc.number_placeholder == "######"
    with pytest.raises(ValueError):
        config.playbook.resolve("no_such_template")
    # "default" tasks serve every profile without its own list.
    assert config.playbook.tasks_for("acfc_prx") is config.playbook.tasks["default"]


def test_config_overlay_deep_merges_and_stays_loud(tmp_path, monkeypatch):
    overlay = tmp_path / "overlay.yaml"
    overlay.write_text(yaml.safe_dump({
        "metadata": {"templates": {"iig_v2": {"template_rows": {
            "DATA_FACTORY_PIPELINE_SCHEDULE": [{"PIPELINE_NAME": "X"}]}}}},
        "rfc": {"number_placeholder": "000000"},
    }), encoding="utf-8")
    merged = load_config(REPO / "config" / "config.yaml", overlays=[overlay])
    rows = merged.metadata.templates["iig_v2"].template_rows
    assert rows["DATA_FACTORY_PIPELINE_SCHEDULE"] == [{"PIPELINE_NAME": "X"}]
    assert rows["EMAIL_TEMPLATE_CONFIG"]  # untouched sibling key survives the merge
    assert merged.rfc.number_placeholder == "000000"
    assert merged.rfc.dir_pattern == "RFC{rfc_number}_{feed_slug}"  # sibling scalar kept
    # env-driven overlay, same effect
    monkeypatch.setenv("CODEGEN_CONFIG_OVERLAYS", str(overlay))
    assert load_config(REPO / "config" / "config.yaml").rfc.number_placeholder == "000000"
    monkeypatch.delenv("CODEGEN_CONFIG_OVERLAYS")
    bad = tmp_path / "bad.yaml"
    bad.write_text("rfc_typo: {x: 1}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="unknown top-level config section"):
        load_config(REPO / "config" / "config.yaml", overlays=[bad])


def test_shipped_config_has_no_client_shaped_template_rows(config, pair1_config):
    shipped = config.metadata.templates["iig_v2"].template_rows
    assert set(shipped) == {"EMAIL_TEMPLATE_CONFIG"}
    text = (REPO / "config" / "config.yaml").read_text(encoding="utf-8")
    assert "PL_VND_P" not in text and "Daily_Ingested_Status_Report" not in text
    overlaid = pair1_config.metadata.templates["iig_v2"].template_rows
    assert len(overlaid["DATA_FACTORY_PIPELINE_SCHEDULE"]) == 4
    assert len(overlaid["DATABRICKS_NOTEBOOK_DETAILS"]) == 3


def test_ui_runner_accepts_the_rfc_mode():
    pytest.importorskip("httpx")
    from ui.backend.demo import DemoRunner

    runner = DemoRunner.__new__(DemoRunner)
    import threading

    runner._lock = threading.Lock()
    runner.state = "idle"
    runner.output_parts = None
    runner.select_output_mode("rfc")
    assert runner.output_mode == "rfc" and runner.output_parts == ["rfc"]
    runner.select_output_mode("all")
    assert runner.output_mode == "all" and runner.output_parts == ["all"]
    # Independent parts map onto one mode; empty = no mode (run refused).
    runner.select_output_parts(["notebook", "rfc"])
    assert runner.output_mode == "all"
    runner.select_output_parts(["framework", "rfc"])
    assert runner.output_mode == "rfc"
    runner.select_output_parts(["notebook", "framework"])
    assert runner.output_mode == "both"
    runner.select_output_parts([])
    assert runner.output_mode is None
    with pytest.raises(ValueError):
        runner.select_output_parts(["zip"])
    with pytest.raises(ValueError):
        runner.select_output_mode("zip")


@needs_golden
def test_all_mode_writes_notebook_framework_and_rfc(pair1_config, pair1_spec, tmp_path):
    """`all` = the notebook tree + framework artefacts + the RFC package,
    the same package `rfc` mode writes."""
    scoped = _scoped(pair1_config, tmp_path)
    gate = cli._generate_feed(pair1_spec, scoped, dry_run=True, skip_tests=True,
                              output_mode="all", conventions_profile="acfc_prx",
                              iig_template="iig_v2", playbook_template="main_single")
    feed_dir = tmp_path / "out" / pair1_spec.feed_slug
    assert (feed_dir / "pipeline").is_dir() and (feed_dir / "framework").is_dir()
    package = _package_dir(tmp_path, pair1_spec.feed_slug)
    assert sorted(p.name for p in package.iterdir()) == sorted([
        "ACCUM_DDL.txt", "ACCUM_IIG.xlsx", *DML_FILES, "MANIFEST.md",
        "RFC######_ACCUM_Deployment_Playbook.xlsx"])
    assert (package / "ACCUM_DDL.txt").read_bytes() == GOLDEN_DDL.read_bytes()
    assert gate.verdict == "PASS_WITH_FLAGS"
    assert any(f.startswith("playbook_blank:") for f in gate.flags)

