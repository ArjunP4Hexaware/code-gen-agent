"""M7 §4 + §5 — artefacts labelled by target system (framework groups,
ADDITION.md block, MANIFEST.md column, review-copy note) and the transcript's
semantics applied to the templates: CREATED_BY / UPDATED_BY from the FAQ
rfc_number in both IIG versions (blank + flagged otherwise), the
CLAIM_TYPE_ID default knob (goldens' blank by default), the DDL header knob
(off for both shipped profiles — goldens are byte-compared)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from codegen.emit.framework import (
    DDL_TARGET_HEADER,
    TARGET_DDL,
    TARGET_DML,
    TARGET_REVIEW,
    artefact_group,
    artefact_groups,
    emit_framework,
)
from codegen.emit.rfc import emit_rfc_package
from codegen.faq import FaqAnswer, LoadPatternFaq

FAQ = LoadPatternFaq(
    rfc_number=FaqAnswer(value="SYN001", source="engineer", evidence="ticket"),
    feed_abbreviation=FaqAnswer(value="ACCUM", source="engineer"),
)


def test_artefact_groups_by_target_system():
    names = ["ACCUM_DDL.txt", "x_stage_table_creation.txt", "config_inserts_q1.sql",
             "Insert_scripts_config_table_q1.py", "config_rows.xlsx", "config_inserts.xlsx",
             "ADDITION.md"]
    assert [artefact_group(n) for n in names] == [
        TARGET_DDL, TARGET_DDL, TARGET_DML, TARGET_DML, TARGET_REVIEW, TARGET_REVIEW, "Notes"]
    groups = artefact_groups([Path(n) for n in names])
    assert list(groups) == [TARGET_DDL, TARGET_DML, TARGET_REVIEW, "Notes"]
    assert groups[TARGET_DML] == ["config_inserts_q1.sql", "Insert_scripts_config_table_q1.py"]


def test_framework_labels_groups_readme_and_addition_block(pair1_config, pair1_spec, tmp_path):
    framework = emit_framework(pair1_spec, FAQ, [], pair1_config, tmp_path,
                               conventions_profile="acfc_prx", iig_template="iig_v2")
    assert set(framework.groups) == {TARGET_DDL, TARGET_DML, TARGET_REVIEW, "Notes"}
    assert len(framework.groups[TARGET_DML]) == 6
    framework_dir = tmp_path / pair1_spec.feed_slug / "framework"
    addition = (framework_dir / "ADDITION.md").read_text(encoding="utf-8")
    assert "## By target system" in addition
    assert f"**{TARGET_DML}**" in addition and f"**{TARGET_DDL}**" in addition
    wb = load_workbook(framework_dir / "config_inserts.xlsx")
    assert wb.sheetnames[0] == "README"
    assert wb["README"]["A1"].value.startswith(
        "review copy — executable script is config_inserts_<env>.sql")
    # the DDL file itself carries NO header line (golden byte-identity; knob off)
    ddl = (framework_dir / "ACCUM_DDL.txt").read_text(encoding="utf-8")
    assert DDL_TARGET_HEADER not in ddl


def test_ddl_target_header_knob_prefixes_the_ddl_when_on(pair1_config, pair1_spec, tmp_path):
    profiles = dict(pair1_config.conventions.profiles)
    profiles["acfc_prx"] = profiles["acfc_prx"].model_copy(update={"target_system_header": True})
    config = pair1_config.model_copy(update={
        "conventions": pair1_config.conventions.model_copy(update={"profiles": profiles})})
    framework = emit_framework(pair1_spec, FAQ, [], config, tmp_path,
                               conventions_profile="acfc_prx", iig_template="iig_v2")
    ddl = next(p for p in framework.files if p.name == "ACCUM_DDL.txt")
    assert ddl.read_text(encoding="utf-8").startswith(DDL_TARGET_HEADER + "\n")


def test_reference_profile_output_carries_no_labels(pair1_config, pair1_spec, tmp_path):
    framework = emit_framework(pair1_spec, LoadPatternFaq(), [], pair1_config, tmp_path,
                               conventions_profile="edo_sfmc", iig_template="iig_v1")
    framework_dir = tmp_path / pair1_spec.feed_slug / "framework"
    assert "## By target system" not in (framework_dir / "ADDITION.md").read_text(encoding="utf-8")
    assert load_workbook(framework_dir / "config_inserts.xlsx").sheetnames[0] != "README"
    assert TARGET_DML not in framework.groups


def test_manifest_groups_files_by_target_system(pair1_config, pair1_spec, tmp_path):
    framework = emit_framework(pair1_spec, FAQ, [], pair1_config, tmp_path,
                               conventions_profile="acfc_prx", iig_template="iig_v2")
    rfc = emit_rfc_package(pair1_spec, FAQ, pair1_config, tmp_path, framework,
                           flags_so_far=list(framework.flags), conventions_profile="acfc_prx",
                           iig_template="iig_v2", playbook_template="main_single")
    text = (rfc.package_dir / "MANIFEST.md").read_text(encoding="utf-8")
    assert "Artefacts by target system:" in text
    assert f"- **{TARGET_DDL}**: `ACCUM_DDL.txt`" in text
    assert f"- **{TARGET_DML}**:" in text and "`config_inserts_prod.sql`" in text
    assert "| File | Target system | Source artefact | Flags that apply |" in text
    assert f"| `config_inserts_q1.sql` | {TARGET_DML} | framework/config_inserts_q1.sql |" in text
    assert "`dml_unassigned` ×" in text
    assert "Review sheets (people)" in text


# -- §5: template semantics ------------------------------------------------------- #


def _cells(framework_dir: Path, sheet: str, header: str) -> list:
    ws = load_workbook(framework_dir / "config_rows.xlsx")[sheet]
    headers = [c.value for c in ws[1]]
    return ["" if row[headers.index(header)] is None else row[headers.index(header)]
            for row in ws.iter_rows(min_row=2, values_only=True)]


def test_created_by_from_the_faq_rfc_number_in_both_iig_versions(pair1_config, pair1_spec,
                                                                 tmp_path):
    v2 = emit_framework(pair1_spec, FAQ, [], pair1_config, tmp_path / "v2",
                        conventions_profile="acfc_prx", iig_template="iig_v2")
    v2_dir = tmp_path / "v2" / pair1_spec.feed_slug / "framework"
    for sheet in ("DATA_FACTORY_PIPELINE_SCHEDULE", "ADLS_DELTA_INGESTION_DETAILS",
                  "EMAIL_TEMPLATE_CONFIG"):
        assert set(_cells(v2_dir, sheet, "CREATED_BY")) == {"RFCSYN001"}, sheet
        assert set(_cells(v2_dir, sheet, "UPDATED_BY")) == {"RFCSYN001"}, sheet
        assert set(_cells(v2_dir, sheet, "CREATED_DATE")) == {""}      # GETDATE() at insert
    blank = [f for f in v2.flags if f.startswith("iig_blank:")]
    assert not any("CREATED_BY" in f or "UPDATED_BY" in f for f in blank)
    # iig_v1: same rule, including the reference sheet's CRETAED_BY header
    emit_framework(pair1_spec, FAQ, [], pair1_config, tmp_path / "v1",
                   conventions_profile="edo_sfmc", iig_template="iig_v1")
    v1_dir = tmp_path / "v1" / pair1_spec.feed_slug / "framework"
    assert set(_cells(v1_dir, "EMAIL_TEMPLATE_CONFIG", "CRETAED_BY")) == {"RFCSYN001"}
    assert set(_cells(v1_dir, "DATA_FACTORY_PIPELINE_SCHEDULE", "UPDATED_BY")) == {"RFCSYN001"}
    # provenance badge says FAQ
    prov = load_workbook(v2_dir / "config_rows.xlsx")["_provenance"]
    rows = [r for r in prov.iter_rows(min_row=2, values_only=True) if r[2] == "CREATED_BY"]
    assert rows and all(r[3] == "from FAQ" for r in rows)


def test_unanswered_rfc_number_leaves_created_by_blank_and_flagged(pair1_config, pair1_spec,
                                                                   tmp_path):
    framework = emit_framework(pair1_spec, LoadPatternFaq(), [], pair1_config, tmp_path,
                               conventions_profile="acfc_prx", iig_template="iig_v2")
    framework_dir = tmp_path / pair1_spec.feed_slug / "framework"
    assert set(_cells(framework_dir, "DATA_QUALITY_RULES", "CREATED_BY")) == {""}
    blank = [f for f in framework.flags if f.startswith("iig_blank:DATA_QUALITY_RULES")]
    assert blank and "CREATED_BY" in blank[0] and "UPDATED_BY" in blank[0]


def test_claim_type_id_default_knob(pair1_config, pair1_spec, tmp_path):
    default_dir = tmp_path / "default"
    emit_framework(pair1_spec, FAQ, [], pair1_config, default_dir,
                   conventions_profile="acfc_prx", iig_template="iig_v2")
    assert set(_cells(default_dir / pair1_spec.feed_slug / "framework",
                      "ADLS_DELTA_INGESTION_DETAILS", "CLAIM_TYPE_ID")) == {""}   # the golden's
    config = pair1_config.model_copy(update={
        "metadata": pair1_config.metadata.model_copy(update={"claim_type_id_default": "NA"})})
    emit_framework(pair1_spec, FAQ, [], config, tmp_path / "na",
                   conventions_profile="acfc_prx", iig_template="iig_v2")
    assert set(_cells(tmp_path / "na" / pair1_spec.feed_slug / "framework",
                      "ADLS_DELTA_INGESTION_DETAILS", "CLAIM_TYPE_ID")) == {"NA"}
