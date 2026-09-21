"""M8.3 — layout resolution in a real workspace: the answers file, the
unresolved-headers report, the live recognizer posture and the rule that a
runtime profile never lands under fixtures/.

Pair 1 (family A) is the documented case: under synonyms alone its stage /
standard table, column and type roles stay open — the failure recorded in
docs/acfc/RETROFIT_LOG.md §2 on the real workbook."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from codegen import cli
from codegen.config import LayoutConfig, load_config
from codegen.layout.answers import (
    AnswersFileError,
    apply_answers,
    load_answers,
    unresolved_report,
)
from codegen.layout.model import FmapiLayoutProvider, MockLayoutProvider, build_layout_provider
from codegen.layout.resolve import _save_runtime, resolve_workbook

REPO = Path(__file__).resolve().parents[1]
PAIR_1 = REPO / "fixtures" / "acfc_shapes" / "sttm" / "pair_1_family_a.xlsx"
TRUTH = REPO / "fixtures" / "layout_profiles" / "sttm_pair_1_family_a.layout.json"
SHEET = "FEED_1_MAPPING"

ANSWERS_YAML = f"""\
answers:
  - {{sheet: {SHEET}, layer: stage, role: schema, column: "Target Schema Name in DL"}}
  - {{document: sttm, sheet: {SHEET}, layer: stage, role: table, column: "Target Table Name in DL"}}
  - {{document: sttm, sheet: {SHEET}, layer: stage, role: column, column: 22}}
  - {{document: sttm, sheet: {SHEET}, layer: stage, role: target_type, column: W}}
  - {{sheet: {SHEET}, layer: standard, role: schema, column: AA}}
  - {{sheet: {SHEET}, layer: standard, role: table, column: "target table name in dl"}}
  - {{sheet: {SHEET}, layer: standard, role: column, column: "Target_Column_Name_in_DL"}}
  - {{sheet: {SHEET}, layer: standard, role: target_type, column: "Target Data Type in DL"}}
  - {{sheet: NO_SUCH_SHEET, role: table, column: 1}}
"""
ANSWERED = {
    f"{SHEET}/stage/schema": 20, f"{SHEET}/stage/table": 21, f"{SHEET}/stage/column": 22,
    f"{SHEET}/stage/target_type": 23, f"{SHEET}/standard/schema": 27,
    f"{SHEET}/standard/table": 28, f"{SHEET}/standard/column": 29,
    f"{SHEET}/standard/target_type": 30}


@pytest.fixture(scope="module")
def shipped_config():
    return load_config(REPO / "config" / "config.yaml")


@pytest.fixture(scope="module")
def config(shipped_config):
    """The PRE-M9 vocabulary (tests/pre_m9_vocabulary.py): the pair-1 sheet
    leaves its eight required target roles open, as inside ACFC on
    2026-09-21 — the state the answers file exists for."""
    import pre_m9_vocabulary

    return pre_m9_vocabulary.strip(shipped_config)


@pytest.fixture
def answers_file(tmp_path):
    path = tmp_path / "answers.yaml"
    path.write_text(ANSWERS_YAML, encoding="utf-8")
    return path


def _open(config):
    doc, _wb = resolve_workbook(PAIR_1, config, provider=None, use_cache=False)
    return doc


def test_answers_place_every_open_role_with_source_user(config, answers_file):
    doc = _open(config)
    assert len(doc.questions) == 8           # M9.1: the schema is required in both bands
    assert sorted(q.key for q in doc.questions) == sorted(ANSWERED)
    answers, notes = apply_answers(load_answers(answers_file), doc.questions,
                                   {"sttm": PAIR_1.name})
    assert answers["sttm"] == ANSWERED
    assert len(notes) == 1 and "NO_SUCH_SHEET" in notes[0]          # reported, never applied

    resolved, _wb = resolve_workbook(PAIR_1, config, provider=None, use_cache=False,
                                     answers=answers["sttm"])
    assert resolved.questions == []
    truth = {(s["name"], b["layer"]): b["roles"] for s in json.loads(
        TRUTH.read_text(encoding="utf-8"))["sheets"] for b in s.get("bands", [])}
    sheet = next(s for s in resolved.profile.sheets if s.name == SHEET)
    for layer in ("stage", "standard"):
        roles = {str(k): v for k, v in sheet.band(layer).roles.items()}
        assert {k: roles[k] for k in ("schema", "table", "column", "target_type")} == {
            k: truth[(SHEET, layer)][k] for k in ("schema", "table", "column", "target_type")}
    assert {resolved.profile.role_sources[key] for key in answers["sttm"]} == {"user"}


def test_an_answer_may_set_a_role_no_question_asks_for_and_it_wins(shipped_config, tmp_path):
    """M9.1 — inside ACFC an answer for a role the profile never listed as
    open was refused ("matches no open question"). Under the shipped tables
    nothing is open on pair 1; the answers still apply: one confirms a
    synonym placement, one OVERRIDES one (and the role that held the column
    gives it up), all recorded source=user."""
    doc, workbook = resolve_workbook(PAIR_1, shipped_config, provider=None, use_cache=False)
    assert doc.questions == [] and doc.profile.source == "synonyms"
    path = tmp_path / "override.yaml"
    path.write_text(
        "answers:\n"
        f"  - {{sheet: {SHEET}, layer: stage, role: schema, column: T}}\n"          # confirms
        f"  - {{sheet: {SHEET}, role: description, column: \"Data Definition\"}}\n"  # overrides
        f"  - {{sheet: {SHEET}, layer: standard, role: nonsense, column: 1}}\n",
        encoding="utf-8")
    with pytest.raises(AnswersFileError, match="is not a layout role"):
        apply_answers(load_answers(path), doc.questions, {"sttm": PAIR_1.name},
                      documents={"sttm": (doc.profile, workbook)})
    path.write_text("\n".join(path.read_text(encoding="utf-8").splitlines()[:3]) + "\n",
                    encoding="utf-8")
    answers, notes = apply_answers(load_answers(path), doc.questions, {"sttm": PAIR_1.name},
                                   documents={"sttm": (doc.profile, workbook)})
    # "Data Definition" sits in the rules band only: the band is inferred.
    assert answers["sttm"] == {f"{SHEET}/stage/schema": 20, f"{SHEET}/rules/description": 11}
    assert any("confirms column 20 (synonyms)" in n for n in notes)
    assert any("places a role no question asked for, at column 11" in n for n in notes)

    resolved, _ = resolve_workbook(PAIR_1, shipped_config, provider=None, use_cache=False,
                                   answers=answers["sttm"])
    profile = resolved.profile
    assert profile.role_sources[f"{SHEET}/stage/schema"] == "user"
    assert profile.confidence[f"{SHEET}/stage/schema"] == 1.0
    rules = profile.sheet(SHEET).band("rules")
    # The answered column was data_definition's (synonyms): it gave the column up.
    assert rules.roles["description"] == 11 and "data_definition" not in rules.roles
    assert any("column 11 given to 'description' by a user answer" in n for n in profile.notes)
    assert resolved.complete                      # data_definition is not a required role

    # An answer that takes a REQUIRED role's column re-opens that role: it is
    # back in the question list, never silently lost.
    moved, _ = resolve_workbook(PAIR_1, shipped_config, provider=None, use_cache=False,
                                answers={f"{SHEET}/stage/schema": 21})
    assert moved.profile.sheet(SHEET).band("stage").roles["schema"] == 21
    assert [q.key for q in moved.questions] == [f"{SHEET}/stage/table"]
    assert any("moved the role from column 20 (synonyms) to column 21" in n
               for n in moved.profile.notes)


def test_an_ambiguous_entry_is_refused_not_guessed(config, tmp_path):
    path = tmp_path / "a.yaml"
    path.write_text(f"answers:\n  - {{sheet: {SHEET}, role: table, column: 21}}\n",
                    encoding="utf-8")
    with pytest.raises(AnswersFileError, match="add `layer:`"):
        apply_answers(load_answers(path), _open(config).questions, {"sttm": PAIR_1.name})


@pytest.mark.parametrize("body", ["answers:\n  - {sheet: S, role: table}\n",
                                  "answers:\n  - {sheet: S, role: t, column: 1, colour: x}\n",
                                  "answers:\n  - {document: frd, sheet: S, role: t, column: 1}\n",
                                  "gaps:\n  k: text\n", "surprise: 1\n"])
def test_a_malformed_answers_file_names_the_entry(tmp_path, body):
    path = tmp_path / "bad.yaml"
    path.write_text(body, encoding="utf-8")
    with pytest.raises(AnswersFileError):
        load_answers(path)


def test_unresolved_report_has_structure_and_no_data(config):
    doc = _open(config)
    report = unresolved_report(doc.questions, {"sttm": PAIR_1.name})
    assert f"`{SHEET}`" in report and "`stage`" in report and "`table`" in report
    assert "`21: Target Table Name in DL`" in report and "`C: Field Name`" in report
    # No cell BELOW the header row may appear: structural labels only.
    sheet_profile = next(s for s in doc.profile.sheets if s.name == SHEET)
    ws = load_workbook(PAIR_1, data_only=True)[SHEET]
    data_values = {str(v).strip() for row in ws.iter_rows(min_row=sheet_profile.header_row + 1,
                                                          values_only=True)
                   for v in row if v is not None and len(str(v).strip()) >= 6}
    header_values = {str(v).strip() for row in ws.iter_rows(max_row=sheet_profile.header_row,
                                                            values_only=True)
                     for v in row if v is not None}
    leaked = sorted(v for v in data_values - header_values if v in report)
    assert leaked == []


def test_cli_layout_answers_and_report(tmp_path, answers_file, monkeypatch, capsys):
    """A workbook no cache / mock answer knows (one renamed sheet = a new
    fingerprint): the CLI reports the open roles, then resolves them from
    the answers file. The mock posture is pinned — no model is reachable."""
    import pre_m9_vocabulary

    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    monkeypatch.chdir(REPO)
    overlay = tmp_path / "pre_m9_overlay.yaml"           # the vocabulary that leaves roles open
    overlay.write_text(pre_m9_vocabulary.overlay_yaml(
        load_config(REPO / "config" / "config.yaml")), encoding="utf-8")
    monkeypatch.setenv("CODEGEN_CONFIG_OVERLAYS", str(overlay))
    workbook = tmp_path / "STTM_new_shape.xlsx"
    wb = load_workbook(PAIR_1)
    wb[SHEET].title = "FEED_X_MAPPING"
    wb.save(workbook)
    answers = tmp_path / "answers_x.yaml"
    answers.write_text(answers_file.read_text(encoding="utf-8").replace(SHEET, "FEED_X_MAPPING"),
                       encoding="utf-8")
    state = tmp_path / "state"
    state.mkdir()
    monkeypatch.setenv("CODEGEN_STORAGE_STATE", f"local:{state.as_posix()}")
    report = tmp_path / "unresolved_headers.md"

    assert cli.main(["layout", "--workbook", str(workbook), "--dry-run",
                     "--report-unresolved", str(report), "--require-complete"]) == 1
    assert "FEED_X_MAPPING" in report.read_text(encoding="utf-8")
    capsys.readouterr()

    assert cli.main(["layout", "--workbook", str(workbook), "--dry-run",
                     "--answers", str(answers), "--require-complete",
                     "--report-unresolved", str(report)]) == 0
    out = capsys.readouterr().out
    assert "8 answer(s) applied" in out and "Nothing is unresolved." in report.read_text(
        encoding="utf-8")
    # The completed profile went to the STATE role — and nowhere under fixtures/.
    cached = list((state / "layout_profiles").glob("*.json"))
    assert len(cached) == 1 and "FEED_X_MAPPING" in cached[0].read_text(encoding="utf-8")
    tracked = "".join(p.read_text(encoding="utf-8")
                      for p in (REPO / "fixtures" / "layout_profiles").rglob("*.json"))
    assert "FEED_X_MAPPING" not in tracked


# -- profiles with real sheet names never land under fixtures/ --------------------------

def test_runtime_profiles_are_never_written_under_fixtures(tmp_path):
    with pytest.raises(ValueError, match="never written under fixtures"):
        _save_runtime({"fingerprint": "x"}, tmp_path / "fixtures" / "layout_profiles", "x.json")
    assert not (tmp_path / "fixtures").exists()
    for bad in ("fixtures/layout_profiles", "./fixtures/layout_profiles/runtime"):
        with pytest.raises(ValueError, match="fixture"):
            LayoutConfig(runtime_cache_dir=bad)


# -- the live recognizer posture ----------------------------------------------------------

def _with_layout(config, **update):
    return config.model_copy(update={"layout": config.layout.model_copy(update=update)})


def test_live_posture_queries_the_configured_endpoint(config, monkeypatch):
    monkeypatch.delenv("CODEGEN_FORCE_MOCK_LAYOUT", raising=False)
    monkeypatch.delenv("CODEGEN_LAYOUT_PROVIDER", raising=False)
    live = _with_layout(config, provider="live", endpoint="databricks-claude-opus-5")
    # Layer 2 mock-locked: the recognizer's own opt-in still stands ...
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_PROVIDER", "1")
    provider = build_layout_provider(live, dry_run=False, base_dir=REPO)
    assert isinstance(provider, FmapiLayoutProvider)
    assert provider._endpoint == "databricks-claude-opus-5"  # noqa: SLF001
    # ... while dry-run and the recognizer's own lock force the mock.
    assert isinstance(build_layout_provider(live, dry_run=True, base_dir=REPO),
                      MockLayoutProvider)
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    assert isinstance(build_layout_provider(live, dry_run=False, base_dir=REPO),
                      MockLayoutProvider)
    monkeypatch.delenv("CODEGEN_FORCE_MOCK_LAYOUT")
    # env selects the posture for an App (app.yaml) without touching the yaml
    monkeypatch.setenv("CODEGEN_LAYOUT_PROVIDER", "live")
    monkeypatch.setenv("CODEGEN_LAYOUT_ENDPOINT", "another-endpoint")
    via_env = build_layout_provider(config, dry_run=False, base_dir=REPO)
    assert isinstance(via_env, FmapiLayoutProvider)
    assert via_env._endpoint == "another-endpoint"  # noqa: SLF001


def test_live_request_is_the_mock_request_and_is_validated(config, monkeypatch):
    """Same request shape as the mock (fingerprint material only — no data
    row reaches the endpoint), same merge + validation of the answer."""
    import codegen.databricks as seam

    sent: list[dict] = []

    def fake_chat(cfg, messages, endpoint=None, max_tokens=0, client=None):
        sent.append({"endpoint": endpoint, "prompt": messages[-1]["content"]})
        return TRUTH.read_text(encoding="utf-8")

    monkeypatch.setattr(seam, "chat", fake_chat)
    live = FmapiLayoutProvider(_with_layout(config, provider="live"), endpoint="ep-under-test")
    mock = MockLayoutProvider([REPO / "fixtures" / "layout_profiles"])
    via_live, _ = resolve_workbook(PAIR_1, config, provider=live, use_cache=False)
    via_mock, _ = resolve_workbook(PAIR_1, config, provider=mock, use_cache=False)

    assert live.requests == mock.requests and len(sent) == 1
    assert sent[0]["endpoint"] == "ep-under-test"
    assert via_live.questions == [] and via_live.profile.sheets == via_mock.profile.sheets
    ws = load_workbook(PAIR_1, data_only=True)[SHEET]
    header_row = next(s for s in via_live.profile.sheets if s.name == SHEET).header_row
    data_values = {str(v) for row in ws.iter_rows(min_row=header_row + 1, values_only=True)
                   for v in row if v is not None and len(str(v)) >= 8}
    header_values = {str(v) for row in ws.iter_rows(max_row=header_row, values_only=True)
                     for v in row if v is not None}
    assert [v for v in data_values - header_values if v in sent[0]["prompt"]] == []
