"""M9.1b (v0.5.2-acfc) — what the first REAL pair-1 run of v0.5.1 showed
(``docs/acfc/PAIR1_REAL_RUN.md`` / ``RUN_v051_pair1.md`` on ``acfc-runs``):

1. the file-pattern chain: FRD → STTM meta rows / File Details → VDD FILES
   sheet → the ``feeds[0].file_patterns`` question; ``extract-sttm`` no longer
   hard-fails, only ``generate`` stops when the chain ends unanswered;
2. the Object Name cell as a BLOCK (stanzas / labelled lines / a nested
   label | value table), read per line / row;
3. ``codegen layout --refresh`` never reports ``source=cache``;
4. the exit code follows the gate verdict, and a check that could not RUN is
   a flag, not a FAIL.
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from acfc_shapes import frd as frd_fixtures
from acfc_shapes import pair1
from acfc_shapes.common import docx_bytes, table
from codegen import cli
from codegen.extract import contract_to_json as sttm_to_json
from codegen.extract import extract_contract
from codegen.extract.frd_docx import (
    contract_to_json,
    extract_frd_contract,
    is_unnamed_feed,
    parse_object_name_block,
)
from codegen.extract.vdd import contract_to_json as vdd_to_json
from codegen.extract.vdd import extract_vdd_contract
from codegen.gate.preflight import GateCheck, _ruff_check
from codegen.gate.verdict import compute_verdict
from codegen.layout.answers import AnswersFile, apply_answers
from codegen.layout.resolve import parse_answers, resolve_pair
from codegen.resolve.resolver import ContractMismatchError
from codegen.resolve.resolver import resolve_pair as resolve_contracts

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PAIR_1 = SHAPES / "sttm" / "pair_1_family_a.xlsx"
PAIR_1_FRD = SHAPES / "frd" / "f1_pair_1.docx"
PAIR_1_VDD = SHAPES / "vdd" / "pair_1_v1_segments.xlsx"
DATE = "2026-01-01"


# ------------------------------------------------------------ 2. the Object Name block


def test_object_name_stanzas_are_the_verified_real_shape(config):
    """PAIR1_REAL_RUN.md §1: '<LOB>:' on one line, the file pattern on the next."""
    frd = config.extractor.frd
    cell = "LOB_A:\nPFX_A_FEED_YYYYMMDD_HHMMSS.txt\nLOB_B:\nPFX_B_FEED_YYYYMMDD_HHMMSS.txt"
    feed_name, files = parse_object_name_block(cell, None, frd)
    assert feed_name is None
    assert [(r.values["label"], r.key) for r in files] == [
        ("LOB_A", "PFX_A_FEED_YYYYMMDD_HHMMSS.txt"), ("LOB_B", "PFX_B_FEED_YYYYMMDD_HHMMSS.txt")]
    # The tracked pair-1 FRD carries that shape: four stanzas -> four patterns.
    contract, _ = extract_frd_contract(PAIR_1_FRD, config, generated_date=DATE)
    (feed,) = contract.feeds
    assert feed.file_name_patterns == pair1.FILE_PATTERNS
    assert is_unnamed_feed(feed.feed_name)
    flag = next(f for f in contract.extraction_flags
                if f.startswith("file_pattern_from_object_name"))
    assert "('Object Name') cell is a block listing 4 file(s) under ['LOB_A outbound'" in flag


def test_object_name_labelled_lines_give_feed_name_and_files_unverified_shape(config, tmp_path):
    """UNVERIFIED shape (plausible, not seen in a client document): a block
    whose lines are labelled — the feed name from its label's value, the file
    names from every 'File Name'-labelled line; never the raw text."""
    frd = config.extractor.frd
    cell = "Object Name: Accumulators Daily\nFile Name: acc_daily_*.txt\nFile Name: acc_adj_*.txt"
    feed_name, files = parse_object_name_block(cell, None, frd)
    assert feed_name == "Accumulators Daily"
    assert [r.key for r in files] == ["acc_daily_*.txt", "acc_adj_*.txt"]
    contract, _ = extract_frd_contract(_frd_with_object_name(tmp_path, cell), config,
                                       generated_date=DATE)
    (feed,) = contract.feeds
    assert feed.feed_name == "Accumulators Daily" and "\n" not in feed.feed_name
    assert feed.file_name_patterns == ["acc_daily_*.txt", "acc_adj_*.txt"]
    kinds = [f.split(":")[0] for f in contract.extraction_flags]
    assert "frd_object_name_block" in kinds and "frd_feed_name_unstated" not in kinds
    # The layout stage keeps a name the block states (it renames only unnamed feeds).
    pair = resolve_pair(PAIR_1, _frd_with_object_name(tmp_path, cell), config, provider=None,
                        cache_dirs=[], generated_date=DATE)
    assert pair.frd_contract.feeds[0].feed_name == "Accumulators Daily"
    assert not any("feeds[0].feed_name source_used:STTM" in f for f in pair.flags)


def test_object_name_two_row_nested_table_unverified_shape(config, tmp_path):
    """UNVERIFIED shape: a nested label | value table — 'Object Name' row and
    'File Name' row — read per ROW like the multiline block."""
    frd = config.extractor.frd
    nested = [["Object Name", "Accumulators Daily"], ["File Name", "acc_daily_*.txt"]]
    assert parse_object_name_block("", nested, frd)[0] == "Accumulators Daily"
    contract, _ = extract_frd_contract(_frd_with_object_name(tmp_path, nested), config,
                                       generated_date=DATE)
    (feed,) = contract.feeds
    assert (feed.feed_name, feed.file_name_patterns) == ("Accumulators Daily",
                                                         ["acc_daily_*.txt"])
    # A header-row table with ONE body row reads the same way …
    headed = [["Object Name", "File Name"], ["Accumulators Daily", "acc_daily_*.txt"]]
    name, files = parse_object_name_block("", headed, frd)
    assert (name, [r.key for r in files]) == ("Accumulators Daily", ["acc_daily_*.txt"])
    # … while a several-row, one-row-per-FILE table keeps the M7 reading.
    per_file = [["Vendor", "FileName"], ["V", "a_*.csv"], ["V", "b_*.csv"], ["V", "c_*.csv"]]
    assert parse_object_name_block("", per_file, frd) is None


def test_unlabelled_or_prose_blocks_are_never_a_value(config):
    frd = config.extractor.frd
    # A flattened table has no label to read by: refused, as before.
    assert parse_object_name_block("Vendor\nFileName\nVENDOR_A\nx_YYYYMMDD.csv", None, frd) is None
    # Prose over two lines: nothing labelled, nothing file-like.
    assert parse_object_name_block("Accumulator balances\nexchanged daily", None, frd) is None
    # A single line is a scalar, not a block.
    assert parse_object_name_block("Accumulators", None, frd) is None


def _frd_with_object_name(tmp_path: Path, cell) -> Path:
    sections = [
        frd_fixtures._section("Descriptive Metadata", "", "d", "Ingest.", frd_fixtures.DESCRIPTIVE,
                              {"Data Source": "VENDOR_A", "Object Name": cell,
                               "Frequency": "Daily", "LOBs": "ALL"}),
        frd_fixtures._section("Structural Metadata", "", "s", "Ingest.",
                              frd_fixtures.STRUCTURAL_P1,
                              {"Object/data Format": pair1.FILE_FORMAT,
                               "Target Table Name": pair1.FRD_TARGET_BLOCK,
                               "Domain and Subdomain": "Pharmacy / Accumulators",
                               "Load Strategy STG": "Append",
                               "Load Strategy STD (View)": "Append"}),
    ]
    path = tmp_path / f"frd_{abs(hash(str(cell))) % 10_000}.docx"
    path.write_bytes(docx_bytes(sections))
    return path


def test_the_docx_table_helper_writes_nested_tables():       # guards the fixture above
    assert "<w:tbl>" in table([["a", [["x", "y"]]]]).split("<w:tc>", 2)[2]


# ------------------------------------------------------------ 1. the file-pattern chain


def _silent_frd(tmp_path: Path) -> Path:
    """The pair-1 FRD with an Object Name that lists NO file (the v0.5.1 run's
    situation: nothing in the FRD yields a pattern)."""
    return _frd_with_object_name(tmp_path, "")


def _vdd_without_patterns(tmp_path: Path) -> Path:
    path = tmp_path / "vdd_no_patterns.xlsx"
    shutil.copyfile(PAIR_1_VDD, path)
    wb = load_workbook(path)
    for row in wb["FILES"].iter_rows(min_row=2):
        row[0].value = None
    wb.save(path)
    return path


def _contracts(tmp_path: Path, config, pair) -> tuple[Path, Path]:
    frd_json, sttm_json = tmp_path / "frd.contract.json", tmp_path / "sttm.contract.json"
    frd_json.write_text(contract_to_json(pair.frd_contract), encoding="utf-8")
    contract = extract_contract(PAIR_1, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    return frd_json, sttm_json


def test_extract_sttm_no_longer_hard_fails_without_a_file_pattern(config, tmp_path):
    pair = resolve_pair(PAIR_1, _silent_frd(tmp_path), config, provider=None, cache_dirs=[],
                        generated_date=DATE)                 # no VDD: nothing states a pattern
    assert pair.frd_contract.feeds[0].file_name_patterns == []
    _frd_json, sttm_json = _contracts(tmp_path, config, pair)
    feed = json.loads(sttm_json.read_text(encoding="utf-8"))["feeds"][0]
    assert feed["source_file"]["name_pattern"] is None and feed["field_count"] == 23
    notes = json.loads(sttm_json.read_text(encoding="utf-8"))["notes"]
    assert any("left open for the VDD FILES sheet / the file_patterns answer" in n for n in notes)


def test_chain_third_link_is_the_vdd_files_sheet_in_the_layout_stage_and_the_resolver(
        config, tmp_path):
    # Layout stage (the UI / `layout --frd-contract-out` path): filled, flagged, shown.
    pair = resolve_pair(PAIR_1, _silent_frd(tmp_path), config, provider=None, cache_dirs=[],
                        vdd_path=PAIR_1_VDD, generated_date=DATE)
    assert pair.questions == []
    assert pair.frd_contract.feeds[0].file_name_patterns == pair1.FILE_PATTERNS
    (flag,) = [f for f in pair.flags if f.startswith("file_pattern_from_vdd")]
    assert "FILES!A2" in flag and "the FRD and the STTM name no file pattern" in flag
    fill = next(f for f in pair.gap_fills if f["field"] == "feeds[0].file_patterns")
    assert fill["source"] == "VDD" and "FILES!A" in fill["cell"]
    # Resolver (the pure-CLI path: extract-frd, no layout stage): the same link.
    no_vdd = resolve_pair(PAIR_1, _silent_frd(tmp_path), config, provider=None, cache_dirs=[],
                          generated_date=DATE)
    frd_json, sttm_json = _contracts(tmp_path, config, no_vdd)
    vdd_json = tmp_path / "vdd.contract.json"
    vdd_contract, _ = extract_vdd_contract(PAIR_1_VDD, config, generated_date=DATE)
    vdd_json.write_text(vdd_to_json(vdd_contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config, vdd_path=vdd_json)
    assert spec.file_name_patterns == pair1.FILE_PATTERNS
    (flag,) = [f for f in spec.provenance_flags if f.startswith("file_pattern_from_vdd")]
    assert "VDD FILES sheet states" in flag and "row(s) [2, 3, 4, 5]" in flag


def test_chain_second_link_is_still_the_sttm(config, tmp_path):
    """STTM meta rows state the names (the documented-shape sheet): taken by the
    extractor / resolver, flagged file_pattern_from_sttm — never the VDD, never asked."""
    sttm = tmp_path / "sttm_names_stated.xlsx"
    shutil.copyfile(PAIR_1, sttm)
    wb = load_workbook(sttm)
    wb["FEED_1_MAPPING"].cell(row=1, column=2, value="; ".join(pair1.FILE_PATTERNS[:2]))
    wb.save(sttm)
    pair = resolve_pair(sttm, _silent_frd(tmp_path), config, provider=None, cache_dirs=[],
                        vdd_path=PAIR_1_VDD, generated_date=DATE)
    assert pair.questions == [] and pair.frd_contract.feeds[0].file_name_patterns == []
    frd_json = tmp_path / "frd.contract.json"
    frd_json.write_text(contract_to_json(pair.frd_contract), encoding="utf-8")
    sttm_json = tmp_path / "sttm.contract.json"
    sttm_json.write_text(sttm_to_json(extract_contract(
        sttm, frd_json, config, generated_date=DATE, layout=pair.sttm.profile)), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config)
    assert spec.file_name_patterns == pair1.FILE_PATTERNS[:2]
    assert [f.split(":")[0] for f in spec.provenance_flags].count("file_pattern_from_sttm") == 1
    assert not any(f.startswith("file_pattern_from_vdd") for f in spec.provenance_flags)


def test_chain_ends_in_the_file_patterns_question_and_the_answer_lands_flagged(config, tmp_path):
    frd = _silent_frd(tmp_path)
    vdd = _vdd_without_patterns(tmp_path)
    pair = resolve_pair(PAIR_1, frd, config, provider=None, cache_dirs=[], vdd_path=vdd,
                        generated_date=DATE)
    (question,) = pair.questions
    assert question.key == "feeds[0].file_patterns" and question.kind == "text"
    assert question.candidates == [] and "no document names a file pattern" in question.reason
    # The answers file addresses it under `gaps:` …
    answers, notes = apply_answers(
        AnswersFile(gaps={"feeds[0].file_patterns": {"value": "a_*.txt; b_*.txt"}}),
        pair.questions, {"sttm": PAIR_1.name})
    assert notes == [] and answers["gaps"]["feeds[0].file_patterns"]["value"] == "a_*.txt; b_*.txt"
    # … and so does the dialog.
    resolved = resolve_pair(PAIR_1, frd, config, provider=None, cache_dirs=[], vdd_path=vdd,
                            generated_date=DATE, answers=parse_answers({"gaps": answers["gaps"]}))
    assert resolved.questions == []
    assert resolved.frd_contract.feeds[0].file_name_patterns == ["a_*.txt", "b_*.txt"]
    (flag,) = [f for f in resolved.flags if f.startswith("file_pattern_from_user")]
    assert "the person stated ['a_*.txt', 'b_*.txt']" in flag
    assert flag in resolved.frd_contract.extraction_flags       # travels with the contract


def test_the_hard_stop_remains_only_at_generate_time(config, tmp_path, capsys, monkeypatch):
    pair = resolve_pair(PAIR_1, _silent_frd(tmp_path), config, provider=None, cache_dirs=[],
                        generated_date=DATE)
    frd_json, sttm_json = _contracts(tmp_path, config, pair)      # extraction succeeded
    with pytest.raises(ContractMismatchError, match="no source names a file pattern") as exc:
        resolve_contracts(frd_json, sttm_json, config)
    assert "feeds[0].file_patterns" in str(exc.value) and "generate --vdd" in str(exc.value)
    monkeypatch.setenv("CODEGEN_STORAGE_OUTPUTS", f"local:{(tmp_path / 'out').as_posix()}")
    (tmp_path / "out").mkdir()
    monkeypatch.chdir(REPO)
    assert cli.main(["generate", "--frd-contract", str(frd_json), "--sttm-contract",
                     str(sttm_json), "--skip-tests", "--dry-run"]) == 1
    assert "no source names a file pattern" in capsys.readouterr().out


# ------------------------------------------------------------ 3. --refresh never says cache


def test_cli_layout_refresh_with_an_answers_file_never_reports_source_cache(tmp_path, monkeypatch,
                                                                            capsys):
    """The v0.5.1 run: `layout --refresh --answers` printed `source=cache
    cache_hit=True` for the STTM and the FRD — the second pass (answers) re-read
    the entries the refresh pass had just written."""
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    monkeypatch.setenv("CODEGEN_STORAGE_STATE", f"local:{(tmp_path / 'state').as_posix()}")
    (tmp_path / "state").mkdir()
    monkeypatch.chdir(REPO)
    answers = tmp_path / "answers.yaml"
    answers.write_text(
        "answers:\n"
        "  - {sheet: FEED_1_MAPPING, layer: stage, role: schema, column: T}\n"
        "  - {sheet: FEED_1_MAPPING, layer: standard, role: schema, column: AA}\n"
        "gaps:\n  \"feeds[0].file_format\": {value: \"Fixed Width\"}\n", encoding="utf-8")
    argv = ["layout", "--workbook", str(PAIR_1), "--frd", str(PAIR_1_FRD), "--vdd",
            str(PAIR_1_VDD), "--dry-run", "--answers", str(answers), "--require-complete"]
    assert cli.main([*argv, "--refresh"]) == 0
    out = capsys.readouterr().out
    assert "REFRESH" in out and "2 answer(s) applied" in out
    assert "source=cache" not in out and "cache_hit=True" not in out
    sttm_line = next(line for line in out.splitlines() if line.startswith("STTM"))
    assert "source=user" in sttm_line and "'user': 2" in sttm_line and "'cache': 0" in sttm_line
    assert "'synonyms': 25" in sttm_line                     # every other role: who placed it
    frd_line = next(line for line in out.splitlines() if line.startswith("FRD "))
    assert "source=synonyms" in frd_line and "'cache': 0" in frd_line
    assert "PROVIDER        mock, 0 call(s)" in out
    # The refresh DID write the entries: a plain run afterwards is a cache hit —
    # and only then is `source=cache` the truth.
    assert cli.main(argv[:-3] + ["--require-complete"]) == 0
    assert "cache_hit=True" in capsys.readouterr().out


def test_second_pass_continues_from_the_first_without_a_second_model_call(config, tmp_path):
    import pre_m9_vocabulary
    from codegen.layout.model import MockLayoutProvider

    profiles = REPO / "fixtures" / "layout_profiles"
    provider = MockLayoutProvider([profiles / "mock", profiles])
    pre_m9 = pre_m9_vocabulary.strip(config)
    kwargs = dict(provider=provider, cache_dirs=[], runtime_cache_dir=tmp_path / "rt",
                  generated_date=DATE)
    first = resolve_pair(PAIR_1, PAIR_1_FRD, pre_m9, refresh=True, **kwargs)
    assert first.provider_calls == 1
    second = resolve_pair(PAIR_1, PAIR_1_FRD, pre_m9, refresh=True, prior=first,
                          answers={"sttm": {"FEED_1_MAPPING/stage/schema": 20}}, **kwargs)
    assert len(provider.requests) == 1                      # no second call
    assert not second.sttm.cache_hit and second.sttm.profile.source == "user"
    sources = second.sttm.sources
    assert sources["cache"] == 0 and sources["user"] == 1 and sources["model"] == 8


# ------------------------------------------------------------ 4. exit code = gate verdict


def _generate(tmp_path, monkeypatch, capsys, *extra: str) -> tuple[int, str]:
    monkeypatch.setenv("CODEGEN_STORAGE_OUTPUTS", f"local:{(tmp_path / 'out').as_posix()}")
    monkeypatch.setenv("CODEGEN_CONFIG_OVERLAYS", str(SHAPES / "pair_1" / "config_overlay.yaml"))
    (tmp_path / "out").mkdir(exist_ok=True)
    monkeypatch.chdir(REPO)
    from codegen.config import load_config

    config = load_config(REPO / "config" / "config.yaml")
    pair = resolve_pair(PAIR_1, PAIR_1_FRD, config, provider=None, cache_dirs=[],
                        vdd_path=PAIR_1_VDD, generated_date=DATE)
    frd_json, sttm_json = _contracts(tmp_path, config, pair)
    capsys.readouterr()
    code = cli.main(["generate", "--frd-contract", str(frd_json), "--sttm-contract",
                     str(sttm_json), "--output-mode", "framework", "--profile", "acfc_prx",
                     "--iig-template", "iig_v2", "--skip-tests", "--dry-run", *extra])
    return code, capsys.readouterr().out


def _report_verdict(tmp_path: Path) -> str:
    (report,) = (tmp_path / "out" / "reports").glob("*.md")
    line = next(li for li in report.read_text(encoding="utf-8").splitlines()
                if li.startswith("**Verdict:"))
    return line.strip("*").split(": ")[1]


def test_exit_code_zero_for_pass_with_flags_and_the_headline_is_the_report_verdict(
        tmp_path, monkeypatch, capsys):
    code, out = _generate(tmp_path, monkeypatch, capsys)
    assert code == 0
    assert out.startswith("PASS_WITH_FLAGS vnd_p_accum_client") and "ruff=ok" in out
    assert _report_verdict(tmp_path) == "PASS_WITH_FLAGS"


def test_exit_code_one_only_when_the_verdict_is_fail_and_the_console_says_why(
        tmp_path, monkeypatch, capsys):
    from codegen import cli as cli_module

    real = cli_module.run_preflight

    def with_a_finding(feed_dir, config):
        return [GateCheck(name="ruff", passed=False,
                          details="1 finding(s)\npipeline/audit.py:3:1: F401 unused import"),
                *[c for c in real(feed_dir, config) if c.name != "ruff"]]

    monkeypatch.setattr(cli_module, "run_preflight", with_a_finding)
    code, out = _generate(tmp_path, monkeypatch, capsys)
    assert code == 1 and out.startswith("FAIL ") and "ruff=FAIL" in out
    assert "CHECK FAILED    ruff — 1 finding(s)" in out
    assert "pipeline/audit.py:3:1: F401 unused import" in out
    assert _report_verdict(tmp_path) == "FAIL"               # headline == report == exit code


def test_a_check_that_could_not_run_is_a_flag_not_a_fail(tmp_path, monkeypatch, capsys):
    """`ruff=FAIL` inside ACFC said nothing about WHY. A tool that does not run
    (no binary in the environment, a config it rejects) found nothing AND
    checked nothing: PASS cannot be claimed — but it is not a finding."""
    from codegen import cli as cli_module

    real = cli_module.run_preflight

    def ruff_missing(feed_dir, config):
        return [GateCheck(name="ruff", passed=True, not_run=True,
                          details="ruff did not run (exit 1): No module named ruff"),
                *[c for c in real(feed_dir, config) if c.name != "ruff"]]

    monkeypatch.setattr(cli_module, "run_preflight", ruff_missing)
    code, out = _generate(tmp_path, monkeypatch, capsys)
    assert code == 0 and out.startswith("PASS_WITH_FLAGS ") and "ruff=not-run" in out
    assert "CHECK NOT RUN   ruff — ruff did not run (exit 1): No module named ruff" in out
    report = next((tmp_path / "out" / "reports").glob("*.md")).read_text(encoding="utf-8")
    assert "| ruff | NOT RUN |" in report
    assert "check_not_run:ruff — ruff did not run (exit 1): No module named ruff" in report


def test_ruff_check_tells_findings_from_a_tool_that_did_not_run(tmp_path, monkeypatch):
    clean = tmp_path / "clean"
    clean.mkdir()
    (clean / "ok.py").write_text('"""Fine."""\n', encoding="utf-8")
    assert _ruff_check(clean) == GateCheck(name="ruff", passed=True, details="ruff clean")
    dirty = tmp_path / "dirty"
    dirty.mkdir()
    (dirty / "bad.py").write_text("import os\n", encoding="utf-8")
    finding = _ruff_check(dirty)
    assert not finding.passed and not finding.not_run
    assert finding.details.splitlines()[0] == "1 finding(s)"
    assert finding.details.splitlines()[1].startswith("bad.py:1:8: F401 ")   # relative path
    # The tool cannot run: not a finding.
    import subprocess

    from codegen.gate import preflight

    def no_ruff(*_a, **_k):
        return subprocess.CompletedProcess([], 1, stdout="",
                                           stderr="/usr/bin/python: No module named ruff")

    monkeypatch.setattr(preflight.subprocess, "run", no_ruff)
    missing = _ruff_check(dirty)
    assert missing.passed and missing.not_run and "No module named ruff" in missing.details
    gate = compute_verdict("feed", [], [], [missing], tests_skipped=False)
    assert gate.verdict == "PASS_WITH_FLAGS"
    assert gate.flags == ["check_not_run:ruff — ruff did not run (exit 1): /usr/bin/python: No "
                          "module named ruff; the generated code was NOT checked — PASS cannot "
                          "be claimed"]


def test_a_lint_dirty_layer2_sketch_is_a_review_artifact_never_a_gate_input(
        tmp_path, monkeypatch, capsys):
    """Layer-2 sketches (mock or live) land in candidates/candidates.json — a
    review artifact the ruff gate never reads: a sketch that would not even
    parse changes neither the ruff check, the verdict nor the exit code."""
    from codegen.reasoning.providers import mock as mock_module

    original = mock_module.MockProvider.complete

    def dirty(self, pack):
        return original(self, pack).model_copy(update={
            "code_candidate": "import os, sys\ndef broken(:\n    print( 'x' )\n"})

    monkeypatch.setattr(mock_module.MockProvider, "complete", dirty)
    code, out = _generate(tmp_path, monkeypatch, capsys)
    assert code == 0 and out.startswith("PASS_WITH_FLAGS ") and "ruff=ok" in out
    candidates = next((tmp_path / "out").rglob("candidates.json")).read_text(encoding="utf-8")
    assert "def broken(:" in candidates                      # the sketch IS in the artifact
