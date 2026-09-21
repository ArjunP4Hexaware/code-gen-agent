"""M9.2 (v0.5.3-acfc) — the fixed-width WIDTH chain.

The first real run of v0.5.2 (``docs/acfc/RUN_v052_pair1.md`` on
``acfc-runs``) stopped before the gate: six Detail amount fields carry
``length='10,2'`` — a precision, which the fixed-width template cannot read as
a byte width. The chain (``codegen.resolve.widths``): STTM integer length →
STTM end−start+1 → VDD end−start+1 (normalized field name + segment) → the
question ``feeds[0].fields[<name>].width``. A width is NEVER derived from a
precision.

The six fields are a VARIANT of the pair-1 documents (``build_pair1
(amounts=True)``), never the tracked acceptance pair. UNVERIFIED: the VDD span
of these fields (13 bytes here) has not been captured from the real dictionary.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from acfc_shapes import pair1
from acfc_shapes import sttm as sttm_fixtures
from acfc_shapes import vdd as vdd_fixtures
from acfc_shapes.common import xlsx_bytes
from codegen import cli
from codegen.emit.context import TemplateGapError, build_context
from codegen.extract import contract_to_json as sttm_to_json
from codegen.extract import extract_contract
from codegen.extract.frd_docx import contract_to_json as frd_to_json
from codegen.extract.vdd import contract_to_json as vdd_to_json
from codegen.extract.vdd import extract_vdd_contract
from codegen.gate.vdd_check import vdd_cross_check
from codegen.layout.answers import AnswersFile, apply_answers, unresolved_report
from codegen.layout.resolve import parse_answers, resolve_pair
from codegen.resolve.resolver import resolve_pair as resolve_contracts
from codegen.resolve.widths import (
    as_integer,
    as_precision,
    parse_width_answers,
    sttm_width,
    width_key,
)

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PAIR_1_FRD = SHAPES / "frd" / "f1_pair_1.docx"
DATE = "2026-01-01"
AMOUNT_LABELS = [" ".join(name.split()) for name, _c, _s in pair1.AMOUNT_FIELDS]


# ------------------------------------------------------------ the rules


@pytest.mark.parametrize("cell,integer,precision", [
    ("13", 13, None), ("013", 13, None), ("13.0", 13, None), (" 8 ", 8, None),
    ("10,2", None, "Decimal(10,2)"), ("10.2", None, "Decimal(10,2)"),
    ("Decimal(10,2)", None, "Decimal(10,2)"), ("NUMBER (10, 2)", None, "Decimal(10,2)"),
    ("ten", None, None), ("", None, None), (None, None, None),
])
def test_a_length_cell_is_a_width_only_when_it_is_an_integer(cell, integer, precision):
    assert as_integer(cell) == integer
    assert as_precision(cell) == precision


def test_a_width_is_never_derived_from_a_precision():
    # "10,2" could occupy 10, 11, 12 or 13 bytes — only a document can say.
    assert sttm_width("10,2", "220", None) == (None, None)
    assert sttm_width("Decimal(10,2)", "220", "") == (None, None)
    assert sttm_width("10,2", "220", "232") == (13, "sttm_span")
    assert sttm_width("13", "220", "240") == (13, "length")          # the length wins
    assert sttm_width(None, "220", "232") == (13, "sttm_span")
    assert sttm_width("10,2", "232", "220") == (None, None)          # an inverted span states none


def test_width_answers_must_be_positive_integers():
    key = width_key(0, "Amount\n(01)")
    assert key == "feeds[0].fields[Amount (01)].width"
    assert parse_width_answers({key: {"value": "13"}, "feeds[0].file_format": {"value": "x"}}) == {
        0: {"amount01": 13}}
    for bad in ("10,2", "0", "thirteen"):
        with pytest.raises(ValueError, match="not a positive integer"):
            parse_width_answers({key: {"value": bad}})
    with pytest.raises(ValueError, match="never a precision"):
        parse_answers({"gaps": {key: {"value": "10,2"}}})            # the UI route's validator


# ------------------------------------------------------------ the variant pair


def _variant(tmp_path: Path, *, amount_end: bool = False, vdd_spans: bool = True,
             vdd_amounts: bool = True) -> tuple[Path, Path]:
    sttm, vdd = tmp_path / "sttm_amounts.xlsx", tmp_path / "vdd_amounts.xlsx"
    sttm.write_bytes(xlsx_bytes(sttm_fixtures.build_pair1(amounts=True, amount_end=amount_end)))
    vdd.write_bytes(xlsx_bytes(vdd_fixtures.build_vdd_pair1(amounts=vdd_amounts,
                                                            amount_spans=vdd_spans)))
    return sttm, vdd


def _spec(tmp_path: Path, config, sttm: Path, vdd: Path, answers: dict | None = None):
    pair = resolve_pair(sttm, PAIR_1_FRD, config, provider=None, cache_dirs=[], vdd_path=vdd,
                        generated_date=DATE, answers=answers)
    frd_json, sttm_json, vdd_json = (tmp_path / n for n in ("frd.json", "sttm.json", "vdd.json"))
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile, width_answers=pair.width_answers)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    vdd_contract, _ = extract_vdd_contract(vdd, config, generated_date=DATE)
    vdd_json.write_text(vdd_to_json(vdd_contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config, vdd_path=vdd_json)
    return pair, contract, spec


def _context(spec, config):
    return build_context(spec, config, allow_duplicate_file_name=True,
                         duplicate_rule_text=None, notification_rule_texts=[])


def _amounts(spec):
    return [f for s in spec.segments for f in s.fields if f.stage_column.startswith("AMOUNT_")]


def test_the_variant_has_the_captured_shape(tmp_path):
    sttm, vdd = _variant(tmp_path)
    ws = load_workbook(sttm)["FEED_1_MAPPING"]
    rows = [r for r in ws.iter_rows(min_row=16, values_only=True)
            if isinstance(r[2], str) and r[2].startswith("Amount")]
    assert [(r[2], r[5], r[6], r[7]) for r in rows] == [
        (name, start, "10,2", None) for name, _col, start in pair1.AMOUNT_FIELDS]
    assert [start for _n, _c, start in pair1.AMOUNT_FIELDS] == [220, 246, 272, 298, 324, 350]
    assert {r[1] for r in rows} == {"DET"}
    dictionary = load_workbook(vdd)["FEED_1 Fields"]
    spans = [(r[1], r[3], r[4]) for r in dictionary.iter_rows(min_row=2, values_only=True)
             if str(r[1]).startswith("AMOUNT")]
    assert spans == [(f"AMOUNT ({n:02d})", 220 + 26 * (n - 1), 232 + 26 * (n - 1))
                     for n in range(1, 7)]                          # end - start + 1 = 13


def test_vdd_span_resolves_the_width_unverified_vdd_shape(config, tmp_path):
    """Link 3. UNVERIFIED: the 13-byte VDD span is a stand-in until the real
    dictionary's rows for these fields are captured."""
    sttm, vdd = _variant(tmp_path)
    pair, contract, spec = _spec(tmp_path, config, sttm, vdd)
    assert pair.questions == []                                     # the VDD answers: never asked
    # The extractor keeps the cell verbatim, the precision as precision, no width.
    extracted = [f for f in contract.feeds[0].fields if f.stage_column.startswith("AMOUNT_")]
    assert [(f.source_length, f.source_precision, f.source_width) for f in extracted] == [
        ("10,2", "Decimal(10,2)", None)] * 6
    flags = contract.feeds[0].extraction_flags
    assert [f.split(" — ")[0] for f in flags if f.startswith("length_is_precision")] == [
        f"length_is_precision:{label}" for label in AMOUNT_LABELS]
    assert "STTM FEED_1_MAPPING!G39 reads '10,2'" in next(
        f for f in flags if f.startswith("length_is_precision:Amount (01)"))
    # The resolver takes the VDD span — by normalized name ("Amount\n(01)" ==
    # "AMOUNT (01)") and segment — and cites its cells.
    assert [(f.byte_width, f.source_width) for f in _amounts(spec)] == [(13, 13)] * 6
    width_flags = [f for f in spec.provenance_flags if f.startswith("width_from_")]
    assert [f.split(" — ")[0] for f in width_flags] == [
        f"width_from_vdd:{label}" for label in AMOUNT_LABELS]
    assert width_flags[0] == (
        "width_from_vdd:Amount (01) — VDD FEED_1 Fields!D25 / FEED_1 Fields!E25: end 232 - start "
        "220 + 1 = 13 (matched by field name + segment 'Detail')")
    # Fields whose length IS an integer are untouched: no flag, no source_width.
    plain = next(f for s in spec.segments for f in s.fields if f.stage_column == "CARDHOLDER_ID")
    assert (plain.source_width, plain.byte_width, plain.source_precision) == (None, 15, None)


def test_template_handler_rows_and_cross_check_use_the_resolved_width(pair1_config, tmp_path):
    config = pair1_config
    sttm, vdd = _variant(tmp_path)
    _pair, _contract, spec = _spec(tmp_path, config, sttm, vdd)
    # 1. the fixed-width template's positions
    context = _context(spec, config)
    detail = next(s for s in context["segments"] if s["name"] == "Detail")
    assert detail["positions"][-6:] == [(name, start, 13)
                                        for name, _col, start in pair1.AMOUNT_FIELDS]
    # 2. the ADLS_FIXED_WIDTH_HANDLER rows: the SAME value, the golden's own untouched
    scoped = config.model_copy(update={"output": config.output.model_copy(update={
        "dir": str(tmp_path / "out"), "reports_dir": str(tmp_path / "reports")})})
    gate = cli._generate_feed(spec, scoped, dry_run=True, skip_tests=True,
                              output_mode="framework", conventions_profile="acfc_prx",
                              iig_template="iig_v2")
    assert gate.verdict == "PASS_WITH_FLAGS" and all(c.passed for c in gate.checks)
    sheet = load_workbook(tmp_path / "out" / spec.feed_slug / "framework" / "config_rows.xlsx")[
        "ADLS_FIXED_WIDTH_HANDLER"]
    header = [c.value for c in sheet[1]]
    rows = {r[header.index("SEGMENT")]: dict(zip(header, r, strict=True))
            for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert rows["DET"]["LEN"] == "2,15,16,17,18,19,20,21,22,23,24,25,13,13,13,13,13,13"
    assert rows["DET"]["start_ind"].endswith(",198,220,246,272,298,324,350")
    assert "10,2" not in rows["DET"]["LEN"]
    assert rows["HDDR"]["LEN"] == "2,1,1,8,6,10,20,2"               # as the golden
    # 3. the VDD cross-check compares WIDTHS: no mismatch here …
    assert not [f for f in gate.flags if f.startswith("vdd_")]
    # … and a real one is caught on the resolved width, never on "10,2".
    narrowed = [f.model_copy(update={"source_width": 12}) if f.stage_column == "AMOUNT_01" else f
                for f in spec.detail_segment.fields]
    segments = [s.model_copy(update={"fields": narrowed}) if s.segment == "Detail" else s
                for s in spec.segments]
    flags, _check = vdd_cross_check(spec.model_copy(update={"segments": segments}), config)
    (mismatch,) = [f for f in flags if f.startswith("vdd_mismatch:length")]
    assert "(length '12 (resolved width)')" in mismatch and "(length 13)" in mismatch


def test_sttm_span_is_the_second_link_and_wins_over_the_vdd(config, tmp_path):
    sttm, vdd = _variant(tmp_path, amount_end=True)
    pair, contract, spec = _spec(tmp_path, config, sttm, vdd)
    assert pair.questions == []
    flags = contract.feeds[0].extraction_flags
    span = next(f for f in flags if f.startswith("width_from_sttm_span:Amount (01)"))
    assert span == ("width_from_sttm_span:Amount (01) — STTM FEED_1_MAPPING!F39 / "
                    "FEED_1_MAPPING!H39: end 232 - start 220 + 1 = 13")
    assert sum(f.startswith("width_from_sttm_span") for f in flags) == 6
    assert sum(f.startswith("length_is_precision") for f in flags) == 6   # still said
    assert [f.byte_width for f in _amounts(spec)] == [13] * 6
    assert not [f for f in spec.provenance_flags if f.startswith("width_from_vdd")]


def test_the_question_is_the_last_link_and_the_answer_lands_flagged(config, tmp_path):
    sttm, vdd = _variant(tmp_path, vdd_spans=False)
    pair = resolve_pair(sttm, PAIR_1_FRD, config, provider=None, cache_dirs=[], vdd_path=vdd,
                        generated_date=DATE)
    keys = [q.key for q in pair.questions]
    assert keys == [f"feeds[0].fields[{label}].width" for label in AMOUNT_LABELS]
    question = pair.questions[0]
    assert (question.kind, question.document, question.candidates) == ("text", "sttm", [])
    assert "FEED_1_MAPPING!G39 reads '10,2'" in question.reason
    assert "a precision is never a width" in question.reason
    # The report that may leave the workspace names the question, not the field.
    report = unresolved_report(pair.questions, {"sttm": sttm.name})
    assert "feeds[0].fields[<field name>].width" in report and "Amount" not in report
    # Answers file (`gaps:`) -> the dialog's payload -> the extractor -> the resolver.
    answers, notes = apply_answers(
        AnswersFile(gaps={key: {"value": "13"} for key in keys}), pair.questions,
        {"sttm": sttm.name})
    assert notes == [] and sorted(answers["gaps"]) == sorted(keys)
    resolved, contract, spec = _spec(tmp_path, config, sttm, vdd,
                                     answers=parse_answers({"gaps": answers["gaps"]}))
    assert resolved.questions == []
    assert {f.width_answer for f in contract.feeds[0].fields
            if f.stage_column.startswith("AMOUNT_")} == {13}
    assert [f.byte_width for f in _amounts(spec)] == [13] * 6
    user = [f for f in spec.provenance_flags if f.startswith("width_from_user")]
    assert len(user) == 6 and "feeds[0].fields[Amount (01)].width = 13" in user[0]
    assert "no VDD span" in user[0]


def test_unanswered_the_stop_is_at_generate_and_names_the_question(config, tmp_path):
    sttm, vdd = _variant(tmp_path, vdd_amounts=False)       # the dictionary lacks the fields
    pair, _contract, spec = _spec(tmp_path, config, sttm, vdd)      # extraction succeeds
    assert len(pair.questions) == 6
    assert [f.byte_width for f in _amounts(spec)] == [None] * 6     # never 10, never 12
    with pytest.raises(TemplateGapError) as exc:
        _context(spec, config)
    message = str(exc.value)
    assert "length='10,2'" in message and "feeds[0].fields[Amount (01)].width" in message


def test_do_not_map_rows_never_ask_for_a_width(config, tmp_path):
    sttm, vdd = _variant(tmp_path, vdd_spans=False)
    wb = load_workbook(sttm)
    ws = wb["FEED_1_MAPPING"]
    for row in ws.iter_rows(min_row=16):
        if row[21].value == "Do Not Map":
            row[6].value = "10,2"                                   # its Length: a precision too
    wb.save(sttm)
    pair = resolve_pair(sttm, PAIR_1_FRD, config, provider=None, cache_dirs=[], vdd_path=vdd,
                        generated_date=DATE)
    assert len(pair.questions) == 6 and not any("Reserved" in q.key for q in pair.questions)


def test_cli_extract_sttm_carries_width_answers_from_the_answers_file(config, tmp_path,
                                                                      monkeypatch, capsys):
    sttm, vdd = _variant(tmp_path, vdd_spans=False)
    pair = resolve_pair(sttm, PAIR_1_FRD, config, provider=None, cache_dirs=[], vdd_path=vdd,
                        generated_date=DATE)
    frd_json = tmp_path / "frd.contract.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    answers = tmp_path / "answers.yaml"
    answers.write_text("gaps:\n" + "".join(
        f'  "feeds[0].fields[{label}].width": {{value: 13}}\n' for label in AMOUNT_LABELS),
        encoding="utf-8")                                           # YAML integers are fine
    monkeypatch.setenv("CODEGEN_STORAGE_STATE", f"local:{(tmp_path / 'state').as_posix()}")
    (tmp_path / "state").mkdir()
    monkeypatch.chdir(REPO)
    out = tmp_path / "sttm.contract.json"
    assert cli.main(["extract-sttm", "--workbook", str(sttm), "--frd-contract", str(frd_json),
                     "--out", str(out), "--answers", str(answers),
                     "--generated-date", DATE]) == 0
    fields = json.loads(out.read_text(encoding="utf-8"))["feeds"][0]["fields"]
    amounts = [f for f in fields if f["stage_column"].startswith("AMOUNT_")]
    assert [(f["source_length"], f["source_precision"], f["width_answer"]) for f in amounts] == [
        ("10,2", "Decimal(10,2)", 13)] * 6
    # … and the tracked pair's contract carries none of the new keys.
    plain = next(f for f in fields if f["stage_column"] == "CARDHOLDER_ID")
    assert not {"source_width", "source_precision", "width_answer"} & set(plain)


def test_the_tracked_pair_is_untouched_by_the_chain(pair1_spec):
    assert not [f for f in pair1_spec.provenance_flags
                if f.startswith(("width_from_", "length_is_precision", "length_not_a_width"))]
    assert all(f.source_width is None and f.source_precision is None
               for s in pair1_spec.segments for f in s.fields)
