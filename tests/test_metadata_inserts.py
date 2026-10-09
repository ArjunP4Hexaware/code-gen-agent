"""Multi-table step 6 — ``framework/metadata_inserts.sql`` and the per-table
CREATE reference text (docs/acfc/MULTI_TABLE_DESIGN.md rules 1, 7, 8; §6 step 6).

The client's "DDL" is the metadata INSERT rows that define the tables in the SQL
Server metadata DB; this file renders them from THE SAME payload cells as the IIG
workbook. Pinned here, for pair 1 (one table) and pair 4 (three tables):

* one block per sheet, in ``dml.table_order``, one INSERT per payload row;
* every open cell (blank, not decided) is its ``<<COLUMN#n>>`` placeholder —
  exactly the blank cells the BSA's review copy lists as open, less the cells a
  DB rule decides (2026-10-08 port of the retired config_inserts_<env>.sql:
  audit-by -> @RFC_NUMBER, audit dates -> GETDATE(), the source connection ->
  @SRC_CONNECTION_ID, ``dml.db_null_columns`` -> NULL) — and every decided
  blank is NULL; every literal is the cell's value through ``dml.db_value_map``
  (ACTIVE_FLAG 'Y' -> 'S');
* catalogs as the cells carry them (mapped), the table index and the row labels
  naming each table by its mapped three-part name;
* a round trip: the script parsed back equals the clean IIG workbook of the same
  run, cell for cell;
* the CREATE reference text: one block per table and layer (pair 4: six), named
  and counted as the golden's TABLE_DEFINITIONS, the standard columns those the
  IIG's STGDELTA row targets.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import pytest
from openpyxl import load_workbook

from codegen import cli
from codegen.emit.framework import (
    TARGET_DDL,
    _combined_ddl_text,
    artefact_group,
    emit_framework,
    table_definitions,
)
from codegen.emit.metadata_inserts import (
    FILE_NAME,
    DbRules,
    cell_sql,
    placeholder,
    placeholders_as_null,
    render_metadata_inserts,
)
from codegen.faq import LoadPatternFaq
from codegen.iig_review import open_cells
from conftest import with_dml
from multi_table_golden import golden_rows
from test_m4_acceptance import GOLDEN_DDL, _scoped

PAIR1_COUNTS = {"DATA_FACTORY_PIPELINE_SCHEDULE": 4, "FILE_ADLS_INGESTION_DETAILS": 1,
                "ADLS_DELTA_INGESTION_DETAILS": 4, "STGDELTA_STDDELTA_INGESTION_DET": 1,
                "ADLS_FIXED_WIDTH_HANDLER": 3, "DATABRICKS_NOTEBOOK_DETAILS": 3,
                "DATA_QUALITY_RULES": 8, "EMAIL_TEMPLATE_CONFIG": 2}
PAIR4_COUNTS = {"DATA_FACTORY_PIPELINE_SCHEDULE": 4, "FILE_ADLS_INGESTION_DETAILS": 1,
                "ADLS_DELTA_INGESTION_DETAILS": 6, "STGDELTA_STDDELTA_INGESTION_DET": 3,
                "ADLS_FIXED_WIDTH_HANDLER": 0, "DATABRICKS_NOTEBOOK_DETAILS": 1,
                "DATA_QUALITY_RULES": 6, "EMAIL_TEMPLATE_CONFIG": 2}
PAIR1_STAGE = "pr_dlk_vnd_p.stg_vnd_p_accum.vnd_p_accum_client"
PAIR1_STANDARD = "pr_std_vnd_p.accum.vnd_p_accum_client"
AUDIT = 3                                  # SRC_FILE_NAME, REC_CREATION_TIME, REC_UPDATED_TIME
ADLS, STGDELTA = "ADLS_DELTA_INGESTION_DETAILS", "STGDELTA_STDDELTA_INGESTION_DET"


# -- a small reader for the script ----------------------------------------------- #


@dataclass(frozen=True)
class Placeholder:
    column: str
    row: int


@dataclass(frozen=True)
class Variable:
    name: str                # @RFC_NUMBER / @SRC_CONNECTION_ID


@dataclass(frozen=True)
class Expression:
    text: str                # GETDATE()


@dataclass
class Insert:
    table: str
    columns: list[str]
    values: list            # str (a literal) | None (NULL) | Placeholder | Variable | Expression
    label: str | None


_HEADER = re.compile(r"^-- ([A-Z0-9_]+): (\d+) row\(s\)(?: — .*)?$")
_INSERT = re.compile(r"^INSERT INTO \[dbo\]\.\[([A-Za-z0-9_]+)\] \((.*?)\) VALUES \((.*)\);$")
_INDEX = re.compile(r"^--   (stage|standard)\s+(\S+) — (\d+) column\(s\) \+ (\d+) audit$")


def _string(text: str, i: int) -> tuple[str, int]:
    """A literal at ``i``: N'…' pieces and NCHAR(10/13) joined by '+'."""
    pieces: list[str] = []
    while True:
        if text.startswith("N'", i):
            j, buf = i + 2, []
            while True:
                if text[j] == "'":
                    if text.startswith("''", j):
                        buf.append("'")
                        j += 2
                        continue
                    break
                buf.append(text[j])
                j += 1
            pieces.append("".join(buf))
            i = j + 1
        elif text.startswith("NCHAR(10)", i):
            pieces.append("\n")
            i += len("NCHAR(10)")
        elif text.startswith("NCHAR(13)", i):
            pieces.append("\r")
            i += len("NCHAR(13)")
        else:
            raise AssertionError(f"not a literal at {text[i:i + 30]!r}")
        k = i
        while k < len(text) and text[k] == " ":
            k += 1
        if not text.startswith("+", k):
            return "".join(pieces), i
        i = k + 1
        while text[i] == " ":
            i += 1


def parse_values(text: str) -> list:
    """The items of one VALUES tuple: str | None (NULL) | Placeholder."""
    items: list = []
    i = 0
    while i < len(text):
        if text.startswith("NULL", i):
            items.append(None)
            i += 4
        elif text.startswith("@", i):
            match = re.match(r"@[A-Za-z0-9_]+", text[i:])
            items.append(Variable(match.group(0)))
            i += match.end()
        elif text.startswith("GETDATE()", i):
            items.append(Expression("GETDATE()"))
            i += len("GETDATE()")
        elif text.startswith("<<", i):
            end = text.index(">>", i)
            column, number = text[i + 2:end].split("#")
            items.append(Placeholder(column, int(number)))
            i = end + 2
        else:
            value, i = _string(text, i)
            items.append(value)
        if i < len(text):
            assert text.startswith(", ", i), text[i:i + 30]
            i += 2
    return items


def _table_names() -> dict[str, str]:
    """Sheet -> the metadata-DB table its INSERTs target (dml.table_names, the
    confirmed names; any other sheet keeps its own name — 2026-10-09)."""
    from codegen.config import load_config

    config = load_config(Path(__file__).resolve().parents[1] / "config" / "config.yaml")
    return {sheet: entry.table for sheet, entry in config.dml.table_names.items()}


TABLE_NAMES = _table_names()


def parse_script(text: str) -> tuple[dict[str, int], dict[str, list[Insert]], list[tuple]]:
    """(block header counts in order, INSERTs per sheet, table index entries)."""
    headers: dict[str, int] = {}
    inserts: dict[str, list[Insert]] = {}
    index: list[tuple] = []
    sheet, label = None, None
    for line in text.splitlines():
        if (match := _INDEX.match(line)) is not None:
            index.append((match.group(1), match.group(2), int(match.group(3)),
                          int(match.group(4))))
        elif (match := _HEADER.match(line)) is not None:
            sheet = match.group(1)
            headers[sheet] = int(match.group(2))
            inserts[sheet] = []
        elif line.startswith("-- table ") and sheet is not None:
            label = line
        elif (match := _INSERT.match(line)) is not None:
            assert match.group(1) == TABLE_NAMES.get(sheet, sheet), (match.group(1), sheet)
            columns = [c.strip()[1:-1] for c in match.group(2).split(",")]
            inserts[sheet].append(Insert(sheet, columns, parse_values(match.group(3)), label))
            label = None
    return headers, inserts, index


def test_the_reader_reads_every_item_kind():
    assert parse_values("N'a', NULL, <<PIPELINE_ID#3>>, N'it''s' + NCHAR(10) + N'x', "
                        "@RFC_NUMBER, GETDATE()") == [
        "a", None, Placeholder("PIPELINE_ID", 3), "it's\nx", Variable("@RFC_NUMBER"),
        Expression("GETDATE()")]


# -- the cell rule ------------------------------------------------------------------ #


def test_cell_rule_literal_placeholder_null():
    assert cell_sql("SOURCE", "Northwind", {"badge": "from_frd"}, 2) == "N'Northwind'"
    assert cell_sql("SEQUENCE_NO", 1, {"badge": "synthetic"}, 2) == "N'1'"
    assert cell_sql("NOTE", "it's", {"badge": "from_frd"}, 1) == "N'it''s'"
    # a line break never sits inside a literal; the value is carried losslessly
    assert cell_sql("NOTE", "a\r\nb", {}, 1) == "N'a' + NCHAR(13) + NCHAR(10) + N'b'"
    assert parse_values(cell_sql("NOTE", "x'\ny", {}, 1)) == ["x'\ny"]
    # open: needs_template, always_blank, or any other blank no input decided
    assert cell_sql("PIPELINE_ID", "", {"badge": "needs_template"}, 3) == "<<PIPELINE_ID#3>>"
    assert cell_sql("SOURCE", None, {"badge": "from_frd"}, 4) == placeholder("SOURCE", 4)
    # decided blank (a family convention) -> NULL
    assert cell_sql("LOB", "", {"badge": "from_frd", "deliberate_blank": True}, 1) == "NULL"


def test_cell_rule_db_values(config):
    """The retired script's DB-side knowledge (config dml.*)."""
    rules = DbRules.from_config(config)
    assert cell_sql("ACTIVE_FLAG", "Y", {}, 1, rules) == "N'S'"          # workbook Y -> DB S
    assert cell_sql("ACTIVE_FLAG", "Y", {}, 1) == "N'Y'"                 # no rules: as written
    assert cell_sql("CREATED_DATE", "", {"badge": "needs_template"}, 2, rules) == "GETDATE()"
    assert cell_sql("CREATED_BY", "", {"badge": "needs_template"}, 2, rules) == "@RFC_NUMBER"
    assert cell_sql("CREATED_BY", "RFC1", {}, 2, rules) == "N'RFC1'"     # answered: the cell
    assert cell_sql("DAY_OF_SCHEDULE", "0", {}, 1, rules) == "NULL"      # §2 not populated
    assert cell_sql("SRC_CONNECTION_ID", "", {}, 1, rules,
                    "FILE_ADLS_INGESTION_DETAILS") == "@SRC_CONNECTION_ID"
    assert cell_sql("SRC_CONNECTION_ID", "", {}, 1, rules, "OTHER") == "<<SRC_CONNECTION_ID#1>>"


# -- payload-level rendering (pair 1 and pair 4) --------------------------------------- #


def _render(config, spec, payload):
    definitions = table_definitions(spec, config.conventions.get("acfc_prx"), [], config)
    text, counts = render_metadata_inserts(spec, payload, definitions, config, banner=[])
    return text, counts


@pytest.fixture(scope="module")
def rendered(pair1_config, pair1_spec, pair1_payload, pair4_config, pair4_spec, pair4_payload):
    return {
        "pair1": (pair1_config, pair1_payload, PAIR1_COUNTS,
                  *_render(pair1_config, pair1_spec, pair1_payload)),
        "pair4": (pair4_config, pair4_payload, PAIR4_COUNTS,
                  *_render(pair4_config, pair4_spec, pair4_payload)),
    }


@pytest.mark.parametrize("pair", ["pair1", "pair4"])
def test_one_block_per_sheet_in_table_order_one_insert_per_row(rendered, pair):
    config, payload, expected, text, counts = rendered[pair]
    headers, inserts, _index = parse_script(text)
    order = [t for t in config.dml.table_order if t in payload["tabs"]]
    assert list(headers) == order == list(counts)
    assert headers == counts == expected
    assert {s: len(tab["rows"]) for s, tab in payload["tabs"].items()} == expected
    for sheet, rows in inserts.items():
        assert len(rows) == expected[sheet], sheet
        for insert in rows:
            assert insert.columns == payload["tabs"][sheet]["headers"]
    if pair == "pair4":     # a sheet without rows: its header line, no INSERT
        assert "-- ADLS_FIXED_WIDTH_HANDLER: 0 row(s)" in text
        assert inserts["ADLS_FIXED_WIDTH_HANDLER"] == []


AUDIT_BY = {"CREATED_BY", "UPDATED_BY", "CRETAED_BY"}


@pytest.mark.parametrize("pair", ["pair1", "pair4"])
def test_placeholders_exactly_for_the_open_cells_null_for_decided_blanks(rendered, pair):
    config, payload, _expected, text, _counts = rendered[pair]
    dml = config.dml
    null_columns, expressions = set(dml.db_null_columns), dml.db_blank_expressions
    _headers, inserts, _index = parse_script(text)
    placeholders, nulls, decided_by_db = set(), set(), set()
    for sheet, rows in inserts.items():
        for number, (insert, row) in enumerate(zip(rows, payload["tabs"][sheet]["rows"],
                                                   strict=True), start=1):
            for column, item in zip(insert.columns, insert.values, strict=True):
                value, entry = row["values"][column], row["badges"][column]
                blank = value in ("", None)
                cell = (sheet, number, column)
                if column in null_columns:                       # NULL whatever it holds
                    assert item is None, cell
                    nulls.add(cell)
                    decided_by_db.add(cell)
                elif isinstance(item, Placeholder):
                    assert item == Placeholder(column, number)
                    assert blank and not entry.get("deliberate_blank"), cell
                    placeholders.add(cell)
                elif item is None:
                    assert blank and entry.get("deliberate_blank"), cell
                    nulls.add(cell)
                elif isinstance(item, Expression):
                    assert blank and expressions[column] == item.text, cell
                    decided_by_db.add(cell)
                elif isinstance(item, Variable):
                    assert blank, cell
                    assert (item.name == "@RFC_NUMBER" and column in AUDIT_BY) or (
                        item.name == "@SRC_CONNECTION_ID" and column == "SRC_CONNECTION_ID"
                        and sheet == "FILE_ADLS_INGESTION_DETAILS"), cell
                    decided_by_db.add(cell)
                else:
                    mapped = dml.db_value_map.get(column, {}).get(str(value).strip())
                    assert item == (mapped if mapped is not None else str(value)), cell
    # the open cells the BSA's review copy lists, restricted to blank ones, less the
    # cells a DB rule writes (variables, GETDATE(), NULL columns)
    review_open = {
        (c.sheet, c.row - 1, c.column)
        for c in open_cells(payload, payload["always_blank"], config, "iig_v2")
        if payload["tabs"][c.sheet]["rows"][c.row - 2]["values"][c.column] in ("", None)}
    assert placeholders == review_open - decided_by_db
    always_blank = set(payload["always_blank"])
    assert {(s, n, c) for s, n, c in placeholders if c in always_blank}   # ids, containers …
    family = {(s, n, c) for s, n, c in nulls if c not in null_columns}
    if pair == "pair1":     # the pair-1 family's LOB is blank by convention: NULL
        assert family == {(ADLS, n, "LOB") for n in range(1, 5)} | {(STGDELTA, 1, "LOB")}
    else:
        assert family == set()
    assert {c for _s, _n, c in nulls if c in null_columns} <= null_columns


def test_pair4_catalogs_are_the_mapped_ones(rendered):
    _config, _payload, _expected, text, _counts = rendered["pair4"]
    assert not re.search(r"pr_dlk|pr_std|syn_fallback", text, re.IGNORECASE)
    _headers, inserts, _index = parse_script(text)
    for insert in inserts[STGDELTA]:
        values = dict(zip(insert.columns, insert.values, strict=True))
        assert (values["SRC_CATALOG_NAME"], values["TGT_CATALOG_NAME"]) == ("d1_dlk", "d1_std")


def test_pair4_table_index_and_labels_name_every_table_by_its_mapped_name(rendered,
                                                                           pair4_golden):
    _config, _payload, _expected, text, _counts = rendered["pair4"]
    _headers, inserts, index = parse_script(text)
    golden = pair4_golden["TABLE_DEFINITIONS"]
    assert sorted(index) == sorted((d["layer"], d["mapped"], d["columns"], AUDIT)
                                   for d in golden)
    by_segment = {(d["layer"], d["segment"]): d["mapped"] for d in golden}
    # per table: stage then standard, in the STTM's table order
    assert [name for _l, name, _c, _a in index] == [
        by_segment[(layer, segment)] for segment in ("Header", "Detail", "Trailer")
        for layer in ("stage", "standard")]
    detail = by_segment[("stage", "Detail")]
    files = [row["SRC_FILE_NAME"] for row in golden_rows(pair4_golden, ADLS)]
    assert [i.label for i in inserts[ADLS]] == [f"-- table {detail} <- file {f}" for f in files]
    assert [i.label for i in inserts[STGDELTA]] == [
        f"-- table {by_segment[('stage', s)]} -> {by_segment[('standard', s)]}"
        for s in ("Header", "Detail", "Trailer")]


def test_pair1_index_and_labels_name_its_one_table_as_the_golden_ddl_does(rendered):
    _config, _payload, _expected, text, _counts = rendered["pair1"]
    _headers, inserts, index = parse_script(text)
    golden_names = re.findall(r"CREATE OR REPLACE TABLE (\S+)",
                              GOLDEN_DDL.read_text(encoding="utf-8"))
    assert golden_names == [PAIR1_STAGE, PAIR1_STANDARD]
    assert index == [("stage", PAIR1_STAGE, 21, AUDIT), ("standard", PAIR1_STANDARD, 21, AUDIT)]
    assert {i.label for i in inserts[ADLS]} == {
        f"-- table {PAIR1_STAGE} <- file {i.values[i.columns.index('SRC_FILE_NAME')]}"
        for i in inserts[ADLS]}
    assert len({i.label for i in inserts[ADLS]}) == 4                  # one per file
    assert [i.label for i in inserts[STGDELTA]] == [
        f"-- table {PAIR1_STAGE} -> {PAIR1_STANDARD}"]


@pytest.mark.parametrize("pair", ["pair1", "pair4"])
def test_with_placeholders_read_as_null_every_statement_parses_as_tsql(rendered, pair):
    pytest.importorskip("sqlglot")
    from codegen.emit.dml import validate_tsql

    _config, _payload, _expected, text, _counts = rendered[pair]
    assert validate_tsql(placeholders_as_null(text)) == []
    # … and not before: an open cell keeps the script from parsing, let alone running
    assert validate_tsql(text)


# -- sheet -> table names (2026-10-09) ------------------------------------------------- #

WALKTHROUGH_TABLES = {"DATA_FACTORY_PIPELINE_SCHEDULE": "§2", "FILE_ADLS_INGESTION_DETAILS": "§5",
                      "ADLS_DELTA_INGESTION_DETAILS": "§7"}


def test_the_confirmed_table_names_are_config_the_walkthrough_and_the_owner_brief(config):
    """Only the walkthrough's IIG tables (METADATA_DB_SEMANTICS §2 / §5 / §7 — its
    other three tables are not IIG sheets) and the owner-stated STGDELTA table are
    confirmed; an Excel tab name (31 characters at most) is not a table name."""
    names = config.dml.table_names
    assert set(names) == {*WALKTHROUGH_TABLES, STGDELTA}
    for sheet, section in WALKTHROUGH_TABLES.items():
        assert names[sheet].table == sheet
        assert f"METADATA_DB_SEMANTICS.md {section}" in names[sheet].citation
    assert names[STGDELTA].table == "stg_delta_stddelta_ingestion_details"
    assert names[STGDELTA].citation == "owner brief 2026-10-09"
    tabs = config.metadata.templates["iig_v2"].tabs
    assert set(names) <= set(tabs)
    assert max(len(t) for t in tabs) == len(STGDELTA) == 31          # the cut tab name


@pytest.mark.parametrize("pair", ["pair1", "pair4"])
def test_insert_targets_are_the_confirmed_names_else_the_sheet_marked_unconfirmed(rendered,
                                                                                  pair):
    config, payload, _expected, text, _counts = rendered[pair]
    lines = text.splitlines()
    for sheet, tab in payload["tabs"].items():
        if not tab["rows"]:
            continue
        start = next(i for i, ln in enumerate(lines) if ln.startswith(f"-- {sheet}: "))
        block = lines[start + 1:lines.index("", start)]          # a block ends at a blank line
        inserts = [ln for ln in block if ln.startswith("INSERT INTO")]
        assert len(inserts) == len(tab["rows"]), sheet
        confirmed = config.dml.table_names.get(sheet)
        table = confirmed.table if confirmed else sheet
        assert inserts and all(ln.startswith(f"INSERT INTO [dbo].[{table}] (") for ln in inserts)
        if confirmed is None:
            assert block[0] == (
                f"-- unconfirmed: target table [dbo].[{sheet}] = the IIG sheet name; a tab name "
                "(Excel cuts it at 31 characters) is not the framework's table name until "
                "confirmed (Friday checklist 12; dml.table_names)"), sheet
        elif table != sheet:
            assert block[0].startswith(f"-- target table [dbo].[{table}], not the IIG sheet ")
            assert block[0].endswith(f"confirmed: {confirmed.citation} (dml.table_names)")
        else:
            assert not block[0].startswith(("-- unconfirmed", "-- target table")), sheet
    # STGDELTA: its INSERTs and its guards name the confirmed table, never the cut tab name
    assert "[dbo].[STGDELTA_STDDELTA_INGESTION_DET]" not in text
    assert ("IF EXISTS (SELECT 1 FROM [dbo].[stg_delta_stddelta_ingestion_details] WHERE "
            "[GROUP_ID] = <<GROUP_ID#1>>) RAISERROR(N'STGDELTA_STDDELTA_INGESTION_DET row 1: "
            "GROUP_ID is already used") in text
    unconfirmed = [s for s, tab in payload["tabs"].items()
                   if tab["rows"] and s not in config.dml.table_names]
    assert set(unconfirmed) <= {"ADLS_FIXED_WIDTH_HANDLER", "DATABRICKS_NOTEBOOK_DETAILS",
                                "DATA_QUALITY_RULES", "EMAIL_TEMPLATE_CONFIG"}
    assert len(unconfirmed) == (4 if pair == "pair1" else 3)   # pair 4: no fixed-width row
    assert text.count("-- unconfirmed: target table ") == len(unconfirmed)


# -- the CREATE reference text, one block per table ------------------------------------ #


def _create_blocks(ddl: str) -> list[tuple[str, list[str]]]:
    return [(name, [line.rstrip(",") for line in body.splitlines()])
            for name, body in re.findall(r"CREATE OR REPLACE TABLE (\S+)\n\(\n(.*?)\n\)",
                                         ddl, re.DOTALL)]


def test_pair4_create_reference_one_block_per_table(pair4_config, pair4_spec, pair4_golden):
    flags: list[str] = []
    ddl = _combined_ddl_text(pair4_spec, pair4_config.conventions.get("acfc_prx"), flags,
                             pair4_config)
    blocks = _create_blocks(ddl)
    golden = pair4_golden["TABLE_DEFINITIONS"]
    assert [name for name, _cols in blocks] == [d["mapped"] for d in golden]     # 6, in order
    assert [len(cols) for _name, cols in blocks] == [d["columns"] + AUDIT for d in golden]
    assert ddl.count("--stage table") == ddl.count("--standard table") == 3
    # stage columns as the stage band types them; standard columns = the ones the
    # IIG's STGDELTA row targets (the Standard band's names), stage ones = the ADLS row's
    stgdelta = {row["TGT_TABLE_NAME"]: row for row in golden_rows(pair4_golden, STGDELTA)}
    adls = golden_rows(pair4_golden, ADLS)[0]
    for name, cols in blocks:
        names = [c.split(" ")[0] for c in cols]
        if name.startswith("d1_std."):
            assert names == stgdelta[name.rsplit(".", 1)[1]]["TGT_COLUMN_NAMES"].split(",")
        else:
            assert all(c.endswith(" String") for c in cols[:-AUDIT])
            if name.endswith("_dtl"):
                assert names == adls["TGT_COLUMN_NAMES"].split(",")
    standard_dtl = dict(blocks)["d1_std.nb_cob.nb_cob_report_dtl"]
    assert "MEMBER_BIRTH_DATE DATE" in standard_dtl            # the Standard band's name + type
    assert not [f for f in flags if f.startswith(("catalog_from_config", "catalog_unstated"))]


def test_pair1_create_reference_is_still_the_golden(pair1_config, pair1_spec):
    ddl = _combined_ddl_text(pair1_spec, pair1_config.conventions.get("acfc_prx"), [],
                             pair1_config)
    assert ddl.encode("utf-8") == GOLDEN_DDL.read_bytes()


# -- end to end: the artefact, its switch, and the round trip --------------------------- #


def _run(config, spec, tmp: Path):
    gate = cli._generate_feed(spec, _scoped(with_dml(config), tmp), dry_run=True,
                              skip_tests=True, output_mode="framework",
                              conventions_profile="acfc_prx", iig_template="iig_v2")
    return gate, tmp / "out" / spec.feed_slug / "framework"


@pytest.fixture(scope="module")
def runs(pair1_config, pair1_spec, pair4_config, pair4_spec, tmp_path_factory):
    return {"pair1": _run(pair1_config, pair1_spec, tmp_path_factory.mktemp("mi_pair1")),
            "pair4": _run(pair4_config, pair4_spec, tmp_path_factory.mktemp("mi_pair4"))}


@pytest.fixture(scope="module")
def runs_config(pair1_config, pair4_config):
    return {"pair1": pair1_config, "pair4": pair4_config}


@pytest.mark.parametrize("pair", ["pair1", "pair4"])
def test_round_trip_the_script_equals_the_clean_iig_cell_for_cell(runs, runs_config, pair):
    _gate, framework_dir = runs[pair]
    _headers, inserts, _index = parse_script(
        (framework_dir / FILE_NAME).read_text(encoding="utf-8"))
    (clean,) = framework_dir.glob("*_IIG.xlsx")
    workbook = load_workbook(clean)
    assert set(inserts) == set(workbook.sheetnames)
    diffs = []
    for sheet, rows in inserts.items():
        table = list(workbook[sheet].iter_rows(values_only=True))
        header, body = list(table[0]), table[1:]
        assert len(rows) == len(body), sheet
        dml = runs_config[pair].dml
        for number, (insert, cells) in enumerate(zip(rows, body, strict=True), start=1):
            assert insert.columns == header, sheet
            for column, item, cell in zip(header, insert.values, cells, strict=True):
                blank = cell is None or cell == ""
                if column in dml.db_null_columns:
                    ok = item is None                     # NULL whatever the workbook holds
                elif isinstance(item, (Placeholder, Variable, Expression)) or item is None:
                    ok = blank
                else:
                    mapped = dml.db_value_map.get(column, {}).get(str(cell).strip())
                    ok = not blank and item == (mapped if mapped is not None else str(cell))
                if not ok:
                    diffs.append((sheet, number, column, item, cell))
    assert diffs == []


def test_the_artefact_is_listed_grouped_and_checked(runs):
    gate, framework_dir = runs["pair4"]
    assert (framework_dir / FILE_NAME).is_file()
    assert gate.verdict == "PASS_WITH_FLAGS", [c for c in gate.checks if not c.passed]
    (check,) = [c for c in gate.checks if c.name == "metadata_inserts"]
    assert check.passed and "one INSERT per IIG row (23)" in check.details
    assert artefact_group(FILE_NAME) == TARGET_DDL
    addition = (framework_dir / "ADDITION.md").read_text(encoding="utf-8")
    assert f"| `{FILE_NAME}` | The client's \"DDL\"" in addition
    assert f"`NB_COB_REPORT_DDL.txt`, `{FILE_NAME}`" in addition         # By target system
    # config_inserts_<env>.sql is retired (2026-10-08); the notebooks run this file
    assert "config_inserts_<env>.sql` is retired (2026-10-08): this file replaces it" in addition
    assert not list(framework_dir.glob("config_inserts_*.sql"))
    assert (framework_dir / "Insert_scripts_config_table_q1.py").is_file()


def test_pair4_with_shipped_settings_writes_metadata_inserts(pair4_config, pair4_spec, tmp_path):
    """2026-10-08: the client's "DDL" is written on every framework run —
    emit_dml / dml.enabled (off as shipped) gate only the runner notebooks."""
    framework = emit_framework(pair4_spec, LoadPatternFaq(), [], pair4_config, tmp_path,
                               conventions_profile="acfc_prx", iig_template="iig_v2")
    assert framework.dml_disabled_reason is not None                   # shipped: switches off
    names = [p.name for p in framework.files]
    assert FILE_NAME in names
    assert (tmp_path / pair4_spec.feed_slug / "framework" / FILE_NAME).is_file()
    assert not any(n.startswith("Insert_scripts_config_table_") for n in names)
    (check,) = [c for c in framework.checks if c.name == "metadata_inserts"]
    assert check.passed and "one INSERT per IIG row (23)" in check.details
