"""codegen.metadata_sheet: row rules per tab, coverage arithmetic, xlsx.

Display-only ACFC metadata-sheet preview. Unit tests build their own FRD
contract in tmp (no fixture files needed); the tests that exercise the
STTM-derived cells resolve the demo fixture pair and skip when it is
absent, like the rest of the suite.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from codegen.contracts.frd import FrdContract
from codegen.metadata_sheet import (
    build_workbook,
    metadata_sheet_payload,
    workbook_bytes,
    workbook_filename,
)
from codegen.resolve.resolver import resolve_pair
from codegen.rules.compiler import compile_rules

REPO = Path(__file__).resolve().parents[1]


def _frd_feed(**overrides) -> dict:
    values = {
        "feed_name": "cv_test_feed",
        "source_system": "Civic Vantage (CV)",
        "file_name_patterns": ["extract_YYYY_MM.csv", "extract_CCYY_MM.csv"],
        "file_format": "csv",
        "delimiter": None,
        "record_segments": [],
        "frequency": "Monthly",
        "load_windows_sla": [],
        "lobs": ["All"],
        "domain": "Care Management",
        "sub_domain": "Public",
        "landing_location": None,
        "stage_target": {
            "catalog": None, "schema": "stg_x", "tables": ["cv_test_feed"],
            "load_strategy": "Truncate and Load",
        },
        "standard_target": {
            "catalog": None, "schema": "x", "tables": ["cv_test_feed"],
            "load_strategy": "Append",
        },
        "validation_rules": [],
        "recycle_rule": None,
        "history_backfill": None,
        "archive_retention": None,
        "phi_pii_notes": None,
        "sttm_reference": None,
        "requirement_ids": [],
    }
    values.update(overrides)
    return values


def _base_dir_with_contract(tmp_path: Path, config, feeds: list[dict]) -> Path:
    contract = {
        "contract_name": "test",
        "generated_from_frd": "test.docx",
        "generated_date": "2026-01-01T00:00:00",
        "generator": "test",
        "status": "PASS",
        "project": {"project_id": None, "project_name": "t", "business_context_summary": None},
        "in_scope": [],
        "out_of_scope": [],
        "assumptions_constraints_dependencies": [],
        "feeds": feeds,
        "system_interfaces": [],
        "open_items": [],
    }
    FrdContract.model_validate(contract)  # fail here, not inside the payload
    target = tmp_path / config.contracts.dir / config.demo.frd
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(contract), encoding="utf-8")
    return tmp_path


def _payload(config, tmp_path, feeds=None, **kwargs) -> dict:
    base = _base_dir_with_contract(tmp_path, config, feeds or [_frd_feed()])
    return metadata_sheet_payload(config, base, **kwargs)


def _cells(row: dict) -> dict:
    return {h: (row["values"][h], row["badges"][h]["badge"]) for h in row["values"]}


# -- ADLS_DELTA_INGESTION_DETAILS (file → stage) ------------------------------- #


def test_adls_delta_frd_sourced_and_synthesized(config, tmp_path):
    payload = _payload(config, tmp_path)
    (row,) = payload["tabs"]["ADLS_DELTA_INGESTION_DETAILS"]["rows"]
    cells = _cells(row)
    assert cells["OBJECT_NAME"] == ("cv_test_feed", "from_frd")
    assert cells["SOURCE"] == ("Civic Vantage (CV)", "from_frd")
    assert cells["DOMAIN"] == ("Care Management", "from_frd")
    assert cells["SUBDOMAIN"] == ("Public", "from_frd")
    # Null landing_location → Part A's synthesis, flagged; split into the
    # container / path columns of the real layout.
    assert cells["SRC_CONTAINER_NAME"] == ("mftlanding", "synthetic")
    assert cells["SRC_ADLS_PATH"] == (
        "/inbound/care_management/public/civic_vantage_cv", "synthetic"
    )
    assert cells["SRC_FILE_NAME"] == (
        "extract_YYYY_MM.csv; extract_CCYY_MM.csv", "from_frd")
    assert cells["SRC_FORMAT"] == ("csv", "from_frd")
    assert cells["FREQUENCY"] == ("Monthly", "from_frd")
    # No header indicator in the FRD contract; no FAQ answer here.
    assert cells["HEADER_FLAG"] == ("", "needs_template")
    # Null delimiter, no run → cannot be known from documents yet.
    assert cells["SRC_FILE_DELIMITER"] == ("", "needs_template")
    # STTM-derived columns are unknowable before a run.
    assert cells["SRC_COLUMNS"] == ("", "needs_template")


def test_adls_delta_contract_landing_and_delimiter(config, tmp_path):
    payload = _payload(
        config, tmp_path,
        feeds=[_frd_feed(
            landing_location="mftlanding\\inbound\\a\\b\\c",
            record_segments=["Header", "Detail", "Trailer"],
            delimiter="|",
        )],
    )
    cells = _cells(payload["tabs"]["ADLS_DELTA_INGESTION_DETAILS"]["rows"][0])
    assert cells["SRC_CONTAINER_NAME"] == ("mftlanding", "from_frd")
    assert cells["SRC_ADLS_PATH"] == ("/inbound/a/b/c", "from_frd")
    assert cells["SRC_FILE_DELIMITER"] == ("|", "from_frd")


def test_always_blank_headers_are_empty_and_flagged(config, tmp_path):
    payload = _payload(config, tmp_path)
    for tab in payload["tabs"].values():
        for row in tab["rows"]:
            for header in config.demo.metadata_sheet.always_blank:
                if header in row["values"]:
                    assert row["values"][header] == ""
                    assert row["badges"][header]["badge"] == "needs_template"
                    assert "ACFC framework" in row["badges"][header]["tooltip"]


# -- DATA_FACTORY_PIPELINE_SCHEDULE / stage-only feeds ------------------------- #


def test_pipeline_schedule_one_row_per_layer(config, tmp_path):
    payload = _payload(config, tmp_path)
    rows = payload["tabs"]["DATA_FACTORY_PIPELINE_SCHEDULE"]["rows"]
    assert [r["values"]["LAYER_NAME"] for r in rows] == ["STAGE", "STANDARD"]
    stage, standard = rows
    assert _cells(stage)["APPLICATION_NAME"] == ("Civic Vantage (CV)", "from_frd")
    assert _cells(stage)["PIPELINE_FREQUENCY"] == ("Monthly", "from_frd")
    # Framework conventions from the reference workbook stay flagged synthetic.
    assert _cells(standard)["ACTIVE_FLAG"] == ("Y", "synthetic")
    assert _cells(standard)["ACTIVE_END_DATE"] == ("9999-12-31", "synthetic")
    # Framework-assigned IDs stay blank.
    assert _cells(stage)["PIPELINE_ID"] == ("", "needs_template")


def test_pipeline_schedule_stage_only_feed(config, tmp_path):
    payload = _payload(
        config, tmp_path,
        feeds=[_frd_feed(standard_target={
            "catalog": None, "schema": None, "tables": [], "load_strategy": "Append",
        })],
    )
    rows = payload["tabs"]["DATA_FACTORY_PIPELINE_SCHEDULE"]["rows"]
    assert [r["values"]["LAYER_NAME"] for r in rows] == ["STAGE"]
    # No standard target → no stage→standard ingestion row either.
    assert payload["tabs"]["STGDELTA_STDDELTA_INGESTION_DET"]["rows"] == []


# -- STTM-derived tabs before a run -------------------------------------------- #


def test_sttm_derived_tabs_empty_before_a_run(config, tmp_path):
    payload = _payload(config, tmp_path)
    # DQ rules come from the mapping contract's column lists — no run, no rows.
    assert payload["tabs"]["DATA_QUALITY_RULES"]["rows"] == []
    assert payload["tabs"]["STGDELTA_STDDELTA_INGESTION_DET"]["rows"] == []


# -- coverage arithmetic ------------------------------------------------------- #


def test_coverage_counts_add_up(config, tmp_path):
    payload = _payload(config, tmp_path)
    coverage = payload["coverage"]
    cells = sum(
        len(row["badges"]) for tab in payload["tabs"].values() for row in tab["rows"]
    )
    assert coverage["total"] == cells
    assert (
        coverage["derived"] + coverage["synthetic"] + coverage["needs_template"]
        == coverage["total"]
    )
    assert coverage["synthetic"] >= 2  # landing_path + the two load strategies


# -- with the demo fixture pair (STTM-derived cells) --------------------------- #

_DEMO_FRD = REPO / "fixtures" / "contracts" / "FRD_demo_cv_golden.contract.json"
_DEMO_STTM = REPO / "fixtures" / "contracts" / "sttm_mapping_contracts_cv_golden.json"
needs_demo_pair = pytest.mark.skipif(
    not (_DEMO_FRD.is_file() and _DEMO_STTM.is_file()),
    reason="demo fixture pair not restored (removed 2026-08-22)",
)


@pytest.fixture(scope="module")
def demo_specs(config):
    if not (_DEMO_FRD.is_file() and _DEMO_STTM.is_file()):
        pytest.skip("demo fixture pair not restored (removed 2026-08-22)")
    return resolve_pair(_DEMO_FRD, _DEMO_STTM, config)


@needs_demo_pair
def test_sttm_cells_populate_from_mapping_contract(config, demo_specs):
    unmapped = {
        spec.feed_slug: {
            o.rule_text for o in compile_rules(spec) if o.classification == "unmapped"
        }
        for spec in demo_specs
    }
    payload = metadata_sheet_payload(
        config, REPO, specs=demo_specs, unmapped_by_slug=unmapped, run_label="test_run"
    )
    adls = payload["tabs"]["ADLS_DELTA_INGESTION_DETAILS"]["rows"]
    assert len(adls) == len(demo_specs)
    for row in adls:
        # Column lists resolve from the STTM after a run.
        assert row["badges"]["SRC_COLUMNS"]["badge"] == "from_sttm"
        assert ":" in row["values"]["SRC_COLUMNS"]
        assert row["badges"]["MANDATORY_FIELD_LIST"]["badge"] == "from_sttm"
        assert row["badges"]["TGT_PRIMARY_KEY"]["badge"] == "from_sttm"
        # CV FRD states no delimiter; with a run it resolves from the STTM.
        assert row["badges"]["SRC_FILE_DELIMITER"]["badge"] == "from_sttm"
        assert row["values"]["SRC_FILE_DELIMITER"]
        # Framework-assigned IDs stay blank even after a run.
        assert row["values"]["GROUP_ID"] == ""
        assert row["badges"]["GROUP_ID"]["badge"] == "needs_template"
    # DQ rules carry the STTM column lists; recycle feeds additionally get a
    # ReferentialCheckRule row sourced from the FRD's DQ requirement.
    dq = payload["tabs"]["DATA_QUALITY_RULES"]["rows"]
    assert dq
    for row in dq:
        if row["values"].get("RULE_CLASS") == "ReferentialCheckRule":
            assert row["badges"]["SOURCE_COLUMN"]["badge"] == "from_frd"
            assert row["values"]["INPUT_PARAM"]  # the reference table
            continue
        assert row["badges"]["SOURCE_COLUMN"]["badge"] == "from_sttm"
        assert row["values"]["SOURCE_COLUMN"] == row["values"]["TARGET_COLUMN"]
    # Standard-layer ingestion rows exist for feeds with a standard target.
    std = payload["tabs"]["STGDELTA_STDDELTA_INGESTION_DET"]["rows"]
    assert std
    for row in std:
        assert row["badges"]["TGT_PRIMARY_KEY"]["badge"] == "from_sttm"
        assert row["values"]["TGT_PRIMARY_KEY"]


def test_notebook_details_refresh_type_no_run_is_synthetic(config, tmp_path):
    # Without a resolved spec the legs carry the config stand-ins, mapped to
    # the framework vocabulary and badged synthetic.
    payload = _payload(config, tmp_path)
    rows = payload["tabs"]["DATABRICKS_NOTEBOOK_DETAILS"]["rows"]
    by_leg = {r["values"]["PROCESS_NAME"].split(" ")[0]: r for r in rows}
    assert by_leg["Stage"]["values"]["TGT_REFRESH_TYPE"] == "Overwrite"
    assert by_leg["Stage"]["badges"]["TGT_REFRESH_TYPE"]["badge"] == "synthetic"


def test_notebook_details_refresh_type_from_frd_load_strategy(config, demo_specs):
    # With a run, TGT_REFRESH_TYPE derives from the FRD's Load Strategy —
    # 'Truncate and Load' (STG) is the framework's 'Overwrite' refresh,
    # 'Append' (STD) passes through — badged from_frd with the citation.
    payload = metadata_sheet_payload(
        config, REPO, specs=demo_specs, unmapped_by_slug={}, run_label="test_run"
    )
    rows = payload["tabs"]["DATABRICKS_NOTEBOOK_DETAILS"]["rows"]
    stage = [r for r in rows if r["values"]["PROCESS_NAME"].startswith("Stage")]
    standard = [r for r in rows
                if r["values"]["PROCESS_NAME"].startswith("Standard")]
    assert stage and standard
    for row in stage:
        cell = row["badges"]["TGT_REFRESH_TYPE"]
        assert row["values"]["TGT_REFRESH_TYPE"] == "Overwrite"
        assert cell["badge"] == "from_frd"
        assert "Truncate and Load" in cell["tooltip"]
        assert "Load Strategy (STG)" in cell["tooltip"]
    for row in standard:
        cell = row["badges"]["TGT_REFRESH_TYPE"]
        assert row["values"]["TGT_REFRESH_TYPE"] == "Append"
        assert cell["badge"] == "from_frd"
        assert "Load Strategy (STD)" in cell["tooltip"]


# -- xlsx ---------------------------------------------------------------------- #


def test_workbook_structure_and_provenance(config, tmp_path):
    payload = _payload(config, tmp_path)
    out = tmp_path / "preview.xlsx"
    build_workbook(payload).save(out)
    workbook = load_workbook(out)
    assert workbook.sheetnames == [*config.demo.metadata_sheet.tabs, "_provenance"]
    for name, tab in config.demo.metadata_sheet.tabs.items():
        headers = [c.value for c in workbook[name][1]]
        assert headers == tab.headers
    provenance = workbook["_provenance"]
    assert provenance.max_row - 1 == payload["coverage"]["total"]
    assert [c.value for c in provenance[1]] == ["tab", "row", "header", "badge", "tooltip"]


def test_workbook_bytes_and_filename(config, tmp_path):
    payload = _payload(config, tmp_path)
    workbook = load_workbook(io.BytesIO(workbook_bytes(payload)))
    assert "_provenance" in workbook.sheetnames
    assert workbook_filename(payload) == "metadata_sheet_preview_pre-run.xlsx"
    payload["run_label"] = "demo_x"
    assert workbook_filename(payload) == "metadata_sheet_preview_demo_x.xlsx"


# -- endpoints ----------------------------------------------------------------- #

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient  # noqa: E402
from ui.backend import main as ui_main  # noqa: E402

needs_demo_frd = pytest.mark.skipif(
    not _DEMO_FRD.is_file(), reason="demo FRD fixture not restored"
)


@pytest.fixture()
def client():
    return TestClient(ui_main.app)


@needs_demo_frd
def test_metadata_sheet_endpoint_no_run_state(client):
    payload = client.get("/api/demo/metadata-sheet").json()
    assert {"layout_note", "tabs", "coverage"} <= set(payload)
    assert "client IIG template" in payload["layout_note"]
    # TestClient skips the startup generate, so no run is loaded and the
    # STTM-derived tabs are empty.
    assert payload["tabs"]["DATA_QUALITY_RULES"]["rows"] == []


@needs_demo_frd
def test_metadata_sheet_xlsx_endpoint(client):
    response = client.get("/api/demo/metadata-sheet.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml"
    )
    assert "metadata_sheet_preview_" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    assert "_provenance" in workbook.sheetnames
