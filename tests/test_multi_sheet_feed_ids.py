"""Family E, one mapping sheet PER SOURCE FILE (2026-10-13, D1 pair_11).

extract-sttm failed with "duplicate feed_id(s) in STTM contract": every sheet
of the workbook mapped a different source file into the SAME stage table under
the same project name, the FRD named that one table, so the layout stage kept
ONE FRD feed and all five sheets matched it. Now a sheet whose feed name would
be shared is its own feed — named from its own file-name cell, else the sheet
name; a workbook whose names are already unique keeps them (pairs 5 / 7 / 9 /
10, the SD / MIDS shape). Chain as the harness runs it: layout
--frd-contract-out -> extract-sttm -> the resolver's FRD ⋈ STTM join.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from acfc_shapes import family_e_multi as fm
from codegen import cli
from codegen.config import load_config
from codegen.contracts.frd import FrdContract
from codegen.contracts.sttm import SttmContract
from codegen.resolve.resolver import resolve_feeds

REPO = Path(__file__).resolve().parents[1]
CONFIG = REPO / "config" / "config.yaml"
DATE = "2026-10-13"


@pytest.fixture()
def env(monkeypatch, tmp_path):
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_PROVIDER", "1")
    monkeypatch.setenv("CODEGEN_NOTIFICATION_EMAILS", "syn.dl.prodsupport@synthetic.example")
    state = tmp_path / "state"
    state.mkdir()
    monkeypatch.setenv("CODEGEN_STORAGE_STATE", f"local:{state.as_posix()}")
    return tmp_path


def _chain(tmp: Path, capsys, *, own_file_rows: bool = True, tag: str = "a"):
    sttm, frd = tmp / f"STTM_nb_care_mgmt_{tag}.xlsx", tmp / f"FRD_nb_care_mgmt_{tag}.docx"
    fm.build_sttm(own_file_rows=own_file_rows).save(sttm)
    frd.write_bytes(fm.build_frd())
    frd_out, sttm_out = tmp / f"frd_{tag}.contract.json", tmp / f"sttm_{tag}.contract.json"
    capsys.readouterr()
    assert cli.main(["layout", "--workbook", str(sttm), "--frd", str(frd), "--dry-run",
                     "--no-cache", "--frd-contract-out", str(frd_out)]) == 0
    layout_out = capsys.readouterr().out
    code = cli.main(["extract-sttm", "--workbook", str(sttm), "--frd-contract", str(frd_out),
                     "--out", str(sttm_out), "--generated-date", DATE])
    out = capsys.readouterr().out
    assert code == 0, layout_out + out
    return frd_out, sttm_out


def test_five_sheets_five_distinct_feed_ids(env, capsys):
    frd_out, sttm_out = _chain(env, capsys)
    sttm = SttmContract.model_validate_json(sttm_out.read_text(encoding="utf-8"))
    ids = [f.feed_id for f in sttm.feeds]
    assert ids == ["nb_cm_north", "nb_cm_south", "nb_cm_east", "nb_cm_west", "nb_cm_additional"]
    # each feed is ITS sheet, with its own file
    by_sheet = {f.mapping_sheet: f for f in sttm.feeds}
    assert set(by_sheet) == set(fm.FILES)
    frd = FrdContract.model_validate_json(frd_out.read_text(encoding="utf-8"))
    patterns = {f.feed_name: f.file_name_patterns for f in frd.feeds}
    for file_name in fm.FILES.values():
        stem = file_name.removesuffix("_YYYYMMDD.txt")
        assert patterns[stem] == [file_name], patterns
    # every feed still lands in the one table the FRD names
    assert {t for f in frd.feeds for t in f.stage_target.tables} == {fm.TABLE}


def test_feed_ids_are_stable_across_runs(env, capsys):
    _f1, first = _chain(env, capsys, tag="a")
    _f2, second = _chain(env, capsys, tag="b")
    strip = lambda p: [{k: v for k, v in f.items() if k not in ("mapping_sheet",)}  # noqa: E731
                       for f in json.loads(p.read_text(encoding="utf-8"))["feeds"]]
    assert [f["feed_id"] for f in strip(first)] == [f["feed_id"] for f in strip(second)]
    assert [f["fields"] for f in strip(first)] == [f["fields"] for f in strip(second)]


def test_the_frd_listing_the_five_files_joins_every_feed(env, capsys):
    frd_out, sttm_out = _chain(env, capsys)
    config = load_config(CONFIG)
    frd = FrdContract.model_validate_json(frd_out.read_text(encoding="utf-8"))
    sttm = SttmContract.model_validate_json(sttm_out.read_text(encoding="utf-8"))
    specs = resolve_feeds(frd, sttm, config, frd_sha256="0" * 64, sttm_sha256="0" * 64)
    assert len(specs) == 5
    assert len({spec.feed_id for spec in specs}) == 5


def test_without_own_file_rows_the_sheet_names_the_feed(env, capsys):
    _frd_out, sttm_out = _chain(env, capsys, own_file_rows=False)
    sttm = SttmContract.model_validate_json(sttm_out.read_text(encoding="utf-8"))
    assert [f.feed_id for f in sttm.feeds] == ["cm_north", "cm_south", "cm_east", "cm_west",
                                               "cm_additional"]


def test_a_layout_split_flag_names_every_renamed_sheet(env, capsys):
    """The per-sheet naming is said, never silent (flag ``feed_names_per_sheet``)."""
    from openpyxl import load_workbook

    from codegen.layout.resolve import _sttm_facts, resolve_workbook, split_frd_feeds_by_sttm

    config = load_config(CONFIG)
    sttm = env / "STTM_nb_care_mgmt.xlsx"
    fm.build_sttm().save(sttm)
    doc, _wb = resolve_workbook(sttm, config, provider=None, base_dir=REPO)
    facts = _sttm_facts(doc.profile, load_workbook(sttm), config)
    frd_path = env / "FRD.docx"
    frd_path.write_bytes(fm.build_frd())
    frd_out = env / "frd.json"
    assert cli.main(["layout", "--workbook", str(sttm), "--frd", str(frd_path), "--dry-run",
                     "--no-cache", "--frd-contract-out", str(frd_out)]) == 0
    capsys.readouterr()
    # Re-split the PRE-split contract: the one-feed FRD the docx reader produces.
    from codegen.extract.frd_docx import read_frd
    from codegen.layout.resolve import resolve_frd

    frd_doc, content = resolve_frd(frd_path, config, provider=None, base_dir=REPO)
    contract = read_frd(content, frd_doc.profile, config, document_name=frd_path.name,
                        generated_date="1970-01-01")
    split, flags, questions = split_frd_feeds_by_sttm(contract, facts, {}, config)
    assert len(split.feeds) == 5 and not questions
    (flag,) = [f for f in flags if f.startswith("feed_names_per_sheet:")]
    for sheet in fm.FILES:
        assert repr(sheet) in flag


def test_an_unsplit_frd_is_split_by_extract_sttm(env, capsys):
    """extract-frd's ONE feed against the five sheets. post38 stopped with a
    named error; since the D1 live check (2026-10-13) extract-sttm splits it
    itself, by the layout stage's rule — five feeds, never a duplicate id
    (tests/test_frd_split_default.py runs it through generate)."""
    sttm, frd = env / "STTM_cm.xlsx", env / "FRD_cm.docx"
    fm.build_sttm().save(sttm)
    frd.write_bytes(fm.build_frd())
    one_feed = env / "frd_one.contract.json"
    assert cli.main(["extract-frd", "--docx", str(frd), "--out", str(one_feed)]) == 0
    contract = json.loads(one_feed.read_text(encoding="utf-8"))
    contract["feeds"][0]["file_name_patterns"] = list(fm.FILES.values())
    one_feed.write_text(json.dumps(contract), encoding="utf-8")
    capsys.readouterr()
    code = cli.main(["extract-sttm", "--workbook", str(sttm), "--frd-contract", str(one_feed),
                     "--out", str(env / "sttm.json"), "--generated-date", DATE])
    out = capsys.readouterr().out
    assert code == 0, out
    assert "frd_feeds_split_at_extract:" in out and "duplicate feed_id(s)" not in out
    feeds = json.loads((env / "sttm.json").read_text(encoding="utf-8"))["feeds"]
    assert len({f["feed_id"] for f in feeds}) == 5
