"""Multi-table Phase C step 4 — DATA_QUALITY_RULES keyed per file.

docs/acfc/MULTI_TABLE_DESIGN.md rule 5 + decision (a) 2026-10-07: one
LoadHeaderAndTrailerToSeparateTablesRule row per FILE when the header /
trailer segments land in a table other than the detail (pair 4: 6 rows; pair
1: none — its three segments are one table). OBJECT_ID = the file's ADLS
OBJECT_ID. Rule classes no input derives are not generated; the review copy
lists one 'additional DQ rules — Engineer' entry per file.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from codegen.iig_review import dq_review_entries
from multi_table_golden import compare_to_golden

SPLIT = "LoadHeaderAndTrailerToSeparateTablesRule"


def test_pair4_dq_is_one_split_row_per_file_and_matches_the_golden(pair4_payload, pair4_golden):
    assert compare_to_golden(pair4_payload, pair4_golden, "DATA_QUALITY_RULES") == []
    rows = pair4_payload["tabs"]["DATA_QUALITY_RULES"]["rows"]
    adls = pair4_payload["tabs"]["ADLS_DELTA_INGESTION_DETAILS"]["rows"]
    assert [r["values"]["INPUT_PARAM"] for r in rows] == [
        r["values"]["SRC_FILE_NAME"] for r in adls]
    assert [r["values"]["OBJECT_ID"] for r in rows] == [r["values"]["OBJECT_ID"] for r in adls]
    tooltip = rows[0]["badges"]["TARGET_COLUMN"]["tooltip"]
    assert "rule 5" in tooltip and "catalog_map PR_DLK → d1_dlk" in tooltip


def test_pair1_gets_no_split_row_and_keeps_its_eight_derived_rules(pair1_payload):
    rows = pair1_payload["tabs"]["DATA_QUALITY_RULES"]["rows"]
    assert len(rows) == 8                                   # 4 files × (date, cast)
    assert SPLIT not in {r["values"]["RULE_CLASS"] for r in rows}
    assert [r["values"]["OBJECT_ID"] for r in rows] == ["1", "1", "2", "2", "3", "3", "4", "4"]
    assert [r["values"]["SEQUENCE_NO"] for r in rows] == [1, 2] * 4


def test_one_additional_dq_rules_entry_per_file(pair4_config, pair4_payload, pair1_payload,
                                                pair1_config):
    entries = dq_review_entries(pair4_payload, pair4_config)
    assert len(entries) == 6
    assert {(e["sheet"], e["reason"], e["owner"], e["cells"]) for e in entries} == {
        ("DATA_QUALITY_RULES", "additional_dq_rules", "engineer", 0)}
    assert "NWB_COB_RPT_110_*.txt (OBJECT_ID 1)" in entries[0]["example_citation"]
    assert len(dq_review_entries(pair1_payload, pair1_config)) == 4


def test_the_review_workbook_lists_the_entries(pair4_config, pair4_payload, tmp_path):
    from codegen.iig_review import write_iig_workbooks

    review, _clean, _cells = write_iig_workbooks(
        pair4_payload, pair4_payload["always_blank"], pair4_config, tmp_path, "nb_cob_report",
        "iig_v2")
    summary = load_workbook(review)["REVIEW_SUMMARY"]
    rows = [r for r in summary.iter_rows(min_row=2, values_only=True)
            if r[2] == "additional DQ rules (not generated)"]
    assert len(rows) == 6 and {r[3] for r in rows} == {"Engineer"}


@pytest.mark.parametrize("feed_fixture", ["pair4_payload", "pair1_payload"])
def test_sequence_numbers_restart_per_file(feed_fixture, request):
    rows = request.getfixturevalue(feed_fixture)["tabs"]["DATA_QUALITY_RULES"]["rows"]
    by_object: dict[str, list] = {}
    for row in rows:
        by_object.setdefault(row["values"]["OBJECT_ID"], []).append(row["values"]["SEQUENCE_NO"])
    assert all(seq == list(range(1, len(seq) + 1)) for seq in by_object.values())
