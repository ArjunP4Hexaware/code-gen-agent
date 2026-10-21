"""post68: the notebook output path goes through the same bookkeeping as the
workbook (codegen.generate_run.GenerateRun). Synthetic pair 11, mock only.

1. A feed with a blocking line open (its own, or a run-level one such as a
   required environment value) writes no notebook. On a notebook-only run it
   keeps no emit artefacts either and its manifest entry says ``written:
   false``. Next to Framework, the workbook is still written for review
   (post58); only the notebook files go. The required-environment check is the
   same for the CLI and the App whatever the output mode.
2. One input set: the answers file's run.environment, the owed inputs and the
   input choices apply on a notebook run exactly as on a framework run — the
   same open lines, the same counts.
3. The manifest lists the notebook output (.ipynb, pipeline/*.py, job/*) as the
   notebook_emit stage; a crash mid-feed removes it like the workbook.
4. The environment profile reaches the notebook: feed_spec states the
   catalogs, the storage account and the landing / archive / reject paths (the
   workbook's ADLS_DELTA cells as abfss URIs), never a None default, no account
   in a widget; a value the profile leaves unset is a blocking line. The
   create_tables=false prerequisites list every standard table.
5. Contact addresses a document's notification requirement quotes go through
   the configured job.notification_emails (CODEGEN_NOTIFICATION_EMAILS): the
   job's workflow.json carries the configured list, never the document's.
"""

from __future__ import annotations

from codegen import cli, run_stages
from codegen.generate_run import run_level_lines
from test_iig_corpus_cells import _pair11


def _run(config, pair, parts, **kw) -> int:
    return cli._run_pairs([pair], config, dry_run=True, skip_tests=True, output_mode=parts,
                          conventions_profile="acfc_prx", iig_template="iig_v2", **kw)


def _v1(config):
    """``config`` with the acfc_prx notebook pinned to v1 (post69: acfc_prx renders
    v2 — these post68 checks are about the v1 module tree)."""
    profiles = dict(config.conventions.profiles)
    profiles["acfc_prx"] = profiles["acfc_prx"].model_copy(update={"notebook_version": "v1"})
    return config.model_copy(update={"conventions": config.conventions.model_copy(
        update={"profiles": profiles})})


def _files(root, suffix: str) -> list:
    return sorted(p for p in root.rglob(f"*{suffix}") if run_stages.QUARANTINE_DIR not in p.parts)


# ------------------------------------------------- 1. blocking line -> no notebook


def test_without_a_blocking_line_the_notebook_is_written(tmp_path):
    config, specs, pair = _pair11(tmp_path)
    code = _run(config, pair, ["notebook"], only_feed=specs[0].feed_id)
    assert code == cli.EXIT_OK
    assert len(_files(tmp_path / "out", ".ipynb")) == 1
    (feed,) = run_stages.read_stage_manifest(tmp_path / "out")["feeds"]
    assert feed["written"] is True and "notebook_withheld" not in feed


def test_a_blocking_line_writes_no_notebook_on_a_notebook_only_run(
        tmp_path, capsys, no_environment_profiles):
    config, specs, pair = _pair11(tmp_path)
    feed_id = specs[0].feed_id
    code = _run(config, pair, ["notebook"], only_feed=feed_id)
    out = capsys.readouterr().out
    assert code == cli.EXIT_NEEDS_ANSWERS
    assert "QUESTION       run.environment.storage_account — " in out
    assert _files(tmp_path / "out", ".ipynb") == []
    feed_dir = tmp_path / "out" / specs[0].feed_slug
    assert not (feed_dir / "pipeline").exists() and not (feed_dir / "job").exists()
    assert _files(feed_dir, ".sql") == []                    # nothing deployable kept
    manifest = run_stages.read_stage_manifest(tmp_path / "out")
    assert manifest["stages"]["generate"]["status"] == "NEEDS_ANSWERS"
    (feed,) = manifest["feeds"]
    assert feed["written"] is False and feed["notebook_withheld"] > 0
    assert (f"NOTE           feeds.{feed_id}.notebook — not written — 1 blocking line(s) open "
            "(run.environment.storage_account); answer them and re-run") in out


def test_next_to_framework_only_the_notebook_is_withheld(tmp_path, no_environment_profiles):
    config, specs, pair = _pair11(tmp_path)
    code = _run(config, pair, ["notebook", "framework"], only_feed=specs[0].feed_id)
    assert code == cli.EXIT_NEEDS_ANSWERS
    out = tmp_path / "out"
    assert _files(out, ".ipynb") == []
    assert len(_files(out, "_IIG.xlsx")) == 1                 # the workbook, for review
    (feed,) = run_stages.read_stage_manifest(out)["feeds"]
    assert feed["written"] is True and feed["notebook_withheld"] > 0


def test_a_feed_level_blocking_line_withholds_that_feed_only(tmp_path):
    from codegen.generate_run import GenerateRun

    config, _specs, _pair = _pair11(tmp_path)
    run = GenerateRun(config, out_root=tmp_path, output_mode=["notebook"])
    run.open_items += [run_stages.open_item("QUESTION", "feeds[0].x", feed="a"),
                       run_stages.open_item("INPUT_NEEDED", "feeds.b.y", feed="b")]
    assert [i["key"] for i in run.blocking_for("a")] == ["feeds[0].x"]
    assert run.blocking_for("b") == []                         # INPUT_NEEDED never blocks


def test_the_required_environment_check_ignores_the_output_mode(tmp_path, no_environment_profiles):
    config, _specs, _pair = _pair11(tmp_path)
    asked = {}
    for parts in (["notebook"], ["framework"], ["notebook", "framework"]):
        questions, _items = run_level_lines(config, {}, output_mode=parts,
                                            iig_template="iig_v2")
        asked[tuple(parts)] = [key for key, _reason in questions]
    assert set(map(tuple, asked.values())) == {("run.environment.storage_account",)}


def test_the_app_lists_the_required_value_on_a_notebook_run(tmp_path, no_environment_profiles):
    from ui.backend.demo import DemoRunner
    from ui.backend.service import GenerationStore

    config, _specs, _pair = _pair11(tmp_path)
    runner = DemoRunner(GenerationStore("config/config.yaml"), work=lambda: None,
                        index_warmup=False)
    runner.output_parts = ["notebook"]
    runner.iig_template = "iig_v2"
    required = [i["key"] for i in runner._run_level_inputs(config) if i.get("required")]
    assert required == ["run.environment.storage_account"]


# ------------------------------------------------------------- 2. one input set


def _open_lines(out) -> tuple[list, list]:
    manifest = run_stages.read_stage_manifest(out)
    return (sorted((q["line"], q["key"]) for q in manifest["questions"]),
            [f["open"] for f in manifest["feeds"]])


def test_notebook_and_framework_runs_show_the_same_open_lines(tmp_path):
    config, specs, pair = _pair11(tmp_path)
    by_mode = {}
    for parts in (["notebook"], ["framework"]):
        out = tmp_path / parts[0]
        run_config = config.model_copy(update={"output": config.output.model_copy(
            update={"dir": str(out), "reports_dir": str(out / "reports")})})
        assert _run(run_config, pair, parts, only_feed=specs[0].feed_id) == cli.EXIT_OK
        by_mode[parts[0]] = _open_lines(out)
    assert by_mode["notebook"] == by_mode["framework"]
    lines, counts = by_mode["notebook"]
    assert ("INPUT_NEEDED", "run.environment") in lines          # the environment assumed
    assert counts[0] == {"blocking": 0, "non_blocking": 0}        # post81: defaults written


def test_the_answers_file_environment_applies_on_a_notebook_run(
        tmp_path, capsys, no_environment_profiles):
    config, specs, pair = _pair11(tmp_path)
    answers = tmp_path / "answers.yaml"
    answers.write_text("inputs:\n  run:\n    environment: {value: D1}\n"
                       '    environment.storage_account: {value: "synstorageacct"}\n',
                       encoding="utf-8")
    code = _run(config, pair, ["notebook"], only_feed=specs[0].feed_id, answers_path=answers)
    out = capsys.readouterr().out
    assert code == cli.EXIT_OK
    assert "QUESTION       run.environment.storage_account" not in out
    assert "NOTE           run.environment — environment: D1" in out
    assert len(_files(tmp_path / "out", ".ipynb")) == 1


# ------------------------------------------------ 3. the manifest's notebook_emit


def test_the_manifest_lists_the_notebook_outputs(tmp_path):
    config, specs, pair = _pair11(tmp_path)
    config = _v1(config)
    assert _run(config, pair, ["notebook"], only_feed=specs[0].feed_id) == cli.EXIT_OK
    manifest = run_stages.read_stage_manifest(tmp_path / "out")
    assert manifest["stages"]["notebook_emit"] == {"status": "PASS"}
    (feed,) = manifest["feeds"]
    listed = feed["artefacts"]["notebook_emit"]
    slug = specs[0].feed_slug
    assert f"{slug}/{slug}.ipynb" in listed
    assert any(p.startswith(f"{slug}/pipeline/") and p.endswith(".py") for p in listed)
    assert any(p.startswith(f"{slug}/job/") for p in listed)
    assert not any(p.startswith(f"{slug}/tests/") for p in listed)


def test_a_framework_run_skips_notebook_emit(tmp_path):
    config, specs, pair = _pair11(tmp_path)
    assert _run(config, pair, ["framework"], only_feed=specs[0].feed_id) == cli.EXIT_OK
    stages = run_stages.read_stage_manifest(tmp_path / "out")["stages"]
    assert stages["notebook_emit"] == {"status": "SKIPPED",
                                       "reason": run_stages.SKIPPED_NOT_WRITTEN}


def test_a_crash_removes_the_partial_notebook_output(tmp_path):
    import pytest

    from codegen.generate_run import GenerateRun

    config, specs, _pair = _pair11(tmp_path)
    out = tmp_path / "crash"
    run = GenerateRun(config, out_root=out, output_mode=["notebook"])
    feed_dir = out / specs[0].feed_slug

    def produce(_faq):
        for rel in ("x.ipynb", "pipeline/reader.py", "job/workflow.json", "ddl/t.sql"):
            (feed_dir / rel).parent.mkdir(parents=True, exist_ok=True)
            (feed_dir / rel).write_text("partial", encoding="utf-8")
        raise RuntimeError("mid-write")

    with pytest.raises(RuntimeError):
        run.feed(specs[0], 1, 1, frd_path=None, produce=produce)
    assert sorted(p.name for p in feed_dir.rglob("*") if p.is_file()) == []
    manifest = run_stages.read_stage_manifest(out)
    assert manifest["stages"]["notebook_emit"]["status"] == "SKIPPED"
    assert manifest["feeds"][0]["removed"] == 4


# --------------------------------------------- 4. the environment in the notebook


def _feed_spec(out) -> dict[str, str]:
    import ast

    (path,) = list(out.rglob("pipeline/feed_spec.py"))
    tree = ast.parse(path.read_text(encoding="utf-8"))
    values = {}
    for node in tree.body:
        target = (node.targets[0] if isinstance(node, ast.Assign) and len(node.targets) == 1
                  else node.target if isinstance(node, ast.AnnAssign) else None)
        if isinstance(target, ast.Name) and node.value is not None:
            try:
                values[target.id] = ast.literal_eval(node.value)
            except ValueError:
                continue
    return values


def test_feed_spec_states_the_environment_profile(tmp_path):
    from codegen import environments as envs

    config, specs, pair = _pair11(tmp_path)
    config = _v1(config)
    parts = ["notebook", "framework"]
    assert _run(config, pair, parts, only_feed=specs[0].feed_id) == cli.EXIT_OK
    out = tmp_path / "out"
    spec_values = _feed_spec(out)
    env = envs.run_environment(config)
    account = env.entry.storage_account
    assert spec_values["DEFAULT_CATALOG"] == envs.default_catalog(config, "stage")
    assert spec_values["STANDARD_CATALOG"] == envs.default_catalog(config, "standard")
    assert spec_values["STORAGE_ACCOUNT"] == account
    for key in ("DEFAULT_CATALOG", "STANDARD_CATALOG", "STORAGE_ACCOUNT", "LANDING_PATH",
                "ARCHIVE_PATH", "REJECT_PATH"):
        assert spec_values[key], key                              # never None / blank
    source = next(out.rglob("pipeline/feed_spec.py")).read_text(encoding="utf-8")
    assert "str | None = None" not in source.split("STAGE_SCHEMA")[0]
    # the paths are the workbook's own ADLS_DELTA cells, on the environment's account
    from openpyxl import load_workbook

    (iig,) = list(out.rglob("*_IIG.xlsx"))
    sheet = load_workbook(iig, read_only=True)["ADLS_DELTA_INGESTION_DETAILS"]
    rows = list(sheet.iter_rows(values_only=True))
    row = dict(zip(rows[0], rows[1], strict=False))
    for key, header, container in (("LANDING", "SRC_ADLS_PATH", "SRC_CONTAINER_NAME"),
                                   ("ARCHIVE", "SRC_ADLS_ARCHVL_PATH", "SRC_CONTAINER_NAME"),
                                   ("REJECT", "TGT_RJT_ADLS_PATH", "TGT_CONTAINER_NAME")):
        assert spec_values[f"{key}_PATH"] == row[header], key
        assert spec_values[f"{key}_CONTAINER"] == row[container], key
    # the abfss:// address is built at run time (adls_uri) ...
    import ast

    (func,) = [n for n in ast.parse(source).body
               if isinstance(n, ast.FunctionDef) and n.name == "adls_uri"]
    namespace = {"STORAGE_ACCOUNT": account, "_ADLS_HOST": spec_values["_ADLS_HOST"]}
    exec(compile(ast.Module(body=[func], type_ignores=[]), "feed_spec", "exec"), namespace)
    uri = namespace["adls_uri"](spec_values["LANDING_CONTAINER"], spec_values["LANDING_PATH"])
    assert uri == (f"abfss://{row['SRC_CONTAINER_NAME']}@{account}.dfs.core.windows.net/"
                   f"{str(row['SRC_ADLS_PATH']).lstrip('/')}")
    # ... so the generated text spells no ADLS host and no address-shaped string
    import sys

    from conftest import REPO
    sys.path.insert(0, str(REPO / "scripts"))
    from scrub_check import scan_paths

    notebook = [p for p in out.rglob("*") if p.is_file() and run_stages.is_notebook_file(p)]
    assert notebook and scan_paths(notebook) == []
    # no account name in a widget: widgets carry the catalog and blanks only
    entry = next(out.rglob("job/notebook_entrypoint.py")).read_text(encoding="utf-8")
    widgets = [ln for ln in entry.splitlines() if "widgets.text(" in ln]
    assert widgets and not any(account in ln for ln in widgets)
    assert f'dbutils.widgets.text("catalog", "{spec_values["DEFAULT_CATALOG"]}")' in entry


def test_a_catalog_the_profile_leaves_unset_withholds_the_notebook(tmp_path, capsys):
    from codegen import environments as envs

    config, specs, pair = _pair11(tmp_path)
    named = envs.run_environment(config).entry
    custom = {"storage_account": named.storage_account,
              **{f"containers.{k}": v for k, v in named.containers.items()}}
    config = envs.with_environment(config, envs.resolve(config, custom=custom))
    code = _run(config, pair, ["notebook"], only_feed=specs[0].feed_id)
    out = capsys.readouterr().out
    assert code == cli.EXIT_NEEDS_ANSWERS
    assert "QUESTION       run.environment.default_catalog.stage — " in out
    assert _files(tmp_path / "out", ".ipynb") == []
    # a framework-only run of the same environment is not blocked by it (its
    # catalog cells are written open: INPUT_NEEDED)
    assert _run(config, pair, ["framework"], only_feed=specs[0].feed_id) == cli.EXIT_OK
    assert "INPUT_NEEDED   run.environment.default_catalog.stage — " in capsys.readouterr().out


def test_without_an_environment_block_the_contract_values_stay(tmp_path):
    from codegen.emit.notebook_env import notebook_environment

    config, specs, _pair = _pair11(tmp_path)
    bare = config.model_copy(update={"environments": {}, "run_environment": None})
    found = notebook_environment(specs[0], bare)
    assert found.environment is None and found.values == {} and found.missing == []


def test_the_prerequisites_list_every_standard_table():
    from codegen.emit.notebook import _prerequisite_cell

    context = {"default_catalog": "cat", "stage_schema": "stg",
               "segments": [{"stage_table": "t"}], "errors_table": "t_errors",
               "processed_files_table": "t_files", "recycle": None,
               "standard": {"catalog": "std", "schema": "s", "table": "a"},
               "standard_tables": [{"catalog": "std", "schema": "s", "table": "a"},
                                   {"catalog": "std", "schema": "s", "table": "b"}]}
    text = "".join(_prerequisite_cell(context)["source"])
    assert "std.s.a" in text and "std.s.b" in text


# ----------------------------------------------- 5. document addresses scrubbed


def test_redact_contacts_uses_the_configured_list():
    from codegen.notification_contacts import NO_EMAILS, redact_contacts

    text = ("Alert the Support team (dl.ops@client-corp.org, cc oncall@client-corp.org) "
            "on failure; see client-corp.org/runbook or @client-corp.")
    out = redact_contacts(text, ["syn.dl@synthetic.example"])
    assert "client-corp" not in out
    assert out.count("syn.dl@synthetic.example") == 2
    assert "synthetic.example/runbook" in out and "@synthetic.example." in out
    assert redact_contacts("no address here", ["a@b.example"]) == "no address here"
    assert NO_EMAILS in redact_contacts("mail x@y.org", [])


def test_the_sfmc_notebook_outputs_carry_no_document_address(tmp_path, monkeypatch):
    import json
    import sys

    from codegen.config import load_config
    from conftest import REPO

    monkeypatch.setenv("CODEGEN_NOTIFICATION_EMAILS", "syn.dl.prodsupport@synthetic.example")
    config = load_config(REPO / "config" / "config.yaml")
    config = config.model_copy(update={"output": config.output.model_copy(
        update={"dir": str(tmp_path / "out"), "reports_dir": str(tmp_path / "reports")})})
    pair = (REPO / "fixtures" / "contracts" / "FRD_sfmc_email_campaign.contract.json",
            REPO / "fixtures" / "contracts" / "sttm_mapping_contracts_sfmc.json")
    from conftest import require_fixture_files

    require_fixture_files(*pair)         # the SFMC contracts are not in every checkout
    code = cli._run_pairs([pair], config, dry_run=True, skip_tests=True,
                          output_mode=["notebook"], only_feed=None)
    assert code in (cli.EXIT_OK, cli.EXIT_NEEDS_ANSWERS)
    (workflow,) = list((tmp_path / "out").rglob("job/workflow.json"))
    job = json.loads(workflow.read_text(encoding="utf-8"))
    rules = job["_notification_rules_from_contract"]
    assert rules and all("syn.dl.prodsupport@synthetic.example" in r for r in rules)
    sys.path.insert(0, str(REPO / "scripts"))
    from scrub_check import scan_paths

    notebook = [p for p in (tmp_path / "out").rglob("*")
                if p.is_file() and run_stages.is_notebook_file(p)]
    assert scan_paths(notebook) == []
