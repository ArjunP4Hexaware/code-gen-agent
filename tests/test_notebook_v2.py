"""post69: notebook v2 — rendered from the IIG rows, selected by the profile.

Four synthetic fixtures (tests/notebook_v2_fixtures.py): SD-shaped (three
delimited feeds), SFMC, a CAQH-shaped single-table TPL feed and a PRX-shaped
fixed-width feed with three segments; plus an unknown RULE_CLASS and an open
blocking input. Every emitted notebook: nbformat-valid, every code cell
compiles, ruff-clean .py, the scrub passes, App and CLI byte-identical.
"""

from __future__ import annotations

import ast
import json
import re
import subprocess
import sys
from pathlib import Path

import pytest

import notebook_v2_fixtures as fx
from codegen import cli, run_stages
from codegen.emit import notebook_v2
from codegen.emit.notebook_files import COMMAND, DATABRICKS_HEADER, check_ruff

PROFILE = {"conventions_profile": "acfc_prx", "iig_template": "iig_v2"}


def _run(config, pairs, parts=("notebook",), **kw) -> int:
    # post70: the SFMC-shaped feeds' merge key, typed as the BSA answers it
    kw.setdefault("answers_path", fx.key_answers(Path(config.output.dir).parent))
    return cli._run_pairs(list(pairs), config, dry_run=True, skip_tests=True,
                          output_mode=list(parts), only_feed=kw.pop("only_feed", None),
                          **PROFILE, **kw)


def _notebooks(out: Path) -> list[Path]:
    return sorted(p for p in out.rglob("notebook/*.ipynb") if "quarantine" not in p.parts)


def _cells(ipynb: Path) -> list[dict]:
    return json.loads(ipynb.read_text(encoding="utf-8"))["cells"]


def _config_values(ipynb: Path) -> dict:
    """The config cell's literal constants (SOURCES, STANDARDS, ...)."""
    (config,) = [c for c in _cells(ipynb) if c["id"] == "config"]
    values = {}
    for node in ast.parse("".join(config["source"])).body:
        if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name):
            try:
                values[node.targets[0].id] = ast.literal_eval(node.value)
            except ValueError:
                continue
    return values


def _assert_valid(out: Path, feeds: int) -> list[Path]:
    import nbformat

    sys.path.insert(0, str(fx.REPO / "scripts"))
    from scrub_check import scan_paths

    notebooks = _notebooks(out)
    assert len(notebooks) == feeds
    for ipynb in notebooks:
        nbformat.validate(nbformat.reads(ipynb.read_text(encoding="utf-8"), as_version=4))
        for cell in _cells(ipynb):
            if cell["cell_type"] == "code":
                compile("".join(cell["source"]), cell["id"], "exec")
        source = ipynb.with_suffix(".py").read_text(encoding="utf-8")
        assert source.startswith(DATABRICKS_HEADER + "\n")
        assert source.count(COMMAND) == len(_cells(ipynb)) - 1
        assert "# MAGIC %md" in source
        assert check_ruff(source).passed, check_ruff(source).details
        assert scan_paths([ipynb, ipynb.with_suffix(".py")]) == []
        # no abfss:// address literal: only adls_uri() builds one, at run time
        assert not re.search(r"abfss://[\w-]+@", source)
    return notebooks


# ------------------------------------------------------------ the four fixtures


def test_sd_shaped_three_delimited_feeds(tmp_path):
    config = fx.acfc_config(tmp_path)
    assert _run(config, [fx.sd_pair()]) == cli.EXIT_OK
    notebooks = _assert_valid(tmp_path / "out", 3)
    for ipynb in notebooks:
        values = _config_values(ipynb)
        assert values["SOURCES"] and values["STANDARDS"]
        assert all(s["delimiter"] for s in values["SOURCES"])
        assert "SEGMENTS" not in values                       # delimited: no handler rows


def test_sfmc_fixture(tmp_path):
    config = fx.acfc_config(tmp_path)
    assert _run(config, [fx.sfmc_pair()]) == cli.EXIT_OK
    _assert_valid(tmp_path / "out", 1)


def test_caqh_shaped_single_table_tpl_feed(tmp_path):
    config = fx.acfc_config(tmp_path)
    assert _run(config, [fx.tpl_pair(tmp_path)]) == cli.EXIT_OK
    (ipynb,) = _assert_valid(tmp_path / "out", 1)
    (source,) = _config_values(ipynb)["SOURCES"]
    assert source["delimiter"] == "|" and len(_config_values(ipynb)["STANDARDS"]) == 1


def test_prx_shaped_fixed_width_three_segments(tmp_path):
    config = fx.acfc_config(tmp_path, fx.PAIR1_OVERLAY)
    frd, sttm, vdd = fx.prx_pair(tmp_path, config)
    assert _run(config, [(frd, sttm)], vdd_path=vdd) == cli.EXIT_OK
    (ipynb,) = _assert_valid(tmp_path / "out", 1)
    values = _config_values(ipynb)
    segments = values["SEGMENTS"]
    assert [s["segment"] for s in segments] == ["HDDR", "DET", "TRLR"]
    assert [s["filter"] for s in segments] == ["value like '00%'", "value like '10%'",
                                               "value like '99%'"]
    assert values["DETAIL_SEGMENT"] == "DET"
    assert all(s["columns"] and all(isinstance(n, int) for _c, n, _l in s["columns"])
               for s in segments)


# ------------------------------------------------- the IIG is the source of truth


def test_every_config_value_is_an_iig_cell(tmp_path):
    from openpyxl import load_workbook

    from codegen import environments as envs

    config = fx.acfc_config(tmp_path)
    assert _run(config, [fx.sfmc_pair()], parts=("notebook", "framework")) == cli.EXIT_OK
    (ipynb,) = _notebooks(tmp_path / "out")
    values = _config_values(ipynb)
    (iig,) = [p for p in (tmp_path / "out").rglob("*_IIG.xlsx")]
    book = load_workbook(iig, read_only=True)

    def rows(sheet):
        data = list(book[sheet].iter_rows(values_only=True))
        return [dict(zip(data[0], r, strict=False)) for r in data[1:]]

    def text(value):
        return "" if value is None else str(value).strip()

    for source, row in zip(values["SOURCES"], rows(notebook_v2.ADLS), strict=True):
        for key, header in (("file_pattern", "SRC_FILE_NAME"),
                            ("delimiter", "SRC_FILE_DELIMITER"),
                            ("landing_container", "SRC_CONTAINER_NAME"),
                            ("landing_path", "SRC_ADLS_PATH"),
                            ("archive_path", "SRC_ADLS_ARCHVL_PATH"),
                            ("reject_container", "TGT_CONTAINER_NAME"),
                            ("reject_path", "TGT_RJT_ADLS_PATH"),
                            ("schema", "TGT_DATABASE_NAME"), ("table", "TGT_TABLE_NAME"),
                            ("load_option", "TGT_LOAD_OPTION"), ("lob", "LOB")):
            assert source[key] == text(row[header]), key
        assert source["target_columns"] == notebook_v2.split_list(row["TGT_COLUMN_NAMES"])
        assert source["mandatory"] == notebook_v2.split_list(row["MANDATORY_FIELD_LIST"])
        assert [c for _s, c in source["src_columns"]] == [
            pair.split(":")[-1] for pair in notebook_v2.split_list(row["SRC_COLUMNS"])]
    for standard, row in zip(values["STANDARDS"], rows(notebook_v2.STD), strict=True):
        for key, header in (("src_catalog", "SRC_CATALOG_NAME"), ("catalog", "TGT_CATALOG_NAME"),
                            ("schema", "TGT_SCHEMA_NAME"), ("table", "TGT_TABLE_NAME"),
                            ("load_option", "TGT_LOAD_OPTION")):
            assert standard[key] == text(row[header]), key
    assert values["STAGE_CATALOG"] == text(rows(notebook_v2.STD)[0]["SRC_CATALOG_NAME"])
    # the one value from the environment profile: the storage account
    assert values["STORAGE_ACCOUNT"] == envs.run_environment(config).entry.storage_account


# ------------------------------------------------------------- the cell layout


def test_the_cell_layout(tmp_path):
    config = fx.acfc_config(tmp_path)
    _run(config, [fx.sfmc_pair()])
    (ipynb,) = _notebooks(tmp_path / "out")
    cells = _cells(ipynb)
    ids = [c["id"] for c in cells]
    sources = _config_values(ipynb)["SOURCES"]
    # post70: drift / recycle only where the IIG states them; the run report last
    drift = ["drift", "drift-provenance"] if any(s["schema_drift"] for s in sources) else []
    recycle = ["recycle", "recycle-provenance"] if any(s["recycle"] for s in sources) else []
    assert recycle                  # the SFMC IIG row: RECYCL_ENBL_FLG = Y
    assert ids[:9 + len(drift) + len(recycle)] == [
        "title", "config", "read", "read-provenance", *drift, "stage", "stage-provenance",
        *recycle, "standard", "standard-provenance", "dq"]
    tail = ids[ids.index("run"):]
    # post76: the review cell is the TRAILING %md (the only cell --layer2 changes)
    assert tail in (["run", "run-report", "run-report-provenance"],
                    ["run", "run-report", "run-report-provenance", "review-candidates"])
    assert "dq-provenance" in ids
    title = "".join(cells[0]["source"])
    assert ("Standalone equivalent of the framework's generic notebooks for this IIG; "
            "the framework path is the deliverable") in title
    config_cell = "".join(cells[1]["source"])
    widgets = [ln for ln in config_cell.splitlines() if "dbutils.widgets." in ln
               and ("text(" in ln or "dropdown(" in ln)]
    assert len(widgets) == 2 and '"run_date"' in widgets[0] and '"dry_run", "true"' in widgets[1]
    run = "".join(cells[ids.index("run")]["source"])
    read = "".join(cells[2]["source"])
    assert "DRY_RUN_ROWS" in read and "limit(DRY_RUN_ROWS)" in read
    assert '"_dryrun"' in config_cell and "if not DRY_RUN:" in run
    for prov in ("read-provenance", "stage-provenance", "standard-provenance", "dq-provenance",
                 "recycle-provenance", "run-report-provenance"):
        text = "".join(cells[ids.index(prov)]["source"])
        assert "row" in text and "| Cell | Badge | From |" in text


def test_review_candidates_are_markdown_and_omitted_when_empty(tmp_path):
    config = fx.acfc_config(tmp_path)
    _run(config, [fx.sfmc_pair()])
    (ipynb,) = _notebooks(tmp_path / "out")
    review = [c for c in _cells(ipynb) if c["id"] == "review-candidates"]
    if review:
        assert review[0]["cell_type"] == "markdown"
        assert "Not executable" in "".join(review[0]["source"])
    context = {"candidates": [], "missing_rules": [], "provenance": {
        k: [] for k in ("read", "stage", "standard", "dq", "drift", "recycle", "run-report")}}
    cells_without = [c for c in notebook_v2.render_cells({**_context(tmp_path), **context})
                     if c.cell_id == "review-candidates"]
    assert cells_without == []


def _context(tmp_path) -> dict:
    """A rendered context of the SFMC feed (for render-level checks)."""
    from codegen.faq import faq_for_spec
    from codegen.metadata_sheet import metadata_sheet_payload
    from codegen.resolve.resolver import resolve_pair

    config = fx.acfc_config(tmp_path / "ctx")
    frd, sttm = fx.sfmc_pair()
    (spec,) = resolve_pair(frd, sttm, config)
    payload = metadata_sheet_payload(
        config, fx.REPO, [spec], {spec.feed_slug: set()}, "test",
        {spec.feed_slug: faq_for_spec(spec, config, base_dir=fx.REPO)}, frd,
        template="iig_v2", conventions_profile="acfc_prx")
    context, _flags = notebook_v2.build_context(spec, config, payload, candidates=[],
                                                frd_path=frd, iig_template="iig_v2")
    return context


# ---------------------------------------------------- unknown rule, blocking line


def test_an_unknown_rule_class_is_markdown_and_a_flag(tmp_path, capsys):
    config = fx.acfc_config(tmp_path)
    tpl = config.metadata.templates["iig_v2"]
    tpl = tpl.model_copy(update={"dq_rule_classes": {**tpl.dq_rule_classes,
                                                    "dedup": "SynUnknownRule"}})
    config = config.model_copy(update={"metadata": config.metadata.model_copy(
        update={"templates": {**config.metadata.templates, "iig_v2": tpl}})})
    assert _run(config, [fx.sfmc_pair()]) == cli.EXIT_OK
    out = capsys.readouterr().out
    (ipynb,) = _assert_valid(tmp_path / "out", 1)
    cells = _cells(ipynb)
    missing = [c for c in cells if c["id"].startswith("dq-not-implemented-")]
    assert missing and missing[0]["cell_type"] == "markdown"
    assert "rule not implemented: SynUnknownRule on " in "".join(missing[0]["source"])
    code = "".join("".join(c["source"]) for c in cells if c["cell_type"] == "code")
    assert "SynUnknownRule" not in code.replace("# DATA_QUALITY_RULES", "")
    reports = "".join(p.read_text(encoding="utf-8") for p in (tmp_path / "reports").glob("*.md"))
    assert "notebook_rule_not_implemented:SynUnknownRule" in reports
    assert "PASS_WITH_FLAGS" in out


def test_an_open_blocking_input_writes_no_notebook(tmp_path, no_environment_profiles):
    config = fx.acfc_config(tmp_path)
    assert _run(config, [fx.sfmc_pair()]) == cli.EXIT_NEEDS_ANSWERS
    assert _notebooks(tmp_path / "out") == []
    (feed,) = run_stages.read_stage_manifest(tmp_path / "out")["feeds"]
    assert feed["written"] is False


def test_a_gate_fail_quarantines_the_v2_notebook(tmp_path, monkeypatch):
    from codegen.emit import notebook_files
    from codegen.gate.preflight import GateCheck

    monkeypatch.setattr(notebook_files, "check_ruff", lambda _source: GateCheck(
        name="notebook_ruff", passed=False, details="forced"))
    monkeypatch.setattr(notebook_v2.nbf, "check_ruff", notebook_files.check_ruff)
    config = fx.acfc_config(tmp_path)
    assert _run(config, [fx.sfmc_pair()]) == cli.EXIT_FAILED
    manifest = run_stages.read_stage_manifest(tmp_path / "out")
    assert manifest["stages"]["generate"]["status"] == "GATE_FAIL"
    assert any("/quarantine/notebook/" in p
               for p in manifest["feeds"][0]["artefacts"]["notebook_emit"])


# ------------------------------------------------------------- v1 stays for others


def test_the_default_profile_keeps_v1(tmp_path):
    config = fx.acfc_config(tmp_path)
    code = cli._run_pairs([fx.sfmc_pair()], config, dry_run=True, skip_tests=True,
                          output_mode=["notebook"], only_feed=None,
                          conventions_profile="edo_sfmc", iig_template="iig_v2",
                          answers_path=fx.key_answers(tmp_path))
    assert code == cli.EXIT_OK
    assert _notebooks(tmp_path / "out") == []
    assert list((tmp_path / "out").rglob("pipeline/feed_spec.py"))       # the v1 tree
    assert notebook_v2.notebook_version(config, "edo_sfmc") == "v1"
    assert notebook_v2.notebook_version(config, "acfc_prx") == "v2"


# ------------------------------------------------------------- App == CLI, bytes


def test_the_app_and_the_cli_emit_byte_identical_notebooks(tmp_path, monkeypatch):
    pytest.importorskip("fastapi")
    from ui.backend.demo import DemoRunner
    from ui.backend.service import GenerationStore

    from codegen.resolve.resolver import pending_answers, resolve_pair

    monkeypatch.setenv("CODEGEN_FORCE_MOCK_PROVIDER", "1")
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    for role in ("state", "outputs"):
        (tmp_path / role).mkdir()
        monkeypatch.setenv(f"CODEGEN_STORAGE_{role.upper()}",
                           f"local:{(tmp_path / role).as_posix()}")
    monkeypatch.chdir(fx.REPO)
    config = fx.acfc_config(tmp_path / "cli")
    frd, sttm = fx.sd_pair()
    # post76: the same Layer-2 mode (the App defers it) — under another mode only
    # the trailing review %md differs (tests/test_layer2_stage.py)
    assert _run(config, [(frd, sttm)], layer2=config.reasoning.layer2_app) == cli.EXIT_OK
    runner = DemoRunner(GenerationStore(str(fx.REPO / "config" / "config.yaml")),
                        index_warmup=False)
    runner.iig_template, runner.conventions_profile = "iig_v2", "acfc_prx"
    runner.output_parts = ["notebook"]
    runner._store.config = config                             # noqa: SLF001
    specs = resolve_pair(frd, sttm, config)
    runner._generate_specs(specs, config, run_env=None, frd_path=frd,  # noqa: SLF001
                           run_root=tmp_path / "app", reports_root=tmp_path / "app_reports",
                           held=list(pending_answers(sttm)), unstated=[], extra_flags=[],
                           layout_usage=None, sttm_path=sttm)
    cli_books = _notebooks(tmp_path / "cli" / "out")
    app_books = _notebooks(tmp_path / "app")
    assert [p.name for p in cli_books] == [p.name for p in app_books] and len(cli_books) == 3
    for ours, theirs in zip(cli_books, app_books, strict=True):
        assert ours.read_bytes() == theirs.read_bytes()
        assert ours.with_suffix(".py").read_bytes() == theirs.with_suffix(".py").read_bytes()


def test_ruff_runs_where_the_suite_runs():
    result = subprocess.run([sys.executable, "-m", "ruff", "--version"], capture_output=True,
                            text=True, check=False)
    assert result.returncode == 0


# ------------------------------------------------------------ tools/notebook_parity.py


def test_the_parity_tool_reports_v1_against_v2(tmp_path, capsys):
    sys.path.insert(0, str(fx.REPO / "tools"))
    import notebook_parity

    frd, sttm = fx.sfmc_pair()
    assert notebook_parity.main(["--frd", str(frd), "--sttm", str(sttm),
                                 "--overlay", str(fx.ACFC_OVERLAY),
                                 "--answers", str(fx.key_answers(tmp_path)),
                                 "--out", str(tmp_path / "parity")]) == 0
    out = capsys.readouterr().out
    assert "generate (v1) exit 0" in out and "generate (v2) exit 0" in out
    assert "== sfmc_email_campaign_tracking" in out
    assert "   cells: v1 " in out and "   v1 only: " in out
    assert "   DELIMITER            same" in out
    # post70: v1's stages an IIG cell drives are IIG-driven, not "only v1"
    assert "recycle (RECYCL_ENBL_FLG = Y: in v2)" in out
    assert "schema drift check (SCHEMA_DRIFT_FLAG = N: not generated)" in out
    assert "run report (always" in out
    assert "   v1 only: PHI masking (no IIG cell, docs/acfc/OPEN_QUESTIONS.md" in out
    assert "FILE_NAME_PATTERNS   expected: v2 writes the IIG's SRC_FILE_NAME" in out
    # values only with --show-values: no constant line carries "v1 <value> | v2 <value>"
    assert not re.search(r"^   [A-Z_]+ +[a-z].*: v1 .* \| v2 ", out, re.M)
