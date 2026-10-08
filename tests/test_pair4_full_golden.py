"""Chunk A (2026-10-08) — pair 4: EVERY cell of EVERY IIG sheet pinned.

The mirror of test_m4_acceptance.test_pair1_every_iig_cell_and_the_ddl_are_pinned
for the multi-table fixture: a real framework run (cli._generate_feed under the
pair-4 environment overlay, acfc_prx / iig_v2) and its config_rows.xlsx compared
cell for cell with the hand-written golden (fixtures/acfc_shapes/pair_4/golden/
IIG_EXPECTED.yaml), all eight sheets, in the template's order. Pair 4's golden is
exact: no open / deviation classes — a golden null or "" means the cell is blank.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from codegen import cli
from multi_table_golden import golden_rows
from test_m4_acceptance import _scoped, _sheet_rows


@pytest.fixture(scope="module")
def pair4_run(pair4_config, pair4_spec, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pair4_run")
    gate = cli._generate_feed(pair4_spec, _scoped(pair4_config, tmp), dry_run=True,
                              skip_tests=True, output_mode="framework",
                              conventions_profile="acfc_prx", iig_template="iig_v2")
    return gate, tmp / "out" / pair4_spec.feed_slug / "framework"


def test_pair4_every_iig_cell_matches_the_golden(pair4_run, pair4_golden, pair4_config):
    _gate, framework_dir = pair4_run
    workbook = load_workbook(framework_dir / "config_rows.xlsx")
    template = pair4_config.metadata.templates["iig_v2"]
    assert [s for s in workbook.sheetnames if not s.startswith("_")] == list(template.tabs)
    diffs = []
    for sheet in template.tabs:
        headers, rows = _sheet_rows(workbook[sheet])
        expected = golden_rows(pair4_golden, sheet)
        assert headers == list(template.tabs[sheet].headers), sheet
        assert len(rows) == len(expected), (sheet, len(rows), len(expected))
        for number, (row, want) in enumerate(zip(rows, expected, strict=True), start=1):
            for column, ours in zip(headers, row, strict=True):
                if ours != (want[column] or ""):
                    diffs.append((sheet, number, column, want[column], ours))
    assert diffs == []


def test_pair4_clean_iig_equals_config_rows(pair4_run):
    _gate, framework_dir = pair4_run
    (clean,) = [p for p in framework_dir.glob("*_IIG.xlsx")]
    clean_wb = load_workbook(clean)
    rows_wb = load_workbook(framework_dir / "config_rows.xlsx")
    for name in clean_wb.sheetnames:
        assert [list(r) for r in clean_wb[name].iter_rows(values_only=True)] == \
               [list(r) for r in rows_wb[name].iter_rows(values_only=True)], name
