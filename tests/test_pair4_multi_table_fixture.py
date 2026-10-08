"""Multi-table Phase B — the pair-4 fixture (fixtures/acfc_shapes/pair_4/).

Pins the synthetic multi-table STTM / FRD contract (built from
tests/acfc_shapes/pair4_nb.py — byte stability is asserted with the rest of
the universe in test_acfc_shapes_fixtures.py) and cross-checks the
HAND-WRITTEN golden ``golden/IIG_EXPECTED.yaml`` against the builder facts,
the environment overlay and the owner map in docs/acfc/MULTI_TABLE_DESIGN.md.
Nothing here runs the generator: Phase C makes it match the golden sheet by
sheet (docs/acfc/MULTI_TABLE_DESIGN.md §6).
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from acfc_shapes import FIXTURE_ROOT, pair4_nb
from codegen.config import load_config
from codegen.contracts.frd import FrdContract

REPO = Path(__file__).resolve().parents[1]
PAIR = FIXTURE_ROOT / "pair_4"
GOLDEN = PAIR / "golden" / "IIG_EXPECTED.yaml"
OVERLAY = PAIR / "config_overlay.yaml"
DESIGN = REPO / "docs" / "acfc" / "MULTI_TABLE_DESIGN.md"
GOLDENED = ["DATA_FACTORY_PIPELINE_SCHEDULE", "ADLS_DELTA_INGESTION_DETAILS",
            "STGDELTA_STDDELTA_INGESTION_DET", "DATA_QUALITY_RULES"]
NOT_SHEETS = {"TABLE_DEFINITIONS"}
ROW_COUNTS = {"DATA_FACTORY_PIPELINE_SCHEDULE": 4, "ADLS_DELTA_INGESTION_DETAILS": 6,
              "STGDELTA_STDDELTA_INGESTION_DET": 3, "DATA_QUALITY_RULES": 6}
CLASSES = {"per_file", "per_table", "per_feed", "environment", "convention", "engineer"}
AUDIT_IIG_TYPES = {"String": "String", "Timestamp": "TIMESTAMP"}   # iig_v2 audit_type_casing


@pytest.fixture(scope="module")
def golden() -> dict:
    return yaml.safe_load(GOLDEN.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def config():
    return load_config(REPO / "config" / "config.yaml", overlays=[OVERLAY])


def _rows(golden: dict, sheet: str) -> list[dict]:
    return golden[sheet]["rows"]


def _cells(row: dict) -> dict:
    return {k: v for k, v in row.items() if not k.startswith("_")}


def _joined(values) -> str:
    return ",".join(values)


def _audit_columns() -> list[str]:
    return [c for c, _t in pair4_nb.AUDIT]


def _audit_types() -> list[str]:
    return [AUDIT_IIG_TYPES[t] for _c, t in pair4_nb.AUDIT]


# ------------------------------------------------------------------ golden shape


def test_golden_sheets_columns_and_row_counts_follow_the_template(golden, config):
    template = config.metadata.templates["iig_v2"]
    assert [k for k in golden if not k.startswith("_") and k not in NOT_SHEETS] == GOLDENED
    for sheet in GOLDENED:
        headers = list(template.tabs[sheet].headers)
        assert list(golden[sheet]["columns"]) == headers, sheet
        assert set(golden[sheet]["columns"].values()) <= CLASSES, sheet
        rows = _rows(golden, sheet)
        assert len(rows) == ROW_COUNTS[sheet], sheet
        for index, row in enumerate(rows):
            assert set(_cells(row)) == set(headers), (sheet, index)
            for column, value in _cells(row).items():
                assert value is None or isinstance(value, str), (sheet, index, column, value)


def test_owner_map_in_the_design_doc_matches_the_golden(golden):
    text = DESIGN.read_text(encoding="utf-8")
    labels = {"per-file": "per_file", "per-table": "per_table", "per-feed": "per_feed",
              "environment": "environment", "convention": "convention",
              "engineer-assigned": "engineer"}
    for sheet in GOLDENED:
        section = text.split(f"### {sheet}\n", 1)[1].split("\n### ", 1)[0]
        found: dict[str, str] = {}
        for label, cls in labels.items():
            match = re.search(rf"^- \*\*{re.escape(label)}:\*\* (.+)$", section, re.MULTILINE)
            assert match, (sheet, label)
            if match.group(1).strip() == "—":
                continue
            for column in match.group(1).split(", "):
                assert column not in found, (sheet, column)
                found[column] = cls
        assert found == golden[sheet]["columns"], sheet


# ------------------------------------------------------------------ golden vs facts


def _file_pattern(lob: str) -> str:
    return pair4_nb.FILE_TEMPLATE.replace("<LOB>", lob).replace("CCYYMMDD", "*")


def test_adls_delta_is_one_row_per_lob_file_all_into_the_detail_table(golden):
    detail = pair4_nb.table_by_segment("Detail")
    stage_columns = [f.column for f in detail.fields]
    rows = _rows(golden, "ADLS_DELTA_INGESTION_DETAILS")
    assert [r["_object"] for r in rows] == list(range(1, len(pair4_nb.LOBS) + 1))
    assert [r["LOB"] for r in rows] == pair4_nb.LOBS
    for row, lob in zip(rows, pair4_nb.LOBS, strict=True):
        assert row["SRC_FILE_NAME"] == _file_pattern(lob)
        assert row["OBJECT_NAME"] == f"NWB_COB_RPT_{lob}"
        assert row["TGT_PARTITION_VALUE"] == lob
        assert row["TGT_PARTITION_COLUMN"] == "LOB" and "LOB" in stage_columns
        assert row["TGT_TABLE_NAME"] == detail.table
        assert row["TGT_DATABASE_NAME"] == pair4_nb.STAGE_SCHEMA
        assert row["SRC_COLUMNS"] == _joined(f"col{i}:{c}"
                                             for i, c in enumerate(stage_columns, start=1))
        assert row["SRC_DATA_TYPE"] == _joined(
            [f"{pair4_nb.SOURCE_TYPE}:{pair4_nb.STAGE_TYPE}"] * len(stage_columns))
        assert row["TGT_COLUMN_NAMES"] == _joined(stage_columns + _audit_columns())
        assert row["TGT_DATA_TYPE"] == _joined(
            [pair4_nb.STAGE_TYPE] * len(stage_columns) + _audit_types())
        assert row["MANDATORY_FIELD_LIST"] == _joined(f.column for f in detail.fields
                                                      if f.mandatory)
        assert row["TGT_PRIMARY_KEY"] == _joined(f.column for f in detail.fields
                                                 if f.primary_key)
        assert row["TGT_RJT_TABLE_NAME"] == f"{detail.table}_reject"
        assert row["SRC_FILE_DELIMITER"] == pair4_nb.DELIMITER
        assert row["FREQUENCY"] == pair4_nb.FREQUENCY
        assert (row["DOMAIN"], row["SUBDOMAIN"]) == (pair4_nb.DOMAIN, pair4_nb.SUB_DOMAIN)
        assert row["SOURCE"] == pair4_nb.VENDOR
        assert row["TGT_LOAD_OPTION"] == pair4_nb.LOAD_STRATEGY
        assert row["OBJECT_ID"] is None and row["GROUP_ID"] is None   # never invented
    header_columns = {f.column for f in pair4_nb.table_by_segment("Header").fields}
    trailer_columns = {f.column for f in pair4_nb.table_by_segment("Trailer").fields}
    only_hdr_trl = (header_columns | trailer_columns) - set(stage_columns)
    assert only_hdr_trl and not only_hdr_trl & set(rows[0]["TGT_COLUMN_NAMES"].split(","))


def test_stg_to_std_is_one_row_per_table_with_the_reference_sheet_on_the_target(golden):
    rows = _rows(golden, "STGDELTA_STDDELTA_INGESTION_DET")
    assert [r["_segment"] for r in rows] == [t.segment for t in pair4_nb.TABLES]
    for row, table in zip(rows, pair4_nb.TABLES, strict=True):
        carried = [f for f in table.fields if f.std_column is not None]
        assert len(carried) < len(table.fields)          # RECORD_TYPE stays in stage
        assert row["SRC_TABLE_NAME"] == row["TGT_TABLE_NAME"] == table.table
        assert row["SRC_SCHEMA_NAME"] == pair4_nb.STAGE_SCHEMA
        assert row["TGT_SCHEMA_NAME"] == pair4_nb.STANDARD_SCHEMA
        assert row["SRC_COLUMNS"] == _joined(
            [f"{f.column}:{f.std_column}" for f in carried]
            + [f"{c}:{c}" for c in _audit_columns()])
        assert row["SRC_DATA_TYPE"] == _joined(
            [f"{pair4_nb.STAGE_TYPE}:{f.std_type}" for f in carried]
            + [f"{t}:{t}" for t in _audit_types()])
        assert row["TGT_COLUMN_NAMES"] == _joined([f.std_column for f in carried]
                                                  + _audit_columns())
        assert row["TGT_DATA_TYPE"] == _joined([f.std_type for f in carried] + _audit_types())
        keys = [f.std_column for f in carried if f.primary_key]
        assert row["TGT_PRIMARY_KEY"] == (_joined(keys) if keys else "NA")
        assert row["TGT_RJT_TABLE_NAME"] == f"{table.table}_reject"
        assert row["TGT_LOAD_OPTION"] == pair4_nb.LOAD_STRATEGY
        assert row["LOB"] == _joined(pair4_nb.LOBS)


def test_dq_split_rule_is_one_row_per_adls_object(golden, config):
    adls = {r["_object"]: r for r in _rows(golden, "ADLS_DELTA_INGESTION_DETAILS")}
    header = pair4_nb.table_by_segment("Header")
    trailer = pair4_nb.table_by_segment("Trailer")
    catalog = config.conventions.catalog_map[pair4_nb.STAGE_CATALOG]
    for row in _rows(golden, "DATA_QUALITY_RULES"):
        assert row["RULE_CLASS"] == "LoadHeaderAndTrailerToSeparateTablesRule"
        assert row["INPUT_PARAM"] == adls[row["_object"]]["SRC_FILE_NAME"]
        assert row["SOURCE_COLUMN"] == _joined(f.column for f in header.fields)
        assert row["TARGET_COLUMN"] == _joined(
            [catalog, pair4_nb.STAGE_SCHEMA, header.table, trailer.table])
        assert row["SEQUENCE_NO"] == "1"
    assert sorted(r["_object"] for r in _rows(golden, "DATA_QUALITY_RULES")) == sorted(adls)


def test_pipeline_schedule_is_the_four_movement_rows(golden):
    rows = _rows(golden, "DATA_FACTORY_PIPELINE_SCHEDULE")
    roles = [r["_role"] for r in rows]
    assert roles == ["grand_master", "master", "file_to_stage", "stage_to_standard"]
    assert [r["_parent"] for r in rows] == [None, "grand_master", "master", "master"]
    assert rows[0]["PARENT_PIPELINE_ID"] == "0"
    assert all(r["PIPELINE_ID"] is None and r["PIPELINE_NAME"] is None for r in rows)
    assert {r["PIPELINE_FREQUENCY"] for r in rows} == {pair4_nb.FREQUENCY}


def test_environment_cells_are_the_overlay_values(golden, config):
    template = config.metadata.templates["iig_v2"]
    for sheet in ("ADLS_DELTA_INGESTION_DETAILS", "STGDELTA_STDDELTA_INGESTION_DET"):
        constants = template.constants[sheet]
        env_columns = [c for c, cls in golden[sheet]["columns"].items() if cls == "environment"]
        for row in _rows(golden, sheet):
            for column in env_columns:
                assert row[column] == constants[column], (sheet, column)
                assert column not in template.always_blank, (sheet, column)


def test_catalog_cells_are_each_bands_logical_catalog_mapped(golden, config):
    catalog_map = config.conventions.catalog_map
    fallback = config.conventions.profiles["acfc_prx"].default_catalog
    stage = catalog_map[pair4_nb.STAGE_CATALOG]
    standard = catalog_map[pair4_nb.STANDARD_CATALOG]
    assert stage != fallback["stage"] and standard != fallback["standard"]
    for row in _rows(golden, "STGDELTA_STDDELTA_INGESTION_DET"):
        assert (row["SRC_CATALOG_NAME"], row["TGT_CATALOG_NAME"]) == (stage, standard)
        assert row["TGT_SCHEMA_NAME"] == pair4_nb.STANDARD_SCHEMA
    definitions = golden["TABLE_DEFINITIONS"]
    assert len(definitions) == 2 * len(pair4_nb.TABLES)
    for entry in definitions:
        table = pair4_nb.table_by_segment(entry["segment"])
        if entry["layer"] == "stage":
            logical = (pair4_nb.STAGE_CATALOG, pair4_nb.STAGE_SCHEMA)
            columns = len(table.fields)
        else:
            logical = (pair4_nb.STANDARD_CATALOG, pair4_nb.STANDARD_SCHEMA)
            columns = len([f for f in table.fields if f.std_column is not None])
        assert entry["logical"] == ".".join([*logical, table.table])
        assert entry["mapped"] == ".".join([catalog_map[logical[0]], logical[1], table.table])
        assert entry["columns"] == columns
    assert pair4_nb.STAGE_SCHEMA.removeprefix("stg_") == pair4_nb.STANDARD_SCHEMA


def test_path_cells_follow_the_overlay_shapes(golden, config):
    template = config.metadata.templates["iig_v2"]
    container = template.constants["ADLS_DELTA_INGESTION_DETAILS"]["SRC_CONTAINER_NAME"]
    assert pair4_nb.FILE_LOCATION.startswith(f"{container}/")
    landing_rel = "/" + pair4_nb.FILE_LOCATION[len(container) + 1:]
    domain_path = landing_rel[len("/inbound/"):]
    for sheet in ("ADLS_DELTA_INGESTION_DETAILS", "STGDELTA_STDDELTA_INGESTION_DET"):
        for row in _rows(golden, sheet):
            for column, shape in template.path_patterns[sheet].items():
                expected = shape.format(landing_rel=landing_rel, domain_path=domain_path,
                                        stage_table=row["TGT_TABLE_NAME"],
                                        reject_table=row["TGT_RJT_TABLE_NAME"])
                assert row[column] == expected, (sheet, column)


# ------------------------------------------------------------------ inputs


def _main_data(ws) -> list[dict]:
    """Data rows as {source header: value, 'stage': {...}, 'standard': {...}}."""
    width = len(pair4_nb.SOURCE_HEADERS)
    band = len(pair4_nb.BAND_HEADERS)
    out = []
    for r in ws.iter_rows(min_row=pair4_nb.TABLE_HEADER_ROW + 1, values_only=True):
        row = dict(zip(pair4_nb.SOURCE_HEADERS, r[:width], strict=True))
        row["stage"] = dict(zip(pair4_nb.BAND_HEADERS, r[width:width + band], strict=True))
        row["standard"] = dict(zip(pair4_nb.BAND_HEADERS, r[width + band:], strict=True))
        out.append(row)
    return out


def test_sttm_main_sheet_carries_a_stage_and_a_standard_band_per_row():
    wb = load_workbook(FIXTURE_ROOT / pair4_nb.STTM_PATH)
    assert wb.sheetnames == [pair4_nb.MAIN_SHEET, pair4_nb.REFERENCE_SHEET]
    ws = wb[pair4_nb.MAIN_SHEET]
    block = [(ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value)
             for r in range(1, len(pair4_nb.HEADER_BLOCK) + 1)]
    assert block == pair4_nb.HEADER_BLOCK
    assert ws.cell(row=4, column=2).value.split(", ") == pair4_nb.LOBS
    band = [c.value for c in ws[pair4_nb.BAND_ROW]]
    assert [v for v in band if v] == ["Source", "Stage Layer", "Standard Layer"]
    assert sorted(str(m) for m in ws.merged_cells.ranges) == ["A10:E10", "F10:L10", "M10:S10"]
    assert [c.value for c in ws[pair4_nb.TABLE_HEADER_ROW]] == pair4_nb.MAIN_HEADERS
    data = _main_data(ws)
    triples: dict[str, set[tuple]] = {"stage": set(), "standard": set()}
    for r in data:
        for layer, found in triples.items():
            b = r[layer]
            if b["TableName"]:
                found.add((b["Catalog"], b["Schema"], b["TableName"]))
    tables = [t.table for t in pair4_nb.TABLES]
    assert triples["stage"] == {(pair4_nb.STAGE_CATALOG, pair4_nb.STAGE_SCHEMA, t) for t in tables}
    assert triples["standard"] == {(pair4_nb.STANDARD_CATALOG, pair4_nb.STANDARD_SCHEMA, t)
                                   for t in tables}
    for table in pair4_nb.TABLES:
        mine = [r for r in data if r["stage"]["TableName"] == table.table]
        assert {r["Segment"] for r in mine} == {table.segment}
        assert len(mine) == len(table.fields) + len(pair4_nb.AUDIT)
        audit = [r for r in mine if r["Field Name"] == "NA"]
        assert [r["stage"]["ColumnName"] for r in audit] == _audit_columns()
        assert [r["standard"]["ColumnName"] for r in audit] == _audit_columns()
        dropped = [r["stage"]["ColumnName"] for r in mine if not r["standard"]["TableName"]]
        assert dropped == ["RECORD_TYPE"]
    assert {len(t.fields) for t in pair4_nb.TABLES} == {8, 20, 6}


def test_reference_sheet_is_a_secondary_source_agreeing_with_the_standard_band():
    wb = load_workbook(FIXTURE_ROOT / pair4_nb.STTM_PATH)
    standard_band = {(r["stage"]["TableName"], r["stage"]["ColumnName"]):
                     (r["standard"]["Schema"], r["standard"]["TableName"],
                      r["standard"]["ColumnName"], r["standard"]["DataType"])
                     for r in _main_data(wb[pair4_nb.MAIN_SHEET])}
    ws = wb[pair4_nb.REFERENCE_SHEET]
    reference = {(r[1], r[2]): tuple(r[4:8]) for r in ws.iter_rows(min_row=3, values_only=True)}
    assert reference == standard_band
    band = [c.value for c in ws[1]]
    assert band[:5] == ["Client - Source", "STG - Dest 1", None, None, "STD - Dest2"]
    assert sorted(str(m) for m in ws.merged_cells.ranges) == ["B1:D1", "E1:H1"]
    assert [c.value for c in ws[2]] == pair4_nb.REFERENCE_HEADERS
    carried = {(r[1], r[2]): r[6] for r in ws.iter_rows(min_row=3, values_only=True)}
    assert carried[("nb_cob_report_dtl", "MEMBER_DOB")] == "MEMBER_BIRTH_DATE"
    assert carried[("nb_cob_report_hdr", "RECORD_TYPE")] is None


def test_frd_contract_validates_and_matches_the_sttm():
    data = json.loads((FIXTURE_ROOT / pair4_nb.FRD_PATH).read_text(encoding="utf-8"))
    contract = FrdContract.model_validate(data)
    (feed,) = contract.feeds
    assert feed.feed_name == pair4_nb.FEED
    assert feed.lobs == pair4_nb.LOBS
    assert feed.file_name_patterns == [pair4_nb.FILE_TEMPLATE]
    assert feed.record_segments == ["Header", "Detail", "Trailer"]
    assert feed.stage_target.tables == [t.table for t in pair4_nb.TABLES]
    assert feed.standard_target.tables == [t.table for t in pair4_nb.TABLES]
    assert feed.standard_target.schema_name == pair4_nb.STANDARD_SCHEMA
    assert feed.stage_target.catalog is None and feed.standard_target.catalog is None
    assert feed.landing_location == pair4_nb.FILE_LOCATION
    assert feed.delimiter == pair4_nb.DELIMITER
