"""post57 (D1 post55 pair_01: GATE_FAIL, ``derivations [path_whitespace 19]``;
the PRX golden's own paths carry double spaces): whitespace in a path the
documents state is no longer a failed check. Under the conventions profile's
``path_whitespace: keep`` (the default) every path is written exactly as the
source states it, byte for byte, and each path cell gets ONE flag naming its
coordinates and the kind of whitespace (double space / leading / trailing /
tab) for the BSA to confirm; ``path_whitespace: fail`` keeps the old check
(tests/test_gate_fail_breakdown.py runs under it)."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from codegen import cli
from codegen.config import load_config
from codegen.extract.extractor import contract_to_json, extract_contract
from codegen.gate.derivations import (
    check_iig_derivations,
    normalise_path,
    path_violations,
    path_whitespace_flags,
    whitespace_kinds,
)
from post51_pairs import family_e_frd, family_e_workbook

REPO = Path(__file__).resolve().parents[1]
CONFIG = REPO / "config" / "config.yaml"
FEED = "feed_5_claims"
LANDING = "inbound/vendor  e/claims daily"


def test_the_kinds_of_whitespace():
    assert whitespace_kinds("/inbound/vendor  e/x") == ["double space"]
    assert whitespace_kinds("/inbound/claims daily/x") == []      # one inner space: a name
    assert whitespace_kinds(" /a/b") == ["leading"]
    assert whitespace_kinds("/a/b ") == ["trailing"]
    assert whitespace_kinds("/a /b") == ["trailing"]
    assert whitespace_kinds("/a/ b") == ["leading"]
    assert whitespace_kinds("/a\tb") == ["tab"]


def test_keep_never_fails_fail_does():
    assert path_violations("/a  b/c", 1024, "keep") == []
    assert path_violations("/a  b/c", 1024, "fail") == ["whitespace"]
    assert path_violations("/a  b//c", 1024, "keep") == ["doubled separator", "empty segment"]


def test_derived_paths_keep_their_whitespace_byte_for_byte():
    assert normalise_path("/inbound/vendor  e/x") == ("/inbound/vendor  e/x", [])
    value, changes = normalise_path("/inbound//vendor  e /x.")
    assert value == "/inbound/vendor  e /x"                        # separators / punctuation only
    assert changes == ["doubled separator collapsed", "trailing punctuation stripped from 'x.'"]


def _payload(value: str) -> dict:
    return {"tabs": {"ADLS_DELTA_INGESTION_DETAILS": {
        "headers": ["OBJECT_NAME", "SRC_ADLS_PATH"],
        "rows": [{"values": {"OBJECT_NAME": "FEED_5", "SRC_ADLS_PATH": value}}]}}}


def test_a_double_space_passes_through_unchanged_with_exactly_one_flag():
    config = load_config(CONFIG)
    path = "/inbound/vendor  e/claims/"
    payload = _payload(path)
    flags = path_whitespace_flags(payload, config.gate.derivations)
    assert flags == ["iig_path_whitespace:ADLS_DELTA_INGESTION_DETAILS!B2 (SRC_ADLS_PATH) — "
                     "double space; kept as the source document states it (BSA to confirm)"]
    assert payload["tabs"]["ADLS_DELTA_INGESTION_DETAILS"]["rows"][0]["values"][
        "SRC_ADLS_PATH"] == path                                    # byte for byte
    assert check_iig_derivations(payload, config, "keep").passed
    failed = check_iig_derivations(payload, config, "fail")
    assert not failed.passed and failed.violations == {"path_whitespace": 1}


def test_the_profile_switch_defaults_to_keep():
    config = load_config(CONFIG)
    assert config.conventions.get("acfc_prx").path_whitespace == "keep"
    assert config.conventions.get("edo_sfmc").path_whitespace == "keep"


@pytest.fixture()
def run(monkeypatch, tmp_path, capsys):
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_PROVIDER", "1")
    monkeypatch.setenv("CODEGEN_NOTIFICATION_EMAILS", "syn.dl.prodsupport@synthetic.example")
    (tmp_path / "out").mkdir()
    monkeypatch.setenv("CODEGEN_STORAGE_OUTPUTS", f"local:{(tmp_path / 'out').as_posix()}")
    monkeypatch.chdir(REPO)
    frd = family_e_frd(tmp_path / "frd.json")
    data = json.loads(frd.read_text(encoding="utf-8"))
    data["feeds"][0]["landing_location"] = LANDING
    frd.write_text(json.dumps(data), encoding="utf-8")
    sttm = tmp_path / "sttm.json"
    sttm.write_text(contract_to_json(extract_contract(
        family_e_workbook(tmp_path / "e.xlsx", [None]), frd, load_config(CONFIG),
        generated_date="2026-10-15")), encoding="utf-8")
    capsys.readouterr()
    code = cli.main(["generate", "--frd-contract", str(frd), "--sttm-contract", str(sttm),
                     "--skip-tests", "--dry-run", "--output-mode", "framework",
                     "--profile", "acfc_prx", "--iig-template", "iig_v2"])
    return code, capsys.readouterr().out, tmp_path / "out"


def test_the_pair01_shape_passes_with_one_flag_per_path_cell(run):
    code, out, out_root = run
    assert code == cli.EXIT_OK, out
    assert "derivations=FAIL" not in out and "GATE_FAIL" not in out
    workbook = next(out_root.rglob("*_IIG.xlsx"))
    book = load_workbook(workbook, read_only=True)
    try:
        cells = [(ws.title, cell) for ws in book.worksheets
                 for row in ws.iter_rows(values_only=True) for cell in row
                 if isinstance(cell, str) and "vendor" in cell and "/" in cell]
    finally:
        book.close()
    assert cells and all("vendor  e" in c for _sheet, c in cells)    # as the FRD states it
    report = next((out_root / "reports").glob(f"{FEED}.md")).read_text(encoding="utf-8")
    flagged = [ln for ln in report.splitlines() if ln.startswith("- iig_path_whitespace:")]
    addresses = [ln.split(" ", 2)[1] for ln in flagged]
    assert flagged and len(addresses) == len(set(addresses))           # one flag per cell
    assert all("— double space;" in ln for ln in flagged)
