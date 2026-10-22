"""post94: two STTM extraction gaps the BSA pass answered by hand (post75 report).

- pair_08: the audit rows sit at the end of the mapping with a description cell in
  the source band ("Source File Name"), so they read as neither audit nor mapping
  rows and the feed asked for its audit columns. A row with NO source field name
  whose stage column is a configured audit name (extractor.audit_column_names) is
  an audit row.
- pairs 08 / 10: a row whose stage type is empty while its standard band types it.
  The blocking question stays and proposes the stage type the SAME sheet's other
  rows unanimously pair with that standard type (pair_08: INT in standard is STRING
  in stage 39 of 39), with the count; never written unasked, never a copy of the
  standard type.

Synthetic Family E workbooks only (tests/post73_pairs.py).
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from codegen.extract import extract_contract
from codegen.extract.extractor import NeedsAnswersError
from post51_pairs import AUDIT, FIELDS
from post73_pairs import FIRST_DATA_ROW, frd, mapping_workbook

DATE = "2026-01-01"
COMMENTS_COL = 11            # the source band's last column ("Comments")


def _with_audit_descriptions(path):
    """The workbook's audit rows carry a description in the source band (pair_08)."""
    book = load_workbook(path)
    ws = book["FEED_5_MAPPING"]
    for k, (column, _t) in enumerate(AUDIT):
        ws.cell(row=FIRST_DATA_ROW + len(FIELDS) + k, column=COMMENTS_COL,
                value=f"{column.title()} description")
    book.save(path)
    return path


def test_audit_rows_with_a_description_in_the_source_band_are_audit_rows(config, tmp_path):
    workbook = _with_audit_descriptions(mapping_workbook(tmp_path / "p8.xlsx"))
    contract = extract_contract(workbook, frd(tmp_path / "frd.json"), config,
                                generated_date=DATE)
    (feed,) = contract.feeds
    assert [a.column for a in feed.audit_columns] == [c for c, _t in AUDIT]
    assert not [q for q in contract.needs_answers if q.key.endswith("audit_columns")]
    assert feed.field_count == len(FIELDS)


def test_without_the_name_list_such_a_row_is_not_an_audit_row(config, tmp_path):
    workbook = _with_audit_descriptions(mapping_workbook(tmp_path / "p8.xlsx"))
    plain = config.model_copy(update={"extractor": config.extractor.model_copy(
        update={"audit_column_names": []})})
    # the pre-post94 behaviour, reproduced: the feed asks for its audit columns
    with pytest.raises(NeedsAnswersError, match=r"feeds\[0\]\.audit_columns"):
        extract_contract(workbook, frd(tmp_path / "frd.json"), plain, generated_date=DATE)


def _int_sheet(tmp_path, overrides):
    """Fields 0, 1 and 3 typed Int in standard, String in stage; ``overrides`` on top."""
    cells = {i: {"standard_type": "Int"} for i in (0, 1, 3)}
    for i, extra in overrides.items():
        cells.setdefault(i, {}).update(extra)
    return mapping_workbook(tmp_path / "p8.xlsx", cells)


def test_a_blank_stage_type_proposes_what_the_sheet_pairs_with_its_standard_type(
        config, tmp_path):
    workbook = _int_sheet(tmp_path, {2: {"stage_type": None, "standard_type": "Int"}})
    contract = extract_contract(workbook, frd(tmp_path / "frd.json"), config,
                                generated_date=DATE)
    (question,) = contract.needs_answers
    assert question.key == "feeds[0].columns[2].target_type"
    assert question.candidates == ["String"]
    assert ("proposed 'String': this sheet's other rows typed INT in the standard band "
            "are 'String' in the stage band (3 of 3)") in question.reason
    # proposed, not written: the column is still out of the contract
    assert "SERVICE_DATE" not in [f.stage_column for f in contract.feeds[0].fields]


def test_no_proposal_when_the_sheet_disagrees_or_has_too_few_rows(config, tmp_path):
    split = _int_sheet(tmp_path, {2: {"stage_type": None, "standard_type": "Int"},
                                  1: {"stage_type": "Int"}})
    (question,) = extract_contract(split, frd(tmp_path / "frd.json"), config,
                                   generated_date=DATE).needs_answers
    assert question.candidates == [] and "proposed" not in question.reason
    few = mapping_workbook(tmp_path / "few.xlsx", {
        0: {"standard_type": "Int"}, 2: {"stage_type": None, "standard_type": "Int"}})
    (question,) = extract_contract(few, frd(tmp_path / "frd2.json"), config,
                                   generated_date=DATE).needs_answers
    assert question.candidates == []                   # 1 row < stage_type_proposal_min_rows
