"""post70: TGT_PRIMARY_KEY is proposed only from an explicit key signal.

The not-null list is never a key: the STTM contract carries the two apart
(``load_rules.not_null_columns`` vs ``load_rules.natural_key``, the latter only
from Primary Key cells). A proposal (STTM Primary Key cells, an FRD key
statement, the FAQ's natural_key_columns, a corpus row) is shown on an
``INPUT_NEEDED feeds.<feed>.primary_key`` line and written only when chosen —
then flagged ``key_from_contract``. No signal on a merge load = ONE blocking
QUESTION whose candidates are the not-null columns, none selected; the v2
notebook stops in its config cell on a merge row whose key is open.
"""

from __future__ import annotations

import json
import shutil
from io import BytesIO
from pathlib import Path

import openpyxl

import notebook_v2_fixtures as fx
from codegen import cli, primary_key, run_stages
from codegen.config import load_config
from codegen.resolve import resolve_pair

PROFILE = {"conventions_profile": "acfc_prx", "iig_template": "iig_v2"}


def _run(config, pairs, answers: Path | None = None, parts=("notebook", "framework")) -> int:
    return cli._run_pairs(list(pairs), config, dry_run=True, skip_tests=True,
                          output_mode=list(parts), only_feed=None, answers_path=answers,
                          **PROFILE)


def _iig_rows(out: Path, tab: str) -> list[dict]:
    (book,) = [p for p in out.rglob("*_IIG.xlsx")
               if "quarantine" not in p.parts and "combined" not in p.parts]
    # from its bytes: openpyxl keeps a read-only workbook's file open (Windows)
    rows = list(openpyxl.load_workbook(BytesIO(book.read_bytes()), read_only=True)[tab]
                .iter_rows(values_only=True))
    return [dict(zip(rows[0], r, strict=False)) for r in rows[1:]]


def _keyed_pair1(tmp_path: Path, config, keyed: tuple[str, ...]) -> tuple[Path, Path]:
    """Pair 1 with the STTM's Primary Key cell (column M) set to Y on the
    ``keyed`` DET rows — the STTM marks a key."""
    workbook = tmp_path / "pair_1_keyed.xlsx"
    shutil.copy(fx.ACFC_SHAPES / "sttm" / "pair_1_family_a.xlsx", workbook)
    book = openpyxl.load_workbook(workbook)
    sheet = book["FEED_1_MAPPING"]
    for row in range(16, sheet.max_row + 1):
        if sheet.cell(row, 2).value == "DET" and sheet.cell(row, 3).value in keyed:
            sheet.cell(row, 13).value = "Y"
    book.save(workbook)
    frd, _sttm, _vdd = fx.prx_pair(tmp_path, config)
    from codegen.extract import contract_to_json, extract_contract
    from codegen.layout.model import MockLayoutProvider
    from codegen.layout.resolve import resolve_pair as resolve_layout_pair

    profiles = fx.REPO / "fixtures" / "layout_profiles"
    layout = resolve_layout_pair(
        workbook, fx.ACFC_SHAPES / "frd" / "f1_pair_1.docx", config,
        provider=MockLayoutProvider([profiles / "mock", profiles]),
        cache_dirs=[tmp_path / "keyed_cache"], generated_date="2026-01-01")
    sttm = tmp_path / "keyed.sttm.contract.json"
    sttm.write_text(contract_to_json(extract_contract(
        workbook, frd, config, generated_date="2026-01-01", layout=layout.sttm.profile)),
        encoding="utf-8")
    return frd, sttm


# ------------------------------------------------- the contract: two lists


def test_required_columns_without_a_key_marker_propose_nothing():
    """The CV golden STTM marks required (NOT NULL) columns and no key."""
    config = load_config(fx.REPO / "config" / "config.yaml", overlays=[fx.ACFC_OVERLAY])
    frd, sttm = fx.sd_pair()
    contract = json.loads(sttm.read_text(encoding="utf-8"))
    for feed in contract["feeds"]:
        assert feed["load_rules"]["not_null_columns"]
        assert "natural_key" not in feed["load_rules"]
    for spec in resolve_pair(frd, sttm, config):
        assert spec.not_null_columns and spec.natural_key_columns == []
        offer = primary_key.feed_proposal(spec)
        assert not offer and offer.columns == ()


def test_a_primary_key_column_proposes_exactly_those_columns(tmp_path, monkeypatch):
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    config = fx.acfc_config(tmp_path, fx.PAIR1_OVERLAY)
    frd, sttm = _keyed_pair1(tmp_path, config, ("CARDHOLDER_ID", "PLAN_ID"))
    (feed,) = json.loads(sttm.read_text(encoding="utf-8"))["feeds"]
    assert feed["load_rules"]["natural_key"] == ["CARDHOLDER_ID", "PLAN_ID"]
    # the not-null list is its own (SEGMENT_IDENTIFIER is required, not a key)
    assert "SEGMENT_IDENTIFIER" in feed["load_rules"]["not_null_columns"]
    (spec,) = resolve_pair(frd, sttm, config)
    offer = primary_key.feed_proposal(spec)
    detail = {f.source_column: f.stage_column for f in spec.detail_segment.fields}
    assert offer.columns == (detail["CARDHOLDER_ID"], detail["PLAN_ID"])
    assert offer.kind == "sttm" and "Primary Key cells" in offer.source
    assert list(offer.columns) == spec.natural_key_columns


def test_the_frd_key_statement_is_a_signal_the_not_null_list_is_not(tmp_path, monkeypatch):
    """Pair 1 unkeyed: the STTM's not-null list (SEGMENT_IDENTIFIER,
    CARDHOLDER_ID, PLAN_ID) is no proposal; the FRD's Business Key row is."""
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    config = fx.acfc_config(tmp_path, fx.PAIR1_OVERLAY)
    frd, sttm, _vdd = fx.prx_pair(tmp_path, config)
    (feed,) = json.loads(frd.read_text(encoding="utf-8"))["feeds"]
    assert feed["key_statement"]
    (spec,) = resolve_pair(frd, sttm, config)
    assert spec.natural_key_columns == [] and spec.not_null_columns
    offer = primary_key.feed_proposal(spec)
    assert offer.kind == "frd"
    assert "SEGMENT_IDENTIFIER" not in offer.value
    assert set(offer.columns) != set(spec.not_null_columns)


def test_no_signal_never_falls_back_to_the_faq_dedup_or_not_null(tmp_path):
    config = fx.acfc_config(tmp_path)
    frd, sttm = fx.sfmc_pair()
    (spec,) = resolve_pair(frd, sttm, config)
    assert not primary_key.feed_proposal(spec)
    reason, candidates = primary_key.question(spec, None)
    assert candidates == list(spec.not_null_columns)
    assert "none selected" in reason


# --------------------------------------------- the run: proposal, question


def _no_autoset(tmp_path: Path) -> Path:
    path = tmp_path / "no_pk_autoset.yaml"
    path.write_text("conventions:\n  profiles:\n    acfc_prx:\n"
                    "      primary_key_autoset: false\n", encoding="utf-8")
    return path


def test_autoset_writes_the_signals_key_unasked(tmp_path, monkeypatch, capsys):
    """post81: acfc_env.yaml sets primary_key_autoset — the STTM's Primary Key
    cells are written as a profile default (a NOTE, not an open input)."""
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    config = fx.acfc_config(tmp_path, fx.PAIR1_OVERLAY)
    assert config.conventions.get("acfc_prx").primary_key_autoset is True
    frd, sttm = _keyed_pair1(tmp_path, config, ("CARDHOLDER_ID", "PLAN_ID"))
    (spec,) = resolve_pair(frd, sttm, config)
    feed = spec.feed_id
    _run(config, [(frd, sttm)], parts=("framework",))
    out = capsys.readouterr().out
    lines = [ln for ln in out.splitlines() if f"feeds.{feed}.primary_key" in ln]
    assert lines and all(ln.startswith("NOTE") for ln in lines)
    detail = {f.source_column: f.stage_column for f in spec.detail_segment.fields}
    wanted = f"{detail['CARDHOLDER_ID']},{detail['PLAN_ID']}"
    assert {r["TGT_PRIMARY_KEY"] for r in _iig_rows(
        tmp_path / "out", "ADLS_DELTA_INGESTION_DETAILS")} == {wanted}
    (record,) = [f["inputs"]["primary_key"] for f in
                 run_stages.read_stage_manifest(tmp_path / "out")["feeds"]]
    assert record["source"] == "profile"


def test_a_proposal_is_shown_and_written_only_when_chosen(tmp_path, monkeypatch, capsys):
    """Without primary_key_autoset (post70): proposed, written only when chosen."""
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    config = fx.acfc_config(tmp_path, fx.PAIR1_OVERLAY, _no_autoset(tmp_path))
    frd, sttm = _keyed_pair1(tmp_path, config, ("CARDHOLDER_ID", "PLAN_ID"))
    (spec,) = resolve_pair(frd, sttm, config)
    feed = spec.feed_id
    assert _run(config, [(frd, sttm)], parts=("framework",)) in (cli.EXIT_OK, 1)
    out = capsys.readouterr().out
    line = next(ln for ln in out.splitlines() if f"feeds.{feed}.primary_key" in ln)
    assert line.startswith("INPUT_NEEDED") and "proposed" in line
    assert all(not r["TGT_PRIMARY_KEY"]
               for r in _iig_rows(tmp_path / "out", "ADLS_DELTA_INGESTION_DETAILS"))

    answers = tmp_path / "answers.yaml"
    answers.write_text(f"inputs:\n  feeds:\n    {feed}:\n      primary_key:\n"
                       "        choice: use_default\n", encoding="utf-8")
    shutil.rmtree(tmp_path / "out")
    _run(config, [(frd, sttm)], answers, parts=("framework",))
    out = capsys.readouterr().out
    assert f"feeds.{feed}.primary_key" not in "".join(
        ln for ln in out.splitlines() if ln.startswith(("INPUT_NEEDED", "QUESTION")))
    detail = {f.source_column: f.stage_column for f in spec.detail_segment.fields}
    wanted = f"{detail['CARDHOLDER_ID']},{detail['PLAN_ID']}"
    assert {r["TGT_PRIMARY_KEY"] for r in _iig_rows(
        tmp_path / "out", "ADLS_DELTA_INGESTION_DETAILS")} == {wanted}
    report = "\n".join(p.read_text(encoding="utf-8")
                       for p in (tmp_path / "reports").rglob("*.md"))
    assert primary_key.FLAG in report


def test_a_merge_load_without_a_signal_is_one_blocking_question(tmp_path, capsys):
    config = fx.acfc_config(tmp_path)
    assert _run(config, [fx.sfmc_pair()]) == cli.EXIT_NEEDS_ANSWERS
    out = capsys.readouterr().out
    asked = [ln for ln in out.splitlines() if ".primary_key" in ln
             and ln.startswith(("QUESTION", "INPUT_NEEDED"))]
    assert len(asked) == 1 and asked[0].startswith("QUESTION")
    assert "none selected" in asked[0]
    template = (tmp_path / "out" / "answers_template.yaml").read_text(encoding="utf-8")
    assert "primary_key" in template
    # the blocking line withholds the notebook (post68); the key cell is open
    assert not list((tmp_path / "out").rglob("notebook/*.ipynb"))
    manifest = run_stages.read_stage_manifest(tmp_path / "out")
    assert manifest["stages"]["generate"] == {"status": "NEEDS_ANSWERS"}


def test_a_typed_key_answers_the_question(tmp_path):
    config = fx.acfc_config(tmp_path)
    frd, sttm = fx.sfmc_pair()
    (spec,) = resolve_pair(frd, sttm, config)
    answers = tmp_path / "answers.yaml"
    answers.write_text(f"inputs:\n  feeds:\n    {spec.feed_id}:\n"
                       "      primary_key: \"MEMBER_ID, EMAIL_DTM\"\n", encoding="utf-8")
    assert _run(config, [(frd, sttm)], answers) == cli.EXIT_OK
    for tab in ("ADLS_DELTA_INGESTION_DETAILS", "STGDELTA_STDDELTA_INGESTION_DET"):
        assert {r["TGT_PRIMARY_KEY"] for r in _iig_rows(tmp_path / "out", tab)} == {
            "MEMBER_ID,EMAIL_DTM"}
    assert list((tmp_path / "out").rglob("notebook/*.ipynb"))


def test_the_v2_notebook_stops_on_an_open_merge_key():
    from codegen.emit import notebook_v2

    opened = notebook_v2._Open()
    row = {"values": {"TGT_LOAD_OPTION": "Upsert", "TGT_PRIMARY_KEY": "NA"}}
    payload = {"tabs": {notebook_v2.STD: {"rows": [row]}}}
    (standard,) = notebook_v2._standards(payload, opened)
    assert standard["primary_key"] == []
    assert any(c.startswith(f"{notebook_v2.STD}.TGT_PRIMARY_KEY") for c in opened.cells)
    opened = notebook_v2._Open()
    row["values"]["TGT_LOAD_OPTION"] = "Append"
    notebook_v2._standards(payload, opened)
    assert not any("TGT_PRIMARY_KEY" in c for c in opened.cells)


# ------------------------------- post108: a key only the Standard band marks


def _standard_keyed_pair1(tmp_path: Path, config) -> tuple[Path, Path, tuple[str, ...]]:
    """Pair 1 with an Upsert standard leg whose key ONLY the Standard band's
    Primary Key cells mark (CARDHOLDER_ID / PLAN_ID: primary_key False,
    standard_primary_key True — what generic.py writes when only that band says
    Y), and no FRD key statement. Returns the keyed standard columns too."""
    frd, sttm, _vdd = fx.prx_pair(tmp_path, config)
    frd_json = json.loads(Path(frd).read_text(encoding="utf-8"))

    def drop_key_statements(node):
        if isinstance(node, dict):
            for key in node:
                if key == "key_statement":
                    node[key] = None
                else:
                    drop_key_statements(node[key])
        elif isinstance(node, list):
            for item in node:
                drop_key_statements(item)
    drop_key_statements(frd_json)
    frd_json["feeds"][0]["standard_target"]["load_strategy"] = "Upsert"
    frd = tmp_path / "std_keyed.frd.contract.json"
    frd.write_text(json.dumps(frd_json), encoding="utf-8")
    sttm_json = json.loads(Path(sttm).read_text(encoding="utf-8"))
    keyed = []
    for field in sttm_json["feeds"][0]["fields"]:
        if (field["source_column"] in ("CARDHOLDER_ID", "PLAN_ID")
                and field.get("record_segment") in (None, "Detail")):
            field["primary_key"] = False
            field["standard_primary_key"] = True
            keyed.append(field["standard_column"])
    sttm = tmp_path / "std_keyed.sttm.contract.json"
    sttm.write_text(json.dumps(sttm_json), encoding="utf-8")
    return frd, sttm, tuple(keyed)


def test_a_standard_only_key_is_written_by_autoset_on_the_standard_row(tmp_path,
                                                                        monkeypatch, capsys):
    # post108 (codeflow audit #11): the feed's key input was built from the STAGE
    # table's signal only, so a key only the Standard band marks had no input,
    # no question, and the Upsert leg shipped with a blank key, exit 0
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    config = fx.acfc_config(tmp_path, fx.PAIR1_OVERLAY)
    assert config.conventions.get("acfc_prx").primary_key_autoset is True
    frd, sttm, keyed = _standard_keyed_pair1(tmp_path, config)
    (spec,) = resolve_pair(frd, sttm, config)
    assert len(keyed) == 2 and primary_key.merge_legs(spec) == ["standard"]
    assert primary_key.feed_proposal(spec).kind == "sttm"
    assert _run(config, [(frd, sttm)], parts=("framework",)) == cli.EXIT_OK
    out = capsys.readouterr().out
    lines = [ln for ln in out.splitlines() if f"feeds.{spec.feed_id}.primary_key" in ln]
    assert lines and all(ln.startswith("NOTE") for ln in lines)
    (std_row,) = _iig_rows(tmp_path / "out", "STGDELTA_STDDELTA_INGESTION_DET")
    assert std_row["TGT_PRIMARY_KEY"] == ",".join(keyed)
    # the stage (ADLS) rows: no stage signal, an append leg — no key written there
    assert all(not r["TGT_PRIMARY_KEY"]
               for r in _iig_rows(tmp_path / "out", "ADLS_DELTA_INGESTION_DETAILS"))


def test_a_standard_only_key_is_an_input_and_written_on_use_default(tmp_path,
                                                                    monkeypatch, capsys):
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    config = fx.acfc_config(tmp_path, fx.PAIR1_OVERLAY, _no_autoset(tmp_path))
    frd, sttm, keyed = _standard_keyed_pair1(tmp_path, config)
    (spec,) = resolve_pair(frd, sttm, config)
    feed = spec.feed_id
    _run(config, [(frd, sttm)], parts=("framework",))
    out = capsys.readouterr().out
    line = next(ln for ln in out.splitlines() if f"feeds.{feed}.primary_key" in ln)
    assert line.startswith("INPUT_NEEDED") and "Standard band" in line
    (std_row,) = _iig_rows(tmp_path / "out", "STGDELTA_STDDELTA_INGESTION_DET")
    assert not std_row["TGT_PRIMARY_KEY"]                        # proposed, not applied

    answers = tmp_path / "answers.yaml"
    answers.write_text(f"inputs:\n  feeds:\n    {feed}:\n      primary_key:\n"
                       "        choice: use_default\n", encoding="utf-8")
    shutil.rmtree(tmp_path / "out")
    _run(config, [(frd, sttm)], answers, parts=("framework",))
    (std_row,) = _iig_rows(tmp_path / "out", "STGDELTA_STDDELTA_INGESTION_DET")
    assert std_row["TGT_PRIMARY_KEY"] == ",".join(keyed)
