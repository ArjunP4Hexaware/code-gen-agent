"""post66: which sheets a feed leaves empty BY RULE and why, the open lines a
feed's IIG review faces, and the report's labelled derived ratio.

post89: the App no longer shows the per-feed summary line; the empty-by-rule
reasons go into the review copy's REVIEW_SUMMARY sheet instead (one row per
empty sheet), where the BSA reviewing the IIG sees why a sheet has no rows.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from codegen import cli
from codegen.emit.framework import FrameworkArtefacts, derived_line, report_section
from codegen.generate_run import GenerateRun
from codegen.metadata_template import EMPTY_BY_RULE_LABELS, fixed_width_handler_empty_reason


def _scoped(config, tmp: Path):
    return config.model_copy(update={"output": config.output.model_copy(update={
        "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")})})


def test_a_delimited_files_fixed_width_handler_is_empty_by_rule(config):
    assert fixed_width_handler_empty_reason(
        SimpleNamespace(delimiter=",", file_format="csv"), config) == "delimited"
    assert fixed_width_handler_empty_reason(
        SimpleNamespace(delimiter="", file_format="fixed width"), config) is None
    assert fixed_width_handler_empty_reason(None, config) is None
    assert EMPTY_BY_RULE_LABELS == {"FILE_ADLS_INGESTION_DETAILS": "FILE_ADLS",
                                    "ADLS_FIXED_WIDTH_HANDLER": "FIXED_WIDTH"}


def _artefacts(spec, config, tmp: Path) -> FrameworkArtefacts:
    result = cli.generate_feed_result(
        spec, _scoped(config, tmp), dry_run=True, skip_tests=True, output_mode="framework",
        conventions_profile="acfc_prx", iig_template="iig_v2", console=False)
    return result.framework_artefacts


def test_empty_by_rule_names_only_empty_sheets_with_their_reason(pair1_config, pair1_spec,
                                                                 tmp_path):
    """Pair 1 is a fixed-width file: its handler sheet has rows, so it is never
    'empty by rule'. Read as a delimited file, the handler is empty and says why."""
    fixed = _artefacts(pair1_spec, pair1_config, tmp_path / "fixed")
    assert fixed.row_counts["ADLS_FIXED_WIDTH_HANDLER"] > 0
    assert "ADLS_FIXED_WIDTH_HANDLER" not in fixed.empty_by_rule
    for sheet in fixed.empty_by_rule:
        assert fixed.row_counts[sheet] == 0

    delimited = pair1_spec.model_copy(update={"delimiter": "|", "file_format": "delimited"})
    artefacts = _artefacts(delimited, pair1_config, tmp_path / "delimited")
    assert artefacts.row_counts["ADLS_FIXED_WIDTH_HANDLER"] == 0
    assert artefacts.empty_by_rule["ADLS_FIXED_WIDTH_HANDLER"] == "delimited"
    assert list(artefacts.empty_by_rule) == [s for s in artefacts.row_counts
                                             if s in artefacts.empty_by_rule]


def test_review_open_counts_the_feeds_own_and_the_run_level_lines():
    gen = GenerateRun.__new__(GenerateRun)
    gen.open_items = [
        {"line": "QUESTION", "key": "run.environment.storage_account", "scope": "run",
         "blocking": True},
        {"line": "INPUT_NEEDED", "key": "run.email_to", "scope": "run", "blocking": False},
        {"line": "INPUT_NEEDED", "key": "feeds.a.active_start_date", "scope": "feed",
         "feed": "a", "blocking": False},
        {"line": "INPUT_NEEDED", "key": "feeds.b.active_start_date", "scope": "feed",
         "feed": "b", "blocking": False},
    ]
    assert gen.review_open("a") == {"open": 3, "blocking": 1}


def test_the_report_carries_the_labelled_derived_ratio_and_the_sheet_rows():
    section = report_section(FrameworkArtefacts(
        files=[Path("framework") / "SYN_DDL.txt"],
        row_counts={"FILE_ADLS_INGESTION_DETAILS": 0, "ADLS_DELTA_INGESTION_DETAILS": 2},
        coverage={"derived": 458, "synthetic": 0, "needs_template": 0, "total": 2095}))
    assert ("cells derived from documents: 458 of 2,095; the rest are profile defaults, "
            "framework-assigned blanks, or open inputs") in section
    assert derived_line({"derived": 458, "total": 2095}) in section
    assert "| IIG rows: FILE_ADLS_INGESTION_DETAILS | 0 |" in section
    assert "| IIG rows: ADLS_DELTA_INGESTION_DETAILS | 2 |" in section
    assert "badge coverage" not in section


def test_the_review_copy_says_why_a_sheet_is_empty(pair1_config, pair1_spec, tmp_path):
    """post89: one REVIEW_SUMMARY row per sheet empty by rule (reason text from
    iig_review.empty_by_rule_reasons); a sheet with rows gets none."""
    review = pair1_config.iig_review
    delimited = pair1_spec.model_copy(update={"delimiter": "|", "file_format": "delimited"})
    artefacts = _artefacts(delimited, pair1_config, tmp_path / "delimited")
    book = load_workbook(next(p for p in artefacts.files if p.name.endswith("_IIG_REVIEW.xlsx")))
    rows = [r for r in book[review.summary_sheet].iter_rows(min_row=2, values_only=True)
            if r[2] == review.empty_by_rule_label]
    assert rows == [(sheet, None, review.empty_by_rule_label, review.empty_by_rule_owner, 0,
                     None, review.empty_by_rule_reasons[reason])
                    for sheet, reason in artefacts.empty_by_rule.items()]
    assert ("ADLS_FIXED_WIDTH_HANDLER", "delimited") in artefacts.empty_by_rule.items()

    fixed = _artefacts(pair1_spec, pair1_config, tmp_path / "fixed")
    book = load_workbook(next(p for p in fixed.files if p.name.endswith("_IIG_REVIEW.xlsx")))
    sheets = [r[0] for r in book[review.summary_sheet].iter_rows(min_row=2, values_only=True)
              if r[2] == review.empty_by_rule_label]
    assert "ADLS_FIXED_WIDTH_HANDLER" not in sheets
