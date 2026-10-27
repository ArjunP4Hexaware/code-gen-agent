"""post91: one set of framework artefacts per run and STTM (codegen.run_artefacts).

``<out>/combined/<STTM>_IIG.xlsx``, ``_IIG_REVIEW.xlsx``, ``_DDL.txt`` and
``_LOG.md`` beside the per-feed files. Rows merge by content only on the
run-level sheets (``combined.dedup_sheets``); a look-alike row anywhere else is
kept (prod writes it as a distinct row once the framework assigns its IDs); a
table block written twice the same way appears once, two different ones for
one table are both kept under a CONFLICT line. Synthetic inputs only.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from codegen import run_stages
from codegen.run_artefacts import merge_ddl, merge_payloads, sttm_name
from test_iig_corpus_cells import _generate
from test_pipeline_name_choice import _pair

STG, STD = "--stage table", "--standard table"


def _row(**values):
    return {"values": values, "badges": {k: {"badge": "synthetic"} for k in values}}


def _payload(rows_by_sheet: dict[str, list[dict]]) -> dict:
    return {"template": "iig_v2", "tabs": {
        sheet: {"headers": sorted(rows[0]["values"]) if rows else ["A"], "rows": rows}
        for sheet, rows in rows_by_sheet.items()}}


# --- names ------------------------------------------------------------------

def test_the_sttm_name_is_the_workbook_stem_made_safe(tmp_path):
    contract = tmp_path / "sttm.contract.json"
    contract.write_text('{"generated_from_workbook": "1787853350398_STTM Syn Feed (2).xlsx"}',
                        encoding="utf-8")
    # post99: the STTM role token at either end is dropped (Soham: confusing on
    # the run's DDL / IIG); a workbook named only "STTM" keeps it
    assert sttm_name(contract) == "Syn_Feed"
    assert sttm_name(tmp_path / "STTM_Syn.xlsx") == "Syn"
    assert sttm_name(tmp_path / "Syn Feed - sttm.xlsx") == "Syn_Feed"
    assert sttm_name(tmp_path / "STTMSyn.xlsx") == "STTMSyn"
    assert sttm_name(tmp_path / "STTM.xlsx") == "STTM"
    assert sttm_name(None) is None
    bad = tmp_path / "bad.json"
    bad.write_text("{not json", encoding="utf-8")
    assert sttm_name(bad) is None


# --- the IIG rows -----------------------------------------------------------

def test_only_run_level_sheets_merge_equal_rows():
    email = _row(TEMPLATE_NAME="SYN", STATUS="Success")
    child = _row(PIPELINE_NAME="PL_CHLD_SYN", PIPELINE_DESCRIPTION="same")
    merged = merge_payloads(
        [("feed_a", _payload({"EMAIL": [email], "SCHEDULE": [child]})),
         ("feed_b", _payload({"EMAIL": [email], "SCHEDULE": [child]}))],
        dedup_sheets=["EMAIL"])
    tabs = merged.payload["tabs"]
    assert len(tabs["EMAIL"]["rows"]) == 1 and merged.merged == {"EMAIL": 1}
    assert len(tabs["SCHEDULE"]["rows"]) == 2 and merged.look_alike == {"SCHEDULE": 1}


def test_a_feeds_own_equal_rows_are_never_merged():
    email = _row(TEMPLATE_NAME="SYN", STATUS="Success")
    merged = merge_payloads([("feed_a", _payload({"EMAIL": [email, email]}))],
                            dedup_sheets=["EMAIL"])
    assert len(merged.payload["tabs"]["EMAIL"]["rows"]) == 2 and not merged.merged


def test_rows_stay_together_per_feed_and_row_references_move():
    def schedule(name):
        parent = _row(PIPELINE_NAME=name, PARENT_PIPELINE_ID="")
        parent["badges"]["PARENT_PIPELINE_ID"]["tooltip"] = (
            "assigned by the engineer: the PIPELINE_ID of row 1 (grand_master)")
        return [_row(PIPELINE_NAME=f"{name}_GM", PARENT_PIPELINE_ID="0"), parent]
    merged = merge_payloads([("feed_a", _payload({"S": schedule("A")})),
                             ("feed_b", _payload({"S": schedule("B")}))], dedup_sheets=[])
    rows = merged.payload["tabs"]["S"]["rows"]
    assert [r["values"]["PIPELINE_NAME"] for r in rows] == ["A_GM", "A", "B_GM", "B"]
    tips = [r["badges"]["PARENT_PIPELINE_ID"].get("tooltip", "") for r in rows]
    assert "row 1 (grand_master)" in tips[1] and "row 3 (grand_master)" in tips[3]
    assert merged.ranges["S"] == [("feed_a", 2, 3), ("feed_b", 4, 5)]


def test_the_masters_are_written_once_and_children_follow_them():
    """post99 (Soham, the pair_11 golden): the grand master / master rows every
    feed repeats are written once; a child row is never merged, and its "row n"
    reference follows its master to the kept row. Children that share a name but
    differ are kept and named."""
    from codegen.config import CombinedRowDedup

    def schedule(child_text):
        gm = _row(PIPELINE_NAME="PL_GMSTR_SYN", PARENT_PIPELINE_ID="0")
        master = _row(PIPELINE_NAME="PL_MSTR_SYN", PARENT_PIPELINE_ID="")
        master["badges"]["PARENT_PIPELINE_ID"]["tooltip"] = (
            "the PIPELINE_ID of row 1 (grand_master)")
        child = _row(PIPELINE_NAME="PL_CHLD_SYN", PARENT_PIPELINE_ID=child_text)
        child["badges"]["PARENT_PIPELINE_ID"]["tooltip"] = "the PIPELINE_ID of row 2 (master)"
        return [gm, master, child]

    rule = {"S": CombinedRowDedup(column="PIPELINE_NAME", prefixes=["PL_GMSTR_", "PL_MSTR_"])}
    merged = merge_payloads([("feed_a", _payload({"S": schedule("")})),
                             ("feed_b", _payload({"S": schedule("")})),
                             ("feed_c", _payload({"S": schedule("x")}))], [], rule)
    rows = merged.payload["tabs"]["S"]["rows"]
    assert [r["values"]["PIPELINE_NAME"] for r in rows] == [
        "PL_GMSTR_SYN", "PL_MSTR_SYN", "PL_CHLD_SYN", "PL_CHLD_SYN", "PL_CHLD_SYN"]
    assert merged.merged == {"S": 4} and merged.look_alike == {"S": 1}
    tips = [r["badges"]["PARENT_PIPELINE_ID"].get("tooltip", "") for r in rows]
    assert all("row 2 (master)" in t for t in tips[2:])
    assert merged.same_name == {"S": ["PL_CHLD_SYN"]}
    assert merged.ranges["S"] == [("feed_a", 2, 4), ("feed_b", 5, 5), ("feed_c", 6, 6)]


def test_object_ids_are_numbered_across_the_run_and_keep_their_links():
    """post99 (Nikshit): OBJECT_ID 1..n across the feeds in the combined IIG; a
    feed's rows on every sheet move alike, so a DQ row still names its object."""
    def feed(objects):
        return _payload({
            "ADLS": [_row(OBJECT_ID=str(i), NAME=f"f{i}") for i in range(1, objects + 1)],
            "DQ": [_row(OBJECT_ID=str(i), RULE="r") for i in range(1, objects + 1)
                   for _ in range(2)]})
    merged = merge_payloads([("a", feed(2)), ("b", feed(1)), ("c", feed(1))], [],
                            object_ids="run")
    tabs = merged.payload["tabs"]
    assert [r["values"]["OBJECT_ID"] for r in tabs["ADLS"]["rows"]] == ["1", "2", "3", "4"]
    assert [r["values"]["OBJECT_ID"] for r in tabs["DQ"]["rows"]] == [
        "1", "1", "2", "2", "3", "3", "4", "4"]
    assert "object 3" in tabs["ADLS"]["rows"][2]["badges"]["OBJECT_ID"]["tooltip"]
    # the default keeps each feed's own numbering
    plain = merge_payloads([("a", feed(1)), ("b", feed(1))], [])
    assert [r["values"]["OBJECT_ID"] for r in plain.payload["tabs"]["ADLS"]["rows"]] == ["1", "1"]


# --- the DDL ----------------------------------------------------------------

def _block(banner: str, table: str, cols: str = "A STRING") -> str:
    return f"{banner}\n\nCREATE OR REPLACE TABLE {table}\n(\n{cols}\n)\n\nUSING delta"


def test_an_identical_table_block_is_written_once():
    a = "\n\n".join([_block(STG, "c.s.t1"), _block(STD, "c.d.t1")])
    b = "\n\n".join([_block(STG, "c.s.t1"), _block(STG, "c.s.t2")])
    text, skipped, conflicts = merge_ddl([("a", a), ("b", b)], (STG, STD))
    assert skipped == 1 and conflicts == []
    assert text.count("CREATE OR REPLACE TABLE c.s.t1") == 1
    assert text.count("CREATE OR REPLACE TABLE") == 3


def test_two_different_blocks_for_one_table_are_both_kept_under_a_conflict():
    text, skipped, conflicts = merge_ddl(
        [("a", _block(STG, "c.s.t1", "A STRING")), ("b", _block(STG, "c.s.t1", "B INT"))],
        (STG, STD))
    assert conflicts == ["c.s.t1"] and skipped == 0
    assert text.count("CREATE OR REPLACE TABLE c.s.t1") == 2
    assert text.index("-- CONFLICT") < text.index("CREATE OR REPLACE TABLE c.s.t1")


# --- a whole run --------------------------------------------------------------

def _combined(out: Path) -> dict:
    (block,) = run_stages.read_stage_manifest(out)["combined"]
    return block


def test_a_run_writes_one_set_for_its_sttm(tmp_path):
    config, _spec, pair = _pair(tmp_path)
    assert _generate(config, pair, only_feed=None) == 0
    out = tmp_path / "out"
    block = _combined(out)
    folder = out / config.combined.dir
    assert sorted(p.name for p in folder.iterdir()) == sorted(block["files"].values())
    assert block["excluded"] == [] and block["failed_checks"] == []
    feeds = block["feeds"]
    assert len(feeds) == 3
    clean = load_workbook(folder / block["files"]["iig"])
    sheet = clean["ADLS_DELTA_INGESTION_DETAILS"]
    # every feed's ADLS row, in feed order (no look-alikes there)
    per_feed = sum(load_workbook(next((out / f / "framework").glob("*_IIG.xlsx")))
                   ["ADLS_DELTA_INGESTION_DETAILS"].max_row - 1 for f in feeds)
    assert sheet.max_row - 1 == per_feed
    log = (folder / block["files"]["log"]).read_text(encoding="utf-8")
    assert log.startswith(f"# {block['sttm']} DDL and IIG")
    assert all(f"`{f}`" in log for f in feeds) and "## Review flow" in log
    ddl = (folder / block["files"]["ddl"]).read_text(encoding="utf-8")
    assert ddl.count("CREATE") == sum(
        next((out / f / "framework").glob("*_DDL.txt")).read_text(encoding="utf-8")
        .count("CREATE") for f in feeds) - block["ddl_blocks_merged"]


def test_the_same_run_writes_the_same_bytes(tmp_path):
    config, _spec, pair = _pair(tmp_path)
    folder = tmp_path / "out" / config.combined.dir
    _generate(config, pair, only_feed=None)
    first = {p.name: p.read_bytes() for p in folder.iterdir()}
    _generate(config, pair, only_feed=None)
    assert {p.name: p.read_bytes() for p in folder.iterdir()} == first


def test_a_run_of_one_feed_writes_no_combined_set(tmp_path):
    config, spec, pair = _pair(tmp_path)
    _generate(config, pair, only_feed=spec.feed_id)
    assert not (tmp_path / "out" / config.combined.dir).exists()
    assert "combined" not in run_stages.read_stage_manifest(tmp_path / "out")


def test_switched_off_there_is_no_combined_set(tmp_path):
    config, _spec, pair = _pair(tmp_path)
    off = config.model_copy(update={"combined": config.combined.model_copy(
        update={"enabled": False})})
    _generate(off, pair, only_feed=None)
    assert not (tmp_path / "out" / config.combined.dir).exists()


def test_all_stage_tables_come_first_then_all_standard_tables():
    """post97: by layer, as one feed's own DDL orders several tables - never feed by
    feed (stage A, standard A, stage B, standard B)."""
    a = "\n\n".join([_block(STG, "c.s.a"), _block(STD, "c.d.a")])
    b = "\n\n".join([_block(STG, "c.s.b"), _block(STD, "c.d.b")])
    shared = _block(STG, "c.s.a")                        # feed C repeats feed A's stage table
    text, skipped, _conflicts = merge_ddl([("a", a), ("b", b), ("c", shared)], (STG, STD))
    tables = [line.split()[-1] for line in text.splitlines() if line.startswith("CREATE")]
    assert tables == ["c.s.a", "c.s.b", "c.d.a", "c.d.b"]
    assert skipped == 1
    # every block keeps its own banner, in the client's format
    assert text.count(STG) == 2 and text.count(STD) == 2


def test_a_rerun_that_keeps_no_feed_leaves_no_stale_combined_set(tmp_path):
    """Audit #14 (post132): the combined folder was written only for STTMs with a
    kept feed and never cleared, so a rerun into the same output folder whose
    every feed now GATE_FAILs kept the previous run's deployable-looking
    <STTM>_IIG.xlsx / _DDL.txt. A full run now clears the earlier set first."""
    config, _spec, pair = _pair(tmp_path)
    assert _generate(config, pair, only_feed=None) == 0
    folder = tmp_path / "out" / config.combined.dir
    first = sorted(p.name for p in folder.iterdir())
    assert any(n.endswith("_IIG.xlsx") for n in first)
    (folder / "notes_by_a_person.txt").write_text("kept", encoding="utf-8")
    import codegen.cli as cli_module
    real = cli_module.generate_feed_result

    def failing(*args, **kwargs):
        result = real(*args, **kwargs)
        gate = result.gate.model_copy(update={"verdict": "FAIL"})
        return result.__class__(**{**result.__dict__, "gate": gate})
    import pytest as _pytest
    with _pytest.MonkeyPatch.context() as mp:
        mp.setattr(cli_module, "generate_feed_result", failing)
        _generate(config, pair, only_feed=None)
    left = sorted(p.name for p in folder.iterdir())
    assert left == ["notes_by_a_person.txt"]          # the old set is gone, nothing else
