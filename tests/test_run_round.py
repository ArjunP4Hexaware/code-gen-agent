"""scripts/run_round.py (2026-10-13): one validation round with no input choice
left to the operator. On the synthetic pair 11 (the one-block / many-files SD
shape): the happy path end to end, a planted blank-column STTM stopping at the
document gate (exit 2), no key (exit 2), a scrub hit in the push dir (exit 3),
the gitignored inputs file and the <FILL refusal."""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
_spec = importlib.util.spec_from_file_location("run_round", ROOT / "scripts" / "run_round.py")
run_round = importlib.util.module_from_spec(_spec)
sys.modules["run_round"] = run_round
_spec.loader.exec_module(run_round)
fp = run_round.fingerprint_mod

STTM = ROOT / "fixtures" / "acfc_shapes" / "sttm" / "pair_11_family_b.xlsx"
FRD = ROOT / "fixtures" / "acfc_shapes" / "frd" / "f1_pair_11_multi_file_catalog.docx"
SFMC = ROOT / "fixtures" / "reference" / "SFMC_IIG.xlsx"
KEY = "syn-run-round-test-key"
GROUPS = ["SYN-GRP-001", "SYN-GRP-002"]
PAIR = "syn_pair"
USER = "someone.real" + "@" + "corp-mail.net"   # assembled: no address here
USERS = "/Workspace/" + "Users/"
PUSHABLE = {"RUN.md", "compare.txt", "fingerprint_structure.json", "golden_structure.json",
            "golden_structure.md", "mismatches.md", "row_keys.md", "SELF_CHECK.md",
            "score_totals.json"}


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setenv(fp.KEY_ENV, KEY)
    monkeypatch.setenv("CODEGEN_NOTIFICATION_EMAILS", "syn.dl.prodsupport@synthetic.example")
    build_goldens = run_round._load("build_goldens", ROOT / "tools" / "build_goldens.py")
    monkeypatch.setitem(build_goldens.PAIR_GROUPS, PAIR, GROUPS)
    docs = tmp_path / "docs"
    docs.mkdir()
    sttm, frd = docs / "STTM title words here.xlsx", docs / ("FRD title words here (" + "1).docx")
    sttm.write_bytes(STTM.read_bytes())
    frd.write_bytes(FRD.read_bytes())
    reference = tmp_path / "reference.json"
    reference.write_text(json.dumps([fp.fingerprint(sttm, 1, KEY.encode()),
                                     fp.fingerprint(frd, 2, KEY.encode())]), encoding="utf-8")
    answers = tmp_path / "answers.yaml"
    answers.write_text("gaps: {}\n", encoding="utf-8")
    golden = tmp_path / "golden_dir" / "goldens"
    golden.mkdir(parents=True)
    src = openpyxl.load_workbook(SFMC, data_only=True)
    for gid in GROUPS:                                   # one export per group
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for ws in src.worksheets:
            target = wb.create_sheet(ws.title)
            for row in ws.iter_rows(values_only=True):
                target.append(list(row))
        wb.save(golden / f"gid_{gid}.xlsx")
    return {"tmp": tmp_path, "sttm": sttm, "frd": frd, "reference": reference,
            "answers": answers, "golden_dir": golden.parent, "out": tmp_path / "push"}


def _argv(e, *, sttm=None, frd=None, extra=()):
    """No --answers: a round runs without one (pass it in ``extra``)."""
    return ["--round", "9", "--pair", PAIR, "--sttm", str(sttm or e["sttm"]), "--frd",
            str(frd or e["frd"]), "--golden-dir", str(e["golden_dir"]),
            "--out", str(e["out"]), "--reference", str(e["reference"]), *extra]


def _assert_no_names(e, text: str) -> None:
    for leak in ("title words", e["sttm"].name, e["frd"].name, str(e["tmp"]),
                 e["tmp"].as_posix(), "pair_11_family_b"):
        assert leak not in text, leak


def test_the_happy_path_end_to_end(env):
    code = run_round.main(_argv(env))
    run_md = (env["out"] / "RUN.md").read_text(encoding="utf-8")
    assert code == 0, run_md
    assert {p.name for p in env["out"].iterdir()} <= PUSHABLE
    for section in ("## Fingerprint", "## Document gate", "## Build", "## CLI",
                    "## Golden (rebuilt on this SHA)", "## Rows per sheet", "## Score"):
        assert section in run_md, section
    assert "**exit 0** — ok" in run_md and "- answers: none given" in run_md
    assert "identical" in run_md and "verdicts: PASS_WITH_FLAGS ×3" in run_md
    assert "`codegen layout --workbook <STTM> --frd <FRD>" in run_md
    assert f"{PAIR}/DATA_FACTORY_PIPELINE_SCHEDULE: dedup " in run_md
    assert "agent " in run_md and "rows (deduped) · paired" in run_md
    _assert_no_names(env, run_md)
    structure = json.loads((env["out"] / "fingerprint_structure.json").read_text("utf-8"))
    assert all("body_hash" not in s and "column_hashes" not in s
               for e in structure for s in e.get("sheets", []))
    totals = json.loads((env["out"] / "score_totals.json").read_text("utf-8"))
    assert "score" in totals["totals"] and totals["sheets"]
    work = env["out"].with_name("push.work")
    assert (work / "fp_keyed.json").is_file() and list((work / "out").rglob("*_IIG.xlsx"))


def _blank_column(src: Path, dst: Path) -> Path:
    wb = openpyxl.load_workbook(src)
    ws = next(s for s in wb.worksheets if s.max_row > 5)
    body = range(3, ws.max_row + 1)
    column = max(range(1, ws.max_column + 1),
                 key=lambda c: sum(ws.cell(r, c).value is not None for r in body))
    for r in body:
        ws.cell(r, column).value = None
    wb.save(dst)
    return dst


def test_a_planted_blank_column_stops_at_the_gate(env):
    planted = _blank_column(env["sttm"], env["tmp"] / "docs" / "STTM planted.xlsx")
    code = run_round.main(_argv(env, sttm=planted))
    run_md = (env["out"] / "RUN.md").read_text(encoding="utf-8")
    assert code == 2
    assert "HARD file#1 vs file#1 sheet[" in run_md and "exit 2 — STOP" in run_md
    assert "## CLI" not in run_md                                  # nothing ran after it
    assert not (env["out"].with_name("push.work") / "frd.json").exists()
    assert "HARD" in (env["out"] / "compare.txt").read_text(encoding="utf-8")
    _assert_no_names(env, run_md)


def test_no_key_is_exit_2(env, monkeypatch, tmp_path):
    monkeypatch.delenv(fp.KEY_ENV)
    monkeypatch.setattr(fp, "KEY_FILE", tmp_path / "absent.key")
    assert run_round.main(_argv(env)) == 2
    assert fp.NO_KEY in (env["out"] / "RUN.md").read_text(encoding="utf-8")


def test_a_scrub_hit_in_the_push_dir_is_exit_3_and_run_md_says_only_classes(env):
    env["out"].mkdir()
    (env["out"] / "stale.txt").write_text("owner jdoe" + "@" + "corp-mail.net", encoding="utf-8")
    planted = _blank_column(env["sttm"], env["tmp"] / "docs" / "STTM planted.xlsx")
    assert run_round.main(_argv(env, sttm=planted)) == 3
    run_md = (env["out"] / "RUN.md").read_text(encoding="utf-8")
    assert "**exit 3** — the scrub check is not clean" in run_md and "email_address: 1" in run_md
    assert "corp-mail" not in run_md


def test_inputs_file_names_the_documents_and_fill_is_refused(env):
    inputs = env["tmp"] / "round_inputs.yaml"
    planted = _blank_column(env["sttm"], env["tmp"] / "docs" / "STTM planted.xlsx")
    inputs.write_text(f"sttm: '{planted.as_posix()}'\nfrd: '{env['frd'].as_posix()}'\n",
                      encoding="utf-8")
    argv = _argv(env, extra=("--answers", str(env["answers"])))
    i = argv.index("--sttm")
    argv[i:i + 4] = ["--inputs", str(inputs)]
    assert run_round.main(argv) == 2                              # the planted copy was read
    env["answers"].write_text('gaps:\n  "feeds[1].file_name_patterns": {value: "<FILL>"}\n',
                              encoding="utf-8")
    assert run_round.main(argv) == 1


def test_display_path_and_redaction_never_show_a_user_or_a_name():
    sttm = Path(f"{USERS}{USER}/frd_sttm_pairs/pair_2/STTM X Y.xlsx")
    assert run_round.display_path(sttm, "STTM") == "/Workspace/Users/<me>/…/<STTM>"
    assert run_round.display_path(Path("/tmp/x/a.xlsx"), "STTM") == "<local>/…/<STTM>"
    redact = run_round.Redactor({"STTM": sttm}, {})
    line = redact(f"read {sttm} · {USERS}{USER}/other/f.py")
    assert line == "read <STTM> · /Workspace/Users/<me>/other/f.py"


def test_a_layout_that_asks_anything_is_exit_3_with_its_keys(env):
    # synthetic pair 2: neither document names a file, so the layout owes two answers
    sttm = env["tmp"] / "docs" / "STTM owes.xlsx"
    frd = env["tmp"] / "docs" / "FRD owes.docx"
    sttm.write_bytes((ROOT / "fixtures/acfc_shapes/sttm/pair_2_family_b.xlsx").read_bytes())
    frd.write_bytes((ROOT / "fixtures/acfc_shapes/frd/f1_pair_2_variant.docx").read_bytes())
    env["reference"].write_text(json.dumps([fp.fingerprint(sttm, 1, KEY.encode()),
                                            fp.fingerprint(frd, 2, KEY.encode())]),
                                encoding="utf-8")
    code = run_round.main(_argv(env, sttm=sttm, frd=frd))
    run_md = (env["out"] / "RUN.md").read_text(encoding="utf-8")
    assert code == 3, run_md
    assert "**exit 3** — layout asks for 2 answer(s)" in run_md
    assert "QUESTION `feeds[0].file_patterns`" in run_md
    assert "QUESTION `feeds[1].file_patterns`" in run_md
    assert "extract-sttm" not in run_md and "## Golden" not in run_md


def test_answer_keys_are_keys_only_and_hide_a_sheet_name():
    text = "\n".join([
        "NOTE           feeds[1].file_name_patterns — by elimination: the only file left",
        "UNRESOLVED     Secret Sheet/stage/schema — no header names it",
        "QUESTION       Secret Sheet/band[2]/layer — which layer",
        "CROSSCHECK     frequency: agree — not an answer line"])
    sheet = run_round.hashlib.sha256(b"Secret Sheet").hexdigest()[:12]
    assert run_round.answer_keys(text) == [
        ("NOTE", "feeds[1].file_name_patterns"),
        ("UNRESOLVED", f"sheet#{sheet}/stage/schema"),
        ("QUESTION", f"sheet#{sheet}/band[2]/layer")]


def test_a_missing_group_export_stops_before_scoring(env):
    (env["golden_dir"] / "goldens" / f"gid_{GROUPS[1]}.xlsx").unlink()
    assert run_round.main(_argv(env)) == 1
    run_md = (env["out"] / "RUN.md").read_text(encoding="utf-8")
    assert "GROUP_ID exports found 1 of 2" in run_md and "1 export(s) missing — STOP" in run_md
    assert "## Score" not in run_md and not (env["out"] / "score_totals.json").exists()
