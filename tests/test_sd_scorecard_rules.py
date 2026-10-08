"""The first real-row scorecard (2026-10-08): the generated IIG rows of the
three SD (vendor) feeds against their actual config rows. Each
finding became a general rule or a feed-FAMILY convention; each is pinned
here on SYNTHETIC inputs (no client data):

* the SD family = pair 11 (the synthetic one-block / many-files analogue of
  the SD pair: FILE_DETAILS, three MAPPING- sheets, a vendor cell with a
  trailing id, a Frequency cell pointing at File Details) generated under the
  committed ACFC environment overlay (``config/overlays/acfc_env.yaml``,
  which carries the SD family conventions);
* the PRX family (pair 1) and the CAQH-style pair 4 keep their full-cell
  goldens (tests/test_m4_acceptance.py, tests/test_pair4_full_golden.py).

Rules: 1 SRC_FILE_DELIMITER is a character; 2 OBJECT_NAME / SRC_FILE_NAME
(one row per pattern); 3 SOURCE = the vendor display name; 4 the audit
columns + casing by family; 5 SRC_DATA_TYPE String:<target> for a file;
6 RECYCL_* derived when stated, else open (never an asserted N); 7 FREQUENCY
delivery statement over narrative, frequency_ambiguous; 8 spark_unavailable /
ruff_formatted in the gate; 9 one vdd_scope_mismatch flag.
"""

from __future__ import annotations

import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from acfc_shapes import sttm as sttm_fixtures
from acfc_shapes.common import xlsx_bytes
from codegen import cli
from codegen.config import load_config
from codegen.extract import contract_to_json as sttm_to_json
from codegen.extract import extract_contract
from codegen.extract.frd_docx import DocxContent, frequency_mentions
from codegen.extract.frd_docx import contract_to_json as frd_to_json
from codegen.gate import tests_runner
from codegen.gate.preflight import _ruff_check
from codegen.gate.verdict import compute_verdict
from codegen.layout.model import MockLayoutProvider
from codegen.layout.resolve import resolve_pair
from codegen.metadata_template import _object_name, _source_cell
from codegen.resolve.resolver import _resolve_delimiter, delimiter_char
from codegen.resolve.resolver import resolve_pair as resolve_contracts

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
ACFC_ENV = REPO / "config" / "overlays" / "acfc_env.yaml"
PAIR_11_FRD = SHAPES / "frd" / "f1_pair_11_multi_file_catalog.docx"
DATE = "2026-01-01"
WORDS = {"comma": ",", "pipe": "|", "tab": "\t", "semicolon": ";"}


@pytest.fixture(scope="module")
def sd_config():
    """The SD family as ACFC runs it: the committed environment overlay."""
    return load_config(REPO / "config" / "config.yaml", overlays=[ACFC_ENV])


RUN_STATEMENT = "'Frequency of data refresh \u2013 Monthly Run' (FRD paragraph 9)"


def _generate(config, workbook, tmp: Path, run_statement: str | None = None):
    sttm = tmp / "sttm.xlsx"
    sttm.write_bytes(xlsx_bytes(workbook))
    pair = resolve_pair(sttm, PAIR_11_FRD, config, provider=MockLayoutProvider([]),
                        cache_dirs=[], use_cache=False, generated_date=DATE)
    assert pair.questions == []
    frd_json = tmp / "frd.contract.json"
    frd = pair.frd_contract
    if run_statement is not None:
        # The synthetic stand-in for the FRD's run-cadence sentence (the real SD
        # FRD says its data refresh runs monthly; pair 11's FRD has no such line).
        frd = frd.model_copy(update={"feeds": [
            f.model_copy(update={"frequency_mentions": [*f.frequency_mentions, run_statement]})
            for f in frd.feeds]})
    frd_json.write_text(frd_to_json(frd), encoding="utf-8")
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    sttm_json = tmp / "sttm.contract.json"
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    scoped = config.model_copy(update={"output": config.output.model_copy(update={
        "dir": str(tmp / "out"), "reports_dir": str(tmp / "reports")})})
    out = {}
    for spec in resolve_contracts(frd_json, sttm_json, config):
        gate = cli._generate_feed(spec, scoped, dry_run=True, skip_tests=True,
                                  output_mode="framework", conventions_profile="acfc_prx",
                                  iig_template="iig_v2")
        book = load_workbook(tmp / "out" / spec.feed_slug / "framework" / "config_rows.xlsx")
        rows = list(book["ADLS_DELTA_INGESTION_DETAILS"].iter_rows(values_only=True))
        out[spec.feed_slug] = (gate, [dict(zip(rows[0], r, strict=True)) for r in rows[1:]])
    return out


@pytest.fixture(scope="module")
def sd_rows(sd_config, tmp_path_factory):
    return _generate(sd_config, sttm_fixtures.build_pair11(), tmp_path_factory.mktemp("sd"),
                     run_statement=RUN_STATEMENT)


def _blank(value) -> bool:
    return value in (None, "")


# ------------------------------------------------------- the SD family, pinned


SD_EXPECTED = {
    "vc_enrollment": {"OBJECT_NAME": "enrollment_package",
                      "SRC_FILE_NAME": "enrollment_package*", "SRC_FILE_DELIMITER": ",",
                      "FREQUENCY": "Monthly"},
    "vc_disenrollment": {"OBJECT_NAME": "disenrollment_package",
                         "SRC_FILE_NAME": "disenrollment_package*", "SRC_FILE_DELIMITER": ",",
                         "FREQUENCY": "Monthly"},
    "vc_individual_risk": {"OBJECT_NAME": "vc_ind_risk_data_package_REGION_A",
                           "SRC_FILE_NAME": "vc_ind_risk_data_package_REGION_A*",
                           "SRC_FILE_DELIMITER": "|", "FREQUENCY": "Monthly"},
}


@pytest.mark.parametrize("slug", sorted(SD_EXPECTED))
def test_the_sd_family_rows(sd_rows, slug):
    gate, rows = sd_rows[slug]
    (row,) = rows                                    # one pattern -> one row, never joined
    for column, value in SD_EXPECTED[slug].items():
        assert row[column] == value, column
    assert row["SOURCE"] == "VENDOR_C"                              # rule 3: the id dropped
    # rule 4: the SD family's audit columns, typed in the data columns' casing
    assert row["TGT_COLUMN_NAMES"].split(",")[-4:] == [
        "LOB", "SRC_FILE_NAME", "REC_CREATION_TIME", "REC_UPDATED_TIME"]
    assert row["TGT_DATA_TYPE"].split(",")[-4:] == ["String", "String", "Timestamp",
                                                    "Timestamp"]
    # rule 5: a file's fields are strings
    assert all(t.startswith("String:") for t in row["SRC_DATA_TYPE"].split(","))
    # rule 6: no document states a recycle -> open, never an asserted N
    for column in ("RECYCL_ENBL_FLG", "RECYCL_TBL_NM", "RECYCL_ADLS_PATH"):
        assert _blank(row[column]), column
    assert gate.verdict == "PASS_WITH_FLAGS"


def test_frequency_is_the_run_cadence_and_a_different_delivery_is_flagged(sd_rows):
    # Files delivered twice a year (File Details), the pipeline runs monthly.
    gate, (row,) = sd_rows["vc_enrollment"]
    assert row["FREQUENCY"] == "Monthly"
    (flag,) = [f for f in gate.flags if f.startswith("frequency_delivery_differs")]
    assert "FREQUENCY 'Monthly' is the pipeline run cadence" in flag
    assert "the files are delivered 'Yearly Twice' (→ 'Yearly'" in flag
    assert "Monthly Run" in flag                                      # the run statement, cited
    gate, (row,) = sd_rows["vc_individual_risk"]                     # delivery Monthly too
    assert row["FREQUENCY"] == "Monthly"
    assert not [f for f in gate.flags if f.startswith("frequency_delivery_differs")]


def _freq_feed(frequency, mentions=()):
    return SimpleNamespace(frequency=frequency, frequency_mentions=list(mentions))


def test_frequency_sources_in_order(config):
    from codegen.metadata_template import _frequency

    tpl = config.metadata.templates["iig_v2"]
    # 1. the FRD's run / schedule statement
    cell = _frequency(_freq_feed("Yearly Twice", [RUN_STATEMENT]), None, tpl)
    assert cell["value"] == "Monthly"
    assert "file delivery: stated as 'Yearly Twice'" in cell["badge_entry"]["tooltip"]
    assert cell["badge_entry"]["note"].startswith("frequency_delivery_differs")
    # a narrative mention without a run word is not a run statement
    plain = _frequency(_freq_feed("Yearly Twice", ["'produces the report monthly' (FRD p 1)"]),
                       None, tpl)
    assert plain["value"] == "Yearly" and "note" not in plain["badge_entry"]
    # 2. the schedule inventory's PIPELINE_FREQUENCY (no run statement)
    inventory = tpl.model_copy(update={"template_rows": {
        "DATA_FACTORY_PIPELINE_SCHEDULE": [{"PIPELINE_NAME": "PL_X",
                                            "PIPELINE_FREQUENCY": "Daily"}]}})
    cell = _frequency(_freq_feed("Monthly"), None, inventory)
    assert cell["value"] == "Daily" and cell["badge_entry"]["badge"] == "synthetic"
    assert "DATA_FACTORY_PIPELINE_SCHEDULE inventory" in cell["badge_entry"]["tooltip"]
    assert cell["badge_entry"]["note"].startswith("frequency_delivery_differs")
    # same cadence both ways: no flag
    same = _frequency(_freq_feed("Monthly", [RUN_STATEMENT]), None, tpl)
    assert same["value"] == "Monthly" and "note" not in same["badge_entry"]
    # two run statements naming different cadences: no single run cadence ->
    # the delivery cadence, unflagged
    split = _frequency(_freq_feed("Yearly Twice", [RUN_STATEMENT, "'runs daily' (FRD p 2)"]),
                       None, tpl)
    assert split["value"] == "Yearly" and "note" not in split["badge_entry"]
    assert "no run / schedule statement: the delivery cadence" in (
        split["badge_entry"]["tooltip"])


def test_open_recycle_cells_offer_the_convention_shape(sd_config, tmp_path):
    from codegen.metadata_template import _recycle_cells

    spec = SimpleNamespace(recycle=None)
    feed = SimpleNamespace(recycle_rule=None, landing_location="/mftlanding/inbound/dom/vc/")
    tpl = sd_config.metadata.templates["iig_v2"]
    cells = _recycle_cells(tpl, "ADLS_DELTA_INGESTION_DETAILS", feed, spec, sd_config,
                           "vc_enrollment")
    assert set(cells) == {"RECYCL_ENBL_FLG", "RECYCL_TBL_NM", "RECYCL_ADLS_PATH"}
    tooltip = cells["RECYCL_TBL_NM"]["badge_entry"]["tooltip"]
    assert all(c["value"] == "" for c in cells.values())
    assert "'vc_enrollment_recycle'" in tooltip and "Recycle/vc_enrollment" in tooltip
    owners = sd_config.iig_review.owners
    assert owners["ADLS_DELTA_INGESTION_DETAILS.RECYCL_ENBL_FLG"] == "engineer_confirms"


def test_a_stated_recycle_derives_every_recycle_cell(sd_config, tmp_path):
    frd = REPO / "fixtures" / "contracts" / "FRD_demo_cv_golden.contract.json"
    sttm = REPO / "fixtures" / "contracts" / "sttm_mapping_contracts_cv_golden.json"
    if not frd.is_file() or not sttm.is_file():
        pytest.skip("CV golden contracts absent")
    (spec,) = [s for s in resolve_contracts(frd, sttm, sd_config) if s.recycle is not None]
    scoped = sd_config.model_copy(update={"output": sd_config.output.model_copy(update={
        "dir": str(tmp_path / "out"), "reports_dir": str(tmp_path / "reports")})})
    cli._generate_feed(spec, scoped, dry_run=True, skip_tests=True, output_mode="framework",
                       conventions_profile="acfc_prx", iig_template="iig_v2")
    book = load_workbook(tmp_path / "out" / spec.feed_slug / "framework" / "config_rows.xlsx")
    rows = list(book["ADLS_DELTA_INGESTION_DETAILS"].iter_rows(values_only=True))
    row = dict(zip(rows[0], rows[1], strict=True))
    stage = spec.detail_segment.stage_table.table
    assert row["RECYCL_ENBL_FLG"] == "Y"
    assert row["RECYCL_TBL_NM"] == spec.recycle.recycle_table.table
    assert row["RECYCL_ADLS_PATH"].endswith(f"Recycle/{stage}")
    assert row["RECYCL_RETN_DAYS"] == str(spec.recycle.spec.recycle_window_days)


def test_the_prx_family_keeps_its_n_when_no_recycle_is_stated(config):
    from codegen.metadata_template import _recycle_cells

    tpl = config.metadata.templates["iig_v2"]
    tpl = tpl.model_copy(update={"family_conventions": tpl.family_conventions.model_copy(
        update={"recycle_unstated": "N"})})
    cells = _recycle_cells(tpl, "ADLS_DELTA_INGESTION_DETAILS",
                           SimpleNamespace(recycle_rule=None, landing_location=None),
                           SimpleNamespace(recycle=None), config, "t")
    assert {k: v["value"] for k, v in cells.items()} == {"RECYCL_ENBL_FLG": "N"}


# ------------------------------------------------------- rule 5: STTM types


def test_a_typed_source_field_is_written_string_and_flagged(sd_config, tmp_path):
    wb = sttm_fixtures.build_pair11()
    ws = wb["MAPPING-VC_ENROLLMENT"]
    ws.cell(row=5, column=5, value="Decimal")          # POVERTY_PCT's source DataType
    out = _generate(sd_config, wb, tmp_path)
    gate, (row,) = out["vc_enrollment"]
    assert all(t.startswith("String:") for t in row["SRC_DATA_TYPE"].split(","))
    (flag,) = [f for f in gate.flags if f.startswith("src_type_coerced_string")]
    assert "1 source field(s)" in flag and "POVERTY_PCT 'Decimal'" in flag


# ------------------------------------------------------- unit rules


@pytest.mark.parametrize("stated,char", [
    ("|", "|"), (",", ","), ("Pipe", "|"), ("comma delimited", ","), ("TAB", "\t"),
    ("\\t", "\t"), ("Pipe (|)", "|"), ("File Data Ingestion", None), (None, None),
])
def test_a_stated_delimiter_is_a_character(stated, char):
    assert delimiter_char(stated, WORDS) == char


def _delimiter_inputs(stated=None, patterns=("x_YYYY.csv",), fmt="File Data Ingestion"):
    frd = SimpleNamespace(delimiter=stated, file_format=fmt, file_name_patterns=list(patterns))
    sttm = SimpleNamespace(source_file=SimpleNamespace(delimiter=None, name_pattern=None))
    return frd, sttm


def test_the_resolver_writes_the_character_and_flags_the_extension(config):
    errors, flags = [], []
    assert _resolve_delimiter(*_delimiter_inputs("Pipe"), errors, config, flags) == "|"
    assert flags == [] and errors == []
    assert _resolve_delimiter(*_delimiter_inputs(patterns=("a.psv",)), errors, config,
                              flags) == "|"
    assert flags[0].startswith("delimiter_from_extension:a.psv -> '|'")
    flags.clear()
    assert _resolve_delimiter(*_delimiter_inputs("File Data Ingestion"), errors, config,
                              flags) == ","
    assert [f.split(" — ")[0] for f in flags] == ["delimiter_unreadable:FRD",
                                                  "delimiter_from_extension:x_YYYY.csv -> ','"]


@pytest.mark.parametrize("pattern,name", [
    ("community_package_YYYY_MM.csv", "community_package"),
    ("community_package_*_MM.csv", "community_package"),
    ("x_ind_risk_data_package_amer_MI_YYYYMMDD_HHMM.psv", "x_ind_risk_data_package_amer"),
    ("NWB_COB_RPT_110_*.txt", "NWB_COB_RPT_110"),                    # pair 4: unchanged
    ("I_ACCUM_*_TO_CLIENT_*.csv", "I_ACCUM_TO_CLIENT"),             # pair 1: unchanged
    ("Northwind_MI_Reports_*.csv", "Northwind_MI_Reports"),     # a token INSIDE: kept
    ("YYYYMMDD.csv", "YYYYMMDD"),                                    # nothing but tokens
])
def test_object_name_drops_the_trailing_date_tokens(pattern, name):
    assert _object_name(pattern, ["CCYY", "YYYY", "YY", "MM", "DD", "HH", "MI", "SS"]) == name


@pytest.mark.parametrize("stated,display", [
    ("VENDOR_C – VC 00000", "VENDOR_C"),
    ("Vendor E – VE 12345", "Vendor E"),
    ("Vendor A - 4471", "Vendor A"),
    ("Northwind Benefits", "Northwind Benefits"),
    ("Smith-Jones Health", "Smith-Jones Health"),                   # a hyphenated name
    ("Vendor B – Northeast Region", "Vendor B – Northeast Region"),   # words
])
def test_source_is_the_vendor_display_name(stated, display):
    cell = _source_cell(SimpleNamespace(source_system=stated))
    assert cell["value"] == display
    if stated != display:
        assert repr(stated) in cell["badge_entry"]["tooltip"]


def test_frequency_mentions_cite_every_cadence_sentence():
    content = DocxContent(
        tables=[[["Frequency", "Please refer to the File Details tab."],
                 ["SLA", "File will be received yearly twice. Between Jan and Feb."]]],
        paragraphs=[(0, "The vendor produces the report monthly."), (0, "Unrelated text.")])
    assert frequency_mentions(content) == [
        "'The vendor produces the report monthly.' (FRD paragraph 0)",
        "'File will be received yearly twice.' (FRD table 0 row 1)"]


# ------------------------------------------------------- rule 8: the gate


def test_no_spark_session_is_check_not_run(monkeypatch, tmp_path):
    def no_spark(*_a, **_k):
        return subprocess.CompletedProcess([], 1, stdout="", stderr=(
            "pyspark.errors.exceptions.base.PySparkRuntimeError: [CANNOT_CONFIGURE_SPARK] "
            "Cannot configure Spark."))

    monkeypatch.setattr(tests_runner.subprocess, "run", no_spark)
    check = tests_runner.run_generated_tests(tmp_path, 20)
    assert check.passed and check.not_run and check.flag == "spark_unavailable"
    gate = compute_verdict("feed", [], [], [check], tests_skipped=False)
    assert gate.verdict == "PASS_WITH_FLAGS"
    assert gate.flags[0].startswith("spark_unavailable — a Spark session could not be "
                                    "configured (CANNOT_CONFIGURE_SPARK)")
    # A real test failure is still a FAIL.
    monkeypatch.setattr(tests_runner.subprocess, "run", lambda *_a, **_k:
                        subprocess.CompletedProcess([], 1, stdout="1 failed", stderr=""))
    assert not tests_runner.run_generated_tests(tmp_path, 20).passed


def test_ruff_format_wraps_a_long_line_instead_of_an_e501_fail(tmp_path):
    feed = tmp_path / "feed_x"
    (feed / "pipeline").mkdir(parents=True)
    (feed / "ruff.toml").write_text('line-length = 100\n[lint]\nselect = ["E", "F"]\n',
                                    encoding="utf-8", newline="\n")
    long_call = "VALUE = dict(" + ", ".join(f"key_{i}={i}" for i in range(20)) + ")\n"
    (feed / "pipeline" / "wide.py").write_text('"""Wide."""\n\n' + long_call,
                                               encoding="utf-8", newline="\n")
    check = _ruff_check(feed)
    assert check.passed, check.details
    assert check.formatted == ["pipeline/wide.py"]
    gate = compute_verdict("feed_x", [], [], [check], tests_skipped=False)
    assert gate.flags == ["ruff_formatted: 1 emitted .py file(s) rewritten by `ruff format` "
                          "before the check — pipeline/wide.py"]


# ------------------------------------------------------- rule 9: VDD scope


def test_a_dictionary_of_another_feed_is_one_scope_flag(config, tmp_path):
    from acfc_shapes import vdd as vdd_fixtures
    from codegen.extract.vdd import extract_vdd_contract
    from codegen.gate.vdd_check import vdd_cross_check

    spec = _pair1_spec(config, tmp_path)
    vdd_path = tmp_path / "vdd.xlsx"
    vdd_path.write_bytes(xlsx_bytes(vdd_fixtures.build_vdd_pair1()))
    vdd, _ = extract_vdd_contract(vdd_path, config, generated_date=DATE)
    flags, _check = vdd_cross_check(spec.model_copy(update={"vdd": vdd}), config)
    assert not [f for f in flags if f.startswith("vdd_scope_mismatch")]   # its own VDD
    other = vdd.model_copy(update={"fields": [
        f.model_copy(update={"name": f"OTHER_FEED_{i}"}) for i, f in enumerate(vdd.fields)]})
    flags, _check = vdd_cross_check(spec.model_copy(update={"vdd": other}), config)
    assert len(flags) == 1 and flags[0].startswith("vdd_scope_mismatch — the VDD sheet(s)")
    assert "per-field VDD flags suppressed" in flags[0]


def _pair1_spec(config, tmp_path: Path):
    sttm = tmp_path / "p1.xlsx"
    sttm.write_bytes(xlsx_bytes(sttm_fixtures.build_pair1()))
    frd = SHAPES / "frd" / "f1_pair_1.docx"
    pair = resolve_pair(sttm, frd, config, provider=None, cache_dirs=[], generated_date=DATE)
    frd_json, sttm_json = tmp_path / "frd.json", tmp_path / "sttm.json"
    frd_json.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
    contract = extract_contract(sttm, frd_json, config, generated_date=DATE,
                                layout=pair.sttm.profile)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config)
    return spec


# ------------------------------------------------------- exit codes


def test_the_exit_codes_are_in_codegen_help(capsys):
    with pytest.raises(SystemExit):
        cli.main(["--help"])
    out = capsys.readouterr().out
    assert "exit codes:" in out
    for code, word in ((cli.EXIT_OK, "ok"), (cli.EXIT_FAILED, "failed"),
                       (cli.EXIT_NEEDS_ANSWERS, "NEEDS_ANSWERS")):
        assert f"  {code}  {word}" in out
    assert (cli.EXIT_OK, cli.EXIT_FAILED, cli.EXIT_NEEDS_ANSWERS) == (0, 1, 3)


def test_the_notebook_fallback_reports_needs_answers():
    text = (REPO / "acfc_run.py").read_text(encoding="utf-8")
    assert "EXIT_NEEDS_ANSWERS: \"needs answers\"" in text
    assert 'line.startswith(("QUESTION", "UNRESOLVED"))' in text
