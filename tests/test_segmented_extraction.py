"""Segmented-dialect extraction — corrected semantics (2026-09-01).

Segments are TABLES in both layers; record identification is DERIVED from
the STTM (trailer static marker + positional header), never assumed; both
layers are scoped by Load Strategy (the STD schema comes from the STTM when
the FRD's Target Schema block names stage targets only); no natural key
exists when the FRD states none (truncate/append — no MERGE). Every
non-layer2 review item and provenance note must cite the exact document
evidence it rests on.

Runs on the synthetic CAQH-shaped golden (tests/segmented_fixture.py — zero
client values; scrub-scanned below). Flat-path byte-identity is guarded by
tests/test_framework_output.py's notebook-mode snapshot.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from codegen import cli
from codegen.extract import ExtractionError, extract_contract
from codegen.reasoning.engine import segmented_review_items
from codegen.resolve.resolver import resolve_pair
from segmented_fixture import (
    FAQ_OVERRIDE_ASSUMED,
    FAQ_OVERRIDE_CONFIRMED,
    FEED_SLUG,
    build_workbook,
    frd_contract_dict,
)

REPO = Path(__file__).resolve().parents[1]


@pytest.fixture()
def pair(tmp_path, config):
    """(workbook path, frd path, scoped config, faq dir)."""
    workbook_path = tmp_path / "synthetic_segmented.xlsx"
    build_workbook().save(workbook_path)
    frd_path = tmp_path / "synthetic_frd.contract.json"
    frd_path.write_text(json.dumps(frd_contract_dict()), encoding="utf-8")
    faq_dir = tmp_path / "faq"
    faq_dir.mkdir()
    scoped = config.model_copy(update={
        "load_pattern_faq": config.load_pattern_faq.model_copy(
            update={"per_feed_dir": str(faq_dir)}),
    })
    return workbook_path, frd_path, scoped, faq_dir


def _extract(pair, workbook=None):
    workbook_path, frd_path, scoped, _ = pair
    return extract_contract(workbook or workbook_path, frd_path, scoped,
                            generated_date="2026-01-01")


# -- extraction: segments are tables, both layers, derived identification ----- #


def test_segments_extract_as_tables_in_both_layers(pair):
    contract = _extract(pair)
    (feed,) = contract.feeds
    seg = feed.segmented
    assert seg.segments_found == ["Header", "Detail", "Trailer"]
    assert seg.row_counts == {"Header": 2, "Detail": 3, "Trailer": 2}
    by_segment = {}
    for field in feed.fields:
        by_segment.setdefault(field.record_segment,
                              (field.stage_table, field.standard_table))
    assert by_segment == {
        "Header": ("EXT_SYN_HDR", "EXT_SYN_HDR"),
        "Detail": ("EXT_SYN_DTL", "EXT_SYN_DTL"),
        "Trailer": ("EXT_SYN_TRL", "EXT_SYN_TRL"),
    }
    # STD schema sourced from the STTM's standard target group, with a cited
    # provenance note quoting Load Strategy STD.
    assert feed.standard is not None
    assert (feed.standard.catalog, feed.standard.schema_name) == ("SYN_STD", "SYN")
    std_note = next(n for n in seg.provenance_notes
                    if "STD target schema sourced from STTM" in n.note)
    assert "Load Strategy STD: 'Append'" in std_note.citation


def test_identification_is_derived_with_sttm_citation(pair):
    contract = _extract(pair)
    ident = contract.feeds[0].segmented.identification
    assert ident.method == "derived_from_sttm"
    assert ident.trailer_marker == "******"
    assert ident.header_rule == "first record of the file"
    assert "Contains the value ******" in ident.citation
    assert "'Record Type'" in ident.citation


def test_no_natural_key_is_frd_consistent_not_an_unknown(pair):
    contract = _extract(pair)
    (feed,) = contract.feeds
    # No refusal, no assumption: empty keys are consistent with the FRD.
    assert feed.load_rules.not_null_columns == []
    note = next(n for n in feed.segmented.provenance_notes
                if "No MERGE key" in n.note)
    assert "Truncate and Load" in note.citation and "'Append'" in note.citation


def test_workbook_mandatory_signal_still_wins(pair, tmp_path):
    workbook_path = tmp_path / "with_mandatory.xlsx"
    build_workbook(mandatory_detail_field=True).save(workbook_path)
    contract = _extract(pair, workbook=workbook_path)
    assert contract.feeds[0].load_rules.not_null_columns == ["Member ID"]


def test_recycle_built_from_frd_dq_with_citation(pair):
    contract = _extract(pair)
    (feed,) = contract.feeds
    recycle = feed.load_rules.recycle
    assert recycle is not None
    assert recycle.applies_to == "Member ID"
    assert recycle.recycle_window_days == 15
    assert "facets_member" in recycle.validation
    note = next(n for n in feed.segmented.provenance_notes
                if "existence checked against" in n.note)
    assert "Member ID validation" in note.citation
    assert "retained for 15 days" in note.citation


def test_missing_trailer_marker_refuses_with_override_remedy(pair, tmp_path):
    workbook = build_workbook()
    ws = workbook.active
    # Blank the trailer comment the derivation rests on (row 19 = Trailer 1).
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "Record Type":
                ws.cell(row=cell.row, column=10).value = "No marker stated here"
    path = tmp_path / "no_marker.xlsx"
    workbook.save(path)
    with pytest.raises(Exception) as excinfo:
        _extract(pair, workbook=path)
    message = str(excinfo.value)
    assert "does not state a static marker value" in message
    assert "record_type_discriminators override" in message


# -- FAQ override: confirmed only --------------------------------------------- #


def _write_faq(faq_dir: Path, text: str) -> None:
    (faq_dir / f"{FEED_SLUG}.faq.yaml").write_text(text, encoding="utf-8")


def test_faq_override_requires_confirmed_status(pair):
    _write_faq(pair[3], FAQ_OVERRIDE_ASSUMED)
    with pytest.raises(ExtractionError) as excinfo:
        _extract(pair)
    assert "requires status 'confirmed'" in str(excinfo.value)


def test_confirmed_faq_override_applies(pair):
    _write_faq(pair[3], FAQ_OVERRIDE_CONFIRMED)
    contract = _extract(pair)
    ident = contract.feeds[0].segmented.identification
    assert ident.method == "declared_override"
    assert ident.trailer_marker == "T"
    assert "status: confirmed" in ident.citation


# -- resolution + review layer ------------------------------------------------ #


def _resolve(pair):
    contract = _extract(pair)
    from codegen.extract import contract_to_json

    sttm_path = pair[0].parent / "extracted.json"
    sttm_path.write_text(contract_to_json(contract), encoding="utf-8")
    (spec,) = resolve_pair(pair[1], sttm_path, pair[2])
    return spec


def test_resolves_three_segments_with_standard_tables(pair):
    spec = _resolve(pair)
    assert [s.segment for s in spec.segments] == ["Header", "Detail", "Trailer"]
    assert spec.is_segmented is True
    for segment in spec.segments:
        assert segment.standard_table is not None
        assert segment.standard_table.catalog == "SYN_STD"
        assert segment.standard_table.table == segment.stage_table.table
    assert spec.natural_key_columns == []
    assert spec.load_as_is is True  # FRD AS-IS rule drives the switch
    assert spec.recycle is not None
    assert spec.recycle.recycle_table.table == "EXT_SYN_DTL_RECYCLE"


def test_review_layer_is_one_cited_confirm_item(pair):
    spec = _resolve(pair)
    items = segmented_review_items(spec)
    assert [i.kind for i in items] == ["confirm"]
    (item,) = items
    assert "positional header/detail identification" in item.rule_text
    assert "'******'" in item.rule_text
    # citation guard: the item quotes the exact STTM cell it rests on.
    assert item.citation and "Contains the value ******" in item.citation
    # No assumption gates and no conflict cards exist any more.
    assert not any("ASSUMPTION" in i.rule_text for i in items)
    assert not any("CONFLICT" in i.rule_text for i in items)


def test_every_provenance_note_carries_a_citation(pair):
    spec = _resolve(pair)
    notes = spec.segmented_extraction.provenance_notes
    assert notes
    for note in notes:
        assert note.note.strip()
        assert note.citation.strip(), f"note without citation: {note.note!r}"


# -- generation --------------------------------------------------------------- #


@pytest.fixture()
def generated(pair, tmp_path):
    spec = _resolve(pair)
    scoped = pair[2].model_copy(update={
        "output": pair[2].output.model_copy(update={
            "dir": str(tmp_path / "out"),
            "reports_dir": str(tmp_path / "reports"),
        })
    })
    gate = cli._generate_feed(spec, scoped, dry_run=True, skip_tests=True,
                              output_mode="both")
    return spec, gate, tmp_path


def test_generation_emits_per_segment_tables_both_layers(generated):
    spec, gate, tmp = generated
    assert gate.verdict == "PASS_WITH_FLAGS"
    framework = tmp / "out" / spec.feed_slug / "framework"
    stage_txt = (framework / f"{spec.feed_slug}_stage_table_creation.txt"
                 ).read_text(encoding="utf-8")
    standard_txt = (framework / f"{spec.feed_slug}_standard_table_creation.txt"
                    ).read_text(encoding="utf-8")
    for table in ("EXT_SYN_HDR", "EXT_SYN_DTL", "EXT_SYN_TRL",
                  "EXT_SYN_DTL_RECYCLE"):
        assert f"CREATE OR REPLACE TABLE SYN_DLK.STG_SYN.{table}" in stage_txt
    for table in ("EXT_SYN_HDR", "EXT_SYN_DTL", "EXT_SYN_TRL"):
        assert f"CREATE OR REPLACE TABLE SYN_STD.SYN.{table}" in standard_txt
    assert "EXT_SYN_DTL_RECYCLE" not in standard_txt
    # STRING everywhere except audit date/timestamp fields — in BOTH layers
    # (FRD acceptance criterion 3), asserted on the segment/recycle target
    # tables (the agent's errors/processed-files bookkeeping tables carry
    # their own typed columns by design).
    target_tables = ("EXT_SYN_HDR", "EXT_SYN_DTL", "EXT_SYN_TRL",
                     "EXT_SYN_DTL_RECYCLE")
    for text in (stage_txt, standard_txt):
        section = None
        for line in text.splitlines():
            stripped = line.strip()
            if stripped.startswith("--") and any(
                    stripped.endswith(t) for t in target_tables):
                section = stripped
                continue
            if stripped.startswith("--") and stripped.startswith("--SYN"):
                section = None
                continue
            if section is None or not stripped or stripped.startswith(
                    ("--", "CREATE", "USING", "COMMENT", "LOCATION",
                     "TBLPROPERTIES", "'", "CLUSTER", "ALTER", ")")):
                continue
            if "REC_CREATION_TIME" in stripped or "REC_UPDATED_TIME" in stripped:
                assert "TIMESTAMP" in stripped, stripped
            elif " STRING" not in stripped and "TIMESTAMP" not in stripped:
                raise AssertionError(
                    f"non-STRING business column in {section}: {stripped}")


def test_generation_review_and_report(generated):
    spec, gate, tmp = generated
    # One cited confirm item + Layer-2 candidates from unmapped rules.
    candidates = json.loads(
        (tmp / "out" / spec.feed_slug / "candidates" / "candidates.json"
         ).read_text(encoding="utf-8"))
    kinds = [c.get("kind", "layer2") for c in candidates]
    assert kinds.count("confirm") == 1
    assert "extraction_assumption" not in kinds
    assert "escalated_conflict" not in kinds
    confirm = next(c for c in candidates if c.get("kind") == "confirm")
    assert confirm["citation"]

    report = (tmp / "reports" / f"{spec.feed_slug}.md").read_text(encoding="utf-8")
    assert "## Segmented extraction" in report
    assert "### Segment tables (both layers)" in report
    assert "### Record identification (derived_from_sttm)" in report
    assert "Evidence: >" in report
    assert "Held back" not in report
    assert "ASSUMED" not in report

    # Framework rows exist for both legs of every segment table.
    from openpyxl import load_workbook

    rows_wb = load_workbook(tmp / "out" / spec.feed_slug / "framework" /
                            "config_rows.xlsx")
    assert rows_wb["ADLS_DELTA_INGESTION_DETAILS"].max_row - 1 == 3
    assert rows_wb["STGDELTA_STDDELTA_INGESTION_DET"].max_row - 1 == 3
    dq_rows = list(rows_wb["DATA_QUALITY_RULES"].iter_rows(min_row=2,
                                                           values_only=True))
    assert any("ReferentialCheckRule" in [str(v) for v in row]
               for row in dq_rows)


def test_synthetic_fixture_passes_the_scrub_denylist(pair):
    import sys

    sys.path.insert(0, str(REPO / "scripts"))
    from scrub_check import scan_paths

    hits = scan_paths([pair[0], pair[1],
                       REPO / "fixtures" / "workbooks" /
                       "synthetic_segmented_golden.xlsx"])
    assert hits == [], hits
