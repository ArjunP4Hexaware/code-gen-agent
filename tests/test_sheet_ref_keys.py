"""post57 (D1 post55 leak report): every one of the run's 10 leak hits was a
sheet name inside the needed-answer keys the harness reads from stdout
(``<sheet>/source/field_name``, ``<sheet>/band[1]/layer``). A printed key now
names the sheet by its workbook position, ``sheets[<i>]`` (0-based, like
``feeds[<i>]``); the name stays in the ACFC-side answers template
(``sheet_name:``) and ``unresolved_headers.md``. An answers file may give the
sheet either way."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from acfc_shapes import bands
from acfc_shapes.common import xlsx_bytes
from codegen import cli
from codegen.contracts.sttm import PendingAnswer
from codegen.layout.sheet_refs import (
    private_key,
    public_key,
    public_note,
    public_role_key,
    sheet_name,
)
from test_harness_answer_lines import _NEEDED_KEY_RE

SENTINEL = "ZZ Secret Mapping"
NAMES = ["Version", SENTINEL]
PRINTED = [f"sheets[1]/band[{n}]/layer" for n in (1, 2, 3)]


def test_public_and_private_keys():
    assert public_key(f"{SENTINEL}/source/field_name", NAMES) == "sheets[1]/source/field_name"
    assert public_key(f"{SENTINEL}/band[1]/layer", NAMES) == "sheets[1]/band[1]/layer"
    assert public_key("feeds[0].file_format", NAMES) == "feeds[0].file_format"
    assert public_key("sheets[1]/stage/table", NAMES) == "sheets[1]/stage/table"
    # a sheet name with a slash: the longest listed name wins
    assert public_key("a/b/stage/table", ["a", "a/b"]) == "sheets[1]/stage/table"
    # a role key whose sheet is not listed never prints the name
    assert public_role_key("Elsewhere/stage/table", NAMES) == "sheets[?]/stage/table"
    assert private_key("sheets[1]/stage/table", NAMES) == f"{SENTINEL}/stage/table"
    assert private_key("sheets[9]/stage/table", NAMES) == "sheets[9]/stage/table"
    assert sheet_name("sheets[0]", NAMES) == "Version"
    assert sheet_name(SENTINEL, NAMES) == SENTINEL
    assert sheet_name("sheets[5]", NAMES) is None
    assert public_note(f"answers[0] ({SENTINEL}/stage/table) confirms column 3", NAMES) \
        == "answers[0] (sheets[1]/stage/table) confirms column 3"


def _workbook(tmp_path: Path) -> Path:
    """The three-look-alike-band sheet (every band's layer is a question),
    renamed to the sentinel, second in the workbook."""
    path = tmp_path / "STTM_sheet_refs.xlsx"
    path.write_bytes(xlsx_bytes(bands.three_lookalike()))
    book = load_workbook(path)
    book[bands.SHEET].title = SENTINEL
    book.create_sheet("Version", 0)["A1"] = "Version"
    book.save(path)
    return path


@pytest.fixture()
def cli_env(monkeypatch, tmp_path):
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_PROVIDER", "1")
    for role in ("state", "outputs"):
        (tmp_path / role).mkdir()
        monkeypatch.setenv(f"CODEGEN_STORAGE_{role.upper()}",
                           f"local:{(tmp_path / role).as_posix()}")
    return tmp_path


def _needed(out: str) -> list[str]:
    return [m.group(1) for ln in out.splitlines() if (m := _NEEDED_KEY_RE.match(ln))]


def test_the_console_never_names_the_sheet_and_the_template_answers_by_position(
        cli_env, capsys):
    tmp_path = cli_env
    workbook = _workbook(tmp_path)
    template = tmp_path / "answers_template.yaml"
    report = tmp_path / "unresolved_headers.md"
    capsys.readouterr()
    assert cli.main(["layout", "--workbook", str(workbook), "--dry-run", "--no-cache",
                     "--require-complete", "--answers-template", str(template),
                     "--report-unresolved", str(report)]) == cli.EXIT_NEEDS_ANSWERS
    out = capsys.readouterr().out
    assert _needed(out) == PRINTED
    assert SENTINEL not in out
    # ACFC-side files keep the name: the template's reading aid, the report
    entries = yaml.safe_load(template.read_text(encoding="utf-8"))["answers"]
    assert [(e["sheet"], e["sheet_name"]) for e in entries] == [("sheets[1]", SENTINEL)] * 3
    assert SENTINEL in report.read_text(encoding="utf-8")
    assert "`sheets[1]/band[1]/layer`" in report.read_text(encoding="utf-8")

    # extract-sttm holds the sheet back: same keys, no name on any line
    frd = tmp_path / "frd.contract.json"
    contract = bands.frd_contract(stage_schema="nb_land")
    contract["feeds"][0]["stage_target"]["catalog"] = "nb_dlk"
    contract["feeds"][0]["standard_target"]["catalog"] = "nb_std"
    frd.write_text(json.dumps(contract), encoding="utf-8")
    sttm_out = tmp_path / "sttm.contract.json"
    assert cli.main(["extract-sttm", "--workbook", str(workbook), "--frd-contract", str(frd),
                     "--out", str(sttm_out), "--generated-date", "2026-10-15"]) \
        == cli.EXIT_NEEDS_ANSWERS
    out = capsys.readouterr().out
    assert _needed(out) == PRINTED and SENTINEL not in out

    # The template filled as written (sheets[1]) settles every question.
    body = template.read_text(encoding="utf-8")
    for layer in ("source", "stage", "standard"):
        body = body.replace('value: ""', f"value: {layer}", 1)
    template.write_text(body, encoding="utf-8")
    assert cli.main(["layout", "--workbook", str(workbook), "--dry-run", "--no-cache",
                     "--require-complete", "--answers", str(template)]) == cli.EXIT_OK
    out = capsys.readouterr().out
    assert _needed(out) == [] and SENTINEL not in out
    assert cli.main(["extract-sttm", "--workbook", str(workbook), "--frd-contract", str(frd),
                     "--out", str(sttm_out), "--generated-date", "2026-10-15",
                     "--answers", str(template)]) == cli.EXIT_OK
    assert json.loads(sttm_out.read_text(encoding="utf-8"))["feeds"][0]["field_count"] \
        == len(bands.FIELDS)


def test_an_unknown_sheet_position_is_reported_never_applied(cli_env, capsys):
    workbook = _workbook(cli_env)
    answers = cli_env / "answers.yaml"
    answers.write_text(yaml.safe_dump({"answers": [
        {"sheet": "sheets[7]", "layer": "band[1]", "role": "layer", "value": "source"}]}),
        encoding="utf-8")
    capsys.readouterr()
    assert cli.main(["layout", "--workbook", str(workbook), "--dry-run", "--no-cache",
                     "--require-complete", "--answers", str(answers)]) \
        == cli.EXIT_NEEDS_ANSWERS
    out = capsys.readouterr().out
    assert "sheets[<i>]) this sttm workbook does not have" in out
    assert _needed(out) == PRINTED


def test_a_held_back_answer_read_from_a_contract_prints_its_position(capsys):
    """generate reads held-back answers from the STTM contract (no workbook at
    hand): the contract carries each sheet's position."""
    pending = PendingAnswer(feed_name=SENTINEL, sheet=SENTINEL, sheet_index=1,
                            key=f"{SENTINEL}/band[2]/layer", reason="no evidence names it")
    again = PendingAnswer.model_validate_json(pending.model_dump_json())
    assert again.sheet_index == 1
    cli._print_needs_answers([again])
    out = capsys.readouterr().out
    assert _needed(out) == ["sheets[1]/band[2]/layer"]
    assert SENTINEL not in out and "feed of sheets[1], sheet sheets[1]" in out
