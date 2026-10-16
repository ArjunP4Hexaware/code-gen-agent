"""post57 (D1 post55 pair_01): the path_whitespace flags came from ONE FRD cell
copied with its label — ``Path : <account>/<container>/<folder>`` — into every
landing / archive / reject path. Two fixes (codegen.resolve.stated), the
whitespace check unchanged:

(a) a cell ``<Label> : <value>`` whose label is a synonym of the field being read
    (the conventions profile's ``field_label_synonyms``) is read as ``<value>``,
    flag ``label_stripped`` with the FRD coordinates; never for a label that is
    not a synonym of that field;
(b) a landing path that starts with a storage account and a container takes both
    from the run's environment profile; a document account / container that
    differs raises ``environment_mismatch`` ("document states another
    environment", never the document's value); the folder stays byte for byte.

The synthetic FRD states a prod-style labelled path; the D1 run (the suite's D1
profile over the ACFC overlay) writes the D1 address with one label_stripped
flag, one environment_mismatch flag and no path_whitespace flag."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from codegen import cli
from codegen.config import load_config
from codegen.contracts.frd import FrdContract
from codegen.extract.extractor import contract_to_json, extract_contract
from codegen.resolve.stated import environment_address, strip_labels
from post51_pairs import family_e_frd, family_e_workbook

REPO = Path(__file__).resolve().parents[1]
CONFIG = REPO / "config" / "config.yaml"
ACFC_OVERLAY = REPO / "config" / "overlays" / "acfc_env.yaml"
FEED = "feed_5_claims"
PROD_ACCOUNT = "synprodacct01"
PROD_CONTAINER = "synprodlanding"
FOLDER = "inbound/vendor_e/claims/"
LABELLED = f"Path : {PROD_ACCOUNT}/{PROD_CONTAINER}/{FOLDER}"


def _frd(path: Path, landing: str) -> Path:
    frd = family_e_frd(path)
    data = json.loads(frd.read_text(encoding="utf-8"))
    data["feeds"][0]["landing_location"] = landing
    data.setdefault("field_provenance", {})["feeds[0].landing_location"] = {
        "table": 2, "row": 6, "col": 2, "label": "ADLS Location", "source": "synonyms"}
    frd.write_text(json.dumps(data), encoding="utf-8")
    return frd


def _feed(tmp_path: Path, landing: str):
    frd = FrdContract.model_validate_json(_frd(tmp_path / "frd.json", landing)
                                          .read_text(encoding="utf-8"))
    return frd, frd.feeds[0]


@pytest.mark.parametrize(("landing", "value", "flagged"), [
    (LABELLED, f"{PROD_ACCOUNT}/{PROD_CONTAINER}/{FOLDER}", True),
    ("ADLS Location: mftlanding/inbound/x/", "mftlanding/inbound/x/", True),
    ("Frequency : Daily", "Frequency : Daily", False),       # not a synonym of the landing
    ("abfss://c@a.dfs.example/x", "abfss://c@a.dfs.example/x", False),
    ("mftlanding/inbound/x/", "mftlanding/inbound/x/", False),
])
def test_a_label_is_stripped_only_when_it_names_the_field(tmp_path, landing, value, flagged):
    config = load_config(CONFIG)
    frd, feed = _feed(tmp_path, landing)
    flags: list[str] = []
    feed = strip_labels(feed, 0, frd, config.conventions.get(None), flags)
    assert feed.landing_location == value
    assert flags == ([f"label_stripped:feeds[0].landing_location — FRD table 2 row 6: the "
                      f"cell's '{'path' if landing == LABELLED else 'adls location'} :' label "
                      "(a synonym of landing_location) taken off its value"] if flagged else [])


def test_the_address_comes_from_the_environment(tmp_path):
    config = load_config(CONFIG, overlays=[ACFC_OVERLAY])
    profile = config.conventions.get(None)
    frd, feed = _feed(tmp_path, f"{PROD_ACCOUNT}/{PROD_CONTAINER}/{FOLDER}")
    flags: list[str] = []
    moved = environment_address(feed, 0, config, profile, flags)
    assert moved.landing_location == f"mftlanding/{FOLDER}"        # D1's container, doc folder
    assert flags == ["environment_mismatch:feeds[0].landing_location — document states another "
                     "environment (storage account, container); the address is environment "
                     "D1's, the folder under the container the document's"]
    assert PROD_ACCOUNT not in flags[0] and PROD_CONTAINER not in flags[0]
    # a container-relative landing (no account) is the document's, untouched
    _, same = _feed(tmp_path, "mftlanding/inbound/x/")
    assert environment_address(same, 0, config, profile, []) == same
    # a URI's account / container from the environment too
    _, uri = _feed(tmp_path, f"abfss://{PROD_CONTAINER}@{PROD_ACCOUNT}.dfs.core.windows.net/x/")
    flags = []
    assert environment_address(uri, 0, config, profile, flags).landing_location == \
        "abfss://mftlanding@synd1teststorage.dfs.core.windows.net/x/"
    assert len(flags) == 1


@pytest.fixture()
def run(monkeypatch, tmp_path, capsys):
    monkeypatch.setenv("CODEGEN_CONFIG_OVERLAYS", str(ACFC_OVERLAY))
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_PROVIDER", "1")
    monkeypatch.setenv("CODEGEN_NOTIFICATION_EMAILS", "syn.dl.prodsupport@synthetic.example")
    (tmp_path / "out").mkdir()
    monkeypatch.setenv("CODEGEN_STORAGE_OUTPUTS", f"local:{(tmp_path / 'out').as_posix()}")
    monkeypatch.chdir(REPO)
    frd = _frd(tmp_path / "frd.json", LABELLED)
    sttm = tmp_path / "sttm.json"
    sttm.write_text(contract_to_json(extract_contract(
        family_e_workbook(tmp_path / "e.xlsx", [None]), frd,
        load_config(CONFIG, overlays=[ACFC_OVERLAY]), generated_date="2026-10-15")),
        encoding="utf-8")
    capsys.readouterr()
    code = cli.main(["generate", "--frd-contract", str(frd), "--sttm-contract", str(sttm),
                     "--skip-tests", "--dry-run", "--output-mode", "framework",
                     "--profile", "acfc_prx", "--iig-template", "iig_v2"])
    return code, capsys.readouterr().out, tmp_path / "out"


def test_the_d1_run_writes_the_d1_address_with_two_flags_and_no_whitespace_flag(run):
    code, out, out_root = run
    assert code == cli.EXIT_OK, out
    book = load_workbook(next(out_root.rglob("*_IIG.xlsx")), read_only=True)
    try:
        ws = book["ADLS_DELTA_INGESTION_DETAILS"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        book.close()
    head = list(rows[0])
    row = dict(zip(head, rows[1], strict=False))
    assert row["SRC_CONTAINER_NAME"] == "mftlanding"                     # D1's container
    assert row["SRC_ADLS_PATH"] == f"/{FOLDER}"                          # the document's folder
    assert row["SRC_ADLS_ARCHVL_PATH"] == f"/{FOLDER}Archive/"
    for path in out_root.rglob("*"):
        if path.is_file() and path.suffix in (".txt", ".sql", ".md", ".json", ".xlsx"):
            data = path.read_bytes()
            assert PROD_ACCOUNT.encode() not in data and b"Path :" not in data, path
    report = next((out_root / "reports").glob(f"{FEED}.md")).read_text(encoding="utf-8")
    flags = [ln for ln in report.splitlines() if ln.startswith("- ")]
    assert sum(ln.startswith("- label_stripped:") for ln in flags) == 1
    assert sum(ln.startswith("- environment_mismatch:") for ln in flags) == 1
    assert not any(ln.startswith(("- iig_path_whitespace:", "- path_whitespace"))
                   for ln in flags)
