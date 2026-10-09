"""Chunk D (owner brief 2026-10-09): the STGDELTA_STDDELTA_INGESTION_DET row
of the SD family, as the committed ACFC environment overlay
(``config/overlays/acfc_env.yaml``) writes it.

The first real ACFC run wrote the stage -> standard row's paths in the
LANDING shape (``/<landing container>/inbound/<domain>/…``) and left its
OBJECT_ID open. Under the overlay now:

a. SRC_ADLS_PATH = the stage side's Processed folder — the folder of the
   ADLS row's TGT_ADLS_PATH, no container segment (the container is its own
   column: SRC_CONTAINER_NAME stays OPEN — no document or real row states
   it; the stage container, the ADLS row's TGT_CONTAINER_NAME, is only an
   inferred candidate, offered in the open cell's tooltip — review D1);
b. TGT_ADLS_PATH / TGT_RJT_ADLS_PATH in the standard container, the same
   container-relative shape (``Processed/<table>``, ``Reject/<table>_reject``);
c. OBJECT_ID is per (group, file) — family convention ``stgdelta_object_id:
   from_file``: the ADLS row's OBJECT_ID of THE file that feeds the table;
   several files feed it -> open, with a note (``stgdelta_object_id_open``).

Pinned on SYNTHETIC inputs only: pair 11 (the SD-shaped synthetic feeds,
one file each) and pair 4 (six LOB files into three tables). The shipped
default (``stgdelta_object_id: open``) keeps the pair-1 / pair-4 goldens.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from acfc_shapes import sttm as sttm_fixtures
from codegen.config import load_config
from codegen.metadata_sheet import metadata_sheet_payload
from codegen.metadata_template import shape_flags
from test_sd_scorecard_rules import RUN_STATEMENT, _generate

REPO = Path(__file__).resolve().parents[1]
ACFC_ENV = REPO / "config" / "overlays" / "acfc_env.yaml"
ADLS, STG = "ADLS_DELTA_INGESTION_DETAILS", "STGDELTA_STDDELTA_INGESTION_DET"
STAGE_CONTAINER = "z-use-d1-dlk-stage-01"          # the overlay's ADLS TGT_CONTAINER_NAME


def _sheet(book, name: str) -> list[dict]:
    rows = list(book[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]


def _tooltips(book, name: str) -> dict[int, dict[str, tuple[str, str]]]:
    out: dict[int, dict] = {}
    for tab, row, header, badge, tooltip in book["_provenance"].iter_rows(min_row=2,
                                                                         values_only=True):
        if tab == name:
            out.setdefault(row, {})[header] = (badge, tooltip or "")
    return out


@pytest.fixture(scope="module")
def sd_run(tmp_path_factory):
    """Pair 11's three feeds generated under the committed ACFC overlay."""
    config = load_config(REPO / "config" / "config.yaml", overlays=[ACFC_ENV])
    tmp = tmp_path_factory.mktemp("stgdelta_sd")
    gates = _generate(config, sttm_fixtures.build_pair11(), tmp, run_statement=RUN_STATEMENT)
    books = {slug: load_workbook(tmp / "out" / slug / "framework" / "config_rows.xlsx")
             for slug in gates}
    return config, {slug: (gate, books[slug]) for slug, (gate, _rows) in gates.items()}


def test_the_overlay_selects_the_sd_family_stage_to_standard_settings(sd_run):
    config, _runs = sd_run
    template = config.metadata.templates["iig_v2"]
    assert template.family_conventions.stgdelta_object_id == "from_file"
    # No STGDELTA constant: the source container is an inference, offered only.
    assert "SRC_CONTAINER_NAME" not in template.constants.get(STG, {})
    assert "SRC_CONTAINER_NAME" not in template.constant_citations.get(STG, {})
    offer = template.open_offers[STG]["SRC_CONTAINER_NAME"]
    assert STAGE_CONTAINER in offer and "Friday checklist 13" in offer
    assert "config/overlays/acfc_env.yaml" in offer
    assert template.path_patterns[STG]["SRC_ADLS_PATH"] == "/{domain_path}Processed/"
    for column in ("SRC_ADLS_PATH", "TGT_ADLS_PATH", "TGT_RJT_ADLS_PATH"):
        assert "config/overlays/acfc_env.yaml" in template.path_citations[STG][column]


def test_stage_to_standard_paths_are_container_relative(sd_run):
    _config, runs = sd_run
    assert len(runs) == 3
    for slug, (gate, book) in runs.items():
        (adls,) = _sheet(book, ADLS)
        (stg,) = _sheet(book, STG)
        landing_container = adls["SRC_CONTAINER_NAME"]
        processed = adls["TGT_ADLS_PATH"].rsplit("/", 1)[0] + "/"   # the stage Processed folder
        assert processed.endswith("/Processed/")
        # a. the source = the stage side's Processed folder, no container segment
        assert stg["SRC_ADLS_PATH"] == processed, slug
        assert adls["TGT_CONTAINER_NAME"] == STAGE_CONTAINER
        # b. target and reject: the standard container, the same relative shape
        assert stg["TGT_ADLS_PATH"] == f"{processed}{stg['TGT_TABLE_NAME']}", slug
        assert stg["TGT_RJT_ADLS_PATH"] == (
            processed.replace("/Processed/", "/Reject/") + stg["TGT_RJT_TABLE_NAME"]), slug
        for column in ("SRC_ADLS_PATH", "TGT_ADLS_PATH", "TGT_RJT_ADLS_PATH"):
            segments = stg[column].strip("/").split("/")
            assert landing_container not in segments and "inbound" not in segments, column
        tips = _tooltips(book, STG)[2]
        for column in ("SRC_ADLS_PATH", "TGT_ADLS_PATH", "TGT_RJT_ADLS_PATH"):
            badge, tooltip = tips[column]
            assert badge == "SYNTHETIC" and "config/overlays/acfc_env.yaml" in tooltip, column
        assert gate.verdict == "PASS_WITH_FLAGS"


def test_no_document_states_this_sheets_containers_or_connections(sd_run):
    _config, runs = sd_run
    for _slug, (gate, book) in runs.items():
        (stg,) = _sheet(book, STG)
        for column in ("SRC_CONTAINER_NAME", "TGT_CONTAINER_NAME", "SRC_ADLS_CONNECTION_ID",
                       "METADATA_CONNECTION_ID", "TGT_CONNECTION_ID"):
            assert stg[column] in (None, ""), column          # open: Friday checklist 2 / 13
        # the source container's candidate (an inference) rides in the open cell's tooltip
        badge, tooltip = _tooltips(book, STG)[2]["SRC_CONTAINER_NAME"]
        assert badge == "NEEDS CLIENT TEMPLATE"
        assert tooltip.startswith("open — no input states it; candidate, NOT written")
        assert STAGE_CONTAINER in tooltip and "Friday checklist 13" in tooltip
        (blank,) = [f for f in gate.flags if f.startswith(f"iig_blank:{STG}: ")]
        assert "SRC_CONTAINER_NAME" in blank.split(": ", 1)[1].split(", ")


def test_an_offered_cell_can_never_also_be_a_constant():
    from pydantic import ValidationError

    from codegen.config import MetadataTemplateConfig

    with pytest.raises(ValidationError, match="SRC_CONTAINER_NAME also a constant"):
        MetadataTemplateConfig(tabs={}, constants={STG: {"SRC_CONTAINER_NAME": "x"}},
                               open_offers={STG: {"SRC_CONTAINER_NAME": "x, by inference"}})


def test_one_file_one_table_carries_the_files_adls_object_id(sd_run):
    _config, runs = sd_run
    for slug, (gate, book) in runs.items():
        (adls,) = _sheet(book, ADLS)
        (stg,) = _sheet(book, STG)
        assert stg["OBJECT_ID"] == adls["OBJECT_ID"] == "1", slug
        badge, tooltip = _tooltips(book, STG)[2]["OBJECT_ID"]
        assert badge == "SYNTHETIC"
        assert "per (group, file)" in tooltip and "stgdelta_object_id: from_file" in tooltip
        assert not [f for f in gate.flags if f.startswith("stgdelta_object_id_open")]


# -- pair 4: six LOB files into three tables ---------------------------------------- #


def _pair4_payload(config, spec, object_id: str):
    from acfc_shapes import FIXTURE_ROOT, pair4_nb

    template = config.metadata.templates["iig_v2"]
    template = template.model_copy(update={"family_conventions": template.family_conventions
                                           .model_copy(update={"stgdelta_object_id": object_id})})
    metadata = config.metadata.model_copy(update={
        "templates": {**config.metadata.templates, "iig_v2": template}})
    return metadata_sheet_payload(config.model_copy(update={"metadata": metadata}), REPO,
                                  specs=[spec], frd_path=FIXTURE_ROOT / pair4_nb.FRD_PATH,
                                  template="iig_v2", conventions_profile="acfc_prx")


def _cells(payload, sheet: str, column: str) -> list:
    return [r["values"][column] for r in payload["tabs"][sheet]["rows"]]


def test_the_shipped_default_leaves_the_stgdelta_object_id_open(pair4_config, pair4_payload):
    assert pair4_config.metadata.templates["iig_v2"].family_conventions \
        .stgdelta_object_id == "open"
    assert _cells(pair4_payload, STG, "OBJECT_ID") == ["", "", ""]
    assert not [f for f in shape_flags(pair4_payload) if "stgdelta_object_id" in f]


def test_several_files_feeding_a_table_leave_it_open_with_a_note(pair4_config, pair4_spec,
                                                                 pair4_payload):
    payload = _pair4_payload(pair4_config, pair4_spec, "from_file")
    assert _cells(payload, ADLS, "OBJECT_ID") == [str(n) for n in range(1, 7)]   # 1..n stays
    assert _cells(payload, STG, "OBJECT_ID") == ["", "", ""]
    notes = [f for f in shape_flags(payload) if f.startswith("stgdelta_object_id_open:")]
    tables = [r["values"]["SRC_TABLE_NAME"] for r in payload["tabs"][STG]["rows"]]
    assert [n.split(" — ")[0] for n in notes] == [f"stgdelta_object_id_open:{t}" for t in tables]
    assert all("6 files feed it" in n for n in notes)
    for row in payload["tabs"][STG]["rows"]:
        entry = row["badges"]["OBJECT_ID"]
        assert entry["badge"] == "needs_template" and "per (group, file)" in entry["tooltip"]
    # nothing else moved
    for sheet, tab in pair4_payload["tabs"].items():
        for before, after in zip(tab["rows"], payload["tabs"][sheet]["rows"], strict=True):
            assert before["values"] == after["values"], sheet


def test_one_file_feeding_three_tables_gives_each_row_that_files_object_id(pair4_config,
                                                                           pair4_spec):
    one = pair4_spec.model_copy(update={"files": pair4_spec.files[:1]})
    payload = _pair4_payload(pair4_config, one, "from_file")
    assert _cells(payload, ADLS, "OBJECT_ID") == ["1"]
    assert _cells(payload, STG, "OBJECT_ID") == ["1", "1", "1"]
    assert not [f for f in shape_flags(payload) if "stgdelta_object_id" in f]
