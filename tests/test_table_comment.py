"""post140 (Soham, 2026-10-28): the DDL table COMMENT is "description + load
strategy" - the STTM's Table Description for the layer's table (the target
band's ``table_description`` role), then the load-strategy text, joined by the
profile's column separator (" — ").

* stated once per table (D1 pair_11's billing sheet states it on one row; the
  synthetic Family E sheet on every row): that text opens the COMMENT;
* a table whose rows state two different texts: none - never a silent pick -
  and a ``table_description_differs`` flag;
* a band that states none: the load strategy alone, one
  ``ddl_table_comment_no_description`` flag per feed.
The pair-1 overlay switches the knob off (its golden: tests/test_m4_acceptance.py).
"""

from __future__ import annotations

import json
import re

import openpyxl

from post51_pairs import family_e_frd, family_e_workbook
from test_iig_cell_counts import _run

# Family E: source band 11 columns, then stage / standard 11 each; Table
# Description is the 10th of a target band (openpyxl columns, 1-based)
STAGE_TABLE_DESCRIPTION, STANDARD_TABLE_DESCRIPTION = 11 + 10, 22 + 10
FIRST_ROW = 12


def _tables(tmp_path, edit=None):
    workbook = family_e_workbook(tmp_path / "e.xlsx", [None])
    if edit is not None:
        book = openpyxl.load_workbook(workbook)
        edit(book["FEED_5_MAPPING"])
        book.save(workbook)
    done, _feed = _run(tmp_path, workbook, family_e_frd(tmp_path / "frd.json"))
    contract = json.loads((tmp_path / "sttm.json").read_text(encoding="utf-8"))
    (ddl,) = [p for p in (tmp_path / "out").rglob("*_DDL.txt") if "combined" not in p.parts]
    text = ddl.read_text(encoding="utf-8")
    comments = re.findall(r"^COMMENT '((?:\\.|[^'\\])*)'$", text, re.M)
    report = "\n".join(p.read_text(encoding="utf-8") for p in (tmp_path / "reports").glob("*.md"))
    return contract["feeds"][0], comments, report, done.stdout


def test_the_table_comment_is_the_description_then_the_load_strategy(tmp_path):
    feed, comments, _report, _out = _tables(tmp_path)
    assert feed["table_descriptions"] == {"stage.feed_5_dtl": "claims",
                                          "standard.feed_5_dtl": "claims"}
    assert len(comments) == 2                                     # stage + standard
    for comment in comments:
        assert comment.startswith("claims — Load strategy: "), comment


def test_rows_stating_two_descriptions_give_none_and_a_flag(tmp_path):
    def differ(sheet):
        sheet.cell(FIRST_ROW + 1, STAGE_TABLE_DESCRIPTION).value = "something else"

    feed, comments, report, _out = _tables(tmp_path, differ)
    assert feed["table_descriptions"] == {"standard.feed_5_dtl": "claims"}
    assert any(f.startswith("table_description_differs:FEED_5_MAPPING stage.feed_5_dtl — 2 ")
               for f in feed["extraction_flags"])
    stage, standard = comments
    assert stage.startswith("Load strategy: ")                    # load strategy alone
    assert standard.startswith("claims — Load strategy: ")
    assert "ddl_table_comment_no_description:" in report


def test_a_band_with_no_table_description_keeps_the_load_strategy_alone(tmp_path):
    def blank(sheet):
        for row in range(FIRST_ROW, sheet.max_row + 1):
            for column in (STAGE_TABLE_DESCRIPTION, STANDARD_TABLE_DESCRIPTION):
                sheet.cell(row, column).value = None

    feed, comments, report, _out = _tables(tmp_path, blank)
    assert "table_descriptions" not in feed                       # absent from the JSON
    assert len(comments) == 2 and all(c.startswith("Load strategy: ") for c in comments)
    assert report.count("ddl_table_comment_no_description:") == 1     # one per feed
