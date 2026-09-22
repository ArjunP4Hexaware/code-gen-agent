"""A sheet's REAL used range replaces a trusted max_row (staging M11 item 11).

A styled EMPTY cell a million rows down makes openpyxl report max_row
1,048,538, and every normal-mode ``iter_rows`` then CREATES a cell per row —
one real STTM took ~120-151 s to classify / read and its selection timed out.
Every document load for discovery and extraction goes through
``codegen.layout.extent.load_document``, which trims the dead tail."""

from __future__ import annotations

import time
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from codegen.layout.classify import classify_workbook
from codegen.layout.discover import discover
from codegen.layout.extent import load_document, trim_worksheet, used_max_row
from codegen.layout.fingerprint import fingerprint
from codegen.layout.resolve import resolve_pair

REPO = Path(__file__).resolve().parents[1]
SHAPES = REPO / "fixtures" / "acfc_shapes"
PAIR_1_STTM = SHAPES / "sttm" / "pair_1_family_a.xlsx"
PAIR_1_FRD = SHAPES / "frd" / "f1_pair_1.docx"
DATE = "2026-01-01"
LAST_EXCEL_ROW_SEEN = 1_048_538          # the row the real bloated sheet reached


@pytest.fixture(autouse=True)
def _mock_only(monkeypatch):
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_PROVIDER", "1")
    monkeypatch.setenv("CODEGEN_FORCE_MOCK_LAYOUT", "1")


def test_used_max_row_stops_at_the_first_long_gap():
    assert used_max_row([], 500) == (0, 0)
    assert used_max_row([1, 2, 3, 300], 500) == (300, 0)          # gaps shorter than 500
    assert used_max_row([1, 2, 3, 2000], 500) == (3, 1)            # 1996 empty rows: stop
    assert used_max_row([1, 5000], 0) == (5000, 0)                 # 0 = trust everything


def test_a_dead_tail_is_trimmed_and_valued_rows_beyond_a_gap_are_reported():
    ws = Workbook().active
    for row in range(1, 11):
        ws.cell(row=row, column=1).value = f"v{row}"
    ws.cell(row=2000, column=1).value = "after the gap"
    ws.cell(row=900_000, column=3).font = Font(bold=True)          # formatting only
    assert trim_worksheet(ws, 500) == (900_000, 10, 1)
    assert ws.max_row == 10


def test_a_few_formatted_empty_rows_keep_max_row_exact():
    """max_row is fingerprint input: a normal sheet must not move it."""
    ws = Workbook().active
    for row in range(1, 11):
        ws.cell(row=row, column=1).value = row
    ws.cell(row=40, column=2).font = Font(bold=True)               # a 30-row tail < 500
    assert trim_worksheet(ws, 500) == (40, 40, 0)
    assert ws.max_row == 40


def test_zero_trusts_max_row_and_an_all_empty_formatted_sheet_is_trimmed():
    ws = Workbook().active
    ws.cell(row=1, column=1).value = "x"
    ws.cell(row=5000, column=1).font = Font(bold=True)
    assert trim_worksheet(ws, 0) == (5000, 5000, 0)                # 0 = the old behaviour
    empty = Workbook().active
    empty.cell(row=LAST_EXCEL_ROW_SEEN, column=1).fill = PatternFill("solid", fgColor="FFFFFF")
    declared, _used, beyond = trim_worksheet(empty, 500)
    assert declared == LAST_EXCEL_ROW_SEEN and beyond == 0
    assert empty.max_row == 1


def test_the_knob_is_in_the_tracked_config(config):
    assert config.extractor.used_range_empty_rows == 500


def _bloated_pair1(tmp_path: Path) -> Path:
    """The pair-1 STTM with one styled EMPTY cell at row 1,048,538 on every sheet."""
    workbook = load_workbook(PAIR_1_STTM)
    for ws in workbook.worksheets:
        ws.cell(row=LAST_EXCEL_ROW_SEEN, column=1).fill = PatternFill("solid", fgColor="FFFFFF")
    path = tmp_path / "bloated.xlsx"
    workbook.save(path)
    return path


def test_the_loader_reads_the_used_range_not_max_row(tmp_path, config):
    path = _bloated_pair1(tmp_path)
    workbook = load_document(path, config.extractor.used_range_empty_rows)
    raw = load_workbook(path, data_only=True)
    for ws in workbook.worksheets:
        assert raw[ws.title].max_row == LAST_EXCEL_ROW_SEEN
        assert ws.max_row < 500
    trimmed = workbook.trim_report.trimmed
    assert set(trimmed) == set(workbook.sheetnames)
    assert all("formatting only" in note for note in workbook.trim_report.notes())


def test_a_clean_sheet_is_untouched_by_the_loader(config):
    """The tracked layout profiles are keyed by fingerprint: a clean document
    must fingerprint exactly as a plain load does."""
    clean = load_document(PAIR_1_STTM, config.extractor.used_range_empty_rows)
    assert clean.trim_report.trimmed == {}
    assert fingerprint(clean) == fingerprint(load_workbook(PAIR_1_STTM, data_only=True))


def test_a_bloated_sttm_classifies_resolves_and_extracts_fast_and_identically(
        tmp_path, pair1_config):
    """Without the trim, this workbook took minutes to classify and read.
    With it: classify + discovery + resolve + extraction well inside 30 s, and
    the STTM contract is byte-identical to the clean workbook's."""
    from codegen.extract import contract_to_json as sttm_to_json
    from codegen.extract import extract_contract
    from codegen.extract.frd_docx import contract_to_json as frd_to_json

    def run(sttm: Path, work: Path) -> tuple[str, str, str]:
        work.mkdir()
        kind = classify_workbook(sttm, pair1_config.extractor).kind
        profile = discover(sttm, pair1_config.extractor).profile
        pair = resolve_pair(sttm, PAIR_1_FRD, pair1_config, provider=None, cache_dirs=[],
                            generated_date=DATE)
        assert pair.questions == []
        frd = work / "frd.json"
        frd.write_text(frd_to_json(pair.frd_contract), encoding="utf-8")
        text = sttm_to_json(extract_contract(sttm, frd, pair1_config, generated_date=DATE,
                                             layout=pair.sttm.profile))
        discovered = profile.model_dump_json(exclude={"fingerprint"})
        return kind, discovered.replace(sttm.name, "<wb>"), text.replace(sttm.name, "<wb>")

    bloated_path = _bloated_pair1(tmp_path)
    started = time.monotonic()
    bloated = run(bloated_path, tmp_path / "bloated")
    elapsed = time.monotonic() - started
    clean = run(PAIR_1_STTM, tmp_path / "clean")
    assert elapsed < 30, f"the bloated workbook took {elapsed:.1f}s"
    assert bloated[0] == clean[0] == "sttm"
    assert bloated[1] == clean[1]
    assert bloated[2] == clean[2]
