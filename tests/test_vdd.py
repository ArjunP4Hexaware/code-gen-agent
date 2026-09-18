"""M3 — the Vendor Data Dictionary through the layout machinery, its
contract, and the STTM-vs-VDD cross-check gate.

Pinned: the three VDD fixtures (V1 segments, V2 per file, V3 per table)
resolve by synonyms alone and extract to the expected tables; the pair-1
STTM and VDD agree with zero vdd flags; a deliberately perturbed VDD yields
exactly five flags citing the right cells; a multi-sheet VDD is narrowed to
a two-table STTM's sheets; a fixed-width FRD with a position-less VDD fails
the gate naming the remedy; the canary cannot reach any VDD output.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from acfc_shapes.layout_truth import CANARY
from codegen.extract import contract_to_json as sttm_to_json
from codegen.extract import extract_contract
from codegen.extract.vdd import contract_to_json, extract_vdd_contract
from codegen.gate.vdd_check import vdd_cross_check
from codegen.layout.discover import discover_vdd
from codegen.layout.model import MockLayoutProvider
from codegen.layout.resolve import resolve_pair, resolve_workbook
from codegen.resolve.resolver import resolve_pair as resolve_contracts

REPO = Path(__file__).resolve().parents[1]
STTM = REPO / "fixtures" / "acfc_shapes" / "sttm"
FRD = REPO / "fixtures" / "acfc_shapes" / "frd"
VDD = REPO / "fixtures" / "acfc_shapes" / "vdd"
PROFILES = REPO / "fixtures" / "layout_profiles"
MOCK = PROFILES / "mock"
DATE = "2026-01-01"

# fixture -> (files rows, field sheets, fields, position rows, canonical segments)
EXPECTED = {
    "pair_1_v1_segments.xlsx": (4, ["FEED_1 Fields"], 23, 23, {"Header", "Detail", "Trailer"}),
    "pair_2_v2_per_file.xlsx": (3, ["Claims Fields", "Members Fields", "Providers Fields"], 32, 0,
                                set()),
    "pair_9_v3_per_table.xlsx": (3, ["ENCOUNTER", "ENCOUNTER_LINE", "ENCOUNTER_STATUS"], 25, 11,
                                 set()),
}


@pytest.fixture()
def no_cache(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    return [empty]


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_vdd_fixtures_resolve_by_synonyms_and_extract_as_pinned(name, config, no_cache):
    mock = MockLayoutProvider([MOCK, PROFILES])
    doc, _ = resolve_workbook(VDD / name, config, provider=mock, cache_dirs=no_cache,
                              document="vdd")
    assert doc.document == "vdd" and doc.profile.strategy == "vdd"
    assert doc.profile.source == "synonyms" and doc.complete
    assert doc.provider_calls == 0 and mock.requests == []
    assert doc.sources["model"] == 0 and doc.sources["synonyms"] > 0
    contract, profile = extract_vdd_contract(VDD / name, config, generated_date=DATE)
    files, sheets, fields, positions, segments = EXPECTED[name]
    assert len(contract.files) == files
    assert contract.field_sheets == sheets
    assert len(contract.fields) == fields
    assert len(contract.position_rows) == positions
    assert {f.segment_canonical for f in contract.fields if f.segment} == segments
    assert all(f.provenance.sheet == f.sheet and f.provenance.source == "synonyms"
               for f in contract.fields)
    assert contract.layout is not None and contract.layout.unresolved == []
    # Every FILES row cites its pattern cell.
    assert all("pattern" in row.cells for row in contract.files)


def test_vdd_pair1_positions_and_files_are_verbatim(config):
    contract, _ = extract_vdd_contract(VDD / "pair_1_v1_segments.xlsx", config,
                                       generated_date=DATE)
    first = contract.fields[0]
    assert (first.name, first.data_type, first.start, first.end, first.length) == (
        "SEGMENT_IDENTIFIER", "String", 1, 2, 2)
    assert first.segment == "HDR" and first.segment_canonical == "Header"
    assert first.cells["start"].a1 == "FEED_1 Fields!D2"
    copay = next(f for f in contract.fields if f.name == "COPAY")
    assert (copay.start, copay.end, copay.length, copay.row) == (69, 87, 19, 15)
    assert contract.has_positions
    files = {f.pattern: f for f in contract.files}
    assert files["I_ACCUM_*_TO_CLIENT_*.csv"].format == "Fixed Width"
    assert files["I_ACCUM_*_TO_CLIENT_*.csv"].record_type_field == "SEGMENT_IDENTIFIER"
    assert files["I_ACCUM_*_TO_CLIENT_*.csv"].delimiter is None
    assert contract_to_json(contract) == contract_to_json(
        extract_vdd_contract(VDD / "pair_1_v1_segments.xlsx", config, generated_date=DATE)[0])


# ----------------------------------------------------------- pair-1 gate


def _pair1_spec(config, tmp_path, vdd_workbook: Path, no_cache):
    """Resolve pair 1 (mock layout), extract the STTM + VDD contracts, and
    resolve them with an FRD contract that names the STTM's file pattern and
    segments (what M4 will do); return the single resolved spec."""
    pair = resolve_pair(STTM / "pair_1_family_a.xlsx", FRD / "f1_pair_1.docx", config,
                        provider=MockLayoutProvider([MOCK, PROFILES]), cache_dirs=no_cache,
                        generated_date=DATE)
    from codegen.extract.frd_docx import contract_to_json as frd_to_json

    frd_json = tmp_path / "frd.contract.json"
    data = json.loads(frd_to_json(pair.frd_contract))
    data["feeds"][0]["file_name_patterns"] = ["I_ACCUM_*_TO_CLIENT_*.csv"]
    data["feeds"][0]["record_segments"] = ["Header", "Detail", "Trailer"]
    frd_json.write_text(json.dumps(data), encoding="utf-8")
    sttm_json = tmp_path / "sttm.contract.json"
    contract = extract_contract(STTM / "pair_1_family_a.xlsx", frd_json, config,
                                generated_date=DATE, layout=pair.sttm.profile)
    sttm_json.write_text(sttm_to_json(contract), encoding="utf-8")
    vdd_json = tmp_path / "vdd.contract.json"
    vdd_contract, _ = extract_vdd_contract(vdd_workbook, config, generated_date=DATE)
    vdd_json.write_text(contract_to_json(vdd_contract), encoding="utf-8")
    (spec,) = resolve_contracts(frd_json, sttm_json, config, vdd_path=vdd_json)
    assert spec.vdd is not None
    return spec


def test_pair1_sttm_and_vdd_agree_with_zero_flags(config, tmp_path, no_cache):
    spec = _pair1_spec(config, tmp_path, VDD / "pair_1_v1_segments.xlsx", no_cache)
    flags, check = vdd_cross_check(spec, config)
    assert flags == []
    assert check is not None and check.name == "vdd_positions" and check.passed
    # STTM fields carry the source band's length / start / end for the check.
    field = spec.detail_segment.fields[1]
    assert (field.source_column, field.source_start, field.source_length, field.source_end) == (
        "CARDHOLDER_ID", "3", "15", "17")


def _perturbed_vdd(tmp_path: Path) -> Path:
    wb = load_workbook(VDD / "pair_1_v1_segments.xlsx")
    ws = wb["FEED_1 Fields"]
    ws["C11"] = "Numeric"          # CARDHOLDER_ID type (STTM says String)
    ws["F12"] = 17                 # PLAN_ID length (STTM says 16)
    ws["D15"] = 70                 # COPAY start (STTM says 69)
    ws.delete_rows(24)             # FILLER_3 (Trailer) removed
    ws.append([24, "EXTRA_FIELD", "String", 33, 40, 8, None, None, None, "TRL", "not in the STTM"])
    path = tmp_path / "pair_1_perturbed.xlsx"
    wb.save(path)
    return path


def test_perturbed_vdd_yields_exactly_the_five_flags_citing_cells(config, tmp_path, no_cache):
    spec = _pair1_spec(config, tmp_path, _perturbed_vdd(tmp_path), no_cache)
    flags, check = vdd_cross_check(spec, config)
    assert check is not None and check.passed
    classes = sorted(f.split(" — ")[0] for f in flags)
    assert classes == ["vdd_mismatch:length", "vdd_mismatch:position", "vdd_mismatch:type",
                       "vdd_missing_in_sttm", "vdd_missing_in_vdd"]
    by_class = {f.split(" — ")[0]: f for f in flags}
    assert "VDD FEED_1 Fields!C11 (data_type 'Numeric')" in by_class["vdd_mismatch:type"]
    assert "STTM FEED_1_MAPPING row 25 (type 'String')" in by_class["vdd_mismatch:type"]
    assert "(field 'CARDHOLDER_ID')" in by_class["vdd_mismatch:type"]
    assert "VDD FEED_1 Fields!F12 (length 17)" in by_class["vdd_mismatch:length"]
    assert "STTM FEED_1_MAPPING row 26 (length '16')" in by_class["vdd_mismatch:length"]
    assert "VDD FEED_1 Fields!D15 (start 70)" in by_class["vdd_mismatch:position"]
    assert "start 69 vs 70" in by_class["vdd_mismatch:position"]
    assert "(field 'COPAY')" in by_class["vdd_mismatch:position"]
    assert "STTM FEED_1_MAPPING row 38 (field 'FILLER_3')" in by_class["vdd_missing_in_vdd"]
    assert "VDD FEED_1 Fields!B24 (field 'EXTRA_FIELD')" in by_class["vdd_missing_in_sttm"]
    # 23 vs 23: no count summary flag.
    assert not any(f.startswith("vdd_field_count") for f in flags)
    from codegen.gate.verdict import compute_verdict

    gate = compute_verdict("accumulators", [], [], [check], True, extra_flags=flags)
    assert gate.verdict == "PASS_WITH_FLAGS" and len([f for f in gate.flags if "vdd_" in f]) == 5


def test_type_equivalence_classes_are_not_mismatches(config):
    from codegen.gate.vdd_check import types_equivalent

    assert types_equivalent("String", "VARCHAR", config)
    assert types_equivalent("Decimal(17,2)", "numeric", config)
    assert types_equivalent("X(10)", "char", config)
    assert types_equivalent("9(6)", "int", config)
    assert not types_equivalent("String", "Numeric", config)
    assert not types_equivalent("Date", "Timestamp", config)


def test_fixed_width_frd_with_a_position_less_vdd_fails_the_gate(config, tmp_path, no_cache):
    spec = _pair1_spec(config, tmp_path, VDD / "pair_2_v2_per_file.xlsx", no_cache)
    flags, check = vdd_cross_check(spec, config)
    assert check is not None and not check.passed
    assert "fixed-width format" in check.details and "Start Position" in check.details
    from codegen.gate.verdict import compute_verdict

    gate = compute_verdict("accumulators", [], [], [check], True, extra_flags=flags)
    assert gate.verdict == "FAIL"
    # Three sheets of an unrelated dictionary: every STTM field is missing there.
    assert any(f.startswith("vdd_field_count") for f in flags)


# ----------------------------------------------------------- multi-sheet


def test_multi_sheet_vdd_is_narrowed_to_the_sttm_tables(config):
    found = discover_vdd(VDD / "pair_9_v3_per_table.xlsx", config.extractor,
                         sttm_tables=["src_sys_d_encounter", "src_sys_d_encounter_line",
                                      "ENCOUNTER", "ENCOUNTER_LINE"])
    kinds = {s.name: s.kind for s in found.profile.sheets}
    assert kinds == {"FILES": "vdd_files", "ENCOUNTER": "vdd_fields",
                     "ENCOUNTER_LINE": "vdd_fields", "ENCOUNTER_STATUS": "ignore"}
    assert any("not among the STTM's tables" in d for d in found.diagnostics)
    contract, _ = extract_vdd_contract(VDD / "pair_9_v3_per_table.xlsx", config,
                                       sttm_tables=["ENCOUNTER", "ENCOUNTER_LINE"],
                                       generated_date=DATE)
    assert contract.field_sheets == ["ENCOUNTER", "ENCOUNTER_LINE"]
    assert contract.ignored_sheets == ["ENCOUNTER_STATUS"]
    assert len(contract.fields) == 22
    # Pair-level: the STTM's stage / source tables drive the narrowing and
    # the VDD is the pair's third document.
    pair = resolve_pair(STTM / "pair_9_family_e.xlsx", None, config, provider=None,
                        cache_dirs=[], vdd_path=VDD / "pair_9_v3_per_table.xlsx")
    assert pair.vdd is not None and pair.vdd.complete
    assert [s.kind for s in pair.vdd.profile.sheets if s.name == "ENCOUNTER_STATUS"] == ["ignore"]
    assert pair.report()["vdd"]["roles_by_source"]["synonyms"] > 0


# ------------------------------------------------------------ pair level


def test_vdd_files_sheet_joins_the_pair_cross_checks(config, no_cache, tmp_path):
    pair = resolve_pair(STTM / "pair_1_family_a.xlsx", FRD / "f1_pair_1.docx", config,
                        provider=MockLayoutProvider([MOCK, PROFILES]), cache_dirs=no_cache,
                        vdd_path=VDD / "pair_1_v1_segments.xlsx", generated_date=DATE)
    assert pair.vdd is not None and pair.vdd.provider_calls == 0
    assert pair.provider_calls == 1
    fmt = next(c for c in pair.cross_checks if c.name == "file_format")
    assert fmt.status == "agree"
    freq = next(c for c in pair.cross_checks if c.name == "frequency")
    assert freq.status == "agree"
    # The VDD's FILES cells are among the cited candidates for cadence / format.
    report = pair.report()
    assert report["vdd"]["source"] == "synonyms"
    assert pair.questions == [] and not [f for f in pair.flags if "vdd" in f and "unresolved" in f]


def test_vdd_canary_and_adversarial_answers(config, no_cache, tmp_path):
    from codegen.layout.resolve import discovery_for

    # A model answer claiming length on the Description column is rejected.
    provider = MockLayoutProvider([MOCK], override=MOCK / "adversarial_vdd_free_text_column.json")
    doc, _ = resolve_workbook(VDD / "pair_1_v1_segments.xlsx", config, provider=provider,
                              cache_dirs=no_cache, document="vdd")
    # Synonyms already resolve every required role: the model is not called.
    assert doc.provider_calls == 0 and doc.complete
    # Force the model path on a copy whose header hides the Length column.
    wb = load_workbook(VDD / "pair_1_v1_segments.xlsx")
    wb["FEED_1 Fields"]["F1"] = "Len (bytes)"
    path = tmp_path / "vdd_needs_model.xlsx"
    wb.save(path)
    doc, workbook = resolve_workbook(path, config, provider=provider, cache_dirs=no_cache,
                                     document="vdd")
    assert doc.provider_calls == 1
    assert any("integer-like" in r.reason and r.role == "length" for r in doc.rejections)
    assert not doc.complete and any(q.role == "length" for q in doc.questions)
    # Canary: a model string never reaches the VDD contract or the report.
    canary = MockLayoutProvider([MOCK], override=MOCK / "adversarial_vdd_canary.json")
    doc, workbook = resolve_workbook(path, config, provider=canary, cache_dirs=no_cache,
                                     document="vdd")
    from codegen.extract.vdd import read_vdd

    contract = read_vdd(discovery_for(doc.profile, workbook), path.name, config,
                        generated_date=DATE)
    for text in (contract_to_json(contract), json.dumps([r.render() for r in doc.rejections]),
                 doc.profile.model_dump_json()):
        assert CANARY not in text
