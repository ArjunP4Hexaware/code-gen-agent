"""P0-5 (2026-10-30, pentest): document text in a written workbook is DATA.

openpyxl binds a str starting ``=`` as a formula cell; an FRD / STTM value such
as ``=WEBSERVICE("http://x")`` copied into the IIG would run when a BSA opens
it. Every workbook goes through ``codegen.xlsx_safe.neutralize_formulas``
(``metadata_sheet.stable_workbook_bytes``), and the gate check
``xlsx_no_formulas`` reopens every emitted workbook and FAILS on a formula.
"""

from __future__ import annotations

import datetime
import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from codegen import cli
from codegen.gate.xlsx_formulas import check_xlsx_no_formulas
from codegen.metadata_sheet import _pin_workbook_properties, stable_workbook_bytes
from test_m4_acceptance import _scoped

TRIGGERED = ["=1+1", '=WEBSERVICE("http://x")', "+1+1", "-1+1", "@SUM(A1)"]
UNTOUCHED = ["plain", "\t=1+1", " =1+1", "a=b", 5, 1.5, None, datetime.date(2026, 1, 1)]


def _book(values) -> Workbook:
    book = Workbook()
    _pin_workbook_properties(book)          # as every writer does: no wall clock
    sheet = book.active
    sheet.append(["header"] * len(values))
    sheet.append(values)
    return book


def _cells(data: bytes):
    book = load_workbook(BytesIO(data))
    return [c for s in book.worksheets for c in s._cells.values()]


def test_trigger_text_is_written_as_a_quote_prefixed_string_value_unchanged():
    data = stable_workbook_bytes(_book([*TRIGGERED, *UNTOUCHED]))
    row = [c for c in _cells(data) if c.row == 2]
    assert [c.data_type for c in row].count("f") == 0
    assert [c.value for c in row[:len(TRIGGERED)]] == TRIGGERED       # byte-identical text
    assert all(c.data_type == "s" and c.quotePrefix for c in row[:len(TRIGGERED)])
    assert not any(c.quotePrefix for c in row[len(TRIGGERED):])
    values = [c.value for c in row[len(TRIGGERED):]]
    assert values[:5] == UNTOUCHED[:5] and values[5] == 1.5
    assert stable_workbook_bytes(_book([*TRIGGERED, *UNTOUCHED])) == data   # byte-stable


def test_a_workbook_without_trigger_text_serializes_to_the_same_bytes(monkeypatch):
    import codegen.xlsx_safe

    neutralized = stable_workbook_bytes(_book(UNTOUCHED))
    monkeypatch.setattr(codegen.xlsx_safe, "neutralize_formulas", lambda workbook: 0)
    assert stable_workbook_bytes(_book(UNTOUCHED)) == neutralized


def test_the_gate_check_fails_on_a_formula_and_names_no_cell_text(tmp_path):
    raw = _book(['=WEBSERVICE("http://x")'])
    bad = tmp_path / "bad.xlsx"
    raw.save(bad)                                           # no neutralizing: a formula
    check = check_xlsx_no_formulas([bad])
    assert not check.passed and check.name == "xlsx_no_formulas"
    assert "bad.xlsx Sheet!A2 (header)" in check.details
    assert "WEBSERVICE" not in check.details
    good = tmp_path / "good.xlsx"
    good.write_bytes(stable_workbook_bytes(_book(['=WEBSERVICE("http://x")'])))
    assert check_xlsx_no_formulas([good]).passed


# ---- pair 4 end to end: FRD + STTM text with trigger characters ----------
VARIANTS = {
    "eq_at": {"source_system": "=1+1", "domain": "@Membership",
              "sub_domain": '=WEBSERVICE("http://x")'},
    "plus_minus": {"source_system": "+Northwind Benefits", "domain": "-Membership",
                   "sub_domain": "=1+1"},
}
_STTM_CELL = {"source_system": "B2", "domain": "B6", "sub_domain": "B7"}


def _pair4_run(pair4_config, tmp, values):
    from acfc_shapes import FIXTURE_ROOT, pair4_nb
    from codegen.extract import contract_to_json, extract_contract
    from codegen.resolve.resolver import resolve_pair
    from codegen.run_artefacts import KeptFeed, write_combined

    frd = json.loads((FIXTURE_ROOT / pair4_nb.FRD_PATH).read_text(encoding="utf-8"))
    frd["feeds"][0].update(values)
    frd["feeds"][0]["validation_rules"] = ['=WEBSERVICE("http://x/rule")', "+1+1", "-2"]
    frd_path = tmp / "frd.contract.json"
    frd_path.write_text(json.dumps(frd, indent=2), encoding="utf-8")
    sttm_book = load_workbook(BytesIO((FIXTURE_ROOT / pair4_nb.STTM_PATH).read_bytes()))
    sheet = sttm_book.worksheets[0]
    for key, coordinate in _STTM_CELL.items():          # the STTM states the same text
        sheet[coordinate] = values[key]
        sheet[coordinate].data_type = "s"
    sttm_path = tmp / "sttm_nb_cob_report.xlsx"
    sttm_book.save(sttm_path)
    sttm = extract_contract(sttm_path, frd_path, pair4_config, generated_date="2026-10-07")
    contract = tmp / "sttm.contract.json"
    contract.write_text(contract_to_json(sttm), encoding="utf-8")
    (spec,) = resolve_pair(frd_path, contract, pair4_config)
    config = _scoped(pair4_config, tmp)
    result = cli.generate_feed_result(spec, config, dry_run=True, skip_tests=True,
                                      output_mode="framework", conventions_profile="acfc_prx",
                                      iig_template="iig_v2")
    out = tmp / "out"
    block = write_combined(out, "nb_cob_report", [KeptFeed(spec.feed_slug, result.gate.verdict,
                                                           result.framework_artefacts)],
                           [], config, [("FRD contract", "x")],
                           config.conventions.get("acfc_prx"))
    return result.gate, block, sorted(out.rglob("*.xlsx"))


@pytest.mark.parametrize("variant", sorted(VARIANTS))
def test_pair4_trigger_text_reaches_every_workbook_as_text(variant, pair4_config,
                                                           tmp_path):
    values = VARIANTS[variant]
    gate, block, books = _pair4_run(pair4_config, tmp_path, values)
    assert {"config_rows.xlsx", "config_inserts.xlsx"} <= {p.name for p in books}
    assert any(p.parent.name == "combined" for p in books)
    found = dict.fromkeys(values.values(), 0)
    for path in books:
        cells = _cells(path.read_bytes())
        assert [c.coordinate for c in cells if c.data_type == "f"] == [], path.name
        for cell in cells:
            if cell.value in found:
                assert cell.data_type == "s" and cell.quotePrefix, path.name
                found[cell.value] += 1
    assert all(found.values()), found                   # each text landed, unchanged
    (check,) = [c for c in gate.checks if c.name == "xlsx_no_formulas"]
    assert check.passed
    assert "xlsx_no_formulas" not in block["failed_checks"]


def test_pair4_gate_fails_when_a_formula_reaches_a_workbook(pair4_config, tmp_path,
                                                           monkeypatch):
    """The gate is the second line: with the writer funnel disabled the formula
    cells are written, and xlsx_no_formulas FAILS the feed and the combined set."""
    import codegen.xlsx_safe

    monkeypatch.setattr(codegen.xlsx_safe, "neutralize_formulas", lambda workbook: 0)
    gate, block, _books = _pair4_run(pair4_config, tmp_path, VARIANTS["eq_at"])
    (check,) = [c for c in gate.checks if c.name == "xlsx_no_formulas"]
    assert not check.passed
    assert "WEBSERVICE" not in check.details and "=1+1" not in check.details
    assert gate.verdict == "FAIL"
    assert "xlsx_no_formulas" in block["failed_checks"]


def test_edh_questions_workbook_quotes_a_trigger_cell_as_text():
    """The EDH E2E questions workbook QUOTES document cells (``cell_text``)."""
    from codegen.config import EdhE2eConfig
    from codegen.layouts import edh_e2e
    from test_edh_e2e_reader import FIXTURES

    contract = edh_e2e.read_dir(FIXTURES, EdhE2eConfig())
    i, feed = next((i, f) for i, f in enumerate(contract.feeds) if f.questions)
    text = '=WEBSERVICE("http://x")'
    questions = [feed.questions[0].model_copy(update={"quoted": text}), *feed.questions[1:]]
    feeds = list(contract.feeds)
    feeds[i] = feed.model_copy(update={"questions": questions})
    data, _rows = edh_e2e.questions_workbook(contract.model_copy(update={"feeds": feeds}))
    cells = _cells(data)
    assert [c.coordinate for c in cells if c.data_type == "f"] == []
    (cell,) = [c for c in cells if c.value == text]
    assert cell.data_type == "s" and cell.quotePrefix
