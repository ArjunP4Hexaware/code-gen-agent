"""Per-column fill rate of an IIG corpus against the agent's IIG workbooks.

    python tools/iig_fill_rate.py --corpus <xlsx | csv | dir> [...] --agent <xlsx | dir> [...]
        [--sheets DATABRICKS_NOTEBOOK_DETAILS,DATA_FACTORY_PIPELINE_SCHEDULE,EMAIL_TEMPLATE_CONFIG]
        [--threshold 90] [--json out.json] [--show-values] [--patterns DIR [--top N]
        [--group-index CSV ...] [--vocabulary tools/iig_vocabulary.yaml]]

2026-10-13: for every column of the chosen sheets, how many corpus rows fill it
(a cell that is not blank and not ``NULL``) against how many agent rows do.
A column the corpus fills in at least ``--threshold`` % of rows and the agent
does not is a GAP; a column the iig_keys.yaml ``needs_assignment`` set names
(ids, audit stamps) is ``assigned`` — the framework / the engineer fills it,
the agent never invents it. The share of the corpus's most common value is
printed (the value itself only with ``--show-values``: corpus values are
production text — D1 only, never committed).

2026-10-14 (D1 post43 item 3): ``--patterns DIR`` writes ``patterns.md`` +
``patterns.json`` — per column of the chosen sheets the dominant value and its
share, and a SLOT pattern (each row's own feed / table / schema / domain /
subdomain / file / source values replaced by ``{feed}``, ``{table}``, …, taken
from its workbook's ingestion sheets) with its share; DAY_OF_SCHEDULE's fill
rate per PIPELINE_FREQUENCY; the grand master / master schedule rows per
GROUP_ID. SCRUBBED: every remaining token outside the fixture vocabulary is a
``G<n>`` placeholder (tools/iig_diff.py's export scrubber) and the written
files are re-scanned with scripts/scrub_check.py (a hit deletes them) — safe
to bring home. Sheets match by name with case and separators ignored (a name
cut to Excel's 31 characters matches its full name); a header row is the
first of the top rows with two or more cells; a ``.csv`` corpus file is one
table, named by its file stem (a suffix allowed: the EMAIL_TEMPLATE_CONFIG
export tools/build_goldens.py reads, email_template_config_by_pipeline.csv —
the 199 per-pair workbooks carry no email sheet). Every corpus workbook's
sheets are listed.

2026-10-14 (D1 post44): ``--top N`` (default: all) caps every value / pattern
list (``top_values``, ``top_patterns``, the per-flow and email lists); a
PER-FLOW breakdown (file_to_stg / stg_to_std / dq / grand_master / master,
else unclassified — the group index by PIPELINE_ID, the master rows, the
workbook's own schedule row, tools/iig_vocabulary.yaml flow_rules, then
iig_keys.yaml flow.rules) of the notebook path / name / process name and the
schedule's PIPELINE_DESCRIPTION; the email sheet's STATUS x SUBJECT pattern x
BODY pattern table on FILE-INGESTION pipelines with rows / pipelines per
status (an email row takes the slots of the workbook its PIPELINE_ID is in);
and a framework-vocabulary allowlist (tools/iig_vocabulary.yaml: notebook
names, DQ rule types / classes, STATUS values, layer names pass unscrubbed —
never a token of a slot value or a person column).

Run on D1 over the 199-IIG corpus (docs/acfc/iig_goldens/*.xlsx); at home
over the SFMC reference IIG, the pair-1 golden and any golden at hand.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

TOOLS = Path(__file__).resolve().parent
DEFAULT_SHEETS = ("DATABRICKS_NOTEBOOK_DETAILS", "DATA_FACTORY_PIPELINE_SCHEDULE",
                  "EMAIL_TEMPLATE_CONFIG")
SCHEDULE = "DATA_FACTORY_PIPELINE_SCHEDULE"
# 2026-10-14 (D1 post43 item 6): the corpus's DQ rule classes (--patterns)
RULES = "DATA_QUALITY_RULES"
RULE_COLUMNS = ("ruletype", "ruleclass")
# The ingestion sheets whose cells name a workbook's feed / tables / files —
# the slot values of --patterns (ALL_FILES_STATIC_INFORMATION is the whole
# inventory, never one feed's).
SLOT_SHEETS = ("ADLS_DELTA_INGESTION_DETAILS", "STGDELTA_STDDELTA_INGESTION_DETAILS",
               "FILE_ADLS_INGESTION_DETAILS", "ADLS_FIXED_WIDTH_HANDLER")
SLOT_COLUMNS = {      # slot -> header ids (case / separators dropped)
    "table": ("tgttablename", "srctablename", "tgttable"),
    "schema": ("tgtschemaname", "srcschemaname"),
    "catalog": ("tgtcatalogname", "srccatalogname", "tgtdatabasename", "srcdatabasename"),
    "domain": ("domain",),
    "subdomain": ("subdomain",),
    "feed": ("objectname",),
    "file": ("srcfilename", "filename"),
    "source": ("source",),
}
NULLS = {"", "null", "none"}
HEADER_SCAN_ROWS = 10
MIN_SLOT_CHARS = 3
VOCABULARY = TOOLS / "iig_vocabulary.yaml"
EMAIL = "EMAIL_TEMPLATE_CONFIG"
GROUP_INDEX_NAMES = ("01_pair_group_index.csv", "01_pair_group_index_addendum.csv")
UNCLASSIFIED = "unclassified"
_NON_ALNUM = re.compile(r"[^a-z0-9]+")
_SHEET_KEY = re.compile(r"[^A-Z0-9]+")
_EXCEL_SHEET_CHARS = 31
_GRAND_MASTER = re.compile(r"grand\s*[-_ ]?\s*master", re.IGNORECASE)
_MASTER = re.compile(r"(?<![a-z])master(?![a-z])", re.IGNORECASE)


def _diff_module():
    if "iig_diff" in sys.modules:
        return sys.modules["iig_diff"]
    spec = importlib.util.spec_from_file_location("iig_diff", TOOLS / "iig_diff.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules["iig_diff"] = module
    spec.loader.exec_module(module)
    return module


def sheet_id(name: str) -> str:
    return name.strip().upper()[:_EXCEL_SHEET_CHARS]


def sheet_key(name: str) -> str:
    """A sheet name as compared: upper case, separators dropped."""
    return _SHEET_KEY.sub("", str(name).upper())


def match_sheet(name: str, wanted) -> str | None:
    """The wanted sheet ``name`` is (case / separators ignored; a name cut to
    Excel's 31 characters matches its full form), else None."""
    key, cut = sheet_key(name), sheet_key(sheet_id(name))
    for target in wanted:
        if key in (sheet_key(target), sheet_key(sheet_id(target))) or cut == sheet_key(target):
            return target
    return None


def header_id(name: object) -> str:
    return _NON_ALNUM.sub("", str(name or "").lower())


def filled(value: object) -> bool:
    return str(value if value is not None else "").strip().lower() not in NULLS


def workbooks(paths: list[Path]) -> list[Path]:
    out: list[Path] = []
    for path in paths:
        if path.is_dir():
            # a run folder: its IIGs only (never config_rows / config_inserts)
            # post91: never the run's combined/ copy (the same rows as the feeds')
            books = sorted(p for p in path.rglob("*.xlsx") if not p.name.startswith("~$")
                           and not p.name.endswith("_REVIEW.xlsx")
                           and "combined" not in p.relative_to(path).parts)
            iigs = [p for p in books if p.name.upper().endswith("_IIG.XLSX")]
            out += iigs or books
            out += sorted(path.rglob("*.csv")) if not iigs else []
        elif path.is_file():
            out.append(path)
    return out


@dataclass
class Book:
    """One corpus / agent workbook: every sheet's name, and the read ones."""

    name: str
    sheetnames: list[str]
    # canonical sheet -> ({header id: header}, [{header id: value}])
    sheets: dict[str, tuple[dict, list[dict]]] = field(default_factory=dict)


def _table(rows) -> tuple[list, list]:
    """(header row, data rows): the header is the first of the top rows with
    two or more cells (a title / blank row above it is skipped)."""
    rows = iter(rows)
    head: tuple = ()
    for _ in range(HEADER_SCAN_ROWS):
        row = next(rows, None)
        if row is None:
            return [], []
        if sum(1 for v in row if filled(v)) >= 2:
            head = row
            break
    return list(head), list(rows)


def _records(head: list, rows: list) -> tuple[dict, list[dict]]:
    ids = [header_id(h) for h in head]
    headers = {i: str(h).strip() for h, i in zip(head, ids, strict=False) if i}
    records = [{i: v for i, v in zip(ids, row, strict=False) if i} for row in rows
               if any(filled(v) for v in row)]
    return headers, records


def read_books(paths: list[Path], wanted) -> list[Book]:
    books: list[Book] = []
    for path in paths:
        if path.suffix.lower() == ".csv":
            # a table export: its stem names the table, possibly with a suffix
            # (build_goldens' email_template_config_by_pipeline.csv)
            target = match_sheet(path.stem, wanted) or next(
                (w for w in wanted if sheet_key(path.stem).startswith(sheet_key(w))), None)
            book = Book(path.name, [path.stem])
            if target:
                with path.open(encoding="utf-8-sig", newline="") as fh:
                    head, rows = _table(csv.reader(fh))
                book.sheets[target] = _records(head, rows)
            books.append(book)
            continue
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            book = Book(path.name, list(wb.sheetnames))
            for name in wb.sheetnames:
                target = match_sheet(name, wanted)
                if target is None:
                    continue
                headers, records = _records(*_table(wb[name].iter_rows(values_only=True)))
                have_h, have_r = book.sheets.setdefault(target, ({}, []))
                for k, v in headers.items():
                    have_h.setdefault(k, v)
                have_r.extend(records)
        finally:
            wb.close()
        books.append(book)
    return books


def merge_books(books: list[Book], sheets) -> dict[str, tuple[dict, list[dict]]]:
    out: dict[str, tuple[dict, list[dict]]] = {}
    for book in books:
        for sid, (headers, records) in book.sheets.items():
            if sid not in sheets:
                continue
            have_h, have_r = out.setdefault(sid, ({}, []))
            for k, v in headers.items():
                have_h.setdefault(k, v)
            have_r.extend(records)
    return out


def sheet_rows(paths: list[Path], wanted: set[str]) -> dict[str, tuple[dict, list[dict]]]:
    """sheet -> ({header id: header as first written}, [{header id: value}])
    over every workbook; a row whose cells are all blank is no row."""
    return merge_books(read_books(paths, wanted), set(wanted))


def always_blank(root: Path = TOOLS.parent, template: str = "iig_v2") -> set[str]:
    """The template's framework-assigned columns (config.yaml, the ACFC overlay
    on top: a list replaces) — assigned like the needs_assignment set."""
    import yaml

    found: list[str] = []
    for path in (root / "config" / "config.yaml", root / "config" / "overlays" / "acfc_env.yaml"):
        try:
            raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        except OSError:
            continue
        listed = (((raw.get("metadata") or {}).get("templates") or {}).get(template)
                  or {}).get("always_blank")
        if listed is not None:
            found = list(listed)
    return {header_id(h) for h in found}


def fill_rates(corpus: dict, agent: dict, sheets: list[str], threshold: float,
               assigned) -> dict:
    """Per sheet, per column: corpus / agent rows and filled counts, the
    dominant corpus value's share, and the status (gap / ok / assigned /
    sparse)."""
    report: dict = {}
    for sid in sheets:
        c_heads, c_rows = corpus.get(sid, ({}, []))
        a_heads, a_rows = agent.get(sid, ({}, []))
        columns = {**{k: v for k, v in a_heads.items()}, **c_heads}
        sheet: dict = {"corpus_rows": len(c_rows), "agent_rows": len(a_rows), "columns": {}}
        for cid, header in columns.items():
            c_filled = [r.get(cid) for r in c_rows if filled(r.get(cid))]
            a_filled = sum(1 for r in a_rows if filled(r.get(cid)))
            c_pct = round(100.0 * len(c_filled) / len(c_rows), 1) if c_rows else None
            a_pct = round(100.0 * a_filled / len(a_rows), 1) if a_rows else None
            common = Counter(str(v).strip() for v in c_filled).most_common(1)
            status = ("assigned" if assigned(header) else
                      "no corpus rows" if c_pct is None else
                      "sparse" if c_pct < threshold else
                      "gap" if (a_pct or 0.0) < threshold else "ok")
            sheet["columns"][header] = {
                "corpus_filled": len(c_filled), "corpus_pct": c_pct,
                "agent_filled": a_filled, "agent_pct": a_pct,
                "dominant_share": (round(100.0 * common[0][1] / len(c_filled), 1)
                                   if common else None),
                "dominant_value": common[0][0] if common else None,
                "in_corpus": cid in c_heads, "in_agent": cid in a_heads, "status": status,
            }
        report[sid] = sheet
    return report


def print_report(report: dict, show_values: bool) -> None:
    for sid, sheet in report.items():
        print(f"{sid}: corpus {sheet['corpus_rows']} row(s), agent {sheet['agent_rows']} "
              "row(s)")
        print(f"  {'column':<34} {'corpus %':>9} {'agent %':>8} {'dominant %':>10}  status")
        for header, col in sheet["columns"].items():
            def pct(v):
                return "-" if v is None else f"{v:.1f}"
            line = (f"  {header:<34} {pct(col['corpus_pct']):>9} {pct(col['agent_pct']):>8} "
                    f"{pct(col['dominant_share']):>10}  {col['status']}")
            if show_values and col["dominant_value"] is not None:
                line += f"  (D1 only: {col['dominant_value'][:60]!r})"
            print(line)
        gaps = [h for h, c in sheet["columns"].items() if c["status"] == "gap"]
        print(f"  GAPS ({len(gaps)}): {', '.join(gaps) or 'none'}")


def print_inventory(books: list[Book], sheets: list[str], label: str = "corpus") -> None:
    """Every workbook's sheets, and which of the chosen sheets each gave rows."""
    print(f"{label} workbooks and their sheets:")
    for n, book in enumerate(books, 1):
        read = ", ".join(f"{s} {len(book.sheets[s][1])}" for s in sheets if s in book.sheets)
        print(f"  {n:>3} {book.name}: {len(book.sheetnames)} sheet(s) "
              f"[{', '.join(book.sheetnames)}] -> read: {read or 'none of the chosen sheets'}")
    for sid in sheets:
        n = sum(1 for b in books if sid in b.sheets)
        print(f"  {sid}: in {n} of {len(books)} {label} workbook(s)")


# ------------------------------------------------------------- --patterns


def _file_tokens(diff) -> tuple:
    try:
        return diff.load_keys(TOOLS / "iig_keys.yaml").file_name_tokens
    except Exception:  # noqa: BLE001 — patterns still work without the date rules
        return ()


def _slot_regex(value: str, tokens: tuple) -> re.Pattern[str] | None:
    """A slot value as a pattern over a cell: its alphanumeric words in order,
    any separators between them, case-insensitive, whole words; wildcards and
    date tokens (iig_keys.yaml file_name_tokens) dropped first."""
    text = str(value).strip().lower().replace("\\", "/").rsplit("/", 1)[-1]
    text = re.sub(r"\.[a-z0-9]{1,5}$", "", text)
    for pattern in tokens:
        text = pattern.sub(" ", text)
    words = [w for w in _NON_ALNUM.split(text) if w]
    if not words or len("".join(words)) < MIN_SLOT_CHARS:
        return None
    body = r"[\W_]*".join(re.escape(w) for w in words)
    return re.compile(rf"(?<![a-z0-9]){body}(?![a-z0-9])", re.IGNORECASE)


def book_slots(book: Book, tokens: tuple) -> list[tuple[str, re.Pattern[str]]]:
    """(slot, regex) for every slot value of the workbook's ingestion sheets,
    the longest value first (a table name before the domain inside it)."""
    found: dict[str, tuple[str, str]] = {}
    for sid, (_headers, records) in book.sheets.items():
        if sid not in SLOT_SHEETS:
            continue
        for row in records:
            for slot, ids in SLOT_COLUMNS.items():
                for cid in ids:
                    value = str(row.get(cid) or "").strip()
                    if not filled(value):
                        continue
                    if slot == "table" and "." in value:
                        schema, value = value.rsplit(".", 1)
                        found.setdefault(schema.split(".")[-1].lower(), ("schema", schema))
                    found.setdefault(value.lower(), (slot, value))
    # one slot per value: a value two slots share (an OBJECT_NAME that is the
    # table's name) takes the first slot of SLOT_COLUMNS
    rank = {slot: n for n, slot in enumerate(SLOT_COLUMNS)}
    chosen: dict[str, tuple[str, re.Pattern[str]]] = {}
    for slot, value in sorted(found.values(), key=lambda sv: rank[sv[0]]):
        regex = _slot_regex(value, tokens)
        if regex is not None:
            chosen.setdefault(regex.pattern, (slot, regex))
    return sorted(chosen.values(), key=lambda sr: -len(sr[1].pattern))


_MARK = "\x00"


def slot_pattern(value: object, slots: list[tuple[str, re.Pattern[str]]]) -> str:
    """The cell with every slot value replaced by a marker (scrubbed later)."""
    text = str(value).strip()
    for slot, regex in slots:
        text = regex.sub(f"{_MARK}{slot}{_MARK}", text)
    return text


def _scrub_pattern(text: str, scr) -> str:
    parts = text.split(_MARK)
    return "".join(f"{{{p}}}" if i % 2 else scr.scrub(p) for i, p in enumerate(parts))


def _share(n: int, total: int) -> float | None:
    return round(100.0 * n / total, 1) if total else None


def _is_master(row: dict) -> str | None:
    text = " ".join(str(row.get(c) or "") for c in ("pipelinedescription", "pipelinename"))
    if _GRAND_MASTER.search(text):
        return "grand_master"
    if _MASTER.search(text):
        return "master"
    return None


def _groups_under(pid: str, seen: set[str], group_of: dict, children: dict) -> set[str]:
    """The GROUP_IDs of ``pid`` and every pipeline below it."""
    if pid in seen:
        return set()
    seen.add(pid)
    out = set(group_of.get(pid, set()))
    for child in children.get(pid, []):
        out |= _groups_under(child, seen, group_of, children)
    return out


def master_counts(books: list[Book]) -> dict:
    """Grand master / master schedule rows (PIPELINE_DESCRIPTION / NAME says so,
    or the row is another row's PARENT_PIPELINE_ID), and the GROUP_IDs under
    each: a child pipeline's GROUP_ID from the workbook's notebook / ingestion
    rows (PIPELINE_ID -> GROUP_ID); a grand master's through its masters."""
    per_book = []
    groups_per_master: Counter[int] = Counter()
    masters_per_group: Counter[int] = Counter()
    kinds: Counter[str] = Counter()
    for book in books:
        sched = book.sheets.get(SCHEDULE, ({}, []))[1]
        if not sched:
            continue
        group_of: dict[str, set[str]] = defaultdict(set)
        for _h, records in book.sheets.values():
            for row in records:
                pid, gid = str(row.get("pipelineid") or "").strip(), str(
                    row.get("groupid") or "").strip()
                if filled(pid) and filled(gid):
                    group_of[pid].add(gid)
        parents = {str(r.get("parentpipelineid") or "").strip() for r in sched}
        children: dict[str, list[str]] = defaultdict(list)
        for row in sched:
            parent = str(row.get("parentpipelineid") or "").strip()
            if filled(parent):
                children[parent].append(str(row.get("pipelineid") or "").strip())

        def groups_under(pid: str, seen: set[str], group_of=group_of,
                         children=children) -> set[str]:
            return _groups_under(pid, seen, group_of, children)

        masters = []
        for row in sched:
            pid = str(row.get("pipelineid") or "").strip()
            kind = _is_master(row) or ("parent" if pid and pid in parents else None)
            if kind is None:
                continue
            kinds[kind] += 1
            under = groups_under(pid, set())
            masters.append((kind, len(under)))
            groups_per_master[len(under)] += 1
        all_groups = {g for gs in group_of.values() for g in gs}
        per_group: Counter[str] = Counter()
        for row in sched:
            pid = str(row.get("pipelineid") or "").strip()
            if _is_master(row) or (pid and pid in parents):
                for gid in groups_under(pid, set()):
                    per_group[gid] += 1
        for gid in all_groups:
            masters_per_group[per_group.get(gid, 0)] += 1
        per_book.append({
            "schedule_rows": len(sched), "groups": len(all_groups),
            "grand_master": sum(1 for k, _ in masters if k == "grand_master"),
            "master": sum(1 for k, _ in masters if k == "master"),
            "parent_only": sum(1 for k, _ in masters if k == "parent")})
    return {"rows_by_kind": dict(kinds),
            "groups_under_a_master": {str(k): v for k, v in sorted(groups_per_master.items())},
            "master_rows_per_group": {str(k): v for k, v in sorted(masters_per_group.items())},
            "per_workbook": per_book}


def _pid(row: dict) -> str:
    return _diff_module()._id(row.get("pipelineid"))


def _ranked(counter: Counter, total: int, top: int | None, render) -> list:
    """[[rendered key, share %, rows], ...] — the most rows first, ``top`` kept
    (None: every one)."""
    return [[render(k), _share(n, total), n] for k, n in counter.most_common(top)]


# ------------------------------------------------- framework vocabulary + flows


@dataclass(frozen=True)
class Vocabulary:
    """tools/iig_vocabulary.yaml (D1 post44)."""

    columns: dict[str, frozenset[str]]                       # sheet key -> header ids
    listing: dict[str, list[str]] = field(default_factory=dict)   # as the file names them
    values: tuple[str, ...] = ()
    person_columns: frozenset[str] = frozenset()
    flow_rules: tuple[tuple[str, re.Pattern[str], str], ...] = ()
    breakdown: dict[str, tuple[str, ...]] = field(default_factory=dict)
    file_flows: frozenset[str] = frozenset()
    file_name: re.Pattern[str] | None = None


_VOCABULARY_KEYS = {"framework_columns", "framework_values", "person_columns", "flow_rules",
                    "flow_breakdown", "file_ingestion_flows", "file_ingestion_pipeline_name"}


def load_vocabulary(path: Path = VOCABULARY, keys=None) -> Vocabulary:
    """The vocabulary file; ``keys`` (iig_keys.yaml) adds its flow.rules after
    the file's own. A missing file = no allowlist, no rules of its own."""
    import yaml

    raw = (yaml.safe_load(path.read_text(encoding="utf-8")) if path.is_file() else None) or {}
    unknown = set(raw) - _VOCABULARY_KEYS
    if unknown:
        raise ValueError(f"{path}: unknown keys {sorted(unknown)}")
    rules = []
    for n, rule in enumerate(raw.get("flow_rules") or []):
        if not isinstance(rule, dict) or set(rule) != {"column", "match", "flow"}:
            raise ValueError(f"{path}: flow_rules[{n}] needs exactly column, match, flow")
        rules.append((header_id(rule["column"]), re.compile(str(rule["match"])),
                      str(rule["flow"])))
    if keys is not None:
        rules += [(header_id(c), rx, flow) for c, rx, flow in keys.flow_rules]
    name = raw.get("file_ingestion_pipeline_name")
    return Vocabulary(
        columns={sheet_key(s): frozenset(header_id(c) for c in cols or [])
                 for s, cols in (raw.get("framework_columns") or {}).items()},
        listing={str(s): [str(c) for c in cols or []]
                 for s, cols in (raw.get("framework_columns") or {}).items()},
        values=tuple(str(v) for v in raw.get("framework_values") or []),
        person_columns=frozenset(header_id(c) for c in raw.get("person_columns") or []),
        flow_rules=tuple(rules),
        breakdown={sheet_key(s): tuple(header_id(c) for c in cols or [])
                   for s, cols in (raw.get("flow_breakdown") or {}).items()},
        file_flows=frozenset(str(f) for f in raw.get("file_ingestion_flows") or []),
        file_name=re.compile(str(name)) if name else None)


def group_index_paths(corpus: list[Path]) -> list[Path]:
    """The group index inside a corpus FOLDER (01_pair_group_index.csv) and its
    round-2 addendum (``<folder>_round2/``), the files that exist — never
    beside a corpus file (a workbook in Downloads is not the corpus)."""
    found: list[Path] = []
    for folder in (p for p in corpus if p.is_dir()):
        for base in (folder, folder.parent / f"{folder.name}_round2"):
            for name in GROUP_INDEX_NAMES:
                if (base / name).is_file() and base / name not in found:
                    found.append(base / name)
    return found


def load_group_index(paths: list[Path]) -> dict[str, str]:
    """PIPELINE_ID -> flow_type (tools/iig_diff.py GroupIndex)."""
    index = _diff_module().GroupIndex(list(paths))
    return {pid: flow for pid, (flow, _table) in index.by_pipeline.items()}


def _rule_flow(row: dict, rules) -> str | None:
    for cid, regex, flow in rules:
        if regex.search(str(row.get(cid) or "").strip()):
            return flow
    return None


class Flows:
    """A row's flow: the group index by PIPELINE_ID; a grand master / master
    schedule row; the flow of its PIPELINE_ID's schedule row in its own
    workbook (else in the workbook that owns the id — an email CSV row); the
    rules; else unclassified. ``of`` returns (flow, basis)."""

    def __init__(self, books: list[Book], index: dict[str, str], rules) -> None:
        self.index, self.rules = index, rules
        self.pipes: dict[int, dict[str, str]] = {}
        self.owner: dict[str, Book] = {}
        for book in books:
            pipes: dict[str, str] = {}
            for row in book.sheets.get(SCHEDULE, ({}, []))[1]:
                pid = _pid(row)
                flow = _is_master(row) or _rule_flow(row, rules)
                if pid and flow:
                    pipes.setdefault(pid, flow)
            self.pipes[id(book)] = pipes
        # a feed workbook first: its slots are the row's
        for book in sorted(books, key=lambda b: not any(s in SLOT_SHEETS for s in b.sheets)):
            if set(book.sheets) <= {EMAIL}:
                continue
            for _h, records in book.sheets.values():
                for row in records:
                    pid = _pid(row)
                    if pid:
                        self.owner.setdefault(pid, book)

    def of(self, book: Book, sheet: str, row: dict) -> tuple[str, str]:
        pid = _pid(row)
        if pid and pid in self.index:
            return self.index[pid], "group index"
        if sheet == SCHEDULE and _is_master(row):
            return _is_master(row), "master row"
        for where in (book, self.owner.get(pid)):
            if where is not None and pid in self.pipes.get(id(where), {}):
                return self.pipes[id(where)][pid], "schedule row"
        flow = _rule_flow(row, self.rules)
        return (flow, "rule") if flow else (UNCLASSIFIED, "none")


@dataclass
class PatternContext:
    slots: dict[int, list]
    flows: Flows
    vocab: Vocabulary

    def slots_of(self, book: Book, row: dict) -> list:
        """The row's workbook's slots, else (an email CSV row) the slots of the
        workbook its PIPELINE_ID is in."""
        own = self.slots.get(id(book))
        if own:
            return own
        owner = self.flows.owner.get(_pid(row))
        return self.slots.get(id(owner), []) if owner is not None else []


def pattern_context(books: list[Book], tokens: tuple, vocab: Vocabulary,
                    index: dict[str, str]) -> PatternContext:
    return PatternContext({id(b): book_slots(b, tokens) for b in books},
                          Flows(books, index, vocab.flow_rules), vocab)


def vocabulary_tokens(books: list[Book], ctx: PatternContext) -> tuple[frozenset[str], dict]:
    """Caseless tokens the scrub keeps: the framework columns' cells (slot
    values taken out first) + framework_values, minus every word of a slot value
    or a person column anywhere in the corpus (framework_values excepted: they
    are listed by name)."""
    diff, vocab = _diff_module(), ctx.vocab
    slot_words: set[str] = set()
    person_words: set[str] = set()
    found: set[str] = set()
    for book in books:
        for sid, (_h, records) in book.sheets.items():
            columns = vocab.columns.get(sheet_key(sid), frozenset())
            for row in records:
                for cid, value in row.items():
                    if not filled(value):
                        continue
                    words = {t.casefold() for t in diff._tokens(str(value))}
                    if cid in vocab.person_columns:
                        person_words |= words
                    if sid in SLOT_SHEETS and any(cid in ids for ids in SLOT_COLUMNS.values()):
                        slot_words |= words
                    if cid in columns:
                        parts = slot_pattern(value, ctx.slots_of(book, row)).split(_MARK)
                        found.update(t.casefold() for p in parts[::2] for t in diff._tokens(p))
    listed = {t.casefold() for v in vocab.values for t in diff._tokens(v)}
    allowed = (found - slot_words - person_words) | listed
    return frozenset(allowed), {"from_columns": len(found), "listed": len(listed),
                                "refused_slot_value": len(found & slot_words),
                                "refused_person": len((found & person_words) - slot_words)}


def flow_breakdown(books: list[Book], sid: str, headers: dict, ctx: PatternContext, scr,
                   assigned, top: int | None) -> dict | None:
    """Per flow: rows, and each breakdown column's values / slot patterns."""
    wanted = ctx.vocab.breakdown.get(sheet_key(sid))
    if not wanted:
        return None
    by_flow: dict[str, list] = defaultdict(list)
    basis: Counter[str] = Counter()
    for book in books:
        for row in book.sheets.get(sid, ({}, []))[1]:
            flow, why = ctx.flows.of(book, sid, row)
            basis[why] += 1
            by_flow[flow].append((book, row))
    out: dict = {"rows_by_flow": {f: len(r) for f, r in sorted(by_flow.items())},
                 "basis": dict(basis.most_common()), "columns": {}}
    for cid in wanted:
        header = headers.get(cid, cid.upper())
        if assigned(header):
            continue
        per: dict = {}
        for flow, rows in sorted(by_flow.items()):
            values = [(b, r) for b, r in rows if filled(r.get(cid))]
            counts = Counter(str(r[cid]).strip() for _b, r in values)
            patterns = Counter(slot_pattern(r[cid], ctx.slots_of(b, r)) for b, r in values)
            per[flow] = {"rows": len(rows), "filled": len(values),
                         "values": _ranked(counts, len(values), top, scr.scrub),
                         "patterns": _ranked(patterns, len(values), top,
                                             lambda p: _scrub_pattern(p, scr))}
        out["columns"][scr.scrub(header)] = per
    return out


def email_table(books: list[Book], ctx: PatternContext, scr, top: int | None) -> dict | None:
    """EMAIL_TEMPLATE_CONFIG on FILE-INGESTION pipelines (the PIPELINE_ID's flow
    in file_ingestion_flows, else an unclassified row whose PIPELINE_NAME
    matches file_ingestion_pipeline_name): rows / pipelines per STATUS, the
    statuses per pipeline, and STATUS x SUBJECT pattern x BODY pattern."""
    rows = [(b, r) for b in books for r in b.sheets.get(EMAIL, ({}, []))[1]]
    if not rows:
        return None
    vocab = ctx.vocab
    by_flow: Counter[str] = Counter()
    basis: Counter[str] = Counter()
    status_all: Counter[str] = Counter()
    status_file: Counter[str] = Counter()
    pipelines: dict[str, set[str]] = defaultdict(set)
    statuses: dict[str, set[str]] = defaultdict(set)
    combos: Counter[tuple] = Counter()
    for book, row in rows:
        flow, why = ctx.flows.of(book, EMAIL, row)
        name = str(row.get("pipelinename") or "").strip()
        is_file = flow in vocab.file_flows
        if not is_file and flow == UNCLASSIFIED and vocab.file_name is not None \
                and vocab.file_name.search(name):
            is_file, why = True, "pipeline name"
        by_flow[flow] += 1
        basis[why] += 1
        status = str(row.get("status") or "").strip()
        status_all[status] += 1
        if not is_file:
            continue
        status_file[status] += 1
        pipeline = _pid(row) or name
        pipelines[status].add(pipeline)
        statuses[pipeline].add(status)
        slots = ctx.slots_of(book, row)
        combos[(status, slot_pattern(row.get("subject") or "", slots),
                slot_pattern(row.get("body") or "", slots))] += 1
    file_rows = sum(status_file.values())
    return {
        "rows": len(rows), "file_ingestion_rows": file_rows,
        "rows_by_flow": dict(by_flow.most_common()), "basis": dict(basis.most_common()),
        "status": [{"status": scr.scrub(s) or "<blank>", "rows": n,
                    "file_ingestion_rows": status_file.get(s, 0),
                    "file_ingestion_pipelines": len(pipelines.get(s, ()))}
                   for s, n in status_all.most_common()],
        "statuses_per_pipeline": {str(k): v for k, v in sorted(
            Counter(len(v) for v in statuses.values()).items())},
        "combinations": [{"status": scr.scrub(s) or "<blank>",
                          "subject": _scrub_pattern(su, scr), "body": _scrub_pattern(bo, scr),
                          "rows": n, "share": _share(n, file_rows)}
                         for (s, su, bo), n in combos.most_common(top)]}


def patterns_report(books: list[Book], sheets: list[str], assigned, withheld, scr,
                    tokens: tuple, top: int | None = None,
                    ctx: PatternContext | None = None) -> dict:
    if ctx is None:
        ctx = pattern_context(books, tokens, load_vocabulary(), {})
    report: dict = {"top": top, "sheets": {}, "flows": {}}
    for sid in sheets:
        rows = [(b, r) for b in books for r in b.sheets.get(sid, ({}, []))[1]]
        headers: dict[str, str] = {}
        for b in books:
            for k, v in b.sheets.get(sid, ({}, []))[0].items():
                headers.setdefault(k, v)
        columns: dict = {}
        for cid, header in headers.items():
            values = [(b, r) for b, r in rows if filled(r.get(cid))]
            col: dict = {"rows": len(rows), "filled": len(values),
                         "filled_pct": _share(len(values), len(rows))}
            if assigned(header) or withheld(header):
                col["values"] = "withheld (assigned / withheld column)"
                columns[scr.scrub(header)] = col
                continue
            counts = Counter(str(r[cid]).strip() for _b, r in values)
            patterns = Counter(slot_pattern(r[cid], ctx.slots_of(b, r)) for b, r in values)
            if counts:
                value, n = counts.most_common(1)[0]
                col.update(dominant_value=scr.scrub(value),
                           dominant_value_share=_share(n, len(values)),
                           distinct_values=len(counts),
                           top_values=_ranked(counts, len(values), top, scr.scrub))
            if patterns:
                first, n = patterns.most_common(1)[0]
                col.update(dominant_pattern=_scrub_pattern(first, scr),
                           dominant_pattern_share=_share(n, len(values)),
                           distinct_patterns=len(patterns),
                           per_feed=_MARK in first,
                           top_patterns=_ranked(patterns, len(values), top,
                                                lambda p: _scrub_pattern(p, scr)))
            columns[scr.scrub(header)] = col
        report["sheets"][sid] = {"rows": len(rows), "columns": columns}
        breakdown = flow_breakdown(books, sid, headers, ctx, scr, assigned, top)
        if breakdown is not None:
            report["flows"][sid] = breakdown
    report["email"] = email_table(books, ctx, scr, top)
    sched = [r for b in books for r in b.sheets.get(SCHEDULE, ({}, []))[1]]
    by_freq: dict[str, list] = defaultdict(list)
    for row in sched:
        by_freq[scr.scrub(str(row.get("pipelinefrequency") or "").strip()) or "<blank>"].append(
            row.get("dayofschedule"))
    report["day_of_schedule_by_frequency"] = {
        freq: {"rows": len(days), "filled": sum(1 for d in days if filled(d)),
               "filled_pct": _share(sum(1 for d in days if filled(d)), len(days)),
               "dominant": (lambda c: [scr.scrub(c[0][0]), _share(c[0][1], len(days))]
                            if c else None)(Counter(str(d).strip() for d in days
                                                    if filled(d)).most_common(1))}
        for freq, days in sorted(by_freq.items(), key=lambda kv: -len(kv[1]))}
    report["masters"] = master_counts(books)
    report["rule_classes"] = rule_classes(books, scr)
    return report


def rule_classes(books: list[Book], scr) -> list[dict]:
    """DATA_QUALITY_RULES rows per (RULE_TYPE, RULE_CLASS) over the corpus,
    and how many workbooks have each (scrubbed), the most rows first."""
    rows: Counter[tuple] = Counter()
    books_with: dict[tuple, set[int]] = defaultdict(set)
    for n, book in enumerate(books):
        for row in book.sheets.get(RULES, ({}, []))[1]:
            key = tuple(str(row.get(c) or "").strip() for c in RULE_COLUMNS)
            rows[key] += 1
            books_with[key].add(n)
    return [{"rule_type": scr.scrub(t), "rule_class": scr.scrub(c), "rows": count,
             "workbooks": len(books_with[(t, c)])}
            for (t, c), count in rows.most_common()]


def _entries(items: list, md, pct) -> list[str]:
    return [f"  - {pct(share)}% ({n}) `{md(text) or '<blank>'}`" for text, share, n in items]


def _vocabulary_md(vocab: dict | None, md) -> list[str]:
    if not vocab:
        return []
    return ["", "## Framework vocabulary (passed through unscrubbed)", "",
            f"- columns: {md(vocab['columns'])}",
            f"- tokens: {vocab['counts']['from_columns']} from the columns, "
            f"{vocab['counts']['listed']} listed (framework_values); refused: "
            f"{vocab['counts']['refused_slot_value']} slot-value word(s), "
            f"{vocab['counts']['refused_person']} person-column word(s)",
            f"- kept: {md(', '.join(vocab['tokens'])) or 'none'}"]


def _flows_md(flows: dict, md, pct) -> list[str]:
    lines: list[str] = []
    for sid, entry in flows.items():
        lines += ["", f"## {sid} per flow", "",
                  f"- rows by flow: {entry['rows_by_flow']}",
                  f"- flow from: {entry['basis']}"]
        for header, per in entry["columns"].items():
            for flow, d in per.items():
                lines += ["", f"### {md(header)} - {flow} ({d['filled']} of {d['rows']} "
                              "filled)", "", "- values:"]
                lines += _entries(d["values"], md, pct)
                lines.append("- patterns:")
                lines += _entries(d["patterns"], md, pct)
    return lines


def _email_md(email: dict | None, md, pct) -> list[str]:
    if not email:
        return []
    lines = ["", f"## {EMAIL} on file-ingestion pipelines", "",
             f"- rows: {email['rows']}; on file-ingestion pipelines: "
             f"{email['file_ingestion_rows']}",
             f"- rows by flow: {email['rows_by_flow']}", f"- flow from: {email['basis']}",
             f"- statuses per file-ingestion pipeline -> pipelines: "
             f"{email['statuses_per_pipeline']}", "",
             "| STATUS | rows | file-ingestion rows | file-ingestion pipelines |",
             "|---|---|---|---|"]
    lines += [f"| {md(s['status'])} | {s['rows']} | {s['file_ingestion_rows']} | "
              f"{s['file_ingestion_pipelines']} |" for s in email["status"]]
    lines += ["", "| STATUS | SUBJECT pattern | BODY pattern | rows | share % |",
              "|---|---|---|---|---|"]
    lines += [f"| {md(c['status'])} | {md(c['subject'])} | {md(c['body'])} | {c['rows']} | "
              f"{pct(c['share'])} |" for c in email["combinations"]]
    return lines


def _lists_md(report: dict, md, pct) -> list[str]:
    top = report.get("top")
    lines = ["", f"## Value and pattern lists ({'all' if top is None else f'top {top}'})"]
    for sid, sheet in report["sheets"].items():
        for header, col in sheet["columns"].items():
            if "top_values" not in col:
                continue
            lines += ["", f"### {sid}.{md(header)}", "", "- values:"]
            lines += _entries(col["top_values"], md, pct)
            lines.append("- patterns:")
            lines += _entries(col.get("top_patterns", []), md, pct)
    return lines


def patterns_markdown(report: dict, books: list[Book], sheets: list[str], scr) -> str:
    def pct(v):
        return "-" if v is None else f"{v:.1f}"

    def md(text: object) -> str:
        return str(text).replace("|", "\\|").replace("\n", " ")

    lines = ["# IIG corpus patterns - scrubbed (placeholders only)", "",
             "Slot placeholders: {feed} {table} {schema} {catalog} {domain} {subdomain} "
             "{file} {source} - the row's own workbook's values (ADLS / STGDELTA / FILE_ADLS"
             " / fixed-width sheets). G<n>: a token outside the fixture vocabulary.", "",
             f"- corpus workbooks: {len(books)}"]
    for sid in sheets:
        n = sum(1 for b in books if sid in b.sheets)
        lines.append(f"- {sid}: in {n} workbook(s), "
                     f"{report['sheets'].get(sid, {}).get('rows', 0)} row(s)")
    shapes = Counter(tuple(sorted(scr.scrub(s) for s in b.sheetnames)) for b in books)
    lines += ["", "## Sheet sets (workbooks with each set)", ""]
    lines += [f"- {c} x [{md(', '.join(s))}]" for s, c in shapes.most_common()]
    for sid, sheet in report["sheets"].items():
        lines += ["", f"## {sid} ({sheet['rows']} corpus rows)", "",
                  "| column | filled % | dominant value | share % | dominant pattern | share % "
                  "| patterns | per-feed |", "|---|---|---|---|---|---|---|---|"]
        for header, col in sheet["columns"].items():
            if "values" in col:
                lines.append(f"| {md(header)} | {pct(col['filled_pct'])} | {col['values']} "
                             "| | | | | |")
                continue
            lines.append(
                f"| {md(header)} | {pct(col['filled_pct'])} | "
                f"{md(col.get('dominant_value', ''))} | {pct(col.get('dominant_value_share'))} | "
                f"{md(col.get('dominant_pattern', ''))} | "
                f"{pct(col.get('dominant_pattern_share'))} | {col.get('distinct_patterns', 0)} | "
                f"{'yes' if col.get('per_feed') else ''} |")
    lines += ["", f"## DAY_OF_SCHEDULE by PIPELINE_FREQUENCY ({SCHEDULE})", "",
              "| frequency | rows | DAY filled | filled % | dominant DAY | share % |",
              "|---|---|---|---|---|---|"]
    for freq, d in report["day_of_schedule_by_frequency"].items():
        dom = d["dominant"] or ["", None]
        lines.append(f"| {md(freq)} | {d['rows']} | {d['filled']} | {pct(d['filled_pct'])} | "
                     f"{md(dom[0])} | {pct(dom[1])} |")
    lines += ["", f"## {RULES} rule classes (RULE_TYPE x RULE_CLASS)", "",
              f"- workbooks with the sheet: {sum(1 for b in books if RULES in b.sheets)} of "
              f"{len(books)}", "", "| RULE_TYPE | RULE_CLASS | rows | workbooks |",
              "|---|---|---|---|"]
    lines += [f"| {md(r['rule_type']) or '<blank>'} | {md(r['rule_class']) or '<blank>'} | "
              f"{r['rows']} | {r['workbooks']} |" for r in report["rule_classes"]]
    lines += _vocabulary_md(report.get("vocabulary"), md)
    lines += _flows_md(report.get("flows") or {}, md, pct)
    lines += _email_md(report.get("email"), md, pct)
    lines += _lists_md(report, md, pct)
    m = report["masters"]
    lines += ["", "## Grand master / master schedule rows", "",
              f"- rows by kind (description / name, else a PARENT_PIPELINE_ID): "
              f"{m['rows_by_kind'] or 'none'}",
              f"- GROUP_IDs under one master row -> master rows: {m['groups_under_a_master']}",
              f"- master rows over one GROUP_ID -> GROUP_IDs: {m['master_rows_per_group']}",
              "", "| workbook | schedule rows | GROUP_IDs | grand master | master | parent only |",
              "|---|---|---|---|---|---|"]
    lines += [f"| {n} | {w['schedule_rows']} | {w['groups']} | {w['grand_master']} | "
              f"{w['master']} | {w['parent_only']} |" for n, w in enumerate(m["per_workbook"], 1)]
    return "\n".join(lines) + "\n"


def write_patterns(out: Path, books: list[Book], sheets: list[str], keys, assigned,
                   root: Path = TOOLS.parent, *, top: int | None = None,
                   index: dict[str, str] | None = None,
                   vocabulary: Path = VOCABULARY) -> dict:
    """patterns.md + patterns.json under ``out``, scrubbed; re-scanned, and
    deleted on a scrub_check hit (the result's ``rescan`` names the classes).
    ``index``: PIPELINE_ID -> flow (load_group_index); ``top``: list cap."""
    diff = _diff_module()
    vocab = load_vocabulary(vocabulary, keys)
    ctx = pattern_context(books, _file_tokens(diff), vocab, index or {})
    allow, allow_counts = vocabulary_tokens(books, ctx)
    sc = diff.load_scrub_check(root)
    terms = diff.load_local_denylist(root / diff.LOCAL_DENYLIST_REL)
    harvested = sc.harvest_raw_denylist()
    universe, _source, _n = diff.fixture_universe(root)
    corpus_tokens: set[str] = set()
    for b in books:
        for _h, records in b.sheets.values():
            for row in records:
                for v in row.values():
                    corpus_tokens.update(t.casefold() for t in diff._tokens(str(v or "")))
    scr = diff.Scrubber(sc, universe, frozenset(corpus_tokens), terms, harvested, allow=allow)

    def withheld(header: str) -> bool:
        return header_id(header) in {header_id(c) for c in keys.withhold_values}

    report = patterns_report(books, sheets, assigned, withheld, scr, _file_tokens(diff),
                             top=top, ctx=ctx)
    report["vocabulary"] = {
        "columns": vocab.listing,
        "counts": allow_counts,
        # only the kept tokens the export shows (outside the fixture universe)
        "tokens": sorted(t for t in allow if t not in universe)}
    out.mkdir(parents=True, exist_ok=True)
    files = [out / "patterns.md", out / "patterns.json"]
    files[0].write_text(patterns_markdown(report, books, sheets, scr), encoding="utf-8")
    files[1].write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n",
                        encoding="utf-8")
    hits = diff._rescan(sc, files, harvested, scr.local)
    if hits:
        for f in files:
            f.unlink(missing_ok=True)
    report["rescan"] = dict(hits)
    report["placeholders"] = len(scr.maps["G"]) + len(scr.maps["A"])
    return report


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--corpus", nargs="+", type=Path, required=True)
    ap.add_argument("--agent", nargs="+", type=Path, required=True)
    ap.add_argument("--sheets", default=",".join(DEFAULT_SHEETS))
    ap.add_argument("--threshold", type=float, default=90.0)
    ap.add_argument("--keys", type=Path, default=TOOLS / "iig_keys.yaml")
    ap.add_argument("--json", type=Path, help="write the report here (values withheld "
                                              "unless --show-values)")
    ap.add_argument("--show-values", action="store_true",
                    help="print / write the dominant corpus value (D1 only: production text)")
    ap.add_argument("--patterns", type=Path, metavar="DIR",
                    help="write the scrubbed patterns.md / patterns.json here")
    ap.add_argument("--top", type=int, default=None, metavar="N",
                    help="--patterns: keep the N most common values / patterns per list "
                         "(default: all)")
    ap.add_argument("--group-index", nargs="*", type=Path, default=None, metavar="CSV",
                    help="--patterns: PIPELINE_ID -> flow_type index(es) (default: "
                         "01_pair_group_index*.csv beside a corpus folder and its _round2)")
    ap.add_argument("--vocabulary", type=Path, default=VOCABULARY,
                    help="--patterns: the framework-vocabulary allowlist + flow rules")
    args = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    keys = _diff_module().load_keys(args.keys)
    blank = always_blank()

    def assigned(header: str) -> bool:
        return keys.is_assignment(header) or header_id(header) in blank
    sheets = [s.strip().upper() for s in args.sheets.split(",") if s.strip()]
    corpus_books, agent_books = workbooks(args.corpus), workbooks(args.agent)
    if not corpus_books or not agent_books:
        print(f"FAIL no workbook: corpus {len(corpus_books)}, agent {len(agent_books)}")
        return 1
    print(f"corpus: {len(corpus_books)} workbook(s); agent: {len(agent_books)} workbook(s); "
          f"threshold {args.threshold:.0f}%")
    read = [*sheets, *(s for s in (*SLOT_SHEETS, RULES, SCHEDULE) if s not in sheets)]
    corpus = read_books(corpus_books, read if args.patterns else sheets)
    agent = read_books(agent_books, sheets)
    print_inventory(corpus, sheets)
    report = fill_rates(merge_books(corpus, set(sheets)), merge_books(agent, set(sheets)),
                        sheets, args.threshold, assigned)
    print_report(report, args.show_values)
    if args.json:
        if not args.show_values:
            for sheet in report.values():
                for col in sheet["columns"].values():
                    col["dominant_value"] = None
        args.json.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n",
                             encoding="utf-8")
        print(f"wrote {args.json}")
    if args.patterns:
        index_files = (args.group_index if args.group_index is not None
                       else group_index_paths(args.corpus))
        index = load_group_index(index_files)
        print(f"group index: {len(index)} pipeline id(s) from "
              f"{', '.join(str(p) for p in index_files) or 'no file'}")
        patterns = write_patterns(args.patterns, corpus, sheets, keys, assigned,
                                  top=args.top, index=index, vocabulary=args.vocabulary)
        if patterns["rescan"]:
            print(f"FAIL patterns not written: the re-scan found "
                  f"{', '.join(f'{k} {v}' for k, v in sorted(patterns['rescan'].items()))}")
            return 1
        print(f"wrote {args.patterns / 'patterns.md'} + patterns.json (scrubbed: "
              f"{patterns['placeholders']} placeholder(s); re-scan 0 hits)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
