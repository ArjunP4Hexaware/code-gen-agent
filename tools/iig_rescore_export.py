"""Re-score a scrubbed iig_diff export (post57) — at home, without the golden.

    python tools/iig_rescore_export.py --export <dir with mismatches.md, row_keys.md,
        sheets.md> --out <dir> [--baseline-tool <older iig_diff.py> --baseline-keys <yaml>]

``tools/iig_diff.py --export-scrubbed`` (D1) writes every mismatch / missing /
extra cell with its (scrubbed) golden and agent values, every row's key and
pairing, and the per-sheet counts. That is enough to rebuild a golden and an
agent workbook that the scorer reads the same way:

* columns — the agent's = the IIG template the export's sheets name (config);
  the golden's = those minus the HEADERS section's ``extra`` and plus its
  ``missing`` (a template-only line of a newer export counts as extra);
* rows — every golden / agent row of ``row_keys.md``, at its row number;
* cells — every listed cell gets its scrubbed value on its side; a cell the
  export does not list (exact / normalized / needs_assignment / convention /
  bsa-owned) gets the SAME placeholder on both sides, so it scores as equal.

The rebuilt pair is then scored by this tree's ``tools/iig_diff.py`` — and,
with ``--baseline-tool``, by an older one — and the per-sheet scores are
printed beside the export's own (``sheets.md``): the export's numbers
reproduce first, then the rule changes show. Values stay as scrubbed; nothing
here reads a client document.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

TOOLS = Path(__file__).resolve().parent
_SPLIT = re.compile(r"(?<!\\)\|")
_ROWS = re.compile(r"\(golden row (?P<g>[-\d]+), agent row (?P<a>[-\d]+)\)\s*$")
EQUAL = "="                    # a cell the export does not list: equal on both sides


@dataclass
class ExportSheet:
    name: str
    golden_rows: list[int] = field(default_factory=list)
    agent_rows: list[int] = field(default_factory=list)
    extra_columns: list[str] = field(default_factory=list)      # agent has, golden lacks
    missing_columns: list[str] = field(default_factory=list)    # golden has, agent lacks
    cells: list[dict] = field(default_factory=list)
    reported: dict = field(default_factory=dict)                 # the export's own row
    key_columns: list[str] = field(default_factory=list)         # identity, as paired
    keys: dict = field(default_factory=dict)                     # (side, row) -> key text


def _cells(line: str) -> list[str]:
    """One markdown table row -> its cells, ``\\|`` unescaped, the one space
    of padding on each side removed (a value's own spaces kept)."""
    parts = _SPLIT.split(line.rstrip("\n"))[1:-1]
    out = []
    for part in parts:
        part = part[1:] if part.startswith(" ") else part
        part = part[:-1] if part.endswith(" ") else part
        out.append(part.replace("\\|", "|").replace("<br>", "\n"))
    return out


def _table(lines: list[str], start: int) -> list[list[str]]:
    """The table rows after the header + separator lines at ``start``."""
    rows = []
    for line in lines[start + 2:]:
        if not line.startswith("|"):
            break
        rows.append(_cells(line))
    return rows


def read_export(folder: Path) -> dict[str, ExportSheet]:
    folder = Path(folder)
    sheets: dict[str, ExportSheet] = {}
    lines = (folder / "sheets.md").read_text(encoding="utf-8").splitlines()
    at = next(i for i, ln in enumerate(lines) if ln.startswith("| sheet | paired"))
    head = _cells(lines[at])
    for row in _table(lines, at):
        sheets[row[0]] = ExportSheet(row[0], reported=dict(zip(head, row, strict=False)))
    for line in lines:
        match = re.match(r"^\s+(\S+) vs golden: (.*)$", line)
        if not match or match.group(1) not in sheets:
            continue
        for part in match.group(2).split("; "):
            kind, _sep, names = part.partition(" ")
            if part.startswith("template-only (not scored) "):
                kind, names = "extra", part[len("template-only (not scored) "):]
            target = {"extra": sheets[match.group(1)].extra_columns,
                      "missing": sheets[match.group(1)].missing_columns}.get(kind)
            if target is not None:
                target += [n.strip() for n in names.split(",") if n.strip()]
    lines = (folder / "row_keys.md").read_text(encoding="utf-8").splitlines()
    current = None
    for i, line in enumerate(lines):
        if line.startswith("## "):
            current = sheets.get(line[3:].strip())
        elif current is not None and line.startswith("- key columns used for pairing:"):
            current.key_columns = [c.strip() for c in line.split(":", 1)[1].split(",")
                                   if c.strip()]
        elif current is not None and line.startswith(("| golden row |", "| agent row |")):
            side = "g" if line.startswith("| golden") else "a"
            rows = _table(lines, i)
            setattr(current, "golden_rows" if side == "g" else "agent_rows",
                    [int(r[0]) for r in rows])
            current.keys.update({(side, int(r[0])): r[1] for r in rows})
    lines = (folder / "mismatches.md").read_text(encoding="utf-8").splitlines()
    at = next(i for i, ln in enumerate(lines) if ln.startswith("| sheet | column |"))
    for sheet, column, status, golden, agent, key in _table(lines, at):
        rows = _ROWS.search(key)
        if rows is None or sheet not in sheets:
            raise ValueError(f"unreadable mismatches.md row: {sheet} {column} {key[-40:]}")
        g = None if rows.group("g") == "-" else int(rows.group("g"))
        a = None if rows.group("a") == "-" else int(rows.group("a"))
        sheets[sheet].cells.append({"column": column, "status": status, "golden": golden,
                                    "agent": agent, "golden_row": g, "agent_row": a})
    return sheets


def template_columns(sheet_names: list[str]) -> dict[str, list[str]]:
    """The agent's columns per sheet: the config template whose sheets these are."""
    diff = load_tool(TOOLS / "iig_diff.py", "iig_diff_current")
    name = diff.detect_template(sheet_names)
    return diff.template_headers().get(name or "", {})


def _value(text: str) -> str | None:
    return text if text != "" else None


def rebuild(sheets: dict[str, ExportSheet], out: Path) -> tuple[Path, Path]:
    """Write ``golden.xlsx`` and ``agent/<feed>_IIG.xlsx`` (+ a build stamp)."""
    out = Path(out)
    templates = template_columns(list(sheets))
    golden_book, agent_book = openpyxl.Workbook(), openpyxl.Workbook()
    for book in (golden_book, agent_book):
        book.remove(book.active)
    for name, sheet in sheets.items():
        agent_cols = list(templates.get(name) or [])
        if not agent_cols:
            raise ValueError(f"sheet {name}: not a sheet of the IIG template")
        golden_cols = [c for c in agent_cols if c not in sheet.extra_columns]
        golden_cols += [c for c in sheet.missing_columns if c not in golden_cols]
        listed: dict[tuple[str, int, str], str] = {}
        # a key column the export does not list: the row's key text (as paired)
        for (side, row), text in sheet.keys.items():
            parts = [p.strip() for p in text.split(" | ")]
            if len(parts) != len(sheet.key_columns):
                continue
            for column, part in zip(sheet.key_columns, parts, strict=True):
                if not column.startswith("@"):
                    listed[(side, row, column)] = "" if part == "<none>" else part
        for cell in sheet.cells:
            if cell["golden_row"] is not None:
                listed[("g", cell["golden_row"], cell["column"])] = cell["golden"]
            if cell["agent_row"] is not None:
                listed[("a", cell["agent_row"], cell["column"])] = cell["agent"]
        for side, book, cols, rows in (("g", golden_book, golden_cols, sheet.golden_rows),
                                       ("a", agent_book, agent_cols, sheet.agent_rows)):
            ws = book.create_sheet(name)
            ws.append(cols)
            for row in rows:
                if row != ws.max_row + 1:
                    raise ValueError(f"sheet {name}: {side} rows are not contiguous at {row}")
                ws.append([_value(listed.get((side, row, c), EQUAL)) for c in cols])
    golden = out / "golden.xlsx"
    agent_dir = out / "agent"
    agent_dir.mkdir(parents=True, exist_ok=True)
    agent = agent_dir / "rebuilt_IIG.xlsx"
    golden_book.save(golden)
    agent_book.save(agent)
    (agent_dir / "CODEGEN_BUILD.json").write_text(json.dumps({
        "line": "rebuilt from a scrubbed iig_diff export (tools/iig_rescore_export.py)",
        "src_tree": "export"}, indent=2) + "\n", encoding="utf-8")
    return golden, agent


def load_tool(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)       # type: ignore[union-attr]
    return module


def score(tool_path: Path, keys: Path, golden: Path, agent: Path, out: Path,
          name: str) -> dict:
    tool = load_tool(tool_path, name)
    return tool.run_pair(golden, [agent], out, keys, group_index=[])


def _row(payload: dict, sheet: str) -> dict:
    s = payload["sheets"].get(sheet) or {}
    c = s.get("counts") or {}
    pairing = s.get("pairing") or {}
    return {"score": s.get("note") or s.get("score"), "paired": sum(pairing.values()),
            "order": pairing.get("order", 0), "exact+normalized":
            c.get("exact", 0) + c.get("normalized", 0), "mismatch": c.get("mismatch", 0),
            "missing": c.get("missing-in-agent", 0), "extra": c.get("extra-in-agent", 0),
            "template-only": c.get("template-only", 0)}


def report(sheets: dict[str, ExportSheet], runs: dict[str, dict]) -> list[str]:
    """One table: the export's score per sheet, then each run's."""
    head = ["sheet", "export score"]
    for label in runs:
        head += [f"{label} score", f"{label} paired (order)",
                 f"{label} exact+norm / mismatch / missing / extra / template-only"]
    lines = ["| " + " | ".join(head) + " |", "|" + " --- |" * len(head)]
    for name, sheet in sheets.items():
        cells = [name, sheet.reported.get("score", "")]
        for payload in runs.values():
            r = _row(payload, name)
            cells += [str(r["score"]), f"{r['paired']} ({r['order']})",
                      f"{r['exact+normalized']} / {r['mismatch']} / {r['missing']} / "
                      f"{r['extra']} / {r['template-only']}"]
        lines.append("| " + " | ".join(cells) + " |")
    total = ["**TOTAL**", ""]
    for payload in runs.values():
        t = payload["totals"]
        total += [f"**{t['score']}**", str(sum(t["pairing"].values())),
                  f"{t['exact'] + t['normalized']} / {t['mismatch']} / "
                  f"{t['missing-in-agent']} / {t['extra-in-agent']} / "
                  f"{t.get('template-only', 0)}"]
    lines.append("| " + " | ".join(total) + " |")
    return lines


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--export", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--keys", type=Path, default=TOOLS / "iig_keys.yaml")
    parser.add_argument("--baseline-tool", type=Path, default=None,
                        help="an older tools/iig_diff.py to score the rebuilt pair with too")
    parser.add_argument("--baseline-keys", type=Path, default=None,
                        help="the iig_keys.yaml that older tool reads")
    args = parser.parse_args(argv)
    sheets = read_export(args.export)
    golden, agent = rebuild(sheets, args.out / "rebuilt")
    runs: dict[str, dict] = {}
    if args.baseline_tool is not None:
        runs["before"] = score(args.baseline_tool, args.baseline_keys or args.keys, golden,
                               agent, args.out / "before", "iig_diff_baseline")
    runs["now"] = score(TOOLS / "iig_diff.py", args.keys, golden, agent, args.out / "now",
                        "iig_diff_now")
    for line in report(sheets, runs):
        print(line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
